"""
Service d'export de fichiers pour le chatbot
Génère des rapports PDF et Excel
"""
import os
import tempfile
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional, BinaryIO
from io import BytesIO

# PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# Excel
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

logger = logging.getLogger(__name__)

# Graphiques
try:
    import matplotlib
    matplotlib.use('Agg')  # Backend non-interactif
    import matplotlib.pyplot as plt
except Exception:
    matplotlib = None
    plt = None
    logger.warning("Matplotlib non disponible. Export de graphiques desactive.")

from django.conf import settings
from django.http import HttpResponse



class FileExporter:
    """
    Service pour exporter des données en PDF et Excel
    """
    
    def __init__(self):
        self.temp_dir = tempfile.gettempdir()
        self.styles = getSampleStyleSheet()
        
        # Style personnalisé pour le titre
        self.title_style = ParagraphStyle(
            'CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            textColor=colors.darkblue,
            alignment=1  # Centré
        )
        
        # Style pour les sous-titres
        self.subtitle_style = ParagraphStyle(
            'CustomSubtitle',
            parent=self.styles['Heading2'],
            fontSize=14,
            spaceAfter=20,
            textColor=colors.darkgreen
        )
    
    def create_pdf_report(self, data: Dict[str, Any], filename: str = None) -> str:
        """
        Crée un rapport PDF basé sur les données
        """
        if filename is None:
            filename = f"rapport_hydria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        filepath = os.path.join(self.temp_dir, filename)
        
        try:
            doc = SimpleDocTemplate(filepath, pagesize=A4)
            story = []
            
            # Titre du rapport
            title = data.get('title', 'Rapport Hydria')
            story.append(Paragraph(title, self.title_style))
            story.append(Spacer(1, 12))
            
            # Date de génération
            date_str = datetime.now().strftime('%d/%m/%Y à %H:%M')
            story.append(Paragraph(f"Généré le {date_str}", self.styles['Normal']))
            story.append(Spacer(1, 20))
            
            # Résumé exécutif
            if 'summary' in data:
                story.append(Paragraph("Résumé exécutif", self.subtitle_style))
                story.append(Paragraph(data['summary'], self.styles['Normal']))
                story.append(Spacer(1, 20))
            
            # Métriques principales
            if 'metrics' in data:
                story.append(Paragraph("Métriques principales", self.subtitle_style))
                metrics_table = self._create_metrics_table(data['metrics'])
                story.append(metrics_table)
                story.append(Spacer(1, 20))
            
            # Données tabulaires
            if 'table_data' in data:
                story.append(Paragraph("Données détaillées", self.subtitle_style))
                data_table = self._create_data_table(data['table_data'])
                story.append(data_table)
                story.append(Spacer(1, 20))
            
            # Graphiques
            if 'charts' in data:
                story.append(Paragraph("Graphiques", self.subtitle_style))
                for chart_data in data['charts']:
                    chart_image = self._create_chart_image(chart_data)
                    if chart_image:
                        story.append(chart_image)
                        story.append(Spacer(1, 12))
            
            # Conclusions
            if 'conclusions' in data:
                story.append(Paragraph("Conclusions", self.subtitle_style))
                story.append(Paragraph(data['conclusions'], self.styles['Normal']))
            
            # Footer
            story.append(Spacer(1, 30))
            footer_text = "Rapport généré par Hydria Assistant IA"
            story.append(Paragraph(footer_text, self.styles['Italic']))
            
            doc.build(story)
            
            logger.info(f"Rapport PDF créé: {filepath}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erreur lors de la création du PDF: {e}")
            raise
    
    def create_excel_report(self, data: Dict[str, Any], filename: str = None) -> str:
        """
        Crée un rapport Excel avec plusieurs feuilles
        """
        if filename is None:
            filename = f"rapport_hydria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        filepath = os.path.join(self.temp_dir, filename)
        
        try:
            wb = Workbook()
            
            # Supprimer la feuille par défaut
            wb.remove(wb.active)
            
            # Feuille de résumé
            summary_ws = wb.create_sheet("Résumé")
            self._create_summary_sheet(summary_ws, data)
            
            # Feuille de données
            if 'table_data' in data:
                data_ws = wb.create_sheet("Données")
                self._create_data_sheet(data_ws, data['table_data'])
            
            # Feuille de graphiques
            if 'charts' in data:
                charts_ws = wb.create_sheet("Graphiques")
                self._create_charts_sheet(charts_ws, data['charts'])
            
            # Feuille de métriques
            if 'metrics' in data:
                metrics_ws = wb.create_sheet("Métriques")
                self._create_metrics_sheet(metrics_ws, data['metrics'])
            
            wb.save(filepath)
            
            logger.info(f"Rapport Excel créé: {filepath}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erreur lors de la création de l'Excel: {e}")
            raise
    
    def _create_metrics_table(self, metrics: List[Dict[str, Any]]) -> Table:
        """
        Crée un tableau de métriques pour PDF
        """
        data = [['Métrique', 'Valeur', 'Description']]
        
        for metric in metrics:
            data.append([
                metric.get('label', ''),
                str(metric.get('value', '')),
                metric.get('description', '')
            ])
        
        table = Table(data, colWidths=[2*inch, 1*inch, 3*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ALIGN', (1, 1), (1, -1), 'RIGHT'),  # Aligner les valeurs à droite
        ]))
        
        return table
    
    def _create_data_table(self, table_data: Dict[str, Any]) -> Table:
        """
        Crée un tableau de données pour PDF
        """
        columns = table_data.get('columns', [])
        rows = table_data.get('data', [])
        
        # Limiter le nombre de lignes pour le PDF
        max_rows = 50
        if len(rows) > max_rows:
            rows = rows[:max_rows]
        
        data = [columns]  # En-têtes
        
        for row in rows:
            if isinstance(row, dict):
                data.append([str(row.get(col, '')) for col in columns])
            elif isinstance(row, list):
                data.append([str(cell) for cell in row])
        
        # Calculer les largeurs des colonnes
        num_cols = len(columns)
        col_width = 6.5 * inch / num_cols if num_cols > 0 else 1 * inch
        col_widths = [col_width] * num_cols
        
        table = Table(data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkgreen),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        
        return table
    
    def _create_chart_image(self, chart_data: Dict[str, Any]) -> Optional[Image]:
        """
        Crée une image de graphique pour insertion dans le PDF
        """
        try:
            if plt is None:
                logger.warning("Matplotlib non disponible. Graphique ignore.")
                return None

            fig, ax = plt.subplots(figsize=(8, 6))
            
            chart_type = chart_data.get('type', 'bar')
            data = chart_data.get('data', [])
            title = chart_data.get('title', 'Graphique')
            
            if chart_type == 'bar' and data:
                x_key = chart_data.get('x_key', 'x')
                y_key = chart_data.get('y_key', 'y')
                
                x_values = [item.get(x_key, '') for item in data]
                y_values = [float(item.get(y_key, 0)) for item in data]
                
                ax.bar(x_values, y_values, color='#3B82F6')
                ax.set_xlabel(x_key.replace('_', ' ').title())
                ax.set_ylabel(y_key.replace('_', ' ').title())
                
            elif chart_type == 'line' and data:
                x_key = chart_data.get('x_key', 'x')
                y_key = chart_data.get('y_key', 'y')
                
                x_values = [item.get(x_key, '') for item in data]
                y_values = [float(item.get(y_key, 0)) for item in data]
                
                ax.plot(x_values, y_values, marker='o', color='#10B981')
                ax.set_xlabel(x_key.replace('_', ' ').title())
                ax.set_ylabel(y_key.replace('_', ' ').title())
                
            elif chart_type == 'pie' and data:
                labels = [item.get('name', '') for item in data]
                values = [float(item.get('value', 0)) for item in data]
                
                ax.pie(values, labels=labels, autopct='%1.1f%%', startangle=90)
                ax.axis('equal')
            
            ax.set_title(title, fontsize=14, fontweight='bold')
            plt.tight_layout()
            
            # Sauvegarder en mémoire
            img_buffer = BytesIO()
            plt.savefig(img_buffer, format='png', dpi=150, bbox_inches='tight')
            img_buffer.seek(0)
            plt.close()
            
            # Créer l'objet Image pour ReportLab
            return Image(img_buffer, width=6*inch, height=4*inch)
            
        except Exception as e:
            logger.warning(f"Erreur lors de la création du graphique: {e}")
            return None
    
    def _create_summary_sheet(self, ws, data: Dict[str, Any]):
        """
        Crée la feuille de résumé Excel
        """
        # Titre
        ws['A1'] = data.get('title', 'Rapport Hydria')
        ws['A1'].font = Font(size=18, bold=True, color='1F4E8C')
        ws.merge_cells('A1:D1')
        
        # Date
        ws['A3'] = f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
        ws['A3'].font = Font(italic=True)
        
        # Résumé
        if 'summary' in data:
            ws['A5'] = "Résumé exécutif"
            ws['A5'].font = Font(size=14, bold=True)
            ws['A6'] = data['summary']
            ws['A6'].alignment = Alignment(wrap_text=True)
        
        # Auto-ajuster les largeurs de colonnes
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_data_sheet(self, ws, table_data: Dict[str, Any]):
        """
        Crée la feuille de données Excel
        """
        columns = table_data.get('columns', [])
        rows = table_data.get('data', [])
        
        # En-têtes
        for col_idx, column in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=column)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='1F4E8C', end_color='1F4E8C', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
        
        # Données
        for row_idx, row in enumerate(rows, 2):
            if isinstance(row, dict):
                for col_idx, column in enumerate(columns, 1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(column, ''))
            elif isinstance(row, list):
                for col_idx, value in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Formatage
        for row in ws.iter_rows(min_row=2, max_row=len(rows)+1):
            for cell in row:
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
    
    def _create_charts_sheet(self, ws, charts_data: List[Dict[str, Any]]):
        """
        Crée la feuille de graphiques Excel
        """
        ws['A1'] = "Graphiques générés"
        ws['A1'].font = Font(size=16, bold=True)
        
        # Pour l'instant, juste lister les graphiques
        # L'implémentation complète nécessiterait d'intégrer les données
        for idx, chart_data in enumerate(charts_data, 2):
            ws[f'A{idx}'] = f"Graphique: {chart_data.get('title', f'Graphique {idx-1}')}"
    
    def _create_metrics_sheet(self, ws, metrics: List[Dict[str, Any]]):
        """
        Crée la feuille de métriques Excel
        """
        # En-têtes
        ws['A1'] = "Métrique"
        ws['B1'] = "Valeur"
        ws['C1'] = "Description"
        
        for cell in [ws['A1'], ws['B1'], ws['C1']]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='1F4E8C', end_color='1F4E8C', fill_type='solid')
        
        # Données
        for idx, metric in enumerate(metrics, 2):
            ws[f'A{idx}'] = metric.get('label', '')
            ws[f'B{idx}'] = metric.get('value', '')
            ws[f'C{idx}'] = metric.get('description', '')
    
    def create_csv_export(self, data: List[Dict[str, Any]], filename: str = None) -> str:
        """
        Crée un export CSV simple
        """
        import csv
        
        if filename is None:
            filename = f"export_hydria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        
        filepath = os.path.join(self.temp_dir, filename)
        
        if not data:
            return filepath
        
        try:
            with open(filepath, 'w', newline='', encoding='utf-8') as csvfile:
                fieldnames = data[0].keys()
                writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
                
                writer.writeheader()
                for row in data:
                    writer.writerow(row)
            
            logger.info(f"Export CSV créé: {filepath}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erreur lors de la création du CSV: {e}")
            raise
    
    def get_file_response(self, filepath: str, content_type: str) -> HttpResponse:
        """
        Retourne une réponse HTTP pour télécharger le fichier
        """
        try:
            with open(filepath, 'rb') as f:
                response = HttpResponse(f.read(), content_type=content_type)
                filename = os.path.basename(filepath)
                response['Content-Disposition'] = f'attachment; filename="{filename}"'
                return response
        except Exception as e:
            logger.error(f"Erreur lors de la lecture du fichier: {e}")
            raise
