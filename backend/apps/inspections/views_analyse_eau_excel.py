"""
Import/Export Excel pour les analyses d'eau.
"""

from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from datetime import datetime, time
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
import pytz

from .models_analyse_eau import AnalyseEau
from apps.infrastructure.models import EquipementConstitutif
from apps.billing.models import PointEau


@require_http_methods(["GET"])
def export_analyses_eau_excel(request):
    """Exporte les analyses d'eau en fichier Excel"""
    try:
        # Récupérer les filtres
        annee = request.GET.get('annee')
        commune_id = request.GET.get('commune_id')
        resultat = request.GET.get('resultat')

        # Récupérer les analyses
        analyses = AnalyseEau.objects.all().select_related(
            'equipement',
            'point_eau',
            'point_eau__parcelle__village__fokontany__commune'
        ).order_by('date_prelevement')

        if annee:
            analyses = analyses.filter(date_prelevement__year=int(annee))

        if resultat == 'positif':
            from django.db.models import Q
            analyses = analyses.filter(
                Q(tnc=True) | Q(nombre_coliformes__gt=0)
            )
        elif resultat == 'zero':
            analyses = analyses.filter(tnc=False, nombre_coliformes=0)

        # Créer le workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Analyses Eau"

        # Styles
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # En-têtes
        headers = [
            "Type de prélèvement",
            "Équipement ou Point d'eau",
            "Description lieu",
            "Commune",
            "Village",
            "Date et heure prélèvement",
            "Date et heure analyse",
            "Date et heure résultat",
            "Volume analysé (mL)",
            "Nombre de coliformes",
            "TNC",
            "Nombre pour 100mL",
            "pH",
            "Observation"
        ]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Données
        for row_idx, analyse in enumerate(analyses, start=2):
            # Filtrer par commune si spécifié
            if commune_id and commune_id != 'all':
                from apps.geography.models import Commune
                try:
                    commune_obj = Commune.objects.get(id=commune_id)
                    if analyse.commune_nom != commune_obj.nom_commune:
                        continue
                except Commune.DoesNotExist:
                    continue

            ws.cell(row=row_idx, column=1, value=analyse.get_type_prelevement_display()).border = thin_border
            ws.cell(row=row_idx, column=2, value=analyse.lieu_affichage).border = thin_border
            ws.cell(row=row_idx, column=3, value=analyse.description_lieu).border = thin_border
            ws.cell(row=row_idx, column=4, value=analyse.commune_nom).border = thin_border
            ws.cell(row=row_idx, column=5, value=analyse.village_nom).border = thin_border
            ws.cell(row=row_idx, column=6, value=analyse.date_prelevement.strftime('%d/%m/%Y %H:%M')).border = thin_border
            ws.cell(row=row_idx, column=7, value=analyse.date_analyse.strftime('%d/%m/%Y %H:%M')).border = thin_border
            ws.cell(row=row_idx, column=8, value=analyse.date_resultat.strftime('%d/%m/%Y %H:%M')).border = thin_border
            ws.cell(row=row_idx, column=9, value=analyse.volume_analyse).border = thin_border
            ws.cell(row=row_idx, column=10, value=analyse.nombre_coliformes if not analyse.tnc else '').border = thin_border
            ws.cell(row=row_idx, column=11, value='Oui' if analyse.tnc else 'Non').border = thin_border

            # Nombre pour 100mL
            if analyse.tnc:
                ws.cell(row=row_idx, column=12, value='TNC').border = thin_border
            elif analyse.nombre_pour_100ml is not None:
                ws.cell(row=row_idx, column=12, value=analyse.nombre_pour_100ml).border = thin_border
            else:
                ws.cell(row=row_idx, column=12, value='N/A').border = thin_border

            ws.cell(row=row_idx, column=13, value=float(analyse.ph) if analyse.ph else '').border = thin_border
            ws.cell(row=row_idx, column=14, value=analyse.observation).border = thin_border

        # Ajuster largeurs
        column_widths = {
            1: 20, 2: 35, 3: 40, 4: 20, 5: 20,
            6: 18, 7: 18, 8: 18, 9: 15, 10: 18,
            11: 8, 12: 18, 13: 8, 14: 40
        }
        for col_idx, width in column_widths.items():
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        # Figer première ligne
        ws.freeze_panes = "A2"

        # Sauvegarder
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # Créer la réponse
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"Analyses_Eau_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@require_http_methods(["GET"])
def template_import_analyses_eau(request):
    """Génère un template Excel pour l'import"""
    try:
        # Créer workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Analyses Eau"

        # Styles
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        example_fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # En-têtes
        headers = [
            "Type prélèvement*",
            "ID Équipement",
            "ID Point d'eau",
            "Description lieu",
            "Date prélèvement*",
            "Heure prélèvement*",
            "Date analyse*",
            "Heure analyse*",
            "Date résultat*",
            "Heure résultat*",
            "Volume (mL)*",
            "Nb coliformes",
            "TNC",
            "pH",
            "Observation"
        ]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Exemples
        exemples = [
            ["equipement", "uuid-equipement", "", "", "15/12/2025", "10:30", "15/12/2025", "14:00",
             "16/12/2025", "09:00", "100", "5", "Non", "7.2", "Eau claire"],
            ["point_eau", "", "uuid-point-eau", "", "16/12/2025", "11:00", "16/12/2025", "15:30",
             "17/12/2025", "10:00", "50", "", "Oui", "6.8", "Trop nombreux"],
            ["point_eau_traditionnel", "", "", "Puits village Ambanja", "17/12/2025", "09:00",
             "17/12/2025", "13:00", "18/12/2025", "08:30", "100", "0", "Non", "7.5", "Conforme"],
            ["bouteille", "", "", "Eau Vive 1.5L", "18/12/2025", "08:00", "18/12/2025", "12:00",
             "19/12/2025", "09:00", "100", "0", "Non", "7.0", ""],
        ]

        for row_idx, exemple in enumerate(exemples, start=2):
            for col_idx, valeur in enumerate(exemple, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=valeur)
                cell.fill = example_fill
                cell.border = thin_border

        # Instructions
        ws_instructions = wb.create_sheet(title="Instructions")
        instructions_text = [
            ("INSTRUCTIONS D'IMPORT", True),
            ("", False),
            ("Colonnes obligatoires (*) :", True),
            ("- Type prélèvement : equipement, point_eau, point_eau_traditionnel, bouteille, autre", False),
            ("- Dates : Format DD/MM/YYYY (ex: 15/12/2025)", False),
            ("- Heures : Format HH:MM (ex: 10:30)", False),
            ("- Volume : 10, 20, 25, 50, 60 ou 100 mL", False),
            ("", False),
            ("Colonnes optionnelles :", True),
            ("- ID Équipement : UUID de l'équipement (si type = equipement)", False),
            ("- ID Point d'eau : UUID du point d'eau (si type = point_eau)", False),
            ("- Description lieu : Texte libre (pour traditionnels, bouteilles, autres)", False),
            ("- Nb coliformes : Nombre entier (laisser vide si TNC)", False),
            ("- TNC : Oui ou Non (Trop Nombreux pour être Compté)", False),
            ("- pH : Valeur décimale entre 0 et 14", False),
            ("- Observation : Texte libre", False),
            ("", False),
            ("Notes importantes :", True),
            ("- Si TNC = Oui, ne pas renseigner Nb coliformes", False),
            ("- Les lignes d'exemple (grisées) doivent être supprimées avant l'import", False),
            ("- Commencer à saisir à partir de la ligne 2", False),
        ]

        for row_idx, (text, bold) in enumerate(instructions_text, start=1):
            cell = ws_instructions.cell(row=row_idx, column=1, value=text)
            if bold:
                cell.font = Font(bold=True, size=12)

        ws_instructions.column_dimensions['A'].width = 80

        # Largeurs colonnes feuille principale
        column_widths = {
            1: 25, 2: 30, 3: 30, 4: 40, 5: 15,
            6: 15, 7: 15, 8: 15, 9: 15, 10: 15,
            11: 12, 12: 12, 13: 8, 14: 8, 15: 40
        }
        for col_idx, width in column_widths.items():
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        # Sauvegarder
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="Template_Import_Analyses_Eau.xlsx"'

        return response

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@csrf_exempt
@require_http_methods(["POST"])
def import_analyses_eau_excel(request):
    """Importe les analyses d'eau depuis un fichier Excel"""
    try:
        if 'file' not in request.FILES:
            return JsonResponse({
                'success': False,
                'error': 'Aucun fichier fourni'
            }, status=400)

        file = request.FILES['file']
        wb = openpyxl.load_workbook(file)
        ws = wb.active

        nb_imported = 0
        nb_errors = 0
        errors = []

        tz_madagascar = pytz.timezone('Indian/Antananarivo')

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                # Ignorer lignes vides
                if not any(row):
                    continue

                type_prelevement = row[0]
                equipement_id = row[1]
                point_eau_id = row[2]
                description_lieu = row[3] or ''
                date_prelevement_str = row[4]
                heure_prelevement_str = row[5]
                date_analyse_str = row[6]
                heure_analyse_str = row[7]
                date_resultat_str = row[8]
                heure_resultat_str = row[9]
                volume_analyse = row[10]
                nombre_coliformes = row[11]
                tnc_str = row[12]
                ph = row[13]
                observation = row[14] or ''

                # Validation champs obligatoires
                if not all([type_prelevement, date_prelevement_str, heure_prelevement_str,
                           date_analyse_str, heure_analyse_str, date_resultat_str,
                           heure_resultat_str, volume_analyse]):
                    errors.append(f"Ligne {row_idx}: Champs obligatoires manquants")
                    nb_errors += 1
                    continue

                # Parser dates - gérer différents formats
                try:
                    def parse_datetime_flexible(date_val, heure_val):
                        """Parse date et heure avec support de multiples formats"""
                        # DEBUG: Log des valeurs reçues
                        print(f"[DEBUG] parse_datetime_flexible - date_val: {date_val} (type: {type(date_val)}), heure_val: {heure_val} (type: {type(heure_val)})")

                        # Si date_val est déjà un objet datetime d'Excel
                        if isinstance(date_val, datetime):
                            # Excel stocke parfois les dates en datetime
                            dt_naive = date_val
                            # Si heure fournie, remplacer l'heure
                            if heure_val:
                                # Gérer différents types d'heure
                                if isinstance(heure_val, time):
                                    # Si heure_val est un objet time Python (openpyxl peut retourner ça)
                                    dt_naive = dt_naive.replace(
                                        hour=heure_val.hour,
                                        minute=heure_val.minute,
                                        second=heure_val.second
                                    )
                                    print(f"[DEBUG] Heure parsée depuis time object: {heure_val.hour:02d}:{heure_val.minute:02d}")
                                elif isinstance(heure_val, datetime):
                                    # Si heure_val est un datetime (Excel stocke parfois l'heure comme ça)
                                    dt_naive = dt_naive.replace(
                                        hour=heure_val.hour,
                                        minute=heure_val.minute,
                                        second=heure_val.second
                                    )
                                    print(f"[DEBUG] Heure parsée depuis datetime: {heure_val.hour:02d}:{heure_val.minute:02d}")
                                elif isinstance(heure_val, str):
                                    heure_parts = str(heure_val).strip().split(':')
                                    if len(heure_parts) >= 2:
                                        dt_naive = dt_naive.replace(
                                            hour=int(heure_parts[0]),
                                            minute=int(heure_parts[1])
                                        )
                                        print(f"[DEBUG] Heure parsée depuis string: {heure_parts[0]}:{heure_parts[1]}")
                                # Si c'est un nombre (fraction de jour Excel)
                                elif isinstance(heure_val, (int, float)):
                                    # Excel stocke les heures comme fraction de jour (ex: 0.5 = 12h00)
                                    total_seconds = int(heure_val * 24 * 3600)
                                    hours = total_seconds // 3600
                                    minutes = (total_seconds % 3600) // 60
                                    dt_naive = dt_naive.replace(hour=hours, minute=minutes)
                                    print(f"[DEBUG] Heure parsée depuis float: {hours:02d}:{minutes:02d}")
                            return dt_naive

                        # Si date_val est une chaîne
                        date_str = str(date_val).strip()
                        heure_str = str(heure_val).strip() if heure_val else '00:00'

                        # Nettoyer les timestamps dupliqués (ex: "2012-10-03 00:00:00 17:10:00")
                        if date_str.count(':') > 1:
                            # Séparer par espace et prendre les 2 premiers éléments (date + time)
                            parts = date_str.split()
                            if len(parts) >= 2:
                                date_str = parts[0]  # Date
                                heure_str = parts[1].split('.')[0]  # Heure (sans microsecondes)

                        # Essayer format DD/MM/YYYY
                        try:
                            return datetime.strptime(f"{date_str} {heure_str}", '%d/%m/%Y %H:%M')
                        except ValueError:
                            pass

                        # Essayer format YYYY-MM-DD
                        try:
                            return datetime.strptime(f"{date_str} {heure_str}", '%Y-%m-%d %H:%M')
                        except ValueError:
                            pass

                        # Essayer format YYYY-MM-DD HH:MM:SS
                        try:
                            return datetime.strptime(f"{date_str} {heure_str}", '%Y-%m-%d %H:%M:%S')
                        except ValueError:
                            pass

                        raise ValueError(f"Format non reconnu: '{date_str}' '{heure_str}'")

                    date_prelevement = parse_datetime_flexible(date_prelevement_str, heure_prelevement_str)
                    date_analyse = parse_datetime_flexible(date_analyse_str, heure_analyse_str)
                    date_resultat = parse_datetime_flexible(date_resultat_str, heure_resultat_str)

                    # Localiser en timezone Madagascar
                    date_prelevement = tz_madagascar.localize(date_prelevement)
                    date_analyse = tz_madagascar.localize(date_analyse)
                    date_resultat = tz_madagascar.localize(date_resultat)
                except ValueError as e:
                    errors.append(f"Ligne {row_idx}: Format de date/heure invalide - {str(e)}")
                    nb_errors += 1
                    continue

                # Créer l'analyse
                analyse = AnalyseEau(
                    type_prelevement=type_prelevement,
                    description_lieu=description_lieu,
                    date_prelevement=date_prelevement,
                    date_analyse=date_analyse,
                    date_resultat=date_resultat,
                    volume_analyse=int(volume_analyse),
                    observation=observation,
                    created_by=request.user if request.user.is_authenticated else None
                )

                # TNC
                if tnc_str and str(tnc_str).strip().lower() in ['oui', 'yes', 'true', '1']:
                    analyse.tnc = True
                else:
                    analyse.tnc = False
                    if nombre_coliformes is not None and str(nombre_coliformes).strip():
                        try:
                            analyse.nombre_coliformes = int(nombre_coliformes)
                        except ValueError:
                            errors.append(f"Ligne {row_idx}: Nombre coliformes invalide")
                            nb_errors += 1
                            continue

                # pH
                if ph is not None and str(ph).strip():
                    try:
                        analyse.ph = Decimal(str(ph))
                    except:
                        errors.append(f"Ligne {row_idx}: pH invalide")
                        nb_errors += 1
                        continue

                # Relations
                if equipement_id and str(equipement_id).strip():
                    try:
                        analyse.equipement = EquipementConstitutif.objects.get(id=equipement_id)
                    except EquipementConstitutif.DoesNotExist:
                        errors.append(f"Ligne {row_idx}: Équipement {equipement_id} introuvable")
                        nb_errors += 1
                        continue

                if point_eau_id and str(point_eau_id).strip():
                    try:
                        analyse.point_eau = PointEau.objects.get(id=point_eau_id)
                    except PointEau.DoesNotExist:
                        errors.append(f"Ligne {row_idx}: Point d'eau {point_eau_id} introuvable")
                        nb_errors += 1
                        continue

                analyse.save()
                nb_imported += 1

            except Exception as e:
                errors.append(f"Ligne {row_idx}: {str(e)}")
                nb_errors += 1

        return JsonResponse({
            'success': True,
            'message': f'{nb_imported} analyse(s) importée(s), {nb_errors} erreur(s)',
            'nb_imported': nb_imported,
            'nb_errors': nb_errors,
            'errors': errors
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)
