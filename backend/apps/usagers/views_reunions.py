"""
Vues pour la gestion des reunions avec les usagers/representants : CRUD reunion,
participants, decisions, pieces jointes, et pointage avec les lignes DiameX.
"""

import io
import json
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Case, IntegerField, Prefetch, Q, Sum, Value, When
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods

from apps.accounting.extracteurs_description import extraire_lieu_reunion
from apps.accounting.models import BaseDiameX
from apps.employees.models import Employe, Poste
from apps.geography.models import Commune, District, Fokontany, Region, Village
from apps.organization.models import Base, Ouvrage
from apps.organization.scoping import (
    filter_ouvrages_by_gestionnaire_period,
    get_user_accessible_gestionnaire_ids,
    get_user_organisation,
    is_partenaire_user,
)

from .models import (
    AvisRepresentant,
    Autorite,
    DecisionReunion,
    LigneDiameXIgnoreeReunion,
    ParticipantAttendu,
    ParticipantReunion,
    PieceJointeReunion,
    PointageReunionDiameX,
    RepresentantUsager,
    Reunion,
    TiersParticipant,
)


def _current_employe(request):
    return request.user.employe if hasattr(request.user, 'employe') and request.user.employe else None


def _gestionnaire_ids_organisation_connectee(user):
    """
    IDs d'organisation a utiliser pour compter les "ouvrages geres" de cette page : toujours
    ceux de l'organisation de l'employe connecte, meme pour un superuser.

    A la difference de get_user_accessible_gestionnaire_ids (qui renvoie None, donc "aucune
    restriction", pour un superuser), cette page doit refleter uniquement le perimetre reel de
    l'organisation connectee : sinon la stat compte les ouvrages geres par toutes les
    organisations de la plateforme, d'ou l'ecart constate (325 au lieu de 217).
    """
    organisation = get_user_organisation(user)
    if organisation is not None and not is_partenaire_user(user):
        return [str(organisation.id)]
    if organisation is not None:
        # Partenaire : ses ouvrages sont ceux des gestionnaires qui lui sont affectés.
        return get_user_accessible_gestionnaire_ids(user)

    # Pas d'employe/organisation rattache (ex: superuser technique sans fiche employe) :
    # on retombe sur le perimetre "accessible" standard plutot que de tout compter.
    return get_user_accessible_gestionnaire_ids(user)


def _ouvrages_geres_qs(gestionnaire_ids, annee=None):
    """
    Ouvrages geres par les organisations gestionnaire_ids sur l'annee donnee, avec la meme
    convention que le dashboard (cf. calculate_dashboard_stats_with_permissions dans
    hydria/views.py) : ouvrage rattache a une ModaliteGestion active/resiliee couvrant la
    periode, hors ouvrages de type "suspendu"/"potentiel", et disposant d'au moins un point
    d'eau. Sans cette coherence, le comptage de cette page ne correspond pas a celui du
    dashboard (https://www.hydria.mg/dashboard/).
    """
    ouvrages_qs = Ouvrage.objects.filter(pointeau__isnull=False).distinct()
    ouvrages_qs = filter_ouvrages_by_gestionnaire_period(ouvrages_qs, gestionnaire_ids, annee)
    ouvrages_qs = ouvrages_qs.exclude(
        Q(type_ouvrage__nom_type_ouvrage__icontains='suspend') |
        Q(type_ouvrage__nom_type_ouvrage__icontains='potentiel')
    )
    return ouvrages_qs.distinct()


def _synchroniser_participants_avis_batch(reunions):
    """
    Rattrapage en masse : les representants ayant un avis lie a une reunion (via la commande
    rattacher_avis_reunions ou une saisie manuelle) doivent apparaitre comme participants,
    meme si personne n'a explicitement ajoute leur fiche dans la section Participants.

    Fait en une seule passe (pas de requete par reunion) pour rester performant sur la page
    liste, contrairement au rattrapage unitaire de reunion_detail_json.
    """
    reunion_ids = [r.id for r in reunions]
    if not reunion_ids:
        return

    avis_par_reunion = {}
    for reunion_id, representant_id in AvisRepresentant.objects.filter(
        reunion_id__in=reunion_ids, representant__isnull=False
    ).values_list('reunion_id', 'representant_id'):
        avis_par_reunion.setdefault(reunion_id, set()).add(representant_id)

    if not avis_par_reunion:
        return

    participants_existants = {}
    for reunion_id, representant_id in ParticipantReunion.objects.filter(
        reunion_id__in=avis_par_reunion.keys(),
        type_participant=ParticipantReunion.TYPE_REPRESENTANT,
    ).values_list('reunion_id', 'representant_id'):
        participants_existants.setdefault(reunion_id, set()).add(representant_id)

    a_creer = []
    for reunion_id, representant_ids in avis_par_reunion.items():
        manquants = representant_ids - participants_existants.get(reunion_id, set())
        for representant_id in manquants:
            a_creer.append(ParticipantReunion(
                reunion_id=reunion_id,
                type_participant=ParticipantReunion.TYPE_REPRESENTANT,
                representant_id=representant_id,
            ))

    if a_creer:
        ParticipantReunion.objects.bulk_create(a_creer)


def _filtrer_ouvrages_par_zone(ouvrages_qs, commune_id=None, district_id=None, region_id=None):
    """
    Applique a un queryset d'ouvrages le filtre geographique du lieu le plus precis fourni
    (commune > district > region). Sans aucun filtre, retourne le queryset inchange (utilise
    aussi bien pour le filtre d'une reunion precise que pour le filtre de la page).
    """
    if commune_id:
        return ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune_id=commune_id)
    if district_id:
        return ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district_id=district_id)
    if region_id:
        return ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district__region_id=region_id)
    return ouvrages_qs


def _nb_ouvrages_geres_zone(gestionnaire_ids, commune_id=None, district_id=None, region_id=None, annee=None):
    """
    Compte les ouvrages geres (cf. _ouvrages_geres_qs) dans la zone geographique du lieu
    le plus precis fourni (commune > district > region), pour l'annee donnee (par defaut,
    celle de la reunion consideree). Sans lieu fourni, retourne 0 (utilise uniquement pour la
    colonne "Ouvrages" du tableau, propre a chaque reunion qui a toujours un lieu).
    """
    if not (commune_id or district_id or region_id):
        return 0

    ouvrages_qs = _filtrer_ouvrages_par_zone(
        _ouvrages_geres_qs(gestionnaire_ids, annee=annee),
        commune_id=commune_id, district_id=district_id, region_id=region_id,
    )
    return ouvrages_qs.distinct().count()


def _calculer_attendus_par_defaut(reunion, gestionnaire_ids):
    """
    Calcule les nombres de participants attendus par type, selon le type de reunion et son
    lieu, au moment de la creation (cf. demande utilisateur) :

    - Restitution communale : representantes = nb villages desservis par un ouvrage gere
      dans la commune (1 representante attendue par village) ; autorite (Maire) = 1 ;
      employe = 2. Les autres types (tiers, abonne) restent a 0 et sont saisis manuellement.
    - Restitution regionale : autorite (Maires) = nb communes de la region ayant au moins un
      ouvrage gere. Les autres types restent a 0, saisis manuellement.
    - Reunion d'equipe : rien de pre-rempli, tout est saisi manuellement.

    Retourne un dict {type_participant: nombre_attendu} ne contenant que les types a
    initialiser (les types absents restent a 0 / non crees, l'utilisateur peut les ajouter).
    """
    annee = reunion.date_reunion.year if reunion.date_reunion else None

    if reunion.type_reunion == Reunion.TYPE_RESTITUTION_COMMUNALE and reunion.commune_id:
        from apps.geography.models import Village

        ouvrages_geres_commune = _ouvrages_geres_qs(gestionnaire_ids, annee=annee).filter(
            pointeau__parcelle__village__fokontany__commune_id=reunion.commune_id
        )
        nb_villages_desservis = Village.objects.filter(
            parcelles__pointeau__ouvrage__in=ouvrages_geres_commune
        ).distinct().count()

        return {
            ParticipantReunion.TYPE_REPRESENTANT: nb_villages_desservis,
            ParticipantReunion.TYPE_AUTORITE: 1,  # Le maire
            ParticipantReunion.TYPE_EMPLOYE: 2,
        }

    if reunion.type_reunion == Reunion.TYPE_RESTITUTION_REGIONALE and reunion.region_id:
        ouvrages_geres_region = _ouvrages_geres_qs(gestionnaire_ids, annee=annee).filter(
            pointeau__parcelle__village__fokontany__commune__district__region_id=reunion.region_id
        )
        nb_communes_avec_ouvrage = Commune.objects.filter(
            fokontanys__villages__parcelles__pointeau__ouvrage__in=ouvrages_geres_region
        ).distinct().count()

        return {
            ParticipantReunion.TYPE_AUTORITE: nb_communes_avec_ouvrage,  # Un maire par commune
        }

    # Reunion d'equipe (ou type non reconnu / lieu manquant) : rien de pre-rempli.
    return {}


def _appliquer_attendus_par_defaut(reunion, gestionnaire_ids):
    """Cree les lignes ParticipantAttendu par defaut pour une reunion nouvellement creee."""
    attendus = _calculer_attendus_par_defaut(reunion, gestionnaire_ids)
    for type_participant, nombre in attendus.items():
        ParticipantAttendu.objects.update_or_create(
            reunion=reunion,
            type_participant=type_participant,
            defaults={'nombre_attendu': nombre},
        )


@login_required
def ouvrages_geres_detail_json(request):
    """
    Liste detaillee des ouvrages comptes dans la stat "Ouvrages geres" de la page reunions,
    pour diagnostiquer les ecarts avec le dashboard (cf. _ouvrages_geres_qs). Renvoie pour
    chaque ouvrage son type, sa/ses modalite(s) de gestion (gestionnaire, statut, dates) afin
    d'identifier ceux qui ne devraient pas etre comptes.
    """
    annee_param = request.GET.get('annee', '').strip()
    try:
        annee = int(annee_param) if annee_param else timezone.localdate().year
    except ValueError:
        annee = timezone.localdate().year

    # Meme filtre geographique que la page (region/district/commune), pour que le detail du
    # modal corresponde exactement au total affiche sur la carte de stat au moment du clic.
    region_id = request.GET.get('region', '').strip()
    district_id = request.GET.get('district', '').strip()
    commune_id = request.GET.get('commune', '').strip()

    gestionnaire_ids = _gestionnaire_ids_organisation_connectee(request.user)
    ouvrages_qs = _filtrer_ouvrages_par_zone(
        _ouvrages_geres_qs(gestionnaire_ids, annee=annee),
        commune_id=commune_id, district_id=district_id, region_id=region_id,
    ).select_related(
        'type_ouvrage'
    ).prefetch_related('modalites_gestion', 'modalites_gestion__gestionnaire').order_by('nom_ouvrage')

    data = []
    for ouvrage in ouvrages_qs:
        modalites = [
            {
                'gestionnaire': m.gestionnaire.nom_organisation if m.gestionnaire else None,
                'gestionnaire_id': str(m.gestionnaire_id) if m.gestionnaire_id else None,
                'statut': m.statut,
                'date_debut': m.date_debut.isoformat() if m.date_debut else None,
                'date_fin': m.date_fin.isoformat() if m.date_fin else None,
            }
            for m in ouvrage.modalites_gestion.all()
        ]
        data.append({
            'id': str(ouvrage.id),
            'nom': ouvrage.nom_ouvrage,
            'type': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else None,
            'modalites': modalites,
        })

    return JsonResponse({
        'success': True,
        'annee': annee,
        'gestionnaire_ids': gestionnaire_ids,
        'total': len(data),
        'ouvrages': data,
    })


@login_required
def reunions_list(request):
    """Page principale de gestion des reunions, avec filtres geographiques.

    La liste est toujours limitee a une seule annee a la fois (annee courante par
    defaut) pour eviter de charger tout l'historique des reunions.
    """
    region_id = request.GET.get('region', '').strip()
    district_id = request.GET.get('district', '').strip()
    commune_id = request.GET.get('commune', '').strip()
    type_reunion = request.GET.get('type_reunion', '').strip()
    search = request.GET.get('search', '').strip()
    annee_param = request.GET.get('annee', '').strip()

    try:
        annee = int(annee_param) if annee_param else timezone.localdate().year
    except ValueError:
        annee = timezone.localdate().year

    from django.db.models import Prefetch

    reunions = Reunion.objects.select_related(
        'commune', 'commune__district', 'district', 'district__region', 'region', 'organise_par'
    ).prefetch_related(
        'participants',
        'attendus',
        'decisions',
        'avis_recueillis',
        Prefetch(
            'pieces_jointes',
            queryset=PieceJointeReunion.objects.filter(type_piece=PieceJointeReunion.TYPE_PHOTO).order_by('created_at'),
            to_attr='photos_couverture',
        ),
    ).filter(date_reunion__year=annee)

    # Une reunion peut etre rattachee a une commune, un district ou une region (au moins un des trois).
    # Le filtre region doit donc remonter la chaine commune->district->region ou district->region.
    if region_id:
        reunions = reunions.filter(
            Q(region_id=region_id)
            | Q(district__region_id=region_id)
            | Q(commune__district__region_id=region_id)
        )
    if district_id:
        reunions = reunions.filter(
            Q(district_id=district_id)
            | Q(commune__district_id=district_id)
        )
    if commune_id:
        reunions = reunions.filter(commune_id=commune_id)
    if type_reunion:
        reunions = reunions.filter(type_reunion=type_reunion)
    if search:
        reunions = reunions.filter(
            Q(objet__icontains=search) |
            Q(compte_rendu__icontains=search) |
            Q(commune__nom_commune__icontains=search) |
            Q(district__nom_district__icontains=search) |
            Q(region__nom_region__icontains=search)
        )

    reunions = list(reunions.order_by('-date_reunion'))

    # Rattrapage : les representants ayant un avis lie a une reunion doivent compter comme
    # participants (cf. _synchroniser_participants_avis_batch), sinon la colonne Participants
    # et les stats de la page restent a 0 tant que personne n'a ouvert le detail de la reunion.
    _synchroniser_participants_avis_batch(reunions)

    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.select_related('region').order_by('nom_district')
    communes = Commune.objects.select_related('district').order_by('nom_commune')

    annee_courante = timezone.localdate().year
    annees_disponibles = list(range(annee_courante, annee_courante - 6, -1))
    if annee not in annees_disponibles:
        annees_disponibles = sorted(set(annees_disponibles + [annee]), reverse=True)

    # Organisation connectee : sert a compter les ouvrages geres (colonne + stats)
    gestionnaire_ids = _gestionnaire_ids_organisation_connectee(request.user)

    # Cache du nombre d'ouvrages geres par (commune, district, region, annee) : plusieurs
    # reunions partagent souvent le meme lieu/annee (ex: plusieurs reunions dans la meme
    # commune), et _nb_ouvrages_geres_zone fait une requete SQL avec jointures lourdes
    # (Ouvrage -> PointEau -> Parcelle -> Village -> Fokontany -> Commune/District/Region).
    # Sans ce cache, une page avec N reunions declenche N requetes de ce type (N+1), ce qui
    # peut faire timeout Nginx/Gunicorn sur une annee avec beaucoup de reunions.
    nb_ouvrages_geres_cache = {}

    def _nb_ouvrages_geres_cache(commune_id, district_id, region_id, annee_reunion):
        cle = (commune_id, district_id, region_id, annee_reunion)
        if cle not in nb_ouvrages_geres_cache:
            nb_ouvrages_geres_cache[cle] = _nb_ouvrages_geres_zone(
                gestionnaire_ids,
                commune_id=commune_id, district_id=district_id, region_id=region_id,
                annee=annee_reunion,
            )
        return nb_ouvrages_geres_cache[cle]

    # Nombre d'ouvrages geres par reunion, selon le lieu de chaque reunion (commune > district > region)
    nb_participants_total = 0
    total_avis = 0
    total_criteres_negatifs = 0
    total_criteres = 0
    for reunion in reunions:
        reunion.nb_ouvrages_geres = _nb_ouvrages_geres_cache(
            str(reunion.commune_id) if reunion.commune_id else None,
            str(reunion.district_id) if reunion.district_id else None,
            str(reunion.region_id) if reunion.region_id else None,
            reunion.date_reunion.year if reunion.date_reunion else annee,
        )
        nb_participants_total += reunion.nb_participants

        attendus_reunion = list(reunion.attendus.all())
        reunion.total_attendu = sum(a.nombre_attendu for a in attendus_reunion)
        reunion.nb_participants_presents = sum(1 for p in reunion.participants.all() if p.present)
        reunion.pct_participation = (
            round((reunion.nb_participants_presents / reunion.total_attendu) * 100)
            if reunion.total_attendu else None
        )

        # Taux de criteres negatifs : agregation sur les 7 criteres de chaque avis (4 booleens
        # + 3 champs probleme), pas sur l'avis dans son ensemble (cf. AvisRepresentant.
        # nb_criteres_negatifs). Un avis avec 2 criteres negatifs sur 7 pese differemment
        # d'un avis avec 6 criteres negatifs sur 7, meme s'ils seraient tous deux "negatifs"
        # avec l'ancienne logique tout-ou-rien.
        avis_reunion = list(reunion.avis_recueillis.all())
        nb_avis = len(avis_reunion)
        reunion_criteres_negatifs = sum(a.nb_criteres_negatifs for a in avis_reunion)
        reunion_criteres_total = nb_avis * AvisRepresentant.NB_CRITERES_TOTAL
        reunion.nb_avis = nb_avis
        reunion.nb_criteres_negatifs = reunion_criteres_negatifs
        reunion.nb_criteres_total = reunion_criteres_total
        reunion.pct_avis_negatifs = (
            round((reunion_criteres_negatifs / reunion_criteres_total) * 100)
            if reunion_criteres_total else None
        )
        total_avis += nb_avis
        total_criteres_negatifs += reunion_criteres_negatifs
        total_criteres += reunion_criteres_total

    nb_reunions = len(reunions)
    nb_reunions_par_type = {
        value: sum(1 for r in reunions if r.type_reunion == value)
        for value, _label in Reunion.TYPE_CHOICES
    }
    nb_reunions_sans_type = sum(1 for r in reunions if not r.type_reunion)

    # Le total "Ouvrages geres" doit suivre le filtre geographique de la page (region/district/
    # commune) : sans filtre, c'est le total de l'organisation ; avec filtre, uniquement la zone
    # selectionnee. Idem pour le modal de detail (cf. ouvrages_geres_detail_json).
    total_ouvrages_geres = _filtrer_ouvrages_par_zone(
        _ouvrages_geres_qs(gestionnaire_ids, annee=annee),
        commune_id=commune_id, district_id=district_id, region_id=region_id,
    ).distinct().count()

    zone_ouvrages_geres_nom = None
    if commune_id:
        commune_obj = next((c for c in communes if str(c.id) == commune_id), None)
        zone_ouvrages_geres_nom = commune_obj.nom_commune if commune_obj else None
    elif district_id:
        district_obj = next((d for d in districts if str(d.id) == district_id), None)
        zone_ouvrages_geres_nom = district_obj.nom_district if district_obj else None
    elif region_id:
        region_obj = next((r for r in regions if str(r.id) == region_id), None)
        zone_ouvrages_geres_nom = region_obj.nom_region if region_obj else None

    # Totaux presents/attendus par type de participant, tous types confondus dans les
    # reunions filtrees. Types masques cote template si nombre_attendu total = 0.
    attendus_par_type_totaux = {value: 0 for value, _label in ParticipantReunion.TYPE_CHOICES}
    presents_par_type_totaux = {value: 0 for value, _label in ParticipantReunion.TYPE_CHOICES}
    for reunion in reunions:
        for a in reunion.attendus.all():
            attendus_par_type_totaux[a.type_participant] = attendus_par_type_totaux.get(a.type_participant, 0) + a.nombre_attendu
        for p in reunion.participants.all():
            if p.present:
                presents_par_type_totaux[p.type_participant] = presents_par_type_totaux.get(p.type_participant, 0) + 1

    attendus_stats = [
        {
            'type_participant': value,
            'type_display': label,
            'present': presents_par_type_totaux.get(value, 0),
            'attendu': attendus_par_type_totaux.get(value, 0),
        }
        for value, label in ParticipantReunion.TYPE_CHOICES
        if attendus_par_type_totaux.get(value, 0) > 0
    ]
    total_attendu_general = sum(attendus_par_type_totaux.values())
    total_present_attendu_general = sum(
        presents_par_type_totaux[v] for v, _l in ParticipantReunion.TYPE_CHOICES if attendus_par_type_totaux.get(v, 0) > 0
    )

    stats = {
        'nb_reunions': nb_reunions,
        'nb_reunions_par_type': nb_reunions_par_type,
        'nb_reunions_sans_type': nb_reunions_sans_type,
        'nb_participants_total': nb_participants_total,
        'moyenne_participants': round(nb_participants_total / nb_reunions, 1) if nb_reunions else 0,
        'total_avis': total_avis,
        'total_criteres_negatifs': total_criteres_negatifs,
        'total_criteres': total_criteres,
        'pct_avis_negatifs': round((total_criteres_negatifs / total_criteres) * 100) if total_criteres else 0,
        'total_ouvrages_geres': total_ouvrages_geres,
        'zone_ouvrages_geres_nom': zone_ouvrages_geres_nom,
        'attendus_par_type': attendus_stats,
        'total_attendu_general': total_attendu_general,
        'total_present_attendu_general': total_present_attendu_general,
    }

    context = {
        'reunions': reunions,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'type_reunion_choices': Reunion.TYPE_CHOICES,
        'autorite_type_choices': Autorite.get_type_choices(),
        'filter_region': region_id,
        'filter_district': district_id,
        'filter_commune': commune_id,
        'filter_type_reunion': type_reunion,
        'filter_search': search,
        'filter_annee': annee,
        'annees_disponibles': annees_disponibles,
        'stats': stats,
        'employes': Employe.objects.filter(salarie=True).order_by('surnom'),
        'current_employe_id': str(request.user.employe.id) if hasattr(request.user, 'employe') and request.user.employe else '',
    }
    return render(request, 'usagers/reunions_list.html', context)


@login_required
@require_http_methods(["POST"])
def reunion_create(request):
    try:
        data = json.loads(request.body)

        commune_id = (data.get('commune_id') or '').strip()
        district_id = (data.get('district_id') or '').strip()
        region_id = (data.get('region_id') or '').strip()
        objet = (data.get('objet') or '').strip()
        date_reunion = (data.get('date_reunion') or '').strip()

        if not (commune_id or district_id or region_id) or not objet or not date_reunion:
            return JsonResponse({
                'success': False,
                'error': 'Commune (ou District ou Region), objet et date sont obligatoires.'
            }, status=400)

        commune = Commune.objects.filter(id=commune_id).first() if commune_id else None
        district = District.objects.filter(id=district_id).first() if district_id else None
        region = Region.objects.filter(id=region_id).first() if region_id else None

        organise_par = None
        organise_par_id = (data.get('organise_par_id') or '').strip()
        if organise_par_id:
            organise_par = Employe.objects.filter(id=organise_par_id).first()

        type_reunion = (data.get('type_reunion') or '').strip()
        if type_reunion and type_reunion not in dict(Reunion.TYPE_CHOICES):
            return JsonResponse({'success': False, 'error': 'Type de reunion invalide.'}, status=400)

        reunion = Reunion.objects.create(
            type_reunion=type_reunion or None,
            commune=commune,
            district=district,
            region=region,
            date_reunion=date_reunion,
            lieu=(data.get('lieu') or '').strip() or None,
            objet=objet,
            compte_rendu=(data.get('compte_rendu') or '').strip() or None,
            recommandations=(data.get('recommandations') or '').strip() or None,
            organise_par=organise_par,
            created_by=_current_employe(request),
        )

        gestionnaire_ids = _gestionnaire_ids_organisation_connectee(request.user)
        _appliquer_attendus_par_defaut(reunion, gestionnaire_ids)

        return JsonResponse({'success': True, 'reunion_id': str(reunion.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def reunion_update(request, reunion_id):
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        data = json.loads(request.body)

        if 'commune_id' in data:
            commune_id = (data.get('commune_id') or '').strip()
            reunion.commune = Commune.objects.filter(id=commune_id).first() if commune_id else None
        if 'district_id' in data:
            district_id = (data.get('district_id') or '').strip()
            reunion.district = District.objects.filter(id=district_id).first() if district_id else None
        if 'region_id' in data:
            region_id = (data.get('region_id') or '').strip()
            reunion.region = Region.objects.filter(id=region_id).first() if region_id else None
        if not (reunion.commune_id or reunion.district_id or reunion.region_id):
            return JsonResponse({
                'success': False,
                'error': 'La reunion doit avoir au moins une Commune, un District ou une Region.'
            }, status=400)
        if 'date_reunion' in data and data['date_reunion']:
            reunion.date_reunion = data['date_reunion']
        if 'lieu' in data:
            reunion.lieu = (data.get('lieu') or '').strip() or None
        if 'objet' in data and data['objet']:
            reunion.objet = data['objet'].strip()
        if 'compte_rendu' in data:
            reunion.compte_rendu = (data.get('compte_rendu') or '').strip() or None
        if 'recommandations' in data:
            reunion.recommandations = (data.get('recommandations') or '').strip() or None
        if 'organise_par_id' in data:
            organise_par_id = (data.get('organise_par_id') or '').strip()
            reunion.organise_par = Employe.objects.filter(id=organise_par_id).first() if organise_par_id else None
        if 'type_reunion' in data:
            type_reunion = (data.get('type_reunion') or '').strip()
            if type_reunion and type_reunion not in dict(Reunion.TYPE_CHOICES):
                return JsonResponse({'success': False, 'error': 'Type de reunion invalide.'}, status=400)
            reunion.type_reunion = type_reunion or None

        reunion.save()

        # Si aucun attendu n'a encore ete saisi/calcule pour cette reunion (ex: type ou lieu
        # renseigne seulement a cette modification), on applique les valeurs par defaut. On ne
        # recalcule jamais si des attendus existent deja, pour ne pas ecraser une saisie
        # manuelle de l'utilisateur.
        if not reunion.attendus.exists():
            gestionnaire_ids = _gestionnaire_ids_organisation_connectee(request.user)
            _appliquer_attendus_par_defaut(reunion, gestionnaire_ids)

        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def reunion_delete(request, reunion_id):
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        reunion.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def reunions_fusionnables_json(request, reunion_id):
    """
    Liste les autres reunions partageant la meme date et la meme commune que la reunion
    donnee, candidates a une fusion (cf. reunion_fusionner). Reserve au directeur cote
    frontend (bouton masque), mais la lecture reste ouverte a tout utilisateur connecte :
    seule l'action de fusion elle-meme est restreinte cote serveur.
    """
    reunion = get_object_or_404(Reunion, id=reunion_id)

    if not reunion.commune_id:
        return JsonResponse({'success': True, 'results': [], 'message': "Cette reunion n'a pas de commune renseignee."})

    candidates_qs = Reunion.objects.filter(
        commune_id=reunion.commune_id,
        date_reunion=reunion.date_reunion,
    ).exclude(id=reunion.id).select_related('organise_par')

    results = []
    for r in candidates_qs.order_by('objet'):
        results.append({
            'id': str(r.id),
            'objet': r.objet,
            'type_reunion_display': r.get_type_reunion_display() if r.type_reunion else '',
            'nb_participants': r.nb_participants,
            'nb_decisions': r.decisions.count(),
            'nb_avis': r.avis_recueillis.count(),
            'nb_pointages': r.pointages_diamex.count(),
        })

    return JsonResponse({'success': True, 'results': results})


@login_required
@require_http_methods(["POST"])
def reunion_fusionner(request, reunion_id):
    """
    Fusionne une reunion secondaire dans la reunion principale (reunion_id) : deplace
    participants (sans doublons), attendus (sommes par type), decisions, avis (sans
    violer la contrainte unique representant+date_avis), pieces jointes et pointages
    DiameX, puis concatene compte-rendu/recommandations si distincts. La reunion
    secondaire est ensuite supprimee.

    Reserve au poste Directeur (cf. is_directeur_poste) : une fusion est destructive
    (suppression de la reunion secondaire) et doit rester une action encadree.
    """
    from apps.employees.permissions import is_directeur_poste

    if not is_directeur_poste(request.user):
        return JsonResponse({
            'success': False,
            'error': "Seul le Directeur peut fusionner des reunions."
        }, status=403)

    try:
        reunion_principale = get_object_or_404(Reunion, id=reunion_id)
        data = json.loads(request.body)
        reunion_secondaire_id = (data.get('reunion_secondaire_id') or '').strip()
        if not reunion_secondaire_id:
            return JsonResponse({'success': False, 'error': 'Reunion a fusionner obligatoire.'}, status=400)
        if reunion_secondaire_id == str(reunion_principale.id):
            return JsonResponse({'success': False, 'error': 'Impossible de fusionner une reunion avec elle-meme.'}, status=400)

        reunion_secondaire = get_object_or_404(Reunion, id=reunion_secondaire_id)

        if reunion_secondaire.commune_id != reunion_principale.commune_id or reunion_secondaire.date_reunion != reunion_principale.date_reunion:
            return JsonResponse({
                'success': False,
                'error': 'Les deux reunions doivent avoir la meme date et la meme commune pour etre fusionnees.'
            }, status=400)

        with transaction.atomic():
            # --- Participants : deplacer en evitant les doublons (meme type + meme fiche liee,
            # ou meme nom_libre pour les tiers historiques sans fiche) ---
            cles_existantes = set()
            for p in reunion_principale.participants.all():
                cle = (
                    p.type_participant,
                    p.representant_id, p.point_eau_id, p.employe_id, p.autorite_id, p.tiers_id,
                    (p.nom_libre or '').strip().lower() if not any([
                        p.representant_id, p.point_eau_id, p.employe_id, p.autorite_id, p.tiers_id
                    ]) else None,
                )
                cles_existantes.add(cle)

            nb_participants_fusionnes = 0
            nb_participants_doublons = 0
            for p in list(reunion_secondaire.participants.all()):
                cle = (
                    p.type_participant,
                    p.representant_id, p.point_eau_id, p.employe_id, p.autorite_id, p.tiers_id,
                    (p.nom_libre or '').strip().lower() if not any([
                        p.representant_id, p.point_eau_id, p.employe_id, p.autorite_id, p.tiers_id
                    ]) else None,
                )
                if cle in cles_existantes:
                    p.delete()
                    nb_participants_doublons += 1
                    continue
                cles_existantes.add(cle)
                p.reunion = reunion_principale
                p.save(update_fields=['reunion'])
                nb_participants_fusionnes += 1

            # --- Participants attendus : sommer par type plutot que dupliquer les lignes ---
            for attendu_secondaire in list(reunion_secondaire.attendus.all()):
                attendu_principal = reunion_principale.attendus.filter(
                    type_participant=attendu_secondaire.type_participant
                ).first()
                if attendu_principal:
                    attendu_principal.nombre_attendu += attendu_secondaire.nombre_attendu
                    attendu_principal.save(update_fields=['nombre_attendu'])
                    attendu_secondaire.delete()
                else:
                    attendu_secondaire.reunion = reunion_principale
                    attendu_secondaire.save(update_fields=['reunion'])

            # --- Decisions : deplacement simple, aucune contrainte d'unicite ---
            nb_decisions = reunion_secondaire.decisions.update(reunion=reunion_principale)

            # --- Avis : rattacher a la principale sauf si le representant a deja un avis a
            # cette date sur la principale (contrainte unique representant+date_avis) : dans
            # ce cas, l'avis de la secondaire est laisse tel quel (reste rattache a une
            # reunion qui va etre supprimee -> on le detache plutot que de le perdre) ---
            representants_avec_avis_principal = set(
                reunion_principale.avis_recueillis.values_list('representant_id', flat=True)
            )
            nb_avis_fusionnes = 0
            nb_avis_conflits = 0
            for avis in list(reunion_secondaire.avis_recueillis.all()):
                if avis.representant_id in representants_avec_avis_principal:
                    # Conflit : on detache l'avis (reunion=None) plutot que de le perdre a la
                    # suppression de la reunion secondaire (reunion a NULL est autorise sur
                    # AvisRepresentant.reunion, cf. modele).
                    avis.reunion = None
                    avis.save(update_fields=['reunion'])
                    nb_avis_conflits += 1
                else:
                    avis.reunion = reunion_principale
                    avis.save(update_fields=['reunion'])
                    representants_avec_avis_principal.add(avis.representant_id)
                    nb_avis_fusionnes += 1

            # --- Pieces jointes : deplacement simple ---
            nb_pieces = reunion_secondaire.pieces_jointes.update(reunion=reunion_principale)

            # --- Pointages DiameX : la contrainte unique porte sur (numero_diamex,
            # base_diamex), pas sur la reunion : un changement de reunion ne peut donc pas
            # entrer en conflit avec un pointage existant de la principale ---
            nb_pointages = reunion_secondaire.pointages_diamex.update(reunion=reunion_principale)

            # --- Compte-rendu / recommandations : concatener si le contenu de la secondaire
            # apporte quelque chose de nouveau (evite de dupliquer si deja identique/vide) ---
            def _concatener(valeur_principale, valeur_secondaire, titre_secondaire):
                valeur_principale = (valeur_principale or '').strip()
                valeur_secondaire = (valeur_secondaire or '').strip()
                if not valeur_secondaire or valeur_secondaire == valeur_principale:
                    return valeur_principale or None
                if not valeur_principale:
                    return valeur_secondaire
                return f"{valeur_principale}\n\n--- {titre_secondaire} ({reunion_secondaire.objet}) ---\n{valeur_secondaire}"

            reunion_principale.compte_rendu = _concatener(
                reunion_principale.compte_rendu, reunion_secondaire.compte_rendu, 'Compte-rendu fusionne'
            )
            reunion_principale.recommandations = _concatener(
                reunion_principale.recommandations, reunion_secondaire.recommandations, 'Recommandations fusionnees'
            )
            reunion_principale.save(update_fields=['compte_rendu', 'recommandations'])

            reunion_secondaire_objet = reunion_secondaire.objet
            reunion_secondaire.delete()

        return JsonResponse({
            'success': True,
            'message': f"Reunion \"{reunion_secondaire_objet}\" fusionnee dans \"{reunion_principale.objet}\".",
            'nb_participants_fusionnes': nb_participants_fusionnes,
            'nb_participants_doublons': nb_participants_doublons,
            'nb_decisions': nb_decisions,
            'nb_avis_fusionnes': nb_avis_fusionnes,
            'nb_avis_conflits': nb_avis_conflits,
            'nb_pieces': nb_pieces,
            'nb_pointages': nb_pointages,
        })
    except Reunion.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Reunion introuvable.'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def reunion_detail_json(request, reunion_id):
    """Detail complet d'une reunion pour affichage dans le modal."""
    reunion = get_object_or_404(
        Reunion.objects.select_related(
            'commune', 'commune__district', 'district', 'district__region', 'region', 'organise_par'
        ),
        id=reunion_id
    )

    # Rattrapage : les representants ayant un avis lie a cette reunion doivent apparaitre
    # comme participants, meme si personne n'a explicitement ajoute leur fiche (cf. fonction
    # partagee avec reunions_list, qui fait le meme rattrapage en masse pour tout le tableau).
    _synchroniser_participants_avis_batch([reunion])

    # Comptage presents/attendus par type de participant : les attendus sont pre-remplis a la
    # creation (cf. _calculer_attendus_par_defaut) mais restent editables manuellement.
    attendus_par_type = {
        a.type_participant: a.nombre_attendu for a in reunion.attendus.all()
    }
    presents_par_type = {}
    for type_value, _label in ParticipantReunion.TYPE_CHOICES:
        presents_par_type[type_value] = reunion.participants.filter(
            type_participant=type_value, present=True
        ).count()

    attendus = [
        {
            'type_participant': type_value,
            'type_display': label,
            'present': presents_par_type.get(type_value, 0),
            'attendu': attendus_par_type.get(type_value, 0),
        }
        for type_value, label in ParticipantReunion.TYPE_CHOICES
    ]

    participants = []
    for p in reunion.participants.select_related('representant', 'point_eau', 'employe', 'autorite').all():
        participants.append({
            'id': str(p.id),
            'type_participant': p.type_participant,
            'type_display': p.get_type_participant_display(),
            'nom_affiche': p.nom_affiche,
            'fonction_libre': p.fonction_libre or '',
            'contact_libre': p.contact_libre or '',
            'present': p.present,
        })

    decisions = []
    for d in reunion.decisions.select_related('responsable').all():
        decisions.append({
            'id': str(d.id),
            'description': d.description,
            'responsable': d.responsable.surnom or d.responsable.nom if d.responsable else '',
            'responsable_id': str(d.responsable_id) if d.responsable_id else '',
            'date_echeance': d.date_echeance.isoformat() if d.date_echeance else None,
            'statut': d.statut,
            'statut_display': d.get_statut_display(),
            'date_realisation': d.date_realisation.isoformat() if d.date_realisation else None,
            'commentaire_suivi': d.commentaire_suivi or '',
            'est_en_retard': d.est_en_retard,
        })

    pieces_jointes = []
    for pj in reunion.pieces_jointes.all():
        pieces_jointes.append({
            'id': str(pj.id),
            'type_piece': pj.type_piece,
            'type_display': pj.get_type_piece_display(),
            'url': pj.fichier.url if pj.fichier else '',
            'legende': pj.legende or '',
        })

    avis_qs = reunion.avis_recueillis.select_related(
        'representant',
        'representant__village',
        'representant__village__fokontany',
        'representant__village__fokontany__commune',
    ).all()

    # Resoudre le(s) nom(s) d'ouvrage lie(s) au village de chaque representant
    # (Village -> Parcelle -> PointEau -> Ouvrage), pour affichage dans le detail des avis.
    village_ids = {
        a.representant.village_id
        for a in avis_qs
        if a.representant and a.representant.village_id
    }
    ouvrages_par_village = {}
    if village_ids:
        from apps.billing.models import PointEau
        rows = PointEau.objects.filter(
            parcelle__village_id__in=village_ids,
            ouvrage__isnull=False,
        ).values_list('parcelle__village_id', 'ouvrage__nom_ouvrage').distinct()
        for village_id, ouvrage_nom in rows:
            if not ouvrage_nom:
                continue
            ouvrages_par_village.setdefault(village_id, set()).add(ouvrage_nom)

    avis = []
    for a in avis_qs:
        village_id = a.representant.village_id if a.representant else None
        ouvrages_noms = sorted(ouvrages_par_village.get(village_id, set()))
        avis.append({
            'id': str(a.id),
            'representant': a.representant.nom if a.representant else '',
            'representant_id': str(a.representant_id) if a.representant_id else '',
            'village': a.representant.village.nom_village if a.representant and a.representant.village else '',
            'ouvrages': ouvrages_noms,
            'ouvrages_display': ', '.join(ouvrages_noms) if ouvrages_noms else '-',
            'date_avis': a.date_avis.isoformat() if a.date_avis else None,
            'score_satisfaction': a.score_satisfaction,
            'pourcentage_satisfaction': a.pourcentage_satisfaction,
            'quantite_ok': a.quantite_ok,
            'qualite_ok': a.qualite_ok,
            'delai_reparation_ok': a.delai_reparation_ok,
            'relation_gestionnaire_ok': a.relation_gestionnaire_ok,
            'probleme_technique': a.probleme_technique or '',
            'probleme_recouvrement': a.probleme_recouvrement or '',
            'probleme_regles_utilisation': a.probleme_regles_utilisation or '',
            'remarque': a.remarque or '',
        })

    pointages = []
    for pt in reunion.pointages_diamex.all():
        pointages.append({
            'id': str(pt.id),
            'entry_id': pt.entry_id,
            'numero_diamex': pt.numero_diamex,
            'base_diamex': pt.base_diamex,
            'date_transaction': pt.date_transaction.isoformat() if pt.date_transaction else None,
            'description': pt.description or '',
            'compte_code': pt.compte_code or '',
            'montant': float(pt.montant or 0),
        })

    # District/region "effectifs" : la reunion ne renseigne parfois que la commune (le
    # district/la region s'en deduisent alors), remontee ici pour pre-remplir correctement le
    # formulaire de creation d'autorite (cf. openCreateFicheModal dans reunions_list.html),
    # qui a besoin de la chaine complete region > district > commune > fokontany.
    district_effectif = reunion.district or (reunion.commune.district if reunion.commune_id and reunion.commune else None)
    region_effective = reunion.region or (district_effectif.region if district_effectif else None)

    return JsonResponse({
        'success': True,
        'reunion': {
            'id': str(reunion.id),
            'type_reunion': reunion.type_reunion or '',
            'type_reunion_display': reunion.get_type_reunion_display() if reunion.type_reunion else '',
            'objet': reunion.objet,
            'date_reunion': reunion.date_reunion.isoformat(),
            'commune_id': str(reunion.commune_id) if reunion.commune_id else '',
            'commune_nom': reunion.commune.nom_commune if reunion.commune else '',
            'district_id': str(reunion.district_id) if reunion.district_id else '',
            'district_nom': reunion.district.nom_district if reunion.district else (reunion.commune.district.nom_district if reunion.commune and reunion.commune.district else ''),
            'region_id': str(reunion.region_id) if reunion.region_id else '',
            'region_nom': reunion.region.nom_region if reunion.region else '',
            'district_effectif_id': str(district_effectif.id) if district_effectif else '',
            'region_effective_id': str(region_effective.id) if region_effective else '',
            'lieu_affiche': reunion.lieu_affiche,
            'lieu': reunion.lieu or '',
            'compte_rendu': reunion.compte_rendu or '',
            'recommandations': reunion.recommandations or '',
            'organise_par': reunion.organise_par.surnom or reunion.organise_par.nom if reunion.organise_par else '',
            'organise_par_id': str(reunion.organise_par_id) if reunion.organise_par_id else '',
            'montant_depense_total': float(reunion.montant_depense_total),
        },
        'participants': participants,
        'attendus': attendus,
        'decisions': decisions,
        'pieces_jointes': pieces_jointes,
        'avis': avis,
        'pointages_diamex': pointages,
    })


@login_required
@require_http_methods(["POST"])
def attendu_update(request, reunion_id):
    """
    Cree ou met a jour le nombre de participants attendus pour un type donne, sur une reunion.
    Permet a l'utilisateur de corriger/completer les valeurs pre-remplies automatiquement
    (cf. _calculer_attendus_par_defaut), ou de saisir les types laisses a 0 (tiers, employes
    supplementaires, etc.).
    """
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        data = json.loads(request.body)

        type_participant = (data.get('type_participant') or '').strip()
        if type_participant not in dict(ParticipantReunion.TYPE_CHOICES):
            return JsonResponse({'success': False, 'error': 'Type de participant invalide.'}, status=400)

        try:
            nombre_attendu = int(data.get('nombre_attendu'))
        except (TypeError, ValueError):
            return JsonResponse({'success': False, 'error': 'Nombre attendu invalide.'}, status=400)

        if nombre_attendu < 0:
            return JsonResponse({'success': False, 'error': 'Le nombre attendu ne peut pas etre negatif.'}, status=400)

        ParticipantAttendu.objects.update_or_create(
            reunion=reunion,
            type_participant=type_participant,
            defaults={'nombre_attendu': nombre_attendu},
        )

        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def export_reunions_excel(request):
    """
    Exporte vers Excel les reunions selectionnees (via ?ids=uuid1,uuid2,...) : une feuille
    "Reunions" (date, nom/objet, compte-rendu, recommandations, participants attendus, taux
    de criteres d'avis negatifs), une feuille "Avis" (un avis de representant par ligne,
    rattache a sa reunion) et une feuille "Recap par region" (totaux agreges par region).
    """
    ids_param = (request.GET.get('ids') or '').strip()
    ids = [v.strip() for v in ids_param.split(',') if v.strip()]
    if not ids:
        return JsonResponse({'success': False, 'error': 'Aucune reunion selectionnee.'}, status=400)

    reunions = Reunion.objects.select_related(
        'commune', 'commune__district', 'district', 'district__region', 'region', 'organise_par'
    ).prefetch_related(
        'participants',
        'attendus',
        Prefetch(
            'avis_recueillis',
            queryset=AvisRepresentant.objects.select_related('representant', 'representant__village', 'employe').order_by('date_avis'),
        ),
    ).filter(id__in=ids).order_by('date_reunion')

    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    wrap_top = Alignment(vertical='top', wrap_text=True)

    def write_header(ws, headers):
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        ws.freeze_panes = 'A2'

    # --- Feuille 1 : Reunions ---
    gestionnaire_ids = _gestionnaire_ids_organisation_connectee(request.user)

    ws_reunions = wb.active
    ws_reunions.title = "Reunions"
    participant_type_labels = [
        (ParticipantReunion.TYPE_REPRESENTANT, 'Representants des usagers'),
        (ParticipantReunion.TYPE_ABONNE, 'Abonnes'),
        (ParticipantReunion.TYPE_EMPLOYE, 'Employes'),
        (ParticipantReunion.TYPE_AUTORITE, 'Autorites'),
        (ParticipantReunion.TYPE_TIERS, 'Tiers'),
    ]
    headers_reunions = [
        'Type', 'Date', 'Lieu', 'Objet', 'Compte-rendu', 'Recommandations',
        'Organise par', 'Total presents',
        *[f'Presents - {label}' for _type, label in participant_type_labels],
        'Total attendus',
        *[f'Attendus - {label}' for _type, label in participant_type_labels],
        'Taux avis negatifs (%)',
        'Nb ouvrages geres', 'Depenses pointees (Ar)',
    ]
    write_header(ws_reunions, headers_reunions)

    # Cumuls pour la feuille "Recap par region" : cle = nom de la region (ou 'Sans region').
    recap_par_region = {}

    for row_num, reunion in enumerate(reunions, 2):
        nb_ouvrages = _nb_ouvrages_geres_zone(
            gestionnaire_ids,
            commune_id=str(reunion.commune_id) if reunion.commune_id else None,
            district_id=str(reunion.district_id) if reunion.district_id else None,
            region_id=str(reunion.region_id) if reunion.region_id else None,
            annee=reunion.date_reunion.year if reunion.date_reunion else None,
        )

        presents_par_type = {type_value: 0 for type_value, _label in participant_type_labels}
        for participant in reunion.participants.all():
            if participant.present and participant.type_participant in presents_par_type:
                presents_par_type[participant.type_participant] += 1
        total_presents = sum(presents_par_type.values())

        attendus_par_type = {type_value: 0 for type_value, _label in participant_type_labels}
        for attendu in reunion.attendus.all():
            if attendu.type_participant in attendus_par_type:
                attendus_par_type[attendu.type_participant] += attendu.nombre_attendu
        total_attendu = sum(attendus_par_type.values())

        # Meme logique que la page liste (reunions_list) : le taux d'avis negatifs s'appuie
        # sur les criteres individuels de chaque avis (4 booleens + 3 champs probleme), pas
        # sur un simple comptage d'avis "negatifs/positifs" (cf. AvisRepresentant.
        # nb_criteres_negatifs et NB_CRITERES_TOTAL).
        avis_reunion = list(reunion.avis_recueillis.all())
        nb_avis = len(avis_reunion)
        criteres_negatifs = sum(a.nb_criteres_negatifs for a in avis_reunion)
        criteres_total = nb_avis * AvisRepresentant.NB_CRITERES_TOTAL
        pct_avis_negatifs = round((criteres_negatifs / criteres_total) * 100) if criteres_total else None

        values = [
            reunion.get_type_reunion_display() if reunion.type_reunion else '',
            reunion.date_reunion,
            reunion.lieu_affiche + (f" - {reunion.lieu}" if reunion.lieu else ''),
            reunion.objet,
            reunion.compte_rendu or '',
            reunion.recommandations or '',
            (reunion.organise_par.surnom or reunion.organise_par.nom) if reunion.organise_par else '',
            total_presents,
            *[presents_par_type[type_value] for type_value, _label in participant_type_labels],
            total_attendu,
            *[attendus_par_type[type_value] for type_value, _label in participant_type_labels],
            pct_avis_negatifs if pct_avis_negatifs is not None else '',
            nb_ouvrages,
            float(reunion.montant_depense_total),
        ]
        for col_num, value in enumerate(values, 1):
            cell = ws_reunions.cell(row=row_num, column=col_num, value=value)
            cell.border = border
            cell.alignment = wrap_top
            if col_num == 2:
                cell.number_format = 'DD/MM/YYYY'

        # Alimente le recap par region (une reunion peut etre rattachee a une region directement,
        # ou via son district, ou via son district/commune)
        if reunion.region_id and reunion.region:
            region_nom = reunion.region.nom_region
        elif reunion.district_id and reunion.district and reunion.district.region:
            region_nom = reunion.district.region.nom_region
        elif reunion.commune_id and reunion.commune and reunion.commune.district and reunion.commune.district.region:
            region_nom = reunion.commune.district.region.nom_region
        else:
            region_nom = 'Sans region'

        cumul = recap_par_region.setdefault(region_nom, {
            'nb_reunions': 0, 'nb_participants': 0, 'nb_attendus': 0,
            'presents_par_type': {type_value: 0 for type_value, _label in participant_type_labels},
            'attendus_par_type': {type_value: 0 for type_value, _label in participant_type_labels},
            'criteres_negatifs': 0, 'criteres_total': 0,
        })
        cumul['nb_reunions'] += 1
        cumul['nb_participants'] += total_presents
        cumul['nb_attendus'] += total_attendu
        for type_value, _label in participant_type_labels:
            cumul['presents_par_type'][type_value] += presents_par_type[type_value]
            cumul['attendus_par_type'][type_value] += attendus_par_type[type_value]
        cumul['criteres_negatifs'] += criteres_negatifs
        cumul['criteres_total'] += criteres_total

    column_widths_reunions = {
        1: 20, 2: 12, 3: 22, 4: 30, 5: 45, 6: 40, 7: 20,
        **{col_num: 20 for col_num in range(8, 20)},
        20: 18, 21: 16, 22: 18,
    }
    for col_num, width in column_widths_reunions.items():
        ws_reunions.column_dimensions[get_column_letter(col_num)].width = width

    # --- Feuille 2 : Avis ---
    ws_avis = wb.create_sheet("Avis")
    headers_avis = [
        'Date reunion', 'Objet reunion', 'Date avis', 'Representant', 'Village',
        'Quantite suffisante', 'Qualite satisfaisante', 'Delai reparation correct', 'Bonne relation gestionnaire',
        'Score satisfaction (/4)', 'Probleme technique', 'Probleme de recouvrement',
        "Probleme sur les regles d'utilisation", 'Remarque', 'Recueilli par',
    ]
    write_header(ws_avis, headers_avis)

    row_num = 2
    oui_non = lambda v: 'Oui' if v else 'Non'
    for reunion in reunions:
        for avis in reunion.avis_recueillis.all():
            values = [
                reunion.date_reunion,
                reunion.objet,
                avis.date_avis,
                avis.representant.nom if avis.representant else '',
                avis.representant.village.nom_village if avis.representant and avis.representant.village else '',
                oui_non(avis.quantite_ok),
                oui_non(avis.qualite_ok),
                oui_non(avis.delai_reparation_ok),
                oui_non(avis.relation_gestionnaire_ok),
                avis.score_satisfaction,
                avis.probleme_technique or '',
                avis.probleme_recouvrement or '',
                avis.probleme_regles_utilisation or '',
                avis.remarque or '',
                (avis.employe.surnom or avis.employe.nom) if avis.employe else '',
            ]
            for col_num, value in enumerate(values, 1):
                cell = ws_avis.cell(row=row_num, column=col_num, value=value)
                cell.border = border
                cell.alignment = wrap_top
                if col_num in (1, 3):
                    cell.number_format = 'DD/MM/YYYY'
            row_num += 1

    column_widths_avis = {
        1: 12, 2: 26, 3: 12, 4: 22, 5: 18,
        6: 16, 7: 16, 8: 16, 9: 16,
        10: 12, 11: 30, 12: 30, 13: 30, 14: 30, 15: 18,
    }
    for col_num, width in column_widths_avis.items():
        ws_avis.column_dimensions[get_column_letter(col_num)].width = width

    # --- Feuille 3 : Recap par region ---
    ws_recap = wb.create_sheet("Recap par region")
    headers_recap = [
        'Region', 'Nb reunions', 'Total presents',
        *[f'Presents - {label}' for _type, label in participant_type_labels],
        'Total attendus',
        *[f'Attendus - {label}' for _type, label in participant_type_labels],
        'Taux avis negatifs (%)',
    ]
    write_header(ws_recap, headers_recap)

    row_num = 2
    for region_nom in sorted(recap_par_region.keys()):
        cumul = recap_par_region[region_nom]
        pct_negatifs = (
            round((cumul['criteres_negatifs'] / cumul['criteres_total']) * 100)
            if cumul['criteres_total'] else ''
        )
        values = [
            region_nom,
            cumul['nb_reunions'],
            cumul['nb_participants'],
            *[
                cumul['presents_par_type'][type_value]
                for type_value, _label in participant_type_labels
            ],
            cumul['nb_attendus'],
            *[
                cumul['attendus_par_type'][type_value]
                for type_value, _label in participant_type_labels
            ],
            pct_negatifs,
        ]
        for col_num, value in enumerate(values, 1):
            cell = ws_recap.cell(row=row_num, column=col_num, value=value)
            cell.border = border
            cell.alignment = wrap_top
        row_num += 1

    # Ligne total general
    total_reunions = sum(c['nb_reunions'] for c in recap_par_region.values())
    total_participants = sum(c['nb_participants'] for c in recap_par_region.values())
    total_attendus = sum(c['nb_attendus'] for c in recap_par_region.values())
    total_presents_par_type = {
        type_value: sum(c['presents_par_type'][type_value] for c in recap_par_region.values())
        for type_value, _label in participant_type_labels
    }
    total_attendus_par_type = {
        type_value: sum(c['attendus_par_type'][type_value] for c in recap_par_region.values())
        for type_value, _label in participant_type_labels
    }
    total_criteres_negatifs = sum(c['criteres_negatifs'] for c in recap_par_region.values())
    total_criteres = sum(c['criteres_total'] for c in recap_par_region.values())
    total_pct_negatifs = round((total_criteres_negatifs / total_criteres) * 100) if total_criteres else ''

    total_values = [
        'TOTAL GENERAL', total_reunions, total_participants,
        *[total_presents_par_type[type_value] for type_value, _label in participant_type_labels],
        total_attendus,
        *[total_attendus_par_type[type_value] for type_value, _label in participant_type_labels],
        total_pct_negatifs,
    ]
    for col_num, value in enumerate(total_values, 1):
        cell = ws_recap.cell(row=row_num, column=col_num, value=value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        cell.border = border
        cell.alignment = wrap_top

    column_widths_recap = {
        1: 24, 2: 14,
        **{col_num: 20 for col_num in range(3, len(headers_recap) + 1)},
    }
    for col_num, width in column_widths_recap.items():
        ws_recap.column_dimensions[get_column_letter(col_num)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=reunions_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    return response


# ---------------------------------------------------------------------------
# Participants
# ---------------------------------------------------------------------------

@login_required
@require_http_methods(["POST"])
def participant_create(request, reunion_id):
    """
    Cree un participant. Pour securiser la saisie et eviter les doublons/erreurs de nom,
    les types representant/abonne/employe/autorite DOIVENT etre rattaches a une fiche
    existante (via son id, choisi dans une liste de recherche cote frontend) : le texte
    libre n'est autorise que pour le type "tiers" (personne sans fiche dans le systeme).
    """
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        data = json.loads(request.body)

        type_participant = (data.get('type_participant') or '').strip()
        if type_participant not in dict(ParticipantReunion.TYPE_CHOICES):
            return JsonResponse({'success': False, 'error': 'Type de participant invalide.'}, status=400)

        participant = ParticipantReunion(
            reunion=reunion,
            type_participant=type_participant,
            present=data.get('present', True),
            fonction_libre=(data.get('fonction_libre') or '').strip() or None,
            contact_libre=(data.get('contact_libre') or '').strip() or None,
        )

        if type_participant == ParticipantReunion.TYPE_REPRESENTANT:
            participant.representant = RepresentantUsager.objects.filter(id=data.get('representant_id')).first()
            if not participant.representant:
                return JsonResponse({'success': False, 'error': 'Selectionnez un representant existant dans la liste.'}, status=400)
        elif type_participant == ParticipantReunion.TYPE_ABONNE:
            from apps.billing.models import PointEau
            participant.point_eau = PointEau.objects.filter(id=data.get('point_eau_id')).first()
            if not participant.point_eau:
                return JsonResponse({'success': False, 'error': 'Selectionnez un abonne (point d\'eau) existant dans la liste.'}, status=400)
        elif type_participant == ParticipantReunion.TYPE_EMPLOYE:
            participant.employe = Employe.objects.filter(id=data.get('employe_id')).first()
            if not participant.employe:
                return JsonResponse({'success': False, 'error': 'Selectionnez un employe existant dans la liste.'}, status=400)
        elif type_participant == ParticipantReunion.TYPE_AUTORITE:
            participant.autorite = Autorite.objects.filter(id=data.get('autorite_id')).first()
            if not participant.autorite:
                return JsonResponse({'success': False, 'error': 'Selectionnez une autorite existante dans la liste.'}, status=400)
        elif type_participant == ParticipantReunion.TYPE_TIERS:
            participant.tiers = TiersParticipant.objects.filter(id=data.get('tiers_id')).first()
            if not participant.tiers:
                return JsonResponse({'success': False, 'error': 'Selectionnez un tiers existant dans la liste (ou creez-le).'}, status=400)

        participant.save()
        return JsonResponse({'success': True, 'participant_id': str(participant.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def participant_search_json(request):
    """
    Recherche dans le referentiel (representant/employe/autorite/abonne) pour peupler
    la liste de selection lors de l'ajout d'un participant a une reunion (saisie
    securisee : on choisit une fiche existante plutot que de taper un nom libre).
    """
    type_participant = (request.GET.get('type') or '').strip()
    search = (request.GET.get('search') or '').strip()
    commune_id = (request.GET.get('commune_id') or '').strip()

    results = []

    if type_participant == ParticipantReunion.TYPE_REPRESENTANT:
        qs = RepresentantUsager.objects.filter(actif=True).select_related(
            'village__fokontany__commune'
        )
        if commune_id:
            qs = qs.filter(village__fokontany__commune_id=commune_id)
        if search:
            qs = qs.filter(Q(nom__icontains=search) | Q(village__nom_village__icontains=search))
        for r in qs.order_by('nom')[:30]:
            village_nom = r.village.nom_village if r.village else ''
            commune_nom = (
                r.village.fokontany.commune.nom_commune
                if r.village and r.village.fokontany_id
                else ''
            )
            results.append({
                'id': str(r.id),
                'label': r.nom or 'Sans nom',
                'detail': ' - '.join(value for value in [village_nom, commune_nom] if value),
            })

    elif type_participant == ParticipantReunion.TYPE_EMPLOYE:
        qs = Employe.objects.filter(date_debauchage__isnull=True).select_related(
            'poste', 'base'
        )
        if search:
            qs = qs.filter(Q(nom__icontains=search) | Q(surnom__icontains=search) | Q(prenom__icontains=search))
        for e in qs.order_by('surnom', 'nom')[:30]:
            poste_nom = e.poste.nom_poste if e.poste else ''
            base_nom = e.base.nom_base if e.base else ''
            results.append({
                'id': str(e.id),
                'label': e.nom_affichage,
                'detail': ' - '.join(value for value in [poste_nom, base_nom] if value),
            })

    elif type_participant == ParticipantReunion.TYPE_AUTORITE:
        qs = Autorite.objects.filter(actif=True).select_related(
            'region', 'district', 'commune', 'fokontany'
        )
        if search:
            qs = qs.filter(Q(nom__icontains=search) | Q(fonction__icontains=search))
        # La commune de la reunion sert a PRIORISER les autorites de cette commune (les plus
        # pertinentes en premier), pas a exclure les autres : une autorite de niveau
        # district/region, ou une fiche pas encore rattachee a une zone precise (commune vide),
        # doit rester choisissable (cf. cas signale : autorite creee sans commune/district).
        if commune_id:
            qs = qs.annotate(
                _priorite_commune=Case(
                    When(commune_id=commune_id, then=Value(0)),
                    default=Value(1),
                    output_field=IntegerField(),
                )
            ).order_by('_priorite_commune', 'nom')
        else:
            qs = qs.order_by('nom')
        for a in qs[:30]:
            if a.fokontany_id:
                rattachement = f"Fokontany {a.fokontany.nom_fokontany}"
            elif a.commune_id:
                rattachement = f"Commune {a.commune.nom_commune}"
            elif a.district_id:
                rattachement = f"District {a.district.nom_district}"
            elif a.region_id:
                rattachement = f"Region {a.region.nom_region}"
            else:
                rattachement = 'Zone non renseignee'
            type_autorite = a.get_type_autorite_display() or 'Type non renseigne'
            results.append({
                'id': str(a.id),
                'label': a.nom,
                'detail': f"{type_autorite} - {rattachement}",
            })

    elif type_participant == ParticipantReunion.TYPE_ABONNE:
        from apps.billing.models import PointEau
        if search:
            qs = PointEau.objects.filter(Q(nomabonne__icontains=search) | Q(idxls__icontains=search))
            if commune_id:
                qs = qs.filter(parcelle__village__fokontany__commune_id=commune_id)
        else:
            qs = PointEau.objects.none()  # Evite de charger tous les points d'eau sans recherche
        for p in qs.order_by('nomabonne')[:30]:
            results.append({
                'id': str(p.id),
                'label': p.forfaitaire_display_label,
                'detail': '',
            })

    elif type_participant == ParticipantReunion.TYPE_TIERS:
        qs = TiersParticipant.objects.filter(actif=True)
        if search:
            qs = qs.filter(Q(nom__icontains=search) | Q(fonction__icontains=search))
        for t in qs.order_by('nom')[:30]:
            results.append({
                'id': str(t.id),
                'label': t.nom,
                'detail': t.fonction or '',
            })

    return JsonResponse({'success': True, 'results': results})


def _avis_non_rattaches_qs(request, exclude_reunion_id=None):
    """Queryset de base des avis non rattaches a une reunion, filtre par perimetre utilisateur."""
    from .views import _filter_avis_queryset_for_user

    qs = AvisRepresentant.objects.filter(reunion__isnull=True).select_related(
        'representant',
        'representant__village',
        'representant__village__fokontany',
        'representant__village__fokontany__commune',
    )
    return _filter_avis_queryset_for_user(qs, request.user)


@login_required
def avis_unlinked_search_json(request):
    """
    Recherche parmi les avis de representants qui ne sont pas encore rattaches a une
    reunion (reunion IS NULL), pour permettre de lier un avis deja saisi (ex: via import
    Excel ou saisi hors reunion) a une reunion existante plutot que d'en recreer un.
    """
    search = (request.GET.get('search') or '').strip()
    qs = _avis_non_rattaches_qs(request)

    if search:
        qs = qs.filter(
            Q(representant__nom__icontains=search) |
            Q(representant__village__nom_village__icontains=search)
        )

    results = []
    for a in qs.order_by('-date_avis')[:30]:
        village_nom = a.representant.village.nom_village if a.representant and a.representant.village else ''
        results.append({
            'id': str(a.id),
            'label': f"{a.representant.nom if a.representant else 'Sans nom'} - {village_nom} ({a.date_avis.strftime('%d/%m/%Y') if a.date_avis else '?'})",
        })

    return JsonResponse({'success': True, 'results': results})


@login_required
def avis_candidats_reunion_json(request, reunion_id):
    """
    Liste les avis non rattaches qui partagent la meme date et la meme commune que la
    reunion (meme critere que la commande rattacher_avis_reunions), pour proposer un
    rattachement rapide en cartes plutot qu'une recherche manuelle avis par avis.
    """
    reunion = get_object_or_404(Reunion, id=reunion_id)

    if not reunion.commune_id:
        return JsonResponse({'success': True, 'results': [], 'message': "Cette reunion n'a pas de commune renseignee."})

    qs = _avis_non_rattaches_qs(request).filter(
        date_avis=reunion.date_reunion,
        representant__village__fokontany__commune_id=reunion.commune_id,
    )

    results = []
    for a in qs.order_by('representant__nom'):
        results.append({
            'id': str(a.id),
            'representant': a.representant.nom if a.representant else 'Sans nom',
            'village': a.representant.village.nom_village if a.representant and a.representant.village else '',
            'date_avis': a.date_avis.strftime('%d/%m/%Y') if a.date_avis else '',
            'score_satisfaction': a.score_satisfaction,
            'pourcentage_satisfaction': a.pourcentage_satisfaction,
            'est_negatif': a.est_negatif,
        })

    return JsonResponse({'success': True, 'results': results})


@login_required
@require_http_methods(["POST"])
def avis_link_reunion(request, avis_id):
    """Rattache un avis existant (non lie) a une reunion, sans toucher a son contenu."""
    from .views import _filter_avis_queryset_for_user

    avis = get_object_or_404(
        _filter_avis_queryset_for_user(AvisRepresentant.objects.all(), request.user),
        id=avis_id,
    )
    reunion_id = request.POST.get('reunion_id')
    if not reunion_id:
        return JsonResponse({'success': False, 'message': 'Reunion manquante'})

    reunion = get_object_or_404(Reunion, id=reunion_id)

    if avis.reunion_id and avis.reunion_id != reunion.id:
        return JsonResponse({
            'success': False,
            'message': "Cet avis est deja rattache a une autre reunion.",
        })

    avis.reunion = reunion
    avis.save(update_fields=['reunion'])

    return JsonResponse({'success': True, 'message': 'Avis rattache a la reunion'})


@login_required
@require_http_methods(["POST"])
def avis_link_reunion_batch(request, reunion_id):
    """
    Rattache en une fois plusieurs avis existants (non lies) a une reunion. Utilise par le
    rattachement rapide en cartes (avis_id[] fournis en POST, cf. avis_candidats_reunion_json).
    """
    from .views import _filter_avis_queryset_for_user

    reunion = get_object_or_404(Reunion, id=reunion_id)
    avis_ids = request.POST.getlist('avis_id')
    if not avis_ids:
        return JsonResponse({'success': False, 'message': 'Aucun avis selectionne'})

    qs = _filter_avis_queryset_for_user(AvisRepresentant.objects.all(), request.user).filter(
        id__in=avis_ids,
        reunion__isnull=True,
    )
    nb = qs.update(reunion=reunion)

    return JsonResponse({
        'success': True,
        'message': f"{nb} avis rattache(s) a la reunion" if nb else "Aucun avis rattache (deja lies entre-temps ?)",
        'nb_rattaches': nb,
    })


@login_required
@require_http_methods(["POST"])
def tiers_create(request):
    """Cree une fiche Tiers reutilisable (referentiel), pour eviter de retaper un nom libre a chaque reunion."""
    try:
        data = json.loads(request.body)
        nom = (data.get('nom') or '').strip()
        if not nom:
            return JsonResponse({'success': False, 'error': 'Le nom est obligatoire.'}, status=400)

        tiers = TiersParticipant.objects.create(
            nom=nom,
            fonction=(data.get('fonction') or '').strip() or None,
            contact=(data.get('contact') or '').strip() or None,
        )
        return JsonResponse({'success': True, 'tiers_id': str(tiers.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def participant_delete(request, participant_id):
    try:
        participant = get_object_or_404(ParticipantReunion, id=participant_id)
        participant.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ---------------------------------------------------------------------------
# Decisions
# ---------------------------------------------------------------------------

@login_required
@require_http_methods(["POST"])
def decision_create(request, reunion_id):
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        data = json.loads(request.body)

        description = (data.get('description') or '').strip()
        if not description:
            return JsonResponse({'success': False, 'error': 'La description de la decision est obligatoire.'}, status=400)

        responsable = None
        if data.get('responsable_id'):
            responsable = Employe.objects.filter(id=data['responsable_id']).first()

        decision = DecisionReunion.objects.create(
            reunion=reunion,
            description=description,
            responsable=responsable,
            date_echeance=data.get('date_echeance') or None,
            statut=data.get('statut') or DecisionReunion.STATUT_A_FAIRE,
        )
        return JsonResponse({'success': True, 'decision_id': str(decision.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def decision_update_statut(request, decision_id):
    """Met a jour le statut de suivi d'une decision (utilise depuis le tableau de suivi transversal)."""
    try:
        decision = get_object_or_404(DecisionReunion, id=decision_id)
        data = json.loads(request.body)

        statut = (data.get('statut') or '').strip()
        if statut not in dict(DecisionReunion.STATUT_CHOICES):
            return JsonResponse({'success': False, 'error': 'Statut invalide.'}, status=400)

        decision.statut = statut
        decision.commentaire_suivi = (data.get('commentaire_suivi') or decision.commentaire_suivi or '').strip() or None
        if statut == DecisionReunion.STATUT_REALISEE and not decision.date_realisation:
            decision.date_realisation = timezone.localdate()
        decision.save()

        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def decision_delete(request, decision_id):
    try:
        decision = get_object_or_404(DecisionReunion, id=decision_id)
        decision.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def decisions_suivi_dashboard(request):
    """Vue transversale de toutes les decisions a suivre, tous commune/reunion confondus."""
    statut_filter = request.GET.get('statut', '').strip()
    commune_id = request.GET.get('commune', '').strip()

    decisions = DecisionReunion.objects.select_related('reunion', 'reunion__commune', 'responsable')

    if statut_filter:
        decisions = decisions.filter(statut=statut_filter)
    else:
        decisions = decisions.exclude(statut=DecisionReunion.STATUT_REALISEE).exclude(statut=DecisionReunion.STATUT_ABANDONNEE)

    if commune_id:
        decisions = decisions.filter(reunion__commune_id=commune_id)

    decisions = decisions.order_by('date_echeance', '-created_at')

    communes = Commune.objects.order_by('nom_commune')

    context = {
        'decisions': decisions,
        'communes': communes,
        'filter_statut': statut_filter,
        'filter_commune': commune_id,
        'statut_choices': DecisionReunion.STATUT_CHOICES,
    }
    return render(request, 'usagers/decisions_suivi.html', context)


# ---------------------------------------------------------------------------
# Pieces jointes
# ---------------------------------------------------------------------------

@login_required
@require_http_methods(["POST"])
def piece_jointe_create(request, reunion_id):
    try:
        reunion = get_object_or_404(Reunion, id=reunion_id)
        fichier = request.FILES.get('fichier')
        if not fichier:
            return JsonResponse({'success': False, 'error': 'Aucun fichier fourni.'}, status=400)

        type_piece = request.POST.get('type_piece') or PieceJointeReunion.TYPE_AUTRE
        if type_piece not in dict(PieceJointeReunion.TYPE_CHOICES):
            type_piece = PieceJointeReunion.TYPE_AUTRE

        piece = PieceJointeReunion.objects.create(
            reunion=reunion,
            type_piece=type_piece,
            fichier=fichier,
            legende=(request.POST.get('legende') or '').strip() or None,
            uploaded_by=_current_employe(request),
        )
        return JsonResponse({
            'success': True,
            'piece_jointe': {
                'id': str(piece.id),
                'type_piece': piece.type_piece,
                'type_display': piece.get_type_piece_display(),
                'url': piece.fichier.url,
                'legende': piece.legende or '',
            }
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def piece_jointe_delete(request, piece_id):
    try:
        piece = get_object_or_404(PieceJointeReunion, id=piece_id)
        piece.fichier.delete(save=False)
        piece.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ---------------------------------------------------------------------------
# Autorites (referentiel)
# ---------------------------------------------------------------------------

@login_required
def autorites_list_json(request):
    commune_id = request.GET.get('commune', '').strip()
    autorites = Autorite.objects.filter(actif=True)
    if commune_id:
        autorites = autorites.filter(commune_id=commune_id)
    autorites = autorites.order_by('nom')

    return JsonResponse({
        'success': True,
        'autorites': [
            {'id': str(a.id), 'nom': a.nom, 'fonction': a.fonction or '', 'commune_id': str(a.commune_id) if a.commune_id else ''}
            for a in autorites
        ]
    })


@login_required
@require_http_methods(["POST"])
def autorite_create(request):
    """
    Creation rapide d'une autorite depuis le modal d'ajout de participant a une reunion.
    La zone geographique (fokontany/commune/district/region) est pre-remplie cote front avec
    le lieu de la reunion en cours (cf. openCreateFicheModal dans reunions_list.html), mais
    reste modifiable : un chef fokontany peut ainsi etre rattache precisement a son fokontany
    plutot qu'a la commune entiere de la reunion.
    """
    try:
        data = json.loads(request.body)
        nom = (data.get('nom') or '').strip()
        if not nom:
            return JsonResponse({'success': False, 'error': 'Le nom est obligatoire.'}, status=400)

        # Type structure (Maire, Chef Fokontany, ...) : evite de melanger cette info avec le
        # champ "fonction" en texte libre, source d'incoherences (fautes de frappe, valeurs
        # non filtrables) constatees quand seule la fonction libre etait saisie ici.
        type_autorite = (data.get('type_autorite') or '').strip()
        if type_autorite and type_autorite not in dict(Autorite.get_type_choices()):
            return JsonResponse({'success': False, 'error': "Type d'autorite invalide."}, status=400)

        fokontany = None
        if data.get('fokontany_id'):
            fokontany = Fokontany.objects.filter(id=data['fokontany_id']).first()

        commune = None
        if data.get('commune_id'):
            commune = Commune.objects.filter(id=data['commune_id']).first()

        district = None
        if data.get('district_id'):
            district = District.objects.filter(id=data['district_id']).first()

        region = None
        if data.get('region_id'):
            region = Region.objects.filter(id=data['region_id']).first()

        autorite = Autorite.objects.create(
            nom=nom,
            type_autorite=type_autorite,
            fonction=(data.get('fonction') or '').strip() or None,
            fokontany=fokontany,
            commune=commune,
            district=district,
            region=region,
            contact=(data.get('contact') or '').strip() or None,
        )
        return JsonResponse({'success': True, 'autorite_id': str(autorite.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def representant_create_quick(request):
    """
    Creation rapide d'un representant des usagers depuis le modal d'ajout de participant a
    une reunion (meme principe que autorite_create). Le village est pre-rempli cote front
    avec le lieu de la reunion en cours si un seul village y correspond, mais reste a choisir
    explicitement par l'utilisateur (contrairement a la commune, le village n'est pas toujours
    deductible du lieu de la reunion).
    """
    try:
        data = json.loads(request.body)
        nom = (data.get('nom') or '').strip()
        if not nom:
            return JsonResponse({'success': False, 'error': 'Le nom est obligatoire.'}, status=400)

        village_id = (data.get('village_id') or '').strip()
        if not village_id:
            return JsonResponse({'success': False, 'error': 'Le village est obligatoire.'}, status=400)

        village = Village.objects.filter(id=village_id).first()
        if not village:
            return JsonResponse({'success': False, 'error': 'Village introuvable.'}, status=400)

        representant = RepresentantUsager.objects.create(
            village=village,
            nom=nom,
            contact=(data.get('contact') or '').strip() or None,
            actif=True,
        )
        return JsonResponse({'success': True, 'representant_id': str(representant.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def employe_create_quick(request):
    """
    Creation rapide (minimale) d'un employe depuis le modal d'ajout de participant a une
    reunion, pour les cas ou la personne (ex: nouvel agent terrain) n'a pas encore de fiche
    RH complete. Seuls nom et base sont obligatoires en base (cf. Employe.base non-nullable) ;
    le poste reste facultatif. Cette fiche minimale est ensuite completable via le module RH
    complet (contrat, salaire, etc.) : creer via ce modal ne remplace pas le workflow RH,
    il debloque juste l'ajout du participant a la reunion sans attendre.
    """
    try:
        data = json.loads(request.body)
        nom = (data.get('nom') or '').strip()
        if not nom:
            return JsonResponse({'success': False, 'error': 'Le nom est obligatoire.'}, status=400)

        base_id = (data.get('base_id') or '').strip()
        if not base_id:
            return JsonResponse({'success': False, 'error': 'La base est obligatoire.'}, status=400)

        # Restreint aux bases de l'organisation de l'utilisateur connecte : evite de rattacher
        # par erreur un nouvel employe a la base d'une autre organisation.
        organisation = get_user_organisation(request.user)
        bases_qs = Base.objects.all()
        if organisation is not None:
            bases_qs = bases_qs.filter(organisation=organisation)
        base = bases_qs.filter(id=base_id).first()
        if not base:
            return JsonResponse({'success': False, 'error': 'Base introuvable.'}, status=400)

        poste = None
        if data.get('poste_id'):
            poste_qs = Poste.objects.all()
            if organisation is not None:
                poste_qs = poste_qs.filter(organisation=organisation)
            poste = poste_qs.filter(id=data['poste_id']).first()

        prenom = (data.get('prenom') or '').strip()
        surnom = (data.get('surnom') or '').strip()
        nom_employe = f"{prenom} {nom}".strip() if prenom else nom

        employe = Employe.objects.create(
            nom_employe=nom_employe,
            nom=nom,
            prenom=prenom or None,
            surnom=surnom or None,
            base=base,
            poste=poste,
            telephone=(data.get('telephone') or '').strip() or None,
        )
        return JsonResponse({'success': True, 'employe_id': str(employe.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def employe_bases_postes_json(request):
    """
    Liste des bases et postes de l'organisation de l'utilisateur connecte, pour peupler les
    selects du formulaire de creation rapide d'employe (cf. employe_create_quick).
    """
    organisation = get_user_organisation(request.user)
    bases_qs = Base.objects.all()
    postes_qs = Poste.objects.all()
    if organisation is not None:
        bases_qs = bases_qs.filter(organisation=organisation)
        postes_qs = postes_qs.filter(organisation=organisation)

    return JsonResponse({
        'success': True,
        'bases': [{'id': str(b.id), 'nom': b.nom_base} for b in bases_qs.order_by('nom_base')],
        'postes': [{'id': str(p.id), 'nom': p.nom_poste} for p in postes_qs.order_by('ordre_affichage', 'nom_poste')],
    })


# ---------------------------------------------------------------------------
# Pointage DiameX (comptes 608004, 625300, 625302) <-> Reunion
# ---------------------------------------------------------------------------

COMPTES_DEPENSES_REUNIONS = ['608004', '625300', '625302']


def _montant_depense_ligne(entry):
    """Les comptes de reunion sont des comptes de charges : le montant utile est en sortie (Out)."""
    if entry.montant_out:
        return float(entry.montant_out)
    return float(entry.montant_in or 0)


@login_required
def pointage_diamex_list(request):
    """
    Page de pointage : liste les lignes BaseDiameX des comptes de reunion non encore
    pointees, avec suggestion automatique de lieu (Commune/District/Region) extraite
    du libelle, pour faciliter le rattachement manuel a une Reunion existante.

    Toujours limitee a une seule annee a la fois (annee courante par defaut) pour
    eviter de charger tout l'historique des lignes DiameX.

    CRITICAL: date_m ("Date mouvement") est souvent calee au 1er du mois comptable et
    n'est PAS fiable pour retrouver la vraie date de la depense. date_rec ("Date recu")
    est la date reelle de la piece/transaction : c'est elle qui doit servir de reference
    pour le filtre annee, le tri, l'affichage et le pre-remplissage de la reunion rapide.
    """
    annee_param = request.GET.get('annee', '').strip()
    montre_pointees = request.GET.get('montre_pointees', '').strip() == '1'
    montre_ignorees = request.GET.get('montre_ignorees', '').strip() == '1'

    try:
        annee = int(annee_param) if annee_param else timezone.localdate().year
    except ValueError:
        annee = timezone.localdate().year

    from django.db.models.functions import Coalesce

    lignes_qs = BaseDiameX.objects.filter(
        compte_code__in=COMPTES_DEPENSES_REUNIONS,
    ).annotate(
        date_reference=Coalesce('date_rec', 'date_m')
    ).filter(
        date_reference__year=annee,
    ).order_by('-date_reference')[:500]

    # Cle (numero_diamex, base_diamex) -> infos du pointage existant, pour permettre de
    # reassigner une ligne pointee par erreur vers la bonne reunion sans repasser par un
    # depointage puis un re-pointage manuels separes (cf. pointage_diamex_reassigner).
    deja_pointes = {
        (numero, base): {'pointage_id': str(pid), 'reunion_id': str(rid), 'reunion_objet': objet, 'reunion_date': date.isoformat() if date else None}
        for pid, numero, base, rid, objet, date in PointageReunionDiameX.objects.select_related('reunion').values_list(
            'id', 'numero_diamex', 'base_diamex', 'reunion_id', 'reunion__objet', 'reunion__date_reunion'
        )
    }
    deja_ignorees = set(
        LigneDiameXIgnoreeReunion.objects.values_list('numero_diamex', 'base_diamex')
    )

    lignes = []
    for entry in lignes_qs:
        numero_diamex = str(entry.id_access or '').strip()
        base_diamex = (entry.prog or '').strip()
        pointage_info = deja_pointes.get((numero_diamex, base_diamex)) if numero_diamex and base_diamex else None
        est_pointe = pointage_info is not None
        est_ignoree = bool(numero_diamex and base_diamex and (numero_diamex, base_diamex) in deja_ignorees)

        if est_pointe and not montre_pointees:
            continue
        if est_ignoree and not montre_ignorees:
            continue

        lieu = extraire_lieu_reunion(entry.description or '')

        # Reunions candidates : meme annee que le filtre courant, meme lieu extrait
        # (commune/district/region correspondant). Sans le filtre annee, une ligne de
        # 2026 pouvait proposer des reunions de 2024/2025 dans la liste de rattachement.
        reunions_candidates_qs = Reunion.objects.none()
        if lieu['commune_id']:
            reunions_candidates_qs = Reunion.objects.filter(commune_id=lieu['commune_id'])
        elif lieu['district_id']:
            reunions_candidates_qs = Reunion.objects.filter(district_id=lieu['district_id'])
        elif lieu['region_id']:
            reunions_candidates_qs = Reunion.objects.filter(region_id=lieu['region_id'])
        reunions_candidates_qs = reunions_candidates_qs.filter(date_reunion__year=annee)

        reunions_candidates = [
            {'id': str(r.id), 'objet': r.objet, 'date_reunion': r.date_reunion.isoformat()}
            for r in reunions_candidates_qs.order_by('-date_reunion')[:10]
        ]

        date_reference = entry.date_rec or entry.date_m

        lignes.append({
            'entry_id': entry.id,
            'numero_diamex': numero_diamex,
            'base_diamex': base_diamex,
            'date_m': entry.date_m.isoformat() if entry.date_m else None,
            'date_rec': entry.date_rec.isoformat() if entry.date_rec else None,
            'date_reference': date_reference.isoformat() if date_reference else None,
            'description': entry.description or '',
            'compte_code': entry.compte_code or '',
            'montant': _montant_depense_ligne(entry),
            'est_pointe': est_pointe,
            'est_ignoree': est_ignoree,
            'lieu_suggere': lieu,
            'reunions_candidates': reunions_candidates,
            'sans_reunion': not est_pointe and not est_ignoree and len(reunions_candidates) == 0,
            'pointage_id': pointage_info['pointage_id'] if pointage_info else None,
            'reunion_actuelle': {
                'id': pointage_info['reunion_id'],
                'objet': pointage_info['reunion_objet'],
                'date_reunion': pointage_info['reunion_date'],
            } if pointage_info else None,
        })

    annee_courante = timezone.localdate().year
    annees_disponibles = list(range(annee_courante, annee_courante - 6, -1))
    if annee not in annees_disponibles:
        annees_disponibles = sorted(set(annees_disponibles + [annee]), reverse=True)

    # Toutes les reunions de l'annee filtree, pour peupler le select de pointage manuel.
    # Ne PAS se baser uniquement sur les "reunions_candidates" par ligne : une reunion deja
    # entierement pointee (plus aucune ligne candidate visible) doit rester choisissable,
    # car d'autres lignes (non detectees automatiquement) peuvent encore lui etre rattachees.
    toutes_reunions_annee = [
        {'id': str(r.id), 'objet': r.objet, 'date_reunion': r.date_reunion.isoformat()}
        for r in Reunion.objects.filter(date_reunion__year=annee).order_by('-date_reunion')
    ]

    context = {
        'lignes_json': json.dumps(lignes, default=str),
        'toutes_reunions_json': json.dumps(toutes_reunions_annee, default=str),
        'filter_annee': annee,
        'annees_disponibles': annees_disponibles,
        'montre_pointees': montre_pointees,
        'montre_ignorees': montre_ignorees,
        'nb_lignes': len(lignes),
        'regions': Region.objects.all().order_by('nom_region'),
        'districts': District.objects.select_related('region').order_by('nom_district'),
        'communes': Commune.objects.select_related('district').order_by('nom_commune'),
    }
    return render(request, 'usagers/pointage_reunions_diamex.html', context)


@login_required
@require_http_methods(["POST"])
def pointage_diamex_create(request):
    """Pointe une ou plusieurs lignes BaseDiameX sur une reunion existante."""
    try:
        data = json.loads(request.body)
        reunion_id = (data.get('reunion_id') or '').strip()
        entry_ids = data.get('entry_ids') or []

        if not reunion_id or not entry_ids:
            return JsonResponse({'success': False, 'error': 'Reunion et lignes a pointer sont obligatoires.'}, status=400)

        reunion = get_object_or_404(Reunion, id=reunion_id)
        employe = _current_employe(request)

        entries = BaseDiameX.objects.filter(id__in=entry_ids, compte_code__in=COMPTES_DEPENSES_REUNIONS)

        nb_pointes = 0
        erreurs = []
        with transaction.atomic():
            for entry in entries:
                numero_diamex = str(entry.id_access or '').strip()
                base_diamex = (entry.prog or '').strip()
                if not numero_diamex or not base_diamex:
                    erreurs.append(f"Ligne {entry.id} sans identifiant DiameX complet, ignoree.")
                    continue

                if PointageReunionDiameX.objects.filter(
                    numero_diamex=numero_diamex, base_diamex=base_diamex
                ).exists():
                    erreurs.append(f"Ligne {numero_diamex}/{base_diamex} deja pointee, ignoree.")
                    continue

                PointageReunionDiameX.objects.create(
                    reunion=reunion,
                    numero_diamex=numero_diamex,
                    base_diamex=base_diamex,
                    entry_id=entry.id,
                    date_transaction=entry.date_m,
                    description=entry.description or '',
                    compte_code=entry.compte_code or '',
                    montant=_montant_depense_ligne(entry),
                    pointe_par=employe,
                )
                nb_pointes += 1

        return JsonResponse({
            'success': True,
            'nb_pointes': nb_pointes,
            'erreurs': erreurs,
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def pointage_diamex_reassigner(request, pointage_id):
    """
    Change la reunion d'un pointage existant, pour corriger un rattachement fait par erreur
    sans devoir depointer puis re-pointer manuellement (cf. page pointage_reunions_diamex,
    lignes deja pointees). Garde le meme PointageReunionDiameX (donc le meme historique
    pointe_par/created_at) : seule la reunion cible change.
    """
    try:
        pointage = get_object_or_404(PointageReunionDiameX, id=pointage_id)
        data = json.loads(request.body)
        reunion_id = (data.get('reunion_id') or '').strip()
        if not reunion_id:
            return JsonResponse({'success': False, 'error': 'Nouvelle reunion obligatoire.'}, status=400)

        nouvelle_reunion = get_object_or_404(Reunion, id=reunion_id)
        if nouvelle_reunion.id == pointage.reunion_id:
            return JsonResponse({'success': False, 'error': 'Cette ligne est deja pointee sur cette reunion.'}, status=400)

        pointage.reunion = nouvelle_reunion
        pointage.save(update_fields=['reunion'])
        return JsonResponse({'success': True, 'reunion_objet': nouvelle_reunion.objet})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def pointage_diamex_delete(request, pointage_id):
    """Retire le pointage d'une ligne (la ligne redevient disponible pour un autre pointage)."""
    try:
        pointage = get_object_or_404(PointageReunionDiameX, id=pointage_id)
        pointage.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def pointage_diamex_ignorer(request):
    """
    Marque une ligne BaseDiameX comme ne correspondant a aucune reunion (ex: depense mal
    categorisee sur ce compte), pour la retirer de la liste de pointage sans creer de
    faux rattachement.
    """
    try:
        data = json.loads(request.body)
        entry_id = data.get('entry_id')
        motif = (data.get('motif') or '').strip() or None
        if not entry_id:
            return JsonResponse({'success': False, 'error': 'Ligne a ignorer obligatoire.'}, status=400)

        entry = get_object_or_404(BaseDiameX, id=entry_id, compte_code__in=COMPTES_DEPENSES_REUNIONS)
        numero_diamex = str(entry.id_access or '').strip()
        base_diamex = (entry.prog or '').strip()
        if not numero_diamex or not base_diamex:
            return JsonResponse({'success': False, 'error': 'Ligne sans identifiant DiameX complet.'}, status=400)

        LigneDiameXIgnoreeReunion.objects.update_or_create(
            numero_diamex=numero_diamex,
            base_diamex=base_diamex,
            defaults={
                'entry_id': entry.id,
                'motif': motif,
                'ignore_par': _current_employe(request),
            },
        )
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["POST"])
def pointage_diamex_reactiver(request):
    """Annule l'ignorance d'une ligne (la ligne redevient disponible pour le pointage)."""
    try:
        data = json.loads(request.body)
        entry_id = data.get('entry_id')
        if not entry_id:
            return JsonResponse({'success': False, 'error': 'Ligne obligatoire.'}, status=400)

        LigneDiameXIgnoreeReunion.objects.filter(entry_id=entry_id).delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)
