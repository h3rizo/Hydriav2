"""Export Excel du tableau de bord Paiements Forfaitaires : paiements prévus
par point d'eau et par mois, pour l'année choisie.

« Prévu » = montant des factures forfaitaires du mois (payées + non payées),
exactement la valeur des cellules mensuelles du tableau de bord
(mois_cells : paye_montant + non_paye_montant). Les mois à venir de l'année
sont inclus. Les points exportés sont ceux de la page après tous ses filtres
(sans pagination).
"""
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MOIS = ['Jan', 'Fév', 'Mar', 'Avr', 'Mai', 'Juin', 'Juil', 'Aoû', 'Sep', 'Oct', 'Nov', 'Déc']


def _montant(value):
    return Decimal(str(value or 0))


def export_paiements_prevus_excel(points_eau_list, annee):
    wb = Workbook()
    ws = wb.active
    ws.title = f'Prévus {annee}'

    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    total_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    coupe_fill = PatternFill(start_color='FDE9E7', end_color='FDE9E7', fill_type='solid')

    ws.cell(row=1, column=1, value=f"Paiements forfaitaires prévus par point d'eau et par mois — {annee}").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=(
        "Prévu = factures forfaitaires du mois (payées + non payées). "
        "Cases rosées : point d'eau coupé ce mois-là."
    )).font = Font(italic=True, color='666666')

    entetes_fixes = ['Commune', 'Ouvrage', 'IDXLS', 'Abonné', "Type de point d'eau",
                     'Tarif forfaitaire (Ar)', 'Statut']
    entetes = entetes_fixes + MOIS + ['Total prévu', 'Dont payé', 'Reste à payer']
    header_row = 4
    for col, label in enumerate(entetes, 1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    premier_mois_col = len(entetes_fixes) + 1
    totaux_mois = [Decimal('0')] * 12
    total_prevu_global = total_paye_global = Decimal('0')

    row_num = header_row + 1
    for point in points_eau_list:
        cells_par_mois = {cell['numero']: cell for cell in point.get('mois_cells', [])}
        valeurs = [
            point.get('commune') or '',
            point.get('ouvrage') or '',
            point.get('idxls') or '',
            point.get('nomabonne') or '',
            point.get('type_point_eau') or '',
            float(_montant(point.get('tarif_forfaitaire'))),
            point.get('statut_label') or '',
        ]
        for col, value in enumerate(valeurs, 1):
            ws.cell(row=row_num, column=col, value=value).border = border
        ws.cell(row=row_num, column=6).number_format = '#,##0'

        total_prevu = total_paye = Decimal('0')
        for index in range(12):
            cell_data = cells_par_mois.get(index + 1, {})
            paye = _montant(cell_data.get('paye_montant'))
            prevu = paye + _montant(cell_data.get('non_paye_montant'))
            total_prevu += prevu
            total_paye += paye
            totaux_mois[index] += prevu
            cell = ws.cell(row=row_num, column=premier_mois_col + index, value=float(prevu) if prevu else None)
            cell.number_format = '#,##0'
            cell.border = border
            if cell_data.get('coupe'):
                cell.fill = coupe_fill

        total_prevu_global += total_prevu
        total_paye_global += total_paye
        for offset, value in enumerate((total_prevu, total_paye, total_prevu - total_paye)):
            cell = ws.cell(row=row_num, column=premier_mois_col + 12 + offset, value=float(value))
            cell.number_format = '#,##0'
            cell.font = Font(bold=offset == 0)
            cell.border = border
        row_num += 1

    # Ligne de total
    ws.cell(row=row_num, column=1, value=f"Total ({len(points_eau_list)} points d'eau)")
    for index, value in enumerate(totaux_mois):
        ws.cell(row=row_num, column=premier_mois_col + index, value=float(value))
    for offset, value in enumerate((total_prevu_global, total_paye_global, total_prevu_global - total_paye_global)):
        ws.cell(row=row_num, column=premier_mois_col + 12 + offset, value=float(value))
    for col in range(1, len(entetes) + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = Font(bold=True)
        cell.fill = total_fill
        cell.border = border
        if col >= premier_mois_col:
            cell.number_format = '#,##0'

    largeurs = [18, 26, 8, 28, 16, 12, 16] + [11] * 12 + [13, 13, 13]
    for col, largeur in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(col)].width = largeur
    ws.freeze_panes = ws.cell(row=header_row + 1, column=premier_mois_col)
    if points_eau_list:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(entetes))}{row_num - 1}"

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"paiements_forfaitaires_prevus_{annee}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
