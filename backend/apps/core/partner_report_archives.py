from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Iterable, List, Optional

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.test import RequestFactory
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone

from apps.employees.permissions import is_visiteur
from apps.organization.models import Organisation
from apps.organization.scoping import get_user_organisation

from .models import PartnerQuarterlyReportArchive


class PartnerReportArchiveError(RuntimeError):
    pass


@dataclass(frozen=True)
class PartnerReportQuarterSpec:
    report_year: int
    report_quarter: int
    due_date: date

    @property
    def label(self):
        return f"T{self.report_quarter} {self.report_year}"


def get_partner_report_quarter_due_date(report_year: int, report_quarter: int) -> date:
    if report_quarter == 1:
        return date(report_year, 4, 15)
    if report_quarter == 2:
        return date(report_year, 7, 15)
    if report_quarter == 3:
        return date(report_year, 10, 15)
    if report_quarter == 4:
        return date(report_year + 1, 1, 15)
    raise ValueError(f"Trimestre invalide: {report_quarter}")


def get_due_partner_report_quarter_specs(today: Optional[date] = None) -> List[PartnerReportQuarterSpec]:
    today = today or timezone.localdate()
    specs = []
    for report_year in (today.year - 1, today.year):
        for quarter in (1, 2, 3, 4):
            due_date = get_partner_report_quarter_due_date(report_year, quarter)
            if due_date <= today:
                specs.append(
                    PartnerReportQuarterSpec(
                        report_year=report_year,
                        report_quarter=quarter,
                        due_date=due_date,
                    )
                )
    specs.sort(key=lambda item: (item.report_year, item.report_quarter))
    return specs


def get_latest_due_partner_report_quarter_spec(today: Optional[date] = None) -> Optional[PartnerReportQuarterSpec]:
    specs = get_due_partner_report_quarter_specs(today=today)
    return specs[-1] if specs else None


def get_partner_report_generation_user(organisation: Organisation):
    user_model = get_user_model()
    return (
        user_model.objects.filter(
            is_active=True,
            employe__base__organisation=organisation,
        )
        .select_related('employe__base__organisation')
        .order_by('id')
        .first()
    )


def _build_partner_report_request(user, query_params):
    request_factory = RequestFactory()
    request = request_factory.get(
        reverse('core:rapport_partenaires_trimestriel'),
        data=query_params,
    )
    request.user = user
    return request


def _collect_archive_scopes(report_user, report_year: int):
    from apps.core.views_rapports import _build_partner_report_filter_context

    root_organisation = get_user_organisation(report_user)
    if not root_organisation:
        return []

    request = _build_partner_report_request(
        report_user,
        {
            'annee': report_year,
            'organisation': str(root_organisation.id),
        },
    )
    filter_context = _build_partner_report_filter_context(request)

    scopes = [
        {
            'scope_level': PartnerQuarterlyReportArchive.SCOPE_ORGANISATION,
            'scope_key': 'organisation:all',
            'region': None,
            'district': None,
            'commune': None,
            'query_params': {
                'annee': report_year,
                'organisation': str(root_organisation.id),
                'print': '1',
            },
        }
    ]

    for region in filter_context['regions']:
        scopes.append(
            {
                'scope_level': PartnerQuarterlyReportArchive.SCOPE_REGION,
                'scope_key': f'region:{region.id}',
                'region': region,
                'district': None,
                'commune': None,
                'query_params': {
                    'annee': report_year,
                    'organisation': str(root_organisation.id),
                    'region': str(region.id),
                    'print': '1',
                },
            }
        )

    for district in filter_context['districts']:
        scopes.append(
            {
                'scope_level': PartnerQuarterlyReportArchive.SCOPE_DISTRICT,
                'scope_key': f'district:{district.id}',
                'region': district.region,
                'district': district,
                'commune': None,
                'query_params': {
                    'annee': report_year,
                    'organisation': str(root_organisation.id),
                    'region': str(district.region_id),
                    'district': str(district.id),
                    'print': '1',
                },
            }
        )

    for commune in filter_context['communes']:
        scopes.append(
            {
                'scope_level': PartnerQuarterlyReportArchive.SCOPE_COMMUNE,
                'scope_key': f'commune:{commune.id}',
                'region': commune.district.region,
                'district': commune.district,
                'commune': commune,
                'query_params': {
                    'annee': report_year,
                    'organisation': str(root_organisation.id),
                    'region': str(commune.district.region_id),
                    'district': str(commune.district_id),
                    'commune': str(commune.id),
                    'print': '1',
                },
            }
        )

    return scopes


def _inject_base_href(html: str, base_url: str) -> str:
    safe_base = (base_url or 'https://www.hydria.mg/').rstrip('/') + '/'
    base_tag = f'<base href="{safe_base}">'
    if base_tag in html:
        return html
    if '<head>' in html:
        return html.replace('<head>', f'<head>{base_tag}', 1)
    return f'{base_tag}{html}'


def _render_partner_report_print_html(report_user, query_params, layout_config=None) -> str:
    from apps.core.views_rapports import build_partner_report_context

    request = _build_partner_report_request(report_user, query_params)
    context = build_partner_report_context(
        request,
        print_mode=True,
        layout_config=layout_config,
    )
    html = render_to_string(
        'core/rapport_partenaires_trimestriel.html',
        context=context,
        request=request,
    )
    base_url = getattr(settings, 'PARTNER_REPORT_ARCHIVE_BASE_URL', 'https://www.hydria.mg/')
    return _inject_base_href(html, base_url)


def _render_pdf_bytes_from_html(html: str) -> bytes:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise PartnerReportArchiveError(
            "Playwright n'est pas installe. Installe le package puis Chromium pour generer les PDF."
        ) from exc

    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1440, 'height': 2400})
            page.emulate_media(media='print')
            page.set_content(html, wait_until='load')
            page.wait_for_function(
                "() => window.__partnerReportChartsReady === true",
                timeout=60000,
            )
            page.wait_for_timeout(750)
            pdf_bytes = page.pdf(
                format='A4',
                print_background=True,
                margin={
                    'top': '12mm',
                    'right': '12mm',
                    'bottom': '12mm',
                    'left': '12mm',
                },
            )
            browser.close()
            return pdf_bytes
    except Exception as exc:
        raise PartnerReportArchiveError(
            f"Generation PDF impossible: {exc}"
        ) from exc


def _build_partner_report_context(report_user, query_params):
    from apps.core.views_rapports import build_partner_report_context

    request = _build_partner_report_request(report_user, query_params)
    return build_partner_report_context(request, print_mode=True)


def _stringify_doc_value(value) -> str:
    if value in (None, ''):
        return '-'
    if isinstance(value, bool):
        return 'Oui' if value else 'Non'
    if isinstance(value, int):
        return f"{value:,}".replace(',', ' ')
    if isinstance(value, float):
        if value.is_integer():
            return f"{int(value):,}".replace(',', ' ')
        return f"{value:,.1f}".replace(',', ' ').replace('.', ',')
    if hasattr(value, 'strftime'):
        try:
            return value.strftime('%d/%m/%Y %H:%M')
        except Exception:
            try:
                return value.strftime('%d/%m/%Y')
            except Exception:
                return str(value)
    return str(value)


def _add_doc_table(document, headers, rows):
    if not rows:
        return None

    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = str(header)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = _stringify_doc_value(value)
    return table


def _add_narrative_block(document, title='Narratif / commentaires'):
    paragraph = document.add_paragraph()
    paragraph.add_run(title).bold = True
    document.add_paragraph('................................................................................................')
    document.add_paragraph('................................................................................................')


def _flatten_dashboard_stat_rows(context):
    rows = []
    for section in context.get('dashboard_data', {}).get('stat_sections', []):
        for card in section.get('cards', []):
            value = card.get('value')
            if card.get('format') == 'float1':
                try:
                    value = round(float(value or 0), 1)
                except Exception:
                    pass
            rows.append([
                section.get('title', ''),
                card.get('label', ''),
                value,
                card.get('desc', ''),
            ])
    return rows


def _flatten_dashboard_commune_rows(context):
    rows = []
    for district_group in context.get('dashboard_data', {}).get('stats_par_commune_groups', []):
        for row in district_group.get('rows', []):
            rows.append([
                district_group.get('district_nom', ''),
                row.get('commune', ''),
                row.get('ouvrages', 0),
                row.get('habitants', 0),
                row.get('points_eau_total', 0),
                row.get('points_eau_coupes', 0),
            ])
    return rows


def _flatten_dashboard_ouvrage_rows(context):
    rows = []
    visible_columns = context.get('dashboard_data', {}).get('ouvrage_detail_visible_columns', {})
    for row in context.get('dashboard_data', {}).get('ouvrage_detail_rows', []):
        values = [row.get('ouvrage', '')]
        headers = ['Ouvrage']
        column_specs = [
            ('modalite_paiement', 'Modalite'),
            ('habitants', 'Habitants'),
            ('habitants_desservis', 'Habitants desservis'),
            ('points_eau_total', 'Points d eau'),
            ('points_eau_prives', 'Points prives'),
            ('points_eau_publics', 'Points publics'),
            ('points_eau_coupes', 'Points coupes'),
            ('ouvrages_forfaitaire', 'Forfaitaire'),
            ('ouvrages_volume', 'Volume'),
            ('ouvrages_mixte', 'Mixte'),
        ]
        for key, label in column_specs:
            if key == 'modalite_paiement' or visible_columns.get(key):
                headers.append(label)
                values.append(row.get(key, ''))
        rows.append((headers, values))
    return rows


def _flatten_grading_distribution_rows(grading_data):
    rows = []
    dist_from = {row.get('value'): row.get('count', 0) for row in grading_data.get('dist_from_rows', [])}
    dist_to = {row.get('value'): row.get('count', 0) for row in grading_data.get('dist_to_rows', [])}
    for row in grading_data.get('dist_from_rows', []):
        grade_value = row.get('value')
        rows.append([
            row.get('label', ''),
            dist_from.get(grade_value, 0),
            dist_to.get(grade_value, 0),
        ])
    return rows


def _flatten_recettes_commune_rows(recettes_data):
    rows = []
    for region_group in recettes_data.get('communes_breakdown', []):
        for district_group in region_group.get('districts', []):
            for row in district_group.get('rows', []):
                rows.append([
                    region_group.get('region_nom', ''),
                    district_group.get('district_nom', ''),
                    row.get('commune_nom', ''),
                    row.get('nb_ouvrages', 0),
                    row.get('total_prevu', 0),
                    row.get('total_paye', 0),
                    row.get('taux', 0),
                ])
    return rows


def _build_partner_report_docx_bytes(context) -> bytes:
    try:
        from docx import Document
    except ImportError:
        return _build_minimal_partner_report_docx_bytes(context)

    document = Document()
    report_identity = context.get('report_identity', {})
    generated_at = context.get('generated_at')
    document.core_properties.title = 'Rapport partenaires trimestriel'
    document.core_properties.subject = report_identity.get('period_label', '')

    document.add_heading('Rapport partenaires trimestriel', level=0)
    document.add_paragraph(
        "Document Word genere automatiquement pour servir de base editable. "
        "Vous pouvez ajouter du narratif, inserer des images et supprimer les blocs inutiles."
    )
    _add_doc_table(
        document,
        ['Champ', 'Valeur'],
        [
            ['Organisation', report_identity.get('organisation_label', '')],
            ['Niveau', report_identity.get('scope_level_label', '')],
            ['Perimetre', report_identity.get('scope_label', '')],
            ['Region', report_identity.get('region_label', '')],
            ['District', report_identity.get('district_label', '')],
            ['Commune', report_identity.get('commune_label', '')],
            ['Periode', report_identity.get('period_label', '')],
            ['Genere le', generated_at],
        ],
    )
    _add_narrative_block(document, 'Introduction / synthese')

    document.add_page_break()
    document.add_heading('1. Dashboard principal', level=1)
    dashboard_rows = _flatten_dashboard_stat_rows(context)
    _add_doc_table(document, ['Bloc', 'Indicateur', 'Valeur', 'Commentaire'], dashboard_rows)
    if context.get('dashboard_data', {}).get('modalite_rows'):
        document.add_paragraph('Repartition par modalite de paiement')
        _add_doc_table(
            document,
            ['Modalite', 'Ouvrages', 'Points d eau'],
            [
                [row.get('label', ''), row.get('ouvrages', 0), row.get('points_eau', 0)]
                for row in context['dashboard_data']['modalite_rows']
            ],
        )
    if context.get('dashboard_data', {}).get('is_commune_detail_mode'):
        document.add_paragraph('Detail des ouvrages de la commune selectionnee')
        ouvrage_rows = _flatten_dashboard_ouvrage_rows(context)
        if ouvrage_rows:
            headers = ouvrage_rows[0][0]
            _add_doc_table(document, headers, [values for _, values in ouvrage_rows])
    else:
        commune_rows = _flatten_dashboard_commune_rows(context)
        if commune_rows:
            document.add_paragraph('Synthese territoriale par commune')
            _add_doc_table(
                document,
                ['District', 'Commune', 'Ouvrages', 'Habitants', 'Points d eau', 'Points coupes'],
                commune_rows,
            )
    _add_narrative_block(document)

    document.add_page_break()
    document.add_heading('2. Plannings visites', level=1)
    _add_doc_table(
        document,
        ['Indicateur', 'Valeur'],
        [
            ['Ouvrages affectes', context.get('visites_data', {}).get('nb_ouvrages', 0)],
            ['Total visites', context.get('visites_data', {}).get('total_visites', 0)],
            ['Visites reparation', context.get('visites_data', {}).get('total_visites_reparation', 0)],
            ['Reparations panne', context.get('visites_data', {}).get('total_visites_reparation_panne', 0)],
            ['Moyenne par ouvrage', context.get('visites_data', {}).get('moyenne_visites', 0)],
            ['Mediane des visites', context.get('visites_data', {}).get('median_visites', 0)],
        ],
    )
    _add_narrative_block(document)

    grading_data = context.get('grading_data', {})
    document.add_page_break()
    document.add_heading('3. Analyse grading annuel', level=1)
    _add_doc_table(
        document,
        ['Indicateur', 'Valeur'],
        [
            ['Ouvrages filtres', grading_data.get('stats', {}).get('ouvrages_filtres', 0)],
            ['Avec grading annee 1', grading_data.get('stats', {}).get('avec_grading_from', 0)],
            ['Avec grading annee 2', grading_data.get('stats', {}).get('avec_grading_to', 0)],
            ['Avec deux annees', grading_data.get('stats', {}).get('avec_deux_annees', 0)],
            ['Uniquement depart', grading_data.get('stats', {}).get('uniquement_depart', 0)],
            ['Uniquement arrivee', grading_data.get('stats', {}).get('uniquement_arrivee', 0)],
        ],
    )
    grading_distribution_rows = _flatten_grading_distribution_rows(grading_data)
    if grading_distribution_rows:
        document.add_paragraph('Distribution par grade')
        _add_doc_table(document, ['Grade', 'Annee 1', 'Annee 2'], grading_distribution_rows)
    if grading_data.get('transition_top'):
        document.add_paragraph('Principales transitions')
        _add_doc_table(
            document,
            ['Depuis', 'Vers', 'Nombre'],
            [
                [row.get('from', ''), row.get('to', ''), row.get('count', 0)]
                for row in grading_data['transition_top']
            ],
        )
    _add_narrative_block(document)

    recettes_data = context.get('recettes_data', {})
    document.add_page_break()
    document.add_heading('4. Recettes globales', level=1)
    _add_doc_table(
        document,
        ['Indicateur', 'Valeur'],
        [
            ['Total prevu', recettes_data.get('stats_global', {}).get('total_prevu', 0)],
            ['Total paye', recettes_data.get('stats_global', {}).get('total_paye', 0)],
            ['Ecart', recettes_data.get('stats_global', {}).get('ecart', 0)],
            ['Taux recouvrement', f"{recettes_data.get('stats_global', {}).get('taux_recouvrement', 0)}%"],
            ['Ouvrages couverts', recettes_data.get('stats_global', {}).get('nb_ouvrages_total', 0)],
        ],
    )
    if recettes_data.get('categories'):
        document.add_paragraph('Synthese par workflow')
        _add_doc_table(
            document,
            ['Workflow', 'Prevision', 'Paye', 'Ecart', 'Taux', 'Ouvrages'],
            [
                [
                    row.get('label', ''),
                    row.get('prevision', 0),
                    row.get('paye', 0),
                    row.get('ecart', 0),
                    f"{row.get('taux', 0)}%",
                    row.get('nb_ouvrages', 0),
                ]
                for row in recettes_data['categories']
            ],
        )
    recettes_commune_rows = _flatten_recettes_commune_rows(recettes_data)
    if recettes_commune_rows:
        document.add_paragraph('Repartition territoriale')
        _add_doc_table(
            document,
            ['Region', 'District', 'Commune', 'Ouvrages', 'Total prevu', 'Total paye', 'Taux'],
            recettes_commune_rows,
        )
    _add_narrative_block(document)

    representants_data = context.get('representants_data', {})
    document.add_page_break()
    document.add_heading('5. Evolution des avis representants', level=1)
    _add_doc_table(
        document,
        ['Indicateur', 'Valeur'],
        [
            ['Representants', representants_data.get('stats', {}).get('nb_representants', 0)],
            ['Avis', representants_data.get('stats', {}).get('total_avis', 0)],
            ['Points negatifs', representants_data.get('stats', {}).get('total_points_negatifs', 0)],
            ['Taux negatif global', f"{representants_data.get('stats', {}).get('taux_negatif_global', 0)}%"],
            ['Satisfaction moyenne', f"{representants_data.get('stats', {}).get('satisfaction_moyenne', 0)}%"],
            ['Score moyen global', representants_data.get('stats', {}).get('score_moyen_global', 0)],
        ],
    )
    if representants_data.get('yearly_rows'):
        document.add_paragraph('Evolution annuelle')
        _add_doc_table(
            document,
            [
                'Annee',
                'Avis',
                'Points negatifs',
                'Taux negatif',
                'Score moyen',
                'Satisfaction moyenne',
            ],
            [
                [
                    row.get('annee', ''),
                    row.get('total_avis', 0),
                    row.get('points_negatifs', 0),
                    f"{row.get('taux_negatif', 0)}%",
                    row.get('score_moyen', 0),
                    f"{row.get('satisfaction_moyenne', 0)}%",
                ]
                for row in representants_data['yearly_rows']
            ],
        )
    _add_narrative_block(document)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_minimal_partner_report_docx_bytes(context) -> bytes:
    import zipfile
    from xml.sax.saxutils import escape

    report_identity = context.get('report_identity', {})
    generated_at = context.get('generated_at') or ''
    lines = [
        'Rapport partenaires trimestriel',
        'Document Word genere automatiquement pour servir de base editable.',
        f"Organisation: {report_identity.get('organisation_label', '')}",
        f"Niveau: {report_identity.get('scope_level_label', '')}",
        f"Perimetre: {report_identity.get('scope_label', '')}",
        f"Periode: {report_identity.get('period_label', '')}",
        f"Genere le: {generated_at}",
    ]
    paragraphs_xml = ''.join(
        f'<w:p><w:r><w:t>{escape(str(line))}</w:t></w:r></w:p>'
        for line in lines
    )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:wp14="http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
        'xmlns:wpi="http://schemas.microsoft.com/office/word/2010/wordprocessingInk" '
        'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'mc:Ignorable="w14 wp14"><w:body>'
        f'{paragraphs_xml}'
        '<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" '
        'w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/>'
        '<w:cols w:space="708"/><w:docGrid w:linePitch="360"/></w:sectPr>'
        '</w:body></w:document>'
    )
    content_types_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>"""
    rels_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""
    app_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
 xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <Application>Hydria</Application>
</Properties>"""
    core_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
 xmlns:dc="http://purl.org/dc/elements/1.1/"
 xmlns:dcterms="http://purl.org/dc/terms/"
 xmlns:dcmitype="http://purl.org/dc/dcmitype/"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:title>Rapport partenaires trimestriel</dc:title>
  <dc:subject>{escape(str(report_identity.get('period_label', '')))}</dc:subject>
</cp:coreProperties>"""

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, mode='w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', content_types_xml)
        archive.writestr('_rels/.rels', rels_xml)
        archive.writestr('docProps/app.xml', app_xml)
        archive.writestr('docProps/core.xml', core_xml)
        archive.writestr('word/document.xml', document_xml)
    return buffer.getvalue()


def _prepare_archive_render_query_params(archive: PartnerQuarterlyReportArchive):
    query_params = dict(archive.filters or {})
    query_params['print'] = '1'
    query_params['archive_period_year'] = str(archive.report_year)
    query_params['archive_period_quarter'] = str(archive.report_quarter)
    return query_params


def regenerate_partner_report_archive_html(
    archive: PartnerQuarterlyReportArchive,
    generated_by=None,
    automatic: bool = False,
):
    report_user = get_partner_report_generation_user(archive.organisation)
    if not report_user:
        raise PartnerReportArchiveError(
            f"Aucun utilisateur actif lie a l'organisation {archive.organisation.nom_organisation}."
        )

    html = _render_partner_report_print_html(
        report_user=report_user,
        query_params=_prepare_archive_render_query_params(archive),
        layout_config=archive.layout_config or {},
    )
    archive.html_content = html
    archive.generated_at = timezone.now()
    if generated_by is not None:
        archive.generated_by = generated_by
    archive.is_automatic = automatic
    archive.save(update_fields=['html_content', 'generated_at', 'generated_by', 'is_automatic', 'updated_at'])
    return archive


def _build_archive_filename(archive: PartnerQuarterlyReportArchive):
    territory_tokens = []
    if archive.region_id:
        territory_tokens.append(archive.region.nom_region.replace(' ', '-'))
    if archive.district_id:
        territory_tokens.append(archive.district.nom_district.replace(' ', '-'))
    if archive.commune_id:
        territory_tokens.append(archive.commune.nom_commune.replace(' ', '-'))
    suffix = '-'.join(token for token in territory_tokens if token) or 'organisation'
    organisation_token = archive.organisation.nom_organisation.replace(' ', '-')
    return f"rapport-partenaires-{organisation_token}-T{archive.report_quarter}-{archive.report_year}-{suffix}.docx"


def generate_partner_report_archive_for_scope(
    organisation: Organisation,
    report_user,
    quarter_spec: PartnerReportQuarterSpec,
    scope_payload,
    generated_by=None,
    overwrite: bool = False,
    automatic: bool = True,
):
    render_query_params = dict(scope_payload['query_params'])
    render_query_params['archive_period_year'] = str(quarter_spec.report_year)
    render_query_params['archive_period_quarter'] = str(quarter_spec.report_quarter)

    archive, created = PartnerQuarterlyReportArchive.objects.get_or_create(
        organisation=organisation,
        report_year=quarter_spec.report_year,
        report_quarter=quarter_spec.report_quarter,
        scope_key=scope_payload['scope_key'],
        defaults={
            'scope_level': scope_payload['scope_level'],
            'region': scope_payload['region'],
            'district': scope_payload['district'],
            'commune': scope_payload['commune'],
            'filters': render_query_params,
            'generated_by': generated_by,
            'is_automatic': automatic,
            'is_user_uploaded': False,
            'generated_at': timezone.now(),
        },
    )

    if not created and not overwrite and archive.document_file and archive.is_word_document:
        if archive.has_html_archive:
            return archive, False

    archive.scope_level = scope_payload['scope_level']
    archive.region = scope_payload['region']
    archive.district = scope_payload['district']
    archive.commune = scope_payload['commune']
    archive.filters = render_query_params
    archive.generated_by = generated_by
    archive.is_automatic = automatic
    archive.is_user_uploaded = False
    archive.uploaded_by = None
    archive.generated_at = timezone.now()

    context = _build_partner_report_context(report_user, render_query_params)
    archive.html_content = _render_partner_report_print_html(
        report_user=report_user,
        query_params=render_query_params,
        layout_config=archive.layout_config or {},
    )
    document_bytes = _build_partner_report_docx_bytes(context)
    filename = _build_archive_filename(archive)
    target_path = PartnerQuarterlyReportArchive._meta.get_field('document_file').generate_filename(archive, filename)

    if archive.document_file and archive.document_file.name and default_storage.exists(archive.document_file.name):
        default_storage.delete(archive.document_file.name)
    if default_storage.exists(target_path):
        default_storage.delete(target_path)

    archive.document_file.save(filename, ContentFile(document_bytes), save=False)
    archive.file_size = len(document_bytes)
    archive.save()
    return archive, True


def regenerate_partner_report_archive(
    archive: PartnerQuarterlyReportArchive,
    generated_by=None,
    automatic: bool = False,
):
    report_user = get_partner_report_generation_user(archive.organisation)
    if not report_user:
        raise PartnerReportArchiveError(
            f"Aucun utilisateur actif lie a l'organisation {archive.organisation.nom_organisation}."
        )

    quarter_spec = PartnerReportQuarterSpec(
        report_year=archive.report_year,
        report_quarter=archive.report_quarter,
        due_date=get_partner_report_quarter_due_date(archive.report_year, archive.report_quarter),
    )
    scope_payload = {
        'scope_level': archive.scope_level,
        'scope_key': archive.scope_key,
        'region': archive.region,
        'district': archive.district,
        'commune': archive.commune,
        'query_params': dict(archive.filters or {}),
    }

    regenerated_archive, _generated = generate_partner_report_archive_for_scope(
        organisation=archive.organisation,
        report_user=report_user,
        quarter_spec=quarter_spec,
        scope_payload=scope_payload,
        generated_by=generated_by,
        overwrite=True,
        automatic=automatic,
    )
    return regenerated_archive


def generate_partner_report_archives_for_organisation(
    organisation: Organisation,
    period_specs: Optional[Iterable[PartnerReportQuarterSpec]] = None,
    generated_by=None,
    overwrite: bool = False,
    automatic: bool = True,
):
    report_user = get_partner_report_generation_user(organisation)
    if not report_user:
        raise PartnerReportArchiveError(
            f"Aucun utilisateur actif lie a l'organisation {organisation.nom_organisation}."
        )

    period_specs = list(period_specs or get_due_partner_report_quarter_specs())
    created_count = 0
    skipped_count = 0
    archives = []

    for quarter_spec in period_specs:
        scopes = _collect_archive_scopes(report_user, quarter_spec.report_year)
        for scope_payload in scopes:
            archive, generated = generate_partner_report_archive_for_scope(
                organisation=organisation,
                report_user=report_user,
                quarter_spec=quarter_spec,
                scope_payload=scope_payload,
                generated_by=generated_by,
                overwrite=overwrite,
                automatic=automatic,
            )
            archives.append(archive)
            if generated:
                created_count += 1
            else:
                skipped_count += 1

    return {
        'organisation': organisation,
        'archives': archives,
        'created_count': created_count,
        'skipped_count': skipped_count,
    }


def generate_due_partner_report_archives(
    organisations: Optional[Iterable[Organisation]] = None,
    today: Optional[date] = None,
    generated_by=None,
    overwrite: bool = False,
    automatic: bool = True,
):
    organisations = list(
        organisations
        or Organisation.objects.order_by('type_organisation', 'nom_organisation')
    )
    period_specs = get_due_partner_report_quarter_specs(today=today)

    results = []
    for organisation in organisations:
        try:
            results.append(
                generate_partner_report_archives_for_organisation(
                    organisation=organisation,
                    period_specs=period_specs,
                    generated_by=generated_by,
                    overwrite=overwrite,
                    automatic=automatic,
                )
            )
        except PartnerReportArchiveError as exc:
            results.append(
                {
                    'organisation': organisation,
                    'archives': [],
                    'created_count': 0,
                    'skipped_count': 0,
                    'error': str(exc),
                }
            )
    return {
        'period_specs': period_specs,
        'results': results,
    }


def build_partner_report_archive_queryset_for_user(user):
    queryset = PartnerQuarterlyReportArchive.objects.select_related(
        'organisation',
        'region',
        'district__region',
        'commune__district__region',
    )
    if getattr(user, 'is_superuser', False) or is_visiteur(user):
        return queryset

    organisation = get_user_organisation(user)
    if not organisation:
        return queryset.none()
    return queryset.filter(organisation=organisation)
