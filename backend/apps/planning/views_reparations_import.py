"""
Import Excel en masse des réparations (page Suivi des Réparations).

Flux en 2 temps, sans état conservé côté serveur entre les deux appels :
1) api_reparations_import_preview : parse le fichier, valide chaque ligne,
   détecte les doublons (stricts = ignorés automatiquement, suspects = à
   confirmer), et renvoie tout au client (rien n'est écrit en base).
2) api_reparations_import_confirm : reçoit du client la liste des lignes
   à importer réellement (lignes sans conflit + doublons suspects cochés),
   les revalide et les insère en une seule transaction.

Une réparation est un Planning dont le fiche_poste appartient aux activités
de réparation (get_reparations_fiches_queryset). Comme le formulaire manuel
"Ajouter une réparation", la cible (point d'eau OU équipement) et ses composants
sont enregistrés sur le Planning (apply_repair_target) et contrôlés selon les
règles de l'activité (verifier_regles_fiche). La Désignation est optionnelle :
si elle est renseignée, une ligne de détail (PlanningDetail) est ajoutée.
Aucune donnée de référence nouvelle n'est créée par cet import.
"""
import io
import json
import traceback
import unicodedata
import uuid
from datetime import datetime, date as date_cls
from collections import defaultdict
from decimal import Decimal, InvalidOperation

import openpyxl
import pytz
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods

from apps.billing.models import PointEau
from apps.employees.models import Employe
from apps.infrastructure.models import EquipementConstitutif
from apps.organization.models import Ouvrage
from apps.inspections.models import ComposantTypeEquipement, ComposantTypePointEau
from apps.planning.models import Designation, Planning, PlanningDetail
from apps.planning.reparation_composants import (
    ComposantsError, apply_repair_target, composants_disponibles, verifier_regles_fiche,
)
from apps.planning.repair_classification import get_reparations_fiches_queryset
from apps.planning.views_reparations import _get_repair_modal_employees_queryset
from apps.planning.views_reparations_apis import _get_scoped_ouvrage_ids_for_request

TZ_MADAGASCAR = pytz.timezone('Indian/Antananarivo')

COLONNES_ATTENDUES = [
    'ID Ouvrage', 'Activité', 'Date',
    'ID Point d\'eau', 'ID Équipement', 'Composants',
    'ID Désignation', 'Quantité', 'Prix unitaire',
    'Description', 'ID Responsable', 'Validée', 'Statut',
]
MAX_LIGNES = 5000
VALIDEE_VRAI = {'oui', 'yes', 'true', '1', 'vrai', 'x'}
STATUT_LABEL_BY_KEY = dict(Planning.STATUS_CHOICES)


def _norm(value):
    """Trim + normalise une valeur de cellule en texte (chaîne vide si vide)."""
    if value is None:
        return ''
    return str(value).strip()


def _norm_key(value):
    """Clé de comparaison insensible à la casse pour les correspondances par nom."""
    return _norm(value).lower()


def _strip_accents(value):
    text = unicodedata.normalize('NFD', _norm(value))
    return ''.join(c for c in text if unicodedata.category(c) != 'Mn')


def _header_key(value):
    """Clé normalisée (minuscule, sans accents) utilisée pour reconnaître les en-têtes de colonnes."""
    return _strip_accents(value).lower()


STATUT_KEY_BY_LABEL = {
    _header_key(label): key for key, label in Planning.STATUS_CHOICES
}


def _parse_date(value):
    """Parse une date Excel (objet datetime natif) ou texte JJ/MM/AAAA / AAAA-MM-JJ."""
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_cls):
        return value
    text = _norm(value)
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _is_uuid(val):
    """Vérifie si une chaîne est un UUID valide."""
    if not val:
        return False
    try:
        uuid.UUID(str(val).strip())
        return True
    except (ValueError, AttributeError, TypeError):
        return False


def _date_to_naive_midnight(d):
    """Date -> datetime timezone-aware Madagascar (minuit), selon les règles timezone du projet."""
    return TZ_MADAGASCAR.localize(datetime.combine(d, datetime.min.time()))


def _planning_local_date(planning_date):
    """DateTimeField Planning.date -> date() en heure Madagascar."""
    if not planning_date:
        return None
    if not isinstance(planning_date, datetime) and isinstance(planning_date, date_cls):
        return planning_date
    try:
        aware = planning_date if timezone.is_aware(planning_date) else timezone.make_aware(planning_date)
        return aware.astimezone(TZ_MADAGASCAR).date()
    except Exception:
        return None


def _parse_validee(value):
    return _norm_key(value) in VALIDEE_VRAI


def _parse_statut(value):
    """Retourne (clé_statut_ou_None, erreur_ou_None)."""
    text = _norm(value)
    if not text:
        return None, None
    key = _header_key(text)
    if key in STATUT_LABEL_BY_KEY:
        return key, None
    if key in STATUT_KEY_BY_LABEL:
        return STATUT_KEY_BY_LABEL[key], None
    valeurs = ', '.join(f"{k} ({v})" for k, v in Planning.STATUS_CHOICES)
    return None, f"Statut invalide : « {text} » (valeurs possibles : {valeurs})"


def _parse_decimal(value, default=None):
    if value in (None, ''):
        return default
    try:
        return Decimal(str(value).replace(',', '.').replace(' ', ''))
    except (InvalidOperation, ValueError):
        return None


def _resoudre_composants(valeur, equipement, point_eau):
    """
    Colonne Composants : noms (ou UUID) séparés par « ; », pris dans le catalogue
    du type de la cible. Retourne (ids, noms, erreur).
    """
    texte = _norm(valeur)
    if not texte:
        return [], [], None
    saisis = [part.strip() for part in texte.replace('\n', ';').replace(',', ';').split(';') if part.strip()]
    if not (equipement or point_eau):
        return [], [], "Composants indiqués sans ID Point d'eau ni ID Équipement"
    disponibles = composants_disponibles(equipement=equipement, point_eau=point_eau)
    if not disponibles:
        return [], [], "Aucun composant n'est défini pour le type de cette cible"
    par_id = {c['id']: c for c in disponibles}
    par_nom = {_header_key(c['nom']): c for c in disponibles}
    ids, noms, inconnus = [], [], []
    for saisi in saisis:
        comp = par_id.get(saisi) or par_nom.get(_header_key(saisi))
        if not comp:
            inconnus.append(saisi)
        elif comp['id'] not in ids:
            ids.append(comp['id'])
            noms.append(comp['nom'])
    if inconnus:
        return [], [], (
            f"Composant(s) inconnu(s) pour cette cible : « {' ; '.join(inconnus)} » "
            f"(possibles : {' ; '.join(c['nom'] for c in disponibles)})"
        )
    return ids, noms, None


def _fiche_label(fiche):
    type_name = getattr(getattr(fiche, 'type_activites', None), 'nomtype', '') or ''
    return f"{type_name} - {fiche.tache}" if type_name else (fiche.tache or '')


def _build_reference_maps():
    """Construit les dictionnaires de correspondance nom -> objet (pour Activité)."""
    fiches = list(get_reparations_fiches_queryset().select_related('type_activites'))
    fiche_by_key = {}
    fiche_ambigus = set()
    for fiche in fiches:
        key = _norm_key(fiche.tache)
        if not key:
            continue
        if key in fiche_by_key:
            fiche_ambigus.add(key)
        else:
            fiche_by_key[key] = fiche
    return fiche_by_key, fiche_ambigus


def _find_employe_autorise(employes_autorises_by_id, employes_autorises_by_key, id_value, nom_value):
    """Cherche un employé autorisé par UUID (priorité) puis par surnom/nom."""
    id_norm = _norm(id_value)
    if id_norm and id_norm in employes_autorises_by_id:
        employe = employes_autorises_by_id[id_norm]
        return employe, (employe.surnom or employe.nom), None
    if id_norm:
        return None, '', f"Responsable introuvable ou non autorisé (ID) : « {id_norm} »"

    nom_norm = _norm(nom_value)
    if not nom_norm:
        return None, '', None
    employe = employes_autorises_by_key.get(_norm_key(nom_norm))
    if employe:
        return employe, (employe.surnom or employe.nom), None
    return None, '', f"Responsable indiqué non reconnu parmi les employés autorisés : « {nom_norm} »"


@login_required
def api_reparations_import_template(request):
    """Télécharge le modèle Excel pour l'import en masse des réparations."""
    _, ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
    ouvrage_ids = list(ouvrage_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Import reparations'
    ws.append(COLONNES_ATTENDUES)

    widths = [38, 30, 12, 38, 38, 35, 38, 10, 12, 40, 38, 10, 14]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

    ws_ouvrages = wb.create_sheet('Ouvrages valables')
    ws_ouvrages.append(['ID Ouvrage', 'Nom'])
    for ouvrage in Ouvrage.objects.filter(id__in=ouvrage_ids, type_ouvrage__existe=True).order_by('nom_ouvrage')[:5000]:
        ws_ouvrages.append([str(ouvrage.id), ouvrage.nom_ouvrage])
    ws_ouvrages.column_dimensions['A'].width = 38
    ws_ouvrages.column_dimensions['B'].width = 40

    composants_par_type_eq = defaultdict(list)
    for comp in ComposantTypeEquipement.objects.select_related('type_equipement'):
        composants_par_type_eq[comp.type_equipement_id].append(comp)
    composants_par_type_pe = defaultdict(list)
    for comp in ComposantTypePointEau.objects.select_related('type_point_eau'):
        composants_par_type_pe[comp.type_point_eau_id].append(comp)

    def _liste_noms(comps):
        return ' ; '.join(c.nom_composant for c in comps)

    ws_pe = wb.create_sheet('Points deau')
    ws_pe.append(['ID Point d\'eau', 'IDXls', 'Nom / Code', 'Ouvrage', 'Type', 'Composants possibles'])
    for pe in PointEau.objects.filter(ouvrage_id__in=ouvrage_ids).select_related('ouvrage', 'type_point_eau').order_by('ouvrage__nom_ouvrage')[:5000]:
        ws_pe.append([
            str(pe.id), pe.idxls or '', pe.nomabonne or pe.code_point_eau or '',
            pe.ouvrage.nom_ouvrage if pe.ouvrage else '',
            pe.type_point_eau.nom_type if pe.type_point_eau else '',
            _liste_noms(composants_par_type_pe.get(pe.type_point_eau_id, [])),
        ])
    ws_pe.column_dimensions['A'].width = 38
    ws_pe.column_dimensions['B'].width = 12
    ws_pe.column_dimensions['C'].width = 30
    ws_pe.column_dimensions['D'].width = 30
    ws_pe.column_dimensions['E'].width = 18
    ws_pe.column_dimensions['F'].width = 70

    ws_eq = wb.create_sheet('Equipements')
    ws_eq.append(['ID Équipement', 'Nom', 'Type d\'équipement', 'Ouvrage', 'Composants possibles'])
    for eq in EquipementConstitutif.objects.filter(ouvrage_id__in=ouvrage_ids).select_related('ouvrage', 'type_equipement').order_by('ouvrage__nom_ouvrage')[:5000]:
        ws_eq.append([
            str(eq.id), eq.nom,
            eq.type_equipement.nom_type if eq.type_equipement else '',
            eq.ouvrage.nom_ouvrage if eq.ouvrage else '',
            _liste_noms(composants_par_type_eq.get(eq.type_equipement_id, [])),
        ])
    ws_eq.column_dimensions['A'].width = 38
    ws_eq.column_dimensions['B'].width = 30
    ws_eq.column_dimensions['C'].width = 25
    ws_eq.column_dimensions['D'].width = 30
    ws_eq.column_dimensions['E'].width = 70

    ws_comp = wb.create_sheet('Composants possibles')
    ws_comp.append([
        'Cible', 'Type',
        'Composant (nom à recopier dans la colonne Composants, plusieurs séparés par « ; »)',
        'ID Composant',
    ])
    for comps in composants_par_type_eq.values():
        for comp in comps:
            ws_comp.append(['Équipement', comp.type_equipement.nom_type, comp.nom_composant, str(comp.id)])
    for comps in composants_par_type_pe.values():
        for comp in comps:
            ws_comp.append(["Point d'eau", comp.type_point_eau.nom_type, comp.nom_composant, str(comp.id)])
    ws_comp.column_dimensions['A'].width = 14
    ws_comp.column_dimensions['B'].width = 28
    ws_comp.column_dimensions['C'].width = 60
    ws_comp.column_dimensions['D'].width = 38

    ws_designations = wb.create_sheet('Designations')
    ws_designations.append(['ID Désignation', 'Nom', 'Prix unitaire (Ar)'])
    for d in Designation.objects.filter(actif=True).order_by('nom')[:5000]:
        ws_designations.append([str(d.id), d.nom, float(d.prix_unitaire) if d.prix_unitaire else None])
    ws_designations.column_dimensions['A'].width = 38
    ws_designations.column_dimensions['B'].width = 45

    ws_employes = wb.create_sheet('Employes autorises')
    ws_employes.append(['ID Responsable', 'Surnom'])
    employes_autorises = _get_repair_modal_employees_queryset(request.user, ouvrage_ids=ouvrage_ids)
    for employe in employes_autorises.order_by('surnom'):
        ws_employes.append([str(employe.id), employe.surnom or employe.nom])
    ws_employes.column_dimensions['A'].width = 38
    ws_employes.column_dimensions['B'].width = 25

    ws_activites = wb.create_sheet('Activites valides')
    ws_activites.append(['Activité (nom exact à recopier dans la colonne Activité)'])
    for fiche in get_reparations_fiches_queryset().select_related('type_activites').order_by('type_activites__nomtype', 'tache'):
        ws_activites.append([fiche.tache])
    ws_activites.column_dimensions['A'].width = 55

    ws_statuts = wb.create_sheet('Statuts valides')
    ws_statuts.append(['Valeur à utiliser dans la colonne Statut', 'Libellé'])
    for key, label in Planning.STATUS_CHOICES:
        ws_statuts.append([key, label])
    ws_statuts.column_dimensions['A'].width = 20
    ws_statuts.column_dimensions['B'].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="modele_import_reparations.xlsx"'
    return response


def _classify_rows(rows_raw, ouvrage_ids, request_user):
    """
    rows_raw: liste de dicts avec les valeurs brutes issues du fichier
    (déjà indexées par en-tête de colonne).

    Retourne (invalides, a_importer, doublons_suspects, doublons_stricts).
    """
    fiche_by_key, fiche_ambigus = _build_reference_maps()

    ouvrage_by_id = {
        str(o.id): o
        for o in Ouvrage.objects.filter(id__in=ouvrage_ids, type_ouvrage__existe=True)
    }

    employes_autorises_qs = _get_repair_modal_employees_queryset(request_user, ouvrage_ids=ouvrage_ids)
    employes_autorises_by_id = {str(e.id): e for e in employes_autorises_qs}
    employes_autorises_by_key = {}
    for employe in employes_autorises_qs:
        key = _norm_key(employe.surnom or employe.nom)
        if key:
            employes_autorises_by_key.setdefault(key, employe)

    fiche_ids_in_file = set()
    ouvrage_ids_in_file = set()

    resolved_rows = []
    invalides = []

    for row in rows_raw:
        ligne = row['ligne']
        motifs = []

        ouvrage_id_saisi = _norm(row.get('ouvrage_id'))
        activite_nom = _norm(row.get('activite'))
        date_val = _parse_date(row.get('date'))
        point_eau_id_saisi = _norm(row.get('point_eau_id'))
        equipement_id_saisi = _norm(row.get('equipement_id'))
        designation_id_saisi = _norm(row.get('designation_id'))
        composants_saisis = row.get('composants')
        quantite_val = _parse_decimal(row.get('quantite'), default=None)
        prix_val = _parse_decimal(row.get('prix_unitaire'), default=None)
        description = _norm(row.get('description'))
        responsable_id_saisi = _norm(row.get('responsable_id'))
        validee = _parse_validee(row.get('validee'))
        statut_key, statut_erreur = _parse_statut(row.get('statut'))
        if statut_erreur:
            motifs.append(statut_erreur)

        if not ouvrage_id_saisi:
            motifs.append("ID Ouvrage manquant")
        if not activite_nom:
            motifs.append("Activité manquante")
        if not row.get('date'):
            motifs.append("Date manquante")
        elif date_val is None:
            motifs.append("Date invalide (attendu JJ/MM/AAAA)")

        ouvrage = ouvrage_by_id.get(ouvrage_id_saisi) if ouvrage_id_saisi else None
        if ouvrage_id_saisi and not ouvrage:
            motifs.append(f"Ouvrage introuvable ou non accessible (ID) : « {ouvrage_id_saisi} »")

        fiche = fiche_by_key.get(_norm_key(activite_nom)) if activite_nom else None
        if activite_nom and not fiche:
            motifs.append(f"Activité introuvable dans la liste des activités de réparation : « {activite_nom} »")
        elif activite_nom and _norm_key(activite_nom) in fiche_ambigus:
            motifs.append(f"Nom d'activité ambigu : « {activite_nom} »")

        if point_eau_id_saisi and equipement_id_saisi:
            motifs.append("Un seul de ID Point d'eau / ID Équipement peut être renseigné, pas les deux")

        point_eau = None
        if point_eau_id_saisi:
            if not _is_uuid(point_eau_id_saisi):
                motifs.append(f"ID Point d'eau invalide (UUID attendu) : « {point_eau_id_saisi} »")
            else:
                point_eau = PointEau.objects.filter(id=point_eau_id_saisi).select_related('ouvrage').first()
                if not point_eau:
                    motifs.append(f"Point d'eau introuvable : « {point_eau_id_saisi} »")
                elif ouvrage and point_eau.ouvrage_id != ouvrage.id:
                    motifs.append("Le point d'eau indiqué n'appartient pas à l'ouvrage de la ligne")

        equipement = None
        if equipement_id_saisi:
            if not _is_uuid(equipement_id_saisi):
                motifs.append(f"ID Équipement invalide (UUID attendu) : « {equipement_id_saisi} »")
            else:
                equipement = EquipementConstitutif.objects.filter(id=equipement_id_saisi).select_related('ouvrage').first()
                if not equipement:
                    motifs.append(f"Équipement introuvable : « {equipement_id_saisi} »")
                elif ouvrage and equipement.ouvrage_id != ouvrage.id:
                    motifs.append("L'équipement indiqué n'appartient pas à l'ouvrage de la ligne")

        designation = None
        if designation_id_saisi:
            if not _is_uuid(designation_id_saisi):
                motifs.append(f"ID Désignation invalide (UUID attendu) : « {designation_id_saisi} »")
            else:
                designation = Designation.objects.filter(id=designation_id_saisi, actif=True).first()
                if not designation:
                    motifs.append(f"Désignation introuvable : « {designation_id_saisi} »")

        # Cible et composants : mêmes règles que le formulaire manuel (Désignation optionnelle).
        composant_ids, composant_noms = [], []
        cible_valide = not (point_eau_id_saisi and equipement_id_saisi) \
            and not (point_eau_id_saisi and not point_eau) \
            and not (equipement_id_saisi and not equipement)
        if cible_valide:
            composant_ids, composant_noms, composants_erreur = _resoudre_composants(
                composants_saisis, equipement, point_eau
            )
            if composants_erreur:
                motifs.append(composants_erreur)
            elif fiche:
                try:
                    verifier_regles_fiche(fiche, equipement, point_eau, composant_ids)
                except ComposantsError as exc:
                    motifs.append(str(exc))

        employe, employe_nom_affiche, employe_erreur = _find_employe_autorise(
            employes_autorises_by_id, employes_autorises_by_key,
            responsable_id_saisi, row.get('responsable_nom')
        )
        # Le responsable est optionnel : on ne bloque pas la ligne, on l'ignore et on le signale.
        responsable_warning = employe_erreur

        if motifs:
            invalides.append({
                'ligne': ligne,
                'motif' : ' ; '.join(motifs),
                'valeurs': {
                    'ouvrage': ouvrage_id_saisi,
                    'activite': activite_nom,
                    'date': _norm(row.get('date')),
                },
            })
            continue

        remarque_parts = []
        if description:
            remarque_parts.append(description)
        if responsable_warning:
            remarque_parts.append(responsable_warning)
        remarque = '\n'.join(remarque_parts)

        fiche_ids_in_file.add(fiche.id)
        ouvrage_ids_in_file.add(ouvrage.id)

        quantite = quantite_val if quantite_val is not None else (Decimal('1') if designation else None)
        prix_unitaire = prix_val if prix_val is not None else (
            designation.prix_unitaire if designation and designation.prix_unitaire is not None else Decimal('0')
        ) if designation else None

        resolved_rows.append({
            'ligne': ligne,
            'ouvrage_id': str(ouvrage.id),
            'ouvrage_nom': ouvrage.nom_ouvrage,
            'fiche_poste_id': str(fiche.id),
            'activite_nom': _fiche_label(fiche),
            'date': date_val.isoformat(),
            'date_affichee': date_val.strftime('%d/%m/%Y'),
            'remarque': remarque,
            'employe_id': str(employe.id) if employe else None,
            'employe_nom': employe_nom_affiche,
            'point_eau_id': str(point_eau.id) if point_eau else None,
            'point_eau_nom': (point_eau.nomabonne or point_eau.code_point_eau) if point_eau else '',
            'equipement_id': str(equipement.id) if equipement else None,
            'equipement_nom': equipement.nom if equipement else '',
            'composants_ids': composant_ids,
            'composants_noms': composant_noms,
            'designation_id': str(designation.id) if designation else None,
            'designation_nom': designation.nom if designation else '',
            'quantite': float(quantite) if quantite is not None else None,
            'prix_unitaire': float(prix_unitaire) if prix_unitaire is not None else None,
            'validee': validee,
            'statut': statut_key,
            'statut_display': STATUT_LABEL_BY_KEY.get(statut_key, '') if statut_key else '',
        })

    # Réparations déjà existantes en base pour les couples (ouvrage, activité) rencontrés dans le fichier
    existing_dates_by_pair = defaultdict(set)
    if fiche_ids_in_file and ouvrage_ids_in_file:
        existing_qs = Planning.objects.filter(
            ouvrage_id__in=ouvrage_ids_in_file,
            fiche_poste_id__in=fiche_ids_in_file,
        ).values_list('ouvrage_id', 'fiche_poste_id', 'date')
        for ouvrage_id, fiche_poste_id, planning_date in existing_qs:
            local_date = _planning_local_date(planning_date)
            if local_date:
                existing_dates_by_pair[(str(ouvrage_id), str(fiche_poste_id))].add(local_date)

    a_importer = []
    doublons_suspects = []
    doublons_stricts = []

    for row in resolved_rows:
        key = (row['ouvrage_id'], row['fiche_poste_id'])
        row_date = date_cls.fromisoformat(row['date'])
        existing_dates = existing_dates_by_pair[key]

        if row_date in existing_dates:
            doublons_stricts.append({
                'ligne': row['ligne'],
                'ouvrage_nom': row['ouvrage_nom'],
                'activite_nom': row['activite_nom'],
                'date_affichee': row['date_affichee'],
            })
        elif existing_dates:
            doublons_suspects.append({
                **row,
                'dates_existantes': sorted(d.strftime('%d/%m/%Y') for d in existing_dates),
            })
        else:
            a_importer.append(row)

        # Les dates déjà vues (dans ce fichier ou en base) alimentent la détection
        # pour les lignes suivantes du même fichier (mêmes règles que pour la base).
        existing_dates_by_pair[key].add(row_date)

    return invalides, a_importer, doublons_suspects, doublons_stricts


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_import_preview(request):
    """Analyse un fichier Excel de réparations : validation + détection des doublons (sans écriture en base)."""
    try:
        uploaded = request.FILES.get('file')
        if not uploaded:
            return JsonResponse({'success': False, 'error': "Aucun fichier fourni"}, status=400)
        if not uploaded.name.lower().endswith(('.xlsx', '.xlsm')):
            return JsonResponse({'success': False, 'error': "Le fichier doit être au format Excel (.xlsx)"}, status=400)

        try:
            wb = openpyxl.load_workbook(uploaded, data_only=True, read_only=True)
        except Exception:
            return JsonResponse({'success': False, 'error': "Fichier Excel illisible ou corrompu"}, status=400)

        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return JsonResponse({'success': False, 'error': "Fichier vide"}, status=400)

        header_index = {}
        for idx, cell in enumerate(header or []):
            label = _header_key(cell)
            if not label:
                continue
            if 'ouvrage' in label:
                header_index['ouvrage_id'] = idx
            elif 'activit' in label:
                header_index['activite'] = idx
            elif label == 'date':
                header_index['date'] = idx
            elif 'composant' in label:
                header_index['composants'] = idx
            elif 'quipement' in label:
                header_index['equipement_id'] = idx
            elif 'signation' in label:
                header_index['designation_id'] = idx
            elif 'eau' in label:
                header_index['point_eau_id'] = idx
            elif 'quantit' in label:
                header_index['quantite'] = idx
            elif 'prix' in label:
                header_index['prix_unitaire'] = idx
            elif 'desc' in label or 'observ' in label:
                header_index['description'] = idx
            elif 'valid' in label:
                header_index['validee'] = idx
            elif 'statut' in label:
                header_index['statut'] = idx
            elif 'respons' in label or 'techni' in label:
                header_index['responsable_id'] = idx
                header_index['responsable_nom'] = idx

        manquantes = [c for c in ('ouvrage_id', 'activite', 'date') if c not in header_index]
        if manquantes:
            return JsonResponse({
                'success': False,
                'error': "Colonnes obligatoires introuvables dans l'en-tête : " + ', '.join(
                    {'ouvrage_id': 'ID Ouvrage', 'activite': 'Activité', 'date': 'Date'}[c] for c in manquantes
                )
            }, status=400)

        def _cell(row_values, key):
            idx = header_index.get(key)
            if idx is None or idx >= len(row_values):
                return None
            return row_values[idx]

        champs = [
            'ouvrage_id', 'activite', 'date', 'point_eau_id', 'equipement_id',
            'composants', 'designation_id', 'quantite', 'prix_unitaire', 'description',
            'responsable_id', 'responsable_nom', 'validee', 'statut',
        ]

        rows_raw = []
        for excel_row_num, row_values in enumerate(rows_iter, start=2):
            if row_values is None or all(v in (None, '') for v in row_values):
                continue
            row_data = {'ligne': excel_row_num}
            for champ in champs:
                row_data[champ] = _cell(row_values, champ)
            rows_raw.append(row_data)
            if len(rows_raw) > MAX_LIGNES:
                return JsonResponse({
                    'success': False,
                    'error': f"Fichier trop volumineux (plus de {MAX_LIGNES} lignes)"
                }, status=400)

        if not rows_raw:
            return JsonResponse({'success': False, 'error': "Aucune ligne de données trouvée dans le fichier"}, status=400)

        _, ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        invalides, a_importer, doublons_suspects, doublons_stricts = _classify_rows(rows_raw, ouvrage_ids, request.user)

        return JsonResponse({
            'success': True,
            'resume': {
                'total_lignes': len(rows_raw),
                'importables_directement': len(a_importer),
                'doublons_stricts_ignores': len(doublons_stricts),
                'doublons_suspects': len(doublons_suspects),
                'lignes_invalides': len(invalides),
            },
            'a_importer': a_importer,
            'doublons_suspects': doublons_suspects,
            'doublons_stricts': doublons_stricts,
            'invalides': invalides,
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f"Erreur lors de l'analyse du fichier : {str(e)}",
            'traceback': traceback.format_exc()
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_import_confirm(request):
    """
    Importe réellement les réparations. Reçoit du client la liste des lignes
    validées lors de la prévisualisation (lignes sans conflit + doublons
    suspects cochés). Chaque ligne est revalidée avant insertion.
    """
    try:
        try:
            payload = json.loads(request.body or '{}')
        except (json.JSONDecodeError, TypeError):
            return JsonResponse({'success': False, 'error': "Corps JSON invalide"}, status=400)

        lignes = payload.get('lignes') or []
        if not isinstance(lignes, list) or not lignes:
            return JsonResponse({'success': False, 'error': "Aucune ligne à importer"}, status=400)
        if len(lignes) > MAX_LIGNES:
            return JsonResponse({'success': False, 'error': "Trop de lignes à importer en une fois"}, status=400)

        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        scoped_ouvrage_ids = {str(v) for v in scoped_ouvrage_ids}
        fiche_ids = {str(v) for v in get_reparations_fiches_queryset().values_list('id', flat=True)}
        employes_autorises_ids = {
            str(v) for v in _get_repair_modal_employees_queryset(request.user, ouvrage_ids=scoped_ouvrage_ids).values_list('id', flat=True)
        }

        crees = 0
        ignores_doublon = []
        erreurs = []

        with transaction.atomic():
            for entree in lignes:
                ligne_no = entree.get('ligne', '?')
                ouvrage_id = str(entree.get('ouvrage_id') or '')
                fiche_poste_id = str(entree.get('fiche_poste_id') or '')
                date_iso = entree.get('date')
                employe_id = entree.get('employe_id') or None
                remarque = entree.get('remarque') or ''
                point_eau_id = entree.get('point_eau_id') or None
                equipement_id = entree.get('equipement_id') or None
                composants_ids = [str(c) for c in (entree.get('composants_ids') or []) if c]
                designation_id = entree.get('designation_id') or None
                quantite = entree.get('quantite')
                prix_unitaire = entree.get('prix_unitaire')
                validee = bool(entree.get('validee'))
                statut = entree.get('statut') or None

                if ouvrage_id not in scoped_ouvrage_ids:
                    erreurs.append(f"Ligne {ligne_no} : ouvrage introuvable ou non accessible")
                    continue
                if fiche_poste_id not in fiche_ids:
                    erreurs.append(f"Ligne {ligne_no} : activité de réparation invalide")
                    continue
                if statut and statut not in STATUT_LABEL_BY_KEY:
                    erreurs.append(f"Ligne {ligne_no} : statut invalide")
                    continue

                row_date = _parse_date(date_iso)
                if not row_date:
                    erreurs.append(f"Ligne {ligne_no} : date invalide")
                    continue

                if employe_id and employe_id not in employes_autorises_ids:
                    employe_id = None  # Responsable non autorisé : on ignore l'assignation, la réparation reste importée

                if designation_id and not _is_uuid(designation_id):
                    erreurs.append(f"Ligne {ligne_no} : ID Désignation invalide (UUID attendu)")
                    continue
                if point_eau_id and not _is_uuid(point_eau_id):
                    erreurs.append(f"Ligne {ligne_no} : ID Point d'eau invalide (UUID attendu)")
                    continue
                if equipement_id and not _is_uuid(equipement_id):
                    erreurs.append(f"Ligne {ligne_no} : ID Équipement invalide (UUID attendu)")
                    continue

                if point_eau_id and equipement_id:
                    erreurs.append(f"Ligne {ligne_no} : un seul de point d'eau / équipement peut être ciblé")
                    continue
                if any(not _is_uuid(c) for c in composants_ids):
                    erreurs.append(f"Ligne {ligne_no} : composant invalide")
                    continue

                point_eau = None
                equipement = None
                designation = None
                if point_eau_id:
                    point_eau = PointEau.objects.filter(id=point_eau_id, ouvrage_id=ouvrage_id).first()
                    if not point_eau:
                        erreurs.append(f"Ligne {ligne_no} : point d'eau introuvable pour cet ouvrage")
                        continue
                if equipement_id:
                    equipement = EquipementConstitutif.objects.filter(id=equipement_id, ouvrage_id=ouvrage_id).first()
                    if not equipement:
                        erreurs.append(f"Ligne {ligne_no} : équipement introuvable pour cet ouvrage")
                        continue
                if designation_id:
                    designation = Designation.objects.filter(id=designation_id, actif=True).first()
                    if not designation:
                        erreurs.append(f"Ligne {ligne_no} : désignation introuvable")
                        continue

                # Revérification anti-doublon au moment de l'écriture (protège contre une
                # double soumission ou un import concurrent survenu entre la prévisualisation
                # et la confirmation). Comparaison en Python après conversion timezone
                # Madagascar (règle du projet : jamais de filtre de date en SQL).
                dates_existantes = {
                    _planning_local_date(d)
                    for d in Planning.objects.filter(
                        ouvrage_id=ouvrage_id,
                        fiche_poste_id=fiche_poste_id,
                    ).values_list('date', flat=True)
                }
                if row_date in dates_existantes:
                    ignores_doublon.append(f"Ligne {ligne_no} (déjà importée)")
                    continue

                try:
                    fiche = get_reparations_fiches_queryset().select_related('type_activites').get(id=fiche_poste_id)
                except Exception:
                    erreurs.append(f"Ligne {ligne_no} : activité introuvable")
                    continue

                try:
                    verifier_regles_fiche(fiche, equipement, point_eau, composants_ids)
                except ComposantsError as exc:
                    erreurs.append(f"Ligne {ligne_no} : {exc}")
                    continue

                # employe_id a déjà été vérifié contre employes_autorises_ids (actif + rôle autorisé)
                employe = None
                if employe_id:
                    employe = Employe.objects.filter(id=employe_id).first()

                # Mêmes valeurs par défaut que la création manuelle desktop (create_planning) :
                # c'est `estvalide` + `suiteadonner`, pas uniquement `statut`, qui pilotent le
                # statut/grading affiché sur cette page (cf. _statut_rep côté export Excel).
                planning = Planning(
                    date=_date_to_naive_midnight(row_date),
                    ouvrage_id=ouvrage_id,
                    fiche_poste=fiche,
                    type_activite_id=fiche.type_activites_id,
                    employe=employe,
                    remarque=remarque or None,
                    suiteadonner='À faire',
                    estvalide=validee,
                    statut=statut,
                )
                planning.save()

                try:
                    # Savepoint : une ligne refusée ne doit pas annuler tout l'import.
                    with transaction.atomic():
                        apply_repair_target(planning, equipement, point_eau, composants_ids)
                except ComposantsError as exc:
                    planning.delete()
                    erreurs.append(f"Ligne {ligne_no} : {exc}")
                    continue

                if designation:
                    try:
                        qte = Decimal(str(quantite)) if quantite is not None else Decimal('1')
                    except (InvalidOperation, ValueError):
                        qte = Decimal('1')
                    try:
                        pu = Decimal(str(prix_unitaire)) if prix_unitaire is not None else Decimal(str(designation.prix_unitaire or 0))
                    except (InvalidOperation, ValueError):
                        pu = Decimal(str(designation.prix_unitaire or 0))

                    PlanningDetail.objects.create(
                        planning=planning,
                        designation=designation,
                        quantity=float(qte),
                        pu=pu,
                        montant=qte * pu,
                        equipement_constitutif=equipement,
                        point_eau=point_eau,
                    )

                crees += 1

        return JsonResponse({
            'success': True,
            'message': f"{crees} réparation(s) importée(s) avec succès.",
            'crees': crees,
            'ignores_doublon': ignores_doublon,
            'erreurs': erreurs,
        })
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f"Erreur lors de l'import : {str(e)}",
            'traceback': traceback.format_exc()
        }, status=500)
