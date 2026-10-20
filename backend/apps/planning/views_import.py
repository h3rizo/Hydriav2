from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
import openpyxl
from openpyxl.styles import Font, PatternFill
import io
from .models import TypeActivite, Activite, ActivitePoste, TypeDesignation, Designation, DesignationCompte, DesignationTypeActivite
from apps.employees.models_evaluation import FichePoste
from apps.organization.models import Compte
import uuid


@csrf_exempt
@require_http_methods(["POST"])
def import_excel_view(request, model_name):
    """Vue pour importer des données Excel selon le modèle spécifié"""

    if 'file' not in request.FILES:
        return JsonResponse({'error': 'Aucun fichier fourni'}, status=400)

    file = request.FILES['file']

    if not file.name.endswith(('.xlsx', '.xls')):
        return JsonResponse({'error': 'Format de fichier non supporté. Utilisez .xlsx ou .xls'}, status=400)

    try:
        # Lire le fichier Excel
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active

        imported_count = 0
        errors = []

        if model_name == 'types_activites':
            imported_count, errors = import_types_activites(sheet)
        elif model_name == 'activites':
            imported_count, errors = import_activites(sheet)
        elif model_name == 'activites_postes':
            imported_count, errors = import_activites_postes(sheet)
        elif model_name == 'types_designations':
            imported_count, errors = import_types_designations(sheet)
        elif model_name == 'designations':
            imported_count, errors = import_designations(sheet)
        elif model_name == 'designations_comptes':
            imported_count, errors = import_designations_comptes(sheet)
        elif model_name == 'designation_type_activites':
            imported_count, errors = import_designation_type_activites(sheet)
        else:
            return JsonResponse({'error': f'Mod?le non support?: {model_name}'}, status=400)

        if errors:
            return JsonResponse({
                'imported_count': imported_count,
                'errors': errors,
                'status': 'partial_success'
            })

        return JsonResponse({
            'imported_count': imported_count,
            'message': f'{imported_count} éléments importés avec succès'
        })

    except Exception as e:
        return JsonResponse({'error': f'Erreur lors de l\'import: {str(e)}'}, status=500)


def import_types_activites(sheet):
    """Importer les types d'activités depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            nomtype = str(row[0]).strip()

            if not nomtype:
                continue

            # Vérifier si le type existe déjà
            if not TypeActivite.objects.filter(nomtype=nomtype).exists():
                TypeActivite.objects.create(nomtype=nomtype)
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Type "{nomtype}" existe déjà')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_activites(sheet):
    """Importer les activités depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            nom = str(row[0]).strip()
            description = str(row[1]).strip() if row[1] else ''
            type_activite_nom = str(row[2]).strip()
            est_corrective = bool(row[3]) if row[3] is not None else False

            if not nom or not type_activite_nom:
                errors.append(f'Ligne {row_num}: Nom et type d\'activité obligatoires')
                continue

            # Trouver le type d'activité
            try:
                type_activite = TypeActivite.objects.get(nomtype=type_activite_nom)
            except TypeActivite.DoesNotExist:
                errors.append(f'Ligne {row_num}: Type d\'activité "{type_activite_nom}" introuvable')
                continue

            # Vérifier si l'activité existe déjà
            if not Activite.objects.filter(nom=nom, type_activite=type_activite).exists():
                Activite.objects.create(
                    nom=nom,
                    description=description,
                    type_activite=type_activite,
                    est_corrective=est_corrective
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Activité "{nom}" existe déjà pour ce type')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_activites_postes(sheet):
    """Importer les activités de postes depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            nom = str(row[0]).strip()
            description = str(row[1]).strip() if row[1] else ''
            activite_nom = str(row[2]).strip()

            if not nom or not activite_nom:
                errors.append(f'Ligne {row_num}: Nom et activité obligatoires')
                continue

            # Trouver l'activité
            try:
                activite = Activite.objects.get(nom=activite_nom)
            except Activite.DoesNotExist:
                errors.append(f'Ligne {row_num}: Activité "{activite_nom}" introuvable')
                continue

            # Vérifier si l'activité de poste existe déjà
            if not ActivitePoste.objects.filter(nom=nom, activite=activite).exists():
                ActivitePoste.objects.create(
                    nom=nom,
                    description=description,
                    activite=activite
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Activité de poste "{nom}" existe déjà')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_types_designations(sheet):
    """Importer les types de désignations depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            nom = str(row[0]).strip()
            description = str(row[1]).strip() if row[1] else ''

            if not nom:
                continue

            # Vérifier si le type existe déjà
            if not TypeDesignation.objects.filter(nom=nom).exists():
                TypeDesignation.objects.create(
                    nom=nom,
                    description=description
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Type de désignation "{nom}" existe déjà')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_designations(sheet):
    """Importer les désignations depuis Excel
    Format: Nom, Description, Type de Désignation, Type Désignation Chantier, Unité, Prix, Date Prix,
            Poids Unitaire (kg), Seuil Minimum, Type Avance Caissière, Code Matériel, Prestation, Poste Acheteur
    """
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            nom = str(row[0]).strip()
            description = str(row[1]).strip() if row[1] else ''
            type_designation_nom = str(row[2]).strip()
            type_designation_chantier = str(row[3]).strip() if row[3] else None
            unite = str(row[4]).strip() if row[4] else None
            prix_unitaire = float(row[5]) if row[5] else None
            date_prix_str = str(row[6]).strip() if row[6] else None
            poids_unitaire = float(row[7]) if row[7] else None
            seuil_minimum = float(row[8]) if row[8] else None
            type_avance_caissiere = str(row[9]).strip() if row[9] else None
            code_materiel = str(row[10]).strip() if row[10] else None
            prestation = str(row[11]).strip().lower() == 'oui' if row[11] else False
            poste_acheteur_nom = str(row[12]).strip() if row[12] else None

            if not nom or not type_designation_nom:
                errors.append(f'Ligne {row_num}: Nom et type de désignation obligatoires')
                continue

            # Mapper les valeurs de type_designation_chantier
            type_chantier_map = {
                'Accessoires': 'accessoires',
                'Matériaux locaux': 'materiaux_locaux',
                'Matériaux de construction': 'materiaux_construction',
                'Tuyau': 'tuyau'
            }
            if type_designation_chantier and type_designation_chantier in type_chantier_map:
                type_designation_chantier = type_chantier_map[type_designation_chantier]

            # Trouver le type de désignation
            try:
                type_designation = TypeDesignation.objects.get(nom=type_designation_nom)
            except TypeDesignation.DoesNotExist:
                errors.append(f'Ligne {row_num}: Type de désignation "{type_designation_nom}" introuvable')
                continue

            # Gérer la date du prix
            date_prix = None
            if date_prix_str:
                from datetime import datetime
                try:
                    date_prix = datetime.strptime(date_prix_str, '%Y-%m-%d').date()
                except ValueError:
                    pass

            # Trouver le poste acheteur si spécifié
            poste_acheteur = None
            if poste_acheteur_nom:
                from apps.employees.models import Poste
                try:
                    poste_acheteur = Poste.objects.get(nom_poste=poste_acheteur_nom)
                except Poste.DoesNotExist:
                    errors.append(f'Ligne {row_num}: Poste "{poste_acheteur_nom}" introuvable (ignoré)')

            # Vérifier si la désignation existe déjà
            if not Designation.objects.filter(nom=nom, type_designation=type_designation).exists():
                Designation.objects.create(
                    nom=nom,
                    description=description,
                    type_designation=type_designation,
                    type_designation_chantier=type_designation_chantier,
                    unite=unite,
                    prix_unitaire=prix_unitaire,
                    date_prix=date_prix,
                    poids_unitaire=poids_unitaire,
                    seuil_minimum=seuil_minimum,
                    type_avance_caissiere=type_avance_caissiere,
                    code_materiel=code_materiel,
                    prestation=prestation,
                    poste_acheteur=poste_acheteur
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Désignation "{nom}" existe déjà pour ce type')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_designations_comptes(sheet):
    """Importer les liaisons désignations-comptes depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            designation_nom = str(row[0]).strip()
            nom_compte = str(row[1]).strip()
            notes = str(row[2]).strip() if row[2] else ''

            if not designation_nom or not nom_compte:
                errors.append(f'Ligne {row_num}: Désignation et nom compte obligatoires')
                continue

            # Trouver la désignation
            try:
                designation = Designation.objects.get(nom=designation_nom)
            except Designation.DoesNotExist:
                errors.append(f'Ligne {row_num}: Désignation "{designation_nom}" introuvable')
                continue

            # Trouver le compte par nom
            try:
                compte = Compte.objects.get(nom_compte=nom_compte)
            except Compte.DoesNotExist:
                errors.append(f'Ligne {row_num}: Compte "{nom_compte}" introuvable')
                continue

            # Vérifier si la liaison existe déjà
            if not DesignationCompte.objects.filter(designation=designation, compte=compte).exists():
                DesignationCompte.objects.create(
                    designation=designation,
                    compte=compte,
                    notes=notes,
                    is_active=True
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Liaison "{designation_nom}" - "{nom_compte}" existe déjà')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def import_designation_type_activites(sheet):
    """Importer les liaisons désignations-types activités depuis Excel"""
    imported_count = 0
    errors = []

    # Ignorer la ligne d'en-tête
    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:  # Si la première colonne est vide, arrêter
            break

        try:
            designation_nom = str(row[0]).strip()
            type_activite_nom = str(row[1]).strip()
            notes = str(row[2]).strip() if row[2] else ''

            if not designation_nom or not type_activite_nom:
                errors.append(f'Ligne {row_num}: Désignation et type d\'activité obligatoires')
                continue

            # Trouver la désignation
            try:
                designation = Designation.objects.get(nom=designation_nom)
            except Designation.DoesNotExist:
                errors.append(f'Ligne {row_num}: Désignation "{designation_nom}" introuvable')
                continue

            # Trouver le type d'activité
            try:
                type_activite = TypeActivite.objects.get(nomtype=type_activite_nom)
            except TypeActivite.DoesNotExist:
                errors.append(f'Ligne {row_num}: Type d\'activité "{type_activite_nom}" introuvable')
                continue

            # Vérifier si la liaison existe déjà
            if not DesignationTypeActivite.objects.filter(designation=designation, type_activite=type_activite).exists():
                DesignationTypeActivite.objects.create(
                    designation=designation,
                    type_activite=type_activite,
                    notes=notes,
                    is_active=True
                )
                imported_count += 1
            else:
                errors.append(f'Ligne {row_num}: Liaison "{designation_nom}" - "{type_activite_nom}" existe déjà')

        except Exception as e:
            errors.append(f'Ligne {row_num}: {str(e)}')

    return imported_count, errors


def download_excel_template(request, model_name):
    """Générer et télécharger un template Excel selon le modèle"""

    # Créer un nouveau workbook
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # Style pour les en-têtes
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    if model_name == 'types_activites':
        sheet.title = "Types d'Activités"
        headers = ['Nom du Type']
        examples = [['Maintenance'], ['Installation'], ['Réparation']]

    elif model_name == 'activites':
        sheet.title = "Activités"
        headers = ['Nom', 'Description', 'Type d\'Activité', 'Est Corrective']
        examples = [
            ['Vérification compteur', 'Contrôle du bon fonctionnement', 'Maintenance', False],
            ['Réparation fuite', 'Réparer une fuite détectée', 'Réparation', True],
            ['Installation nouveau compteur', 'Poser un compteur neuf', 'Installation', False]
        ]

    elif model_name == 'activites_postes':
        sheet.title = "Activités de Postes"
        headers = ['Nom', 'Description', 'Activité']
        examples = [
            ['Technicien maintenance', 'Technicien spécialisé maintenance', 'Vérification compteur'],
            ['Chef équipe réparation', 'Responsable équipe de réparation', 'Réparation fuite']
        ]

    elif model_name == 'types_designations':
        sheet.title = "Types de Désignations"
        headers = ['Nom', 'Description']
        examples = [
            ['Équipements', 'Types d\'équipements utilisés'],
            ['Personnel', 'Types de personnel affecté'],
            ['Matériel', 'Types de matériel nécessaire']
        ]

    elif model_name == 'designations':
        sheet.title = "Désignations"
        headers = [
            'Nom', 'Description', 'Type de Désignation', 'Type Désignation Chantier', 'Unité', 'Prix', 'Date Prix',
            'Poids Unitaire (kg)', 'Seuil Minimum', 'Type Avance Caissière',
            'Code Matériel', 'Prestation', 'Poste Acheteur'
        ]
        examples = [
            ['Compteur principal', 'Compteur d\'eau principal', 'Équipements', 'Accessoires', 'unité', '50000', '2025-01-01', '2.5', '10', 'Bon de livraison', 'MAT001', 'Non', 'Acheteur'],
            ['Tuyau PVC 32mm', 'Tuyau PVC diamètre 32mm', 'Matériel', 'Tuyau', 'mètre', '15000', '2025-01-01', '0.8', '50', 'Bon de livraison', 'TUY032', 'Non', 'Acheteur'],
            ['Sable', 'Sable pour construction', 'Matériel', 'Matériaux locaux', 'tonne', '45000', '2025-01-01', '1500', '5', 'Bon de livraison', 'SAB001', 'Non', 'Acheteur'],
            ['Ciment CPJ 42.5', 'Ciment Portland', 'Matériel', 'Matériaux de construction', 'sac', '28000', '2025-01-01', '50', '20', 'Bon de livraison', 'CIM001', 'Non', 'Acheteur']
        ]

    elif model_name == 'designations_comptes':
        sheet.title = "Liaisons Désignations-Comptes"
        headers = ['Désignation', 'Nom Compte', 'Notes']
        examples = [
            ['Compteur principal', 'Achats matériels', 'Liaison pour compteurs'],
            ['Tuyau PVC 32mm', 'Achats tuyauterie', 'Liaison pour tuyauterie'],
            ['Sable', 'Achats matériaux locaux', 'Liaison pour matériaux']
        ]
    elif model_name == 'designation_type_activites':
        sheet.title = "Liaisons Désignations-Types Activités"
        headers = ['Désignation', 'Type Activité', 'Notes']
        examples = [
            ['Compteur principal', 'Installation', 'Nécessite installation'],
            ['Tuyau PVC 32mm', 'Maintenance', 'Maintenance régulière'],
            ['Pompe hydraulique', 'Réparation', 'Réparations courantes']
        ]
    elif model_name == 'fiches_poste':
        sheet.title = "Fiches de Poste"
        headers = [
            'Tache', 'Description', "Type d'Activite", 'Outils', 'Ordre',
            'A evaluer', 'Dans planning', 'Preuve obligatoire', 'Type de preuve'
        ]
        examples = [
            ['Verifier compteur', 'Controle visuel et test', 'Maintenance', 'Cle, tournevis', 1, 'Oui', 'Oui', 'Non', 'piece_jointe'],
            ['Reparer fuite', "Reparation d'urgence", 'Reparation', 'Cle, joint', 2, 'Oui', 'Oui', 'Oui', 'evaluation']
        ]

    else:
        return JsonResponse({'error': f'Modèle non supporté: {model_name}'}, status=400)

    # Ajouter les en-têtes
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    # Ajouter les exemples
    for row, example in enumerate(examples, 2):
        for col, value in enumerate(example, 1):
            sheet.cell(row=row, column=col, value=value)

    # Ajuster la largeur des colonnes
    for column in sheet.columns:
        max_length = 0
        column = [cell for cell in column]
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column[0].column_letter].width = adjusted_width

    # Préparer la réponse
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=template_{model_name}.xlsx'

    return response


def export_excel_data(request, model_name):
    """Exporter les données vers Excel selon le modèle"""

    # Créer un nouveau workbook
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # Style pour les en-têtes
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    if model_name == 'types_activites':
        sheet.title = "Types d'Activités"
        headers = ['ID', 'Nom du Type', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        types = TypeActivite.objects.all().order_by('nomtype')
        for row, type_activite in enumerate(types, 2):
            sheet.cell(row=row, column=1, value=str(type_activite.id))
            sheet.cell(row=row, column=2, value=type_activite.nomtype)
            sheet.cell(row=row, column=3, value=type_activite.createdat.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'activites':
        sheet.title = "Activités"
        headers = ['ID', 'Nom', 'Description', 'Type d\'Activité', 'Est Corrective', 'Active', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        activites = Activite.objects.select_related('type_activite').order_by('type_activite__nomtype', 'nom')
        for row, activite in enumerate(activites, 2):
            sheet.cell(row=row, column=1, value=str(activite.id))
            sheet.cell(row=row, column=2, value=activite.nom)
            sheet.cell(row=row, column=3, value=activite.description or '')
            sheet.cell(row=row, column=4, value=activite.type_activite.nomtype if activite.type_activite else '')
            sheet.cell(row=row, column=5, value=activite.est_corrective)
            sheet.cell(row=row, column=6, value=activite.is_active)
            sheet.cell(row=row, column=7, value=activite.created_at.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'activites_postes':
        sheet.title = "Activités de Postes"
        headers = ['ID', 'Nom', 'Description', 'Activité', 'Active', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        activites_postes = ActivitePoste.objects.select_related('activite').order_by('activite__nom', 'nom')
        for row, activite_poste in enumerate(activites_postes, 2):
            sheet.cell(row=row, column=1, value=str(activite_poste.id))
            sheet.cell(row=row, column=2, value=activite_poste.nom)
            sheet.cell(row=row, column=3, value=activite_poste.description or '')
            sheet.cell(row=row, column=4, value=activite_poste.activite.nom if activite_poste.activite else '')
            sheet.cell(row=row, column=5, value=activite_poste.is_active)
            sheet.cell(row=row, column=6, value=activite_poste.created_at.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'types_designations':
        sheet.title = "Types de Désignations"
        headers = ['ID', 'Nom', 'Description', 'Actif', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        types_designations = TypeDesignation.objects.all().order_by('nom')
        for row, type_designation in enumerate(types_designations, 2):
            sheet.cell(row=row, column=1, value=str(type_designation.id))
            sheet.cell(row=row, column=2, value=type_designation.nom)
            sheet.cell(row=row, column=3, value=type_designation.description or '')
            sheet.cell(row=row, column=4, value=type_designation.is_active)
            sheet.cell(row=row, column=5, value=type_designation.created_at.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'designations':
        sheet.title = "Désignations"
        headers = [
            'ID', 'Nom', 'Description', 'Type de Désignation', 'Type Désignation Chantier', 'Unité', 'Prix', 'Date Prix',
            'Poids Unitaire (kg)', 'Seuil Minimum', 'Type Avance Caissière',
            'Code Matériel', 'Prestation', 'Poste Acheteur', 'Actif', 'Date Création'
        ]

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        designations = Designation.objects.select_related('type_designation', 'poste_acheteur').order_by('type_designation__nom', 'nom')
        for row, designation in enumerate(designations, 2):
            sheet.cell(row=row, column=1, value=str(designation.id))
            sheet.cell(row=row, column=2, value=designation.nom)
            sheet.cell(row=row, column=3, value=designation.description or '')
            sheet.cell(row=row, column=4, value=designation.type_designation.nom if designation.type_designation else '')
            sheet.cell(row=row, column=5, value=designation.get_type_designation_chantier_display() if designation.type_designation_chantier else '')
            sheet.cell(row=row, column=6, value=designation.unite or '')
            sheet.cell(row=row, column=7, value=float(designation.prix_unitaire) if designation.prix_unitaire else '')
            sheet.cell(row=row, column=8, value=designation.date_prix.strftime('%Y-%m-%d') if designation.date_prix else '')
            sheet.cell(row=row, column=9, value=float(designation.poids_unitaire) if designation.poids_unitaire else '')
            sheet.cell(row=row, column=10, value=float(designation.seuil_minimum) if designation.seuil_minimum else '')
            sheet.cell(row=row, column=11, value=designation.get_type_avance_caissiere_display() if designation.type_avance_caissiere else '')
            sheet.cell(row=row, column=12, value=designation.code_materiel or '')
            sheet.cell(row=row, column=13, value='Oui' if designation.prestation else 'Non')
            sheet.cell(row=row, column=14, value=designation.poste_acheteur.nom_poste if designation.poste_acheteur else '')
            sheet.cell(row=row, column=15, value=designation.actif)
            sheet.cell(row=row, column=16, value=designation.created_at.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'designations_comptes':
        sheet.title = "Liaisons Désignations-Comptes"
        headers = ['ID', 'Désignation', 'Nom Compte', 'Code Compte', 'Notes', 'Actif', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        liaisons = DesignationCompte.objects.select_related('designation', 'compte').order_by('designation__nom', 'compte__nom_compte')
        for row, liaison in enumerate(liaisons, 2):
            sheet.cell(row=row, column=1, value=str(liaison.id))
            sheet.cell(row=row, column=2, value=liaison.designation.nom if liaison.designation else '')
            sheet.cell(row=row, column=3, value=liaison.compte.nom_compte if liaison.compte else '')
            sheet.cell(row=row, column=4, value=liaison.compte.compte if liaison.compte else '')
            sheet.cell(row=row, column=5, value=liaison.notes or '')
            sheet.cell(row=row, column=6, value=liaison.is_active)
            sheet.cell(row=row, column=7, value=liaison.created_at.strftime('%Y-%m-%d %H:%M'))

    elif model_name == 'designation_type_activites':
        sheet.title = "Liaisons Désignations-Types Activités"
        headers = ['ID', 'Désignation', 'Type Activité', 'Notes', 'Actif', 'Date Création']

        # Ajouter les en-têtes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Ajouter les données
        liaisons = DesignationTypeActivite.objects.select_related('designation', 'type_activite').order_by('designation__nom', 'type_activite__nomtype')
        for row, liaison in enumerate(liaisons, 2):
            sheet.cell(row=row, column=1, value=str(liaison.id))
            sheet.cell(row=row, column=2, value=liaison.designation.nom if liaison.designation else '')
            sheet.cell(row=row, column=3, value=liaison.type_activite.nomtype if liaison.type_activite else '')
            sheet.cell(row=row, column=4, value=liaison.notes or '')
            sheet.cell(row=row, column=5, value=liaison.is_active)
            sheet.cell(row=row, column=6, value=liaison.created_at.strftime('%Y-%m-%d %H:%M'))
    elif model_name == 'fiches_poste':
        sheet.title = "Fiches de Poste"
        headers = [
            'ID', 'Tache', 'Description', "Type d'Activite", 'Poste concerne', 'Outils', 'Ordre',
            'A evaluer', 'Dans planning', 'Preuve obligatoire', 'Type de preuve', 'Date creation'
        ]

        # Ajouter les en-tetes
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        # Postes concernés : via ActivitePoste (même source que la page).
        # Une ligne par (fiche, poste) ; une fiche sans poste garde une ligne, poste vide.
        postes_par_fiche = {}
        for fiche_id, nom_poste in ActivitePoste.objects.filter(
            fiche_poste__isnull=False, poste__isnull=False,
        ).values_list('fiche_poste_id', 'poste__nom_poste').distinct():
            postes_par_fiche.setdefault(fiche_id, set()).add(nom_poste)

        fiches = FichePoste.objects.select_related('type_activites').order_by('ordre', 'tache')
        lignes = (
            (fiche, nom_poste)
            for fiche in fiches
            for nom_poste in (sorted(postes_par_fiche.get(fiche.id, ()), key=str.casefold) or [''])
        )
        for row, (fiche, nom_poste) in enumerate(lignes, 2):
            sheet.cell(row=row, column=1, value=str(fiche.id))
            sheet.cell(row=row, column=2, value=fiche.tache)
            sheet.cell(row=row, column=3, value=fiche.description or '')
            sheet.cell(row=row, column=4, value=fiche.type_activites.nomtype if fiche.type_activites else '')
            sheet.cell(row=row, column=5, value=nom_poste)
            sheet.cell(row=row, column=6, value=fiche.outils or '')
            sheet.cell(row=row, column=7, value=fiche.ordre)
            sheet.cell(row=row, column=8, value='Oui' if fiche.a_evaluer else 'Non')
            sheet.cell(row=row, column=9, value='Oui' if fiche.dans_planning else 'Non')
            sheet.cell(row=row, column=10, value='Oui' if fiche.preuve_obligatoire else 'Non')
            sheet.cell(row=row, column=11, value=fiche.preuve_type or '')
            sheet.cell(row=row, column=12, value=fiche.date_creation.strftime('%Y-%m-%d %H:%M') if fiche.date_creation else '')

    else:
        return JsonResponse({'error': f'Modèle non supporté: {model_name}'}, status=400)

    # Ajuster la largeur des colonnes
    for column in sheet.columns:
        max_length = 0
        column = [cell for cell in column]
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column[0].column_letter].width = adjusted_width

    # Préparer la réponse
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=export_{model_name}.xlsx'

    return response