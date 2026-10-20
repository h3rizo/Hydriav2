"""Rapport mensuel de pointage GPS par employé (page web + export Excel)."""
import re
from datetime import date

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render
from django.utils import timezone

from apps.employees.models import Employe

from . import services_pointage as svc

MOIS_FR = ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet',
           'août', 'septembre', 'octobre', 'novembre', 'décembre']

COLONNES = ['Date', 'Village(s)', 'Activité', 'Pointage Matinée', 'Pointage Après-midi',
            'Pointage Nuitée', 'Total pointages validés']


def _employes_autorises(user):
    """Soi-même + tous ses subordonnés ; le staff voit son organisation, le superuser tout."""
    if user.is_superuser:
        return Employe.objects.all()
    me = getattr(user, 'employe', None)
    if me is None:
        return Employe.objects.none()
    ids = {me.id, *svc.subordonnes_ids(me)}
    if user.is_staff and getattr(me, 'base_id', None):
        return Employe.objects.filter(base__organisation_id=me.base.organisation_id) | Employe.objects.filter(id__in=ids)
    return Employe.objects.filter(id__in=ids)


def grouper_employes(employes, selectionne_id=None):
    """Employés du sélecteur regroupés par base puis poste, sans les employés débauchés.

    Un employé débauché n'est gardé que s'il est celui actuellement affiché, pour que le
    sélecteur reste cohérent avec la page. Retourne [{'label': 'Base › Poste', 'employes': [...]}].
    """
    groupes = {}
    for e in employes:
        if e.date_debauchage is not None and str(e.id) != str(selectionne_id):
            continue
        base = e.base.nom_base if getattr(e, 'base', None) else None
        poste = e.poste.nom_poste if getattr(e, 'poste', None) else None
        groupes.setdefault((base, poste), []).append({'id': str(e.id), 'nom': svc._nom(e)})
    # « Sans base » / « Sans poste » à la fin, le reste par ordre alphabétique.
    ordre = sorted(groupes, key=lambda k: (k[0] is None, (k[0] or '').lower(), k[1] is None, (k[1] or '').lower()))
    return [
        {'label': f"{k[0] or 'Sans base'} › {k[1] or 'Sans poste'}",
         'employes': sorted(groupes[k], key=lambda x: x['nom'].lower())}
        for k in ordre
    ]


def _parse_mois(valeur):
    """'AAAA-MM' -> (annee, mois) ; défaut : mois courant à Madagascar."""
    if valeur and re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', valeur):
        return int(valeur[:4]), int(valeur[5:])
    aujourdhui = timezone.now().astimezone(svc.tz_madagascar).date()
    return aujourdhui.year, aujourdhui.month


def _decaler_mois(annee, mois, delta):
    index = annee * 12 + (mois - 1) + delta
    return index // 12, index % 12 + 1


def _totaux(lignes):
    total = {'valides': sum(l['total_valides'] for l in lignes)}
    for cle in ('matin', 'aprem', 'nuit'):
        total[cle + '_ok'] = sum(1 for l in lignes if l[cle] == 'OK')
        total[cle + '_non_ok'] = sum(1 for l in lignes if l[cle] == 'Non OK')
    total['prevus'] = sum(total[c + '_ok'] + total[c + '_non_ok'] for c in ('matin', 'aprem', 'nuit'))
    return total


def _resoudre(request):
    """Employé, mois et lignes du rapport, ou None si l'employé demandé n'est pas autorisé."""
    autorises = _employes_autorises(request.user)
    me = getattr(request.user, 'employe', None)
    demande = request.GET.get('employe')
    if demande:
        employe = autorises.filter(id=demande).first() if re.fullmatch(r'[0-9a-fA-F-]{32,36}', demande) else None
        if employe is None:
            return None
    else:
        employe = me or autorises.first()
    if employe is None:
        return None
    annee, mois = _parse_mois(request.GET.get('mois'))
    return autorises, employe, annee, mois, svc.rapport_mensuel(employe, annee, mois)


@login_required
def rapport_pointage_mensuel(request):
    resolu = _resoudre(request)
    if resolu is None:
        return HttpResponseForbidden("Vous n'avez pas accès au pointage de cet employé.")
    autorises, employe, annee, mois, lignes = resolu
    prec = _decaler_mois(annee, mois, -1)
    suiv = _decaler_mois(annee, mois, 1)
    return render(request, 'planning/rapport_pointage.html', {
        'employe': employe,
        'employe_nom': svc._nom(employe),
        'groupes': grouper_employes(autorises.select_related('base', 'poste'), employe.id),
        'lignes': lignes,
        'totaux': _totaux(lignes),
        'annee': annee,
        'mois': mois,
        'mois_value': f"{annee}-{mois:02d}",
        'mois_libelle': f"{MOIS_FR[mois - 1].capitalize()} {annee}",
        'mois_prec': f"{prec[0]}-{prec[1]:02d}",
        'mois_suiv': f"{suiv[0]}-{suiv[1]:02d}",
    })


def construire_classeur(lignes, employe_nom, annee, mois):
    """Classeur Excel du rapport (fonction pure, testable sans base)."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{annee}-{mois:02d}"
    fin = Side(style='thin', color='BBBBBB')
    bordure = Border(left=fin, right=fin, top=fin, bottom=fin)

    ws['A1'] = f"Pointage GPS - {employe_nom} - {MOIS_FR[mois - 1].capitalize()} {annee}"
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLONNES))

    for col, titre in enumerate(COLONNES, 1):
        cell = ws.cell(row=3, column=col, value=titre)
        cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = bordure

    vert = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    rouge = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    ligne = 4
    for l in lignes:
        valeurs = [l['date'].strftime('%d/%m/%Y'), l['villages'], l['activite'],
                   l['matin'], l['aprem'], l['nuit'], l['total_valides']]
        for col, valeur in enumerate(valeurs, 1):
            cell = ws.cell(row=ligne, column=col, value=valeur)
            cell.border = bordure
            if 4 <= col <= 6:
                cell.alignment = Alignment(horizontal='center')
                if valeur == 'OK':
                    cell.fill = vert
                elif valeur == 'Non OK':
                    cell.fill = rouge
            elif col == 7:
                cell.alignment = Alignment(horizontal='center')
        ligne += 1

    totaux = _totaux(lignes)
    total_ligne = ['TOTAL', '', '',
                   f"{totaux['matin_ok']} OK / {totaux['matin_non_ok']} non OK",
                   f"{totaux['aprem_ok']} OK / {totaux['aprem_non_ok']} non OK",
                   f"{totaux['nuit_ok']} OK / {totaux['nuit_non_ok']} non OK",
                   totaux['valides']]
    for col, valeur in enumerate(total_ligne, 1):
        cell = ws.cell(row=ligne, column=col, value=valeur)
        cell.font = Font(bold=True)
        cell.border = bordure
        if col >= 4:
            cell.alignment = Alignment(horizontal='center')

    for col, largeur in enumerate([12, 34, 34, 20, 20, 20, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = largeur
    ws.freeze_panes = 'A4'
    return wb


@login_required
def fiche_indemnite_pointage(request):
    """Fiche d'indemnité formelle (à imprimer) d'un employé pour un mois, d'après ses pointages validés."""
    resolu = _resoudre(request)
    if resolu is None:
        return HttpResponseForbidden("Vous n'avez pas accès au pointage de cet employé.")
    _autorises, employe, annee, mois, lignes = resolu
    employe = Employe.objects.select_related('poste', 'base__organisation').get(pk=employe.pk)
    poste = employe.poste
    fiche = svc.calculer_fiche_indemnite(
        lignes,
        poste.tarif_indemnite_demi_journee if poste else 0,
        poste.tarif_indemnite_nuitee if poste else 0,
    )
    superieur = employe.superieur_hierarchique
    return render(request, 'planning/fiche_indemnite_pointage.html', {
        'employe': employe,
        'employe_nom': employe.nom_complet or svc._nom(employe),
        'surnom': employe.surnom,
        'poste': poste,
        'base': employe.base,
        'organisation': employe.base.organisation if employe.base_id else None,
        'superieur_nom': (superieur.nom_complet or svc._nom(superieur)) if superieur else '',
        'fiche': fiche,
        'mois_libelle': f"{MOIS_FR[mois - 1].capitalize()} {annee}",
        'mois_value': f"{annee}-{mois:02d}",
        'date_edition': timezone.now().astimezone(svc.tz_madagascar).date(),
    })


@login_required
def rapport_pointage_mensuel_export(request):
    resolu = _resoudre(request)
    if resolu is None:
        return HttpResponseForbidden("Vous n'avez pas accès au pointage de cet employé.")
    _autorises, employe, annee, mois, lignes = resolu
    nom = svc._nom(employe)
    wb = construire_classeur(lignes, nom, annee, mois)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    fichier = re.sub(r'[^A-Za-z0-9_-]+', '_', nom) or 'employe'
    response['Content-Disposition'] = f'attachment; filename=Pointage_{fichier}_{annee}-{mois:02d}.xlsx'
    wb.save(response)
    return response
