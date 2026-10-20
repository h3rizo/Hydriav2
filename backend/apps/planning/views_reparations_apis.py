"""
APIs pour le suivi des reparations (nouveau workflow sans modele Reparation)
"""
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from django.db.models import Q
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import connection, transaction
import pytz
from decimal import Decimal
from dateutil.relativedelta import relativedelta
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from apps.planning.models import Planning, PlanningDetail, Activite, BonLivraison, Designation, PlanningEmploye
from apps.planning.views_reparations import (
    _filter_ouvrages_queryset_by_responsable_employe,
    _get_default_repair_organisation_id,
    _get_scoped_modalites_queryset,
    _get_user_scoped_ouvrage_ids,
)
from apps.core.views_mobile_auth import mobile_auth_required
from apps.planning.views import (
    _get_debit_status_label,
    _get_ouvrage_debit_context,
    DebitSourceMesure,
)
from apps.planning.repair_classification import get_reparations_fiches_queryset
from apps.planning.repair_workflow_grading import (
    get_repair_workflow_grading_from_plannings,
)
from apps.employees.models import Employe
from apps.organization.models import Ouvrage
from apps.organization.scoping import get_user_accessible_gestionnaire_ids
from apps.billing.models import PointEau, ModaliteGestion
from apps.infrastructure.models import EquipementConstitutif
from apps.geography.models import Commune
from apps.messaging.repair_sync import ensure_repair_conversation, close_repair_conversation
from apps.planning.doublons_reparations import calculer_doublons_securise
from apps.planning.reparation_composants import (
    ComposantsError, apply_repair_target, composants_disponibles,
    parse_composants_ids, serialize_repair_target, verifier_regles_fiche, regles_fiche,
)


def _repair_activity_mappings():
    fiches_reparation = list(get_reparations_fiches_queryset())

    fiche_poste_ids = [fiche.id for fiche in fiches_reparation]
    fiche_poste_to_activite_id = {}
    fiche_poste_to_activite_nom = {}

    for fiche in fiches_reparation:
        fiche_id = fiche.id
        type_name = getattr(getattr(fiche, 'type_activites', None), 'nomtype', '') or ''
        fiche_poste_to_activite_id[fiche_id] = fiche_id
        fiche_poste_to_activite_nom[fiche_id] = f"{type_name} {fiche.tache or ''}".strip()

    return (
        [],
        fiche_poste_ids,
        fiche_poste_to_activite_id,
        fiche_poste_to_activite_nom,
        {},
    )


def _as_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in ('1', 'true', 'vrai', 'yes', 'oui', 'on')
    return False


def _parse_planning_photo_payload(value):
    if not value:
        return []

    if isinstance(value, list):
        items = value
    else:
        try:
            items = json.loads(str(value))
        except (TypeError, ValueError, json.JSONDecodeError):
            return []

    if not isinstance(items, list):
        return []

    normalized = []
    for item in items:
        if not isinstance(item, dict):
            continue
        data_url = str(item.get('data') or '').strip()
        if not data_url.startswith('data:image/'):
            continue
        normalized.append({
            'id': item.get('id'),
            'name': str(item.get('name') or '').strip(),
            'data': data_url,
            'size': item.get('size') or 0,
            'originalSize': item.get('originalSize') or 0,
            'type': str(item.get('type') or '').strip(),
            'date': str(item.get('date') or '').strip(),
        })
    return normalized


def _serialize_planning_photo_payload(value):
    items = _parse_planning_photo_payload(value)
    return json.dumps(items, ensure_ascii=True) if items else ''


def _request_actor(request):
    user = getattr(request, "user", None)
    if not user or not getattr(user, "is_authenticated", False):
        return None
    return getattr(user, "employe", None)


def _normalize_participant_ids(values):
    if values is None:
        return []
    if isinstance(values, str):
        raw_values = values.split(',')
    else:
        raw_values = values

    participant_ids = []
    seen = set()
    for value in raw_values:
        candidate = str(value or '').strip()
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        participant_ids.append(candidate)
    return participant_ids


def _get_selected_organisation_id(request, gestionnaire_ids):
    selected = (request.GET.get('organisation') or '').strip()
    if selected and selected.lower() == 'all':
        selected = ''
    if selected:
        return selected

    return _get_default_repair_organisation_id(request.user, gestionnaire_ids)


def _get_scoped_ouvrage_ids_for_request(request):
    gestionnaire_ids = get_user_accessible_gestionnaire_ids(request.user)
    organisation_id = _get_selected_organisation_id(request, gestionnaire_ids)

    annee_filter = (request.GET.get('annee') or '').strip()
    try:
        annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
    except (TypeError, ValueError):
        annee_int = None

    ouvrage_ids = _get_user_scoped_ouvrage_ids(
        organisation_id,
        gestionnaire_ids,
        request.user,
        annee=annee_int,
    )
    return organisation_id, ouvrage_ids


@csrf_exempt
@require_http_methods(["POST"])
def api_supprimer_reparation(request, planning_id):
    """API pour supprimer une reparation (planning)

    Note: Utilise SQL brut car planning_employees et indemnites ont des planning_id
    en INTEGER alors que plannings.id est UUID, causant des erreurs ORM.
    """

    try:
        planning = Planning.objects.get(id=planning_id)

        # Sauvegarder les infos pour le message
        description = planning.fiche_poste.tache if planning.fiche_poste else "Reparation"
        planning_uuid = str(planning.id)
        close_repair_conversation(
            planning,
            actor=_request_actor(request),
            reason="Reparation supprimee. Discussion archivee automatiquement.",
            notify_completion=False,
        )

        # Utiliser SQL brut pour tout supprimer car Django ORM echoue
        # a cause des types incompatibles (INTEGER vs UUID)
        with connection.cursor() as cursor:
            # 1. Supprimer planningdetails
            cursor.execute(
                "DELETE FROM planningdetails WHERE planning_id = %s",
                [planning_uuid]
            )

            # 2. Recuperer reparations liees (si le modele Reparation existe)
            cursor.execute(
                "SELECT id FROM reparations WHERE planning_id = %s",
                [planning_uuid]
            )
            reparation_ids = [row[0] for row in cursor.fetchall()]

            if reparation_ids:
                # Supprimer historiques lies
                placeholders = ','.join(['%s'] * len(reparation_ids))
                cursor.execute(
                    f"DELETE FROM historiquereparations WHERE reparation_id IN ({placeholders})",
                    reparation_ids
                )

            # 3. Supprimer reparations
            cursor.execute(
                "DELETE FROM reparations WHERE planning_id = %s",
                [planning_uuid]
            )

            # 4. Supprimer les bons de livraison lies (FK bons_livraison.planning_id)
            cursor.execute(
                "DELETE FROM bons_livraison WHERE planning_id = %s",
                [planning_uuid]
            )

            # 4b. Supprimer les pieces jointes liees (FK planning_pieces.planning_id)
            cursor.execute(
                "DELETE FROM planning_pieces WHERE planning_id = %s",
                [planning_uuid]
            )

            # 5. Supprimer le planning
            cursor.execute(
                "DELETE FROM plannings WHERE id = %s",
                [planning_uuid]
            )

        return JsonResponse({
            'success': True,
            'message': f'{description} supprimee avec succes'
        })

    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Reparation non trouvee'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


def _point_eau_label(point_eau):
    """
    « [Type][IDXLS] - Nom d'abonné », le type étant collé à l'IDXLS (ex. « BP12 - RAKOTO Jean »).
    Repli sur le code du point d'eau sans nom d'abonné ; sans IDXLS, le type précède le nom.
    """
    nom = point_eau.nomabonne or point_eau.code_point_eau
    idxls = str(point_eau.idxls).strip() if getattr(point_eau, 'idxls', None) else ''
    type_obj = getattr(point_eau, 'type_point_eau', None)
    type_nom = (getattr(type_obj, 'nom_type', '') or '').strip()
    reference = f"{type_nom}{idxls}"
    if reference and nom:
        return f"{reference} - {nom}"
    return nom or reference or None


def _cible_affichage(planning, repli):
    """Nom et type de la cible d'une réparation : cible enregistrée, sinon déduite des détails/bons."""
    if planning.cible_equipement_id:
        return {'cible_type': 'equipement', 'cible_nom': planning.cible_equipement.nom}
    if planning.cible_point_eau_id:
        return {'cible_type': 'point_eau', 'cible_nom': _point_eau_label(planning.cible_point_eau)}
    if repli:
        noms = list(repli['equipements'].values()) + list(repli['points_eau'].values())
        noms = [n for n in noms if n]
        if noms:
            return {
                'cible_type': 'equipement' if repli['equipements'] else 'point_eau',
                'cible_nom': ' · '.join(dict.fromkeys(noms)),
            }
    return {'cible_type': None, 'cible_nom': None}


def _recurrence_composant(entry, tz_madagascar, today):
    """Compteur de récurrence (même ouvrage, même composant) pour l'affichage."""
    if not entry:
        return {'recurrence_count': 0, 'recurrence_derniere_date': None, 'recurrence_recente': False}
    derniere_date = entry['derniere_date']
    derniere_date_str = None
    recente = False
    if derniere_date:
        date_aware = derniere_date if derniere_date.tzinfo else timezone.make_aware(derniere_date)
        derniere_date_locale = date_aware.astimezone(tz_madagascar).date()
        derniere_date_str = derniere_date_locale.strftime('%Y-%m-%d')
        recente = derniere_date_locale >= (today - relativedelta(months=3))
    return {
        'recurrence_count': entry['count'],
        'recurrence_derniere_date': derniere_date_str,
        'recurrence_recente': recente,
    }


@mobile_auth_required
@require_http_methods(["GET"])
def api_reparations_list(request):
    """API pour lister les reparations avec filtres"""

    try:
        # Recuperer les filtres
        planning_id = request.GET.get('planning_id')
        ouvrage_id = request.GET.get('ouvrage_id')
        activite_id = (request.GET.get('activite_id') or request.GET.get('activite') or '').strip()
        annee_filter = request.GET.get('annee', 'all')
        limit_raw = request.GET.get('limit')
        offset_raw = request.GET.get('offset')
        summary_only = _as_bool(request.GET.get('summary') or request.GET.get('summary_only'))
        try:
            annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
        except (TypeError, ValueError):
            annee_int = None
        try:
            limit = int(limit_raw) if limit_raw else 100
        except (TypeError, ValueError):
            limit = 100
        limit = max(1, min(limit, 1000))
        try:
            offset = max(0, int(offset_raw)) if offset_raw else 0
        except (TypeError, ValueError):
            offset = 0

        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({
                'success': True,
                'reparations': [],
                'total': 0
            })

        # Recuperer les activites de reparation/rehabilitation
        repair_activities, fiche_poste_ids, fiche_poste_to_activite_id, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()

        # Requete de base
        plannings = Planning.objects.filter(
            ouvrage_id__in=scoped_ouvrage_ids,
        ).filter(
            Q(fiche_poste_id__in=fiche_poste_ids)
            | Q(remarque__startswith='[Saisie depuis la fiche ouvrage]')
        ).select_related(
            'fiche_poste',
            'fiche_poste__type_activites',
            'employe',
            'ouvrage',
            'cible_equipement',
            'cible_point_eau',
            'cible_point_eau__type_point_eau'
        ).prefetch_related(
            'planning_employees__employee',
            'composants_equipement',
            'composants_point_eau'
        ).order_by('-date')

        if planning_id:
            plannings = plannings.filter(id=planning_id)
        # Filtre par ouvrage
        if ouvrage_id:
            plannings = plannings.filter(ouvrage_id=ouvrage_id)
        if activite_id and activite_id != 'all':
            plannings = plannings.filter(fiche_poste_id=activite_id)
        if annee_int:
            plannings = plannings.filter(date__year=annee_int)

        # Construire la reponse
        data = []
        tz_madagascar = pytz.timezone('Indian/Antananarivo')
        today = timezone.now().astimezone(tz_madagascar).date()

        total_disponible = plannings.count()
        plannings_page = list(plannings[offset:offset + limit])
        doublons_par_planning = calculer_doublons_securise([p.id for p in plannings_page])

        # Récurrence : combien de fois ce composant a déjà été traité sur ce même ouvrage,
        # toutes activités de réparation confondues (utile pour repérer un composant qui
        # casse souvent). Cherché sur l'historique complet des ouvrages présents dans la page.
        recurrence_par_planning = {}
        ouvrage_ids_page = {p.ouvrage_id for p in plannings_page if p.ouvrage_id}
        if ouvrage_ids_page:
            historique_qs = Planning.objects.filter(
                ouvrage_id__in=ouvrage_ids_page,
            ).filter(
                Q(fiche_poste_id__in=fiche_poste_ids)
                | Q(remarque__startswith='[Saisie depuis la fiche ouvrage]')
            ).prefetch_related('composants_equipement', 'composants_point_eau').only(
                'id', 'ouvrage_id', 'date',
            )
            # index[(ouvrage_id, nom_composant)] = liste de (planning_id, date)
            occurrences_index = {}
            for h in historique_qs:
                noms = {c.nom_composant for c in h.composants_equipement.all()} \
                    | {c.nom_composant for c in h.composants_point_eau.all()}
                for nom in noms:
                    occurrences_index.setdefault((h.ouvrage_id, nom), []).append((h.id, h.date))

            for planning in plannings_page:
                if not planning.ouvrage_id:
                    continue
                noms_ligne = {c.nom_composant for c in planning.composants_equipement.all()} \
                    | {c.nom_composant for c in planning.composants_point_eau.all()}
                meilleur_count, meilleure_date = 0, None
                for nom in noms_ligne:
                    autres = [
                        (pid, d) for pid, d in occurrences_index.get((planning.ouvrage_id, nom), [])
                        if pid != planning.id
                    ]
                    if len(autres) > meilleur_count or (len(autres) == meilleur_count and autres and (
                            meilleure_date is None or max(d for _, d in autres) > meilleure_date)):
                        meilleur_count = len(autres)
                        meilleure_date = max(d for _, d in autres) if autres else None
                if meilleur_count:
                    recurrence_par_planning[planning.id] = {
                        'count': meilleur_count,
                        'derniere_date': meilleure_date,
                    }

        # Cible de repli (comme dans le formulaire « Modifier ») : quand la réparation n'a pas de
        # cible enregistrée, on la déduit de ses détails puis de ses bons de livraison.
        # Deux requêtes groupées, valables aussi en mode résumé (sans détails chargés).
        cibles_repli = {}

        def _ajouter_cible_repli(planning_id, equipement, point_eau):
            entry = cibles_repli.setdefault(str(planning_id), {'equipements': {}, 'points_eau': {}})
            if equipement is not None:
                entry['equipements'].setdefault(str(equipement.id), equipement.nom)
            if point_eau is not None:
                entry['points_eau'].setdefault(str(point_eau.id), _point_eau_label(point_eau))

        sans_cible_ids = [p.id for p in plannings_page
                          if not p.cible_equipement_id and not p.cible_point_eau_id]
        if sans_cible_ids:
            for d in PlanningDetail.objects.filter(planning_id__in=sans_cible_ids).select_related(
                    'equipement_constitutif', 'point_eau', 'point_eau__type_point_eau'):
                _ajouter_cible_repli(d.planning_id, d.equipement_constitutif, d.point_eau)
            for b in BonLivraison.objects.filter(planning_id__in=sans_cible_ids).select_related(
                    'equipement', 'point_eau', 'point_eau__type_point_eau'):
                _ajouter_cible_repli(b.planning_id, b.equipement, b.point_eau)

        for planning in plannings_page:
            if not planning.estvalide:
                ensure_repair_conversation(planning, actor=_request_actor(request), notify=False)

            # Convertir la date en timezone Madagascar
            date_prevue_str = None
            date_locale = None
            if planning.date:
                date_aware = planning.date if planning.date.tzinfo else timezone.make_aware(planning.date)
                date_madagascar = date_aware.astimezone(tz_madagascar)
                date_locale = date_madagascar.date()
                date_prevue_str = date_locale.strftime('%Y-%m-%d')

            # Determiner le statut en fonction de estvalide ET de la date
            if planning.estvalide:
                statut_val = 'validee'
                statut_display = 'Validée'
            else:
                # Si la date est dans le futur: planifiée
                # Si la date est passée ou aujourd'hui: en cours
                if date_locale and date_locale > today:
                    statut_val = 'planifiee'
                    statut_display = 'Planifiée'
                else:
                    statut_val = 'en_cours'
                    statut_display = 'En cours'

            planning_photos = _parse_planning_photo_payload(planning.photo)
            details = []
            if not summary_only:
                planning_details = PlanningDetail.objects.filter(
                    planning=planning
                ).select_related(
                    'designation',
                    'designation__type_designation',
                    'equipement_constitutif',
                    'point_eau'
                )

                for detail in planning_details:
                    # Construire les infos du detail
                    detail_info = {
                        'id': str(detail.id),
                        'designation': detail.designation.nom if detail.designation else 'Sans désignation',
                        'designation_id': str(detail.designation.id) if detail.designation else None,
                        'type_designation': detail.designation.type_designation.nom if detail.designation and detail.designation.type_designation else None,
                        'quantite': detail.quantity or 0,  # Corrigé: quantity au lieu de quantite
                        'prix_unitaire': float(detail.pu) if detail.pu else 0.0,
                        'montant': float(detail.montant) if detail.montant else 0.0,
                        'equipement': detail.equipement_constitutif.nom if detail.equipement_constitutif else None,
                        'equipement_id': str(detail.equipement_constitutif.id) if detail.equipement_constitutif else None,
                        'point_eau': detail.point_eau.nomabonne or detail.point_eau.code_point_eau if detail.point_eau else None,
                        'point_eau_id': str(detail.point_eau.id) if detail.point_eau else None,
                        'retour_caissiere': detail.retour,  # Corrigé: retour au lieu de retour_caissiere
                        'date_retour': detail.date_retour.strftime('%Y-%m-%d') if detail.date_retour else None,
                    }
                    details.append(detail_info)

            bons_livraison = []
            if not summary_only:
                try:
                    bons_qs = BonLivraison.objects.filter(planning=planning).select_related(
                        'designation',
                        'equipement',
                        'point_eau',
                    ).order_by('-created_at')
                    for bon in bons_qs:
                        if bon.retour:
                            bon_statut = 'retour'
                        elif bon.utilise:
                            bon_statut = 'livre'
                        else:
                            bon_statut = 'en_cours'

                        magasin_display = bon.magasin or ''
                        try:
                            from apps.planning.views_livraisons import _resolve_magasin_from_magasin_text
                            magasin_obj = _resolve_magasin_from_magasin_text(bon.magasin)
                            if magasin_obj:
                                magasin_display = magasin_obj.nom_magasin or ''
                                if getattr(magasin_obj, 'base', None) and magasin_obj.base.nom_base:
                                    magasin_display = f"{magasin_display} ({magasin_obj.base.nom_base})"
                        except Exception:
                            pass

                        bons_livraison.append({
                            'id': str(bon.id),
                            'designation': bon.designation.nom if bon.designation else None,
                            'designation_id': str(bon.designation.id) if bon.designation else None,
                            'quantite': bon.quantite or 0,
                            'magasin': bon.magasin or '',
                            'magasin_display': magasin_display,
                            'equipement': bon.equipement.nom if bon.equipement else None,
                            'equipement_id': str(bon.equipement.id) if bon.equipement else None,
                            'point_eau': bon.point_eau.nomabonne or bon.point_eau.code_point_eau if bon.point_eau else None,
                            'point_eau_id': str(bon.point_eau.id) if bon.point_eau else None,
                            'numero_bon': bon.numero_bon or '',
                            'utilise': bool(bon.utilise),
                            'statut': bon_statut,
                        })
                except Exception:
                    pass

            data.append({
                'id': str(planning.id),
                'description': planning.fiche_poste.tache if planning.fiche_poste else 'Reparation',
                'activite_nom': planning.fiche_poste.tache if planning.fiche_poste else None,
                'activite_id': str(planning.fiche_poste_id) if planning.fiche_poste_id else None,
                'fiche_poste_nom': planning.fiche_poste.tache if planning.fiche_poste else None,
                'fiche_poste_id': str(planning.fiche_poste_id) if planning.fiche_poste_id else None,
                'type_activite_nom': getattr(getattr(planning.fiche_poste, 'type_activites', None), 'nomtype', None),
                'remarque': planning.remarque or '',
                'bilan': planning.bilan or '',
                'suite_a_donner': planning.suiteadonner or '',
                'priorite_suivi': planning.priorite_suivi or '',
                'priorite_suivi_display': planning.get_priorite_suivi_display() if planning.priorite_suivi else '',
                'statut': statut_val,
                'statut_display': statut_display,
                'date_prevue': date_prevue_str,
                'assigne_a': {
                    'id': str(planning.employe.id),
                    'surnom': planning.employe.surnom
                } if planning.employe else None,
                'participants': [
                    {
                        'id': str(participant.employee.id),
                        'surnom': participant.employee.surnom,
                    }
                    for participant in planning.planning_employees.select_related('employee').all()
                ],
                'ouvrage_id': str(planning.ouvrage.id) if planning.ouvrage else None,
                'ouvrage_nom': planning.ouvrage.nom_ouvrage if planning.ouvrage else 'Sans ouvrage',
                'point_eau_id': (str(planning.cible_point_eau_id) if planning.cible_point_eau_id else None)
                    or next((d.get('point_eau_id') for d in details if d.get('point_eau_id')), None)
                    or next((b.get('point_eau_id') for b in bons_livraison if b.get('point_eau_id')), None),
                'equipement_id': (str(planning.cible_equipement_id) if planning.cible_equipement_id else None)
                    or next((d.get('equipement_id') for d in details if d.get('equipement_id')), None)
                    or next((b.get('equipement_id') for b in bons_livraison if b.get('equipement_id')), None),
                **serialize_repair_target(planning),
                **_cible_affichage(planning, cibles_repli.get(str(planning.id))),
                **_recurrence_composant(recurrence_par_planning.get(planning.id), tz_madagascar, today),
                'doublons': doublons_par_planning.get(str(planning.id)),
                'summary_only': bool(summary_only),
                'photos': [] if summary_only else planning_photos,
                'photo_count': len(planning_photos),
                'details': details,
                'bons_livraison': bons_livraison,
                'total_details': len(details),
                'total_bons_livraison': len(bons_livraison),
            })

        return JsonResponse({
            'success': True,
            'reparations': data,
            'total': total_disponible,
            'offset': offset,
            'limit': limit,
            'has_more': offset + len(data) < total_disponible,
        })

    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_modifier_reparation(request, planning_id):
    """API pour modifier l'activité, l'employé, la date et la remarque d'une réparation"""
    try:
        import json
        import pytz
        from datetime import datetime
        data = json.loads(request.body)

        planning = Planning.objects.get(id=planning_id)

        # Récupérer les données
        activite_id = data.get('activite_id')
        fiche_poste_id = data.get('fiche_poste_id')
        employe_id = data.get('employe_id')
        date_prevue_str = data.get('date_prevue')
        remarque = data.get('remarque', '')
        detail_payload = data.get('detail') or {}
        details_payload = data.get('details')
        bons_payload = data.get('bons_livraison')
        repair_point_eau_id = (data.get('point_eau_id') or '').strip()
        repair_equipement_id = (data.get('equipement_id') or '').strip()
        has_photos_payload = ('photos' in data) or ('photo' in data)
        photos_payload = data.get('photos', data.get('photo'))
        participant_ids = _normalize_participant_ids(
            data.get('participants')
            or data.get('participants_ids')
            or data.get('participants[]')
        )
        suite_a_donner = data.get('suite_a_donner', data.get('suiteadonner', ''))
        priorite_suivi = (data.get('priorite_suivi') or '').strip().lower() or None
        has_validation_flag = ('validation' in data) or ('estvalide' in data)
        validation_requested = _as_bool(data.get('validation') or data.get('estvalide'))

        valid_priorites = {'basse', 'normale', 'haute', 'critique'}
        if priorite_suivi and priorite_suivi not in valid_priorites:
            return JsonResponse({
                'success': False,
                'error': 'Priorité invalide'
            }, status=400)
        if repair_point_eau_id and repair_equipement_id:
            return JsonResponse({
                'success': False,
                'error': "La réparation ne peut pas cibler à la fois un point d'eau et un équipement"
            }, status=400)

        repair_point_eau = None
        repair_equipement = None
        if repair_point_eau_id:
            try:
                repair_point_eau = PointEau.objects.get(id=repair_point_eau_id)
            except PointEau.DoesNotExist:
                return JsonResponse({
                    'success': False,
                    'error': "Point d'eau introuvable"
                }, status=404)
        if repair_equipement_id:
            try:
                repair_equipement = EquipementConstitutif.objects.get(id=repair_equipement_id)
            except EquipementConstitutif.DoesNotExist:
                return JsonResponse({
                    'success': False,
                    'error': 'Equipement introuvable'
                }, status=404)

        # Valider la fiche de poste ou l'ancien activite_id de compatibilite
        if fiche_poste_id:
            from apps.employees.models import FichePoste
            try:
                fiche_poste = FichePoste.objects.select_related('type_activites').get(id=fiche_poste_id)
            except FichePoste.DoesNotExist:
                return JsonResponse({
                    'success': False,
                    'error': 'Fiche de poste introuvable'
                }, status=404)

            planning.fiche_poste = fiche_poste
            planning.type_activite = fiche_poste.type_activites
        elif activite_id:
            from apps.employees.models import FichePoste
            try:
                fiche_poste = FichePoste.objects.select_related('type_activites').get(id=activite_id)
                planning.fiche_poste = fiche_poste
                planning.type_activite = fiche_poste.type_activites
            except FichePoste.DoesNotExist:
                try:
                    activite = Activite.objects.select_related('type_activite').get(id=activite_id)
                    planning.type_activite = activite.type_activite
                except Activite.DoesNotExist:
                    return JsonResponse({
                        'success': False,
                        'error': 'Fiche de poste ou activite introuvable'
                    }, status=404)

                _, _, _, _, activite_to_fiche_poste_id = _repair_activity_mappings()
                fiche_poste_id = (
                    activite_to_fiche_poste_id.get(activite.id)
                    or activite_to_fiche_poste_id.get(str(activite.id))
                )
                if fiche_poste_id:
                    try:
                        fiche_poste = FichePoste.objects.get(id=fiche_poste_id)
                        planning.fiche_poste = fiche_poste
                    except FichePoste.DoesNotExist:
                        pass

        # Valider et mettre à  jour l'employé
        if employe_id:
            try:
                employe = Employe.objects.get(id=employe_id)
                planning.employe = employe
            except Employe.DoesNotExist:
                return JsonResponse({
                    'success': False,
                    'error': 'Employé introuvable'
                }, status=404)

        # Valider et mettre à  jour la date
        if date_prevue_str:
            try:
                # Parser la date (format YYYY-MM-DD)
                date_naive = datetime.strptime(date_prevue_str, '%Y-%m-%d')
                # Localiser en timezone Madagascar
                tz_madagascar = pytz.timezone('Indian/Antananarivo')
                date_aware = tz_madagascar.localize(date_naive)
                planning.date = date_aware
            except ValueError:
                return JsonResponse({
                    'success': False,
                    'error': 'Format de date invalide (YYYY-MM-DD attendu)'
                }, status=400)

        # Mettre à  jour la remarque
        # Règles de l'activité (cible / composants obligatoires ou restreints), avant toute écriture
        if 'point_eau_id' in data or 'equipement_id' in data:
            composants_recus = parse_composants_ids(data)
            if composants_recus is None:
                cible_identique = (
                    (planning.cible_equipement_id or None) == (repair_equipement.id if repair_equipement else None)
                    and (planning.cible_point_eau_id or None) == (repair_point_eau.id if repair_point_eau else None)
                )
                composants_recus = (
                    [str(c.id) for c in planning.composants_equipement.all()]
                    + [str(c.id) for c in planning.composants_point_eau.all()]
                ) if cible_identique else []
            try:
                verifier_regles_fiche(planning.fiche_poste, repair_equipement, repair_point_eau, composants_recus)
            except ComposantsError as exc:
                return JsonResponse({'success': False, 'error': str(exc)}, status=400)

        planning.remarque = remarque
        planning.suiteadonner = suite_a_donner or ''
        planning.priorite_suivi = priorite_suivi
        if has_photos_payload:
            planning.photo = _serialize_planning_photo_payload(photos_payload)
        if has_validation_flag:
            planning.estvalide = validation_requested
        planning.save()

        if 'point_eau_id' in data or 'equipement_id' in data:
            try:
                apply_repair_target(
                    planning, repair_equipement, repair_point_eau, parse_composants_ids(data)
                )
            except ComposantsError as exc:
                return JsonResponse({'success': False, 'error': str(exc)}, status=400)

        secondary_participant_ids = []
        owner_id = str(planning.employe_id or '').strip()
        seen_participants = set()
        for participant_id in participant_ids:
            participant_str = str(participant_id or '').strip()
            if not participant_str or participant_str == owner_id or participant_str in seen_participants:
                continue
            seen_participants.add(participant_str)
            secondary_participant_ids.append(participant_str)

        current_participants = {
            str(row.employee_id): row
            for row in PlanningEmploye.objects.filter(planning=planning)
        }
        for participant_id, participant_row in current_participants.items():
            if participant_id not in secondary_participant_ids:
                participant_row.delete()
        for participant_id in secondary_participant_ids:
            PlanningEmploye.objects.get_or_create(
                planning=planning,
                employee_id=participant_id,
            )

        actor = _request_actor(request)
        if planning.estvalide:
            close_repair_conversation(
                planning,
                actor=actor,
                reason="Reparation terminee depuis le suivi reparations.",
            )
        else:
            ensure_repair_conversation(planning, actor=actor, notify=True)

        details_explicit = isinstance(details_payload, list)
        if details_explicit:
            detail_rows = details_payload
        elif detail_payload:
            detail_rows = [detail_payload]
        else:
            detail_rows = []

        has_material_rows = False
        has_bon_rows = False
        if detail_rows or details_explicit:
            existing_details = {
                str(d.id): d for d in PlanningDetail.objects.filter(planning=planning)
            }
            seen_detail_ids = set()

            for row in detail_rows:
                designation_id = row.get('designation_id')
                if not designation_id:
                    continue

                try:
                    designation = Designation.objects.get(id=designation_id)
                except Designation.DoesNotExist:
                    return JsonResponse({
                        'success': False,
                        'error': 'Designation detail introuvable'
                    }, status=404)

                try:
                    quantity = float(row.get('quantity') or 1)
                except (TypeError, ValueError):
                    quantity = 1

                try:
                    pu = Decimal(str(designation.prix_unitaire or 0))
                except Exception:
                    pu = Decimal('0')

                detail_id = str(row.get('detail_id') or '').strip()

                detail = existing_details.get(detail_id) if detail_id else None
                if detail:
                    detail.designation = designation
                    detail.quantity = quantity
                    detail.pu = pu
                    detail.equipement_constitutif = repair_equipement
                    detail.point_eau = repair_point_eau
                    detail.save()
                    seen_detail_ids.add(str(detail.id))
                else:
                    detail = PlanningDetail.objects.create(
                        planning=planning,
                        designation=designation,
                        quantity=quantity,
                        pu=pu,
                        montant=Decimal('0'),
                        equipement_constitutif=repair_equipement,
                        point_eau=repair_point_eau,
                    )
                    seen_detail_ids.add(str(detail.id))
                    has_material_rows = True

                if detail.designation_id:
                    has_material_rows = True

            # Supprimer les details retires dans le modal
            for current_id, current_detail in existing_details.items():
                if current_id not in seen_detail_ids:
                    current_detail.delete()

        if isinstance(bons_payload, list):
            has_bon_rows = bool(bons_payload)
            existing_bons = {
                str(b.id): b for b in BonLivraison.objects.filter(planning=planning)
            }
            seen_bon_ids = set()

            for row in bons_payload:
                designation_id = row.get('designation_id')
                if not designation_id:
                    continue

                try:
                    designation = Designation.objects.get(id=designation_id)
                except Designation.DoesNotExist:
                    return JsonResponse({
                        'success': False,
                        'error': 'Designation bon de livraison introuvable'
                    }, status=404)

                try:
                    quantite = float(row.get('quantite') or row.get('quantity') or 1)
                except (TypeError, ValueError):
                    quantite = 1
                magasin = (row.get('magasin') or '').strip() or None
                bon_id = str(row.get('bon_id') or row.get('id') or '').strip()
                bon = existing_bons.get(bon_id) if bon_id else None
                if bon:
                    bon.designation = designation
                    bon.quantite = quantite
                    bon.magasin = magasin
                    bon.equipement = repair_equipement
                    bon.point_eau = repair_point_eau
                    bon.save()
                    seen_bon_ids.add(str(bon.id))
                else:
                    bon = BonLivraison.objects.create(
                        planning=planning,
                        designation=designation,
                        quantite=quantite,
                        magasin=magasin,
                        equipement=repair_equipement,
                        point_eau=repair_point_eau,
                    )
                    seen_bon_ids.add(str(bon.id))

            for current_id, current_bon in existing_bons.items():
                if current_id not in seen_bon_ids:
                    current_bon.delete()

        return JsonResponse({
            'success': True,
            'message': 'Réparation modifiée avec succès'
        })
    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Planning introuvable'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


BULK_UPDATE_MAX = 500


def _peut_supprimer_reparations(user):
    from apps.employees.permissions import is_directeur_poste
    return bool(getattr(user, 'is_superuser', False) or is_directeur_poste(user))


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_bulk_delete(request):
    """
    Suppression (définitive) de réparations depuis la vue consolidée.
    Réservée au directeur (ou superuser) ; limitée au périmètre de l'utilisateur.
    Corps : {"ids": [...]}
    """
    if not _peut_supprimer_reparations(request.user):
        return JsonResponse(
            {'success': False, 'error': 'Suppression réservée au directeur'}, status=403)

    try:
        payload = json.loads(request.body or '{}')
    except (TypeError, ValueError):
        return JsonResponse({'success': False, 'error': 'Corps JSON invalide'}, status=400)
    ids = payload.get('ids') if isinstance(payload, dict) else None
    if not isinstance(ids, list) or not ids:
        return JsonResponse({'success': False, 'error': 'ids est requis'}, status=400)
    ids = list(dict.fromkeys(str(i) for i in ids))
    if len(ids) > BULK_UPDATE_MAX:
        return JsonResponse(
            {'success': False, 'error': f'Maximum {BULK_UPDATE_MAX} réparations par opération'}, status=400)

    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        _, fiche_poste_ids, _, _, _ = _repair_activity_mappings()
        try:
            plannings = list(
                Planning.objects.filter(id__in=ids, ouvrage_id__in=scoped_ouvrage_ids)
                .filter(Q(fiche_poste_id__in=fiche_poste_ids)
                        | Q(remarque__startswith='[Saisie depuis la fiche ouvrage]'))
            )
        except (DjangoValidationError, ValueError):
            return JsonResponse({'success': False, 'error': 'Identifiant invalide'}, status=400)

        actor = _request_actor(request)
        for planning in plannings:
            close_repair_conversation(
                planning, actor=actor,
                reason="Reparation supprimee. Discussion archivee automatiquement.",
                notify_completion=False,
            )

        # Même suppression en SQL brut que api_supprimer_reparation (types INTEGER/UUID incompatibles)
        deleted_ids = []
        with transaction.atomic():
            with connection.cursor() as cursor:
                for planning in plannings:
                    pid = str(planning.id)
                    cursor.execute("DELETE FROM planningdetails WHERE planning_id = %s", [pid])
                    cursor.execute("SELECT id FROM reparations WHERE planning_id = %s", [pid])
                    reparation_ids = [row[0] for row in cursor.fetchall()]
                    if reparation_ids:
                        placeholders = ','.join(['%s'] * len(reparation_ids))
                        cursor.execute(
                            f"DELETE FROM historiquereparations WHERE reparation_id IN ({placeholders})",
                            reparation_ids)
                    cursor.execute("DELETE FROM reparations WHERE planning_id = %s", [pid])
                    cursor.execute("DELETE FROM bons_livraison WHERE planning_id = %s", [pid])
                    cursor.execute("DELETE FROM planning_pieces WHERE planning_id = %s", [pid])
                    cursor.execute("DELETE FROM plannings WHERE id = %s", [pid])
                    deleted_ids.append(pid)

        return JsonResponse({
            'success': True,
            'deleted': len(deleted_ids),
            'deleted_ids': deleted_ids,
            'ignored': [i for i in ids if i not in set(deleted_ids)],
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_bulk_update(request):
    """
    Mise à jour groupée (vue consolidée) : date, activité (fiche de poste) et/ou validation.

    Corps : {"ids": [...], "changes": {"date_prevue": "YYYY-MM-DD", "fiche_poste_id": "...", "estvalide": bool}}
    Seuls les champs présents dans "changes" sont modifiés ; remarque, participants, détails
    et cible ne sont pas touchés.
    """
    from datetime import datetime
    from apps.employees.models import FichePoste

    try:
        payload = json.loads(request.body or '{}')
    except (TypeError, ValueError):
        return JsonResponse({'success': False, 'error': 'Corps JSON invalide'}, status=400)

    ids = payload.get('ids') if isinstance(payload, dict) else None
    changes = payload.get('changes') if isinstance(payload, dict) else None
    if not isinstance(ids, list) or not ids or not isinstance(changes, dict):
        return JsonResponse({'success': False, 'error': 'ids et changes sont requis'}, status=400)
    ids = list(dict.fromkeys(str(i) for i in ids))
    if len(ids) > BULK_UPDATE_MAX:
        return JsonResponse(
            {'success': False, 'error': f'Maximum {BULK_UPDATE_MAX} réparations par opération'}, status=400)

    new_date = None
    if changes.get('date_prevue'):
        try:
            naive = datetime.strptime(str(changes['date_prevue']), '%Y-%m-%d')
        except ValueError:
            return JsonResponse(
                {'success': False, 'error': 'Format de date invalide (YYYY-MM-DD attendu)'}, status=400)
        new_date = pytz.timezone('Indian/Antananarivo').localize(naive)

    new_fiche = None
    if changes.get('fiche_poste_id'):
        new_fiche = get_reparations_fiches_queryset().select_related('type_activites').filter(
            id=changes['fiche_poste_id']).first()
        if not new_fiche:
            return JsonResponse(
                {'success': False, 'error': "Activité introuvable parmi les activités de réparation"}, status=400)

    has_validation = 'estvalide' in changes
    validation = _as_bool(changes.get('estvalide')) if has_validation else None

    if new_date is None and new_fiche is None and not has_validation:
        return JsonResponse({'success': False, 'error': 'Aucune modification demandée'}, status=400)

    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        _, fiche_poste_ids, _, _, _ = _repair_activity_mappings()
        try:
            plannings = list(
                Planning.objects.filter(id__in=ids, ouvrage_id__in=scoped_ouvrage_ids)
                .filter(Q(fiche_poste_id__in=fiche_poste_ids)
                        | Q(remarque__startswith='[Saisie depuis la fiche ouvrage]'))
                .select_related('cible_equipement', 'cible_point_eau')
                .prefetch_related('composants_equipement', 'composants_point_eau')
            )
        except (DjangoValidationError, ValueError):
            return JsonResponse({'success': False, 'error': 'Identifiant invalide'}, status=400)

        actor = _request_actor(request)

        # Changer d'activité ne doit pas contourner ses règles (cible / composants) :
        # les réparations dont la cible actuelle ne convient pas à la nouvelle activité sont refusées.
        refused = []
        if new_fiche is not None:
            accepted = []
            for planning in plannings:
                planning_composants = (
                    [str(c.id) for c in planning.composants_equipement.all()]
                    + [str(c.id) for c in planning.composants_point_eau.all()]
                )
                try:
                    verifier_regles_fiche(
                        new_fiche, planning.cible_equipement, planning.cible_point_eau, planning_composants)
                    accepted.append(planning)
                except ComposantsError as exc:
                    refused.append({'id': str(planning.id), 'reason': str(exc)})
            plannings = accepted

        updated = []
        with transaction.atomic():
            for planning in plannings:
                fields = ['updated_at']
                if new_date is not None:
                    planning.date = new_date
                    fields.append('date')
                if new_fiche is not None:
                    planning.fiche_poste = new_fiche
                    planning.type_activite = new_fiche.type_activites
                    fields += ['fiche_poste', 'type_activite']
                if has_validation:
                    planning.estvalide = validation
                    planning.statut = 'termine' if validation else 'en_cours'
                    planning.suiteadonner = 'Termine' if validation else 'A suivre'
                    fields += ['estvalide', 'statut', 'suiteadonner']
                planning.save(update_fields=fields)
                updated.append(planning)

        # Conversations de suivi : hors transaction, sans notifier en masse
        if has_validation:
            for planning in updated:
                if validation:
                    close_repair_conversation(
                        planning, actor=actor,
                        reason="Reparation marquee terminee depuis la vue consolidee.")
                else:
                    ensure_repair_conversation(planning, actor=actor, notify=False)

        updated_ids = {str(p.id) for p in updated}
        refused_ids = {r['id'] for r in refused}
        return JsonResponse({
            'success': True,
            'updated': len(updated),
            'ignored': [i for i in ids if i not in updated_ids and i not in refused_ids],
            'refused': refused,
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


# --------------------------------------------------------------------------- #
# Composants en masse (vue consolidée)
# --------------------------------------------------------------------------- #

def _fold(text):
    import unicodedata
    text = unicodedata.normalize('NFKD', text or '')
    return ''.join(c for c in text if not unicodedata.combining(c)).strip().lower()


def _cibles_deduites(planning_ids):
    """{planning_id: {'eq': {id: objet}, 'pe': {id: objet}}} déduit des détails puis des bons de livraison."""
    out = {}

    def add(pid, equipement, point_eau):
        entry = out.setdefault(str(pid), {'eq': {}, 'pe': {}})
        if equipement is not None:
            entry['eq'].setdefault(equipement.id, equipement)
        if point_eau is not None:
            entry['pe'].setdefault(point_eau.id, point_eau)

    for d in PlanningDetail.objects.filter(planning_id__in=planning_ids).select_related(
            'equipement_constitutif', 'point_eau'):
        add(d.planning_id, d.equipement_constitutif, d.point_eau)
    for b in BonLivraison.objects.filter(planning_id__in=planning_ids).select_related('equipement', 'point_eau'):
        add(b.planning_id, b.equipement, b.point_eau)
    return out


def _preparer_selection_composants(request, ids):
    """
    Plannings de la sélection (dans le périmètre de l'utilisateur) avec leur cible et le catalogue
    de composants du type de cette cible.

    Retourne (lignes, sans_cible) : lignes = [{'planning', 'equipement', 'point_eau', 'type_nom',
    'catalogue': {nom_normalisé: (id, nom)}}] ; sans_cible = [{'id', 'reason'}].
    La cible enregistrée est prioritaire ; à défaut, on prend celle des détails / bons si elle est sans ambiguïté.
    """
    from apps.inspections.models import ComposantTypeEquipement, ComposantTypePointEau

    _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
    _, fiche_poste_ids, _, _, _ = _repair_activity_mappings()
    plannings = list(
        Planning.objects.filter(id__in=ids, ouvrage_id__in=scoped_ouvrage_ids)
        .filter(Q(fiche_poste_id__in=fiche_poste_ids) | Q(remarque__startswith='[Saisie depuis la fiche ouvrage]'))
        .select_related('fiche_poste', 'cible_equipement__type_equipement', 'cible_point_eau__type_point_eau')
        .prefetch_related('composants_equipement', 'composants_point_eau')
    )
    deduites = _cibles_deduites([p.id for p in plannings])

    resolus, sans_cible = [], []
    for planning in plannings:
        equipement = point_eau = None
        if planning.cible_equipement_id:
            equipement = planning.cible_equipement
        elif planning.cible_point_eau_id:
            point_eau = planning.cible_point_eau
        else:
            found = deduites.get(str(planning.id), {'eq': {}, 'pe': {}})
            if len(found['eq']) == 1 and not found['pe']:
                equipement = next(iter(found['eq'].values()))
            elif len(found['pe']) == 1 and not found['eq']:
                point_eau = next(iter(found['pe'].values()))
        if equipement is None and point_eau is None:
            sans_cible.append({'id': str(planning.id), 'reason': 'cible non précisée (ou ambiguë)'})
            continue
        resolus.append((planning, equipement, point_eau))

    # Catalogues des types concernés : une requête par famille
    type_eq_ids = {e.type_equipement_id for _, e, _ in resolus if e is not None and e.type_equipement_id}
    type_pe_ids = {p.type_point_eau_id for _, _, p in resolus if p is not None and p.type_point_eau_id}
    catalogues = {'eq': {}, 'pe': {}}
    for c in ComposantTypeEquipement.objects.filter(type_equipement_id__in=type_eq_ids):
        catalogues['eq'].setdefault(c.type_equipement_id, {})[_fold(c.nom_composant)] = (str(c.id), c.nom_composant)
    for c in ComposantTypePointEau.objects.filter(type_point_eau_id__in=type_pe_ids):
        catalogues['pe'].setdefault(c.type_point_eau_id, {})[_fold(c.nom_composant)] = (str(c.id), c.nom_composant)

    lignes = []
    for planning, equipement, point_eau in resolus:
        if equipement is not None:
            type_id, type_obj, cat = equipement.type_equipement_id, equipement.type_equipement, catalogues['eq']
        else:
            type_id, type_obj, cat = point_eau.type_point_eau_id, point_eau.type_point_eau, catalogues['pe']
        lignes.append({
            'planning': planning,
            'equipement': equipement,
            'point_eau': point_eau,
            'type_nom': getattr(type_obj, 'nom_type', '') or '',
            'catalogue': cat.get(type_id, {}) if type_id else {},
        })
    return lignes, sans_cible


def _lire_ids_selection(request):
    try:
        payload = json.loads(request.body or '{}')
    except (TypeError, ValueError):
        return None, None, JsonResponse({'success': False, 'error': 'Corps JSON invalide'}, status=400)
    ids = payload.get('ids') if isinstance(payload, dict) else None
    if not isinstance(ids, list) or not ids:
        return None, None, JsonResponse({'success': False, 'error': 'ids est requis'}, status=400)
    ids = list(dict.fromkeys(str(i) for i in ids))
    if len(ids) > BULK_UPDATE_MAX:
        return None, None, JsonResponse(
            {'success': False, 'error': f'Maximum {BULK_UPDATE_MAX} réparations par opération'}, status=400)
    return payload, ids, None


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_composants_selection(request):
    """
    Composants proposables pour une sélection : noms du catalogue des types des cibles,
    avec le nombre de réparations auxquelles chacun s'applique.
    """
    _, ids, erreur = _lire_ids_selection(request)
    if erreur:
        return erreur
    try:
        lignes, sans_cible = _preparer_selection_composants(request, ids)
    except (DjangoValidationError, ValueError):
        return JsonResponse({'success': False, 'error': 'Identifiant invalide'}, status=400)

    noms = {}
    for ligne in lignes:
        for cle, (_, nom) in ligne['catalogue'].items():
            entry = noms.setdefault(cle, {'nom': nom, 'count': 0, 'types': set()})
            entry['count'] += 1
            entry['types'].add(ligne['type_nom'] or '—')
    return JsonResponse({
        'success': True,
        'total': len(lignes) + len(sans_cible),
        'avec_cible': len(lignes),
        'sans_cible': len(sans_cible),
        'composants': sorted(
            ({'cle': k, 'nom': v['nom'], 'count': v['count'], 'types': sorted(v['types'])} for k, v in noms.items()),
            key=lambda x: x['nom'].lower()),
    })


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_reparations_bulk_composants(request):
    """
    Applique des composants (par nom) à une sélection de réparations.

    Corps : {"ids": [...], "noms": ["Plaque", ...], "mode": "ajouter" | "remplacer" | "vider"}
    Chaque nom n'est appliqué que s'il existe dans le catalogue du type de la cible de la réparation.
    Les règles de l'activité (composants obligatoires, un seul composant) sont respectées ;
    les réparations refusées sont renvoyées avec leur motif.
    """
    payload, ids, erreur = _lire_ids_selection(request)
    if erreur:
        return erreur
    mode = payload.get('mode') or 'ajouter'
    if mode not in ('ajouter', 'remplacer', 'vider'):
        return JsonResponse({'success': False, 'error': 'Mode invalide'}, status=400)
    noms = payload.get('noms') or []
    if mode != 'vider' and (not isinstance(noms, list) or not noms):
        return JsonResponse({'success': False, 'error': 'Choisissez au moins un composant'}, status=400)
    voulus = list(dict.fromkeys(_fold(n) for n in noms if _fold(n)))

    try:
        lignes, sans_cible = _preparer_selection_composants(request, ids)
    except (DjangoValidationError, ValueError):
        return JsonResponse({'success': False, 'error': 'Identifiant invalide'}, status=400)

    refuses = list(sans_cible)
    updated = []
    with transaction.atomic():
        for ligne in lignes:
            planning = ligne['planning']
            catalogue = ligne['catalogue']
            existants = (
                [str(c.id) for c in planning.composants_equipement.all()]
                + [str(c.id) for c in planning.composants_point_eau.all()]
            )
            trouves = [catalogue[f][0] for f in voulus if f in catalogue]

            if mode == 'vider':
                finaux = []
            elif not trouves:
                refuses.append({
                    'id': str(planning.id),
                    'reason': f"aucun de ces composants dans le catalogue du type « {ligne['type_nom'] or '—'} »",
                })
                continue
            elif mode == 'remplacer':
                finaux = trouves
            else:  # ajouter
                finaux = list(dict.fromkeys(existants + trouves))

            try:
                verifier_regles_fiche(planning.fiche_poste, ligne['equipement'], ligne['point_eau'], finaux)
                # Enregistre aussi la cible déduite des détails quand elle n'était pas encore enregistrée
                apply_repair_target(planning, ligne['equipement'], ligne['point_eau'], finaux)
            except ComposantsError as exc:
                refuses.append({'id': str(planning.id), 'reason': str(exc)})
                continue
            updated.append(str(planning.id))

    return JsonResponse({
        'success': True,
        'updated': len(updated),
        'updated_ids': updated,
        'refused': refuses,
    })


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_marquer_reparation_terminee(request, planning_id):
    """API pour marquer une réparation comme terminée"""
    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        planning = Planning.objects.get(
            id=planning_id,
            ouvrage_id__in=scoped_ouvrage_ids,
        )

        # Garder les deux représentations historiques du statut cohérentes.
        planning.estvalide = True
        planning.statut = 'termine'
        planning.suiteadonner = 'Termine'
        planning.save(update_fields=['estvalide', 'statut', 'suiteadonner', 'updated_at'])
        close_repair_conversation(
            planning,
            actor=_request_actor(request),
            reason="Reparation marquee terminee depuis le suivi reparations.",
        )

        return JsonResponse({
            'success': True,
            'message': 'Réparation marquée comme terminée'
        })
    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Planning introuvable'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@mobile_auth_required
@csrf_exempt
@require_http_methods(["POST"])
def api_marquer_reparation_non_terminee(request, planning_id):
    '''API pour marquer une reparation comme non terminee'''
    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        planning = Planning.objects.get(
            id=planning_id,
            ouvrage_id__in=scoped_ouvrage_ids,
        )

        # Garder les deux représentations historiques du statut cohérentes.
        planning.estvalide = False
        planning.statut = 'en_cours'
        planning.suiteadonner = 'A suivre'
        planning.save(update_fields=['estvalide', 'statut', 'suiteadonner', 'updated_at'])
        ensure_repair_conversation(planning, actor=_request_actor(request), notify=True)

        return JsonResponse({
            'success': True,
            'message': 'Reparation marquee comme non terminee'
        })
    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Planning introuvable'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@csrf_exempt
@require_http_methods(["POST"])
def api_livrer_reparation(request, planning_id):
    """Livre les bons d'une réparation et génère le PDF de bon de livraison."""
    try:
        planning = Planning.objects.get(id=planning_id)

        if not planning.estvalide:
            return JsonResponse({
                'success': False,
                'error': "La réparation doit être validée avant livraison."
            }, status=400)

        payload = {}
        if request.body:
            try:
                payload = json.loads(request.body)
            except Exception:
                payload = {}

        bon_ids = payload.get('bon_ids') if isinstance(payload, dict) else None
        bons_qs = BonLivraison.objects.filter(
            planning=planning,
            utilise=False,
            retour=False
        ).select_related(
            'planning',
            'planning__employe',
            'planning__ouvrage',
            'designation',
            'point_eau',
            'equipement',
        )

        if isinstance(bon_ids, list) and bon_ids:
            bons_qs = bons_qs.filter(id__in=bon_ids)

        bons_list = list(bons_qs)
        if not bons_list:
            return JsonResponse({
                'success': False,
                'error': "Aucun bon de livraison non livré à traiter."
            }, status=400)

        from apps.inventory.models import FicheStock
        from apps.planning.utils import generer_numero_bon
        from apps.planning.views_livraisons import (
            _compute_stock_avant_pour_bon_livraison,
            _resolve_base_from_magasin_text,
            generer_bon_livraison_pdf,
        )

        stocks_avant = {}
        for bon in bons_list:
            stock_actuel = _compute_stock_avant_pour_bon_livraison(bon)
            stocks_avant[str(bon.id)] = stock_actuel

        for bon in bons_list:
            bon.utilise = True
            bon.date_utilisation = timezone.now()
            if not bon.numero_bon:
                bon.numero_bon = generer_numero_bon(type_bon='livraison')
            bon.save()

            if bon.designation and bon.quantite and bon.magasin:
                base = _resolve_base_from_magasin_text(bon.magasin)
                if base:
                    FicheStock.objects.create(
                        designation=bon.designation,
                        base=base,
                        quantity=bon.quantite,
                        operation_type='livraison',
                        operation_date=timezone.now().date(),
                        notes=f'BonLivraison ID: {bon.id} - Livraison depuis réparations - BL {bon.numero_bon or "N/A"}',
                        reference=bon.numero_bon
                    )

        pdf_buffer = generer_bon_livraison_pdf(bons_list, stocks_avant)
        response = HttpResponse(pdf_buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename=bon_livraison_reparation_{timezone.now().strftime("%Y%m%d_%H%M")}.pdf'
        return response
    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Planning introuvable'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@csrf_exempt
@require_http_methods(["POST"])
def api_utiliser_pieces_transferees_reparation(request, planning_id):
    """Marque directement les bons de pièces déjà transférées comme livrés et utilisés, sans PDF."""
    try:
        planning = Planning.objects.select_related('employe', 'employe__base').get(id=planning_id)
        actor = _request_actor(request)

        payload = {}
        if request.body:
            try:
                payload = json.loads(request.body)
            except Exception:
                payload = {}

        bon_ids = payload.get('bon_ids') if isinstance(payload, dict) else None
        bons_qs = BonLivraison.objects.filter(
            planning=planning,
            utilise=False,
            retour=False,
        ).select_related('designation')

        if isinstance(bon_ids, list) and bon_ids:
            bons_qs = bons_qs.filter(id__in=bon_ids)

        bons_list = list(bons_qs)
        if not bons_list:
            return JsonResponse({
                'success': False,
                'error': "Aucun bon de pièce transférée à enregistrer sur cette réparation."
            }, status=400)

        with transaction.atomic():
            from apps.planning.utils import generer_numero_bon

            processed_bon_ids = []
            for bon in bons_list:
                bon.utilise = True
                if not bon.date_utilisation:
                    bon.date_utilisation = timezone.now()
                if not bon.numero_bon:
                    bon.numero_bon = generer_numero_bon(type_bon='livraison')
                bon.save()
                processed_bon_ids.append(str(bon.id))

            if not planning.estvalide:
                planning.estvalide = True
                planning.save()

        close_repair_conversation(
            planning,
            actor=actor,
            reason="Pieces transferees enregistrees comme utilisees depuis le suivi reparations.",
        )

        if processed_bon_ids:
            message = (
                f"{len(processed_bon_ids)} pièce(s) transférée(s) marquée(s) livrée(s) "
                f"et utilisée(s) sans PDF. Réparation validée et terminée."
            )
        else:
            message = "Les pièces transférées étaient déjà enregistrées. Réparation validée et terminée."

        return JsonResponse({
            'success': True,
            'message': message,
            'processed_bon_ids': processed_bon_ids,
            'processed_count': len(processed_bon_ids),
        })
    except Planning.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Planning introuvable'
        }, status=404)
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_export_reparations_excel(request):
    """Export Excel des réparations (vue tableau) avec activité et remarque."""
    try:
        # Filtres
        region_id = request.GET.get('region_id', 'all')
        district_id = request.GET.get('district_id', 'all')
        commune_id = request.GET.get('commune_id', 'all')
        employe_id = (request.GET.get('employe_id') or 'all').strip()
        activite_id = (request.GET.get('activite_id') or request.GET.get('activite') or 'all').strip()
        recherche = request.GET.get('recherche', '').strip()
        grading_filter = request.GET.get('grading', 'all')
        statut_filter = request.GET.get('statut', 'all')
        annee_filter = request.GET.get('annee', 'all')

        try:
            annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
        except (TypeError, ValueError):
            annee_int = None

        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({'success': False, 'error': "Aucun ouvrage accessible pour cette organisation"}, status=404)

        # Mapping activités de réparation
        _, fiche_poste_ids, _, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()
        if not fiche_poste_ids:
            return JsonResponse({'success': False, 'error': "Aucune activité de réparation configurée"}, status=400)

        # Ouvrages filtrés (même logique que la vue)
        ouvrages_qs = Ouvrage.objects.filter(
            id__in=scoped_ouvrage_ids,
            type_ouvrage__existe=True
        ).select_related('type_ouvrage').distinct()

        if commune_id and commune_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune_id=commune_id
            ).distinct()
        elif district_id and district_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district_id=district_id
            ).distinct()
        elif region_id and region_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
            ).distinct()

        if recherche:
            ouvrages_qs = ouvrages_qs.filter(nom_ouvrage__icontains=recherche)
        ouvrages_qs = _filter_ouvrages_queryset_by_responsable_employe(
            ouvrages_qs,
            employe_id,
        )

        ouvrages = list(ouvrages_qs[:500])
        ouvrage_ids = [o.id for o in ouvrages]
        if not ouvrage_ids:
            return JsonResponse({'success': False, 'error': "Aucun ouvrage trouvé pour les filtres"}, status=404)

        gestionnaire_ids = get_user_accessible_gestionnaire_ids(request.user)
        organisation_id = _get_selected_organisation_id(request, gestionnaire_ids)

        # Localisation par ouvrage (premier point d'eau)
        point_eaux = PointEau.objects.filter(ouvrage_id__in=ouvrage_ids).select_related(
            'parcelle__village__fokontany__commune__district__region'
        )
        loc_by_ouvrage = {}
        for pe in point_eaux:
            if pe.ouvrage_id in loc_by_ouvrage:
                continue
            commune_nom = district_nom = region_nom = 'Non defini'
            if pe.parcelle and pe.parcelle.village and pe.parcelle.village.fokontany and pe.parcelle.village.fokontany.commune:
                c = pe.parcelle.village.fokontany.commune
                commune_nom = c.nom_commune
                if c.district:
                    district_nom = c.district.nom_district
                    if c.district.region:
                        region_nom = c.district.region.nom_region
            loc_by_ouvrage[pe.ouvrage_id] = {
                'commune': commune_nom,
                'district': district_nom,
                'region': region_nom,
            }

        # Modalité active par ouvrage
        modalites = _get_scoped_modalites_queryset(
            organisation_id,
            gestionnaire_ids,
            annee=annee_int,
        ).filter(ouvrage_id__in=ouvrage_ids)
        modalite_by_ouvrage = {}
        for m in modalites:
            label = 'Non défini'
            if m.modalitepaiement == 'forfaitaire':
                label = 'Forfaitaire'
            elif m.modalitepaiement == 'au_volume':
                label = 'Au volume'
            elif m.modalitepaiement == 'mixte':
                label = 'Mixte'
            elif m.modalitepaiement == 'cotisation_annuelle':
                label = 'Cotisation annuelle'
            modalite_by_ouvrage[m.ouvrage_id] = label

        # Réparations base (pour grading et export)
        tz_madagascar = pytz.timezone('Indian/Antananarivo')
        today = timezone.now().astimezone(tz_madagascar).date()
        reparations_qs = Planning.objects.filter(
            ouvrage_id__in=ouvrage_ids,
            fiche_poste_id__in=fiche_poste_ids
        ).select_related('ouvrage', 'employe')

        if annee_int:
            reparations_qs = reparations_qs.filter(date__year=annee_int)

        reparations = list(reparations_qs)

        # Calcul grading par ouvrage: uniquement sur les reparations non validees
        stats_by_ouvrage = {}
        for o in ouvrages:
            stats_by_ouvrage[o.id] = []

        for rep in reparations:
            if rep.ouvrage_id in stats_by_ouvrage:
                stats_by_ouvrage[rep.ouvrage_id].append(rep)

        grading_by_ouvrage = {}
        for ouvrage_id, ouvrage_reparations in stats_by_ouvrage.items():
            grading_by_ouvrage[ouvrage_id] = get_repair_workflow_grading_from_plannings(
                ouvrage_reparations,
                fiche_poste_to_activite_nom=fiche_poste_to_activite_nom,
            )['code']

        # Filtre grading (sur les ouvrages)
        if grading_filter and grading_filter != 'all':
            allowed_ouvrage_ids = {oid for oid, g in grading_by_ouvrage.items() if g == grading_filter}
            ouvrages = [o for o in ouvrages if o.id in allowed_ouvrage_ids]
            ouvrage_ids = [o.id for o in ouvrages]

        # Filtre statut (pour déterminer le "dernier" planning retenu)
        def _statut_rep(rep):
            if rep.estvalide:
                return 'validee', 'Validée'
            if rep.date:
                date_aware = rep.date if rep.date.tzinfo else timezone.make_aware(rep.date)
                date_obj = date_aware.astimezone(tz_madagascar).date()
                if date_obj > today:
                    return 'planifiee', 'Planifiée'
            return 'en_cours', 'En cours'

        reparations_filtrees = []
        for rep in reparations:
            statut_val, _ = _statut_rep(rep)
            if statut_filter and statut_filter != 'all' and statut_val != statut_filter:
                continue
            if rep.ouvrage_id not in ouvrage_ids:
                continue
            reparations_filtrees.append(rep)

        # Dernière réparation par ouvrage (date la plus récente)
        latest_rep_by_ouvrage = {}
        for rep in reparations_filtrees:
            if not rep.ouvrage_id:
                continue
            current = latest_rep_by_ouvrage.get(rep.ouvrage_id)
            if current is None:
                latest_rep_by_ouvrage[rep.ouvrage_id] = rep
                continue
            rep_date = rep.date or timezone.now()
            cur_date = current.date or timezone.now()
            if rep_date > cur_date:
                latest_rep_by_ouvrage[rep.ouvrage_id] = rep

        # Une seule ligne par ouvrage (feuille 1)
        rows = []
        for ouvrage in ouvrages:
            loc = loc_by_ouvrage.get(ouvrage.id, {'region': 'Non defini', 'district': 'Non defini', 'commune': 'Non defini'})
            latest_rep = latest_rep_by_ouvrage.get(ouvrage.id)

            date_dernier = ''
            activite = ''
            remarque = ''
            explication_grading = ''
            statut_dernier = ''
            assigne_a = ''

            if latest_rep:
                if latest_rep.date:
                    date_aware = latest_rep.date if latest_rep.date.tzinfo else timezone.make_aware(latest_rep.date)
                    date_dernier = date_aware.astimezone(tz_madagascar).strftime('%d/%m/%Y')
                activite = fiche_poste_to_activite_nom.get(latest_rep.fiche_poste_id, '') if latest_rep.fiche_poste_id else ''
                remarque = latest_rep.remarque or ''
                if activite and remarque:
                    explication_grading = f"{activite} - {remarque}"
                else:
                    explication_grading = activite or remarque or ''
                _, statut_dernier = _statut_rep(latest_rep)
                assigne_a = latest_rep.employe.surnom if latest_rep.employe else ''

            rows.append({
                'region': loc.get('region', 'Non defini'),
                'district': loc.get('district', 'Non defini'),
                'commune': loc.get('commune', 'Non defini'),
                'ouvrage': ouvrage.nom_ouvrage,
                'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
                'modalite': modalite_by_ouvrage.get(ouvrage.id, 'Non défini'),
                'grading': grading_by_ouvrage.get(ouvrage.id, ''),
                'date_dernier': date_dernier,
                'statut_dernier': statut_dernier,
                'activite': activite,
                'remarque': remarque,
                'explication_grading': explication_grading,
                'assigne_a': assigne_a,
            })

        rows.sort(key=lambda r: (r['region'], r['district'], r['commune'], r['ouvrage']))

        # Feuille 2: toutes les réparations (terminées ou non)
        reparations_detail_rows = []
        reparations_detail_all = [
            r for r in reparations
            if r.ouvrage_id in ouvrage_ids
        ]
        for rep in reparations_detail_all:
            ouvrage = rep.ouvrage
            if not ouvrage:
                continue

            loc = loc_by_ouvrage.get(ouvrage.id, {'region': 'Non defini', 'district': 'Non defini', 'commune': 'Non defini'})
            activite_nom = fiche_poste_to_activite_nom.get(rep.fiche_poste_id, '') if rep.fiche_poste_id else ''
            statut_val, statut_label = _statut_rep(rep)
            date_prevue = ''
            if rep.date:
                date_aware = rep.date if rep.date.tzinfo else timezone.make_aware(rep.date)
                date_prevue = date_aware.astimezone(tz_madagascar).strftime('%d/%m/%Y')

            reparations_detail_rows.append({
                'region': loc.get('region', 'Non defini'),
                'district': loc.get('district', 'Non defini'),
                'commune': loc.get('commune', 'Non defini'),
                'ouvrage': ouvrage.nom_ouvrage,
                'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
                'modalite': modalite_by_ouvrage.get(ouvrage.id, 'Non défini'),
                'grading': grading_by_ouvrage.get(ouvrage.id, ''),
                'statut': statut_label,
                'date_prevue': date_prevue,
                'activite': activite_nom or '',
                'remarque': rep.remarque or '',
                'assigne_a': rep.employe.surnom if rep.employe else '',
                'bilan': rep.bilan or '',
                'suite_a_donner': rep.suiteadonner or '',
            })

        reparations_detail_rows.sort(
            key=lambda r: (r['region'], r['district'], r['commune'], r['ouvrage'], r['date_prevue'])
        )

        # Génération Excel (2 feuilles)
        wb = Workbook()
        ws = wb.active
        ws.title = "Synthese Ouvrages"

        headers = [
            'Région', 'District', 'Commune', 'Ouvrage', 'Type ouvrage', 'Modalité',
            'Dernier grading', 'Date dernier planning', 'Statut dernier planning',
            'Explication grading', 'Activité', 'Remarque', 'Assigné à'
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for r in rows:
            ws.append([
                r['region'],
                r['district'],
                r['commune'],
                r['ouvrage'],
                r['type_ouvrage'],
                r['modalite'],
                r['grading'],
                r['date_dernier'],
                r['statut_dernier'],
                r['explication_grading'],
                r['activite'],
                r['remarque'],
                r['assigne_a'],
            ])

        # Largeurs utiles
        widths = {
            1: 20, 2: 22, 3: 22, 4: 30, 5: 18, 6: 16, 7: 14, 8: 18,
            9: 20, 10: 55, 11: 30, 12: 45, 13: 16
        }
        for col_idx, width in widths.items():
            ws.column_dimensions[chr(64 + col_idx)].width = width

        # Feuille 2: détail réparations
        ws2 = wb.create_sheet(title="Reparations Detail")
        headers2 = [
            'Région', 'District', 'Commune', 'Ouvrage', 'Type ouvrage', 'Modalité',
            'Grading', 'Statut', 'Date prévue', 'Activité', 'Remarque',
            'Assigné à', 'Bilan', 'Suite à donner'
        ]
        ws2.append(headers2)
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=1, column=col_idx)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for r in reparations_detail_rows:
            ws2.append([
                r['region'],
                r['district'],
                r['commune'],
                r['ouvrage'],
                r['type_ouvrage'],
                r['modalite'],
                r['grading'],
                r['statut'],
                r['date_prevue'],
                r['activite'],
                r['remarque'],
                r['assigne_a'],
                r['bilan'],
                r['suite_a_donner'],
            ])

        widths2 = {
            1: 20, 2: 22, 3: 22, 4: 30, 5: 18, 6: 16, 7: 12, 8: 14,
            9: 14, 10: 28, 11: 45, 12: 16, 13: 30, 14: 30
        }
        for col_idx, width in widths2.items():
            ws2.column_dimensions[chr(64 + col_idx)].width = width

        filename = f"reparations_tableau_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_stats_reparations(request):
    """API pour les statistiques de reparations par statut"""
    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({
                'success': True,
                'stats': {
                    'signalees': 0,
                    'planifiees': 0,
                    'en_cours': 0,
                    'terminees': 0,
                    'validees': 0,
                }
            })

        # Recuperer les filtres de la requete
        region_id = request.GET.get('region_id', 'all')
        district_id = request.GET.get('district_id', 'all')
        commune_id = request.GET.get('commune_id', 'all')
        ouvrage_id = request.GET.get('ouvrage_id', 'all')
        employe_id = (request.GET.get('employe_id') or 'all').strip()
        grading = request.GET.get('grading', 'all')
        annee_filter = request.GET.get('annee', 'all')
        try:
            annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
        except (TypeError, ValueError):
            annee_int = None

        # Recuperer tous les ouvrages de l'organisation avec modalite active
        # et type_ouvrage.existe = true
        ouvrages_qs = Ouvrage.objects.filter(
            id__in=scoped_ouvrage_ids,
            type_ouvrage__existe=True  # Seulement les ouvrages qui existent
        ).distinct()

        # Appliquer les filtres geographiques via PointEau
        if ouvrage_id and ouvrage_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(id=ouvrage_id)
        elif commune_id and commune_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune_id=commune_id
            ).distinct()
        elif district_id and district_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district_id=district_id
            ).distinct()
        elif region_id and region_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
            ).distinct()
        ouvrages_qs = _filter_ouvrages_queryset_by_responsable_employe(
            ouvrages_qs,
            employe_id,
        )

        # Si filtre par grading, calculer le grading de chaque ouvrage et filtrer
        if grading and grading != 'all':
            repair_activities, fiche_poste_ids, _, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()

            ouvrages_filtres = []
            for ouvrage_id in ouvrages_qs.values_list('id', flat=True):
                reparations_ouvrage = Planning.objects.filter(
                    ouvrage_id=ouvrage_id,
                    fiche_poste_id__in=fiche_poste_ids
                )
                if annee_int:
                    reparations_ouvrage = reparations_ouvrage.filter(date__year=annee_int)

                ouvrage_grading = get_repair_workflow_grading_from_plannings(
                    reparations_ouvrage,
                    fiche_poste_to_activite_nom=fiche_poste_to_activite_nom,
                )['code']

                # Filtrer selon le grading demande
                if ouvrage_grading == grading:
                    ouvrages_filtres.append(ouvrage_id)

            ouvrages_ids = ouvrages_filtres
        else:
            ouvrages_ids = ouvrages_qs.values_list('id', flat=True)

        # Recuperer toutes les reparations (Planning avec activite reparation)
        repair_activities, fiche_poste_ids, _, _, _ = _repair_activity_mappings()
        reparations = Planning.objects.filter(
            ouvrage_id__in=ouvrages_ids,
            fiche_poste_id__in=fiche_poste_ids
        )
        if annee_int:
            reparations = reparations.filter(date__year=annee_int)

        # Pour simplifier:
        # - validees: estvalide=True
        # - en_cours: estvalide=False et date passee
        # - signalees/planifiees: estvalide=False et date future
        today = timezone.now().date()
        tz_madagascar = pytz.timezone('Indian/Antananarivo')

        validees = 0
        en_cours = 0
        planifiees = 0

        for rep in reparations:
            if rep.estvalide:
                validees += 1
            elif rep.date:
                date_aware = rep.date if rep.date.tzinfo else timezone.make_aware(rep.date)
                date_mg = date_aware.astimezone(tz_madagascar)
                if date_mg.date() < today:
                    en_cours += 1
                else:
                    planifiees += 1
            else:
                en_cours += 1

        return JsonResponse({
            'success': True,
            'stats': {
                'signalees': 0,  # Non utilise dans ce workflow
                'planifiees': planifiees,
                'en_cours': en_cours,
                'terminees': 0,  # Non utilise dans ce workflow
                'validees': validees,
            }
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_stats_grading(request):
    """API pour les statistiques de grading"""
    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({
                'success': True,
                'stats': {
                    'grade_a': 0,
                    'grade_b': 0,
                    'grade_b_panne': 0,
                    'grade_c': 0,
                    'grade_d': 0,
                    'grade_d_panne': 0,
                }
            })

        # Recuperer les filtres de la requete
        region_id = request.GET.get('region_id', 'all')
        district_id = request.GET.get('district_id', 'all')
        commune_id = request.GET.get('commune_id', 'all')
        ouvrage_id = request.GET.get('ouvrage_id', 'all')
        employe_id = (request.GET.get('employe_id') or 'all').strip()
        activite_id = (request.GET.get('activite_id') or request.GET.get('activite') or 'all').strip()
        grading = request.GET.get('grading', 'all')
        annee_filter = request.GET.get('annee', 'all')
        try:
            annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
        except (TypeError, ValueError):
            annee_int = None

        # Recuperer tous les ouvrages de l'organisation avec modalite active
        # et type_ouvrage.existe = true
        ouvrages_qs = Ouvrage.objects.filter(
            id__in=scoped_ouvrage_ids,
            type_ouvrage__existe=True  # Seulement les ouvrages qui existent
        ).distinct()

        # Appliquer les filtres geographiques via PointEau
        if ouvrage_id and ouvrage_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(id=ouvrage_id)
        elif commune_id and commune_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune_id=commune_id
            ).distinct()
        elif district_id and district_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district_id=district_id
            ).distinct()
        elif region_id and region_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
            ).distinct()
        ouvrages_qs = _filter_ouvrages_queryset_by_responsable_employe(
            ouvrages_qs,
            employe_id,
        )

        repair_activities, fiche_poste_ids, _, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()
        today = timezone.now().date()
        tz_madagascar = pytz.timezone('Indian/Antananarivo')

        # Compter par grading
        grade_a = 0
        grade_b = 0
        grade_b_panne = 0
        grade_c = 0
        grade_d = 0
        grade_d_panne = 0

        for ouvrage in ouvrages_qs:
            # Calculer le grading de cet ouvrage
            reparations = Planning.objects.filter(
                ouvrage=ouvrage,
                fiche_poste_id__in=fiche_poste_ids
            )
            if activite_id and activite_id != 'all':
                reparations = reparations.filter(fiche_poste_id=activite_id)
            if annee_int:
                reparations = reparations.filter(date__year=annee_int)

            if activite_id and activite_id != 'all' and not reparations.exists():
                continue

            ouvrage_grading = get_repair_workflow_grading_from_plannings(
                reparations,
                fiche_poste_to_activite_nom=fiche_poste_to_activite_nom,
            )['code']

            # Si filtre par grading, ne compter que les ouvrages correspondants
            if grading and grading != 'all' and ouvrage_grading != grading:
                continue

            # Incrementer le compteur correspondant
            if ouvrage_grading == 'A':
                grade_a += 1
            elif ouvrage_grading == 'D_panne':
                grade_d_panne += 1
            elif ouvrage_grading == 'D':
                grade_d += 1
            elif ouvrage_grading == 'B_panne':
                grade_b_panne += 1
            elif ouvrage_grading == 'B':
                grade_b += 1
            elif ouvrage_grading == 'C':
                grade_c += 1

        return JsonResponse({
            'success': True,
            'stats': {
                'grade_a': grade_a,
                'grade_b': grade_b,
                'grade_b_panne': grade_b_panne,
                'grade_c': grade_c,
                'grade_d': grade_d,
                'grade_d_panne': grade_d_panne,
            }
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_ouvrages_avec_modalites(request):
    """API pour lister les ouvrages avec leurs modalites et reparations"""

    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({
                'success': True,
                'ouvrages': [],
                'total': 0
            })

        # Recuperer les filtres de la requete
        region_id = request.GET.get('region_id', 'all')
        district_id = request.GET.get('district_id', 'all')
        commune_id = request.GET.get('commune_id', 'all')
        employe_id = (request.GET.get('employe_id') or 'all').strip()
        recherche = request.GET.get('recherche', '').strip()
        annee_filter = request.GET.get('annee', 'all')
        try:
            annee_int = int(annee_filter) if annee_filter and annee_filter != 'all' else None
        except (TypeError, ValueError):
            annee_int = None

        gestionnaire_ids = get_user_accessible_gestionnaire_ids(request.user)
        organisation_id = _get_selected_organisation_id(request, gestionnaire_ids)

        # Requete de base - ouvrages avec modalites actives de l'organisation
        # et type_ouvrage.existe = true
        ouvrages_qs = Ouvrage.objects.filter(
            id__in=scoped_ouvrage_ids,
            type_ouvrage__existe=True  # Seulement les ouvrages qui existent
        ).select_related('type_ouvrage').distinct()

        # Appliquer les filtres geographiques via PointEau
        if commune_id and commune_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune_id=commune_id
            ).distinct()
        elif district_id and district_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district_id=district_id
            ).distinct()
        elif region_id and region_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
            ).distinct()

        # Appliquer filtre de recherche par nom
        if recherche:
            ouvrages_qs = ouvrages_qs.filter(nom_ouvrage__icontains=recherche)
        ouvrages_qs = _filter_ouvrages_queryset_by_responsable_employe(
            ouvrages_qs,
            employe_id,
        )

        # Limiter a 500 resultats
        ouvrages_qs = ouvrages_qs[:500]

        ouvrages_data = []
        tz_madagascar = pytz.timezone('Indian/Antananarivo')

        # Recuperer les activites de reparation
        repair_activities, fiche_poste_ids, _, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()
        modalite_map = {}
        for modalite in _get_scoped_modalites_queryset(
            organisation_id,
            gestionnaire_ids,
            annee=annee_int,
        ).filter(
            ouvrage_id__in=ouvrages_qs.values_list('id', flat=True)
        ).order_by('ouvrage_id', 'date_debut', 'created_at'):
            modalite_map.setdefault(modalite.ouvrage_id, modalite)

        ouvrage_ids = list(ouvrages_qs.values_list('id', flat=True))
        latest_measures_qs = (
            DebitSourceMesure.objects.filter(ouvrage_id__in=ouvrage_ids)
            .select_related('ouvrage')
            .order_by('ouvrage_id', '-date_visite', '-created_at')
        )
        latest_measure_by_ouvrage = {}
        for measure in latest_measures_qs:
            latest_measure_by_ouvrage.setdefault(measure.ouvrage_id, measure)

        for ouvrage in ouvrages_qs:
            # Récupérer la modalité de gestion active
            modalite_active = modalite_map.get(ouvrage.id)

            modalite_paiement_label = 'Non défini'
            if modalite_active:
                # Récupérer le libellé de la modalité de paiement
                if modalite_active.modalitepaiement == 'forfaitaire':
                    modalite_paiement_label = 'Forfaitaire'
                elif modalite_active.modalitepaiement == 'au_volume':
                    modalite_paiement_label = 'Au volume'
                elif modalite_active.modalitepaiement == 'mixte':
                    modalite_paiement_label = 'Mixte'
                elif modalite_active.modalitepaiement == 'cotisation_annuelle':
                    modalite_paiement_label = 'Cotisation annuelle'

            # Commune via point d'eau
            point_eau = PointEau.objects.filter(ouvrage=ouvrage).select_related(
                'parcelle__village__fokontany__commune__district'
            ).first()

            commune_nom = district_nom = region_nom = None
            if point_eau and point_eau.parcelle and point_eau.parcelle.village:
                v = point_eau.parcelle.village
                if v.fokontany and v.fokontany.commune:
                    commune_nom = v.fokontany.commune.nom_commune
                    if v.fokontany.commune.district:
                        district_nom = v.fokontany.commune.district.nom_district
                        if v.fokontany.commune.district.region:
                            region_nom = v.fokontany.commune.district.region.nom_region

            # Calculer statistiques de reparations pour cet ouvrage
            today = timezone.now().date()
            reparations = Planning.objects.filter(
                ouvrage=ouvrage,
                fiche_poste_id__in=fiche_poste_ids
            )
            if annee_int:
                reparations = reparations.filter(date__year=annee_int)

            grading_info = get_repair_workflow_grading_from_plannings(
                reparations,
                fiche_poste_to_activite_nom=fiche_poste_to_activite_nom,
            )
            grading = grading_info['code']
            grading_badge = grading_info['badge']

            stats = {
                'total': reparations.count(),
                'signalee': 0,
                'en_cours': reparations.filter(estvalide=False).count(),
                'terminee': 0,
                'validee': reparations.filter(estvalide=True).count()
            }

            latest_debit_measure = latest_measure_by_ouvrage.get(ouvrage.id)
            debit_context = _get_ouvrage_debit_context(
                ouvrage,
                latest_debit_measure.date_visite if latest_debit_measure else timezone.localdate(),
            )
            debit_threshold = debit_context.get('seuil_minimum')
            debit_value = latest_debit_measure.debit_l_s if latest_debit_measure else None
            debit_delta = None
            debit_alert = False
            debit_status = 'inconnu'

            if latest_debit_measure:
                previous_measure = (
                    DebitSourceMesure.objects.filter(
                        ouvrage_id=ouvrage.id,
                        date_visite__lt=latest_debit_measure.date_visite,
                    )
                    .order_by('-date_visite', '-created_at')
                    .first()
                )
                if previous_measure and previous_measure.debit_l_s is not None and latest_debit_measure.debit_l_s is not None:
                    debit_delta = latest_debit_measure.debit_l_s - previous_measure.debit_l_s
                debit_status = _get_debit_status_label(latest_debit_measure, debit_context, debit_delta)
                debit_alert = bool(
                    debit_threshold is not None
                    and debit_value is not None
                    and Decimal(str(debit_value)) < Decimal(str(debit_threshold))
                )

            ouvrages_data.append({
                'id': str(ouvrage.id),
                'nom': ouvrage.nom_ouvrage,
                'type': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
                'commune': commune_nom or 'Non defini',
                'district': district_nom or 'Non defini',
                'region': region_nom or 'Non defini',
                'modalite': modalite_paiement_label,
                'grading': grading,
                'grading_badge': grading_badge,
                'reparations_stats': stats,
                'est_saisonnier': ouvrage.estsaisonnier if hasattr(ouvrage, 'estsaisonnier') else False,
                'debit_source': {
                    'has_measure': bool(latest_debit_measure),
                    'debit_l_s': float(debit_value) if debit_value is not None else None,
                    'date_visite': latest_debit_measure.date_visite.isoformat() if latest_debit_measure and latest_debit_measure.date_visite else '',
                    'threshold_l_s': float(debit_threshold) if debit_threshold is not None else None,
                    'delta_vs_prev': float(debit_delta) if debit_delta is not None else None,
                    'status': debit_status,
                    'is_alert': debit_alert,
                    'comparative_label': debit_context.get('comparative_label', ''),
                    'page_url': f"/plannings/debits-source/?ouvrage={ouvrage.id}",
                },
            })

        # Filtrer par grading si specifie
        grading_filter = request.GET.get('grading', 'all')
        if grading_filter and grading_filter != 'all':
            ouvrages_data = [o for o in ouvrages_data if o['grading'] == grading_filter]

        return JsonResponse({
            'success': True,
            'ouvrages': ouvrages_data,
            'total': len(ouvrages_data)
        })

    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_communes_reparations_summary(request):
    """API pour lister les statistiques reparations par commune"""
    try:
        _, scoped_ouvrage_ids = _get_scoped_ouvrage_ids_for_request(request)
        if not scoped_ouvrage_ids:
            return JsonResponse({'success': True, 'communes': [], 'total': 0})

        # Recuperer les filtres de la requete
        region_id = request.GET.get('region_id', 'all')
        district_id = request.GET.get('district_id', 'all')
        commune_id = request.GET.get('commune_id', 'all')
        employe_id = (request.GET.get('employe_id') or 'all').strip()
        recherche = request.GET.get('recherche', '').strip()
        grading_filter = request.GET.get('grading', 'all')
        statut_filter = request.GET.get('statut', 'all')

        # Requete de base - ouvrages avec modalites actives de l'organisation
        # et type_ouvrage.existe = true
        ouvrages_qs = Ouvrage.objects.filter(
            id__in=scoped_ouvrage_ids,
            type_ouvrage__existe=True
        ).distinct()

        # Appliquer les filtres geographiques via PointEau
        if commune_id and commune_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune_id=commune_id
            ).distinct()
        elif district_id and district_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district_id=district_id
            ).distinct()
        elif region_id and region_id != 'all':
            ouvrages_qs = ouvrages_qs.filter(
                pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
            ).distinct()

        # Appliquer filtre de recherche par nom
        if recherche:
            ouvrages_qs = ouvrages_qs.filter(nom_ouvrage__icontains=recherche)
        ouvrages_qs = _filter_ouvrages_queryset_by_responsable_employe(
            ouvrages_qs,
            employe_id,
        )

        ouvrages_ids = list(ouvrages_qs.values_list('id', flat=True))
        if not ouvrages_ids:
            return JsonResponse({'success': True, 'communes': [], 'total': 0})

        # Mapping ouvrage -> commune/district/region (premier point d'eau)
        point_eaux = PointEau.objects.filter(
            ouvrage_id__in=ouvrages_ids
        ).select_related(
            'parcelle__village__fokontany__commune__district__region'
        ).order_by('id')

        ouvrage_location = {}
        for pe in point_eaux:
            if pe.ouvrage_id in ouvrage_location:
                continue
            commune = None
            district = None
            region = None
            if pe.parcelle and pe.parcelle.village and pe.parcelle.village.fokontany:
                commune = pe.parcelle.village.fokontany.commune
                if commune and commune.district:
                    district = commune.district
                    if district.region:
                        region = district.region

            ouvrage_location[pe.ouvrage_id] = {
                'commune_id': str(commune.id) if commune else None,
                'commune_nom': commune.nom_commune if commune else 'Non defini',
                'district_nom': district.nom_district if district else 'Non defini',
                'region_nom': region.nom_region if region else 'Non defini'
            }

        # Recuperer les activites de reparation
        _, fiche_poste_ids, _, fiche_poste_to_activite_nom, _ = _repair_activity_mappings()
        tz_madagascar = pytz.timezone('Indian/Antananarivo')
        today = timezone.now().date()

        plannings = Planning.objects.filter(
            ouvrage_id__in=ouvrages_ids,
            fiche_poste_id__in=fiche_poste_ids
        ).values('ouvrage_id', 'date', 'estvalide', 'fiche_poste_id')

        # Preparer stats par ouvrage
        ouvrage_stats = {}
        for ouvrage_id in ouvrages_ids:
            ouvrage_stats[ouvrage_id] = {
                'total_reparations': 0,
                'validees': 0,
                'non_validees': 0,
                'plannings': []
            }

        for rep in plannings:
            oid = rep['ouvrage_id']
            stats = ouvrage_stats.get(oid)
            if not stats:
                continue

            stats['total_reparations'] += 1
            if rep['estvalide']:
                stats['validees'] += 1
            else:
                stats['non_validees'] += 1
            stats['plannings'].append(rep)

        # Agreger par commune
        communes_data = {}
        for ouvrage_id in ouvrages_ids:
            stats = ouvrage_stats.get(ouvrage_id) or {}

            grading = get_repair_workflow_grading_from_plannings(
                stats.get('plannings', []),
                fiche_poste_to_activite_nom=fiche_poste_to_activite_nom,
            )['code']

            # Filtre grading si present
            if grading_filter and grading_filter != 'all' and grading != grading_filter:
                continue

            # Filtre statut si present
            if statut_filter and statut_filter != 'all':
                if statut_filter == 'signalee' and stats.get('total_reparations', 0) == 0:
                    continue
                if statut_filter == 'planifiee' and stats.get('non_validees', 0) == 0:
                    continue
                if statut_filter == 'en_cours' and stats.get('non_validees', 0) == 0:
                    continue
                if statut_filter == 'terminee' and stats.get('validees', 0) == 0:
                    continue
                if statut_filter == 'validee' and stats.get('validees', 0) == 0:
                    continue

            loc = ouvrage_location.get(ouvrage_id) or {
                'commune_id': None,
                'commune_nom': 'Non defini',
                'district_nom': 'Non defini',
                'region_nom': 'Non defini'
            }
            commune_key = loc['commune_id'] or f"none-{loc['commune_nom']}"

            if commune_key not in communes_data:
                communes_data[commune_key] = {
                    'commune_id': loc['commune_id'],
                    'commune': loc['commune_nom'],
                    'district': loc['district_nom'],
                    'region': loc['region_nom'],
                    'total_ouvrages': 0,
                    'grade_a': 0,
                    'grade_b': 0,
                    'grade_b_panne': 0,
                    'grade_c': 0,
                    'grade_d': 0,
                    'grade_d_panne': 0,
                    'reparations_terminees': 0,
                    'reparations_autres': 0,
                    'reparations_total': 0
                }

            row = communes_data[commune_key]
            row['total_ouvrages'] += 1
            if grading == 'A':
                row['grade_a'] += 1
            elif grading == 'B':
                row['grade_b'] += 1
            elif grading == 'B_panne':
                row['grade_b_panne'] += 1
            elif grading == 'D':
                row['grade_d'] += 1
            elif grading == 'D_panne':
                row['grade_d_panne'] += 1
            elif grading == 'C':
                row['grade_c'] += 1

            row['reparations_terminees'] += stats.get('validees', 0)
            row['reparations_autres'] += stats.get('non_validees', 0)
            row['reparations_total'] += stats.get('total_reparations', 0)

        communes_list = sorted(
            communes_data.values(),
            key=lambda item: (item.get('region') or '', item.get('district') or '', item.get('commune') or '')
        )

        return JsonResponse({
            'success': True,
            'communes': communes_list,
            'total': len(communes_list)
        })
    except Exception as e:
        import traceback
        return JsonResponse({
            'success': False,
            'error': str(e),
            'traceback': traceback.format_exc()
        }, status=500)


@require_http_methods(["GET"])
def api_composants_cible_reparation(request):
    """Composants proposables pour un équipement ou un point d'eau (catalogue de son type)."""
    equipement_id = (request.GET.get('equipement_id') or '').strip()
    point_eau_id = (request.GET.get('point_eau_id') or '').strip()
    if bool(equipement_id) == bool(point_eau_id):
        return JsonResponse(
            {'success': False, 'error': "Indiquez soit equipement_id, soit point_eau_id"}, status=400)
    # Activité facultative : restreint la liste à ses composants admis (ex. anti-rouille).
    fiche_poste_id = (request.GET.get('fiche_poste_id') or '').strip()
    try:
        from apps.employees.models_evaluation import FichePoste
        fiche = FichePoste.objects.filter(id=fiche_poste_id).first() if fiche_poste_id else None
        if equipement_id:
            cible = EquipementConstitutif.objects.select_related('type_equipement').get(id=equipement_id)
            composants = composants_disponibles(equipement=cible, fiche=fiche)
        else:
            cible = PointEau.objects.select_related('type_point_eau').get(id=point_eau_id)
            composants = composants_disponibles(point_eau=cible, fiche=fiche)
    except (EquipementConstitutif.DoesNotExist, PointEau.DoesNotExist):
        return JsonResponse({'success': False, 'error': 'Cible introuvable'}, status=404)
    except DjangoValidationError:
        return JsonResponse({'success': False, 'error': 'Identifiant invalide'}, status=400)
    type_obj = (cible.type_equipement if equipement_id else cible.type_point_eau)
    return JsonResponse({
        'success': True,
        'composants': composants,
        # Le catalogue dépend du TYPE de la cible : on le nomme pour que l'écran l'affiche
        'type_defini': bool(type_obj),
        'type_nom': getattr(type_obj, 'nom_type', '') or '',
    })
