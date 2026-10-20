# ==================== EXPORT/IMPORT EXCEL DÉPLACEMENTS ====================

from django.shortcuts import redirect
from django.http import HttpResponse
from django.contrib import messages
from apps.employees.permissions import employe_required


@employe_required
def export_excel_deplacements(request):
    """Exporter les déplacements vers Excel"""
    from .models import Deplacement
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    # Créer le workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Déplacements"

    # En-têtes
    headers = ['ID Point A*', 'Point A (Village)', 'Fokontany A', 'ID Point B*', 'Point B (Village)', 'Fokontany B',
               'Moyen de Transport', 'Prix Unitaire (Ar)', 'Date du Prix', 'Créé le', 'Modifié le']
    ws.append(headers)

    # Style des en-têtes
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Récupérer les déplacements
    deplacements = Deplacement.objects.select_related('point_a__fokontany', 'point_b__fokontany').order_by('-date_prix')

    # Remplir les données
    for dep in deplacements:
        ws.append([
            str(dep.point_a.id) if dep.point_a else '',
            dep.point_a.nom_village if dep.point_a else '',
            dep.point_a.fokontany.nom_fokontany if dep.point_a and dep.point_a.fokontany else '',
            str(dep.point_b.id) if dep.point_b else '',
            dep.point_b.nom_village if dep.point_b else '',
            dep.point_b.fokontany.nom_fokontany if dep.point_b and dep.point_b.fokontany else '',
            dep.get_moyen_transport_display(),
            float(dep.prix_unitaire),
            dep.date_prix.strftime('%Y-%m-%d'),
            dep.created_at.strftime('%Y-%m-%d %H:%M'),
            dep.updated_at.strftime('%Y-%m-%d %H:%M'),
        ])

    # Ajuster les largeurs de colonnes
    ws.column_dimensions['A'].width = 12  # ID Point A
    ws.column_dimensions['B'].width = 20  # Point A (Village)
    ws.column_dimensions['C'].width = 20  # Fokontany A
    ws.column_dimensions['D'].width = 12  # ID Point B
    ws.column_dimensions['E'].width = 20  # Point B (Village)
    ws.column_dimensions['F'].width = 20  # Fokontany B
    ws.column_dimensions['G'].width = 18  # Moyen de Transport
    ws.column_dimensions['H'].width = 15  # Prix Unitaire
    ws.column_dimensions['I'].width = 15  # Date du Prix
    ws.column_dimensions['J'].width = 18  # Créé le
    ws.column_dimensions['K'].width = 18  # Modifié le

    # Créer la réponse HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=deplacements_export.xlsx'
    wb.save(response)
    return response


@employe_required
def template_excel_deplacements(request):
    """Télécharger le template Excel pour l'import de déplacements"""
    from apps.geography.models import Village
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    # Feuille principale
    ws = wb.active
    ws.title = "Déplacements"

    # En-têtes avec instructions
    headers = ['ID Point A*', 'ID Point B*', 'Moyen de Transport*',
               'Prix Unitaire (Ar)*', 'Date du Prix (YYYY-MM-DD)*']
    ws.append(headers)

    # Style des en-têtes
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajouter des exemples (avec des IDs fictifs)
    ws.append(['12345', '67890', 'taxi-brousse', 5000, '2025-10-20'])
    ws.append(['11111', '22222', 'pirogue', 3000, '2025-10-20'])

    # Instructions
    ws_info = wb.create_sheet("Instructions")
    ws_info.append(["INSTRUCTIONS POUR L'IMPORT"])
    ws_info.append([])
    ws_info.append(["Colonnes obligatoires (*)"])
    ws_info.append(["- ID Point A*: ID du village de départ (voir feuille Villages)"])
    ws_info.append(["- ID Point B*: ID du village d'arrivée (voir feuille Villages)"])
    ws_info.append(["- Moyen de Transport*: taxi-brousse, pirogue, ou bac"])
    ws_info.append(["- Prix Unitaire (Ar)*: Montant en Ariary (nombre)"])
    ws_info.append(["- Date du Prix*: Format YYYY-MM-DD (ex: 2025-10-20)"])
    ws_info.append([])
    ws_info.append(["Notes importantes:"])
    ws_info.append(["- Utilisez les IDs de la feuille Villages (pas les noms)"])
    ws_info.append(["- Les IDs évitent les problèmes de doublons de noms"])
    ws_info.append(["- Les doublons (même trajet + moyen + date) seront ignorés"])
    ws_info.append([])
    ws_info.append(["Comment trouver les IDs:"])
    ws_info.append(["1. Allez dans la feuille 'Villages'"])
    ws_info.append(["2. Recherchez le village par son nom (CTRL+F)"])
    ws_info.append(["3. Copiez l'ID dans la colonne A"])

    # Feuille de référence Villages
    ws_villages = wb.create_sheet("Villages")
    ws_villages.append(['ID Village', 'Nom Village', 'Fokontany', 'Commune'])
    villages = Village.objects.select_related('fokontany__commune').all()[:1000]
    for village in villages:
        ws_villages.append([
            str(village.id),
            village.nom_village,
            village.fokontany.nom_fokontany if village.fokontany else '',
            village.fokontany.commune.nom_commune if village.fokontany and village.fokontany.commune else ''
        ])

    # Ajuster les largeurs
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws.column_dimensions[col].width = 25
    ws_info.column_dimensions['A'].width = 80
    ws_villages.column_dimensions['A'].width = 15  # ID Village
    ws_villages.column_dimensions['B'].width = 25  # Nom Village
    ws_villages.column_dimensions['C'].width = 25  # Fokontany
    ws_villages.column_dimensions['D'].width = 25  # Commune

    # Créer la réponse
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=template_deplacements.xlsx'
    wb.save(response)
    return response


@employe_required
def import_excel_deplacements(request):
    """Importer des déplacements depuis Excel"""
    if request.method != 'POST':
        return redirect('planning:deplacements_list')

    from .models import Deplacement
    from apps.geography.models import Village
    from openpyxl import load_workbook
    from datetime import datetime

    try:
        file = request.FILES.get('file')
        if not file:
            messages.error(request, "Aucun fichier fourni")
            return redirect('planning:deplacements_list')

        wb = load_workbook(file)
        ws = wb.active

        success_count = 0
        error_count = 0
        errors = []

        # Parcourir les lignes (skip header)
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not any(row):
                continue

            try:
                # Extraire les données (colonnes: ID Point A, ID Point B, Moyen, Prix, Date)
                point_a_id, point_b_id, moyen_transport, prix_unitaire, date_prix_str = (list(row) + [None] * 5)[:5]

                if not all([point_a_id, point_b_id, moyen_transport, prix_unitaire, date_prix_str]):
                    errors.append(f"Ligne {idx}: Données manquantes")
                    error_count += 1
                    continue

                # Récupérer les villages par ID (UUID)
                try:
                    point_a = Village.objects.get(id=str(point_a_id).strip())
                except (Village.DoesNotExist, ValueError, TypeError):
                    errors.append(f"Ligne {idx}: Village ID '{point_a_id}' introuvable")
                    error_count += 1
                    continue

                try:
                    point_b = Village.objects.get(id=str(point_b_id).strip())
                except (Village.DoesNotExist, ValueError, TypeError):
                    errors.append(f"Ligne {idx}: Village ID '{point_b_id}' introuvable")
                    error_count += 1
                    continue

                # Valider le moyen de transport
                moyen_transport = str(moyen_transport).strip().lower()
                if moyen_transport not in ['taxi-brousse', 'pirogue', 'bac']:
                    errors.append(f"Ligne {idx}: Moyen de transport invalide '{moyen_transport}'")
                    error_count += 1
                    continue

                # Parser la date
                if isinstance(date_prix_str, str):
                    date_prix = datetime.strptime(date_prix_str.strip(), '%Y-%m-%d').date()
                else:
                    date_prix = date_prix_str

                # Créer ou ignorer si existe
                deplacement, created = Deplacement.objects.get_or_create(
                    point_a=point_a,
                    point_b=point_b,
                    moyen_transport=moyen_transport,
                    date_prix=date_prix,
                    defaults={'prix_unitaire': float(prix_unitaire)}
                )

                if created:
                    success_count += 1
                else:
                    errors.append(f"Ligne {idx}: Déplacement déjà existant (ignoré)")

            except Exception as e:
                errors.append(f"Ligne {idx}: {str(e)}")
                error_count += 1

        # Messages de résultat
        if success_count > 0:
            messages.success(request, f"{success_count} déplacement(s) importé(s) avec succès")
        if error_count > 0:
            messages.warning(request, f"{error_count} erreur(s) rencontrée(s)")
            for error in errors[:10]:  # Limiter à 10 erreurs affichées
                messages.error(request, error)

    except Exception as e:
        messages.error(request, f"Erreur lors de l'import: {str(e)}")

    return redirect('planning:deplacements_list')
