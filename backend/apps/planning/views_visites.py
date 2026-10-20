"""
Vues pour le comptage des visites par ouvrage.
Une visite = séquence de dates consécutives dans le planning.
"""
from collections import defaultdict

from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Exists, OuterRef, Q
from django.utils import timezone
import pytz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import Planning
from apps.organization.models import Ouvrage, Organisation, PartenaireGestionnaireAffectation
from apps.organization.scoping import (
    filter_modalites_by_gestionnaire_period,
    filter_ouvrages_by_gestionnaire_period,
    get_user_accessible_gestionnaire_ids,
    get_user_organisation,
)
from apps.billing.models import PointEau, ModaliteGestion
from apps.billing.payment_modalities import build_payment_modality_scope
from apps.geography.models import Region, District, Commune


tz_madagascar = pytz.timezone('Indian/Antananarivo')

def _get_user_gestionnaire_ids(request):
    # Priorite aux affectations explicites partenaire -> gestionnaires.
    # Cela evite les ecarts quand le type_organisation est mal renseigne.
    if getattr(request.user, 'is_superuser', False):
        return None

    organisation = get_user_organisation(request.user)
    if not organisation:
        return []

    affectations_ids = list(
        PartenaireGestionnaireAffectation.objects.filter(
            partenaire=organisation
        ).values_list('gestionnaire_id', flat=True)
    )
    if affectations_ids:
        return affectations_ids

    scoped_ids = get_user_accessible_gestionnaire_ids(request.user)
    if scoped_ids is None:
        return None
    if scoped_ids:
        return list(scoped_ids)

    return [organisation.id]


def _get_organisation_filter_state(request):
    gestionnaire_ids = _get_user_gestionnaire_ids(request)
    if gestionnaire_ids is None:
        raw_ids = None
        normalized_ids = None
        default_org_id = 'all'
        show_organisation_filter = True
        organisations = Organisation.objects.filter(
            type_organisation=Organisation.TYPE_GESTIONNAIRE
        )
    else:
        raw_ids = list(gestionnaire_ids)
        normalized_ids = [str(_id) for _id in raw_ids]

        organisations = Organisation.objects.filter(id__in=raw_ids)
        if normalized_ids:
            default_org_id = normalized_ids[0] if len(normalized_ids) == 1 else 'all'
        else:
            organisations = organisations.none()
            default_org_id = 'all'
        show_organisation_filter = True

    organisation_filter = request.GET.get('organisation', default_org_id) or default_org_id
    if organisation_filter != 'all' and normalized_ids is not None:
        if organisation_filter not in set(normalized_ids):
            organisation_filter = default_org_id

    return {
        'gestionnaire_ids': raw_ids if gestionnaire_ids is not None else None,
        'organisations': organisations.order_by('nom_organisation'),
        'organisation_filter': organisation_filter,
        'show_organisation_filter': show_organisation_filter,
    }


def _apply_organisation_filter(queryset, organisation_filter, annee=None):
    if organisation_filter and organisation_filter != 'all':
        queryset = filter_ouvrages_by_gestionnaire_period(
            queryset,
            organisation_filter,
            annee,
        )
    return queryset


def _base_ouvrages_queryset(request, gestionnaire_ids=None, annee=None):
    qs = Ouvrage.objects.select_related('type_ouvrage')

    qs = qs.annotate(
        _has_point_eau=Exists(
            PointEau.objects.filter(ouvrage_id=OuterRef('pk'))
        )
    ).filter(_has_point_eau=True)

    if gestionnaire_ids is not None:
        if gestionnaire_ids:
            qs = filter_ouvrages_by_gestionnaire_period(
                qs,
                gestionnaire_ids,
                annee,
            )
        else:
            return Ouvrage.objects.none()

    return qs.exclude(
        type_ouvrage__nom_type_ouvrage__iexact='Potentiel'
    ).exclude(
        type_ouvrage__nom_type_ouvrage__iexact='Suspendu'
    )


def _apply_geography_filters(queryset, region_id, district_id, commune_id):
    if not any([region_id, district_id, commune_id]):
        return queryset

    point_eaux = PointEau.objects.filter(ouvrage_id=OuterRef('pk'))
    if commune_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune_id=commune_id
        )
    if district_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune__district_id=district_id
        )
    if region_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune__district__region_id=region_id
        )

    queryset = queryset.annotate(
        _matches_geo_filters=Exists(point_eaux)
    ).filter(_matches_geo_filters=True)
    return queryset


def _apply_annee_realisation_filter(queryset, annee):
    """
    Aligne le scope annuel sur le dashboard:
    - annee_realisation <= annee
    - ou, si annee_realisation est vide, annee_realisation_avant <= annee
    """
    if not annee:
        return queryset
    try:
        annee_int = int(annee)
    except (TypeError, ValueError):
        return queryset

    return queryset.filter(
        Q(annee_realisation__lte=annee_int) |
        Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee_int)
    )


def _get_filtered_ouvrages_queryset(
    request,
    gestionnaire_ids,
    organisation_filter,
    region_filter='',
    district_filter='',
    commune_filter='',
    annee=None,
):
    ouvrages_qs = _base_ouvrages_queryset(
        request,
        gestionnaire_ids,
        annee=annee,
    )
    ouvrages_qs = _apply_organisation_filter(
        ouvrages_qs,
        organisation_filter,
        annee=annee,
    )
    ouvrages_qs = _apply_geography_filters(
        ouvrages_qs,
        region_filter,
        district_filter,
        commune_filter,
    )
    ouvrages_qs = _apply_annee_realisation_filter(ouvrages_qs, annee)
    return ouvrages_qs.order_by('nom_ouvrage')


def _get_modalite_map(ouvrage_ids, organisation_filter='', gestionnaire_ids=None, annee=None):
    if not ouvrage_ids:
        return {}

    modalites_qs = filter_modalites_by_gestionnaire_period(
        ModaliteGestion.objects.filter(ouvrage_id__in=ouvrage_ids),
        organisation_filter if organisation_filter and organisation_filter != 'all' else gestionnaire_ids,
        annee,
    )

    modalite_map = {}
    for row in modalites_qs.values('ouvrage_id', 'modalitepaiement'):
        modalite_map.setdefault(row['ouvrage_id'], row['modalitepaiement'])
    return modalite_map


def _get_ouvrage_geo_map(ouvrage_ids):
    if not ouvrage_ids:
        return {}

    geo_map = {}
    pe_geo = PointEau.objects.filter(
        ouvrage_id__in=ouvrage_ids,
        parcelle__village__fokontany__commune__isnull=False,
    ).values(
        'ouvrage_id',
        'parcelle__village__fokontany__commune__nom_commune',
        'parcelle__village__fokontany__commune__district__nom_district',
        'parcelle__village__fokontany__commune__district__region__nom_region',
    ).distinct()

    for row in pe_geo:
        if row['ouvrage_id'] not in geo_map:
            geo_map[row['ouvrage_id']] = {
                'commune': row['parcelle__village__fokontany__commune__nom_commune'],
                'district': row['parcelle__village__fokontany__commune__district__nom_district'],
                'region': row['parcelle__village__fokontany__commune__district__region__nom_region'],
            }
    return geo_map


def _get_points_map(ouvrage_ids):
    if not ouvrage_ids:
        return {}

    pe_counts = PointEau.objects.filter(
        ouvrage_id__in=ouvrage_ids
    ).values('ouvrage_id').annotate(nb=Count('id'))
    return {row['ouvrage_id']: row['nb'] for row in pe_counts}


def _get_dashboard_points_scope_by_ouvrage(
    ouvrage_ids,
    annee=None,
    region_filter='',
    district_filter='',
    commune_filter='',
):
    if not ouvrage_ids:
        return {
            'modalite_by_ouvrage': {},
            'counts_by_ouvrage': {},
        }

    from hydria.views import _build_dashboard_service_scope

    ouvrages_qs = Ouvrage.objects.filter(id__in=ouvrage_ids)
    points_qs = PointEau.objects.filter(ouvrage_id__in=ouvrage_ids)

    if commune_filter:
        points_qs = points_qs.filter(
            parcelle__village__fokontany__commune_id=commune_filter
        )
    if district_filter:
        points_qs = points_qs.filter(
            parcelle__village__fokontany__commune__district_id=district_filter
        )
    if region_filter:
        points_qs = points_qs.filter(
            parcelle__village__fokontany__commune__district__region_id=region_filter
        )

    if annee:
        try:
            annee_int = int(annee)
            from apps.billing.points_eau_statut import filtrer_points_annee
            points_qs = filtrer_points_annee(points_qs, annee_int)
        except (TypeError, ValueError):
            pass

    _, _, points_actifs_qs, _, _ = _build_dashboard_service_scope(
        points_qs.distinct(),
        ouvrages_qs,
        annee=annee,
    )
    payment_scope = build_payment_modality_scope(points_actifs_qs, annee=annee)

    counts_by_ouvrage = {
        ouvrage_id: {
            'total': 0,
            'forfaitaire': 0,
            'au_volume': 0,
            'sans_modalite': 0,
        }
        for ouvrage_id in ouvrage_ids
    }

    point_modality_map = {}
    for modalite_key in ('forfaitaire', 'au_volume', 'sans_modalite'):
        for point_id in payment_scope['point_ids'][modalite_key]:
            point_modality_map[point_id] = modalite_key

    for point_id, ouvrage_id in points_actifs_qs.values_list('id', 'ouvrage_id'):
        counts = counts_by_ouvrage.setdefault(
            ouvrage_id,
            {
                'total': 0,
                'forfaitaire': 0,
                'au_volume': 0,
                'sans_modalite': 0,
            }
        )
        counts['total'] += 1
        modalite_key = point_modality_map.get(point_id) or 'sans_modalite'
        counts[modalite_key] += 1

    modalite_labels = {
        'forfaitaire': 'Forfaitaire',
        'au_volume': 'Au volume',
        'mixte': 'Mixte',
        'cotisation_annuelle': 'Cotisation annuelle',
        'sans_modalite': 'Sans modalite',
    }
    modalite_by_ouvrage = {}
    for modalite_key, label in modalite_labels.items():
        for ouvrage_id in payment_scope['ouvrage_ids'].get(modalite_key, set()):
            modalite_by_ouvrage[ouvrage_id] = label

    return {
        'modalite_by_ouvrage': modalite_by_ouvrage,
        'counts_by_ouvrage': counts_by_ouvrage,
    }


def _get_filtered_plannings_queryset(ouvrage_ids, annee=None, mois=None, date_debut=None, date_fin=None):
    """
    Filtre les plannings par ouvrage_ids et éventuellement par année/mois ou par période personnalisée.

    Args:
        ouvrage_ids: Liste d'IDs d'ouvrages
        annee: Année (optionnel, ignoré si date_debut/date_fin fournis)
        mois: Mois (optionnel, ignoré si date_debut/date_fin fournis)
        date_debut: Date de début de la période personnalisée (format YYYY-MM-DD)
        date_fin: Date de fin de la période personnalisée (format YYYY-MM-DD)

    Returns:
        QuerySet de Planning triés par ouvrage, date, id
    """
    if not ouvrage_ids:
        return Planning.objects.none()

    queryset = Planning.objects.filter(
        ouvrage_id__in=ouvrage_ids,
        date__isnull=False,
    ).select_related(
        'employe',
        'fiche_poste',
        'type_activite',
    ).order_by('ouvrage_id', 'date', 'id')

    # Priorité à la période personnalisée (date_debut/date_fin)
    if date_debut or date_fin:
        try:
            from datetime import datetime
            if date_debut:
                # Parser la date de début et filtrer date >= date_debut
                date_debut_obj = datetime.strptime(date_debut, '%Y-%m-%d').date()
                queryset = queryset.filter(date__date__gte=date_debut_obj)
            if date_fin:
                # Parser la date de fin et filtrer date <= date_fin
                date_fin_obj = datetime.strptime(date_fin, '%Y-%m-%d').date()
                queryset = queryset.filter(date__date__lte=date_fin_obj)
        except (TypeError, ValueError):
            # En cas d'erreur de format, ignorer les filtres de période personnalisée
            pass
    else:
        # Si pas de période personnalisée, utiliser annee/mois
        try:
            if annee:
                queryset = queryset.filter(date__year=int(annee))
        except (TypeError, ValueError):
            pass

        try:
            if mois:
                queryset = queryset.filter(date__month=int(mois))
        except (TypeError, ValueError):
            pass

    return queryset


def _get_plannings_grouped_by_ouvrage(ouvrage_ids, annee=None, mois=None, date_debut=None, date_fin=None):
    """
    Retourne un dictionnaire {ouvrage_id: [liste de Planning]} trié.

    Args:
        ouvrage_ids: Liste d'IDs d'ouvrages
        annee: Année (optionnel, ignoré si date_debut/date_fin fournis)
        mois: Mois (optionnel, ignoré si date_debut/date_fin fournis)
        date_debut: Date de début de la période personnalisée (format YYYY-MM-DD)
        date_fin: Date de fin de la période personnalisée (format YYYY-MM-DD)
    """
    grouped = defaultdict(list)
    queryset = _get_filtered_plannings_queryset(
        ouvrage_ids,
        annee=annee,
        mois=mois,
        date_debut=date_debut,
        date_fin=date_fin,
    )
    for planning in queryset:
        grouped[planning.ouvrage_id].append(planning)
    return grouped


def _build_ouvrage_stats(
    ouvrages,
    annee=None,
    mois=None,
    organisation_filter='',
    gestionnaire_ids=None,
    include_visites=False,
    nonzero_only=False,
    region_filter='',
    district_filter='',
    commune_filter='',
    date_debut=None,
    date_fin=None,
):
    ouvrage_list = list(ouvrages)
    ouvrage_ids = [ouvrage.id for ouvrage in ouvrage_list]

    geo_map = _get_ouvrage_geo_map(ouvrage_ids)
    points_scope = _get_dashboard_points_scope_by_ouvrage(
        ouvrage_ids,
        annee=annee,
        region_filter=region_filter,
        district_filter=district_filter,
        commune_filter=commune_filter,
    )
    points_counts_by_ouvrage = points_scope['counts_by_ouvrage']
    modalite_map = points_scope['modalite_by_ouvrage']
    plannings_by_ouvrage = _get_plannings_grouped_by_ouvrage(
        ouvrage_ids,
        annee=annee,
        mois=mois,
        date_debut=date_debut,
        date_fin=date_fin,
    )

    stats_ouvrages = []
    for ouvrage in ouvrage_list:
        visites_source = plannings_by_ouvrage.get(ouvrage.id, [])
        nb_visites, visites = compter_visites(visites_source)
        if nonzero_only and nb_visites <= 0:
            continue

        nb_visites_reparation = sum(1 for v in visites if v.get('has_reparation'))
        nb_visites_reparation_panne = sum(1 for v in visites if v.get('has_reparation_panne'))

        geo = geo_map.get(ouvrage.id, {})
        point_counts = points_counts_by_ouvrage.get(
            ouvrage.id,
            {
                'total': 0,
                'forfaitaire': 0,
                'au_volume': 0,
                'sans_modalite': 0,
            }
        )
        nb_points = point_counts.get('total', 0)
        nb_points_forfaitaire = point_counts.get('forfaitaire', 0)
        nb_points_volume = point_counts.get('au_volume', 0)
        nb_points_sans_modalite = point_counts.get('sans_modalite', 0)
        modalite_paiement = modalite_map.get(ouvrage.id) or 'Aucun point actif'
        if nb_points_volume > 10:
            categorie_key = 'au_volume_plus_10'
            categorie_label = "Au volume > 10 points d'eau"
        else:
            categorie_key = 'autres'
            categorie_label = "Autres (au volume <= 10 ou forfaitaire)"

        stat = {
            'ouvrage': ouvrage,
            'nb_visites': nb_visites,
            'nb_visites_reparation': nb_visites_reparation,
            'nb_visites_reparation_panne': nb_visites_reparation_panne,
            'derniere_visite': visites[-1]['date_fin'] if visites else None,
            'premiere_visite': visites[0]['date_debut'] if visites else None,
            'region_nom': geo.get('region') or 'Sans region',
            'district_nom': geo.get('district') or 'Sans district',
            'commune_nom': geo.get('commune') or 'Sans commune',
            'modalite_paiement': modalite_paiement,
            'nb_points_eau': nb_points,
            'nb_points_eau_forfaitaire': nb_points_forfaitaire,
            'nb_points_eau_au_volume': nb_points_volume,
            'nb_points_eau_sans_modalite': nb_points_sans_modalite,
            'categorie_paiement_key': categorie_key,
            'categorie_paiement_label': categorie_label,
        }
        if include_visites:
            stat['visites'] = visites

        stats_ouvrages.append(stat)

    return stats_ouvrages


def compter_visites(plannings_source):
    """
    Compte le nombre de visites distinctes.
    Une visite = séquence de dates consécutives.

    Args:
        plannings_source: QuerySet ou iterable de Planning déjà filtrés

    Returns:
        Tuple (nombre_visites, liste_visites)
        Chaque visite contient: date_debut, date_fin, nb_jours, employes
    """
    # Extraire les dates uniques et trier
    if hasattr(plannings_source, 'order_by'):
        plannings = list(plannings_source.order_by('date', 'id'))
    else:
        plannings = list(plannings_source)

    if not plannings:
        return 0, []

    visites = []

    def _normalize(text):
        if text is None:
            return ''
        nom = str(text).lower()
        import unicodedata
        nom = unicodedata.normalize('NFKD', nom)
        nom = ''.join(ch for ch in nom if not unicodedata.combining(ch))
        return nom

    def _has_reparation(activites):
        for activite in activites:
            nom = _normalize(activite)
            if not nom:
                continue
            if 'au cas ou reparation' in nom:
                continue
            if 'reparation panne' in nom:
                continue
            if 'reparation' in nom or 'rehabilitation' in nom:
                return True
        return False

    def _has_reparation_panne(activites):
        for activite in activites:
            nom = _normalize(activite)
            if not nom:
                continue
            if 'reparation panne' in nom:
                return True
        return False

    def _get_activite_nom(planning):
        if planning.fiche_poste and getattr(planning.fiche_poste, 'tache', None):
            return planning.fiche_poste.tache
        if planning.type_activite and getattr(planning.type_activite, 'nomtype', None):
            return planning.type_activite.nomtype
        return None

    # Initialiser la première visite
    current_visit = {
        'date_debut': plannings[0].date.date() if plannings[0].date else None,
        'date_fin': plannings[0].date.date() if plannings[0].date else None,
        'employes': set(),
        'activites': set(),
        'has_reparation': False,
        'plannings': [plannings[0]]
    }

    if plannings[0].employe:
        current_visit['employes'].add(plannings[0].employe.surnom or plannings[0].employe.prenom or plannings[0].employe.nom)
    activite_nom = _get_activite_nom(plannings[0])
    if activite_nom:
        current_visit['activites'].add(activite_nom)

    previous_date = current_visit['date_debut']

    for planning in plannings[1:]:
        current_date = planning.date.date() if planning.date else None

        if current_date is None:
            continue

        # Vérifier si la date est consécutive (ou même jour)
        if previous_date and (current_date - previous_date).days <= 1:
            # Même visite - étendre
            current_visit['date_fin'] = current_date
            current_visit['plannings'].append(planning)
            if planning.employe:
                current_visit['employes'].add(planning.employe.surnom or planning.employe.prenom or planning.employe.nom)
            activite_nom = _get_activite_nom(planning)
            if activite_nom:
                current_visit['activites'].add(activite_nom)
        else:
            # Nouvelle visite - sauvegarder l'ancienne
            current_visit['nb_jours'] = (current_visit['date_fin'] - current_visit['date_debut']).days + 1 if current_visit['date_debut'] and current_visit['date_fin'] else 1
            current_visit['employes'] = list(current_visit['employes'])
            current_visit['activites'] = list(current_visit['activites'])
            current_visit['has_reparation'] = _has_reparation(current_visit['activites'])
            current_visit['has_reparation_panne'] = _has_reparation_panne(current_visit['activites'])
            visites.append(current_visit)

            # Démarrer nouvelle visite
            current_visit = {
                'date_debut': current_date,
                'date_fin': current_date,
                'employes': set(),
                'activites': set(),
                'plannings': [planning]
            }
            if planning.employe:
                current_visit['employes'].add(planning.employe.surnom or planning.employe.prenom or planning.employe.nom)
            activite_nom = _get_activite_nom(planning)
            if activite_nom:
                current_visit['activites'].add(activite_nom)

        previous_date = current_date

    # Ne pas oublier la dernière visite
    current_visit['nb_jours'] = (current_visit['date_fin'] - current_visit['date_debut']).days + 1 if current_visit['date_debut'] and current_visit['date_fin'] else 1
    current_visit['employes'] = list(current_visit['employes'])
    current_visit['activites'] = list(current_visit['activites'])
    current_visit['has_reparation'] = _has_reparation(current_visit['activites'])
    current_visit['has_reparation_panne'] = _has_reparation_panne(current_visit['activites'])
    visites.append(current_visit)

    return len(visites), visites


@login_required
def visites_par_ouvrage(request):
    """
    Page principale affichant le nombre de visites par ouvrage.
    """
    # Filtres
    annee = request.GET.get('annee')
    mois = request.GET.get('mois')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')
    categorie_paiement_filter = request.GET.get('categorie_paiement', '')

    # Récupérer les ouvrages avec leurs plannings
    scope_info = _get_organisation_filter_state(request)
    gestionnaire_ids = scope_info['gestionnaire_ids']
    organisation_filter = scope_info['organisation_filter']
    organisations = scope_info['organisations']
    show_organisation_filter = scope_info['show_organisation_filter']
    ouvrages = _get_filtered_ouvrages_queryset(
        request,
        gestionnaire_ids,
        organisation_filter,
        region_filter,
        district_filter,
        commune_filter,
        annee,
    )

    stats_ouvrages = _build_ouvrage_stats(
        ouvrages,
        annee=annee,
        mois=mois,
        date_debut=date_debut,
        date_fin=date_fin,
        organisation_filter=organisation_filter,
        gestionnaire_ids=gestionnaire_ids,
        region_filter=region_filter,
        district_filter=district_filter,
        commune_filter=commune_filter,
    )

    if categorie_paiement_filter:
        stats_ouvrages = [
            stat for stat in stats_ouvrages
            if stat.get('categorie_paiement_key') == categorie_paiement_filter
        ]

    total_visites = sum(stat['nb_visites'] for stat in stats_ouvrages)
    total_visites_reparation = sum(stat['nb_visites_reparation'] for stat in stats_ouvrages)
    total_visites_reparation_panne = sum(stat['nb_visites_reparation_panne'] for stat in stats_ouvrages)
    total_points_eau = sum(int(stat.get('nb_points_eau') or 0) for stat in stats_ouvrages)
    total_points_eau_forfaitaire = sum(int(stat.get('nb_points_eau_forfaitaire') or 0) for stat in stats_ouvrages)
    total_points_eau_au_volume = sum(int(stat.get('nb_points_eau_au_volume') or 0) for stat in stats_ouvrages)
    total_points_eau_sans_modalite = sum(int(stat.get('nb_points_eau_sans_modalite') or 0) for stat in stats_ouvrages)

    # Trier pour le regroupement Region -> District -> Commune -> Categorie
    stats_ouvrages.sort(
        key=lambda x: (
            (x.get('region_nom') or '').lower(),
            (x.get('district_nom') or '').lower(),
            (x.get('commune_nom') or '').lower(),
            (x.get('categorie_paiement_label') or '').lower(),
            (x.get('ouvrage').nom_ouvrage or '').lower(),
        )
    )
    max_visites = max([stat['nb_visites'] for stat in stats_ouvrages], default=0)
    max_visites_reparation = max([stat['nb_visites_reparation'] for stat in stats_ouvrages], default=0)
    max_visites_reparation_panne = max([stat['nb_visites_reparation_panne'] for stat in stats_ouvrages], default=0)
    max_visites_global = max(max_visites, max_visites_reparation, max_visites_reparation_panne)

    chart_mode = 'ouvrage' if commune_filter else 'commune'
    chart_title = 'Histogramme des visites par ouvrage'
    visites_chart_data = []

    commune_stats = {}
    commune_ouvrages = {}
    for stat in stats_ouvrages:
        commune_name = stat.get('commune_nom') or 'Sans commune'
        commune_stats[commune_name] = commune_stats.get(commune_name, 0) + stat['nb_visites']
        commune_ouvrages.setdefault(commune_name, set()).add(stat['ouvrage'].id)

    visites_commune_totals = [
        {'nom': name, 'nb_visites': count}
        for name, count in sorted(commune_stats.items(), key=lambda item: item[1], reverse=True)
    ]

    visites_commune_avg_chart_data = [
        {
            'nom': name,
            'nb_visites': round((commune_stats.get(name, 0) / max(len(commune_ouvrages.get(name, [])), 1)), 2),
        }
        for name in sorted(commune_stats.keys(), key=lambda n: commune_stats[n], reverse=True)
    ]

    if chart_mode == 'ouvrage':
        visites_chart_data = [
            {
                'nom': stat['ouvrage'].nom_ouvrage,
                'nb_visites': stat['nb_visites'],
            }
            for stat in stats_ouvrages
        ]
    else:
        chart_title = 'Histogramme des visites par commune'
        visites_chart_data = visites_commune_totals
    visites_counts = [stat['nb_visites'] for stat in stats_ouvrages]
    if visites_counts:
        sorted_counts = sorted(visites_counts)
        mid = len(sorted_counts) // 2
        if len(sorted_counts) % 2 == 1:
            median_visites = sorted_counts[mid]
        else:
            median_visites = (sorted_counts[mid - 1] + sorted_counts[mid]) / 2
    else:
        median_visites = 0

    # Récupérer les années disponibles pour le filtre
    annees_disponibles = Planning.objects.filter(
        date__isnull=False
    ).dates('date', 'year', order='DESC')
    annees_list = [d.year for d in annees_disponibles]

    moyenne_visites = (total_visites / len(stats_ouvrages)) if stats_ouvrages else 0

    def _group_stats(items, key):
        grouped = {}
        for stat in items:
            name = stat.get(key) or 'Non defini'
            if name not in grouped:
                grouped[name] = {
                    'nom': name,
                    'nb_ouvrages': 0,
                    'nb_visites': 0,
                    'nb_visites_reparation': 0,
                    'nb_visites_reparation_panne': 0,
                }
            grouped[name]['nb_ouvrages'] += 1
            grouped[name]['nb_visites'] += stat['nb_visites']
            grouped[name]['nb_visites_reparation'] += stat['nb_visites_reparation']
            grouped[name]['nb_visites_reparation_panne'] += stat['nb_visites_reparation_panne']
        return sorted(grouped.values(), key=lambda x: x['nb_visites'], reverse=True)

    stats_par_region = _group_stats(stats_ouvrages, 'region_nom')
    stats_par_district = _group_stats(stats_ouvrages, 'district_nom')
    stats_par_commune = _group_stats(stats_ouvrages, 'commune_nom')
    stats_par_categorie = _group_stats(stats_ouvrages, 'categorie_paiement_label')

    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.select_related('region').order_by('region__nom_region', 'nom_district')
    communes = Commune.objects.select_related('district').order_by('district__nom_district', 'nom_commune')

    context = {
        'stats_ouvrages': stats_ouvrages,
        'total_visites': total_visites,
        'total_visites_reparation': total_visites_reparation,
        'total_visites_reparation_panne': total_visites_reparation_panne,
        'total_points_eau': total_points_eau,
        'total_points_eau_forfaitaire': total_points_eau_forfaitaire,
        'total_points_eau_au_volume': total_points_eau_au_volume,
        'total_points_eau_sans_modalite': total_points_eau_sans_modalite,
        'nb_ouvrages': len(stats_ouvrages),
        'annee_selectionnee': annee,
        'mois_selectionne': mois,
        'date_debut_selectionnee': date_debut or '',
        'date_fin_selectionnee': date_fin or '',
        'periode_personnalisee_active': bool(date_debut or date_fin),
        'organisation_filter': organisation_filter,
        'organisations': organisations,
        'show_organisation_filter': show_organisation_filter,
        'region_selectionnee': region_filter,
        'district_selectionne': district_filter,
        'commune_selectionnee': commune_filter,
        'categorie_paiement_selectionnee': categorie_paiement_filter,
        'annees_disponibles': annees_list,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'moyenne_visites': moyenne_visites,
        'median_visites': median_visites,
        'visites_chart_data': visites_chart_data,
        'visites_commune_avg_chart_data': visites_commune_avg_chart_data,
        'max_visites': max_visites,
        'max_visites_reparation': max_visites_reparation,
        'max_visites_reparation_panne': max_visites_reparation_panne,
        'max_visites_global': max_visites_global,
        'chart_title': chart_title,
        'stats_par_region': stats_par_region,
        'stats_par_district': stats_par_district,
        'stats_par_commune': stats_par_commune,
        'stats_par_categorie': stats_par_categorie,
    }

    return render(request, 'planning/visites_par_ouvrage.html', context)


@login_required
def visites_global(request):
    """Vue globale des visites, regroupement dynamique par zone ou categorie."""
    annee = request.GET.get('annee')
    mois = request.GET.get('mois')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')
    categorie_paiement_filter = request.GET.get('categorie_paiement', '')
    group_by = request.GET.get('group_by', 'commune')

    scope_info = _get_organisation_filter_state(request)
    gestionnaire_ids = scope_info['gestionnaire_ids']
    organisation_filter = scope_info['organisation_filter']
    organisations = scope_info['organisations']
    show_organisation_filter = scope_info['show_organisation_filter']
    ouvrages = _get_filtered_ouvrages_queryset(
        request,
        gestionnaire_ids,
        organisation_filter,
        region_filter,
        district_filter,
        commune_filter,
        annee,
    )

    stats_ouvrages = _build_ouvrage_stats(
        ouvrages,
        annee=annee,
        mois=mois,
        date_debut=date_debut,
        date_fin=date_fin,
        organisation_filter=organisation_filter,
        gestionnaire_ids=gestionnaire_ids,
        nonzero_only=True,
        region_filter=region_filter,
        district_filter=district_filter,
        commune_filter=commune_filter,
    )

    if categorie_paiement_filter:
        stats_ouvrages = [
            stat for stat in stats_ouvrages
            if stat.get('categorie_paiement_key') == categorie_paiement_filter
        ]

    stats_ouvrages.sort(
        key=lambda x: (
            (x.get('region_nom') or '').lower(),
            (x.get('district_nom') or '').lower(),
            (x.get('commune_nom') or '').lower(),
            (x.get('categorie_paiement_label') or '').lower(),
            (x.get('ouvrage').nom_ouvrage or '').lower(),
        )
    )

    group_key_map = {
        'region': 'region_nom',
        'district': 'district_nom',
        'commune': 'commune_nom',
        'categorie_paiement': 'categorie_paiement_label',
    }
    group_key = group_key_map.get(group_by, 'commune_nom')

    grouped = {}
    for stat in stats_ouvrages:
        name = stat.get(group_key) or 'Non defini'
        if name not in grouped:
            grouped[name] = {
                'nom': name,
                'nb_ouvrages': 0,
                'nb_visites': 0,
                'nb_visites_reparation': 0,
                'nb_visites_reparation_panne': 0,
                'ouvrages': [],
            }
        grouped[name]['nb_ouvrages'] += 1
        grouped[name]['nb_visites'] += stat['nb_visites']
        grouped[name]['nb_visites_reparation'] += stat['nb_visites_reparation']
        grouped[name]['nb_visites_reparation_panne'] += stat['nb_visites_reparation_panne']
        grouped[name]['ouvrages'].append(stat)

    groups = sorted(grouped.values(), key=lambda x: x['nom'])
    for group in groups:
        nb_ouvrages = group['nb_ouvrages'] or 1
        group['avg_visites'] = round(group['nb_visites'] / nb_ouvrages, 2)
        group['avg_reparation'] = round(group['nb_visites_reparation'] / nb_ouvrages, 2)
        group['avg_panne'] = round(group['nb_visites_reparation_panne'] / nb_ouvrages, 2)

    max_visites = max([stat['nb_visites'] for stat in stats_ouvrages], default=0)
    max_visites_reparation = max([stat['nb_visites_reparation'] for stat in stats_ouvrages], default=0)
    max_visites_reparation_panne = max([stat['nb_visites_reparation_panne'] for stat in stats_ouvrages], default=0)

    annees_disponibles = Planning.objects.filter(
        date__isnull=False
    ).dates('date', 'year', order='DESC')
    annees_list = [d.year for d in annees_disponibles]

    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.select_related('region').order_by('region__nom_region', 'nom_district')
    communes = Commune.objects.select_related('district').order_by('district__nom_district', 'nom_commune')

    context = {
        'groups': groups,
        'annees_disponibles': annees_list,
        'annee_selectionnee': annee,
        'mois_selectionne': mois,
        'date_debut_selectionnee': date_debut or '',
        'date_fin_selectionnee': date_fin or '',
        'periode_personnalisee_active': bool(date_debut or date_fin),
        'organisation_filter': organisation_filter,
        'organisations': organisations,
        'show_organisation_filter': show_organisation_filter,
        'region_selectionnee': region_filter,
        'district_selectionne': district_filter,
        'commune_selectionnee': commune_filter,
        'categorie_paiement_selectionnee': categorie_paiement_filter,
        'group_by': group_by,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'max_visites': max_visites or 1,
        'max_visites_reparation': max_visites_reparation or 1,
        'max_visites_reparation_panne': max_visites_reparation_panne or 1,
    }

    return render(request, 'planning/visites_global.html', context)


@login_required
def api_visites_par_ouvrage(request):
    """
    API retournant les statistiques de visites par ouvrage.
    """
    annee = request.GET.get('annee')
    mois = request.GET.get('mois')
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')
    ouvrage_id = request.GET.get('ouvrage_id')

    # Construire le queryset de base
    scope_info = _get_organisation_filter_state(request)
    gestionnaire_ids = scope_info['gestionnaire_ids']
    organisation_filter = scope_info['organisation_filter']
    if ouvrage_id:
        ouvrages = _get_filtered_ouvrages_queryset(
            request,
            gestionnaire_ids,
            organisation_filter,
            region_filter,
            district_filter,
            commune_filter,
            annee,
        ).filter(id=ouvrage_id)
    else:
        ouvrages = _get_filtered_ouvrages_queryset(
            request,
            gestionnaire_ids,
            organisation_filter,
            region_filter,
            district_filter,
            commune_filter,
            annee,
        )

    raw_stats = _build_ouvrage_stats(
        ouvrages,
        annee=annee,
        mois=mois,
        organisation_filter=organisation_filter,
        gestionnaire_ids=gestionnaire_ids,
        include_visites=True,
        nonzero_only=True,
        region_filter=region_filter,
        district_filter=district_filter,
        commune_filter=commune_filter,
    )

    stats_ouvrages = []
    total_visites = 0
    total_visites_reparation = 0
    total_visites_reparation_panne = 0

    for stat in raw_stats:
        visites = stat.get('visites', [])
        visites_serialized = []
        for visite in visites:
            visites_serialized.append({
                'date_debut': visite['date_debut'].strftime('%d/%m/%Y') if visite['date_debut'] else None,
                'date_fin': visite['date_fin'].strftime('%d/%m/%Y') if visite['date_fin'] else None,
                'nb_jours': visite['nb_jours'],
                'employes': visite['employes'],
                'activites': visite['activites'],
            })

        ouvrage = stat['ouvrage']
        stats_ouvrages.append({
            'ouvrage_id': str(ouvrage.id),
            'ouvrage_nom': ouvrage.nom_ouvrage,
            'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else None,
            'nb_visites': stat['nb_visites'],
            'nb_visites_reparation': stat['nb_visites_reparation'],
            'nb_visites_reparation_panne': stat['nb_visites_reparation_panne'],
            'visites': visites_serialized,
            'derniere_visite': stat['derniere_visite'].strftime('%d/%m/%Y') if stat['derniere_visite'] else None,
            'premiere_visite': stat['premiere_visite'].strftime('%d/%m/%Y') if stat['premiere_visite'] else None,
        })
        total_visites += stat['nb_visites']
        total_visites_reparation += stat['nb_visites_reparation']
        total_visites_reparation_panne += stat['nb_visites_reparation_panne']

    # Trier par nombre de visites décroissant
    stats_ouvrages.sort(key=lambda x: x['nb_visites'], reverse=True)

    return JsonResponse({
        'success': True,
        'stats_ouvrages': stats_ouvrages,
        'total_visites': total_visites,
        'total_visites_reparation': total_visites_reparation,
        'total_visites_reparation_panne': total_visites_reparation_panne,
        'nb_ouvrages': len(stats_ouvrages),
    })


@login_required
def api_detail_visites_ouvrage(request, ouvrage_id):
    """
    API retournant le détail des visites pour un ouvrage spécifique.
    """
    annee = request.GET.get('annee')
    mois = request.GET.get('mois')
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')

    scope_info = _get_organisation_filter_state(request)
    gestionnaire_ids = scope_info['gestionnaire_ids']
    organisation_filter = scope_info['organisation_filter']

    ouvrage = _get_filtered_ouvrages_queryset(
        request,
        gestionnaire_ids,
        organisation_filter,
        region_filter,
        district_filter,
        commune_filter,
        annee,
    ).filter(id=ouvrage_id).first()
    if not ouvrage:
        return JsonResponse({'success': False, 'error': 'Ouvrage non trouvé'}, status=404)

    plannings_by_ouvrage = _get_plannings_grouped_by_ouvrage(
        [ouvrage.id],
        annee=annee,
        mois=mois,
    )
    nb_visites, visites = compter_visites(plannings_by_ouvrage.get(ouvrage.id, []))

    visites_detail = []
    for i, v in enumerate(visites, 1):
        visites_detail.append({
            'numero': i,
            'date_debut': v['date_debut'].strftime('%d/%m/%Y') if v['date_debut'] else None,
            'date_fin': v['date_fin'].strftime('%d/%m/%Y') if v['date_fin'] else None,
            'nb_jours': v['nb_jours'],
            'employes': v['employes'],
            'activites': v['activites'],
            'has_reparation': v.get('has_reparation', False),
            'has_reparation_panne': v.get('has_reparation_panne', False),
        })

    return JsonResponse({
        'success': True,
        'ouvrage': {
            'id': str(ouvrage.id),
            'nom': ouvrage.nom_ouvrage,
            'type': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else None,
        },
        'nb_visites': nb_visites,
        'nb_visites_reparation': sum(1 for v in visites if v.get('has_reparation')),
        'nb_visites_reparation_panne': sum(1 for v in visites if v.get('has_reparation_panne')),
        'visites': visites_detail,
    })


@login_required
def api_export_visites_excel(request):
    """
    Export Excel exploitable des statistiques visites par ouvrage.
    """
    annee = request.GET.get('annee')
    mois = request.GET.get('mois')
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')
    categorie_paiement_filter = request.GET.get('categorie_paiement', '')

    scope_info = _get_organisation_filter_state(request)
    gestionnaire_ids = scope_info['gestionnaire_ids']
    organisation_filter = scope_info['organisation_filter']
    ouvrages = _get_filtered_ouvrages_queryset(
        request,
        gestionnaire_ids,
        organisation_filter,
        region_filter,
        district_filter,
        commune_filter,
        annee,
    )

    stats_ouvrages = _build_ouvrage_stats(
        ouvrages,
        annee=annee,
        mois=mois,
        organisation_filter=organisation_filter,
        gestionnaire_ids=gestionnaire_ids,
        region_filter=region_filter,
        district_filter=district_filter,
        commune_filter=commune_filter,
    )

    if categorie_paiement_filter:
        stats_ouvrages = [
            stat for stat in stats_ouvrages
            if stat['categorie_paiement_key'] == categorie_paiement_filter
        ]

    rows = []
    for stat in stats_ouvrages:
        ouvrage = stat['ouvrage']
        rows.append({
            'region': stat['region_nom'],
            'district': stat['district_nom'],
            'commune': stat['commune_nom'],
            'ouvrage': ouvrage.nom_ouvrage,
            'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
            'modalite_paiement': stat['modalite_paiement'],
            'categorie_paiement': stat['categorie_paiement_label'],
            'nb_points_eau': stat['nb_points_eau'],
            'nb_points_eau_forfaitaire': stat['nb_points_eau_forfaitaire'],
            'nb_points_eau_au_volume': stat['nb_points_eau_au_volume'],
            'nb_points_eau_sans_modalite': stat['nb_points_eau_sans_modalite'],
            'nb_visites': stat['nb_visites'],
            'nb_visites_reparation': stat['nb_visites_reparation'],
            'nb_visites_reparation_panne': stat['nb_visites_reparation_panne'],
            'premiere_visite': stat['premiere_visite'].strftime('%d/%m/%Y') if stat['premiere_visite'] else '',
            'derniere_visite': stat['derniere_visite'].strftime('%d/%m/%Y') if stat['derniere_visite'] else '',
        })

    rows.sort(key=lambda r: (r['region'].lower(), r['district'].lower(), r['commune'].lower(), r['ouvrage'].lower()))

    synthese_modalites = {
        'ouvrages_forfaitaire': sum(1 for stat in stats_ouvrages if stat['modalite_paiement'] == 'Forfaitaire'),
        'ouvrages_au_volume': sum(1 for stat in stats_ouvrages if stat['modalite_paiement'] == 'Au volume'),
        'ouvrages_mixte': sum(1 for stat in stats_ouvrages if stat['modalite_paiement'] == 'Mixte'),
        'ouvrages_sans_point_actif': sum(1 for stat in stats_ouvrages if stat['modalite_paiement'] == 'Aucun point actif'),
        'points_forfaitaire': sum(int(stat.get('nb_points_eau_forfaitaire') or 0) for stat in stats_ouvrages),
        'points_au_volume': sum(int(stat.get('nb_points_eau_au_volume') or 0) for stat in stats_ouvrages),
        'points_sans_modalite': sum(int(stat.get('nb_points_eau_sans_modalite') or 0) for stat in stats_ouvrages),
        'points_total_actifs': sum(int(stat.get('nb_points_eau') or 0) for stat in stats_ouvrages),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Visites Ouvrages"

    headers = [
        'Region', 'District', 'Commune', 'Ouvrage', 'Type ouvrage',
        'Modalite paiement', 'Categorie paiement', 'Nb points eau',
        'Nb points forfaitaire', 'Nb points au volume', 'Nb points sans modalite',
        'Nb visites', 'Visites reparation', 'Reparation panne',
        'Premiere visite', 'Derniere visite'
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r in rows:
        ws.append([
            r['region'],
            r['district'],
            r['commune'],
            r['ouvrage'],
            r['type_ouvrage'],
            r['modalite_paiement'],
            r['categorie_paiement'],
            r['nb_points_eau'],
            r['nb_points_eau_forfaitaire'],
            r['nb_points_eau_au_volume'],
            r['nb_points_eau_sans_modalite'],
            r['nb_visites'],
            r['nb_visites_reparation'],
            r['nb_visites_reparation_panne'],
            r['premiere_visite'],
            r['derniere_visite'],
        ])

    widths = {
        1: 20, 2: 22, 3: 22, 4: 30, 5: 18, 6: 20, 7: 35,
        8: 14, 9: 18, 10: 18, 11: 20, 12: 12, 13: 16, 14: 16, 15: 16, 16: 16
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = width

    ws_summary = wb.create_sheet(title="Synthese Modalites")
    ws_summary.append(['Indicateur', 'Valeur'])
    ws_summary['A1'].font = Font(bold=True)
    ws_summary['B1'].font = Font(bold=True)
    ws_summary.append(["Ouvrages forfaitaire", synthese_modalites['ouvrages_forfaitaire']])
    ws_summary.append(["Ouvrages au volume", synthese_modalites['ouvrages_au_volume']])
    ws_summary.append(["Ouvrages mixte", synthese_modalites['ouvrages_mixte']])
    ws_summary.append(["Ouvrages sans point actif", synthese_modalites['ouvrages_sans_point_actif']])
    ws_summary.append(["Points d'eau forfaitaire", synthese_modalites['points_forfaitaire']])
    ws_summary.append(["Points d'eau au volume", synthese_modalites['points_au_volume']])
    ws_summary.append(["Points d'eau sans modalite", synthese_modalites['points_sans_modalite']])
    ws_summary.append(["Points d'eau actifs total", synthese_modalites['points_total_actifs']])
    ws_summary.append([])
    ws_summary.append([
        "Note",
        "La synthese utilise le meme scope annee/modalite que la repartition par modalite de paiement du dashboard."
    ])
    ws_summary.column_dimensions['A'].width = 34
    ws_summary.column_dimensions['B'].width = 110

    filename = f"visites_ouvrages_{annee or 'toutes_annees'}_{region_filter or 'toutes_regions'}_{district_filter or 'tous_districts'}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
