from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.db import models
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
import json
import math
from decimal import Decimal, InvalidOperation
from .models import ExProvince, Region, District, Commune, Fokontany, Village, LiaisonVillage, Parcelle
from .forms import ExProvinceForm, RegionForm, DistrictForm, CommuneForm, FokontanyForm, VillageForm, ParcelleForm
from .services.habitants import habitants_par_zone
from apps.billing.models import PointEau, TypePointEau
from django.utils import timezone
from datetime import datetime
import pytz
from django.conf import settings


def user_is_director(user):
    """Détermine si l'utilisateur actif occupe le poste 'Directeur'."""
    if not user.is_authenticated:
        return False
    employe = getattr(user, 'employe', None)
    if not employe or not hasattr(employe, 'poste') or not employe.poste:
        return False
    poste_nom = getattr(employe.poste, 'nom_poste', '') or ''
    return poste_nom.strip().lower() == 'directeur'


def get_local_timezone():
    """Retourne la timezone locale configurée dans Django"""
    return pytz.timezone(settings.TIME_ZONE)


def parse_local_date(date_string):
    """Parse une date string et la convertit en datetime avec timezone locale"""
    if not date_string:
        return None

    # Parse la date (format YYYY-MM-DD du frontend)
    naive_datetime = datetime.fromisoformat(date_string)

    # Si c'est juste une date, on considère le début de journée
    if naive_datetime.time() == datetime.min.time():
        naive_datetime = naive_datetime.replace(hour=8, minute=0, second=0)  # 8h du matin local

    # Localiser dans la timezone de Madagascar
    local_tz = get_local_timezone()
    localized_datetime = local_tz.localize(naive_datetime)

    return localized_datetime


def get_local_now():
    """Retourne l'heure actuelle dans la timezone locale"""
    return timezone.now().astimezone(get_local_timezone())


def _parse_optional_decimal(value, field_label):
    """Convertit une valeur numerique optionnelle en Decimal."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        value = value.replace(',', '.')

    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValueError(f"Valeur invalide pour {field_label}.") from exc


def _get_geography_photo(data):
    photo = (data.get('photo') or '').strip()
    if not photo:
        return None
    if not photo.startswith('data:image/') or ';base64,' not in photo[:100]:
        raise ValueError("Le format de la photo n'est pas valide.")
    if len(photo) > 3_000_000:
        raise ValueError("La photo est trop volumineuse après compression.")
    return photo


def _get_village_payload(data):
    """Extrait et valide les champs modifiables d'un village."""
    nom_village = (data.get('name') or '').strip()
    if not nom_village:
        raise ValueError("Le nom du village est obligatoire.")

    accessibilite = (data.get('accessibilite') or '').strip() or None
    valid_accessibilites = {choice[0] for choice in Village.ACCESSIBILITE_CHOICES}
    if accessibilite and accessibilite not in valid_accessibilites:
        raise ValueError("Accessibilite invalide.")

    distance_dernier_point = _parse_optional_decimal(
        data.get('distance_dernier_point'),
        "la distance depuis le dernier point"
    )
    if distance_dernier_point is not None and distance_dernier_point < 0:
        raise ValueError("La distance depuis le dernier point ne peut pas etre negative.")

    return {
        'nom_village': nom_village,
        'latitude': _parse_optional_decimal(data.get('latitude'), "la latitude"),
        'longitude': _parse_optional_decimal(data.get('longitude'), "la longitude"),
        'altitude': _parse_optional_decimal(data.get('altitude'), "l'altitude"),
        'accessibilite': accessibilite,
        'distance_dernier_point': distance_dernier_point,
    }


def _serialize_village(village):
    """Serialise un village pour les APIs JSON du module geography."""
    parcelles_count = getattr(village, 'parcelles_count', village.parcelles.count())
    habitants_count = getattr(village, 'habitants_count', None)
    if habitants_count is None:
        habitants_count = (
            village.parcelles.aggregate(total=models.Sum('nombre_habitants'))['total'] or 0
        )

    return {
        'id': str(village.id),
        'nom_village': village.nom_village,
        'photo': village.photo or '',
        'fokontany_id': str(village.fokontany_id) if village.fokontany_id else None,
        'fokontany_name': village.fokontany.nom_fokontany if village.fokontany_id else None,
        'commune_id': str(village.fokontany.commune_id) if village.fokontany_id else None,
        'commune_name': village.fokontany.commune.nom_commune if village.fokontany_id else None,
        'parcelles_count': parcelles_count,
        'habitants_count': habitants_count,
        'latitude': float(village.latitude) if village.latitude is not None else None,
        'longitude': float(village.longitude) if village.longitude is not None else None,
        'altitude': float(village.altitude) if village.altitude is not None else None,
        'accessibilite': village.accessibilite,
        'accessibilite_display': village.get_accessibilite_display() if village.accessibilite else None,
        'distance_dernier_point': (
            float(village.distance_dernier_point)
            if village.distance_dernier_point is not None else None
        ),
        'enclavement': village.enclavement,
        'enclavement_badge_class': village.get_enclavement_badge_class(),
    }


def _serialize_liaison_village(liaison, current_village_id=None):
    """Serialise une liaison pour le detail village."""
    is_outgoing = liaison.village_source_id == current_village_id
    other_village = liaison.village_cible if is_outgoing else liaison.village_source
    air_distance_km = _calculate_air_distance_km(
        liaison.village_source.latitude,
        liaison.village_source.longitude,
        liaison.village_cible.latitude,
        liaison.village_cible.longitude,
    )
    return {
        'id': str(liaison.id),
        'village_source_id': str(liaison.village_source_id),
        'village_cible_id': str(liaison.village_cible_id),
        'type_liaison': liaison.type_liaison,
        'type_liaison_display': liaison.get_type_liaison_display(),
        'type_route': liaison.type_route,
        'type_route_display': liaison.get_type_route_display() if liaison.type_route else None,
        'distance_km': float(liaison.distance_km) if liaison.distance_km is not None else None,
        'distance_vol_oiseau_km': air_distance_km,
        'temps_acces_minutes': liaison.temps_acces_minutes,
        'distance_route_km': liaison.distance_route_km,
        'duree_route_min': liaison.duree_route_min,
        'route_statut': liaison.route_statut,
        'moyen_acces': liaison.moyen_acces,
        'moyen_acces_display': liaison.get_moyen_acces_display() if liaison.moyen_acces else None,
        'ordre': liaison.ordre,
        'sens': 'sortant' if is_outgoing else 'entrant',
        'autre_village': {
            'id': str(other_village.id),
            'nom_village': other_village.nom_village,
            'fokontany': other_village.fokontany.nom_fokontany if other_village.fokontany_id else None,
            'commune': other_village.commune.nom_commune if other_village.fokontany_id else None,
        },
    }


def _calculate_air_distance_km(lat1, lon1, lat2, lon2):
    """Calcule la distance a vol d'oiseau entre deux points GPS."""
    if None in (lat1, lon1, lat2, lon2):
        return None

    lat1 = math.radians(float(lat1))
    lon1 = math.radians(float(lon1))
    lat2 = math.radians(float(lat2))
    lon2 = math.radians(float(lon2))

    delta_lat = lat2 - lat1
    delta_lon = lon2 - lon1
    a = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat1) * math.cos(lat2) * math.sin(delta_lon / 2) ** 2
    )
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return round(6371 * c, 2)


def _build_village_organigram(village, incoming_links, outgoing_links):
    """Construit une vue simple nodes/edges autour d'un village pivot."""
    nodes = [{
        'id': str(village.id),
        'nom': village.nom_village,
        'commune': village.commune.nom_commune if village.fokontany_id else None,
        'role': 'current',
    }]
    seen_node_ids = {str(village.id)}
    edges = []

    for liaison in incoming_links:
        source = liaison.village_source
        node_id = str(source.id)
        if node_id not in seen_node_ids:
            nodes.append({
                'id': node_id,
                'nom': source.nom_village,
                'commune': source.commune.nom_commune if source.fokontany_id else None,
                'role': 'incoming',
            })
            seen_node_ids.add(node_id)
        edges.append({
            'id': str(liaison.id),
            'source': node_id,
            'target': str(village.id),
            'source_name': source.nom_village,
            'target_name': village.nom_village,
            'label': liaison.get_type_liaison_display(),
        })

    for liaison in outgoing_links:
        target = liaison.village_cible
        node_id = str(target.id)
        if node_id not in seen_node_ids:
            nodes.append({
                'id': node_id,
                'nom': target.nom_village,
                'commune': target.commune.nom_commune if target.fokontany_id else None,
                'role': 'outgoing',
            })
            seen_node_ids.add(node_id)
        edges.append({
            'id': str(liaison.id),
            'source': str(village.id),
            'target': node_id,
            'source_name': village.nom_village,
            'target_name': target.nom_village,
            'label': liaison.get_type_liaison_display(),
        })

    return {'nodes': nodes, 'edges': edges}


def _get_linked_village_ids(village, incoming_links=None, outgoing_links=None):
    """Retourne les ids des villages deja lies au village courant."""
    if incoming_links is None:
        incoming_links = LiaisonVillage.objects.filter(village_cible=village, actif=True)
    if outgoing_links is None:
        outgoing_links = LiaisonVillage.objects.filter(village_source=village, actif=True)

    return {liaison.village_source_id for liaison in incoming_links} | {
        liaison.village_cible_id for liaison in outgoing_links
    }


def _get_nearest_unlinked_villages(village, limit=8):
    """Retourne les villages les plus proches non encore lies."""
    linked_village_ids = _get_linked_village_ids(village)
    villages_candidats = list(
        Village.objects.select_related('fokontany__commune')
        .exclude(pk=village.pk)
        .exclude(pk__in=linked_village_ids)
    )

    for candidate in villages_candidats:
        candidate.air_distance_km = _calculate_air_distance_km(
            village.latitude,
            village.longitude,
            candidate.latitude,
            candidate.longitude,
        )

    villages_candidats = [item for item in villages_candidats if item.air_distance_km is not None]
    villages_candidats.sort(
        key=lambda item: (
            item.air_distance_km,
            item.nom_village.lower(),
        )
    )
    return villages_candidats[:limit]


def _autolink_nearest_villages(village, limit=3, type_liaison='proximite'):
    """
    Cree des liaisons sortantes vers les villages les plus proches.
    Cette premiere passe utilise la distance a vol d'oiseau comme distance_km.
    """
    nearest_villages = _get_nearest_unlinked_villages(village, limit=limit)
    created_count = 0

    for ordre, target_village in enumerate(nearest_villages, start=1):
        if LiaisonVillage.objects.filter(
            village_source=village,
            village_cible=target_village,
            type_liaison=type_liaison,
        ).exists():
            continue

        LiaisonVillage.objects.create(
            village_source=village,
            village_cible=target_village,
            type_liaison=type_liaison,
            distance_km=Decimal(str(target_village.air_distance_km)),
            ordre=ordre,
            actif=True,
        )
        created_count += 1

    return created_count, nearest_villages


def _autolink_all_villages(limit_per_village=3, type_liaison='proximite'):
    """
    Cree en lot des liaisons de proximite pour tous les villages disposant de GPS.
    Pour chaque village, on cree des liaisons sortantes vers ses plus proches voisins non deja lies.
    """
    villages = list(
        Village.objects.select_related('fokontany__commune')
        .filter(latitude__isnull=False, longitude__isnull=False)
    )
    if not villages:
        return {
            'eligible_villages': 0,
            'created_links': 0,
            'skipped_villages': 0,
        }

    existing_links = set(
        LiaisonVillage.objects.filter(type_liaison=type_liaison)
        .values_list('village_source_id', 'village_cible_id')
    )

    villages_by_id = {v.id: v for v in villages}
    to_create = []
    created_links = 0
    skipped_villages = 0

    for village in villages:
        ranked_candidates = []
        for candidate in villages:
            if candidate.id == village.id:
                continue

            if (village.id, candidate.id) in existing_links:
                continue

            air_distance_km = _calculate_air_distance_km(
                village.latitude,
                village.longitude,
                candidate.latitude,
                candidate.longitude,
            )
            if air_distance_km is None:
                continue

            ranked_candidates.append((air_distance_km, candidate.nom_village.lower(), candidate.id))

        ranked_candidates.sort()
        selected_candidates = ranked_candidates[:limit_per_village]
        if not selected_candidates:
            skipped_villages += 1
            continue

        for ordre, (air_distance_km, _, candidate_id) in enumerate(selected_candidates, start=1):
            if (village.id, candidate_id) in existing_links:
                continue

            to_create.append(
                LiaisonVillage(
                    village_source=village,
                    village_cible=villages_by_id[candidate_id],
                    type_liaison=type_liaison,
                    distance_km=Decimal(str(air_distance_km)),
                    ordre=ordre,
                    actif=True,
                )
            )
            existing_links.add((village.id, candidate_id))
            created_links += 1

    if to_create:
        LiaisonVillage.objects.bulk_create(to_create, batch_size=500)

    return {
        'eligible_villages': len(villages),
        'created_links': created_links,
        'skipped_villages': skipped_villages,
    }


def _authority_candidates_for_zones(zones, zone_type):
    """Retourne les autorites directes ou hierarchiquement liees aux zones."""
    from apps.usagers.models_reunions import Autorite

    zones = list(zones)
    if not zones:
        return Autorite.objects.none()

    zone_ids = [zone.id for zone in zones]
    if zone_type == 'region':
        relation_filter = (
            Q(region_id__in=zone_ids)
            | Q(district__region_id__in=zone_ids)
            | Q(commune__district__region_id__in=zone_ids)
            | Q(fokontany__commune__district__region_id__in=zone_ids)
        )
    elif zone_type == 'district':
        region_ids = [zone.region_id for zone in zones]
        relation_filter = (
            Q(region_id__in=region_ids)
            | Q(district_id__in=zone_ids)
            | Q(commune__district_id__in=zone_ids)
            | Q(fokontany__commune__district_id__in=zone_ids)
        )
    elif zone_type == 'commune':
        district_ids = [zone.district_id for zone in zones]
        region_ids = [zone.district.region_id for zone in zones]
        relation_filter = (
            Q(region_id__in=region_ids)
            | Q(district_id__in=district_ids)
            | Q(commune_id__in=zone_ids)
            | Q(fokontany__commune_id__in=zone_ids)
        )
    elif zone_type == 'fokontany':
        commune_ids = [zone.commune_id for zone in zones]
        district_ids = [zone.commune.district_id for zone in zones]
        region_ids = [zone.commune.district.region_id for zone in zones]
        relation_filter = (
            Q(region_id__in=region_ids)
            | Q(district_id__in=district_ids)
            | Q(commune_id__in=commune_ids)
            | Q(fokontany_id__in=zone_ids)
        )
    else:
        return Autorite.objects.none()

    return Autorite.objects.filter(relation_filter).select_related(
        'region', 'district', 'commune', 'fokontany'
    ).distinct().order_by('nom')


def _authority_relation_to_zone(authority, zone_type, zone):
    """Retourne 'direct', 'indirect' ou None selon le rattachement hierarchique."""
    # Une fiche peut contenir toute la chaine geographique (region + district +
    # commune + fokontany). Le rattachement direct est toujours le niveau le
    # plus precis renseigne; les niveaux parents ne sont que des liens indirects.
    most_specific_level = (
        'fokontany' if authority.fokontany_id else
        'commune' if authority.commune_id else
        'district' if authority.district_id else
        'region' if authority.region_id else
        None
    )

    if most_specific_level == 'fokontany':
        target_region_id = authority.fokontany.commune.district.region_id
        target_district_id = authority.fokontany.commune.district_id
        target_commune_id = authority.fokontany.commune_id
        target_fokontany_id = authority.fokontany_id
    elif most_specific_level == 'commune':
        target_region_id = authority.commune.district.region_id
        target_district_id = authority.commune.district_id
        target_commune_id = authority.commune_id
        target_fokontany_id = None
    elif most_specific_level == 'district':
        target_region_id = authority.district.region_id
        target_district_id = authority.district_id
        target_commune_id = None
        target_fokontany_id = None
    elif most_specific_level == 'region':
        target_region_id = authority.region_id
        target_district_id = None
        target_commune_id = None
        target_fokontany_id = None
    else:
        return None

    if zone_type == 'region':
        same_hierarchy = target_region_id == zone.id
    elif zone_type == 'district':
        same_hierarchy = (
            target_district_id == zone.id
            if target_district_id
            else target_region_id == zone.region_id
        )
    elif zone_type == 'commune':
        if target_commune_id:
            same_hierarchy = target_commune_id == zone.id
        elif target_district_id:
            same_hierarchy = target_district_id == zone.district_id
        else:
            same_hierarchy = target_region_id == zone.district.region_id
    elif zone_type == 'fokontany':
        if target_fokontany_id:
            same_hierarchy = target_fokontany_id == zone.id
        elif target_commune_id:
            same_hierarchy = target_commune_id == zone.commune_id
        elif target_district_id:
            same_hierarchy = target_district_id == zone.commune.district_id
        else:
            same_hierarchy = target_region_id == zone.commune.district.region_id
    else:
        return None

    if not same_hierarchy:
        return None
    return 'direct' if most_specific_level == zone_type else 'indirect'


def _add_authority_scope_counts(zones, zone_type):
    zones = list(zones)
    authorities = list(_authority_candidates_for_zones(zones, zone_type))
    for zone in zones:
        direct_ids = set()
        indirect_ids = set()
        for authority in authorities:
            relation = _authority_relation_to_zone(authority, zone_type, zone)
            if relation == 'direct':
                direct_ids.add(authority.id)
            elif relation == 'indirect':
                indirect_ids.add(authority.id)
        zone.autorites_directes_count = len(direct_ids)
        zone.autorites_indirectes_count = len(indirect_ids - direct_ids)
    return zones


def _authority_attachment(authority):
    if authority.fokontany_id:
        return 'Fokontany', authority.fokontany.nom_fokontany
    if authority.commune_id:
        return 'Commune', authority.commune.nom_commune
    if authority.district_id:
        return 'District', authority.district.nom_district
    if authority.region_id:
        return 'Region', authority.region.nom_region
    return 'Non precise', ''


def geography_list(request):
    """Page principale pour la gestion de la géographie avec liste des régions"""
    ex_provinces = list(ExProvince.objects.all().order_by('nom_ex_province'))
    regions_queryset = Region.objects.select_related('ex_province').order_by('nom_region')

    # Compteurs « habitants » (somme des parcelles) affiches au-dessus des tableaux.
    habitants_ex_provinces = habitants_par_zone(ExProvince.objects.all(), 'ex_province')
    habitants_regions = habitants_par_zone(regions_queryset, 'region')

    for ex_province in ex_provinces:
        ex_province.habitants_count = habitants_ex_provinces.get(ex_province.id, 0)

    regions = _add_authority_scope_counts(regions_queryset, 'region')
    for region in regions:
        region.habitants_count = habitants_regions.get(region.id, 0)

    est_directeur = user_is_director(request.user)
    return render(request, 'geography/geography_list.html', {
        'ex_provinces': ex_provinces,
        'regions': regions,
        'ex_provinces_habitants_total': sum(habitants_ex_provinces.values()),
        'regions_habitants_total': sum(habitants_regions.values()),
        'title': 'Gestion Géographique',
        'est_directeur': est_directeur,
    })


# CRUD pour les Régions
def region_list(request):
    regions = Region.objects.all().order_by('nom_region')
    return render(request, 'geography/region_list.html', {
        'regions': regions,
        'title': 'Régions'
    })


def region_create(request):
    if request.method == 'POST':
        form = RegionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Région créée avec succès!')
            return redirect('geography:region_list')
        else:
            messages.error(request, 'Erreur dans le formulaire. Veuillez vérifier les champs.')
    else:
        form = RegionForm()

    return render(request, 'geography/region_form.html', {
        'form': form,
        'title': 'Nouvelle Région'
    })


def region_detail(request, pk):
    region = get_object_or_404(Region, pk=pk)
    districts = region.districts.all().order_by('nom_district')
    return render(request, 'geography/region_detail.html', {
        'region': region,
        'districts': districts,
        'title': f'Région {region.nom_region}'
    })


def region_update(request, pk):
    region = get_object_or_404(Region, pk=pk)
    if request.method == 'POST':
        form = RegionForm(request.POST, instance=region)
        if form.is_valid():
            form.save()
            messages.success(request, 'Région modifiée avec succès!')
            return redirect('geography:region_detail', pk=pk)
    else:
        form = RegionForm(instance=region)

    return render(request, 'geography/region_form.html', {
        'form': form,
        'region': region,
        'title': 'Modifier Région'
    })


def region_delete(request, pk):
    region = get_object_or_404(Region, pk=pk)
    if request.method == 'POST':
        region.delete()
        messages.success(request, 'Région supprimée avec succès!')
        return redirect('geography:region_list')
    return render(request, 'geography/region_confirm_delete.html', {
        'region': region,
        'title': 'Supprimer Région'
    })


# CRUD pour les Districts
def district_list(request):
    districts = District.objects.select_related('region').all().order_by('region__nom_region', 'nom_district')
    return render(request, 'geography/district_list.html', {
        'districts': districts,
        'title': 'Districts'
    })


def district_create(request):
    if request.method == 'POST':
        form = DistrictForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'District créé avec succès!')
            return redirect('geography:district_list')
        else:
            messages.error(request, 'Erreur dans le formulaire. Veuillez vérifier les champs.')
    else:
        form = DistrictForm()

    return render(request, 'geography/district_form.html', {
        'form': form,
        'title': 'Nouveau District'
    })


def district_detail(request, pk):
    district = get_object_or_404(District, pk=pk)
    communes = district.communes.all().order_by('nom_commune')
    # Contour du district et de ses communes (importés des shapefiles), s'ils existent.
    try:
        from .services.contours_communes import contours_district
        carte_contours = {'type': 'FeatureCollection', 'features': contours_district(district)}
    except Exception:
        carte_contours = {'type': 'FeatureCollection', 'features': []}
    return render(request, 'geography/district_detail.html', {
        'district': district,
        'communes': communes,
        'carte_contours': carte_contours,
        'title': f'District {district.nom_district}'
    })


def district_update(request, pk):
    district = get_object_or_404(District, pk=pk)
    if request.method == 'POST':
        form = DistrictForm(request.POST, instance=district)
        if form.is_valid():
            form.save()
            messages.success(request, 'District modifié avec succès!')
            return redirect('geography:district_detail', pk=pk)
    else:
        form = DistrictForm(instance=district)

    return render(request, 'geography/district_form.html', {
        'form': form,
        'district': district,
        'title': 'Modifier District'
    })


def district_delete(request, pk):
    district = get_object_or_404(District, pk=pk)
    if request.method == 'POST':
        district.delete()
        messages.success(request, 'District supprimé avec succès!')
        return redirect('geography:district_list')
    return render(request, 'geography/district_confirm_delete.html', {
        'district': district,
        'title': 'Supprimer District'
    })


# CRUD pour les Communes
def commune_list(request):
    from django.core.paginator import Paginator
    from django.db.models import Q, Count

    communes = Commune.objects.select_related('district__region').annotate(
        villages_count=Count('fokontanys__villages', distinct=True)
    )

    # Filtres de recherche
    search = request.GET.get('search', '').strip()
    district = request.GET.get('district', '').strip()

    if search:
        communes = communes.filter(nom_commune__icontains=search)
    if district:
        communes = communes.filter(district__nom_district__icontains=district)

    communes = communes.order_by('district__region__nom_region', 'district__nom_district', 'nom_commune')

    # Pagination
    paginator = Paginator(communes, 50)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'geography/commune_list.html', {
        'communes': page_obj,
        'page_obj': page_obj,
        'title': 'Communes'
    })


def commune_create(request):
    if request.method == 'POST':
        form = CommuneForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Commune créée avec succès!')
            return redirect('geography:commune_list')
        else:
            messages.error(request, 'Erreur dans le formulaire. Veuillez vérifier les champs.')
    else:
        form = CommuneForm()

    return render(request, 'geography/commune_form.html', {
        'form': form,
        'title': 'Nouvelle Commune'
    })


def commune_detail(request, pk):
    commune = get_object_or_404(Commune, pk=pk)
    villages = commune.villages.all().order_by('nom_village')
    return render(request, 'geography/commune_detail.html', {
        'commune': commune,
        'villages': villages,
        'title': f'Commune {commune.nom_commune}'
    })


def commune_update(request, pk):
    commune = get_object_or_404(Commune, pk=pk)
    if request.method == 'POST':
        form = CommuneForm(request.POST, instance=commune)
        if form.is_valid():
            form.save()
            messages.success(request, 'Commune modifiée avec succès!')
            return redirect('geography:commune_detail', pk=pk)
    else:
        form = CommuneForm(instance=commune)

    return render(request, 'geography/commune_form.html', {
        'form': form,
        'commune': commune,
        'title': 'Modifier Commune'
    })


def commune_delete(request, pk):
    commune = get_object_or_404(Commune, pk=pk)
    if request.method == 'POST':
        commune.delete()
        messages.success(request, 'Commune supprimée avec succès!')
        return redirect('geography:commune_list')
    return render(request, 'geography/commune_confirm_delete.html', {
        'commune': commune,
        'title': 'Supprimer Commune'
    })


# CRUD pour les Villages
def village_list(request):
    from django.core.paginator import Paginator
    from django.db.models import Q, Count

    villages = Village.objects.select_related('fokontany__commune__district__region').annotate(
        parcelles_count=Count('parcelles', distinct=True)
    )

    # Filtres de recherche
    search = request.GET.get('search', '').strip()
    commune = request.GET.get('commune', '').strip()
    district = request.GET.get('district', '').strip()

    if search:
        villages = villages.filter(nom_village__icontains=search)
    if commune:
        villages = villages.filter(fokontany__commune__nom_commune__icontains=commune)
    if district:
        villages = villages.filter(fokontany__commune__district__nom_district__icontains=district)

    villages = villages.order_by('fokontany__commune__district__region__nom_region', 'fokontany__commune__district__nom_district', 'fokontany__commune__nom_commune', 'nom_village')

    # Pagination
    paginator = Paginator(villages, 100)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    villages_with_gps_count = Village.objects.filter(latitude__isnull=False, longitude__isnull=False).count()
    villages_without_gps_count = Village.objects.filter(
        Q(latitude__isnull=True) | Q(longitude__isnull=True)
    ).count()

    return render(request, 'geography/village_list.html', {
        'villages': page_obj,
        'page_obj': page_obj,
        'title': 'Villages',
        'est_directeur': user_is_director(request.user),
        'villages_with_gps_count': villages_with_gps_count,
        'villages_without_gps_count': villages_without_gps_count,
    })


def village_create(request):
    if request.method == 'POST':
        form = VillageForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Village créé avec succès!')
            return redirect('geography:village_list')
        else:
            messages.error(request, 'Erreur dans le formulaire. Veuillez vérifier les champs.')
    else:
        form = VillageForm()

    return render(request, 'geography/village_form.html', {
        'form': form,
        'title': 'Nouveau Village'
    })


def village_detail(request, pk):
    village = get_object_or_404(
        Village.objects.select_related('fokontany__commune__district__region'),
        pk=pk,
    )
    parcelles = village.parcelles.all().order_by('reference_parcelle')
    liaisons_sortantes = list(
        LiaisonVillage.objects.filter(village_source=village, actif=True)
        .select_related('village_source__fokontany__commune', 'village_cible__fokontany__commune')
        .order_by('ordre', 'village_cible__nom_village')
    )
    liaisons_entrantes = list(
        LiaisonVillage.objects.filter(village_cible=village, actif=True)
        .select_related('village_source__fokontany__commune', 'village_cible__fokontany__commune')
        .order_by('ordre', 'village_source__nom_village')
    )
    villages_candidats = list(
        Village.objects.select_related('fokontany__commune').exclude(pk=village.pk)
    )
    for candidate in villages_candidats:
        candidate.air_distance_km = _calculate_air_distance_km(
            village.latitude,
            village.longitude,
            candidate.latitude,
            candidate.longitude,
        )
    villages_candidats.sort(
        key=lambda item: (
            item.air_distance_km is None,
            item.air_distance_km if item.air_distance_km is not None else float('inf'),
            item.nom_village.lower(),
        )
    )
    villages_proches_non_lies = _get_nearest_unlinked_villages(village, limit=8)
    village_map_points_by_id = {}

    def add_linked_village_to_map(liaison, direction):
        other_village = liaison.village_cible if direction == 'sortant' else liaison.village_source
        if other_village.latitude is None or other_village.longitude is None:
            return
        point = village_map_points_by_id.setdefault(str(other_village.id), {
            'id': str(other_village.id),
            'name': other_village.nom_village,
            'commune': other_village.commune.nom_commune if other_village.fokontany_id else '',
            'latitude': float(other_village.latitude),
            'longitude': float(other_village.longitude),
            'status': 'linked',
            'directions': [],
            'relation_labels': [],
            'distance_km': _calculate_air_distance_km(
                village.latitude,
                village.longitude,
                other_village.latitude,
                other_village.longitude,
            ),
        })
        if direction not in point['directions']:
            point['directions'].append(direction)
        relation_label = liaison.get_type_liaison_display()
        if relation_label not in point['relation_labels']:
            point['relation_labels'].append(relation_label)

    for liaison in liaisons_entrantes:
        add_linked_village_to_map(liaison, 'entrant')
    for liaison in liaisons_sortantes:
        add_linked_village_to_map(liaison, 'sortant')

    for nearby_village in villages_proches_non_lies:
        if nearby_village.latitude is None or nearby_village.longitude is None:
            continue
        village_map_points_by_id.setdefault(str(nearby_village.id), {
            'id': str(nearby_village.id),
            'name': nearby_village.nom_village,
            'commune': nearby_village.commune.nom_commune if nearby_village.fokontany_id else '',
            'latitude': float(nearby_village.latitude),
            'longitude': float(nearby_village.longitude),
            'status': 'nearby',
            'directions': [],
            'relation_labels': [],
            'distance_km': nearby_village.air_distance_km,
        })

    organigram = _build_village_organigram(village, liaisons_entrantes, liaisons_sortantes)
    regions = Region.objects.all().order_by('nom_region')
    from .views_liaisons_api import liaison_properties

    return render(request, 'geography/village_detail.html', {
        'village': village,
        'village_liaisons_json': [liaison_properties(item) for item in liaisons_entrantes + liaisons_sortantes],
        'parcelles': parcelles,
        'regions': regions,
        'current_region_id': str(village.region.id) if village.fokontany_id else '',
        'current_district_id': str(village.district.id) if village.fokontany_id else '',
        'current_commune_id': str(village.commune.id) if village.fokontany_id else '',
        'current_fokontany_id': str(village.fokontany.id) if village.fokontany_id else '',
        'current_village_id': str(village.id),
        'liaisons_sortantes': [_serialize_liaison_village(item, village.id) for item in liaisons_sortantes],
        'liaisons_entrantes': [_serialize_liaison_village(item, village.id) for item in liaisons_entrantes],
        'villages_candidats': villages_candidats,
        'villages_proches_non_lies': villages_proches_non_lies,
        'village_map_points': list(village_map_points_by_id.values()),
        'suggested_auto_link_count': min(3, len(villages_proches_non_lies)),
        'liaison_type_choices': LiaisonVillage.TYPE_LIAISON_CHOICES,
        'liaison_route_choices': LiaisonVillage.TYPE_ROUTE_CHOICES,
        'liaison_moyen_choices': LiaisonVillage.MOYEN_ACCES_CHOICES,
        'liaisons_total': len(liaisons_entrantes) + len(liaisons_sortantes),
        'organigram_nodes': organigram['nodes'],
        'organigram_edges': organigram['edges'],
        'title': f'Village {village.nom_village}'
    })


def village_update(request, pk):
    village = get_object_or_404(Village, pk=pk)
    if request.method == 'POST':
        form = VillageForm(request.POST, instance=village)
        if form.is_valid():
            form.save()
            messages.success(request, 'Village modifié avec succès!')
            return redirect('geography:village_detail', pk=pk)
    else:
        form = VillageForm(instance=village)

    return render(request, 'geography/village_form.html', {
        'form': form,
        'village': village,
        'title': 'Modifier Village'
    })


def village_delete(request, pk):
    village = get_object_or_404(Village, pk=pk)
    if request.method == 'POST':
        village.delete()
        messages.success(request, 'Village supprimé avec succès!')
        return redirect('geography:village_list')
    return render(request, 'geography/village_confirm_delete.html', {
        'village': village,
        'title': 'Supprimer Village'
    })


@require_http_methods(["POST"])
def village_liaison_create(request, pk):
    """Cree une liaison orientee depuis ou vers le village courant."""
    village = get_object_or_404(Village, pk=pk)
    target_id = request.POST.get('target_village_id')
    direction = (request.POST.get('direction') or 'sortant').strip().lower()
    type_liaison = (request.POST.get('type_liaison') or 'desserte').strip().lower()
    type_route = (request.POST.get('type_route') or '').strip().lower() or None
    moyen_acces = (request.POST.get('moyen_acces') or '').strip().lower() or None
    ordre_raw = (request.POST.get('ordre') or '').strip()
    distance_raw = (request.POST.get('distance_km') or '').strip()
    temps_raw = (request.POST.get('temps_acces_minutes') or '').strip()

    if not target_id:
        messages.error(request, "Village cible obligatoire.")
        return redirect('geography:village_detail', pk=pk)

    target_village = get_object_or_404(Village, pk=target_id)

    if direction not in {'sortant', 'entrant'}:
        messages.error(request, "Sens de liaison invalide.")
        return redirect('geography:village_detail', pk=pk)

    valid_types = {choice[0] for choice in LiaisonVillage.TYPE_LIAISON_CHOICES}
    if type_liaison not in valid_types:
        messages.error(request, "Type de liaison invalide.")
        return redirect('geography:village_detail', pk=pk)

    valid_routes = {choice[0] for choice in LiaisonVillage.TYPE_ROUTE_CHOICES}
    if type_route and type_route not in valid_routes:
        messages.error(request, "Type de route invalide.")
        return redirect('geography:village_detail', pk=pk)

    valid_moyens = {choice[0] for choice in LiaisonVillage.MOYEN_ACCES_CHOICES}
    if moyen_acces and moyen_acces not in valid_moyens:
        messages.error(request, "Moyen d'acces invalide.")
        return redirect('geography:village_detail', pk=pk)

    village_source = village if direction == 'sortant' else target_village
    village_cible = target_village if direction == 'sortant' else village

    if village_source.pk == village_cible.pk:
        messages.error(request, "Un village ne peut pas etre relie a lui-meme.")
        return redirect('geography:village_detail', pk=pk)

    try:
        ordre = int(ordre_raw) if ordre_raw else 0
    except ValueError:
        messages.error(request, "Ordre invalide.")
        return redirect('geography:village_detail', pk=pk)

    try:
        distance_km = _parse_optional_decimal(distance_raw, "la distance")
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect('geography:village_detail', pk=pk)

    if distance_km is not None and distance_km < 0:
        messages.error(request, "La distance ne peut pas etre negative.")
        return redirect('geography:village_detail', pk=pk)

    try:
        temps_acces_minutes = int(temps_raw) if temps_raw else None
    except ValueError:
        messages.error(request, "Temps d'acces invalide.")
        return redirect('geography:village_detail', pk=pk)

    if temps_acces_minutes is not None and temps_acces_minutes < 0:
        messages.error(request, "Le temps d'acces ne peut pas etre negatif.")
        return redirect('geography:village_detail', pk=pk)

    if LiaisonVillage.objects.filter(
        village_source=village_source,
        village_cible=village_cible,
        type_liaison=type_liaison,
    ).exists():
        messages.warning(request, "Cette liaison existe deja.")
        return redirect('geography:village_detail', pk=pk)

    nouvelle = LiaisonVillage.objects.create(
        village_source=village_source,
        village_cible=village_cible,
        type_liaison=type_liaison,
        type_route=type_route,
        distance_km=distance_km,
        temps_acces_minutes=temps_acces_minutes,
        moyen_acces=moyen_acces,
        ordre=max(0, ordre),
    )
    from .services.liaisons_routes import calculer_une
    calculer_une(nouvelle)
    messages.success(request, "Liaison entre villages ajoutee.")
    return redirect('geography:village_detail', pk=pk)


@require_http_methods(["POST"])
def village_liaisons_autofill(request, pk):
    """Cree automatiquement des liaisons de proximite vers les villages les plus proches."""
    village = get_object_or_404(Village, pk=pk)

    limit_raw = (request.POST.get('limit') or '').strip()
    try:
        limit = int(limit_raw) if limit_raw else 3
    except ValueError:
        messages.error(request, "Nombre de liaisons a creer invalide.")
        return redirect('geography:village_detail', pk=pk)

    if limit <= 0:
        messages.error(request, "Le nombre de liaisons a creer doit etre superieur a zero.")
        return redirect('geography:village_detail', pk=pk)

    if village.latitude is None or village.longitude is None:
        messages.error(request, "Le village courant doit avoir des coordonnees GPS pour lancer l'auto-liaison.")
        return redirect('geography:village_detail', pk=pk)

    created_count, nearest_villages = _autolink_nearest_villages(village, limit=limit, type_liaison='proximite')

    if created_count == 0:
        if nearest_villages:
            messages.info(request, "Aucune nouvelle liaison creee. Les villages proches proposes sont deja lies.")
        else:
            messages.warning(request, "Aucun village proche avec coordonnees GPS disponibles pour une auto-liaison.")
        return redirect('geography:village_detail', pk=pk)

    messages.success(
        request,
        f"{created_count} liaison(s) de proximite creee(s) automatiquement a partir de la distance a vol d'oiseau."
    )
    return redirect('geography:village_detail', pk=pk)


@require_http_methods(["POST"])
def village_liaisons_autofill_global(request):
    """Lance un peuplement global des liaisons de proximite entre tous les villages geolocalises."""
    if not user_is_director(request.user):
        messages.error(request, "Seul un directeur peut lancer le peuplement global des liaisons.")
        return redirect('geography:village_list')

    limit_raw = (request.POST.get('limit_per_village') or '').strip()
    try:
        limit_per_village = int(limit_raw) if limit_raw else 3
    except ValueError:
        messages.error(request, "Nombre de voisins par village invalide.")
        return redirect('geography:village_list')

    if limit_per_village <= 0:
        messages.error(request, "Le nombre de voisins par village doit etre superieur a zero.")
        return redirect('geography:village_list')

    with transaction.atomic():
        result = _autolink_all_villages(limit_per_village=limit_per_village, type_liaison='proximite')

    if result['eligible_villages'] == 0:
        messages.warning(request, "Aucun village geolocalise disponible pour le peuplement global.")
        return redirect('geography:village_list')

    messages.success(
        request,
        (
            f"Peuplement global termine : {result['created_links']} liaison(s) creee(s) "
            f"sur {result['eligible_villages']} village(s) geolocalise(s). "
            f"{result['skipped_villages']} village(s) sans voisin disponible."
        )
    )
    return redirect('geography:village_list')


@require_http_methods(["POST"])
def village_liaison_update(request, liaison_id):
    """Met a jour une liaison entre villages, y compris le sens et le village lie."""
    liaison = get_object_or_404(LiaisonVillage, pk=liaison_id)
    redirect_village_id = request.POST.get('redirect_village_id') or liaison.village_source_id
    target_id = request.POST.get('target_village_id')
    direction = (request.POST.get('direction') or 'sortant').strip().lower()
    type_liaison = (request.POST.get('type_liaison') or liaison.type_liaison).strip().lower()
    type_route = (request.POST.get('type_route') or '').strip().lower() or None
    moyen_acces = (request.POST.get('moyen_acces') or '').strip().lower() or None
    ordre_raw = (request.POST.get('ordre') or '').strip()
    distance_raw = (request.POST.get('distance_km') or '').strip()
    temps_raw = (request.POST.get('temps_acces_minutes') or '').strip()

    if not target_id:
        messages.error(request, "Village cible obligatoire.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    target_village = get_object_or_404(Village, pk=target_id)

    if direction not in {'sortant', 'entrant'}:
        messages.error(request, "Sens de liaison invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    valid_types = {choice[0] for choice in LiaisonVillage.TYPE_LIAISON_CHOICES}
    if type_liaison not in valid_types:
        messages.error(request, "Type de liaison invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    valid_routes = {choice[0] for choice in LiaisonVillage.TYPE_ROUTE_CHOICES}
    if type_route and type_route not in valid_routes:
        messages.error(request, "Type de route invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    valid_moyens = {choice[0] for choice in LiaisonVillage.MOYEN_ACCES_CHOICES}
    if moyen_acces and moyen_acces not in valid_moyens:
        messages.error(request, "Moyen d'acces invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    try:
        ordre = int(ordre_raw) if ordre_raw else 0
    except ValueError:
        messages.error(request, "Ordre invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    try:
        distance_km = _parse_optional_decimal(distance_raw, "la distance")
    except ValueError as exc:
        messages.error(request, str(exc))
        return redirect('geography:village_detail', pk=redirect_village_id)

    if distance_km is not None and distance_km < 0:
        messages.error(request, "La distance ne peut pas etre negative.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    try:
        temps_acces_minutes = int(temps_raw) if temps_raw else None
    except ValueError:
        messages.error(request, "Temps d'acces invalide.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    if temps_acces_minutes is not None and temps_acces_minutes < 0:
        messages.error(request, "Le temps d'acces ne peut pas etre negatif.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    pivot_village = get_object_or_404(Village, pk=redirect_village_id)
    village_source = pivot_village if direction == 'sortant' else target_village
    village_cible = target_village if direction == 'sortant' else pivot_village

    if village_source.pk == village_cible.pk:
        messages.error(request, "Un village ne peut pas etre relie a lui-meme.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    if LiaisonVillage.objects.filter(
        village_source=village_source,
        village_cible=village_cible,
        type_liaison=type_liaison,
    ).exclude(pk=liaison.pk).exists():
        messages.warning(request, "Une liaison identique existe deja avec ce type.")
        return redirect('geography:village_detail', pk=redirect_village_id)

    liaison.village_source = village_source
    liaison.village_cible = village_cible
    liaison.type_liaison = type_liaison
    liaison.type_route = type_route
    liaison.moyen_acces = moyen_acces
    liaison.distance_km = distance_km
    liaison.temps_acces_minutes = temps_acces_minutes
    liaison.ordre = max(0, ordre)
    liaison.save(update_fields=[
        'village_source',
        'village_cible',
        'type_liaison',
        'type_route',
        'moyen_acces',
        'distance_km',
        'temps_acces_minutes',
        'ordre',
        'updated_at',
    ])
    from .services.liaisons_routes import calculer_une
    calculer_une(liaison)

    messages.success(request, "Liaison mise a jour.")
    return redirect('geography:village_detail', pk=redirect_village_id)


@require_http_methods(["POST"])
def village_liaison_delete(request, liaison_id):
    """Supprime une liaison entre villages."""
    liaison = get_object_or_404(LiaisonVillage, pk=liaison_id)
    redirect_village_id = request.POST.get('redirect_village_id') or liaison.village_source_id
    liaison.delete()
    messages.success(request, "Liaison supprimee.")
    return redirect('geography:village_detail', pk=redirect_village_id)


# CRUD pour les Parcelles
def parcelle_list(request):
    from django.core.paginator import Paginator
    from django.db.models import Q

    parcelles = Parcelle.objects.select_related('village__fokontany__commune__district__region').all()

    # Filtres de recherche
    search = request.GET.get('search', '').strip()
    village = request.GET.get('village', '').strip()
    commune = request.GET.get('commune', '').strip()

    if search:
        parcelles = parcelles.filter(reference_parcelle__icontains=search)
    if village:
        parcelles = parcelles.filter(village__nom_village__icontains=village)
    if commune:
        parcelles = parcelles.filter(village__fokontany__commune__nom_commune__icontains=commune)

    parcelles = parcelles.order_by('village__fokontany__commune__district__region__nom_region', 'village__fokontany__commune__district__nom_district', 'village__fokontany__commune__nom_commune', 'village__nom_village', 'reference_parcelle')

    # Pagination
    paginator = Paginator(parcelles, 100)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'geography/parcelle_list.html', {
        'parcelles': page_obj,
        'page_obj': page_obj,
        'title': 'Parcelles'
    })


def parcelle_create(request):
    if request.method == 'POST':
        form = ParcelleForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Parcelle créée avec succès!')
            return redirect('geography:parcelle_list')
        else:
            messages.error(request, 'Erreur dans le formulaire. Veuillez vérifier les champs.')
    else:
        form = ParcelleForm()

    return render(request, 'geography/parcelle_form.html', {
        'form': form,
        'title': 'Nouvelle Parcelle'
    })


def parcelle_detail(request, pk):
    parcelle = get_object_or_404(Parcelle, pk=pk)
    est_directeur = user_is_director(request.user)
    return render(request, 'geography/parcelle_detail.html', {
        'parcelle': parcelle,
        'est_directeur': est_directeur,
        'title': f'Parcelle {parcelle.reference_parcelle}'
    })


def parcelle_update(request, pk):
    parcelle = get_object_or_404(Parcelle, pk=pk)
    if request.method == 'POST':
        form = ParcelleForm(request.POST, instance=parcelle)
        if form.is_valid():
            form.save()
            messages.success(request, 'Parcelle modifiée avec succès!')
            return redirect('geography:parcelle_detail', pk=pk)
    else:
        form = ParcelleForm(instance=parcelle)

    return render(request, 'geography/parcelle_form.html', {
        'form': form,
        'parcelle': parcelle,
        'title': 'Modifier Parcelle'
    })


def parcelle_delete(request, pk):
    parcelle = get_object_or_404(Parcelle, pk=pk)
    if request.method == 'POST':
        parcelle.delete()
        messages.success(request, 'Parcelle supprimée avec succès!')
        return redirect('geography:parcelle_list')
    return render(request, 'geography/parcelle_confirm_delete.html', {
        'parcelle': parcelle,
        'title': 'Supprimer Parcelle'
    })


# API Views pour les opérations CRUD en cascade

def api_exprovinces(request):
    """API pour récupérer toutes les ex-provinces."""
    ex_provinces = ExProvince.objects.all().order_by('nom_ex_province')
    habitants = habitants_par_zone(ex_provinces, 'ex_province')
    ex_provinces = ex_provinces.annotate(regions_count=models.Count('regions'))
    return JsonResponse([
        {
            'id': str(ex_province.id),
            'nom_ex_province': ex_province.nom_ex_province,
            'regions_count': ex_province.regions_count,
            'habitants_count': habitants.get(ex_province.id, 0),
        }
        for ex_province in ex_provinces
    ], safe=False)


def api_regions(request):
    """API pour récupérer toutes les régions"""
    ex_province_id = request.GET.get('ex_province_id')
    regions = Region.objects.select_related('ex_province')
    if ex_province_id:
        regions = regions.filter(ex_province_id=ex_province_id)
    habitants = habitants_par_zone(regions, 'region')
    regions = _add_authority_scope_counts(
        regions.annotate(districts_count=models.Count('districts', distinct=True)).order_by('nom_region'),
        'region',
    )
    data = [
        {
            'id': str(region.id),
            'nom_region': region.nom_region,
            'photo': region.photo,
            'districts_count': region.districts_count,
            'habitants_count': habitants.get(region.id, 0),
            'autorites_directes_count': region.autorites_directes_count,
            'autorites_indirectes_count': region.autorites_indirectes_count,
            'ex_province_id': str(region.ex_province_id) if region.ex_province_id else None,
            'ex_province__nom_ex_province': region.ex_province.nom_ex_province if region.ex_province else None,
        }
        for region in regions
    ]
    return JsonResponse(data, safe=False)


def api_districts(request, region_id):
    """API pour récupérer les districts d'une région"""
    districts = District.objects.select_related('region').filter(region_id=region_id)
    habitants = habitants_par_zone(districts, 'district')
    districts = _add_authority_scope_counts(
        districts.annotate(communes_count=models.Count('communes', distinct=True)).order_by('nom_district'),
        'district',
    )
    return JsonResponse([
        {
            'id': str(district.id),
            'nom_district': district.nom_district,
            'photo': district.photo,
            'communes_count': district.communes_count,
            'habitants_count': habitants.get(district.id, 0),
            'autorites_directes_count': district.autorites_directes_count,
            'autorites_indirectes_count': district.autorites_indirectes_count,
        }
        for district in districts
    ], safe=False)


def api_communes(request, district_id):
    """API pour récupérer les communes d'un district"""
    communes = Commune.objects.select_related('district__region').filter(district_id=district_id)
    habitants = habitants_par_zone(communes, 'commune')
    communes = _add_authority_scope_counts(
        communes.annotate(fokontanys_count=models.Count('fokontanys', distinct=True)).order_by('nom_commune'),
        'commune',
    )
    return JsonResponse([
        {
            'id': str(commune.id),
            'nom_commune': commune.nom_commune,
            'photo': commune.photo,
            'fokontanys_count': commune.fokontanys_count,
            'habitants_count': habitants.get(commune.id, 0),
            'autorites_directes_count': commune.autorites_directes_count,
            'autorites_indirectes_count': commune.autorites_indirectes_count,
        }
        for commune in communes
    ], safe=False)


def api_fokontanys(request, commune_id):
    """API pour récupérer les fokontanys d'une commune"""
    fokontanys = Fokontany.objects.select_related('commune__district__region').filter(commune_id=commune_id)
    habitants = habitants_par_zone(fokontanys, 'fokontany')
    fokontanys = _add_authority_scope_counts(
        fokontanys.annotate(villages_count=models.Count('villages', distinct=True)).order_by('nom_fokontany'),
        'fokontany',
    )
    return JsonResponse([
        {
            'id': str(fokontany.id),
            'nom_fokontany': fokontany.nom_fokontany,
            'photo': fokontany.photo,
            'villages_count': fokontany.villages_count,
            'habitants_count': habitants.get(fokontany.id, 0),
            'autorites_directes_count': fokontany.autorites_directes_count,
            'autorites_indirectes_count': fokontany.autorites_indirectes_count,
        }
        for fokontany in fokontanys
    ], safe=False)


def api_autorites_zone(request, zone_type, zone_id):
    """
    API pour recuperer les autorites (Maire, Chef Fokontany, ...) rattachees a une zone
    geographique donnee (region, district, commune ou fokontany), utilisee par le modal
    "Voir detail" de la page geography_list.
    """
    model_par_type = {
        'region': Region.objects.all(),
        'district': District.objects.select_related('region'),
        'commune': Commune.objects.select_related('district__region'),
        'fokontany': Fokontany.objects.select_related('commune__district__region'),
    }
    zone_queryset = model_par_type.get(zone_type)
    if zone_queryset is None:
        return JsonResponse({'success': False, 'error': 'Type de zone invalide.'}, status=400)

    zone = get_object_or_404(zone_queryset, id=zone_id)
    autorites = _authority_candidates_for_zones([zone], zone_type)

    direct_results = []
    indirect_results = []
    for authority in autorites:
        relation = _authority_relation_to_zone(authority, zone_type, zone)
        if relation not in {'direct', 'indirect'}:
            continue
        attachment_type, attachment_name = _authority_attachment(authority)
        item = {
            'id': str(authority.id),
            'nom': authority.nom,
            'type_autorite_display': authority.get_type_autorite_display(),
            'fonction': authority.fonction or '',
            'contact': authority.contact or '',
            'email': authority.email or '',
            'actif': authority.actif,
            'rattachement_type': attachment_type,
            'rattachement_nom': attachment_name,
        }
        if relation == 'direct':
            direct_results.append(item)
        else:
            indirect_results.append(item)

    results = direct_results + indirect_results
    return JsonResponse({
        'success': True,
        'results': results,
        'direct_results': direct_results,
        'indirect_results': indirect_results,
        'direct_total': len(direct_results),
        'indirect_total': len(indirect_results),
        'total': len(results),
    })


def api_villages(request, fokontany_id):
    """API pour récupérer les villages d'un fokontany"""
    villages = Village.objects.select_related('fokontany__commune').filter(fokontany_id=fokontany_id).annotate(
        parcelles_count=models.Count('parcelles'),
        habitants_count=models.Sum('parcelles__nombre_habitants'),
    )

    data = [_serialize_village(village) for village in villages]
    return JsonResponse(data, safe=False)


def api_villages_commune(request, commune_id):
    """Villages d'une commune (tous fokontany confondus), pour déplacer une parcelle d'un village à l'autre."""
    villages = (
        Village.objects.filter(fokontany__commune_id=commune_id)
        .select_related('fokontany')
        .only('id', 'nom_village', 'fokontany__nom_fokontany')
        .order_by('nom_village')
    )
    return JsonResponse([
        {'id': str(v.id), 'nom_village': v.nom_village, 'fokontany_name': v.fokontany.nom_fokontany if v.fokontany_id else ''}
        for v in villages
    ], safe=False)


def api_village_detail(request, village_id):
    """API pour recuperer le detail d'un village."""
    village = get_object_or_404(
        Village.objects.select_related('fokontany__commune'),
        id=village_id,
    )
    return JsonResponse(_serialize_village(village))


def api_parcelles(request, village_id):
    """API pour récupérer les parcelles d'un village"""
    parcelles = Parcelle.objects.filter(village_id=village_id).values(
        'id', 'reference_parcelle', 'photo', 'nombre_habitants', 'nombre_menages', 'source_donnees'
    )
    return JsonResponse(list(parcelles), safe=False)


def api_pointeaus(request, parcelle_id):
    """API pour récupérer les points d'eau d'une parcelle"""
    pointeaus = PointEau.objects.filter(parcelle_id=parcelle_id).select_related('type_point_eau')

    data = []
    for pointeau in pointeaus:
        data.append({
            'id': str(pointeau.id),
            'nomabonne': pointeau.nomabonne,
            'numeroteleph': pointeau.numeroteleph,
            'contactabonne': pointeau.contactabonne,
            'datebranchement': pointeau.datebranchement.isoformat() if pointeau.datebranchement else None,
            'datearretbranchement': pointeau.datearretbranchement.isoformat() if pointeau.datearretbranchement else None,
            'typeabonne': pointeau.typeabonne,
            'type_point_eau_nom': pointeau.type_point_eau.nom_type if pointeau.type_point_eau else None,
            'is_active': pointeau.is_active,
            'tarifforfaitaire': pointeau.tarifforfaitaire,
            'latitude': pointeau.latitude,
            'longitude': pointeau.longitude
        })

    return JsonResponse(data, safe=False)


@csrf_exempt
@require_http_methods(["POST"])
def api_create(request):
    """API pour créer un nouvel élément"""
    try:
        if (request.content_type or '').startswith('application/json'):
            data = json.loads(request.body)
        else:
            data = request.POST

        element_type = data.get('type')
        name = data.get('name')
        parent_id = data.get('parent_id')

        if element_type == 'exprovince':
            ex_province = ExProvince.objects.create(nom_ex_province=name)
            return JsonResponse({'success': True, 'id': ex_province.id})

        if element_type == 'region':
            ex_province_id = data.get('ex_province_id') or parent_id
            ex_province = get_object_or_404(ExProvince, id=ex_province_id) if ex_province_id else None
            region = Region.objects.create(nom_region=name, ex_province=ex_province, photo=_get_geography_photo(data))
            return JsonResponse({'success': True, 'id': region.id})

        elif element_type == 'district':
            region = get_object_or_404(Region, id=parent_id)
            district = District.objects.create(nom_district=name, region=region, photo=_get_geography_photo(data))
            return JsonResponse({'success': True, 'id': district.id})

        elif element_type == 'commune':
            district = get_object_or_404(District, id=parent_id)
            commune = Commune.objects.create(nom_commune=name, district=district, photo=_get_geography_photo(data))
            return JsonResponse({'success': True, 'id': commune.id})

        elif element_type == 'fokontany':
            commune = get_object_or_404(Commune, id=parent_id)
            fokontany = Fokontany.objects.create(nom_fokontany=name, commune=commune, photo=_get_geography_photo(data))
            return JsonResponse({'success': True, 'id': fokontany.id})

        elif element_type == 'village':
            fokontany = get_object_or_404(Fokontany, id=parent_id)
            village = Village.objects.create(
                fokontany=fokontany,
                photo=_get_geography_photo(data),
                **_get_village_payload(data)
            )
            return JsonResponse({'success': True, 'id': village.id, 'village': _serialize_village(village)})

        elif element_type == 'parcelle':
            village = get_object_or_404(Village, id=parent_id)
            nombre_habitants = data.get('nombre_habitants')
            nombre_menages = data.get('nombre_menages')
            source_donnees = data.get('source_donnees')

            parcelle = Parcelle.objects.create(
                reference_parcelle=name,
                village=village,
                nombre_habitants=int(nombre_habitants) if nombre_habitants else None,
                nombre_menages=int(nombre_menages) if nombre_menages else None,
                source_donnees=source_donnees if source_donnees else None,
                photo=_get_geography_photo(data),
            )
            return JsonResponse({'success': True, 'id': parcelle.id})

        return JsonResponse({'success': False, 'error': 'Type non reconnu'})

    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@csrf_exempt
@require_http_methods(["POST"])
def api_update(request):
    """API pour modifier un élément"""
    try:
        if (request.content_type or '').startswith('application/json'):
            data = json.loads(request.body)
        else:
            data = request.POST

        element_type = data.get('type')
        element_id = data.get('id')
        name = data.get('name')

        if element_type == 'exprovince':
            ex_province = get_object_or_404(ExProvince, id=element_id)
            ex_province.nom_ex_province = name
            ex_province.save()

        elif element_type == 'region':
            region = get_object_or_404(Region, id=element_id)
            region.nom_region = name
            region.photo = _get_geography_photo(data)
            ex_province_id = data.get('ex_province_id')
            region.ex_province = get_object_or_404(ExProvince, id=ex_province_id) if ex_province_id else None
            region.save()

        elif element_type == 'district':
            district = get_object_or_404(District, id=element_id)
            district.nom_district = name
            district.photo = _get_geography_photo(data)
            district.save()

        elif element_type == 'commune':
            commune = get_object_or_404(Commune, id=element_id)
            commune.nom_commune = name
            commune.photo = _get_geography_photo(data)
            commune.save()

        elif element_type == 'fokontany':
            fokontany = get_object_or_404(Fokontany, id=element_id)
            fokontany.nom_fokontany = name
            fokontany.photo = _get_geography_photo(data)
            fokontany.save()

        elif element_type == 'village':
            village = get_object_or_404(
                Village.objects.select_related('fokontany__commune'),
                id=element_id,
            )
            fokontany_id = data.get('fokontany_id')
            if fokontany_id:
                nouveau_fokontany = get_object_or_404(
                    Fokontany.objects.select_related('commune'),
                    id=fokontany_id,
                )
                commune_actuelle_id = village.fokontany.commune_id if village.fokontany_id else None
                if commune_actuelle_id and nouveau_fokontany.commune_id != commune_actuelle_id:
                    return JsonResponse({
                        'success': False,
                        'error': 'Le fokontany sélectionné doit appartenir à la commune du village.',
                    }, status=400)
                village.fokontany = nouveau_fokontany
            village_data = _get_village_payload(data)
            for field_name, field_value in village_data.items():
                setattr(village, field_name, field_value)
            village.photo = _get_geography_photo(data)
            village.save()

        elif element_type == 'parcelle':
            parcelle = get_object_or_404(Parcelle, id=element_id)
            parcelle.reference_parcelle = name
            parcelle.photo = _get_geography_photo(data)

            nombre_habitants = data.get('nombre_habitants')
            nombre_menages = data.get('nombre_menages')
            source_donnees = data.get('source_donnees')

            if nombre_habitants:
                parcelle.nombre_habitants = int(nombre_habitants)
            if nombre_menages:
                parcelle.nombre_menages = int(nombre_menages)
            if source_donnees:
                parcelle.source_donnees = source_donnees

            # Changement de village : uniquement vers un village de la même commune
            nouveau_village_id = data.get('village_id')
            if nouveau_village_id and str(nouveau_village_id) != str(parcelle.village_id):
                nouveau_village = get_object_or_404(Village.objects.select_related('fokontany'), id=nouveau_village_id)
                commune_actuelle = parcelle.village.fokontany.commune_id
                if nouveau_village.fokontany.commune_id != commune_actuelle:
                    return JsonResponse({
                        'success': False,
                        'error': 'Le village choisi doit appartenir à la commune de la parcelle.',
                    }, status=400)
                if Parcelle.objects.filter(village=nouveau_village, reference_parcelle=parcelle.reference_parcelle).exclude(pk=parcelle.pk).exists():
                    return JsonResponse({
                        'success': False,
                        'error': f'Le village {nouveau_village.nom_village} a déjà une parcelle « {parcelle.reference_parcelle} ».',
                    }, status=400)
                parcelle.village = nouveau_village

            parcelle.save()

        return JsonResponse({'success': True})

    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@csrf_exempt
@require_http_methods(["POST"])
def api_delete(request):
    """API pour supprimer un élément"""
    try:
        if request.content_type == 'application/json':
            data = json.loads(request.body)
        else:
            data = request.POST

        element_type = data.get('type')
        element_id = data.get('id')

        if element_type == 'exprovince':
            ex_province = get_object_or_404(ExProvince, id=element_id)
            ex_province.delete()

        elif element_type == 'region':
            region = get_object_or_404(Region, id=element_id)
            region.delete()

        elif element_type == 'district':
            district = get_object_or_404(District, id=element_id)
            district.delete()

        elif element_type == 'commune':
            commune = get_object_or_404(Commune, id=element_id)
            commune.delete()

        elif element_type == 'fokontany':
            fokontany = get_object_or_404(Fokontany, id=element_id)
            fokontany.delete()

        elif element_type == 'village':
            village = get_object_or_404(Village, id=element_id)
            village.delete()

        elif element_type == 'parcelle':
            parcelle = get_object_or_404(Parcelle, id=element_id)
            parcelle.delete()

        return JsonResponse({'success': True})

    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


# ===== APIs spécifiques pour les Points d'eau =====

def api_types_point_eau(request):
    """API pour récupérer tous les types de points d'eau"""
    types = TypePointEau.objects.all().values('id', 'nom_type')
    return JsonResponse(list(types), safe=False)


def api_pointeau_detail(request, pointeau_id):
    """API pour récupérer les détails d'un point d'eau"""
    try:
        pointeau = get_object_or_404(PointEau, id=pointeau_id)
        data = {
            'id': str(pointeau.id),
            'nomabonne': pointeau.nomabonne,
            'numeroteleph': pointeau.numeroteleph,
            'contactabonne': pointeau.contactabonne,
            'typeabonne': pointeau.typeabonne,
            'type_point_eau_id': str(pointeau.type_point_eau.id) if pointeau.type_point_eau else None,
            'datebranchement': pointeau.datebranchement.isoformat() if pointeau.datebranchement else None,
            'datearretbranchement': pointeau.datearretbranchement.isoformat() if pointeau.datearretbranchement else None,
            'tarifforfaitaire': pointeau.tarifforfaitaire,
            'latitude': pointeau.latitude,
            'longitude': pointeau.longitude,
            'estbranchementinstitutionnel': pointeau.estbranchementinstitutionnel,
            'parcelle_id': str(pointeau.parcelle.id) if pointeau.parcelle else None
        }
        return JsonResponse(data)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def api_pointeau_create(request):
    """API pour créer un nouveau point d'eau"""
    try:
        data = json.loads(request.body)

        # Créer le point d'eau
        pointeau_data = {
            'nomabonne': data.get('nomabonne'),
            'numeroteleph': data.get('numeroteleph'),
            'contactabonne': data.get('contactabonne'),
            'typeabonne': data.get('typeabonne'),
            'tarifforfaitaire': data.get('tarifforfaitaire') if data.get('tarifforfaitaire') else None,
            'latitude': data.get('latitude') if data.get('latitude') else None,
            'longitude': data.get('longitude') if data.get('longitude') else None,
            'estbranchementinstitutionnel': data.get('estbranchementinstitutionnel', False),
        }

        # Relations optionnelles
        if data.get('parcelle_id'):
            parcelle = get_object_or_404(Parcelle, id=data.get('parcelle_id'))
            pointeau_data['parcelle'] = parcelle

        if data.get('type_point_eau_id'):
            type_point_eau = get_object_or_404(TypePointEau, id=data.get('type_point_eau_id'))
            pointeau_data['type_point_eau'] = type_point_eau

        # Dates avec timezone locale
        if data.get('datebranchement'):
            pointeau_data['datebranchement'] = parse_local_date(data.get('datebranchement'))

        if data.get('datearretbranchement'):
            pointeau_data['datearretbranchement'] = parse_local_date(data.get('datearretbranchement'))

        # Timestamps de création avec heure locale
        pointeau_data['createdat'] = get_local_now()

        pointeau = PointEau.objects.create(**pointeau_data)

        return JsonResponse({
            'success': True,
            'message': 'Point d\'eau créé avec succès',
            'id': str(pointeau.id)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Erreur lors de la création: {str(e)}'
        }, status=400)


@csrf_exempt
@require_http_methods(["PUT"])
def api_pointeau_update(request, pointeau_id):
    """API pour modifier un point d'eau"""
    try:
        pointeau = get_object_or_404(PointEau, id=pointeau_id)
        data = json.loads(request.body)

        # Mise à jour des champs
        pointeau.nomabonne = data.get('nomabonne', pointeau.nomabonne)
        pointeau.numeroteleph = data.get('numeroteleph', pointeau.numeroteleph)
        pointeau.contactabonne = data.get('contactabonne', pointeau.contactabonne)
        pointeau.typeabonne = data.get('typeabonne', pointeau.typeabonne)
        pointeau.estbranchementinstitutionnel = data.get('estbranchementinstitutionnel', pointeau.estbranchementinstitutionnel)

        # Champs numériques
        if data.get('tarifforfaitaire'):
            pointeau.tarifforfaitaire = data.get('tarifforfaitaire')
        if data.get('latitude'):
            pointeau.latitude = data.get('latitude')
        if data.get('longitude'):
            pointeau.longitude = data.get('longitude')

        # Relations
        if data.get('type_point_eau_id'):
            type_point_eau = get_object_or_404(TypePointEau, id=data.get('type_point_eau_id'))
            pointeau.type_point_eau = type_point_eau

        # Dates avec timezone locale
        if data.get('datebranchement'):
            pointeau.datebranchement = parse_local_date(data.get('datebranchement'))
        if data.get('datearretbranchement'):
            pointeau.datearretbranchement = parse_local_date(data.get('datearretbranchement'))

        # Timestamp de mise à jour avec heure locale
        pointeau.updatedat = get_local_now()
        employe = getattr(request.user, 'employe', None)
        if employe:
            pointeau.updated_by = employe
        pointeau.save()

        return JsonResponse({
            'success': True,
            'message': 'Point d\'eau modifié avec succès'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'Erreur lors de la modification: {str(e)}'
        }, status=400)

# ===== TABLEAU DE BORD GÉOGRAPHIQUE =====

def dashboard_geography(request):
    """Tableau de bord géographique avec statistiques et graphiques"""
    from django.db.models import Count, Sum, Q
    
    # Récupérer les filtres
    region_id = request.GET.get('region')
    district_id = request.GET.get('district')
    commune_id = request.GET.get('commune')
    village_id = request.GET.get('village')
    
    # Préparer les filtres pour les querysets
    filters = {}
    if village_id:
        filters['village_id'] = village_id
    elif commune_id:
        filters['village__fokontany__commune_id'] = commune_id
    elif district_id:
        filters['village__fokontany__commune__district_id'] = district_id
    elif region_id:
        filters['village__fokontany__commune__district__region_id'] = region_id
    
    # Statistiques globales filtrées
    if village_id:
        total_regions = 1 if Village.objects.filter(id=village_id).exists() else 0
        total_districts = 1
        total_communes = 1
        total_fokontanys = 1
        total_villages = 1
    elif commune_id:
        total_regions = 1 if Commune.objects.filter(id=commune_id).exists() else 0
        total_districts = 1
        total_communes = 1
        total_fokontanys = Fokontany.objects.filter(commune_id=commune_id).count()
        total_villages = Village.objects.filter(fokontany__commune_id=commune_id).count()
    elif district_id:
        total_regions = 1 if District.objects.filter(id=district_id).exists() else 0
        total_districts = 1
        total_communes = Commune.objects.filter(district_id=district_id).count()
        total_fokontanys = Fokontany.objects.filter(commune__district_id=district_id).count()
        total_villages = Village.objects.filter(fokontany__commune__district_id=district_id).count()
    elif region_id:
        total_regions = 1
        total_districts = District.objects.filter(region_id=region_id).count()
        total_communes = Commune.objects.filter(district__region_id=region_id).count()
        total_fokontanys = Fokontany.objects.filter(commune__district__region_id=region_id).count()
        total_villages = Village.objects.filter(fokontany__commune__district__region_id=region_id).count()
    else:
        total_regions = Region.objects.count()
        total_districts = District.objects.count()
        total_communes = Commune.objects.count()
        total_fokontanys = Fokontany.objects.count()
        total_villages = Village.objects.count()
    
    # Statistiques des parcelles
    parcelles_qs = Parcelle.objects.filter(**filters) if filters else Parcelle.objects.all()
    total_parcelles = parcelles_qs.count()
    parcelles_avec_population = parcelles_qs.filter(nombre_habitants__isnull=False).count()
    parcelles_sans_population = total_parcelles - parcelles_avec_population
    
    # Population totale
    population_stats = parcelles_qs.aggregate(
        total_habitants=Sum('nombre_habitants'),
        total_menages=Sum('nombre_menages')
    )
    total_habitants = population_stats['total_habitants'] or 0
    total_menages = population_stats['total_menages'] or 0
    
    # Listes pour les filtres en cascade
    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.filter(region_id=region_id).order_by('nom_district') if region_id else []
    communes = Commune.objects.filter(district_id=district_id).order_by('nom_commune') if district_id else []
    villages = Village.objects.filter(fokontany__commune_id=commune_id).order_by('nom_village') if commune_id else []
    
    context = {
        'title': 'Tableau de Bord Géographique',
        'total_regions': total_regions,
        'total_districts': total_districts,
        'total_communes': total_communes,
        'total_fokontanys': total_fokontanys,
        'total_villages': total_villages,
        'total_parcelles': total_parcelles,
        'parcelles_avec_population': parcelles_avec_population,
        'parcelles_sans_population': parcelles_sans_population,
        'total_habitants': total_habitants,
        'total_menages': total_menages,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'villages': villages,
        'selected_region': region_id,
        'selected_district': district_id,
        'selected_commune': commune_id,
        'selected_village': village_id,
    }
    
    return render(request, 'geography/dashboard.html', context)


def api_dashboard_chart_data(request):
    """API pour récupérer les données des graphiques du dashboard"""
    from django.db.models import Count, Sum
    
    region_id = request.GET.get('region')
    district_id = request.GET.get('district')
    commune_id = request.GET.get('commune')
    
    # Graphique 1: Répartition des parcelles par village
    if commune_id:
        villages_data = Village.objects.filter(fokontany__commune_id=commune_id).annotate(
            parcelles_count=Count('parcelles')
        ).values('nom_village', 'parcelles_count').order_by('-parcelles_count')[:10]
    elif district_id:
        villages_data = Village.objects.filter(fokontany__commune__district_id=district_id).annotate(
            parcelles_count=Count('parcelles')
        ).values('nom_village', 'parcelles_count').order_by('-parcelles_count')[:10]
    elif region_id:
        villages_data = Village.objects.filter(fokontany__commune__district__region_id=region_id).annotate(
            parcelles_count=Count('parcelles')
        ).values('nom_village', 'parcelles_count').order_by('-parcelles_count')[:10]
    else:
        villages_data = Village.objects.annotate(
            parcelles_count=Count('parcelles')
        ).values('nom_village', 'parcelles_count').order_by('-parcelles_count')[:10]
    
    # Graphique 2: Population par village
    if commune_id:
        population_data = Village.objects.filter(fokontany__commune_id=commune_id).annotate(
            total_population=Sum('parcelles__nombre_habitants')
        ).values('nom_village', 'total_population').order_by('-total_population')[:10]
    elif district_id:
        population_data = Village.objects.filter(fokontany__commune__district_id=district_id).annotate(
            total_population=Sum('parcelles__nombre_habitants')
        ).values('nom_village', 'total_population').order_by('-total_population')[:10]
    elif region_id:
        population_data = Village.objects.filter(fokontany__commune__district__region_id=region_id).annotate(
            total_population=Sum('parcelles__nombre_habitants')
        ).values('nom_village', 'total_population').order_by('-total_population')[:10]
    else:
        population_data = Village.objects.annotate(
            total_population=Sum('parcelles__nombre_habitants')
        ).values('nom_village', 'total_population').order_by('-total_population')[:10]
    
    # Graphique 3: Communes par district
    if region_id:
        districts_data = District.objects.filter(region_id=region_id).annotate(
            communes_count=Count('communes')
        ).values('nom_district', 'communes_count').order_by('-communes_count')
    else:
        districts_data = District.objects.annotate(
            communes_count=Count('communes')
        ).values('nom_district', 'communes_count').order_by('-communes_count')[:10]
    
    return JsonResponse({
        'villages_parcelles': list(villages_data),
        'villages_population': list(population_data),
        'districts_communes': list(districts_data)
    })


def api_all_communes(request):
    """API pour récupérer toutes les communes avec leur district"""
    from django.db.models import Count
    
    region_id = request.GET.get('region')
    district_id = request.GET.get('district')
    
    communes_qs = Commune.objects.select_related('district__region').annotate(
        fokontanys_count=Count('fokontanys')
    )
    
    # Appliquer les filtres si nécessaires
    if district_id:
        communes_qs = communes_qs.filter(district_id=district_id)
    elif region_id:
        communes_qs = communes_qs.filter(district__region_id=region_id)
    
    communes_data = list(communes_qs.values(
        'id',
        'nom_commune',
        'district__nom_district',
        'district__region__nom_region',
        'fokontanys_count'
    ).order_by('district__region__nom_region', 'district__nom_district', 'nom_commune'))
    
    return JsonResponse(communes_data, safe=False)


def api_all_districts(request):
    """API pour récupérer tous les districts avec leur région"""
    from django.db.models import Count
    
    region_id = request.GET.get('region')
    
    districts_qs = District.objects.select_related('region').annotate(
        communes_count=Count('communes')
    )
    
    if region_id:
        districts_qs = districts_qs.filter(region_id=region_id)
    
    districts_data = list(districts_qs.values(
        'id',
        'nom_district',
        'region__nom_region',
        'communes_count'
    ).order_by('region__nom_region', 'nom_district'))
    
    return JsonResponse(districts_data, safe=False)


def api_all_villages(request):
    """API pour récupérer tous les villages avec leur commune et district"""
    from django.db.models import Count

    commune_id = request.GET.get('commune')
    district_id = request.GET.get('district')
    region_id = request.GET.get('region')
    query = (request.GET.get('q') or '').strip()

    villages_qs = Village.objects.select_related(
        'fokontany__commune__district__region'
    ).annotate(
        parcelles_count=Count('parcelles')
    )

    if commune_id:
        villages_qs = villages_qs.filter(fokontany__commune_id=commune_id)
    elif district_id:
        villages_qs = villages_qs.filter(fokontany__commune__district_id=district_id)
    elif region_id:
        villages_qs = villages_qs.filter(fokontany__commune__district__region_id=region_id)
    elif query:
        villages_qs = villages_qs.filter(
            Q(nom_village__icontains=query) |
            Q(fokontany__nom_fokontany__icontains=query) |
            Q(fokontany__commune__nom_commune__icontains=query)
        )

    villages_data = list(villages_qs.values(
        'id',
        'nom_village',
        'fokontany__nom_fokontany',
        'fokontany__commune__nom_commune',
        'fokontany__commune__district__nom_district',
        'parcelles_count'
    ).order_by('fokontany__commune__district__nom_district', 'fokontany__commune__nom_commune', 'nom_village')[:50])

    return JsonResponse(villages_data, safe=False)



# ============================================
# EXPORT/IMPORT EXCEL
# ============================================

from django.contrib import messages
from django.shortcuts import redirect
from openpyxl import load_workbook
from .excel_utils import export_to_excel, create_excel_template


def export_excel(request, model_name):
    """Vue pour exporter les données en Excel"""
    try:
        return export_to_excel(model_name)
    except ValueError as e:
        messages.error(request, str(e))
        return redirect('geography:list')


def excel_template(request, model_name):
    """Vue pour télécharger un template Excel vide"""
    try:
        return create_excel_template(model_name)
    except ValueError as e:
        messages.error(request, str(e))
        return redirect('geography:list')


def import_excel_update_villages_gps(request):
    """Vue pour mettre à jour les coordonnées GPS des villages depuis Excel"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            wb = load_workbook(excel_file)
            ws = wb.active

            updated_count = 0
            not_found_count = 0
            errors = []

            # Format attendu : id_village | nom_village | latitude | longitude | altitude (optionnel)
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:  # Ignorer les lignes vides
                    continue

                id_village = str(row[0]).strip() if row[0] else None
                nom_village = str(row[1]).strip() if len(row) > 1 and row[1] else None
                latitude = row[2] if len(row) > 2 and row[2] else None
                longitude = row[3] if len(row) > 3 and row[3] else None
                altitude = row[4] if len(row) > 4 and row[4] else None

                if not id_village:
                    errors.append(f'Ligne {row_num}: ID du village manquant')
                    continue

                if latitude is None or longitude is None:
                    errors.append(f'Ligne {row_num}: Coordonnées GPS manquantes pour ID "{id_village}"')
                    continue

                try:
                    # Chercher le village par ID
                    try:
                        village = Village.objects.get(id=id_village)
                    except Village.DoesNotExist:
                        not_found_count += 1
                        errors.append(f'Ligne {row_num}: Village avec ID "{id_village}" introuvable')
                        continue

                    # Mettre à jour les coordonnées
                    village.latitude = float(latitude)
                    village.longitude = float(longitude)

                    if altitude is not None:
                        try:
                            village.altitude = float(altitude)
                        except (ValueError, TypeError):
                            pass

                    village.save()
                    updated_count += 1

                except (ValueError, TypeError) as e:
                    errors.append(f'Ligne {row_num}: Coordonnées invalides pour ID "{id_village}" - {str(e)}')
                except Exception as e:
                    errors.append(f'Ligne {row_num}: Erreur pour ID "{id_village}" - {str(e)}')

            # Messages de résultat
            if updated_count > 0:
                messages.success(request, f'✅ {updated_count} village(s) mis à jour avec succès !')

            if not_found_count > 0:
                messages.warning(request, f'⚠️ {not_found_count} village(s) non trouvé(s) dans la base')

            if errors and len(errors) <= 10:
                for error in errors:
                    messages.error(request, error)
            elif errors:
                messages.error(request, f'❌ {len(errors)} erreurs détectées. Premières erreurs : {"; ".join(errors[:5])}')

        except Exception as e:
            messages.error(request, f'❌ Erreur lors de l\'import : {str(e)}')

    return redirect('geography:list')


def import_excel_update_villages_accessibilite(request):
    """Vue pour mettre à jour l'accessibilité des villages depuis Excel"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            wb = load_workbook(excel_file)
            ws = wb.active

            updated_count = 0
            not_found_count = 0
            errors = []

            # Format attendu : id_village | nom_village | accessibilite | distance_dernier_point
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not row[0]:  # Ignorer les lignes vides
                    continue

                id_village = str(row[0]).strip() if row[0] else None
                nom_village = str(row[1]).strip() if len(row) > 1 and row[1] else None
                accessibilite = str(row[2]).strip().lower() if len(row) > 2 and row[2] else None
                distance = row[3] if len(row) > 3 and row[3] else None

                if not id_village:
                    errors.append(f'Ligne {row_num}: ID du village manquant')
                    continue

                try:
                    # Chercher le village par ID
                    try:
                        village = Village.objects.get(id=id_village)
                    except Village.DoesNotExist:
                        not_found_count += 1
                        errors.append(f'Ligne {row_num}: Village avec ID "{id_village}" introuvable')
                        continue

                    # Valider et convertir l'accessibilité
                    valid_choices = ['pied', 'moto', 'voiture', 'pirogue', 'goudron']
                    if accessibilite:
                        # Normaliser l'entrée (accepter variations)
                        accessibilite_map = {
                            'pied': 'pied',
                            'a pied': 'pied',
                            'à pied': 'pied',
                            'marche': 'pied',
                            'moto': 'moto',
                            'motocyclette': 'moto',
                            'voiture': 'voiture',
                            'auto': 'voiture',
                            'vehicule': 'voiture',
                            'véhicule': 'voiture',
                            'pirogue': 'pirogue',
                            'bateau': 'pirogue',
                            'goudron': 'goudron',
                            'route': 'goudron',
                            'route goudronnée': 'goudron',
                            'route goudronee': 'goudron',
                            'bitume': 'goudron',
                        }

                        accessibilite_normalized = accessibilite_map.get(accessibilite, None)

                        if accessibilite_normalized:
                            village.accessibilite = accessibilite_normalized
                        else:
                            errors.append(f'Ligne {row_num}: Accessibilité invalide "{accessibilite}" pour village "{nom_village}". Valeurs acceptées : {", ".join(valid_choices)}')
                            continue

                    # Mettre à jour la distance
                    if distance is not None and distance != '':
                        try:
                            village.distance_dernier_point = float(distance)
                        except (ValueError, TypeError):
                            errors.append(f'Ligne {row_num}: Distance invalide "{distance}" pour village "{nom_village}"')
                            continue
                    else:
                        village.distance_dernier_point = None

                    village.save()
                    updated_count += 1

                except Exception as e:
                    errors.append(f'Ligne {row_num}: Erreur pour village "{nom_village}" - {str(e)}')

            # Messages de résultat
            if updated_count > 0:
                messages.success(request, f'✅ {updated_count} village(s) mis à jour avec succès !')

            if not_found_count > 0:
                messages.warning(request, f'⚠️ {not_found_count} village(s) non trouvé(s) dans la base')

            if errors and len(errors) <= 10:
                for error in errors:
                    messages.error(request, error)
            elif errors:
                messages.error(request, f'❌ {len(errors)} erreurs détectées. Premières erreurs : {"; ".join(errors[:5])}')

        except Exception as e:
            messages.error(request, f'❌ Erreur lors de l\'import : {str(e)}')

    return redirect('geography:village_list')


def import_excel(request, model_name):
    """Vue pour importer des données depuis Excel"""
    if request.method == 'POST' and request.FILES.get('excel_file'):
        try:
            excel_file = request.FILES['excel_file']
            wb = load_workbook(excel_file)
            ws = wb.active

            imported_count = 0
            errors = []

            if model_name == 'regions':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nom_region = str(row[0]).strip() if row[0] else None
                    nom_ex_province = str(row[1]).strip() if len(row) > 1 and row[1] else None
                    if nom_region:
                        ex_province = None
                        if nom_ex_province:
                            try:
                                ex_province = ExProvince.objects.get(nom_ex_province=nom_ex_province)
                            except ExProvince.DoesNotExist:
                                errors.append(f'Ex-province "{nom_ex_province}" introuvable pour la région "{nom_region}"')
                                continue

                        region, created = Region.objects.get_or_create(
                            nom_region=nom_region,
                            defaults={'ex_province': ex_province}
                        )
                        if not created and region.ex_province_id != (ex_province.id if ex_province else None):
                            region.ex_province = ex_province
                            region.save(update_fields=['ex_province'])
                        imported_count += 1

            elif model_name == 'districts':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nom_district, nom_region = row[0], row[1]
                    if nom_district and nom_region:
                        try:
                            region = Region.objects.get(nom_region=nom_region)
                            District.objects.get_or_create(
                                nom_district=nom_district,
                                region=region
                            )
                            imported_count += 1
                        except Region.DoesNotExist:
                            errors.append(f'Région "{nom_region}" introuvable pour {nom_district}')

            elif model_name == 'communes':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nom_commune, nom_district = row[0], row[1]
                    if nom_commune and nom_district:
                        try:
                            district = District.objects.get(nom_district=nom_district)
                            Commune.objects.get_or_create(
                                nom_commune=nom_commune,
                                district=district
                            )
                            imported_count += 1
                        except District.DoesNotExist:
                            errors.append(f'District "{nom_district}" introuvable pour {nom_commune}')

            elif model_name == 'fokontanys':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nom_fokontany, commune_id = row[0], row[1]
                    if nom_fokontany and commune_id:
                        try:
                            commune = Commune.objects.get(id=commune_id)
                            Fokontany.objects.get_or_create(
                                nom_fokontany=nom_fokontany,
                                commune=commune
                            )
                            imported_count += 1
                        except Commune.DoesNotExist:
                            errors.append(f'Commune avec ID "{commune_id}" introuvable pour {nom_fokontany}')

            elif model_name == 'villages':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    nom_village, fokontany_id = row[0], row[1]
                    if nom_village and fokontany_id:
                        try:
                            fokontany = Fokontany.objects.get(id=fokontany_id)
                            Village.objects.get_or_create(
                                nom_village=nom_village,
                                fokontany=fokontany
                            )
                            imported_count += 1
                        except Fokontany.DoesNotExist:
                            errors.append(f'Fokontany avec ID "{fokontany_id}" introuvable pour {nom_village}')

            elif model_name == 'parcelles':
                for row in ws.iter_rows(min_row=2, values_only=True):
                    # Format attendu : reference_parcelle, village_id, nombre_habitants, nombre_menages, source_donnees
                    reference = row[0] if len(row) > 0 else None
                    village_id = row[1] if len(row) > 1 else None
                    habitants = row[2] if len(row) > 2 and row[2] else None
                    menages = row[3] if len(row) > 3 and row[3] else None
                    source = row[4] if len(row) > 4 else None

                    if reference and village_id:
                        try:
                            village = Village.objects.get(id=village_id)

                            # Convertir les nombres si nécessaire
                            try:
                                habitants = int(habitants) if habitants else None
                            except (ValueError, TypeError):
                                habitants = None

                            try:
                                menages = int(menages) if menages else None
                            except (ValueError, TypeError):
                                menages = None

                            Parcelle.objects.get_or_create(
                                reference_parcelle=str(reference).strip(),
                                village=village,
                                defaults={
                                    'nombre_habitants': habitants,
                                    'nombre_menages': menages,
                                    'source_donnees': str(source).strip() if source else None
                                }
                            )
                            imported_count += 1
                        except Village.DoesNotExist:
                            errors.append(f'Village avec ID "{village_id}" introuvable pour parcelle {reference}')

            if errors:
                messages.warning(request, f'{imported_count} importés. Erreurs : {"; ".join(errors[:5])}')
            else:
                messages.success(request, f'{imported_count} éléments importés avec succès !')

        except Exception as e:
            messages.error(request, f'Erreur lors de l\'import : {str(e)}')

    return redirect('geography:list')


def geography_consolidated(request):
    """Page consolidée avec toutes les données géographiques dans un tableau unique"""
    from django.db.models import Q, Sum, Count
    from django.core.paginator import Paginator

    est_directeur = user_is_director(request.user)

    # Récupération des filtres
    search_query = request.GET.get('search', '').strip()
    region_filter = request.GET.get('region', '')
    district_filter = request.GET.get('district', '')
    commune_filter = request.GET.get('commune', '')
    fokontany_filter = request.GET.get('fokontany', '')
    village_filter = request.GET.get('village', '')

    # Requête de base avec toutes les relations et comptages
    parcelles = Parcelle.objects.select_related(
        'village__fokontany__commune__district__region'
    ).annotate(
        nb_points_eau=Count('pointeau', distinct=True)
    ).order_by(
        'village__fokontany__commune__district__region__nom_region',
        'village__fokontany__commune__district__nom_district',
        'village__fokontany__commune__nom_commune',
        'village__fokontany__nom_fokontany',
        'village__nom_village',
        'reference_parcelle'
    )

    # Récupérer les colonnes de recherche sélectionnées
    search_columns = request.GET.getlist('search_columns')

    # Si aucune colonne n'est sélectionnée, rechercher dans toutes
    if not search_columns:
        search_columns = ['region', 'district', 'commune', 'fokontany', 'village', 'parcelle']

    # Appliquer la recherche sur les colonnes sélectionnées
    if search_query:
        search_filters = Q()

        if 'region' in search_columns:
            search_filters |= Q(village__fokontany__commune__district__region__nom_region__icontains=search_query)
        if 'district' in search_columns:
            search_filters |= Q(village__fokontany__commune__district__nom_district__icontains=search_query)
        if 'commune' in search_columns:
            search_filters |= Q(village__fokontany__commune__nom_commune__icontains=search_query)
        if 'fokontany' in search_columns:
            search_filters |= Q(village__fokontany__nom_fokontany__icontains=search_query)
        if 'village' in search_columns:
            search_filters |= Q(village__nom_village__icontains=search_query)
        if 'parcelle' in search_columns:
            search_filters |= Q(reference_parcelle__icontains=search_query)

        parcelles = parcelles.filter(search_filters)

    # Appliquer les filtres spécifiques
    if region_filter and region_filter != 'all':
        parcelles = parcelles.filter(village__fokontany__commune__district__region_id=region_filter)

    if district_filter and district_filter != 'all':
        parcelles = parcelles.filter(village__fokontany__commune__district_id=district_filter)

    if commune_filter and commune_filter != 'all':
        parcelles = parcelles.filter(village__fokontany__commune_id=commune_filter)

    if fokontany_filter and fokontany_filter != 'all':
        parcelles = parcelles.filter(village__fokontany_id=fokontany_filter)

    if village_filter and village_filter != 'all':
        parcelles = parcelles.filter(village_id=village_filter)

    # Statistiques globales
    stats = parcelles.aggregate(
        total_habitants=Sum('nombre_habitants'),
        total_menages=Sum('nombre_menages')
    )

    # Compter les villages uniques géolocalisés (ayant latitude ET longitude)
    villages_geolocalises = Village.objects.filter(
        parcelles__in=parcelles,
        latitude__isnull=False,
        longitude__isnull=False
    ).distinct().count()

    # Compter le total de villages uniques dans les parcelles filtrées
    total_villages = parcelles.values('village').distinct().count()

    # Compter les parcelles sans données d'habitants
    parcelles_sans_habitants = parcelles.filter(
        nombre_habitants__isnull=True
    ).count()

    # Compter les villages uniques sans données d'accessibilité
    villages_sans_accessibilite = Village.objects.filter(
        parcelles__in=parcelles,
        accessibilite__isnull=True
    ).distinct().count()

    # Pagination - 50 éléments par page
    paginator = Paginator(parcelles, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    # Listes pour les filtres
    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.all().order_by('nom_district')
    communes = Commune.objects.all().order_by('nom_commune')
    fokontanys = Fokontany.objects.all().order_by('nom_fokontany')
    villages = Village.objects.all().order_by('nom_village')

    context = {
        'page_obj': page_obj,
        'parcelles': page_obj,
        'search_query': search_query,
        'region_filter': region_filter,
        'district_filter': district_filter,
        'commune_filter': commune_filter,
        'fokontany_filter': fokontany_filter,
        'village_filter': village_filter,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'fokontanys': fokontanys,
            'villages': villages,
            'total_parcelles': paginator.count,
            'total_habitants': stats['total_habitants'] or 0,
            'total_menages': stats['total_menages'] or 0,
            'villages_geolocalises': villages_geolocalises,
            'total_villages': total_villages,
            'parcelles_sans_habitants': parcelles_sans_habitants,
            'villages_sans_accessibilite': villages_sans_accessibilite,
            'est_directeur': est_directeur,
        }

    # Passer les colonnes de recherche sélectionnées au template
    search_columns = request.GET.getlist('search_columns')
    context['search_columns'] = search_columns

    return render(request, 'geography/consolidated.html', context)


def update_village_gps_api(request, village_id):
    """API pour mettre à jour les coordonnées GPS d'un village"""
    import json
    from django.http import JsonResponse
    from django.views.decorators.csrf import csrf_exempt
    from django.views.decorators.http import require_http_methods

    if request.method == 'PUT':
        try:
            data = json.loads(request.body)
            village = Village.objects.get(id=village_id)

            # Mettre à jour les coordonnées GPS
            village.latitude = data.get('latitude')
            village.longitude = data.get('longitude')
            village.altitude = data.get('altitude')
            village.save()

            return JsonResponse({
                'success': True,
                'message': 'Coordonnées GPS mises à jour avec succès'
            })

        except Village.DoesNotExist:
            return JsonResponse({
                'success': False,
                'error': 'Village introuvable'
            }, status=404)

        except Exception as e:
            return JsonResponse({
                'success': False,
                'error': str(e)
            }, status=400)

    return JsonResponse({'success': False, 'error': 'Méthode non autorisée'}, status=405)


# ========== APIs CRUD POUR LES PARCELLES ==========

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from django.views.decorators.csrf import csrf_exempt
import json

@api_view(['PUT'])
@permission_classes([AllowAny])
def update_parcelle_api(request, parcelle_id):
    """API pour mettre à jour les données d'une parcelle (habitants, ménages)"""
    try:
        from apps.geography.models import Parcelle

        parcelle = Parcelle.objects.get(id=parcelle_id)

        # Récupérer les données du body JSON
        data = json.loads(request.body)

        # Mettre à jour les champs
        if 'nombre_habitants' in data:
            parcelle.nombre_habitants = data['nombre_habitants']
        if 'nombre_menages' in data:
            parcelle.nombre_menages = data['nombre_menages']

        parcelle.save()

        return Response({
            'success': True,
            'message': 'Parcelle mise à jour avec succès',
            'data': {
                'id': str(parcelle.id),
                'nombre_habitants': parcelle.nombre_habitants,
                'nombre_menages': parcelle.nombre_menages
            }
        })

    except Parcelle.DoesNotExist:
        return Response({'success': False, 'error': 'Parcelle non trouvée'}, status=404)
    except Exception as e:
        return Response({'success': False, 'error': str(e)}, status=500)


@api_view(['DELETE'])
@permission_classes([AllowAny])
def delete_parcelle_api(request, parcelle_id):
    """API pour supprimer une parcelle"""
    try:
        from apps.geography.models import Parcelle

        parcelle = Parcelle.objects.get(id=parcelle_id)
        reference = parcelle.reference_parcelle

        parcelle.delete()

        return Response({
            'success': True,
            'message': f'Parcelle {reference} supprimée avec succès'
        })

    except Parcelle.DoesNotExist:
        return Response({'success': False, 'error': 'Parcelle non trouvée'}, status=404)
    except Exception as e:
        return Response({'success': False, 'error': str(e)}, status=500)


# ========== PAGE CARTE GÉOGRAPHIQUE ==========

def geography_map(request):
    """Page avec carte interactive Leaflet"""
    from django.db.models import Count, Sum

    # Statistiques globales
    stats = {
        'total_regions': Region.objects.count(),
        'total_districts': District.objects.count(),
        'total_fokontany': Fokontany.objects.count(),
        'total_population': Parcelle.objects.aggregate(Sum('nombre_habitants'))['nombre_habitants__sum']
    }

    return render(request, 'geography/map.html', {'stats': stats})


def geography_liaisons_map(request):
    """Page organigramme zoomable des liaisons villages."""
    from django.db.models import Sum

    stats = {
        'total_regions': Region.objects.count(),
        'total_districts': District.objects.count(),
        'total_communes': Commune.objects.count(),
        'total_fokontany': Fokontany.objects.count(),
        'total_villages': Village.objects.count(),
        'total_liaisons': LiaisonVillage.objects.filter(actif=True).count(),
        'total_population': Parcelle.objects.aggregate(Sum('nombre_habitants'))['nombre_habitants__sum'],
    }
    return render(request, 'geography/liaisons_map.html', {'stats': stats})


# ========== APIs GEOJSON POUR LA CARTE ==========

def _tolerer_tables_shp_absentes(vue):
    """Tables de géométries (region_shp, district_shp, fokontany_shp) non chargées : carte vide
    avec avertissement au lieu d'une erreur 500."""
    from functools import wraps
    from django.db import DatabaseError, transaction

    niveau = {'api_geojson_regions': 'region', 'api_geojson_districts': 'district',
              'api_geojson_fokontany': 'fokontany'}.get(vue.__name__)

    @wraps(vue)
    def enveloppe(request, *args, **kwargs):
        try:
            with transaction.atomic():          # point de sauvegarde : une requête SQL en échec n'invalide pas la transaction en cours
                return vue(request, *args, **kwargs)
        except DatabaseError:
            # Pas de PostGIS : on se rabat sur les contours importés des shapefiles (ContourAdmin), s'il y en a.
            from .views_contours import _lire_bbox, contours_features
            bbox = _lire_bbox(request.GET.get('bbox', '')) if request.GET.get('bbox') else None
            if niveau and (niveau != 'fokontany' or bbox):       # les 17 000 fokontany ne se servent que par zone
                try:
                    zoom = float(request.GET['zoom']) if 'zoom' in request.GET else None
                except ValueError:
                    zoom = None
                features, tronque = contours_features(niveau, bbox, zoom)
                if features:
                    return Response({'type': 'FeatureCollection', 'features': features, 'tronque': tronque})
            return Response({
                'type': 'FeatureCollection',
                'features': [],
                'warning': "Géométries non disponibles : table des contours (shapefile) non chargée.",
            })
    return enveloppe


@api_view(['GET'])
@permission_classes([AllowAny])
@_tolerer_tables_shp_absentes
def api_geojson_regions(request):
    """API GeoJSON pour les régions avec leurs géométries PostGIS"""
    from django.db import connection
    from django.db.models import Count, Sum

    # Récupérer les statistiques des régions depuis les tables métier
    regions_stats = Region.objects.annotate(
        nb_districts=Count('districts'),
        population=Sum('districts__communes__fokontanys__villages__parcelles__nombre_habitants')
    ).values('id', 'nom_region', 'nb_districts', 'population')

    # Créer un dictionnaire pour lookup rapide
    stats_dict = {str(r['id']): r for r in regions_stats}

    # Récupérer les géométries depuis region_shp
    with connection.cursor() as cursor:
        cursor.execute("""
            SELECT
                r.id,
                r.nom_region,
                ST_AsGeoJSON(ST_Transform(ST_SetSRID(rs.geom, 32738), 4326)) as geojson,
                rs.reg_pcode
            FROM regions r
            LEFT JOIN region_shp rs ON LOWER(TRIM(r.nom_region)) = LOWER(TRIM(rs.region))
            WHERE rs.geom IS NOT NULL
            ORDER BY r.nom_region
        """)

        features = []
        for row in cursor.fetchall():
            region_id = str(row[0])
            stats = stats_dict.get(region_id, {})

            feature = {
                'type': 'Feature',
                'properties': {
                    'id': region_id,
                    'nom': row[1],
                    'code': row[3],
                    'nb_districts': stats.get('nb_districts', 0),
                    'population': stats.get('population')
                },
                'geometry': json.loads(row[2]) if row[2] else None
            }
            features.append(feature)

    geojson = {
        'type': 'FeatureCollection',
        'features': features
    }

    return Response(geojson)


@api_view(['GET'])
@permission_classes([AllowAny])
@_tolerer_tables_shp_absentes
def api_geojson_districts(request):
    """API GeoJSON pour les districts"""
    from django.db import connection
    from django.db.models import Sum

    # Récupérer les statistiques
    districts_stats = District.objects.select_related('region').annotate(
        population=Sum('communes__fokontanys__villages__parcelles__nombre_habitants')
    ).values('id', 'nom_district', 'region__nom_region', 'population')

    stats_dict = {str(d['id']): d for d in districts_stats}

    # Récupérer les géométries
    with connection.cursor() as cursor:
        cursor.execute("""
            SELECT
                d.id,
                d.nom_district,
                r.nom_region,
                ST_AsGeoJSON(ST_Transform(ST_SetSRID(ds.geom, 32738), 4326)) as geojson
            FROM districts d
            INNER JOIN regions r ON d.region_id = r.id
            LEFT JOIN district_shp ds ON LOWER(TRIM(d.nom_district)) = LOWER(TRIM(ds.district))
            WHERE ds.geom IS NOT NULL
            ORDER BY d.nom_district
        """)

        features = []
        for row in cursor.fetchall():
            district_id = str(row[0])
            stats = stats_dict.get(district_id, {})

            feature = {
                'type': 'Feature',
                'properties': {
                    'id': district_id,
                    'nom': row[1],
                    'region': row[2],
                    'population': stats.get('population')
                },
                'geometry': json.loads(row[3]) if row[3] else None
            }
            features.append(feature)

    return Response({'type': 'FeatureCollection', 'features': features})


@api_view(['GET'])
@permission_classes([AllowAny])
@_tolerer_tables_shp_absentes
def api_geojson_fokontany(request):
    """API GeoJSON pour les fokontany"""
    from django.db import connection
    from django.db.models import Sum

    # Statistiques
    fokontany_stats = Fokontany.objects.select_related('commune').annotate(
        population=Sum('villages__parcelles__nombre_habitants')
    ).values('id', 'nom_fokontany', 'commune__nom_commune', 'population')

    stats_dict = {str(f['id']): f for f in fokontany_stats}

    # Géométries
    with connection.cursor() as cursor:
        cursor.execute("""
            SELECT
                f.id,
                f.nom_fokontany,
                c.nom_commune,
                ST_AsGeoJSON(ST_Transform(ST_SetSRID(fs.geom, 32738), 4326)) as geojson
            FROM fokontanys f
            INNER JOIN communes c ON f.commune_id = c.id
            LEFT JOIN fokontany_shp fs ON LOWER(TRIM(f.nom_fokontany)) = LOWER(TRIM(fs.fokontany))
            WHERE fs.geom IS NOT NULL
            ORDER BY f.nom_fokontany
        """)

        features = []
        for row in cursor.fetchall():
            fokontany_id = str(row[0])
            stats = stats_dict.get(fokontany_id, {})

            feature = {
                'type': 'Feature',
                'properties': {
                    'id': fokontany_id,
                    'nom': row[1],
                    'commune': row[2],
                    'population': stats.get('population')
                },
                'geometry': json.loads(row[3]) if row[3] else None
            }
            features.append(feature)

    return Response({'type': 'FeatureCollection', 'features': features})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_villages(request):
    """API GeoJSON pour les villages avec coordonnées GPS"""
    villages = Village.objects.select_related(
        'fokontany__commune__district__region'
    ).filter(
        latitude__isnull=False,
        longitude__isnull=False
    ).values(
        'id', 'nom_village',
        'latitude', 'longitude',
        'fokontany__nom_fokontany',
        'fokontany__commune__nom_commune',
        'fokontany__commune__district__nom_district',
        'fokontany__commune__district__region__nom_region'
    )

    features = []
    for village in villages:
        feature = {
            'type': 'Feature',
            'properties': {
                'id': str(village['id']),
                'nom': village['nom_village'],
                'fokontany': village['fokontany__nom_fokontany'],
                'commune': village['fokontany__commune__nom_commune'],
                'district': village['fokontany__commune__district__nom_district'],
                'region': village['fokontany__commune__district__region__nom_region']
            },
            'geometry': {
                'type': 'Point',
                'coordinates': [float(village['longitude']), float(village['latitude'])]
            }
        }
        features.append(feature)

    return Response({'type': 'FeatureCollection', 'features': features})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_communes(request):
    """API GeoJSON pour les communes sous forme de points centroides calcules depuis les villages GPS."""
    communes = Commune.objects.select_related('district__region').prefetch_related('fokontanys__villages')

    features = []
    for commune in communes:
        villages = [
            village
            for fokontany in commune.fokontanys.all()
            for village in fokontany.villages.all()
            if village.latitude is not None and village.longitude is not None
        ]
        if not villages:
            continue

        avg_lat = sum(float(v.latitude) for v in villages) / len(villages)
        avg_lon = sum(float(v.longitude) for v in villages) / len(villages)
        sample_villages = sorted(v.nom_village for v in villages)[:4]

        features.append({
            'type': 'Feature',
            'properties': {
                'id': str(commune.id),
                'nom': commune.nom_commune,
                'district': commune.district.nom_district,
                'region': commune.district.region.nom_region,
                'villages_count': len(villages),
                'sample_villages': sample_villages,
            },
            'geometry': {
                'type': 'Point',
                'coordinates': [avg_lon, avg_lat],
            }
        })

    return Response({'type': 'FeatureCollection', 'features': features})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_liaisons_villages(request):
    """API GeoJSON des liaisons entre villages pour visualisation cartographique.

    `?bbox=ouest,sud,est,nord` : ne renvoie que les liaisons dont au moins un village est dans la zone
    (la couche complète pèse ~15 Mo : les cartes la chargent par zone visible).
    """
    liaisons = LiaisonVillage.objects.filter(
        actif=True,
        village_source__latitude__isnull=False,
        village_source__longitude__isnull=False,
        village_cible__latitude__isnull=False,
        village_cible__longitude__isnull=False,
    ).select_related(
        'village_source__fokontany__commune',
        'village_cible__fokontany__commune',
    )

    bbox = request.GET.get('bbox')
    if bbox:
        try:
            ouest, sud, est, nord = (float(v) for v in bbox.split(','))
        except ValueError:
            return Response({'error': 'bbox invalide (ouest,sud,est,nord).'}, status=400)
        dans_zone = lambda champ: Q(**{
            f'{champ}__longitude__gte': ouest, f'{champ}__longitude__lte': est,
            f'{champ}__latitude__gte': sud, f'{champ}__latitude__lte': nord,
        })
        liaisons = liaisons.filter(dans_zone('village_source') | dans_zone('village_cible'))

    from .views_liaisons_api import liaison_feature
    features = [liaison_feature(liaison) for liaison in liaisons]

    return Response({'type': 'FeatureCollection', 'features': features})


# Catégories affichées selon le zoom, simplification du tracé (degrés) et plafond de tronçons par réponse.
_RESEAU_PAR_ZOOM = [
    (12, None, 0.00004),                                   # zoom >= 12 : tout le réseau
    (10, ('nationale', 'provinciale', 'communale', 'route'), 0.0004),
    (8, ('nationale', 'provinciale'), 0.002),
]
_RESEAU_MAX_TRONCONS = 6000


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_routes_reseau(request):
    """GeoJSON des routes du réseau routier dans une zone : ?bbox=ouest,sud,est,nord&zoom=N.

    Plus on est dézoomé, moins il y a de catégories et plus le tracé est simplifié ; en dessous du zoom 8
    rien n'est renvoyé. Les routes les plus importantes passent en premier si le plafond est atteint.
    """
    from shapely.geometry import box

    from .services import index_spatial
    from .services.voies_config import RANG_CATEGORIE

    try:
        ouest, sud, est, nord = (float(v) for v in request.GET.get('bbox', '').split(','))
        zoom = int(float(request.GET.get('zoom', 12)))
    except ValueError:
        return Response({'error': 'bbox (ouest,sud,est,nord) et zoom requis.'}, status=400)

    regle = next(((cats, tol) for z, cats, tol in _RESEAU_PAR_ZOOM if zoom >= z), None)
    if regle is None:
        return Response({'type': 'FeatureCollection', 'features': [], 'zoom_min': _RESEAU_PAR_ZOOM[-1][0]})
    categories, tolerance = regle

    ix = index_spatial.index()
    candidats = [int(i) for i in ix.arbre.query(box(ouest, sud, est, nord))]
    if categories is not None:
        candidats = [i for i in candidats if ix.categories[i] in categories]
    rang = {c: n for n, c in enumerate(RANG_CATEGORIE)}
    candidats.sort(key=lambda i: rang.get(ix.categories[i], len(rang)))
    tronque = len(candidats) > _RESEAU_MAX_TRONCONS
    features = []
    for i in candidats[:_RESEAU_MAX_TRONCONS]:
        ligne = ix.lignes[i].simplify(tolerance, preserve_topology=False) if tolerance else ix.lignes[i]
        coords = [[round(x, 5), round(y, 5)] for x, y in ligne.coords]
        if len(coords) < 2:
            continue
        features.append({
            'type': 'Feature',
            'properties': {'categorie': ix.categories[i], 'source': ix.sources[i]},
            'geometry': {'type': 'LineString', 'coordinates': coords},
        })
    return Response({'type': 'FeatureCollection', 'features': features, 'tronque': tronque})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_villages_non_relies(request):
    """Villages mal reliés au réseau routier : ?seuil=<mètres> (défaut 1000).

    Un village est « non relié » s'il est sur un morceau de réseau isolé du réseau principal (aucun itinéraire
    possible vers le reste du pays) et/ou à plus de `seuil` mètres de la route la plus proche (long trajet à pied).
    Pour ces derniers, un segment relie le village à son point d'accès à la route.
    """
    from .models_routes import VillageAccesRoute

    try:
        seuil = max(0.0, float(request.GET.get('seuil', 1000)))
    except ValueError:
        return Response({'error': 'seuil invalide.'}, status=400)

    acces = (VillageAccesRoute.objects
             .filter(Q(reseau_principal=False) | Q(distance_m__gt=seuil),
                     village__latitude__isnull=False, village__longitude__isnull=False)
             .select_related('village__fokontany__commune'))
    features, nb_isoles, nb_eloignes = [], 0, 0
    for a in acces:
        v = a.village
        motifs = []
        if not a.reseau_principal:
            motifs.append('isole')
            nb_isoles += 1
        if a.distance_m > seuil:
            motifs.append('eloigne')
            nb_eloignes += 1
        lon, lat = float(v.longitude), float(v.latitude)
        features.append({
            'type': 'Feature',
            'properties': {
                'type': 'village', 'id': str(v.id), 'nom': v.nom_village, 'motifs': motifs,
                'commune': v.fokontany.commune.nom_commune if v.fokontany_id else None,
                'distance_route_m': round(a.distance_m),
            },
            'geometry': {'type': 'Point', 'coordinates': [lon, lat]},
        })
        if 'eloigne' in motifs:
            features.append({
                'type': 'Feature',
                'properties': {'type': 'acces', 'id': str(v.id)},
                'geometry': {'type': 'LineString', 'coordinates': [[lon, lat], [float(a.longitude), float(a.latitude)]]},
            })
    return Response({'type': 'FeatureCollection', 'features': features, 'seuil_m': seuil,
                     'nb_isoles': nb_isoles, 'nb_eloignes': nb_eloignes})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_geojson_bases(request):
    """GeoJSON des bases placées sur la carte (lieu de départ des déplacements de leurs agents).

    Les bases sans lieu ne sont pas dessinées : elles sont comptées dans `nb_non_placees`.
    """
    from apps.organization.models import Base

    features, non_placees = [], []
    for base in Base.objects.select_related('organisation', 'village_depart').order_by('nom_base'):
        lieu = base.lieu_depart
        if lieu is None:
            non_placees.append(base.nom_base)
            continue
        features.append({
            'type': 'Feature',
            'properties': {
                'id': str(base.id), 'nom': base.nom_base, 'organisation': base.organisation.nom_organisation,
                'lieu': lieu[2], 'sur_village': base.latitude_depart is None,
            },
            'geometry': {'type': 'Point', 'coordinates': [lieu[0], lieu[1]]},
        })
    return Response({'type': 'FeatureCollection', 'features': features,
                     'nb_non_placees': len(non_placees), 'non_placees': non_placees})


@api_view(['GET'])
@permission_classes([AllowAny])
def api_search_zones(request):
    """API pour rechercher des zones géographiques"""
    query = request.GET.get('q', '').strip()

    if len(query) < 2:
        return Response([])

    results = []

    # Rechercher dans régions
    regions = Region.objects.filter(nom_region__icontains=query)[:5]
    for r in regions:
        results.append({
            'id': str(r.id),
            'nom': r.nom_region,
            'type': 'Région'
        })

    # Rechercher dans districts
    districts = District.objects.filter(nom_district__icontains=query).select_related('region')[:5]
    for d in districts:
        results.append({
            'id': str(d.id),
            'nom': f"{d.nom_district} ({d.region.nom_region})",
            'type': 'District'
        })

    # Rechercher dans les communes (contours importés : 1 585 communes, toutes ne sont pas en base)
    from .models_contours import ContourAdmin
    for k in ContourAdmin.objects.filter(niveau='commune', nom__icontains=query).order_by('nom')[:5]:
        results.append({'id': str(k.objet_id or k.pk), 'nom': f"{k.nom} ({k.district_nom})", 'type': 'Commune'})

    # Rechercher dans fokontany
    fokontany = Fokontany.objects.filter(nom_fokontany__icontains=query).select_related('commune')[:5]
    for f in fokontany:
        results.append({
            'id': str(f.id),
            'nom': f"{f.nom_fokontany} ({f.commune.nom_commune})",
            'type': 'Fokontany'
        })

    return Response(results)


@api_view(['GET'])
@permission_classes([AllowAny])
def api_zone_bounds(request, zone_type, zone_id):
    """API pour récupérer les bounds d'une zone"""
    from django.db import connection

    table_map = {
        'Région': 'region_shp',
        'District': 'district_shp',
        'Fokontany': 'fokontany_shp'
    }

    join_map = {
        'Région': ('regions', 'nom_region', 'region'),
        'District': ('districts', 'nom_district', 'district'),
        'Fokontany': ('fokontanys', 'nom_fokontany', 'fokontany')
    }

    # Contours importés des shapefiles (ContourAdmin) : aucune dépendance à PostGIS.
    from django.core.exceptions import ValidationError
    from .models_contours import ContourAdmin
    niveau = {'Région': 'region', 'District': 'district', 'Commune': 'commune', 'Fokontany': 'fokontany'}.get(zone_type)
    if niveau is None:
        return Response({'error': 'Type de zone invalide'}, status=400)
    try:
        contour = (ContourAdmin.objects.filter(niveau=niveau, objet_id=zone_id).first()
                   or ContourAdmin.objects.filter(niveau=niveau, pk=zone_id).first())
    except (ValidationError, ValueError):
        contour = None
    if contour is not None:
        return Response({'bounds': {'west': contour.ouest, 'south': contour.sud, 'east': contour.est, 'north': contour.nord}})

    shp_table = table_map.get(zone_type)
    model_table, model_col, shp_col = join_map.get(zone_type, (None, None, None))

    if not shp_table or not model_table:
        return Response({'error': 'Zone non trouvée'}, status=404)

    from django.db import DatabaseError
    try:
        return _bounds_postgis(connection, model_table, shp_table, model_col, shp_col, zone_id)
    except DatabaseError:                       # tables *_shp absentes (pas de PostGIS)
        return Response({'error': 'Zone non trouvée'}, status=404)


def _bounds_postgis(connection, model_table, shp_table, model_col, shp_col, zone_id):
    with connection.cursor() as cursor:
        cursor.execute(f"""
            SELECT
                ST_XMin(ST_Extent(ST_Transform(ST_SetSRID(s.geom, 32738), 4326))) as west,
                ST_YMin(ST_Extent(ST_Transform(ST_SetSRID(s.geom, 32738), 4326))) as south,
                ST_XMax(ST_Extent(ST_Transform(ST_SetSRID(s.geom, 32738), 4326))) as east,
                ST_YMax(ST_Extent(ST_Transform(ST_SetSRID(s.geom, 32738), 4326))) as north
            FROM {model_table} m
            LEFT JOIN {shp_table} s ON LOWER(TRIM(m.{model_col})) = LOWER(TRIM(s.{shp_col}))
            WHERE m.id = %s AND s.geom IS NOT NULL
        """, [zone_id])

        row = cursor.fetchone()
        if row and row[0]:
            return Response({
                'bounds': {
                    'west': row[0],
                    'south': row[1],
                    'east': row[2],
                    'north': row[3]
                }
            })

    return Response({'error': 'Zone non trouvée'}, status=404)
