from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from .models import Region, District, Commune, Fokontany, Village, Parcelle
from apps.billing.models import PointEau
from apps.organization.models import Ouvrage

# Import pour Excel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter






def export_excel_data(request, model_name):
    """Export complet des données géographiques en Excel"""
    models_export = {
        'regions': {
            'model': Region,
            'fields': ['id', 'nom_region'],
            'filename': 'export_regions.xlsx',
            'sheet_name': 'Régions'
        },
        'districts': {
            'model': District,
            'fields': ['id', 'nom_district', 'region_id'],
            'filename': 'export_districts.xlsx',
            'sheet_name': 'Districts'
        },
        'communes': {
            'model': Commune,
            'fields': ['id', 'nom_commune', 'district__nom_district', 'district_id', 'district__region__nom_region'],
            'filename': 'export_communes.xlsx',
            'sheet_name': 'Communes'
        },
        'villages': {
            'model': Village,
            'fields': ['id', 'nom_village', 'fokontany__nom_fokontany', 'fokontany_id', 'fokontany__commune__nom_commune', 'fokontany__commune__district__nom_district'],
            'filename': 'export_villages.xlsx',
            'sheet_name': 'Villages'
        },
        'parcelles': {
            'model': Parcelle,
            'fields': ['id', 'reference_parcelle', 'village_id', 'village__nom_village', 'nombre_habitants', 'nombre_menages', 'source_donnees'],
            'filename': 'export_parcelles.xlsx',
            'sheet_name': 'Parcelles'
        },
        'points_eau': {
            'model': PointEau,
            'fields': ['id', 'nomabonne', 'parcelle_id', 'parcelle__reference_parcelle', 'ouvrage_id', 'ouvrage__nom', 'typeabonne', 'numeroteleph', 'contactabonne', 'latitude', 'longitude'],
            'filename': 'export_points_eau.xlsx',
            'sheet_name': 'Points d\'Eau'
        }
    }

    if model_name not in models_export:
        return JsonResponse({'error': 'Type d\'export non reconnu'}, status=404)

    config = models_export[model_name]
    model_class = config['model']
    fields = config['fields']

    # Créer un workbook Excel
    wb = Workbook()
    ws = wb.active
    ws.title = config['sheet_name']

    # Créer les en-têtes avec des noms plus clairs
    headers = []
    for field in fields:
        if field == 'id':
            headers.append('ID')
        elif '__' in field:
            if field.endswith('__nom_region'):
                headers.append('Région')
            elif field.endswith('__nom_district'):
                headers.append('District')
            elif field.endswith('__nom_commune'):
                headers.append('Commune')
            elif field.endswith('__nom_village'):
                headers.append('Village')
            else:
                headers.append(field.replace('__', ' - ').replace('_', ' ').title())
        elif field.endswith('_id'):
            field_name = field.replace('_id', '').replace('_', ' ').title()
            headers.append(f'{field_name} (ID)')
        else:
            headers.append(field.replace('_', ' ').title())

    # Styliser les en-têtes
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Ajouter les en-têtes
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Exporter toutes les données
    try:
        if model_name == 'regions':
            queryset = model_class.objects.all()
        elif model_name == 'districts':
            queryset = model_class.objects.select_related('region').all()
        elif model_name == 'communes':
            queryset = model_class.objects.select_related('district', 'district__region').all()
        elif model_name == 'villages':
            queryset = model_class.objects.select_related('fokontany', 'fokontany__commune', 'fokontany__commune__district').all()
        elif model_name == 'parcelles':
            queryset = model_class.objects.select_related('village', 'village__commune', 'village__commune__district').all()
        elif model_name == 'points_eau':
            queryset = model_class.objects.select_related('parcelle', 'ouvrage').all()

        row_num = 2
        for obj in queryset:
            for col_num, field in enumerate(fields, 1):
                try:
                    if '__' in field:
                        field_parts = field.split('__')
                        value = obj
                        for part in field_parts:
                            if value is None:
                                break
                            value = getattr(value, part, None)
                        ws.cell(row=row_num, column=col_num, value=value if value is not None else '')
                    else:
                        value = getattr(obj, field, '')
                        if value is None:
                            ws.cell(row=row_num, column=col_num, value='')
                        else:
                            ws.cell(row=row_num, column=col_num, value=str(value))
                except Exception:
                    ws.cell(row=row_num, column=col_num, value='')
            row_num += 1

        # Auto-ajuster la largeur des colonnes
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    except Exception as e:
        return JsonResponse({'error': f'Erreur lors de l\'export: {str(e)}'}, status=500)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{config["filename"]}"'
    wb.save(response)
    return response


def download_excel_template(request, model_name):
    """Téléchargement des templates Excel pour la géographie"""
    templates = {
        'regions': {
            'fields': ['nom_region'],
            'examples': [
                ['Analamanga'],
                ['Vakinankaratra'],
                ['Haute Matsiatra']
            ]
        },
        'districts': {
            'fields': ['nom_district', 'region'],
            'examples': [
                ['Antananarivo Renivohitra', 'Analamanga'],
                ['Antananarivo Atsimondrano', 'Analamanga'],
                ['Antsirabe I', 'Vakinankaratra']
            ]
        },
        'communes': {
            'fields': ['nom_commune', 'district_id'],
            'examples': [
                ['Ambato', '0973a37b-e757-4a3a-9154-78a5143fb7e8'],
                ['Antsahafilo', '0973a37b-e757-4a3a-9154-78a5143fb7e8']
            ]
        },
        'villages': {
            'fields': ['nom_village', 'fokontany_id'],
            'examples': [
                ['Ambato', 'uuid-du-fokontany'],
                ['Antanifotsy', 'uuid-du-fokontany']
            ]
        },
        'parcelles': {
            'fields': ['reference_parcelle', 'village_id', 'nombre_habitants', 'nombre_menages', 'source_donnees'],
            'examples': [
                ['PARC-001', 'f4b21cc8-1e0a-40ca-8ce7-c0a73baa9c73', '150', '30', 'Recensement 2023'],
                ['PARC-002', 'f4b21cc8-1e0a-40ca-8ce7-c0a73baa9c73', '200', '42', 'Enquête locale 2024']
            ]
        },
        'points_eau': {
            'fields': ['nom_abonne', 'parcelle_id', 'ouvrage_id', 'type_abonne', 'numero_tel', 'contact', 'latitude', 'longitude'],
            'examples': [
                ['Rakoto Jean', 'f4b21cc8-1e0a-40ca-8ce7-c0a73baa9c73', 'a1b2c3d4-5e6f-7g8h-9i0j-k1l2m3n4o5p6', 'particulier', '0340123456', 'Rakoto', '-18.8792', '47.5079'],
                ['SARL Madagascar', 'f4b21cc8-1e0a-40ca-8ce7-c0a73baa9c73', 'a1b2c3d4-5e6f-7g8h-9i0j-k1l2m3n4o5p6', 'commercial', '0331234567', 'Directeur', '-18.8800', '47.5100']
            ]
        }
    }

    if model_name not in templates:
        return JsonResponse({'error': 'Template non trouvé'}, status=404)

    template = templates[model_name]

    wb = Workbook()
    ws = wb.active
    ws.title = f'Template {model_name.title()}'

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = []
    for field in template['fields']:
        headers.append(field.replace('_', ' ').title())

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_num, example in enumerate(template['examples'], 2):
        for col_num, value in enumerate(example, 1):
            ws.cell(row=row_num, column=col_num, value=value)

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="template_{model_name}.xlsx"'
    wb.save(response)
    return response


def import_excel_view(request, model_name):
    """Vue générique pour l'import Excel de la géographie"""
    models = {
        'regions': Region,
        'districts': District,
        'communes': Commune,
        'villages': Village,
        'parcelles': Parcelle,
        'points_eau': PointEau
    }

    if model_name not in models:
        messages.error(request, f'Type de données non reconnu: {model_name}')
        return redirect('geography:list')

    if request.method == 'POST':
        if 'excel_file' not in request.FILES:
            messages.error(request, 'Aucun fichier sélectionné.')
            return redirect(request.META.get('HTTP_REFERER', 'geography:list'))

        excel_file = request.FILES['excel_file']

        if not (excel_file.name.endswith('.xlsx') or excel_file.name.endswith('.xls')):
            messages.error(request, 'Le fichier doit être un fichier Excel.')
            return redirect(request.META.get('HTTP_REFERER', 'geography:list'))

        try:
            wb = load_workbook(excel_file)
            ws = wb.active

            headers = []
            for cell in ws[1]:
                headers.append(cell.value.lower().replace(' ', '_') if cell.value else '')

            created_count = 0
            error_count = 0

            for row_num in range(2, ws.max_row + 1):
                try:
                    row_data = {}
                    for col_num, header in enumerate(headers, 1):
                        cell_value = ws.cell(row=row_num, column=col_num).value
                        row_data[header] = str(cell_value).strip() if cell_value is not None else ''

                    if model_name == 'regions':
                        Region.objects.get_or_create(
                            nom_region=row_data.get('nom_region', '').strip()
                        )
                    elif model_name == 'districts':
                        # Accepter soit le nom de la région, soit le region_id (UUID)
                        region_id = row_data.get('region_(id)', '').strip() or row_data.get('region_id', '').strip()
                        region_nom = row_data.get('region', '').strip() or row_data.get('région', '').strip()

                        # Rechercher la région soit par ID soit par nom
                        if region_id:
                            region = Region.objects.get(id=region_id)
                        elif region_nom:
                            region = Region.objects.get(nom_region=region_nom)
                        else:
                            raise ValueError("Le nom de la région ou region_id est requis")

                        District.objects.get_or_create(
                            nom_district=row_data.get('nom_district', '').strip(),
                            defaults={'region': region}
                        )
                    elif model_name == 'communes':
                        # Utiliser district_id (UUID) pour éviter les ambiguïtés
                        district_id = row_data.get('district_(id)', '').strip() or row_data.get('district_id', '').strip()
                        if not district_id:
                            raise ValueError("district_id est requis")
                        district = District.objects.get(id=district_id)
                        Commune.objects.get_or_create(
                            nom_commune=row_data.get('nom_commune', '').strip(),
                            defaults={'district': district}
                        )
                    elif model_name == 'villages':
                        # Utiliser fokontany_id (UUID) pour éviter les doublons de noms
                        fokontany_id = row_data.get('fokontany_(id)', '').strip() or row_data.get('fokontany_id', '').strip()
                        if not fokontany_id:
                            raise ValueError("fokontany_id est requis")
                        fokontany = Fokontany.objects.get(id=fokontany_id)
                        Village.objects.get_or_create(
                            nom_village=row_data.get('nom_village', '').strip(),
                            defaults={'fokontany': fokontany}
                        )
                    elif model_name == 'parcelles':
                        # Utiliser village_id (UUID) pour éviter les doublons de noms
                        village_id = row_data.get('village_(id)', '').strip() or row_data.get('village_id', '').strip() or row_data.get('id_village_(uuid)', '').strip()
                        if not village_id:
                            raise ValueError("village_id est requis")
                        village = Village.objects.get(id=village_id)

                        # Récupérer les données démographiques optionnelles
                        nombre_habitants = row_data.get('nombre_habitants', '').strip() or row_data.get('nombre_habitants_(optionnel)', '').strip()
                        nombre_menages = row_data.get('nombre_menages', '').strip() or row_data.get('nombre_ménages_(optionnel)', '').strip()
                        source_donnees = row_data.get('source_donnees', '').strip() or row_data.get('source_données_(optionnel)', '').strip()

                        # Convertir en entiers si présents
                        nombre_habitants = int(nombre_habitants) if nombre_habitants and nombre_habitants.isdigit() else None
                        nombre_menages = int(nombre_menages) if nombre_menages and nombre_menages.isdigit() else None
                        source_donnees = source_donnees if source_donnees else None

                        Parcelle.objects.get_or_create(
                            reference_parcelle=row_data.get('reference_parcelle', '').strip() or row_data.get('référence_parcelle', '').strip(),
                            village=village,
                            defaults={
                                'nombre_habitants': nombre_habitants,
                                'nombre_menages': nombre_menages,
                                'source_donnees': source_donnees
                            }
                        )
                    elif model_name == 'points_eau':
                        # Format: nom_abonne, parcelle_id, ouvrage_id, type_abonne, numero_tel, contact, latitude, longitude
                        nom_abonne = row_data.get('nom_abonne', '').strip() or row_data.get('nom_abonné', '').strip() or row_data.get('nom_abonné_*', '').strip()
                        parcelle_id = row_data.get('parcelle_id', '').strip() or row_data.get('id_parcelle_(uuid)', '').strip() or row_data.get('id_parcelle_(uuid)_*', '').strip()
                        ouvrage_id = row_data.get('ouvrage_id', '').strip() or row_data.get('id_ouvrage_(uuid)', '').strip() or row_data.get('id_ouvrage_(uuid)_*', '').strip()
                        type_abonne = row_data.get('type_abonne', '').strip() or row_data.get('type_abonné_(particulier/commercial/institution)', '').strip() or row_data.get('type_abonné_*_(particulier/commercial/institution)', '').strip()
                        numero_tel = row_data.get('numero_tel', '').strip() or row_data.get('numéro_tél_(optionnel)', '').strip()
                        contact = row_data.get('contact', '').strip() or row_data.get('contact_(optionnel)', '').strip()
                        latitude = row_data.get('latitude', '').strip() or row_data.get('latitude_(optionnel)', '').strip()
                        longitude = row_data.get('longitude', '').strip() or row_data.get('longitude_(optionnel)', '').strip()

                        # Validation des champs requis
                        if not nom_abonne:
                            raise ValueError("Nom abonné est requis")
                        if not parcelle_id:
                            raise ValueError("ID Parcelle est requis")
                        if not ouvrage_id:
                            raise ValueError("ID Ouvrage est requis")
                        if not type_abonne:
                            raise ValueError("Type abonné est requis")

                        # Valider le type d'abonné
                        type_abonne_lower = type_abonne.lower()
                        if type_abonne_lower not in ['particulier', 'commercial', 'institution']:
                            raise ValueError(f"Type abonné invalide: '{type_abonne}'. Valeurs acceptées: particulier, commercial, institution")

                        # Récupérer parcelle et ouvrage (maintenant requis)
                        parcelle = Parcelle.objects.get(id=parcelle_id)
                        ouvrage = Ouvrage.objects.get(id=ouvrage_id)

                        # Convertir latitude/longitude
                        try:
                            latitude = float(latitude) if latitude else None
                        except (ValueError, TypeError):
                            latitude = None

                        try:
                            longitude = float(longitude) if longitude else None
                        except (ValueError, TypeError):
                            longitude = None

                        PointEau.objects.get_or_create(
                            nomabonne=nom_abonne,
                            parcelle=parcelle,
                            defaults={
                                'ouvrage': ouvrage,
                                'typeabonne': type_abonne_lower,
                                'numeroteleph': numero_tel if numero_tel else None,
                                'contactabonne': contact if contact else None,
                                'latitude': latitude,
                                'longitude': longitude
                            }
                        )

                    created_count += 1

                except Exception as e:
                    error_count += 1

            if created_count > 0:
                messages.success(request, f'{created_count} enregistrement(s) importé(s) avec succès depuis Excel.')
            if error_count > 0:
                messages.warning(request, f'{error_count} erreur(s) rencontrée(s) lors de l\'import Excel.')

        except Exception as e:
            messages.error(request, f'Erreur lors de la lecture du fichier Excel: {str(e)}')

    return redirect(request.META.get('HTTP_REFERER', 'geography:list'))