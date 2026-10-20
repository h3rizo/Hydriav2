"""Utilitaires pour l'export et import Excel des données géographiques"""
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .models import ExProvince, Region, District, Commune, Fokontany, Village, Parcelle
from apps.billing.models import PointEau


def export_to_excel(model_name):
    """Exporte les données géographiques vers Excel avec formatage"""

    # Configuration des modèles
    model_config = {
        'regions': {
            'model': Region,
            'fields': ['id', 'nom_region', 'ex_province_id', 'ex_province__nom_ex_province'],
            'headers': ['ID Région (UUID)', 'Nom de la Région'],
            'headers': ['ID Region (UUID)', 'Nom de la Region', 'ID Ex-province (UUID)', 'Ex-province'],
            'filename': 'regions.xlsx',
            'sheet_name': 'Régions'
        },
        'districts': {
            'model': District,
            'fields': ['id', 'nom_district', 'region_id', 'region__nom_region'],
            'headers': ['ID District (UUID)', 'Nom du District', 'ID Région (UUID)', 'Région'],
            'filename': 'districts.xlsx',
            'sheet_name': 'Districts'
        },
        'communes': {
            'model': Commune,
            'fields': ['id', 'nom_commune', 'district_id', 'district__nom_district', 'district__region__nom_region'],
            'headers': ['ID Commune (UUID)', 'Nom de la Commune', 'ID District (UUID)', 'District', 'Région'],
            'filename': 'communes.xlsx',
            'sheet_name': 'Communes'
        },
        'fokontanys': {
            'model': Fokontany,
            'fields': ['id', 'nom_fokontany', 'commune_id', 'commune__nom_commune'],
            'headers': ['ID Fokontany (UUID)', 'Nom du Fokontany', 'ID Commune (UUID)', 'Nom Commune (référence)'],
            'filename': 'fokontanys.xlsx',
            'sheet_name': 'Fokontanys'
        },
        'villages': {
            'model': Village,
            'fields': ['id', 'nom_village', 'fokontany_id', 'fokontany__nom_fokontany', 'fokontany__commune__nom_commune'],
            'headers': ['ID Village (UUID)', 'Nom du Village', 'ID Fokontany (UUID)', 'Nom Fokontany (référence)', 'Commune'],
            'filename': 'villages.xlsx',
            'sheet_name': 'Villages'
        },
        'parcelles': {
            'model': Parcelle,
            'fields': ['id', 'reference_parcelle', 'village_id', 'village__nom_village', 'village__fokontany__commune__nom_commune', 'nombre_habitants', 'nombre_menages', 'source_donnees'],
            'headers': ['ID Parcelle (UUID)', 'Référence Parcelle', 'ID Village (UUID)', 'Nom Village (référence)', 'Commune', 'Nombre Habitants', 'Nombre Ménages', 'Source Données'],
            'filename': 'parcelles.xlsx',
            'sheet_name': 'Parcelles'
        },
        'points_eau': {
            'model': PointEau,
            'fields': ['id', 'nomabonne', 'parcelle_id', 'parcelle__reference_parcelle', 'ouvrage_id', 'ouvrage__nom', 'typeabonne', 'numeroteleph', 'contactabonne', 'latitude', 'longitude'],
            'headers': ['ID Point Eau (UUID)', 'Nom Abonné', 'ID Parcelle (UUID)', 'Référence Parcelle', 'ID Ouvrage (UUID)', 'Nom Ouvrage', 'Type Abonné', 'Numéro Tél', 'Contact', 'Latitude', 'Longitude'],
            'filename': 'points_eau.xlsx',
            'sheet_name': 'Points d\'Eau'
        }
    }

    if model_name not in model_config:
        raise ValueError(f"Modèle {model_name} non supporté")

    config = model_config[model_name]

    # Créer le workbook
    wb = Workbook()
    ws = wb.active
    ws.title = config['sheet_name']

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-têtes
    for col_num, header in enumerate(config['headers'], 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Données
    queryset = config['model'].objects.all()

    # Pour les modèles avec relations, utiliser select_related
    if 'ex_province__nom_ex_province' in config['fields']:
        queryset = queryset.select_related('ex_province')
    if 'region__nom_region' in config['fields']:
        queryset = queryset.select_related('region')
    if 'district__nom_district' in config['fields'] or 'district__region__nom_region' in config['fields']:
        queryset = queryset.select_related('district__region')
    if 'commune__nom_commune' in config['fields']:
        queryset = queryset.select_related('commune__district')
    if 'fokontany__nom_fokontany' in config['fields'] or 'fokontany__commune__nom_commune' in config['fields']:
        queryset = queryset.select_related('fokontany__commune')
    if 'village__nom_village' in config['fields']:
        queryset = queryset.select_related('village__fokontany')
    if 'village__fokontany__commune__nom_commune' in config['fields']:
        queryset = queryset.select_related('village__fokontany__commune')
    if 'parcelle__reference_parcelle' in config['fields'] or 'ouvrage__nom' in config['fields']:
        queryset = queryset.select_related('parcelle', 'ouvrage')

    for row_num, obj in enumerate(queryset, 2):
        for col_num, field in enumerate(config['fields'], 1):
            # Utiliser getattr avec navigation dans les relations
            value = obj
            for attr in field.split('__'):
                value = getattr(value, attr, None)
                if value is None:
                    value = ''
                    break

            # Si c'est un objet modèle, prendre son attribut de nom spécifique
            if hasattr(value, 'nom_region'):
                value = value.nom_region
            elif hasattr(value, 'nom_ex_province'):
                value = value.nom_ex_province
            elif hasattr(value, 'nom_district'):
                value = value.nom_district
            elif hasattr(value, 'nom_commune'):
                value = value.nom_commune
            elif hasattr(value, 'nom_fokontany'):
                value = value.nom_fokontany
            elif hasattr(value, 'nom_village'):
                value = value.nom_village

            cell = ws.cell(row=row_num, column=col_num, value=str(value) if value else '')
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Ajuster la largeur des colonnes
    for col_num in range(1, len(config['headers']) + 1):
        ws.column_dimensions[get_column_letter(col_num)].width = 25

    # Créer la réponse HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{config["filename"]}"'
    response['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'

    wb.save(response)
    return response


def create_excel_template(model_name):
    """Crée un template Excel vide pour l'import"""

    templates_config = {
        'regions': {
            'headers': ['Nom de la Région'],
            'headers': ['Nom de la Region', 'Ex-province (exactement comme dans la base, optionnel)'],
            'filename': 'template_regions.xlsx',
            'sheet_name': 'Régions'
        },
        'districts': {
            'headers': ['Nom du District', 'Région (exactement comme dans la base)'],
            'filename': 'template_districts.xlsx',
            'sheet_name': 'Districts'
        },
        'communes': {
            'headers': ['Nom de la Commune', 'District (exactement comme dans la base)'],
            'filename': 'template_communes.xlsx',
            'sheet_name': 'Communes'
        },
        'fokontanys': {
            'headers': ['Nom du Fokontany', 'ID Commune (UUID)'],
            'filename': 'template_fokontanys.xlsx',
            'sheet_name': 'Fokontanys'
        },
        'villages': {
            'headers': ['Nom du Village', 'ID Fokontany (UUID)'],
            'filename': 'template_villages.xlsx',
            'sheet_name': 'Villages'
        },
        'parcelles': {
            'headers': ['Référence Parcelle', 'ID Village (UUID)', 'Nombre Habitants (optionnel)', 'Nombre Ménages (optionnel)', 'Source Données (optionnel)'],
            'filename': 'template_parcelles.xlsx',
            'sheet_name': 'Parcelles'
        },
        'points_eau': {
            'headers': ['Nom Abonné *', 'ID Parcelle (UUID) *', 'ID Ouvrage (UUID) *', 'Type Abonné * (particulier/commercial/institution)', 'Numéro Tél (optionnel)', 'Contact (optionnel)', 'Latitude (optionnel)', 'Longitude (optionnel)'],
            'filename': 'template_points_eau.xlsx',
            'sheet_name': 'Points d\'Eau'
        }
    }

    if model_name not in templates_config:
        raise ValueError(f"Template {model_name} non supporté")

    config = templates_config[model_name]

    # Créer le workbook
    wb = Workbook()
    ws = wb.active
    ws.title = config['sheet_name']

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # En-têtes
    for col_num, header in enumerate(config['headers'], 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        ws.column_dimensions[get_column_letter(col_num)].width = 35

    # Créer la réponse HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{config["filename"]}"'

    wb.save(response)
    return response
