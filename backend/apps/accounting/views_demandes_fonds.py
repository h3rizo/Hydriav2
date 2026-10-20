"""
Vues pour la gestion des demandes d'envoi de fonds.

Workflow:
1. Liste des budgets
2. Création/modification budget
3. Liste des demandes de fonds
4. Créer une demande (calculs automatiques)
5. Compléter les explications
6. Valider et envoyer
"""
from django.shortcuts import render, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
from django.db.models import Sum, Q
from django.utils import timezone
from decimal import Decimal
from datetime import datetime, date
import json
import re
import unicodedata
from collections import defaultdict

from openpyxl import load_workbook, Workbook

from .models_demandes_fonds import Budget, LigneBudget, DemandeFonds, LigneDemandeFonds
from .models import BaseDiameX, SoldesReels
from apps.billing.models_previsions import PrevisionBudget
from apps.employees.models import Employe, Contrat, GrilleSalaire, FraisMedicaux, FraisScolarite, FamilleEmploye
from apps.employees.views_verification_paiements import extraire_info_description, detecter_type_paiement, normaliser_texte
from apps.planning.models import PlanningDetail
from apps.inventory.models import Purchase


# ========== BUDGETS ==========

MONTH_NAME_TO_NUM = {
    'jan': 1, 'janvier': 1,
    'fev': 2, 'fevr': 2, 'fevrier': 2, 'feb': 2, 'february': 2,
    'mar': 3, 'mars': 3,
    'avr': 4, 'avril': 4, 'apr': 4, 'april': 4,
    'mai': 5, 'may': 5,
    'jun': 6, 'juin': 6, 'june': 6,
    'jul': 7, 'juil': 7, 'juillet': 7, 'july': 7,
    'aou': 8, 'aout': 8, 'ao': 8, 'aug': 8, 'august': 8,
    'sep': 9, 'sept': 9, 'septembre': 9, 'september': 9,
    'oct': 10, 'octobre': 10, 'october': 10,
    'nov': 11, 'novembre': 11, 'november': 11,
    'dec': 12, 'decembre': 12, 'december': 12,
}

HEADER_ALIASES = {
    'compte': {
        'compte', 'code compte', 'numero compte', 'num compte',
        'n compte', 'ncompte', 'compte comptable'
    },
    'libelle': {
        'libelle', 'intitule', 'designation', 'description', 'nom compte'
    },
    'mois': {
        'mois', 'month', 'periode', 'period'
    },
    'montant': {
        'montant', 'valeur', 'somme', 'budget', 'montant ar', 'montant ariary'
    },
}

SALARY_BUDGET_CODES = {
    'salaire': ('SALAIRE', 'Salaires'),
    'irsa': ('IRSA', 'IRSA'),
    'cnaps': ('CNAPS', 'CNaPS'),
    'fmfp': ('FMFP', 'FMFP'),
    'sante': ('SANTE', 'Santé'),
}

EMPLOYEE_BENEFITS_BUDGET_CODES = {
    'frais_medicaux': ('FRAIS_MEDICAUX', 'Frais médicaux'),
    'frais_scolarite': ('FRAIS_SCOLARITE', 'Frais de scolarité'),
}

EXCLUDED_OTHER_ACCOUNT_CODES = {
    '641613', '641612', '641611', '641200',
    '631411', '631305',
    '646311', '646312', '646313',
}


def _normalize_excel_text(value):
    if value is None:
        return ''
    text = str(value).strip().lower()
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace('°', '').replace('º', '')
    text = text.replace('-', ' ').replace('_', ' ')
    return ' '.join(text.split())


def _normalize_compte(value):
    if value is None:
        return ''
    if isinstance(value, Decimal):
        if value == value.to_integral():
            return str(int(value))
        return format(value, 'f').rstrip('0').rstrip('.')
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, 'f').rstrip('0').rstrip('.')
    return str(value).strip()


def _parse_decimal(value):
    if value is None or value == '':
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))

    raw = str(value).strip()
    if not raw:
        return None

    raw = raw.replace('\xa0', '').replace(' ', '')
    raw = raw.replace('Ar', '').replace('AR', '').replace('ariary', '')
    raw = raw.replace("'", '')

    if ',' in raw and '.' in raw:
        if raw.rfind(',') > raw.rfind('.'):
            raw = raw.replace('.', '').replace(',', '.')
        else:
            raw = raw.replace(',', '')
    elif ',' in raw:
        raw = raw.replace(',', '.')

    raw = re.sub(r'[^0-9.\-]', '', raw)
    if raw in {'', '-', '.', '-.'}:
        return None

    return Decimal(raw)


def _safe_decimal(value, default='0'):
    try:
        if value in (None, ''):
            return Decimal(str(default))
        return Decimal(str(value))
    except Exception:
        return Decimal(str(default))


def _get_user_organisation_id(user):
    if getattr(user, 'is_superuser', False):
        return None
    try:
        return user.employe.base.organisation_id
    except Exception:
        return None


def _filter_diamex_by_scope(queryset, user):
    if getattr(user, 'is_superuser', False):
        return queryset

    organisation_id = _get_user_organisation_id(user)
    if not organisation_id:
        return queryset.none()

    return queryset.filter(imported_by__base__organisation_id=organisation_id)


def _filter_soldes_by_scope(queryset, user):
    if getattr(user, 'is_superuser', False):
        return queryset

    organisation_id = _get_user_organisation_id(user)
    if not organisation_id:
        return queryset.none()

    allowed_fichiers = list(
        BaseDiameX.objects.filter(imported_by__base__organisation_id=organisation_id)
        .exclude(fichier_source__isnull=True)
        .exclude(fichier_source='')
        .values_list('fichier_source', flat=True)
        .distinct()
    )

    return queryset.filter(
        Q(imported_by__base__organisation_id=organisation_id) |
        Q(fichier_source__in=allowed_fichiers)
    ).distinct()


def _filter_planning_details_by_scope(queryset, user):
    if getattr(user, 'is_superuser', False):
        return queryset

    organisation_id = _get_user_organisation_id(user)
    if not organisation_id:
        return queryset.none()

    return queryset.filter(planning__employe__base__organisation_id=organisation_id)


def _filter_purchases_by_scope(queryset, user):
    if getattr(user, 'is_superuser', False):
        return queryset

    organisation_id = _get_user_organisation_id(user)
    if not organisation_id:
        return queryset.none()

    return queryset.filter(
        Q(base__organisation_id=organisation_id) |
        Q(magasin__base__organisation_id=organisation_id)
    ).distinct()


def _month_bounds(annee, mois):
    start_date = date(int(annee), int(mois), 1)
    if int(mois) == 12:
        end_date = date(int(annee) + 1, 1, 1)
    else:
        end_date = date(int(annee), int(mois) + 1, 1)
    return start_date, end_date


def _get_demande_journal_suggestions(user):
    journaux = list(
        _filter_soldes_by_scope(SoldesReels.objects.all(), user)
        .exclude(journal__isnull=True)
        .exclude(journal='')
        .values_list('journal', flat=True)
        .distinct()
        .order_by('journal')
    )
    if journaux:
        return journaux[:20]

    return list(
        _filter_diamex_by_scope(BaseDiameX.objects.all(), user)
        .exclude(journal_code__isnull=True)
        .exclude(journal_code='')
        .values_list('journal_code', flat=True)
        .distinct()
        .order_by('journal_code')[:20]
    )


def _parse_journal_balances(raw_rows):
    rows = []
    total = Decimal('0')
    for raw in raw_rows or []:
        journal = str((raw or {}).get('journal') or '').strip()
        montant = _parse_decimal((raw or {}).get('montant'))
        if not journal and montant in (None, Decimal('0')):
            continue
        if not journal:
            raise ValueError("Chaque ligne de solde doit avoir un journal.")
        montant = montant or Decimal('0')
        total += montant
        rows.append({
            'journal': journal,
            'montant': montant,
        })
    return rows, total


def _serialize_future_months(future_months_breakdown):
    serialized = []
    for item in future_months_breakdown:
        serialized.append({
            'mois': item['mois'],
            'annee': item['annee'],
            'mois_nom': item['mois_nom'],
            'total': float(item['total']),
        })
    return serialized


def _build_salary_prevision_rows(annee):
    employees_payload = [_serialize_salary_contract(contrat, annee) for contrat in _active_contracts_for_year(annee)]
    preview = _build_salary_budget_preview(annee, employees_payload)

    rows = {}
    for source_key, (compte_code, libelle) in SALARY_BUDGET_CODES.items():
        monthly = {month: Decimal('0') for month in range(1, 13)}
        annual_total = Decimal('0')
        for month in range(1, 13):
            amount = _safe_decimal(preview['monthly_totals'][str(month)][source_key])
            monthly[month] = amount
            annual_total += amount

        rows[compte_code] = {
            'compte': compte_code,
            'libelle_compte': libelle,
            'budget_annuel': annual_total,
            'monthly': monthly,
            'source_types': {'salaires'},
        }
    return rows


def _build_planning_advances_prevision_rows(user, annee):
    details = _filter_planning_details_by_scope(
        PlanningDetail.objects.select_related('planning', 'compte', 'designation'),
        user
    ).filter(
        planning__date__year=annee
    ).filter(
        Q(designation__type_avance_caissiere='avance') |
        Q(designation__type_avance_caissiere__iexact='Avance') |
        Q(designation__type_avance_caissiere__icontains='avance')
    )

    rows = {}
    for detail in details:
        compte_code = _normalize_compte(getattr(getattr(detail, 'compte', None), 'compte', '') or '')
        if not compte_code:
            continue
        month = detail.planning.date.month if detail.planning and detail.planning.date else None
        if not month:
            continue

        row = rows.setdefault(compte_code, {
            'compte': compte_code,
            'libelle_compte': getattr(getattr(detail, 'compte', None), 'nom_compte', '') or 'Avances planning',
            'budget_annuel': Decimal('0'),
            'monthly': {m: Decimal('0') for m in range(1, 13)},
            'source_types': {'avances_planning'},
        })
        montant = _safe_decimal(detail.montant)
        row['monthly'][month] += montant
        row['budget_annuel'] += montant
    return rows


def _build_purchase_prevision_rows(user, annee):
    purchases = _filter_purchases_by_scope(
        Purchase.objects.select_related('compte', 'magasin', 'magasin__base', 'base'),
        user
    ).filter(
        purchase_date__year=annee
    ).exclude(
        purchase_status='annule'
    )

    rows = {}
    for purchase in purchases:
        compte_code = _normalize_compte(getattr(getattr(purchase, 'compte', None), 'compte', '') or '')
        if not compte_code:
            continue
        month = purchase.purchase_date.month if purchase.purchase_date else None
        if not month:
            continue

        row = rows.setdefault(compte_code, {
            'compte': compte_code,
            'libelle_compte': getattr(getattr(purchase, 'compte', None), 'nom_compte', '') or 'Achats',
            'budget_annuel': Decimal('0'),
            'monthly': {m: Decimal('0') for m in range(1, 13)},
            'source_types': {'achats'},
        })
        montant = _safe_decimal(purchase.total_amount)
        row['monthly'][month] += montant
        row['budget_annuel'] += montant
    return rows


def _merge_prevision_rows(*row_maps):
    merged = {}
    for row_map in row_maps:
        for compte_code, row in row_map.items():
            target = merged.setdefault(compte_code, {
                'compte': row['compte'],
                'libelle_compte': row['libelle_compte'],
                'budget_annuel': Decimal('0'),
                'monthly': {m: Decimal('0') for m in range(1, 13)},
                'source_types': set(),
            })
            target['budget_annuel'] += _safe_decimal(row['budget_annuel'])
            target['source_types'].update(row.get('source_types') or set())
            if not target['libelle_compte']:
                target['libelle_compte'] = row['libelle_compte']
            for month, amount in row['monthly'].items():
                target['monthly'][int(month)] += _safe_decimal(amount)
    return merged


def _build_demande_prevision_dataset(user, annee):
    return _merge_prevision_rows(
        _build_salary_prevision_rows(annee),
        _build_planning_advances_prevision_rows(user, annee),
        _build_purchase_prevision_rows(user, annee),
    )


def _build_demande_fonds_projection(user, budget, mois, annee, date_demande, journal_balances):
    if budget.annee != int(annee):
        raise ValueError("Le budget sélectionné doit correspondre à l'année de la demande.")

    current_start, current_end = _month_bounds(annee, mois)
    next_mois = mois + 1
    next_annee = annee
    if next_mois == 13:
        next_mois = 1
        next_annee = annee + 1

    diamex_base_qs = _filter_diamex_by_scope(BaseDiameX.objects.all(), user)
    prevision_rows = _build_demande_prevision_dataset(user, annee)
    if not prevision_rows:
        raise ValueError("Aucune prévision trouvée à partir des tables sources pour cette année.")

    rows = []
    total_current_remaining = Decimal('0')
    total_next_month = Decimal('0')
    total_future_months = Decimal('0')
    total_gross_need = Decimal('0')

    future_months_breakdown = []
    for future_mois in range(mois + 2, 13):
        future_months_breakdown.append({
            'mois': future_mois,
            'annee': annee,
            'mois_nom': date(annee, future_mois, 1).strftime('%B').capitalize(),
            'total': Decimal('0'),
        })

    for compte in sorted(prevision_rows.keys()):
        prevision_row = prevision_rows[compte]
        current_budget = _safe_decimal(prevision_row['monthly'].get(mois))
        next_budget = _safe_decimal(prevision_row['monthly'].get(next_mois)) if next_annee == annee else Decimal('0')

        current_realized = diamex_base_qs.filter(
            compte_code=compte,
            date_m__gte=current_start,
            date_m__lt=current_end,
            date_m__lte=date_demande
        ).aggregate(total=Sum('montant_out'))['total'] or Decimal('0')
        current_realized = _safe_decimal(current_realized)

        current_remaining = current_budget - current_realized
        if current_remaining < 0:
            current_remaining = Decimal('0')

        future_after_next = Decimal('0')
        for future_item in future_months_breakdown:
            future_amount = _safe_decimal(prevision_row['monthly'].get(future_item['mois']))
            future_item['total'] += future_amount
            future_after_next += future_amount

        gross_need = current_remaining + next_budget
        annual_total = _safe_decimal(prevision_row['budget_annuel'])
        projected_total = current_realized + current_remaining + next_budget + future_after_next
        depassement = projected_total - annual_total
        if depassement < 0:
            depassement = Decimal('0')

        budget_restant = annual_total - projected_total

        rows.append({
            'compte': compte,
            'libelle_compte': prevision_row['libelle_compte'],
            'budget_annuel': annual_total,
            'depenses_mois_a_date': current_realized,
            'reste_mois_courant': current_remaining,
            'depenses_mois_suivant': next_budget,
            'previsions_mois_suivants': future_after_next,
            'besoin_brut': gross_need,
            'solde_impute': Decimal('0'),
            'montant_demande': gross_need,
            'budget_restant': budget_restant,
            'depassement': depassement,
            'sources': sorted(prevision_row['source_types']),
        })

        total_current_remaining += current_remaining
        total_next_month += next_budget
        total_future_months += future_after_next
        total_gross_need += gross_need

    total_balance = sum((item['montant'] for item in journal_balances), Decimal('0'))
    total_requested = total_gross_need - total_balance
    if total_requested < 0:
        total_requested = Decimal('0')

    if total_gross_need > 0 and total_balance > 0:
        balance_to_allocate = min(total_balance, total_gross_need)
        allocated_total = Decimal('0')
        rows_with_need = [row for row in rows if row['besoin_brut'] > 0]

        for index, row in enumerate(rows_with_need):
            if index == len(rows_with_need) - 1:
                allocation = balance_to_allocate - allocated_total
            else:
                allocation = (balance_to_allocate * row['besoin_brut'] / total_gross_need).quantize(Decimal('0.01'))
                allocated_total += allocation
            row['solde_impute'] = allocation
            row['montant_demande'] = row['besoin_brut'] - allocation
            if row['montant_demande'] < 0:
                row['montant_demande'] = Decimal('0')

    return {
        'summary': {
            'total_journal_balances': total_balance,
            'total_current_remaining': total_current_remaining,
            'total_next_month': total_next_month,
            'total_future_months': total_future_months,
            'total_gross_need': total_gross_need,
            'total_requested': total_requested,
            'date_demande': date_demande.isoformat(),
            'next_month_label': date(next_annee, next_mois, 1).strftime('%B %Y').capitalize(),
            'prevision_mode': 'tables_sources',
        },
        'rows': rows,
        'future_months_breakdown': future_months_breakdown,
    }


def _demande_fonds_commentaire_payload(preview, journal_balances):
    summary = preview['summary']
    return json.dumps({
        'workflow_version': 2,
        'date_demande': summary['date_demande'],
        'journal_balances': [
            {
                'journal': row['journal'],
                'montant': float(row['montant']),
            }
            for row in journal_balances
        ],
        'current_month_remaining': float(summary['total_current_remaining']),
        'next_month_expenses': float(summary['total_next_month']),
        'future_months_total': float(summary['total_future_months']),
        'gross_need': float(summary['total_gross_need']),
        'next_month_label': summary['next_month_label'],
        'future_months_breakdown': _serialize_future_months(preview['future_months_breakdown']),
    }, ensure_ascii=False)


def _parse_demande_workflow_commentaire(commentaire):
    if not commentaire:
        return {}
    try:
        return json.loads(commentaire)
    except Exception:
        return {}


def _active_contracts_for_year(annee):
    start_year = date(int(annee), 1, 1)
    end_year = date(int(annee), 12, 31)
    return Contrat.objects.select_related(
        'employe',
        'employe__poste',
        'employe__base',
        'employe__base__organisation',
        'poste',
    ).filter(
        employe__date_debauchage__isnull=True,
        debut_contrat__lte=end_year
    ).filter(
        Q(fin_contrat__isnull=True) | Q(fin_contrat__gte=start_year)
    )


def _active_employees_for_year(annee):
    contracts = _active_contracts_for_year(annee)
    employee_ids = contracts.values_list('employe_id', flat=True).distinct()
    return Employe.objects.select_related(
        'poste',
        'base',
        'base__organisation',
    ).filter(
        id__in=employee_ids,
        date_debauchage__isnull=True,
    ).order_by('nom', 'prenom')


def _serialize_salary_contract(contrat, annee):
    employe = contrat.employe
    start_month = 1
    end_month = 12

    if contrat.debut_contrat and contrat.debut_contrat.year == int(annee):
        start_month = contrat.debut_contrat.month
    if contrat.fin_contrat and contrat.fin_contrat.year == int(annee):
        end_month = contrat.fin_contrat.month
    if employe.date_debauchage and employe.date_debauchage.year == int(annee):
        end_month = min(end_month, employe.date_debauchage.month)

    return {
        'source': 'active_contract',
        'employe_id': str(employe.id),
        'nom': employe.nom_affichage,
        'poste': getattr(contrat.poste, 'nom_poste', '') or getattr(employe.poste, 'nom_poste', '') or '',
        'base': getattr(getattr(employe, 'base', None), 'nom_base', '') or '',
        'type_contrat': contrat.type_contrat or 'temps_plein',
        'salaire_brut_base': float(contrat.salaire_brut_base or 0),
        'personnes_charge': int(getattr(employe, 'personnes_charge', 0) or 0),
        'mutuelle_sante': getattr(employe, 'mutuelle_sante', '') or '',
        'start_month': max(1, min(12, start_month)),
        'end_month': max(1, min(12, end_month)),
    }


def _build_salary_preview_row(item, annee):
    start_month = max(1, min(12, int(item.get('start_month') or 1)))
    end_month = max(1, min(12, int(item.get('end_month') or 12)))
    if end_month < start_month:
        raise ValueError("La période de l'employé est invalide")

    salaire_brut_base = _safe_decimal(item.get('salaire_brut_base'))
    personnes_charge = int(item.get('personnes_charge') or 0)
    type_contrat = (item.get('type_contrat') or 'temps_plein').strip() or 'temps_plein'

    dummy_contract = Contrat(
        type_contrat=type_contrat,
        salaire_brut_base=salaire_brut_base,
    )
    dummy_contract.employe = Employe(personnes_charge=personnes_charge)

    monthly = {}
    totals = {key: Decimal('0') for key in SALARY_BUDGET_CODES.keys()}

    for month in range(1, 13):
        if month < start_month or month > end_month:
            monthly[str(month)] = {key: 0.0 for key in totals.keys()}
            continue

        grille = GrilleSalaire(
            contrat=dummy_contract,
            date_salaire=date(int(annee), month, 1),
            personnes_charge=personnes_charge,
            salaire_brut_base=salaire_brut_base,
            nb_jours_presence=30 if type_contrat == 'journalier' else 0,
            est_prorata=False,
            nb_jours_feries_gardien=0,
            nb_heures_remplacement_gardien=0,
        )
        grille.calculer_tous_les_champs()

        month_values = {
            'salaire': grille.salaire_net or Decimal('0'),
            'irsa': grille.irsa or Decimal('0'),
            'cnaps': (grille.charge_cnaps_employe or Decimal('0')) * Decimal('14'),
            'fmfp': grille.charge_cnaps_employe or Decimal('0'),
            'sante': grille.mutuelle_sante or Decimal('0'),
        }

        for key, value in month_values.items():
            totals[key] += value
        monthly[str(month)] = {key: float(value) for key, value in month_values.items()}

    return {
        'nom': item.get('nom') or 'Sans nom',
        'poste': item.get('poste') or '',
        'base': item.get('base') or '',
        'type_contrat': type_contrat,
        'salaire_brut_base': float(salaire_brut_base),
        'personnes_charge': personnes_charge,
        'mutuelle_sante': item.get('mutuelle_sante') or '',
        'start_month': start_month,
        'end_month': end_month,
        'monthly': monthly,
        'totals': {key: float(value) for key, value in totals.items()},
    }


def _build_salary_budget_preview(annee, employees_payload):
    rows = []
    monthly_totals = {
        str(month): {key: Decimal('0') for key in SALARY_BUDGET_CODES.keys()}
        for month in range(1, 13)
    }
    annual_totals = {key: Decimal('0') for key in SALARY_BUDGET_CODES.keys()}

    for item in employees_payload:
        row = _build_salary_preview_row(item, annee)
        rows.append(row)
        for month in range(1, 13):
            for key in annual_totals.keys():
                monthly_totals[str(month)][key] += _safe_decimal(row['monthly'][str(month)][key])
        for key in annual_totals.keys():
            annual_totals[key] += _safe_decimal(row['totals'][key])

    return {
        'rows': rows,
        'monthly_totals': {
            month: {key: float(value) for key, value in values.items()}
            for month, values in monthly_totals.items()
        },
        'annual_totals': {key: float(value) for key, value in annual_totals.items()},
    }


def _serialize_employee_benefits_row(employe, reference_year, medical_total, school_total):
    return {
        'employe_id': str(employe.id),
        'nom': employe.nom_affichage,
        'poste': getattr(getattr(employe, 'poste', None), 'nom_poste', '') or '',
        'base': getattr(getattr(employe, 'base', None), 'nom_base', '') or '',
        'reference_year': int(reference_year),
        'frais_medicaux_annuel': float(medical_total),
        'frais_scolarite_annuel': float(school_total),
    }


def _build_employee_benefits_preview(annee, employees_payload):
    rows = []
    monthly_totals = {
        str(month): {key: Decimal('0') for key in EMPLOYEE_BENEFITS_BUDGET_CODES.keys()}
        for month in range(1, 13)
    }
    annual_totals = {key: Decimal('0') for key in EMPLOYEE_BENEFITS_BUDGET_CODES.keys()}

    for item in employees_payload:
        medical_annual = _safe_decimal(item.get('frais_medicaux_annuel'))
        school_annual = _safe_decimal(item.get('frais_scolarite_annuel'))
        monthly_medical = medical_annual / Decimal('12')
        monthly_school = school_annual / Decimal('12')

        row = {
            'employe_id': item.get('employe_id'),
            'nom': item.get('nom') or 'Sans nom',
            'poste': item.get('poste') or '',
            'base': item.get('base') or '',
            'frais_medicaux_annuel': float(medical_annual),
            'frais_scolarite_annuel': float(school_annual),
            'frais_medicaux_mensuel': float(monthly_medical),
            'frais_scolarite_mensuel': float(monthly_school),
        }
        rows.append(row)

        for month in range(1, 13):
            monthly_totals[str(month)]['frais_medicaux'] += monthly_medical
            monthly_totals[str(month)]['frais_scolarite'] += monthly_school
        annual_totals['frais_medicaux'] += medical_annual
        annual_totals['frais_scolarite'] += school_annual

    return {
        'rows': rows,
        'monthly_totals': {
            month: {key: float(value) for key, value in values.items()}
            for month, values in monthly_totals.items()
        },
        'annual_totals': {key: float(value) for key, value in annual_totals.items()},
    }


def _load_other_accounts_reference_rows(annee):
    reference_year = int(annee) - 1
    entries = BaseDiameX.objects.select_related('compte').filter(
        date_m__year=reference_year,
    ).filter(
        Q(compte_code__startswith='2') | Q(compte_code__startswith='6') | Q(compte_code__startswith='7')
    ).exclude(
        compte_code__in=EXCLUDED_OTHER_ACCOUNT_CODES
    ).order_by('compte_code')

    totals_by_account = {}
    for entry in entries:
        compte_code = (entry.compte_code or '').strip()
        if not compte_code:
            continue
        montant_in = _safe_decimal(entry.montant_in)
        montant_out = _safe_decimal(entry.montant_out)
        amount = montant_out if montant_out > 0 else montant_in
        if amount <= 0:
            continue
        row = totals_by_account.setdefault(compte_code, {
            'compte': compte_code,
            'libelle_compte': getattr(getattr(entry, 'compte', None), 'nom_compte', '') or entry.get_categorie_parente() or compte_code,
            'montant_annuel': Decimal('0'),
        })
        row['montant_annuel'] += amount

    return reference_year, [
        {
            'compte': row['compte'],
            'libelle_compte': row['libelle_compte'],
            'montant_annuel': float(row['montant_annuel']),
        }
        for row in totals_by_account.values()
    ]


def _build_other_accounts_preview(rows_payload):
    rows = []
    monthly_totals = {str(month): Decimal('0') for month in range(1, 13)}
    annual_total = Decimal('0')

    for item in rows_payload:
        compte = (item.get('compte') or '').strip()
        if not compte:
            continue
        annual_amount = _safe_decimal(item.get('montant_annuel'))
        monthly_amount = annual_amount / Decimal('12')
        rows.append({
            'compte': compte,
            'libelle_compte': item.get('libelle_compte') or compte,
            'montant_annuel': float(annual_amount),
            'montant_mensuel': float(monthly_amount),
        })
        for month in range(1, 13):
            monthly_totals[str(month)] += monthly_amount
        annual_total += annual_amount

    return {
        'rows': rows,
        'monthly_totals': {month: float(value) for month, value in monthly_totals.items()},
        'annual_total': float(annual_total),
    }


def _default_salary_employees_payload(annee):
    contracts = _active_contracts_for_year(annee)
    employees = [_serialize_salary_contract(contrat, annee) for contrat in contracts]
    employees.sort(key=lambda item: ((item.get('nom') or '').lower(), (item.get('poste') or '').lower()))
    return employees


def _default_employee_benefits_payload(annee):
    reference_year = annee - 1
    employees = _active_employees_for_year(annee)
    employees_list = list(employees)
    medical_totals = {employe.id: Decimal('0') for employe in employees_list}
    school_totals = {employe.id: Decimal('0') for employe in employees_list}
    employee_ids = [employe.id for employe in employees_list]
    family_members = list(
        FamilleEmploye.objects.filter(
            employe_id__in=employee_ids,
            actif=True,
        ).select_related('employe')
    )

    def _match_employee_from_description(nom_employe, raw_description='', raw_detail=''):
        description_norm = normaliser_texte(raw_description or '')
        detail_norm = normaliser_texte(raw_detail or '')
        name = (nom_employe or '').strip()
        owner_segment = ''

        if raw_description:
            cleaned = raw_description.lstrip('_')
            parts = cleaned.split('_')
            if parts:
                owner_segment = parts[0].strip()
                owner_segment = re.sub(r'^[A-Z]{2,6}\s+', '', owner_segment).strip()

        def _token_score(employee, candidate_text):
            candidate_norm = normaliser_texte(candidate_text or '')
            if not candidate_norm:
                return 0
            employee_tokens = [
                normaliser_texte(employee.surnom or ''),
                normaliser_texte(employee.nom or ''),
                normaliser_texte(employee.prenom or ''),
                normaliser_texte(employee.nom_complet or ''),
                normaliser_texte(employee.nom_affichage or ''),
            ]
            score = 0
            for token in employee_tokens:
                if token and len(token) > 2 and token in candidate_norm:
                    score = max(score, len(token))
            return score

        if detail_norm:
            for employe in employees_list:
                if detail_norm == normaliser_texte(employe.surnom or ''):
                    return employe

            best_detail_match = None
            best_detail_score = 0
            for employe in employees_list:
                score = _token_score(employe, detail_norm)
                if score > best_detail_score:
                    best_detail_score = score
                    best_detail_match = employe
            if best_detail_match:
                return best_detail_match

        if name:
            for employe in employees_list:
                if (employe.surnom or '').strip().lower() == name.lower():
                    return employe

            for employe in employees_list:
                if (
                    name.lower() in (employe.surnom or '').lower() or
                    name.lower() in (employe.nom or '').lower() or
                    name.lower() in (employe.prenom or '').lower()
                ):
                    return employe

            parts = [part for part in name.split() if len(part) > 3]
            for part in parts:
                for employe in employees_list:
                    if (
                        part.lower() in (employe.surnom or '').lower() or
                        part.lower() in (employe.nom or '').lower() or
                        part.lower() in (employe.prenom or '').lower()
                    ):
                        return employe

        if owner_segment:
            best_owner_match = None
            best_owner_score = 0
            for employe in employees_list:
                score = _token_score(employe, owner_segment)
                if score > best_owner_score:
                    best_owner_score = score
                    best_owner_match = employe
            if best_owner_match:
                return best_owner_match

        if description_norm:
            best_match = None
            best_score = 0
            for employe in employees_list:
                score = _token_score(employe, description_norm)
                if score > best_score:
                    best_score = score
                    best_match = employe
            if best_match:
                return best_match

            best_family_match = None
            best_family_score = 0
            for membre in family_members:
                member_name = (membre.nom or '').strip()
                member_name_norm = normaliser_texte(member_name)
                if member_name_norm and len(member_name_norm) > 2 and member_name_norm in description_norm:
                    score = len(member_name_norm)
                    if score > best_family_score:
                        best_family_score = score
                        best_family_match = membre.employe
            if best_family_match:
                return best_family_match
        return None

    diamex_entries = BaseDiameX.objects.filter(
        Q(date_m__year=reference_year) | Q(date_rec__year=reference_year),
    ).exclude(
        description__isnull=True,
    ).exclude(
        description='',
    ).order_by('date_m', 'id')

    found_diamex_employee_benefits = False
    for entry in diamex_entries:
        description = entry.description or ''
        info = extraire_info_description(description)
        type_paiement = info.get('type_paiement') or detecter_type_paiement(description)
        employe = _match_employee_from_description(info.get('nom_employe'), description, entry.detail or '')
        montant = _safe_decimal(entry.montant_out if (entry.montant_out or 0) > 0 else entry.montant_in)
        if montant <= 0:
            continue

        if type_paiement == 'frais_medicaux':
            if not employe:
                continue
            medical_totals[employe.id] += montant
            found_diamex_employee_benefits = True
        elif type_paiement == 'frais_scolarite':
            if not employe:
                continue
            school_totals[employe.id] += montant
            found_diamex_employee_benefits = True

    if not found_diamex_employee_benefits:
        employee_ids = [employe.id for employe in employees_list]

        fallback_medical_totals = {
            row['employe_id']: row['total'] or Decimal('0')
            for row in FraisMedicaux.objects.filter(
                statut__in=['imprime', 'exporte'],
                date__year=reference_year,
            ).filter(
                Q(membre_famille__employe_id__in=employee_ids) |
                Q(employe_beneficiaire_id__in=employee_ids)
            ).values('membre_famille__employe_id', 'employe_beneficiaire_id').annotate(
                total=Sum('montant_rembourse')
            ).iterator()
            for row in [{
                'employe_id': row['membre_famille__employe_id'] or row['employe_beneficiaire_id'],
                'total': row['total'],
            }]
        }

        fallback_school_totals = {
            row['membre_famille__employe_id']: row['total'] or Decimal('0')
            for row in FraisScolarite.objects.filter(
                statut__in=['imprime', 'exporte'],
                date__year=reference_year,
                membre_famille__employe_id__in=employee_ids,
            ).values('membre_famille__employe_id').annotate(
                total=Sum('montant_rembourse')
            )
        }

        for employe in employees_list:
            medical_totals[employe.id] = fallback_medical_totals.get(employe.id, Decimal('0'))
            school_totals[employe.id] = fallback_school_totals.get(employe.id, Decimal('0'))

    return [
        _serialize_employee_benefits_row(
            employe,
            reference_year,
            medical_totals.get(employe.id, Decimal('0')),
            school_totals.get(employe.id, Decimal('0')),
        )
        for employe in employees
    ]


def _build_full_budget_preview(annee):
    salary_preview = _build_salary_budget_preview(annee, _default_salary_employees_payload(annee))
    employee_benefits_preview = _build_employee_benefits_preview(annee, _default_employee_benefits_payload(annee))
    other_reference_year, other_rows = _load_other_accounts_reference_rows(annee)
    other_accounts_preview = _build_other_accounts_preview(other_rows)

    total_general = (
        sum(_safe_decimal(value) for value in salary_preview['annual_totals'].values()) +
        sum(_safe_decimal(value) for value in employee_benefits_preview['annual_totals'].values()) +
        _safe_decimal(other_accounts_preview['annual_total'])
    )

    return {
        'salary': salary_preview,
        'employee_benefits': employee_benefits_preview,
        'other_accounts': other_accounts_preview,
        'other_reference_year': other_reference_year,
        'total_general': float(total_general),
    }


def _parse_month(value):
    if value is None or value == '':
        return None

    if isinstance(value, (datetime, date)):
        return value.month

    if isinstance(value, (int, float, Decimal)):
        month_num = int(value)
        return month_num if 1 <= month_num <= 12 else None

    text = _normalize_excel_text(value)
    if not text:
        return None

    if text.isdigit():
        month_num = int(text)
        return month_num if 1 <= month_num <= 12 else None

    # Formats de type 2025-11, 11/2025, 2025/11
    for token in re.split(r'[\s/._-]+', text):
        if token.isdigit():
            num = int(token)
            if 1 <= num <= 12:
                return num
        if token in MONTH_NAME_TO_NUM:
            return MONTH_NAME_TO_NUM[token]

    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%Y', '%Y/%m', '%m-%Y', '%Y-%m'):
        try:
            parsed = datetime.strptime(str(value).strip(), fmt)
            return parsed.month
        except Exception:
            continue

    return None


def _detect_import_columns(header_values):
    columns = {}

    for idx, header in enumerate(header_values):
        normalized = _normalize_excel_text(header)
        if not normalized:
            continue

        for target, aliases in HEADER_ALIASES.items():
            if target in columns:
                continue
            if normalized in aliases:
                columns[target] = idx
                continue
            if any(alias in normalized for alias in aliases):
                columns[target] = idx

    return columns

@login_required
def budgets_list(request):
    """Page de liste des budgets."""
    return render(request, 'accounting/budgets_list.html')


@login_required
@require_http_methods(["GET"])
def api_budgets_list(request):
    """API pour récupérer la liste des budgets."""
    try:
        annee = request.GET.get('annee')

        queryset = Budget.objects.all()

        if annee:
            queryset = queryset.filter(annee=annee)

        budgets = []
        for budget in queryset:
            nb_lignes = budget.lignes.count()
            total_prevu = budget.lignes.aggregate(total=Sum('montant_prevu'))['total'] or 0

            budgets.append({
                'id': str(budget.id),
                'annee': budget.annee,
                'nom': budget.nom,
                'description': budget.description or '',
                'valide': budget.valide,
                'date_creation': budget.date_creation.strftime('%Y-%m-%d %H:%M'),
                'date_validation': budget.date_validation.strftime('%Y-%m-%d %H:%M') if budget.date_validation else '',
                'nb_lignes': nb_lignes,
                'total_prevu': float(total_prevu),
            })

        return JsonResponse({
            'success': True,
            'budgets': budgets,
            'count': len(budgets)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_create(request):
    """API pour créer un nouveau budget."""
    try:
        data = json.loads(request.body)

        annee = data.get('annee')
        nom = data.get('nom')
        description = data.get('description', '')

        if not annee or not nom:
            return JsonResponse({
                'success': False,
                'error': 'Année et nom requis'
            }, status=400)

        budget = Budget.objects.create(
            annee=annee,
            nom=nom,
            description=description,
            cree_par=request.user
        )

        return JsonResponse({
            'success': True,
            'message': 'Budget créé avec succès',
            'budget_id': str(budget.id)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@require_http_methods(["GET"])
def api_budget_salary_employees(request):
    try:
        annee = int(request.GET.get('annee') or timezone.now().year)
        contracts = _active_contracts_for_year(annee)
        employees = [_serialize_salary_contract(contrat, annee) for contrat in contracts]
        employees.sort(key=lambda item: ((item.get('nom') or '').lower(), (item.get('poste') or '').lower()))
        return JsonResponse({
            'success': True,
            'annee': annee,
            'employees': employees,
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_salary_preview(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or timezone.now().year)
        employees = data.get('employees') or []
        preview = _build_salary_budget_preview(annee, employees)
        return JsonResponse({'success': True, 'preview': preview})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_create_salary(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or 0)
        nom = (data.get('nom') or '').strip()
        description = (data.get('description') or '').strip()
        employees = data.get('employees') or []

        if not annee or not nom:
            return JsonResponse({'success': False, 'error': 'Année et nom requis'}, status=400)
        if not employees:
            return JsonResponse({'success': False, 'error': 'Aucun employé à budgéter'}, status=400)

        preview = _build_salary_budget_preview(annee, employees)

        with transaction.atomic():
            budget = Budget.objects.create(
                annee=annee,
                nom=nom,
                description=description or 'Budget généré depuis le workflow charges salariales',
                cree_par=request.user,
            )

            for key, (compte, libelle) in SALARY_BUDGET_CODES.items():
                monthly = preview['monthly_totals']
                ligne = LigneBudget.objects.create(
                    budget=budget,
                    compte=compte,
                    libelle_compte=libelle,
                    montant_prevu=_safe_decimal(preview['annual_totals'][key]),
                    janvier=_safe_decimal(monthly['1'][key]),
                    fevrier=_safe_decimal(monthly['2'][key]),
                    mars=_safe_decimal(monthly['3'][key]),
                    avril=_safe_decimal(monthly['4'][key]),
                    mai=_safe_decimal(monthly['5'][key]),
                    juin=_safe_decimal(monthly['6'][key]),
                    juillet=_safe_decimal(monthly['7'][key]),
                    aout=_safe_decimal(monthly['8'][key]),
                    septembre=_safe_decimal(monthly['9'][key]),
                    octobre=_safe_decimal(monthly['10'][key]),
                    novembre=_safe_decimal(monthly['11'][key]),
                    decembre=_safe_decimal(monthly['12'][key]),
                    notes='Généré automatiquement depuis le workflow charges salariales',
                )
                ligne.save()

        return JsonResponse({
            'success': True,
            'budget_id': str(budget.id),
            'message': 'Budget charges salariales créé avec succès',
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@require_http_methods(["GET"])
def api_budget_employee_benefits_employees(request):
    try:
        annee = int(request.GET.get('annee') or timezone.now().year)
        reference_year = annee - 1
        employees = _active_employees_for_year(annee)

        employee_ids = list(employees.values_list('id', flat=True))

        medical_totals = {
            row['employe_id']: row['total'] or Decimal('0')
            for row in FraisMedicaux.objects.filter(
                statut__in=['imprime', 'exporte'],
                date__year=reference_year,
            ).filter(
                Q(membre_famille__employe_id__in=employee_ids) |
                Q(employe_beneficiaire_id__in=employee_ids)
            ).values('membre_famille__employe_id', 'employe_beneficiaire_id').annotate(
                total=Sum('montant_rembourse')
            ).iterator()
            for row in [{
                'employe_id': row['membre_famille__employe_id'] or row['employe_beneficiaire_id'],
                'total': row['total'],
            }]
        }

        school_totals = {
            row['membre_famille__employe_id']: row['total'] or Decimal('0')
            for row in FraisScolarite.objects.filter(
                statut__in=['imprime', 'exporte'],
                date__year=reference_year,
                membre_famille__employe_id__in=employee_ids,
            ).values('membre_famille__employe_id').annotate(
                total=Sum('montant_rembourse')
            )
        }

        rows = [
            _serialize_employee_benefits_row(
                employe,
                reference_year,
                medical_totals.get(employe.id, Decimal('0')),
                school_totals.get(employe.id, Decimal('0')),
            )
            for employe in employees
        ]

        return JsonResponse({
            'success': True,
            'annee': annee,
            'reference_year': reference_year,
            'employees': rows,
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_employee_benefits_preview(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or timezone.now().year)
        employees = data.get('employees') or []
        preview = _build_employee_benefits_preview(annee, employees)
        return JsonResponse({'success': True, 'preview': preview})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_create_employee_benefits(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or 0)
        nom = (data.get('nom') or '').strip()
        description = (data.get('description') or '').strip()
        employees = data.get('employees') or []

        if not annee or not nom:
            return JsonResponse({'success': False, 'error': 'Année et nom requis'}, status=400)
        if not employees:
            return JsonResponse({'success': False, 'error': 'Aucun employé actif à budgéter'}, status=400)

        preview = _build_employee_benefits_preview(annee, employees)

        with transaction.atomic():
            budget = Budget.objects.create(
                annee=annee,
                nom=nom,
                description=description or 'Budget généré depuis le workflow frais employés',
                cree_par=request.user,
            )

            for key, (compte, libelle) in EMPLOYEE_BENEFITS_BUDGET_CODES.items():
                monthly = preview['monthly_totals']
                ligne = LigneBudget.objects.create(
                    budget=budget,
                    compte=compte,
                    libelle_compte=libelle,
                    montant_prevu=_safe_decimal(preview['annual_totals'][key]),
                    janvier=_safe_decimal(monthly['1'][key]),
                    fevrier=_safe_decimal(monthly['2'][key]),
                    mars=_safe_decimal(monthly['3'][key]),
                    avril=_safe_decimal(monthly['4'][key]),
                    mai=_safe_decimal(monthly['5'][key]),
                    juin=_safe_decimal(monthly['6'][key]),
                    juillet=_safe_decimal(monthly['7'][key]),
                    aout=_safe_decimal(monthly['8'][key]),
                    septembre=_safe_decimal(monthly['9'][key]),
                    octobre=_safe_decimal(monthly['10'][key]),
                    novembre=_safe_decimal(monthly['11'][key]),
                    decembre=_safe_decimal(monthly['12'][key]),
                    notes='Généré depuis la moyenne de l année complète précédente des frais employés',
                )
                ligne.save()

        return JsonResponse({
            'success': True,
            'budget_id': str(budget.id),
            'message': 'Budget frais employés créé avec succès',
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@require_http_methods(["GET"])
def api_budget_other_accounts_rows(request):
    try:
        annee = int(request.GET.get('annee') or timezone.now().year)
        reference_year, rows = _load_other_accounts_reference_rows(annee)
        return JsonResponse({
            'success': True,
            'annee': annee,
            'reference_year': reference_year,
            'rows': rows,
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@require_http_methods(["GET"])
def api_budget_full_workflow_preview(request):
    try:
        annee = int(request.GET.get('annee') or timezone.now().year)
        preview = _build_full_budget_preview(annee)
        return JsonResponse({
            'success': True,
            'annee': annee,
            'preview': preview,
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_create_full_workflow(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or 0)
        nom = (data.get('nom') or '').strip()
        description = (data.get('description') or '').strip()
        salary_employees = data.get('salary_employees')
        employee_benefits_employees = data.get('employee_benefits_employees')
        other_account_rows = data.get('other_account_rows')

        if not annee or not nom:
            return JsonResponse({'success': False, 'error': 'Année et nom requis'}, status=400)

        if salary_employees is None and employee_benefits_employees is None and other_account_rows is None:
            preview = _build_full_budget_preview(annee)
        else:
            salary_payload = salary_employees if salary_employees is not None else _default_salary_employees_payload(annee)
            benefits_payload = employee_benefits_employees if employee_benefits_employees is not None else _default_employee_benefits_payload(annee)
            if other_account_rows is None:
                other_reference_year, other_rows = _load_other_accounts_reference_rows(annee)
            else:
                other_reference_year = annee - 1
                other_rows = other_account_rows

            salary_preview = _build_salary_budget_preview(annee, salary_payload)
            employee_benefits_preview = _build_employee_benefits_preview(annee, benefits_payload)
            other_accounts_preview = _build_other_accounts_preview(other_rows)
            total_general = (
                sum(_safe_decimal(value) for value in salary_preview['annual_totals'].values()) +
                sum(_safe_decimal(value) for value in employee_benefits_preview['annual_totals'].values()) +
                _safe_decimal(other_accounts_preview['annual_total'])
            )
            preview = {
                'salary': salary_preview,
                'employee_benefits': employee_benefits_preview,
                'other_accounts': other_accounts_preview,
                'other_reference_year': other_reference_year,
                'total_general': float(total_general),
            }

        with transaction.atomic():
            budget = Budget.objects.create(
                annee=annee,
                nom=nom,
                description=description or 'Budget généré depuis le workflow combiné',
                cree_par=request.user,
            )

            for key, (compte, libelle) in SALARY_BUDGET_CODES.items():
                monthly = preview['salary']['monthly_totals']
                LigneBudget.objects.create(
                    budget=budget,
                    compte=compte,
                    libelle_compte=libelle,
                    montant_prevu=_safe_decimal(preview['salary']['annual_totals'][key]),
                    janvier=_safe_decimal(monthly['1'][key]),
                    fevrier=_safe_decimal(monthly['2'][key]),
                    mars=_safe_decimal(monthly['3'][key]),
                    avril=_safe_decimal(monthly['4'][key]),
                    mai=_safe_decimal(monthly['5'][key]),
                    juin=_safe_decimal(monthly['6'][key]),
                    juillet=_safe_decimal(monthly['7'][key]),
                    aout=_safe_decimal(monthly['8'][key]),
                    septembre=_safe_decimal(monthly['9'][key]),
                    octobre=_safe_decimal(monthly['10'][key]),
                    novembre=_safe_decimal(monthly['11'][key]),
                    decembre=_safe_decimal(monthly['12'][key]),
                    notes='Workflow combiné - charges salariales',
                )

            for key, (compte, libelle) in EMPLOYEE_BENEFITS_BUDGET_CODES.items():
                monthly = preview['employee_benefits']['monthly_totals']
                LigneBudget.objects.create(
                    budget=budget,
                    compte=compte,
                    libelle_compte=libelle,
                    montant_prevu=_safe_decimal(preview['employee_benefits']['annual_totals'][key]),
                    janvier=_safe_decimal(monthly['1'][key]),
                    fevrier=_safe_decimal(monthly['2'][key]),
                    mars=_safe_decimal(monthly['3'][key]),
                    avril=_safe_decimal(monthly['4'][key]),
                    mai=_safe_decimal(monthly['5'][key]),
                    juin=_safe_decimal(monthly['6'][key]),
                    juillet=_safe_decimal(monthly['7'][key]),
                    aout=_safe_decimal(monthly['8'][key]),
                    septembre=_safe_decimal(monthly['9'][key]),
                    octobre=_safe_decimal(monthly['10'][key]),
                    novembre=_safe_decimal(monthly['11'][key]),
                    decembre=_safe_decimal(monthly['12'][key]),
                    notes='Workflow combiné - frais employés',
                )

            for row in preview['other_accounts']['rows']:
                monthly_amount = _safe_decimal(row['montant_mensuel'])
                LigneBudget.objects.create(
                    budget=budget,
                    compte=row['compte'],
                    libelle_compte=row['libelle_compte'],
                    montant_prevu=_safe_decimal(row['montant_annuel']),
                    janvier=monthly_amount,
                    fevrier=monthly_amount,
                    mars=monthly_amount,
                    avril=monthly_amount,
                    mai=monthly_amount,
                    juin=monthly_amount,
                    juillet=monthly_amount,
                    aout=monthly_amount,
                    septembre=monthly_amount,
                    octobre=monthly_amount,
                    novembre=monthly_amount,
                    decembre=monthly_amount,
                    notes='Workflow combiné - autres comptes',
                )

        return JsonResponse({
            'success': True,
            'budget_id': str(budget.id),
            'message': 'Budget complet créé avec succès',
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_other_accounts_preview(request):
    try:
        data = json.loads(request.body or '{}')
        rows = data.get('rows') or []
        preview = _build_other_accounts_preview(rows)
        return JsonResponse({'success': True, 'preview': preview})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_create_other_accounts(request):
    try:
        data = json.loads(request.body or '{}')
        annee = int(data.get('annee') or 0)
        nom = (data.get('nom') or '').strip()
        description = (data.get('description') or '').strip()
        rows = data.get('rows') or []

        if not annee or not nom:
            return JsonResponse({'success': False, 'error': 'Année et nom requis'}, status=400)
        if not rows:
            return JsonResponse({'success': False, 'error': 'Aucun autre compte à budgéter'}, status=400)

        preview = _build_other_accounts_preview(rows)

        with transaction.atomic():
            budget = Budget.objects.create(
                annee=annee,
                nom=nom,
                description=description or 'Budget généré depuis le workflow autres comptes',
                cree_par=request.user,
            )

            for row in preview['rows']:
                monthly_amount = _safe_decimal(row['montant_mensuel'])
                ligne = LigneBudget.objects.create(
                    budget=budget,
                    compte=row['compte'],
                    libelle_compte=row['libelle_compte'],
                    montant_prevu=_safe_decimal(row['montant_annuel']),
                    janvier=monthly_amount,
                    fevrier=monthly_amount,
                    mars=monthly_amount,
                    avril=monthly_amount,
                    mai=monthly_amount,
                    juin=monthly_amount,
                    juillet=monthly_amount,
                    aout=monthly_amount,
                    septembre=monthly_amount,
                    octobre=monthly_amount,
                    novembre=monthly_amount,
                    decembre=monthly_amount,
                    notes='Généré depuis le total de la dernière année complète, réparti sur 12 mois',
                )
                ligne.save()

        return JsonResponse({
            'success': True,
            'budget_id': str(budget.id),
            'message': 'Budget autres comptes créé avec succès',
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_import_excel(request):
    """
    Importer un budget depuis Excel avec structure "mois en ligne":
    colonnes attendues (noms flexibles): compte, libellé (optionnel), mois, montant.
    """
    try:
        if 'file' not in request.FILES:
            return JsonResponse({
                'success': False,
                'error': 'Aucun fichier Excel fourni'
            }, status=400)

        annee = request.POST.get('annee')
        nom = (request.POST.get('nom') or '').strip()
        description = (request.POST.get('description') or '').strip()

        if not annee or not nom:
            return JsonResponse({
                'success': False,
                'error': 'Annee et nom du budget sont requis'
            }, status=400)

        try:
            annee = int(annee)
        except Exception:
            return JsonResponse({
                'success': False,
                'error': 'Annee invalide'
            }, status=400)

        excel_file = request.FILES['file']

        try:
            workbook = load_workbook(excel_file, data_only=True)
            worksheet = workbook.active
        except Exception as e:
            return JsonResponse({
                'success': False,
                'error': f'Fichier Excel invalide: {str(e)}'
            }, status=400)

        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            return JsonResponse({
                'success': False,
                'error': 'Le fichier est vide'
            }, status=400)

        columns = _detect_import_columns(header)
        missing = [field for field in ('compte', 'mois', 'montant') if field not in columns]
        if missing:
            return JsonResponse({
                'success': False,
                'error': (
                    "Colonnes manquantes dans l'entete. Requis: compte, mois, montant "
                    "(libelle optionnel)."
                ),
                'missing_columns': missing
            }, status=400)

        lignes_par_compte = {}
        errors = []
        warnings = []
        nb_rows_valides = 0

        for row_num, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(cell is not None and str(cell).strip() != '' for cell in row):
                continue

            compte_raw = row[columns['compte']]
            mois_raw = row[columns['mois']]
            montant_raw = row[columns['montant']]
            libelle_raw = row[columns['libelle']] if 'libelle' in columns else None

            compte = _normalize_compte(compte_raw)
            if not compte:
                errors.append(f"Ligne {row_num}: compte vide")
                continue

            mois = _parse_month(mois_raw)
            if not mois:
                errors.append(f"Ligne {row_num}: mois invalide ({mois_raw})")
                continue

            montant = _parse_decimal(montant_raw)
            if montant is None:
                errors.append(f"Ligne {row_num}: montant invalide ({montant_raw})")
                continue

            libelle = str(libelle_raw).strip() if libelle_raw is not None else ''
            if not libelle:
                libelle = compte

            if compte not in lignes_par_compte:
                lignes_par_compte[compte] = {
                    'libelle': libelle,
                    'mois': {m: Decimal('0') for m in range(1, 13)}
                }
            elif libelle and lignes_par_compte[compte]['libelle'] != libelle:
                warnings.append(
                    f"Compte {compte}: plusieurs libelles detectes, "
                    f"libelle conserve '{lignes_par_compte[compte]['libelle']}'"
                )

            lignes_par_compte[compte]['mois'][mois] += montant
            nb_rows_valides += 1

        if errors:
            return JsonResponse({
                'success': False,
                'error': (
                    f'Import bloque: {len(errors)} ligne(s) invalide(s). '
                    'Corrigez le fichier puis reessayez.'
                ),
                'row_errors': errors[:50]
            }, status=400)

        if not lignes_par_compte:
            return JsonResponse({
                'success': False,
                'error': 'Aucune ligne exploitable trouvee dans le fichier'
            }, status=400)

        with transaction.atomic():
            budget = Budget.objects.create(
                annee=annee,
                nom=nom,
                description=description,
                cree_par=request.user
            )

            nb_lignes_creees = 0
            for compte, data_ligne in sorted(lignes_par_compte.items()):
                mois_data = data_ligne['mois']
                montant_total = sum(mois_data.values())

                LigneBudget.objects.create(
                    budget=budget,
                    compte=compte,
                    libelle_compte=data_ligne['libelle'][:255],
                    montant_prevu=montant_total,
                    janvier=mois_data[1],
                    fevrier=mois_data[2],
                    mars=mois_data[3],
                    avril=mois_data[4],
                    mai=mois_data[5],
                    juin=mois_data[6],
                    juillet=mois_data[7],
                    aout=mois_data[8],
                    septembre=mois_data[9],
                    octobre=mois_data[10],
                    novembre=mois_data[11],
                    decembre=mois_data[12],
                )
                nb_lignes_creees += 1

        return JsonResponse({
            'success': True,
            'message': (
                f'Budget importe avec succes: {nb_lignes_creees} compte(s) '
                f'construit(s) depuis {nb_rows_valides} ligne(s) Excel'
            ),
            'budget_id': str(budget.id),
            'nb_lignes_creees': nb_lignes_creees,
            'nb_rows_valides': nb_rows_valides,
            'warnings': warnings[:20]
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@require_http_methods(["GET"])
def api_budget_import_excel_template(request):
    """Télécharger un template Excel pour l'import budget (mois en ligne)."""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Import Budget"

        ws.append(['Compte', 'Libellé', 'Mois', 'Montant'])
        ws.append(['6211', 'Salaires', 'Janvier', 1500000])
        ws.append(['6211', 'Salaires', 'Février', 1500000])
        ws.append(['6130', 'Carburant', 1, 300000])
        ws.append(['6130', 'Carburant', 2, 250000])

        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 36
        ws.column_dimensions['C'].width = 14
        ws.column_dimensions['D'].width = 16

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=template_import_budget_mois_en_ligne.xlsx'
        wb.save(response)
        return response

    except Exception as e:
        return HttpResponse(f'Erreur: {str(e)}', status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_creer_depuis_previsions(request):
    """
    API pour creer un budget a partir des previsions existantes.
    Permet des options de construction plus fine:
    - mode de montant (depenses, recettes, net)
    - source (all, auto, manual)
    - filtrage par types de prevision
    - filtrage par prefixes de comptes
    - granularite (detail par type)
    - repartition (respect previsions, egalisee_12_mois)
    """
    try:
        data = json.loads(request.body)

        annee = data.get('annee')
        nom = data.get('nom')
        description = data.get('description', '')

        if not annee or not nom:
            return JsonResponse({
                'success': False,
                'error': 'Annee et nom requis'
            }, status=400)

        montant_mode = (data.get('montant_mode') or 'depenses').strip().lower()
        source_mode = (data.get('source_mode') or 'all').strip().lower()
        repartition_mode = (data.get('repartition_mode') or 'respect_previsions').strip().lower()
        detail_raw = data.get('detail_par_type', False)
        if isinstance(detail_raw, str):
            detail_par_type = detail_raw.strip().lower() in {'1', 'true', 'yes', 'on'}
        else:
            detail_par_type = bool(detail_raw)
        type_previsions = data.get('type_previsions') or []
        compte_prefixes_raw = data.get('compte_prefixes', '')

        valid_montant_modes = {'depenses', 'recettes', 'net'}
        valid_source_modes = {'all', 'auto', 'manual'}
        valid_repartition_modes = {'respect_previsions', 'egalisee_12_mois'}

        if montant_mode not in valid_montant_modes:
            return JsonResponse({'success': False, 'error': 'montant_mode invalide'}, status=400)
        if source_mode not in valid_source_modes:
            return JsonResponse({'success': False, 'error': 'source_mode invalide'}, status=400)
        if repartition_mode not in valid_repartition_modes:
            return JsonResponse({'success': False, 'error': 'repartition_mode invalide'}, status=400)

        previsions_qs = PrevisionBudget.objects.filter(annee=annee)

        if source_mode == 'auto':
            previsions_qs = previsions_qs.filter(est_calculee_auto=True)
        elif source_mode == 'manual':
            previsions_qs = previsions_qs.filter(est_calculee_auto=False)

        if type_previsions:
            previsions_qs = previsions_qs.filter(type_prevision__in=type_previsions)

        compte_prefixes = [p.strip() for p in str(compte_prefixes_raw).split(',') if p.strip()]
        if compte_prefixes:
            q_compte = Q()
            for prefix in compte_prefixes:
                q_compte |= Q(cat__startswith=prefix)
            previsions_qs = previsions_qs.filter(q_compte)

        previsions_count = previsions_qs.count()
        if previsions_count == 0:
            return JsonResponse({
                'success': False,
                'error': f'Aucune prevision trouvee pour l\'annee {annee} avec les filtres choisis'
            }, status=400)

        budget = Budget.objects.create(
            annee=annee,
            nom=nom,
            description=description,
            cree_par=request.user
        )

        previsions = previsions_qs.select_related('base', 'ouvrage', 'district', 'commune')

        comptes_data = {}
        for prev in previsions:
            compte = prev.cat or 'SANS_COMPTE'
            mois = prev.mois if prev.mois else 1
            type_label = prev.get_type_prevision_display()

            if detail_par_type:
                compte_key = f"{compte}__{prev.type_prevision}"
                libelle_compte = f"[{type_label}] {prev.libelle or compte}"
            else:
                compte_key = compte
                libelle_compte = prev.libelle or compte

            if compte_key not in comptes_data:
                comptes_data[compte_key] = {
                    'compte': compte,
                    'libelle': libelle_compte,
                    'montant_total': Decimal('0'),
                    'mois': {m: Decimal('0') for m in range(1, 13)}
                }

            out_amount = prev.out_amount or Decimal('0')
            in_amount = prev.in_amount or Decimal('0')
            if montant_mode == 'depenses':
                montant = out_amount
            elif montant_mode == 'recettes':
                montant = in_amount
            else:
                montant = out_amount - in_amount

            comptes_data[compte_key]['montant_total'] += montant
            if 1 <= mois <= 12:
                comptes_data[compte_key]['mois'][mois] += montant

        nb_lignes_creees = 0
        for _, compte_info in comptes_data.items():
            montant_total = compte_info['montant_total']
            if montant_mode == 'net':
                keep_ligne = (montant_total != 0)
            else:
                keep_ligne = (montant_total > 0)

            if keep_ligne:
                mois_values = compte_info['mois']
                if repartition_mode == 'egalisee_12_mois':
                    mensualite = (montant_total / Decimal('12')).quantize(Decimal('0.01'))
                    mois_values = {m: mensualite for m in range(1, 13)}
                    somme_11 = mensualite * Decimal('11')
                    mois_values[12] = montant_total - somme_11

                LigneBudget.objects.create(
                    budget=budget,
                    compte=compte_info['compte'],
                    libelle_compte=compte_info['libelle'][:255],
                    montant_prevu=montant_total,
                    janvier=mois_values[1],
                    fevrier=mois_values[2],
                    mars=mois_values[3],
                    avril=mois_values[4],
                    mai=mois_values[5],
                    juin=mois_values[6],
                    juillet=mois_values[7],
                    aout=mois_values[8],
                    septembre=mois_values[9],
                    octobre=mois_values[10],
                    novembre=mois_values[11],
                    decembre=mois_values[12]
                )
                nb_lignes_creees += 1

        return JsonResponse({
            'success': True,
            'message': f'Budget cree avec succes depuis {previsions_count} prevision(s)',
            'budget_id': str(budget.id),
            'nb_lignes_creees': nb_lignes_creees,
            'nb_previsions_utilisees': previsions_count,
            'options_utilisees': {
                'montant_mode': montant_mode,
                'source_mode': source_mode,
                'repartition_mode': repartition_mode,
                'detail_par_type': detail_par_type,
                'type_previsions': type_previsions,
                'compte_prefixes': compte_prefixes,
            }
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)
@login_required
def budget_detail(request, budget_id):
    """Page de détail d'un budget avec ses lignes."""
    budget = get_object_or_404(Budget, id=budget_id)

    context = {
        'budget': budget,
    }

    return render(request, 'accounting/budget_detail.html', context)


@login_required
@require_http_methods(["GET"])
def api_budget_lignes(request, budget_id):
    """API pour récupérer les lignes d'un budget."""
    try:
        budget = get_object_or_404(Budget, id=budget_id)

        lignes = []
        for ligne in budget.lignes.all():
            lignes.append({
                'id': str(ligne.id),
                'compte': ligne.compte,
                'libelle_compte': ligne.libelle_compte,
                'montant_prevu': float(ligne.montant_prevu),
                'janvier': float(ligne.janvier or 0),
                'fevrier': float(ligne.fevrier or 0),
                'mars': float(ligne.mars or 0),
                'avril': float(ligne.avril or 0),
                'mai': float(ligne.mai or 0),
                'juin': float(ligne.juin or 0),
                'juillet': float(ligne.juillet or 0),
                'aout': float(ligne.aout or 0),
                'septembre': float(ligne.septembre or 0),
                'octobre': float(ligne.octobre or 0),
                'novembre': float(ligne.novembre or 0),
                'decembre': float(ligne.decembre or 0),
                'notes': ligne.notes or '',
            })

        return JsonResponse({
            'success': True,
            'lignes': lignes,
            'count': len(lignes)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_ligne_save(request, budget_id):
    """API pour ajouter/modifier une ligne budgétaire."""
    try:
        budget = get_object_or_404(Budget, id=budget_id)

        if budget.valide:
            return JsonResponse({
                'success': False,
                'error': 'Budget déjÃ  validé, modification impossible'
            }, status=400)

        data = json.loads(request.body)

        compte = data.get('compte')
        libelle_compte = data.get('libelle_compte')
        montant_prevu = data.get('montant_prevu', 0)

        if not compte or not libelle_compte:
            return JsonResponse({
                'success': False,
                'error': 'Compte et libellé requis'
            }, status=400)

        ligne, created = LigneBudget.objects.update_or_create(
            budget=budget,
            compte=compte,
            defaults={
                'libelle_compte': libelle_compte,
                'montant_prevu': Decimal(str(montant_prevu)),
                'janvier': Decimal(str(data.get('janvier', 0))),
                'fevrier': Decimal(str(data.get('fevrier', 0))),
                'mars': Decimal(str(data.get('mars', 0))),
                'avril': Decimal(str(data.get('avril', 0))),
                'mai': Decimal(str(data.get('mai', 0))),
                'juin': Decimal(str(data.get('juin', 0))),
                'juillet': Decimal(str(data.get('juillet', 0))),
                'aout': Decimal(str(data.get('aout', 0))),
                'septembre': Decimal(str(data.get('septembre', 0))),
                'octobre': Decimal(str(data.get('octobre', 0))),
                'novembre': Decimal(str(data.get('novembre', 0))),
                'decembre': Decimal(str(data.get('decembre', 0))),
                'notes': data.get('notes', ''),
            }
        )

        action = 'créée' if created else 'modifiée'

        return JsonResponse({
            'success': True,
            'message': f'Ligne budgétaire {action} avec succès',
            'ligne_id': str(ligne.id)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_budget_valider(request, budget_id):
    """API pour valider (figer) un budget."""
    try:
        budget = get_object_or_404(Budget, id=budget_id)

        if budget.valide:
            return JsonResponse({
                'success': False,
                'error': 'Ce budget est déjà validé'
            }, status=400)

        if not budget.lignes.exists():
            return JsonResponse({
                'success': False,
                'error': 'Impossible de valider un budget sans ligne budgétaire'
            }, status=400)

        budget.valide = True
        budget.date_validation = timezone.now()
        budget.valide_par = request.user
        budget.save(update_fields=['valide', 'date_validation', 'valide_par'])

        return JsonResponse({
            'success': True,
            'message': 'Budget validé avec succès'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


# ========== DEMANDES DE FONDS ==========

@login_required
def demandes_fonds_list(request):
    """Page de liste des demandes de fonds."""
    budgets = Budget.objects.filter(valide=True).order_by('-annee', '-date_creation')

    context = {
        'budgets': budgets,
        'journaux_suggeres': _get_demande_journal_suggestions(request.user),
    }

    return render(request, 'accounting/demandes_fonds_list.html', context)


@login_required
@require_http_methods(["GET"])
def api_demandes_fonds_list(request):
    """API pour récupérer la liste des demandes de fonds."""
    try:
        budget_id = request.GET.get('budget_id')
        annee = request.GET.get('annee')
        statut = request.GET.get('statut')

        queryset = DemandeFonds.objects.all()

        if budget_id:
            queryset = queryset.filter(budget_id=budget_id)
        if annee:
            queryset = queryset.filter(annee=annee)
        if statut:
            queryset = queryset.filter(statut=statut)

        demandes = []
        for demande in queryset:
            nb_lignes = demande.lignes.count()
            workflow_meta = _parse_demande_workflow_commentaire(demande.commentaire)
            current_month_remaining = workflow_meta.get('current_month_remaining', float(demande.depenses_reelles))
            next_month_expenses = workflow_meta.get('next_month_expenses', float(demande.previsions_futures))
            future_months_total = workflow_meta.get('future_months_total', 0)
            date_demande = workflow_meta.get('date_demande', '')

            demandes.append({
                'id': str(demande.id),
                'budget_nom': demande.budget.nom,
                'mois': demande.mois,
                'mois_nom': demande.mois_nom,
                'annee': demande.annee,
                'periode': f"{demande.mois_nom} {demande.annee}",
                'date_demande': date_demande,
                'solde_initial': float(demande.solde_initial),
                'depenses_reelles': current_month_remaining,
                'previsions_futures': next_month_expenses,
                'future_months_total': future_months_total,
                'montant_demande': float(demande.montant_demande),
                'statut': demande.statut,
                'statut_display': demande.get_statut_display(),
                'date_creation': demande.date_creation.strftime('%Y-%m-%d %H:%M'),
                'nb_lignes': nb_lignes,
            })

        return JsonResponse({
            'success': True,
            'demandes': demandes,
            'count': len(demandes)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_demande_fonds_preview(request):
    """Aperçu du nouveau workflow de demande de fonds."""
    try:
        data = json.loads(request.body)

        budget_id = data.get('budget_id')
        mois = int(data.get('mois'))
        annee = int(data.get('annee'))
        date_demande_str = data.get('date_demande')

        if not date_demande_str:
            return JsonResponse({
                'success': False,
                'error': "La date de demande est obligatoire."
            }, status=400)

        date_demande = datetime.strptime(date_demande_str, '%Y-%m-%d').date()
        budget = get_object_or_404(Budget, id=budget_id)

        journal_balances, _ = _parse_journal_balances(data.get('journal_balances'))
        preview = _build_demande_fonds_projection(
            request.user,
            budget,
            mois,
            annee,
            date_demande,
            journal_balances,
        )

        return JsonResponse({
            'success': True,
            'summary': {
                'total_journal_balances': float(preview['summary']['total_journal_balances']),
                'total_current_remaining': float(preview['summary']['total_current_remaining']),
                'total_next_month': float(preview['summary']['total_next_month']),
                'total_future_months': float(preview['summary']['total_future_months']),
                'total_gross_need': float(preview['summary']['total_gross_need']),
                'total_requested': float(preview['summary']['total_requested']),
                'date_demande': preview['summary']['date_demande'],
                'next_month_label': preview['summary']['next_month_label'],
            },
            'rows': [
                {
                    'compte': row['compte'],
                    'libelle_compte': row['libelle_compte'],
                    'budget_annuel': float(row['budget_annuel']),
                    'depenses_mois_a_date': float(row['depenses_mois_a_date']),
                    'reste_mois_courant': float(row['reste_mois_courant']),
                    'depenses_mois_suivant': float(row['depenses_mois_suivant']),
                    'previsions_mois_suivants': float(row['previsions_mois_suivants']),
                    'besoin_brut': float(row['besoin_brut']),
                    'solde_impute': float(row['solde_impute']),
                    'montant_demande': float(row['montant_demande']),
                    'budget_restant': float(row['budget_restant']),
                    'depassement': float(row['depassement']),
                    'sources': row['sources'],
                }
                for row in preview['rows']
            ],
            'future_months_breakdown': _serialize_future_months(preview['future_months_breakdown']),
        })

    except ValueError as exc:
        return JsonResponse({
            'success': False,
            'error': str(exc)
        }, status=400)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_demande_fonds_create(request):
    """Crée une demande de fonds à partir du nouveau workflow avec aperçu."""
    try:
        data = json.loads(request.body)

        budget_id = data.get('budget_id')
        mois = int(data.get('mois'))
        annee = int(data.get('annee'))
        date_demande_str = data.get('date_demande')

        budget = get_object_or_404(Budget, id=budget_id)

        if not budget.valide:
            return JsonResponse({
                'success': False,
                'error': 'Le budget doit être validé avant de créer une demande de fonds'
            }, status=400)

        # Vérifier si demande existe déjÃ 
        if DemandeFonds.objects.filter(budget=budget, annee=annee, mois=mois).exists():
            return JsonResponse({
                'success': False,
                'error': f'Une demande existe déjÃ  pour {annee}-{mois:02d}'
            }, status=400)

        if not date_demande_str:
            return JsonResponse({
                'success': False,
                'error': "La date de demande est obligatoire."
            }, status=400)

        date_demande = datetime.strptime(date_demande_str, '%Y-%m-%d').date()
        journal_balances, total_journal_balance = _parse_journal_balances(data.get('journal_balances'))
        preview = _build_demande_fonds_projection(
            request.user,
            budget,
            mois,
            annee,
            date_demande,
            journal_balances,
        )

        demande = DemandeFonds.objects.create(
            budget=budget,
            mois=mois,
            annee=annee,
            cree_par=request.user,
            statut='brouillon',
            solde_initial=total_journal_balance,
            depenses_reelles=preview['summary']['total_current_remaining'],
            previsions_futures=preview['summary']['total_next_month'],
            montant_demande=preview['summary']['total_requested'],
            commentaire=_demande_fonds_commentaire_payload(preview, journal_balances),
        )

        for row in preview['rows']:
            LigneDemandeFonds.objects.create(
                demande=demande,
                compte=row['compte'],
                libelle_compte=row['libelle_compte'],
                budget_annuel=row['budget_annuel'],
                prevision_mois=row['reste_mois_courant'],
                previsions_passees=row['depenses_mois_a_date'],
                previsions_futures=row['depenses_mois_suivant'],
                solde_reel=row['solde_impute'],
                depenses_reelles=row['previsions_mois_suivants'],
                budget_restant=row['budget_restant'],
                depassement=row['depassement'],
                montant_demande=row['montant_demande'],
            )

        return JsonResponse({
            'success': True,
            'message': 'Demande de fonds créée avec succès',
            'demande_id': str(demande.id),
            'montant_demande': float(demande.montant_demande)
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
def demande_fonds_detail(request, demande_id):
    """Page de détail d'une demande de fonds."""
    demande = get_object_or_404(DemandeFonds, id=demande_id)

    context = {
        'demande': demande,
    }

    return render(request, 'accounting/demande_fonds_detail.html', context)


@login_required
@require_http_methods(["GET"])
def api_demande_fonds_lignes(request, demande_id):
    """API pour récupérer les lignes d'une demande de fonds."""
    try:
        demande = get_object_or_404(DemandeFonds, id=demande_id)
        workflow_meta = _parse_demande_workflow_commentaire(demande.commentaire)

        lignes = []
        for ligne in demande.lignes.all():
            lignes.append({
                'id': str(ligne.id),
                'compte': ligne.compte,
                'libelle_compte': ligne.libelle_compte,
                'budget_annuel': float(ligne.budget_annuel),
                'prevision_mois': float(ligne.prevision_mois),  # reste mois courant
                'previsions_passees': float(ligne.previsions_passees),  # dépenses à date
                'previsions_futures': float(ligne.previsions_futures),  # mois suivant
                'solde_reel': float(ligne.solde_reel),  # solde imputé
                'depenses_reelles': float(ligne.depenses_reelles),  # mois ultérieurs
                'budget_restant': float(ligne.budget_restant),
                'depassement': float(ligne.depassement),
                'sources': [],
                'montant_demande': float(ligne.montant_demande),
                'explication_ecart': ligne.explication_ecart or '',
                'explication_depassement': ligne.explication_depassement or '',
                'a_depassement': ligne.depassement > 0,
            })

        return JsonResponse({
            'success': True,
            'lignes': lignes,
            'count': len(lignes),
            'workflow_meta': workflow_meta,
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_demande_fonds_update_explication(request, ligne_id):
    """API pour mettre Ã  jour les explications d'une ligne."""
    try:
        ligne = get_object_or_404(LigneDemandeFonds, id=ligne_id)

        if ligne.demande.statut not in ['brouillon', 'en_cours']:
            return JsonResponse({
                'success': False,
                'error': 'Demande déjÃ  validée, modification impossible'
            }, status=400)

        data = json.loads(request.body)

        ligne.explication_ecart = data.get('explication_ecart', '')
        ligne.explication_depassement = data.get('explication_depassement', '')
        ligne.save()

        return JsonResponse({
            'success': True,
            'message': 'Explications mises Ã  jour'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
@csrf_exempt
@require_http_methods(["POST"])
def api_demande_fonds_valider(request, demande_id):
    """API pour valider une demande de fonds."""
    try:
        demande = get_object_or_404(DemandeFonds, id=demande_id)

        if demande.statut != 'brouillon':
            return JsonResponse({
                'success': False,
                'error': 'Seules les demandes en brouillon peuvent être validées'
            }, status=400)

        # Vérifier que tous les dépassements ont une explication
        lignes_sans_explication = demande.lignes.filter(
            depassement__gt=0,
            explication_depassement__isnull=True
        ).count()

        if lignes_sans_explication > 0:
            return JsonResponse({
                'success': False,
                'error': f'{lignes_sans_explication} ligne(s) avec dépassement sans explication'
            }, status=400)

        demande.statut = 'validee'
        demande.date_validation = timezone.now()
        demande.valide_par = request.user
        demande.save()

        return JsonResponse({
            'success': True,
            'message': 'Demande de fonds validée avec succès'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)
