from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.billing.models import PaiementForfaitaire, PointEau, Releve
from .extracteurs_description import extraire_ouvrage_et_mois_generique
from .models import BaseDiameX
from .views import _get_responsable_scoped_ouvrage_ids
from .views_dashboard_forfaitaire import _get_scoped_forfaitaire_ouvrage_ids
from .views_verification_ouvrages import diviser_paiement_volume_multi


def _resolve_period(request):
    today = timezone.localdate()
    raw_start = (request.GET.get('date_debut') or f'{today.year}-01-01').strip()
    raw_end = (request.GET.get('date_fin') or today.isoformat()).strip()
    try:
        start = date.fromisoformat(raw_start)
        end = date.fromisoformat(raw_end)
        if start > end:
            raise ValueError
        return start, end, raw_start, raw_end, ''
    except (TypeError, ValueError):
        return None, None, raw_start, raw_end, 'La période sélectionnée est invalide.'


def _month_keys(start, end):
    keys = set()
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        keys.add((year, month))
        if month == 12:
            year, month = year + 1, 1
        else:
            month += 1
    return keys


def _previous_month_start(value):
    if value.month == 1:
        return date(value.year - 1, 12, 1)
    return date(value.year, value.month - 1, 1)


def _month_end(value):
    if value.month == 12:
        return date(value.year, 12, 31)
    return date(value.year, value.month + 1, 1) - timedelta(days=1)


def _billing_period(releve_date):
    if releve_date.month == 12:
        return releve_date.year + 1, 1
    return releve_date.year, releve_date.month + 1


def _accessible_volume_ouvrage_ids(user):
    return set(_get_responsable_scoped_ouvrage_ids(
        user,
        effective_gestionnaire_ids=None,
        modalites=['au_volume', 'mixte'],
    ))


def _accessible_forfaitaire_ouvrage_ids(user, start, end):
    ouvrage_ids = set()
    for year in range(start.year, end.year + 1):
        ouvrage_ids.update(_get_scoped_forfaitaire_ouvrage_ids(user, annee=year))
    return ouvrage_ids


def _location_map(ouvrage_ids):
    locations = {}
    points = PointEau.objects.filter(
        ouvrage_id__in=ouvrage_ids,
    ).select_related(
        'parcelle__village__fokontany__commune__district__region'
    ).defer('photos').order_by('ouvrage_id', 'id')  # photos base64 inutiles ici (très lourdes)
    for point in points:
        if point.ouvrage_id in locations:
            continue
        parcelle = point.parcelle
        village = getattr(parcelle, 'village', None)
        fokontany = getattr(village, 'fokontany', None)
        commune = getattr(fokontany, 'commune', None)
        district = getattr(commune, 'district', None)
        region = getattr(district, 'region', None)
        locations[point.ouvrage_id] = {
            'region': getattr(region, 'nom_region', '') or 'Sans région',
            'district': getattr(district, 'nom_district', '') or 'Sans district',
            'commune': getattr(commune, 'nom_commune', '') or 'Sans commune',
        }
    return locations


def _finalize_rows(rows_by_id):
    locations = _location_map(rows_by_id.keys())
    rows = []
    for ouvrage_id, item in rows_by_id.items():
        expected = Decimal(str(item['expected'] or 0))
        paid = Decimal(str(item['paid'] or 0))
        item['gap'] = expected - paid
        item['rate'] = round(float((paid / expected) * 100), 1) if expected else 0
        item.update(locations.get(ouvrage_id, {
            'region': 'Sans région', 'district': 'Sans district', 'commune': 'Sans commune'
        }))
        rows.append(item)
    return sorted(rows, key=lambda row: (
        row['region'].casefold(), row['district'].casefold(),
        row['commune'].casefold(), row['ouvrage_nom'].casefold()
    ))


def _stats(rows, paid_point_count=0, billed_point_count=0):
    expected = sum((row['expected'] for row in rows), Decimal('0'))
    paid = sum((row['paid'] for row in rows), Decimal('0'))
    return {
        'expected': expected,
        'paid': paid,
        'gap': expected - paid,
        'rate': round(float((paid / expected) * 100), 1) if expected else 0,
        'ouvrage_count': len(rows),
        'paid_point_count': paid_point_count,
        'billed_point_count': billed_point_count,
    }


def _build_volume_rows(request, start, end):
    rows_by_id = defaultdict(lambda: {'expected': Decimal('0'), 'paid': Decimal('0')})
    point_payment_status = {}
    period_keys = _month_keys(start, end)
    allowed_ids = _accessible_volume_ouvrage_ids(request.user)
    source_start = _previous_month_start(start)
    source_end = _month_end(_previous_month_start(end))
    releves = Releve.objects.filter(
        point_eau__ouvrage_id__in=allowed_ids,
        datereleve__gte=source_start,
        datereleve__lte=source_end,
        montant__gt=0,
    ).exclude(remarque__icontains='coupé').select_related('point_eau__ouvrage').defer('point_eau__photos')
    for releve in releves:
        if _billing_period(releve.datereleve) not in period_keys:
            continue
        ouvrage = releve.point_eau.ouvrage
        item = rows_by_id[ouvrage.id]
        item.update({'ouvrage_id': ouvrage.id, 'ouvrage_nom': ouvrage.nom_ouvrage})
        item['expected'] += Decimal(str(releve.montant or 0))
        point_status = point_payment_status.setdefault(
            releve.point_eau_id,
            {'has_invoice': False, 'all_paid': True},
        )
        point_status['has_invoice'] = True
        if not (
            str(releve.numeroDiameX or '').strip()
            and str(releve.baseDiameX or '').strip()
        ):
            point_status['all_paid'] = False

    relevant_ids = set(rows_by_id)
    for entry in BaseDiameX.objects.filter(compte_code='707401'):
        segments = []
        if entry.description and ';' in entry.description:
            segments = diviser_paiement_volume_multi(
                entry.description, entry.compte_code, entry.montant_in or 0
            )
        parsed = [
            (segment.get('extraction') or {}, Decimal(str(segment.get('montant') or 0)))
            for segment in segments
        ] if segments else [(
            extraire_ouvrage_et_mois_generique(entry.description or '', entry.compte_code),
            Decimal(str(entry.montant_in or 0)),
        )]
        for extraction, amount in parsed:
            ouvrage = extraction.get('ouvrage_obj')
            key = (extraction.get('annee'), extraction.get('mois'))
            if not ouvrage or ouvrage.id not in relevant_ids or key not in period_keys:
                continue
            rows_by_id[ouvrage.id]['paid'] += amount

    return rows_by_id, point_payment_status


def _build_forfaitaire_rows(request, start, end):
    rows_by_id = defaultdict(lambda: {'expected': Decimal('0'), 'paid': Decimal('0')})
    point_payment_status = {}
    allowed_ids = _accessible_forfaitaire_ouvrage_ids(request.user, start, end)
    payments = PaiementForfaitaire.objects.filter(
        point_eau__ouvrage_id__in=allowed_ids,
        date_facture__gte=start,
        date_facture__lte=end,
    ).select_related('point_eau__ouvrage').defer('point_eau__photos')
    for payment in payments:
        ouvrage = payment.point_eau.ouvrage
        item = rows_by_id[ouvrage.id]
        item.update({'ouvrage_id': ouvrage.id, 'ouvrage_nom': ouvrage.nom_ouvrage})
        item['expected'] += Decimal(str(payment.montant or 0))
        point_status = point_payment_status.setdefault(
            payment.point_eau_id,
            {'has_invoice': False, 'all_paid': True},
        )
        point_status['has_invoice'] = True
        if not payment.est_paye:
            point_status['all_paid'] = False

    relevant_ids = set(rows_by_id)
    transactions = BaseDiameX.objects.filter(
        compte_code='748411',
        date_m__gte=start,
        date_m__lte=end,
    )
    for entry in transactions:
        extraction = extraire_ouvrage_et_mois_generique(entry.description or '', '748411')
        ouvrage = extraction.get('ouvrage_obj')
        if not ouvrage or ouvrage.id not in relevant_ids:
            continue
        rows_by_id[ouvrage.id]['paid'] += Decimal(str(entry.montant_in or 0))

    return rows_by_id, point_payment_status


def _payment_point_counts(point_payment_status):
    paid_point_count = sum(
        1 for status in point_payment_status.values()
        if status['has_invoice'] and status['all_paid']
    )
    billed_point_count = sum(
        1 for status in point_payment_status.values() if status['has_invoice']
    )
    return paid_point_count, billed_point_count


@login_required
def dashboard_paiements_volume_periode(request):
    start, end, raw_start, raw_end, error = _resolve_period(request)
    rows_by_id = defaultdict(lambda: {'expected': Decimal('0'), 'paid': Decimal('0')})
    point_payment_status = {}
    if not error:
        rows_by_id, point_payment_status = _build_volume_rows(request, start, end)

    rows = _finalize_rows(rows_by_id)
    paid_point_count, billed_point_count = _payment_point_counts(point_payment_status)
    return render(request, 'accounting/dashboard_paiements_periode.html', {
        'dashboard_kind': 'volume',
        'title': 'Paiements au volume — période multi-années',
        'account_code': '707401',
        'date_debut': raw_start,
        'date_fin': raw_end,
        'period_error': error,
        'rows': rows,
        'stats': _stats(rows, paid_point_count, billed_point_count),
        'back_url_name': 'accounting:dashboard_paiements_volume',
        'export_url_name': 'accounting:dashboard_paiements_volume_periode_export',
        'expected_label': 'À collecter',
    })


@login_required
def dashboard_paiements_forfaitaire_periode(request):
    start, end, raw_start, raw_end, error = _resolve_period(request)
    rows_by_id = defaultdict(lambda: {'expected': Decimal('0'), 'paid': Decimal('0')})
    point_payment_status = {}
    if not error:
        rows_by_id, point_payment_status = _build_forfaitaire_rows(request, start, end)

    rows = _finalize_rows(rows_by_id)
    paid_point_count, billed_point_count = _payment_point_counts(point_payment_status)
    return render(request, 'accounting/dashboard_paiements_periode.html', {
        'dashboard_kind': 'forfaitaire',
        'title': 'Paiements forfaitaires — période multi-années',
        'account_code': '748411',
        'date_debut': raw_start,
        'date_fin': raw_end,
        'period_error': error,
        'rows': rows,
        'stats': _stats(rows, paid_point_count, billed_point_count),
        'back_url_name': 'accounting:dashboard_paiements_forfaitaire',
        'export_url_name': 'accounting:dashboard_paiements_forfaitaire_periode_export',
        'expected_label': 'Paiement forfaitaire',
    })


def _export_rows_to_excel(rows, stats, title, account_code, raw_start, raw_end, filename_prefix, expected_label='À collecter'):
    wb = Workbook()
    ws = wb.active
    ws.title = "Paiements"

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws.merge_cells('A1:H1')
    cell = ws['A1']
    cell.value = title
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:H2')
    cell = ws['A2']
    cell.value = f"Compte {account_code} — période du {raw_start} au {raw_end}"
    cell.font = Font(size=10, italic=True)
    cell.alignment = Alignment(horizontal='center')

    current_row = 4
    stats_labels = [
        (expected_label, stats['expected']),
        ('Collecté', stats['paid']),
        ('Écart', stats['gap']),
        ('Taux (%)', stats['rate']),
        ('Ouvrages', stats['ouvrage_count']),
        ('Points d\'eau facturés', stats['billed_point_count']),
        ('Points d\'eau ayant tout payé', stats['paid_point_count']),
    ]
    for label, value in stats_labels:
        ws.cell(row=current_row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=current_row, column=2, value=float(value) if isinstance(value, Decimal) else value)
        current_row += 1

    current_row += 1
    headers = ['Région', 'District', 'Commune', 'Ouvrage', expected_label, 'Collecté', 'Écart', 'Taux (%)']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    current_row += 1
    for row in rows:
        ws.cell(row=current_row, column=1, value=row['region']).border = border
        ws.cell(row=current_row, column=2, value=row['district']).border = border
        ws.cell(row=current_row, column=3, value=row['commune']).border = border
        ws.cell(row=current_row, column=4, value=row['ouvrage_nom']).border = border
        ws.cell(row=current_row, column=5, value=float(row['expected'])).border = border
        ws.cell(row=current_row, column=6, value=float(row['paid'])).border = border
        ws.cell(row=current_row, column=7, value=float(row['gap'])).border = border
        ws.cell(row=current_row, column=8, value=row['rate']).border = border
        current_row += 1

    for col_num in range(1, 9):
        ws.column_dimensions[get_column_letter(col_num)].width = 18

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"{filename_prefix}_{raw_start}_a_{raw_end}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def dashboard_paiements_volume_periode_export(request):
    start, end, raw_start, raw_end, error = _resolve_period(request)
    if error:
        return HttpResponse(error, status=400)

    rows_by_id, point_payment_status = _build_volume_rows(request, start, end)
    rows = _finalize_rows(rows_by_id)
    paid_point_count, billed_point_count = _payment_point_counts(point_payment_status)
    stats = _stats(rows, paid_point_count, billed_point_count)
    return _export_rows_to_excel(
        rows, stats,
        'Paiements au volume — période multi-années',
        '707401', raw_start, raw_end,
        'paiements_volume_periode',
        expected_label='À collecter',
    )


@login_required
def dashboard_paiements_forfaitaire_periode_export(request):
    start, end, raw_start, raw_end, error = _resolve_period(request)
    if error:
        return HttpResponse(error, status=400)

    rows_by_id, point_payment_status = _build_forfaitaire_rows(request, start, end)
    rows = _finalize_rows(rows_by_id)
    paid_point_count, billed_point_count = _payment_point_counts(point_payment_status)
    stats = _stats(rows, paid_point_count, billed_point_count)
    return _export_rows_to_excel(
        rows, stats,
        'Paiements forfaitaires — période multi-années',
        '748411', raw_start, raw_end,
        'paiements_forfaitaire_periode',
        expected_label='Paiement forfaitaire',
    )
