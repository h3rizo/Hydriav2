"""
Export Excel et impression pour vérification des scans
"""
from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Q, Count, Case, When, IntegerField
from collections import defaultdict
from .models import BaseDiameX, ScanVerificationCache
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


@login_required
def export_excel_scans(request):
    """
    Export Excel avec détails des liens morts et sans scan pour un fichier spécifique
    """
    fichier_source = request.GET.get('fichier')

    if not fichier_source:
        return HttpResponse("Fichier source requis", status=400)

    # Récupérer les entrées du fichier
    # Exclure:
    # - 00BOOK (transferts internes)
    # - 00PERS (comptes de tiers)
    # - 00FND1 SANS "EDF" dans description/info (transferts inter-fichiers)
    entries_base = BaseDiameX.objects.exclude(
        Q(compte_code__icontains='00BOOK') |
        Q(compte_code__icontains='00PERS') |
        (Q(compte_code__icontains='00FND1') & ~Q(description__icontains='EDF') & ~Q(info__icontains='EDF'))
    ).filter(fichier_source=fichier_source)

    # Récupérer le cache de vérification
    cache = ScanVerificationCache.objects.first()
    scans_valides_ids = set(cache.scans_valides_ids) if cache and cache.scans_valides_ids else set()
    liens_morts_ids = set(cache.liens_morts_ids) if cache and cache.liens_morts_ids else set()

    # Groupement par mois
    recap_data = entries_base.values('date_m').annotate(
        total=Count('id'),
        scans_valides=Count(
            Case(
                When(id__in=scans_valides_ids, then=1),
                output_field=IntegerField()
            )
        ) if scans_valides_ids else Count('id', filter=Q(id__in=[])),
        liens_morts=Count(
            Case(
                When(id__in=liens_morts_ids, then=1),
                output_field=IntegerField()
            )
        ) if liens_morts_ids else Count('id', filter=Q(id__in=[])),
        sans_scan=Count('id', filter=Q(scan__isnull=True) | Q(scan=''))
    ).order_by('-date_m')

    # Organiser par mois
    recap_par_mois = []
    for entry in recap_data:
        mois_key = entry['date_m'].strftime('%Y-%m') if entry['date_m'] else 'Sans date'
        recap_par_mois.append({
            'mois': mois_key,
            'total': entry['total'],
            'scans_valides': entry['scans_valides'],
            'liens_morts': entry['liens_morts'],
            'sans_scan': entry['sans_scan'],
            'taux_valides': int((entry['scans_valides'] / entry['total'] * 100) if entry['total'] > 0 else 0)
        })

    # Créer le workbook Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Vérification Scans"

    # Styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    subheader_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    subheader_font = Font(bold=True, size=10)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # En-tête du rapport
    ws.merge_cells('A1:F1')
    cell = ws['A1']
    cell.value = f"Vérification des scans - {fichier_source}"
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:F2')
    cell = ws['A2']
    cell.value = f"Date d'export: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    cell.font = Font(size=10, italic=True)
    cell.alignment = Alignment(horizontal='center')

    # Ligne vide
    current_row = 4

    # En-têtes du tableau récapitulatif
    headers = ['Mois', 'Total', 'Valides', 'Liens morts', 'Sans scan', 'Taux valides']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    current_row += 1

    # Données du tableau récapitulatif
    totaux = {'total': 0, 'valides': 0, 'liens_morts': 0, 'sans_scan': 0}

    for mois_data in recap_par_mois:
        ws.cell(row=current_row, column=1, value=mois_data['mois']).border = border
        ws.cell(row=current_row, column=2, value=mois_data['total']).border = border
        ws.cell(row=current_row, column=3, value=mois_data['scans_valides']).border = border
        ws.cell(row=current_row, column=4, value=mois_data['liens_morts']).border = border
        ws.cell(row=current_row, column=5, value=mois_data['sans_scan']).border = border
        ws.cell(row=current_row, column=6, value=f"{mois_data['taux_valides']}%").border = border

        totaux['total'] += mois_data['total']
        totaux['valides'] += mois_data['scans_valides']
        totaux['liens_morts'] += mois_data['liens_morts']
        totaux['sans_scan'] += mois_data['sans_scan']

        current_row += 1

    # Ligne de total
    cell = ws.cell(row=current_row, column=1, value="TOTAL")
    cell.font = Font(bold=True)
    cell.border = border
    ws.cell(row=current_row, column=2, value=totaux['total']).border = border
    ws.cell(row=current_row, column=3, value=totaux['valides']).border = border
    ws.cell(row=current_row, column=4, value=totaux['liens_morts']).border = border
    ws.cell(row=current_row, column=5, value=totaux['sans_scan']).border = border
    taux_global = int((totaux['valides'] / totaux['total'] * 100) if totaux['total'] > 0 else 0)
    ws.cell(row=current_row, column=6, value=f"{taux_global}%").border = border

    current_row += 2

    # Détails des liens morts ET sans scan (exclure les valides)
    for mois_data in recap_par_mois:
        if mois_data['liens_morts'] == 0 and mois_data['sans_scan'] == 0:
            continue  # Skip les mois sans problème

        # Récupérer les entrées problématiques pour ce mois
        year, month = None, None
        if mois_data['mois'] != 'Sans date':
            year, month = map(int, mois_data['mois'].split('-'))

        entries_mois = entries_base
        if year and month:
            entries_mois = entries_mois.filter(date_m__year=year, date_m__month=month)
        else:
            entries_mois = entries_mois.filter(date_m__isnull=True)

        # Filtrer pour ne garder que liens morts et sans scan
        entries_problemes = entries_mois.filter(
            Q(id__in=liens_morts_ids) | Q(scan__isnull=True) | Q(scan='')
        ).exclude(id__in=scans_valides_ids)

        if entries_problemes.count() == 0:
            continue

        # En-tête du mois
        ws.merge_cells(f'A{current_row}:J{current_row}')
        cell = ws.cell(row=current_row, column=1)
        cell.value = f"Détails - {mois_data['mois']} ({entries_problemes.count()} entrées)"
        cell.fill = subheader_fill
        cell.font = subheader_font
        cell.border = border
        current_row += 1

        # En-têtes des colonnes de détail
        detail_headers = ['Date Rec', 'Rec', 'Journal', 'Description', 'Compte', 'Montant IN', 'Montant OUT', 'PERS', 'Détail', 'Statut']
        for col_num, header in enumerate(detail_headers, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = header
            cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            cell.font = Font(bold=True, size=9)
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        current_row += 1

        # Données des entrées problématiques
        for entry in entries_problemes.order_by('-date_rec')[:1000]:  # Limite 1000 par mois
            # Déterminer le statut
            if entry.id in liens_morts_ids:
                statut = "Lien mort"
            elif not entry.scan or entry.scan.strip() == '':
                statut = "Sans scan"
            else:
                statut = "Non vérifié"

            ws.cell(row=current_row, column=1, value=entry.date_rec.strftime('%d/%m/%Y') if entry.date_rec else '-').border = border
            ws.cell(row=current_row, column=2, value=entry.rec or '-').border = border
            ws.cell(row=current_row, column=3, value=entry.journal_code or '-').border = border
            ws.cell(row=current_row, column=4, value=entry.description or '-').border = border
            ws.cell(row=current_row, column=5, value=entry.compte_code or '-').border = border
            ws.cell(row=current_row, column=6, value=float(entry.montant_in) if entry.montant_in else 0).border = border
            ws.cell(row=current_row, column=7, value=float(entry.montant_out) if entry.montant_out else 0).border = border
            ws.cell(row=current_row, column=8, value=entry.pers or '-').border = border
            ws.cell(row=current_row, column=9, value=entry.detail or '-').border = border
            ws.cell(row=current_row, column=10, value=statut).border = border

            current_row += 1

        current_row += 1  # Ligne vide entre les mois

    # Ajuster la largeur des colonnes
    for col_num in range(1, 11):
        column_letter = get_column_letter(col_num)
        if col_num == 4:  # Description
            ws.column_dimensions[column_letter].width = 50
        elif col_num == 9:  # Détail
            ws.column_dimensions[column_letter].width = 40
        else:
            ws.column_dimensions[column_letter].width = 15

    # Préparer la réponse
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"verification_scans_{fichier_source[:30]}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response


@login_required
def impression_scans(request):
    """
    Vue d'impression avec tableau et détails pour un fichier spécifique
    """
    fichier_source = request.GET.get('fichier')

    if not fichier_source:
        return HttpResponse("Fichier source requis", status=400)

    # Récupérer les entrées du fichier
    # Exclure:
    # - 00BOOK (transferts internes)
    # - 00PERS (comptes de tiers)
    # - 00FND1 SANS "EDF" dans description/info (transferts inter-fichiers)
    entries_base = BaseDiameX.objects.exclude(
        Q(compte_code__icontains='00BOOK') |
        Q(compte_code__icontains='00PERS') |
        (Q(compte_code__icontains='00FND1') & ~Q(description__icontains='EDF') & ~Q(info__icontains='EDF'))
    ).filter(fichier_source=fichier_source)

    # Récupérer le cache de vérification
    cache = ScanVerificationCache.objects.first()
    scans_valides_ids = set(cache.scans_valides_ids) if cache and cache.scans_valides_ids else set()
    liens_morts_ids = set(cache.liens_morts_ids) if cache and cache.liens_morts_ids else set()

    # Groupement par mois
    recap_data = entries_base.values('date_m').annotate(
        total=Count('id'),
        scans_valides=Count(
            Case(
                When(id__in=scans_valides_ids, then=1),
                output_field=IntegerField()
            )
        ) if scans_valides_ids else Count('id', filter=Q(id__in=[])),
        liens_morts=Count(
            Case(
                When(id__in=liens_morts_ids, then=1),
                output_field=IntegerField()
            )
        ) if liens_morts_ids else Count('id', filter=Q(id__in=[])),
        sans_scan=Count('id', filter=Q(scan__isnull=True) | Q(scan=''))
    ).order_by('-date_m')

    # Organiser par mois avec détails
    recap_par_mois = []
    for entry in recap_data:
        mois_key = entry['date_m'].strftime('%Y-%m') if entry['date_m'] else 'Sans date'

        # Récupérer les entrées problématiques pour ce mois
        year, month = None, None
        if mois_key != 'Sans date':
            year, month = map(int, mois_key.split('-'))

        entries_mois = entries_base
        if year and month:
            entries_mois = entries_mois.filter(date_m__year=year, date_m__month=month)
        else:
            entries_mois = entries_mois.filter(date_m__isnull=True)

        # Filtrer pour ne garder que liens morts et sans scan
        entries_problemes = entries_mois.filter(
            Q(id__in=liens_morts_ids) | Q(scan__isnull=True) | Q(scan='')
        ).exclude(id__in=scans_valides_ids).order_by('-date_rec')[:500]  # Limite 500 par mois

        # Enrichir avec statut
        details = []
        for e in entries_problemes:
            if e.id in liens_morts_ids:
                statut = "Lien mort"
                statut_class = "badge-warning"
            elif not e.scan or e.scan.strip() == '':
                statut = "Sans scan"
                statut_class = "badge-error"
            else:
                statut = "Non vérifié"
                statut_class = "badge-ghost"

            details.append({
                'entry': e,
                'statut': statut,
                'statut_class': statut_class
            })

        recap_par_mois.append({
            'mois': mois_key,
            'total': entry['total'],
            'scans_valides': entry['scans_valides'],
            'liens_morts': entry['liens_morts'],
            'sans_scan': entry['sans_scan'],
            'taux_valides': int((entry['scans_valides'] / entry['total'] * 100) if entry['total'] > 0 else 0),
            'details': details
        })

    # Calculer les totaux
    totaux = {
        'total': sum(m['total'] for m in recap_par_mois),
        'scans_valides': sum(m['scans_valides'] for m in recap_par_mois),
        'liens_morts': sum(m['liens_morts'] for m in recap_par_mois),
        'sans_scan': sum(m['sans_scan'] for m in recap_par_mois),
    }
    totaux['taux_valides'] = int((totaux['scans_valides'] / totaux['total'] * 100) if totaux['total'] > 0 else 0)

    context = {
        'fichier_source': fichier_source,
        'recap_par_mois': recap_par_mois,
        'totaux': totaux,
        'date_export': datetime.now(),
    }

    return render(request, 'accounting/impression_scans.html', context)
