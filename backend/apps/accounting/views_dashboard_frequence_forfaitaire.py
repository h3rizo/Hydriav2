"""
Dashboard "Fréquence de paiement des tranches forfaitaires".

Classe chaque point d'eau forfaitaire (regroupé par Région > District > Commune > Ouvrage)
selon deux axes distincts :

- La FACTURATION : le nombre de tranches trimestrielles effectivement facturées
  (PaiementForfaitaire, quelle que soit la date) doit correspondre aux tranches
  attendues sur l'année = 4 (ou moins si la modalité de gestion a démarré en cours
  d'année) MOINS la décote pour coupures d'eau. Un écart est une "Anomalie
  facturation", indépendamment du fait que ce soit payé ou non.

- Le PAIEMENT : une fois la facturation jugée correcte, on vérifie que tout ce qui
  est déjà échu (date_facture <= aujourd'hui) est payé. Un manque sur une tranche
  déjà échue est une "Anomalie paiement". Un manque sur une tranche dont la date
  est future n'est PAS une anomalie (elle n'est simplement pas encore due) : le
  point d'eau est alors classé "En cours".

Si tout est facturé conforme ET intégralement payé, le point d'eau est classé
selon le nombre de paiements distincts utilisés pour régler l'année : 1, 2, 3 ou 4 fois.
"""

from collections import defaultdict
from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from apps.billing.tarifs_forfaitaires import tarif_pour_annee
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.billing.models import PaiementForfaitaire
from apps.accounting.views_dashboard_forfaitaire import (
    FORFAITAIRE_MODALITES_PAIEMENT,
    _apply_forfaitaire_point_eau_geo_filters,
    _forfait_points_queryset,
    _get_forfaitaire_geo_querysets,
    _get_modalite_active_map,
    _get_organisations_filtre_queryset,
    _get_scoped_forfaitaire_ouvrage_ids,
    _resolve_gestionnaire_scope,
    _trimestres_planifies_depuis_date_debut,
    calculer_tranches_attendues,
)
from apps.organization.scoping import get_user_accessible_gestionnaire_ids


BUCKET_LABELS = {
    1: 'Payé en 1 fois',
    2: 'Payé en 2 fois',
    3: 'Payé en 3 fois',
    4: 'Payé en 4 fois',
    'en_cours': 'En cours (à jour)',
    'anomalie_facturation': 'Anomalie facturation',
    'anomalie_paiement': 'Anomalie paiement',
}

BUCKET_LABELS_COURTS = {
    1: '1x',
    2: '2x',
    3: '3x',
    4: '4x',
    'en_cours': 'En cours',
    'anomalie_facturation': 'Anom. factu.',
    'anomalie_paiement': 'Anom. paiement',
}

BUCKET_BADGES = {
    1: 'badge-success',
    2: 'badge-info',
    3: 'badge-warning',
    4: 'badge-primary',
    'en_cours': 'badge-info',
    'anomalie_facturation': 'badge-error',
    'anomalie_paiement': 'badge-secondary',
}

# Ordre d'affichage des cartes de synthèse
BUCKET_ORDER = [1, 2, 3, 4, 'en_cours', 'anomalie_facturation', 'anomalie_paiement']

BUCKET_META = [
    {
        'bucket': bucket,
        'label': BUCKET_LABELS[bucket],
        'label_court': BUCKET_LABELS_COURTS[bucket],
        'badge': BUCKET_BADGES[bucket],
    }
    for bucket in BUCKET_ORDER
]


def _empty_stats():
    return {bucket: 0 for bucket in BUCKET_ORDER}


def _new_node(nom):
    return {
        'nom': nom,
        'stats': _empty_stats(),
        'total': 0,
        'children': {},
        'points_eau': [],
    }


def _add_to_stats(stats, bucket):
    stats[bucket] = stats.get(bucket, 0) + 1


def _pct_int(count, total):
    if not total:
        return 0
    return int(round((count / total) * 100))


def _parse_filters(request):
    filter_annee = request.GET.get('annee', str(datetime.now().year))
    try:
        annee = int(filter_annee)
    except (TypeError, ValueError):
        annee = datetime.now().year

    return {
        'annee': annee,
        'region': request.GET.get('region', ''),
        'district': request.GET.get('district', ''),
        'commune': request.GET.get('commune', ''),
        'ouvrage': request.GET.get('ouvrage', ''),
        'organisation': request.GET.get('organisation', ''),
        'frequence': request.GET.get('frequence', '').strip(),
    }


def _classer_point_eau(pe, annee, today, modalite_active_map, paiements_par_point):
    """
    Calcule tous les indicateurs de facturation/paiement pour un point d'eau et
    retourne (bucket, indicateurs) ou (None, None) si le point d'eau n'est pas
    concerné cette année (aucune tranche attendue, aucune facture).
    """
    modalite_active = modalite_active_map.get(pe.ouvrage_id)
    date_debut_modalite = getattr(modalite_active, 'date_debut', None)
    tranches_dues_modalite = _trimestres_planifies_depuis_date_debut(annee, date_debut_modalite)

    _, trimestres_sans_eau = calculer_tranches_attendues(pe, annee)
    tranches_attendues_facturation = max(0, tranches_dues_modalite - trimestres_sans_eau)

    paiements_point = paiements_par_point.get(pe.id, [])
    tranches_facturees_total = sum(int(p.quantite or 0) for p in paiements_point)

    if tranches_attendues_facturation <= 0 and tranches_facturees_total <= 0:
        return None, None

    tranches_facturees_echues = sum(
        int(p.quantite or 0)
        for p in paiements_point
        if p.date_facture and p.date_facture <= today
    )
    paiements_payes = [p for p in paiements_point if p.paye]
    nb_factures_payees = len(paiements_payes)
    tranches_payees = sum(int(p.quantite or 0) for p in paiements_payes)
    montant_total_paye = sum(float(p.montant or 0) for p in paiements_payes)
    montant_total_facture = sum(float(p.montant or 0) for p in paiements_point)

    if tranches_facturees_total != tranches_attendues_facturation:
        bucket = 'anomalie_facturation'
    elif tranches_payees < tranches_facturees_echues:
        bucket = 'anomalie_paiement'
    elif tranches_payees >= tranches_facturees_total:
        bucket = nb_factures_payees if nb_factures_payees in (1, 2, 3, 4) else 'anomalie_paiement'
    else:
        bucket = 'en_cours'

    indicateurs = {
        'tranches_attendues_facturation': tranches_attendues_facturation,
        'trimestres_sans_eau': trimestres_sans_eau,
        'tranches_facturees_total': tranches_facturees_total,
        'tranches_facturees_echues': tranches_facturees_echues,
        'tranches_payees': tranches_payees,
        'nb_factures_payees': nb_factures_payees,
        'montant_total_paye': montant_total_paye,
        'montant_total_facture': montant_total_facture,
    }
    return bucket, indicateurs


def _build_dataset(request, filters):
    """
    Construit la liste plate des points d'eau classés (avec leur localisation)
    ainsi que les querysets de filtres. Partagé par la page HTML et l'export Excel.
    """
    annee = filters['annee']
    today = timezone.localdate()

    accessible_gestionnaire_ids = get_user_accessible_gestionnaire_ids(request.user)
    gestionnaire_ids, organisation_resolue = _resolve_gestionnaire_scope(
        accessible_gestionnaire_ids,
        filters['organisation'],
    )
    filters['organisation'] = organisation_resolue
    organisations = _get_organisations_filtre_queryset(accessible_gestionnaire_ids)

    ouvrages_org = _get_scoped_forfaitaire_ouvrage_ids(
        request.user,
        annee=annee,
        gestionnaire_ids=gestionnaire_ids,
    )

    regions, districts_all, communes_all = _get_forfaitaire_geo_querysets(ouvrages_org)

    points_eau_qs = _forfait_points_queryset().filter(ouvrage_id__in=ouvrages_org)
    points_eau_qs = _apply_forfaitaire_point_eau_geo_filters(
        points_eau_qs,
        filter_region=filters['region'],
        filter_district=filters['district'],
        filter_commune=filters['commune'],
    )
    if filters['ouvrage']:
        points_eau_qs = points_eau_qs.filter(ouvrage_id=filters['ouvrage'])

    points_eau_qs = points_eau_qs.select_related(
        'ouvrage',
        'parcelle__village__fokontany__commune__district__region',
    ).order_by(
        'parcelle__village__fokontany__commune__district__region__nom_region',
        'parcelle__village__fokontany__commune__district__nom_district',
        'parcelle__village__fokontany__commune__nom_commune',
        'ouvrage__nom_ouvrage',
        'nomabonne',
        'code_point_eau',
    )

    points_eau = list(points_eau_qs)
    point_eau_ids = [pe.id for pe in points_eau]

    modalite_active_map = _get_modalite_active_map(
        list(ouvrages_org),
        annee=annee,
        gestionnaire_ids=gestionnaire_ids,
        modalites_paiement=FORFAITAIRE_MODALITES_PAIEMENT,
    )

    paiements_qs = PaiementForfaitaire.objects.filter(
        point_eau_id__in=point_eau_ids,
        date_facture__year=annee,
        quantite__gt=0,
    ).order_by('point_eau_id', 'date_facture', 'created_at')

    paiements_par_point = defaultdict(list)
    for paiement in paiements_qs:
        paiements_par_point[paiement.point_eau_id].append(paiement)

    lignes = []
    stats_globales = _empty_stats()
    nb_points_eau_total = 0

    for pe in points_eau:
        bucket, indicateurs = _classer_point_eau(
            pe, annee, today, modalite_active_map, paiements_par_point
        )
        if bucket is None:
            continue

        nb_points_eau_total += 1
        _add_to_stats(stats_globales, bucket)

        # --- Localisation ---
        region_nom = 'Sans région'
        district_nom = 'Sans district'
        commune_nom = 'Sans commune'
        if pe.parcelle and pe.parcelle.village:
            village = pe.parcelle.village
            if village.fokontany and village.fokontany.commune:
                commune = village.fokontany.commune
                commune_nom = commune.nom_commune or commune_nom
                if commune.district:
                    district_nom = commune.district.nom_district or district_nom
                    if commune.district.region:
                        region_nom = commune.district.region.nom_region or region_nom

        ouvrage_nom = pe.ouvrage.nom_ouvrage if pe.ouvrage else 'Sans ouvrage'
        ouvrage_id = str(pe.ouvrage_id) if pe.ouvrage_id else 'sans-ouvrage'

        lignes.append({
            'region': region_nom,
            'district': district_nom,
            'commune': commune_nom,
            'ouvrage_id': ouvrage_id,
            'ouvrage_nom': ouvrage_nom,
            'id': str(pe.id),
            'nom': pe.nomabonne or pe.code_point_eau or '-',
            'code': pe.code_point_eau or '-',
            'tarif_trimestriel': float(tarif_pour_annee(pe, annee)),
            'tranches_attendues_facturation': indicateurs['tranches_attendues_facturation'],
            'trimestres_sans_eau': indicateurs['trimestres_sans_eau'],
            'tranches_facturees_total': indicateurs['tranches_facturees_total'],
            'tranches_facturees_echues': indicateurs['tranches_facturees_echues'],
            'tranches_payees': indicateurs['tranches_payees'],
            'nb_factures_payees': indicateurs['nb_factures_payees'],
            'montant_total_paye': indicateurs['montant_total_paye'],
            'montant_total_facture': indicateurs['montant_total_facture'],
            'bucket': bucket,
            'bucket_label': BUCKET_LABELS.get(bucket, str(bucket)),
            'bucket_badge': BUCKET_BADGES.get(bucket, 'badge-neutral'),
        })

    return {
        'lignes': lignes,
        'stats_globales': stats_globales,
        'nb_points_eau_total': nb_points_eau_total,
        'organisations': organisations,
        'regions': regions,
        'districts_all': districts_all,
        'communes_all': communes_all,
    }


def _build_hierarchy(lignes, filter_frequence):
    racine = {}
    nb_points_eau_affiches = 0

    for ligne in lignes:
        bucket = ligne['bucket']
        if filter_frequence and str(bucket) != filter_frequence:
            continue

        nb_points_eau_affiches += 1

        region_node = racine.setdefault(ligne['region'], _new_node(ligne['region']))
        _add_to_stats(region_node['stats'], bucket)
        region_node['total'] += 1

        district_node = region_node['children'].setdefault(
            ligne['district'], _new_node(ligne['district'])
        )
        _add_to_stats(district_node['stats'], bucket)
        district_node['total'] += 1

        commune_node = district_node['children'].setdefault(
            ligne['commune'], _new_node(ligne['commune'])
        )
        _add_to_stats(commune_node['stats'], bucket)
        commune_node['total'] += 1

        ouvrage_node = commune_node['children'].setdefault(
            ligne['ouvrage_id'], _new_node(ligne['ouvrage_nom'])
        )
        _add_to_stats(ouvrage_node['stats'], bucket)
        ouvrage_node['total'] += 1

        ouvrage_node['points_eau'].append(ligne)

    def _finalize_ouvrages(commune_node):
        ouvrages_list = list(commune_node['children'].values())
        ouvrages_list.sort(key=lambda o: o['nom'])
        return ouvrages_list

    def _finalize_communes(district_node):
        communes_list = list(district_node['children'].values())
        for commune_node in communes_list:
            commune_node['ouvrages'] = _finalize_ouvrages(commune_node)
        communes_list.sort(key=lambda c: c['nom'])
        return communes_list

    def _finalize_districts(region_node):
        districts_list = list(region_node['children'].values())
        for district_node in districts_list:
            district_node['communes'] = _finalize_communes(district_node)
        districts_list.sort(key=lambda d: d['nom'])
        return districts_list

    regions_data = list(racine.values())
    for region_node in regions_data:
        region_node['districts'] = _finalize_districts(region_node)
    regions_data.sort(key=lambda r: r['nom'])

    return regions_data, nb_points_eau_affiches


@login_required
def dashboard_frequence_paiements_forfaitaire(request):
    """
    Page de classification des points d'eau forfaitaires selon la conformité de
    leur facturation et de leur paiement pour l'année sélectionnée.
    """
    filters = _parse_filters(request)
    dataset = _build_dataset(request, filters)

    regions_data, nb_points_eau_affiches = _build_hierarchy(
        dataset['lignes'], filters['frequence']
    )

    nb_points_eau_total = dataset['nb_points_eau_total']
    stats_globales = dataset['stats_globales']

    stats_cartes = [
        {
            'bucket': bucket,
            'label': BUCKET_LABELS.get(bucket, str(bucket)),
            'badge': BUCKET_BADGES.get(bucket, 'badge-neutral'),
            'count': stats_globales.get(bucket, 0),
            'pct': _pct_int(stats_globales.get(bucket, 0), nb_points_eau_total),
        }
        for bucket in BUCKET_ORDER
    ]

    annees_disponibles = list(range(datetime.now().year, datetime.now().year - 5, -1))

    context = {
        'annee': filters['annee'],
        'annees_disponibles': annees_disponibles,
        'regions': dataset['regions'],
        'districts_all': dataset['districts_all'],
        'communes_all': dataset['communes_all'],
        'organisations': dataset['organisations'],
        'filter_region': filters['region'],
        'filter_district': filters['district'],
        'filter_commune': filters['commune'],
        'filter_ouvrage': filters['ouvrage'],
        'filter_organisation': filters['organisation'],
        'filter_frequence': filters['frequence'],
        'stats_cartes': stats_cartes,
        'stats_globales': stats_globales,
        'nb_points_eau_total': nb_points_eau_total,
        'nb_points_eau_affiches': nb_points_eau_affiches,
        'regions_data': regions_data,
        'bucket_order': BUCKET_ORDER,
        'bucket_meta': BUCKET_META,
    }

    return render(
        request,
        'accounting/dashboard_frequence_paiements_forfaitaire.html',
        context,
    )


@login_required
def export_frequence_paiements_forfaitaire_excel(request):
    """
    Export Excel de l'analyse de fréquence de paiement des tranches forfaitaires,
    avec les mêmes filtres que la page (année, organisation, région, district,
    commune, ouvrage, fréquence).
    """
    filters = _parse_filters(request)
    dataset = _build_dataset(request, filters)
    regions_data, nb_points_eau_affiches = _build_hierarchy(
        dataset['lignes'], filters['frequence']
    )
    nb_points_eau_total = dataset['nb_points_eau_total']
    stats_globales = dataset['stats_globales']

    wb = Workbook()

    # ---- Feuille Synthèse ----
    ws_synthese = wb.active
    ws_synthese.title = "Synthèse"

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws_synthese.merge_cells('A1:D1')
    cell = ws_synthese['A1']
    cell.value = "Fréquence de paiement des tranches forfaitaires"
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal='center')

    ws_synthese.merge_cells('A2:D2')
    cell = ws_synthese['A2']
    cell.value = f"Année {filters['annee']} — {nb_points_eau_total} point(s) d'eau concerné(s)"
    cell.font = Font(size=10, italic=True)
    cell.alignment = Alignment(horizontal='center')

    row = 4
    headers = ['Statut', 'Nombre de points d\'eau', 'Pourcentage']
    for col_num, header in enumerate(headers, 1):
        c = ws_synthese.cell(row=row, column=col_num, value=header)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = Alignment(horizontal='center')
    row += 1

    for bucket in BUCKET_ORDER:
        count = stats_globales.get(bucket, 0)
        pct = _pct_int(count, nb_points_eau_total)
        ws_synthese.cell(row=row, column=1, value=BUCKET_LABELS.get(bucket, str(bucket))).border = border
        ws_synthese.cell(row=row, column=2, value=count).border = border
        c = ws_synthese.cell(row=row, column=3, value=pct / 100)
        c.number_format = '0%'
        c.border = border
        row += 1

    for col_num, width in zip(range(1, 4), (28, 22, 16)):
        ws_synthese.column_dimensions[get_column_letter(col_num)].width = width

    # ---- Feuille Détail ----
    ws = wb.create_sheet("Détail")

    headers = [
        'Région', 'District', 'Commune', 'Ouvrage', 'Point d\'eau',
        'Tarif / trimestre (Ar)', 'Trimestres sans eau (décote)',
        'Tranches attendues (facturation)', 'Tranches facturées',
        'Tranches facturées échues', 'Tranches payées', 'Nb paiements',
        'Statut', 'Montant facturé (Ar)', 'Montant payé (Ar)',
    ]
    for col_num, header in enumerate(headers, 1):
        c = ws.cell(row=1, column=col_num, value=header)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = Alignment(horizontal='center', wrap_text=True)

    row = 2
    for region_node in regions_data:
        for district_node in region_node['districts']:
            for commune_node in district_node['communes']:
                for ouvrage_node in commune_node['ouvrages']:
                    for ligne in ouvrage_node['points_eau']:
                        values = [
                            ligne['region'],
                            ligne['district'],
                            ligne['commune'],
                            ligne['ouvrage_nom'],
                            ligne['nom'],
                            ligne['tarif_trimestriel'],
                            ligne['trimestres_sans_eau'],
                            ligne['tranches_attendues_facturation'],
                            ligne['tranches_facturees_total'],
                            ligne['tranches_facturees_echues'],
                            ligne['tranches_payees'],
                            ligne['nb_factures_payees'],
                            ligne['bucket_label'],
                            ligne['montant_total_facture'],
                            ligne['montant_total_paye'],
                        ]
                        for col_num, value in enumerate(values, 1):
                            ws.cell(row=row, column=col_num, value=value).border = border
                        row += 1

    widths = (16, 16, 18, 24, 26, 16, 14, 16, 14, 16, 14, 12, 20, 16, 16)
    for col_num, width in zip(range(1, len(widths) + 1), widths):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    ws.freeze_panes = 'A2'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"frequence_paiements_forfaitaires_{filters['annee']}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
