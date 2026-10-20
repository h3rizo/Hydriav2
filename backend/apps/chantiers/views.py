from django.shortcuts import render, get_object_or_404
from django.http import JsonResponse
from django.db.models import Q, Count
from django.db import transaction
from django.db import connection
from django.views.decorators.http import require_http_methods
from django.views.decorators.csrf import csrf_exempt
import json
from datetime import datetime
from openpyxl import load_workbook

from .models import DemandeVillage, ValidationSource, CaracteristiqueChantier, Tacheron
from .models_workflow import (
    BaremeMainOeuvre,
    ChantierBesoinLigne,
    ChantierSelectionEquipement,
    ChantierWorkflow,
    CiterneDimensionnement,
    CiterneParametrageReference,
    PosteMainOeuvre,
    WorkflowModeleEquipement,
    WorkflowModeleEquipementLigne,
    generer_besoins_chantier,
)
from apps.geography.models import Region, District, Commune, Fokontany, Village
from apps.organization.models import TypeOuvrage, Ouvrage
from apps.employees.models import Employe
from apps.filemanager.models import FileUpload
from apps.infrastructure.models import TypeEquipement
from apps.billing.models import TypePointEau


# ==================== DEMANDES VILLAGES ====================

def demandes_list(request):
    """Page consolidée des demandes de villages avec filtres et stats"""

    # Paramètres de filtrage
    search = request.GET.get('search', '')
    region_filter = request.GET.get('region', 'all')
    district_filter = request.GET.get('district', 'all')
    commune_filter = request.GET.get('commune', 'all')
    fokontany_filter = request.GET.get('fokontany', 'all')
    statut_filter = request.GET.get('statut', 'all')
    type_travaux_filter = request.GET.get('type_travaux', 'all')

    # Requête SQL brute pour performance (car relations complexes via properties)
    cursor = connection.cursor()

    query = """
        SELECT
            d.id,
            d.statut,
            d.type_travaux,
            d.detail_travaux,
            d.type_main_oeuvre,
            d.semestre_prevision,
            v.nom_village,
            f.nom_fokontany,
            c.nom_commune,
            di.nom_district,
            r.nom_region,
            typ.nom_type_ouvrage,
            v.id as village_id,
            COUNT(vs.id) as nb_validations
        FROM demandes_villages d
        JOIN villages v ON d.village_id = v.id
        JOIN fokontanys f ON v.fokontany_id = f.id
        JOIN communes c ON f.commune_id = c.id
        JOIN districts di ON c.district_id = di.id
        JOIN regions r ON di.region_id = r.id
        JOIN type_ouvrages typ ON d.type_ouvrage_id = typ.id
        LEFT JOIN validations_sources vs ON v.id = vs.village_id
        WHERE 1=1
    """

    params = []

    # Recherche
    if search:
        query += " AND (v.nom_village ILIKE %s OR d.detail_travaux ILIKE %s)"
        params.extend([f'%{search}%', f'%{search}%'])

    # Filtres géographiques
    if region_filter != 'all':
        query += " AND r.id = %s"
        params.append(region_filter)
    if district_filter != 'all':
        query += " AND di.id = %s"
        params.append(district_filter)
    if commune_filter != 'all':
        query += " AND c.id = %s"
        params.append(commune_filter)
    if fokontany_filter != 'all':
        query += " AND f.id = %s"
        params.append(fokontany_filter)

    # Filtres métier
    if statut_filter != 'all':
        query += " AND d.statut = %s"
        params.append(statut_filter)
    if type_travaux_filter != 'all':
        query += " AND d.type_travaux = %s"
        params.append(type_travaux_filter)

    query += " GROUP BY d.id, v.id, v.nom_village, f.nom_fokontany, c.nom_commune, di.nom_district, r.nom_region, typ.nom_type_ouvrage"
    query += " ORDER BY d.created_at DESC"

    cursor.execute(query, params)
    demandes = list(cursor.fetchall())

    # Desserte de chaque village par un ouvrage géré / non géré par l'organisation de l'utilisateur
    from django.utils import timezone
    from apps.billing.models import ModaliteGestion, PointEau
    from apps.organization.scoping import get_user_organisation

    organisation = get_user_organisation(request.user)
    village_ids = {row[12] for row in demandes}
    ouvrages_par_village = {}
    for village_id, ouvrage_id in PointEau.objects.filter(
        parcelle__village_id__in=village_ids, datebranchement__isnull=False, ouvrage_id__isnull=False,
    ).values_list('parcelle__village_id', 'ouvrage_id').distinct():
        ouvrages_par_village.setdefault(str(village_id), set()).add(ouvrage_id)

    ouvrages_geres = set()
    if organisation and ouvrages_par_village:
        today = timezone.localdate()
        ouvrages_geres = set(
            ModaliteGestion.objects.filter(
                ouvrage_id__in={o for ids in ouvrages_par_village.values() for o in ids},
                gestionnaire_id=organisation.id, statut='actif', date_debut__lte=today,
            ).filter(Q(date_fin__isnull=True) | Q(date_fin__gte=today))
            .exclude(type_gestion='pas_de_gestion')
            .values_list('ouvrage_id', flat=True)
        )

    # Statut fonctionnel des ouvrages desservant chaque village (défini dans Ouvrage,
    # voir apps.organization.services_fonctionnalite)
    from apps.organization.services_fonctionnalite import get_statuts_fonctionnel_bulk

    tous_ouvrage_ids = {o for ids in ouvrages_par_village.values() for o in ids}
    statuts_fonctionnel_ouvrages = get_statuts_fonctionnel_bulk(tous_ouvrage_ids)

    def desserte_info(village_id):
        ouvrages = ouvrages_par_village.get(str(village_id), set())
        statuts_ouvrages = [statuts_fonctionnel_ouvrages[o] for o in ouvrages if o in statuts_fonctionnel_ouvrages]
        non_fonctionnels = [s for s in statuts_ouvrages if not s['fonctionnel']]
        return {
            'dessert': bool(ouvrages),
            'gere': bool(ouvrages & ouvrages_geres),
            'non_gere': bool(ouvrages - ouvrages_geres),
            'organisation_connue': organisation is not None,
            'ouvrage_non_fonctionnel': bool(non_fonctionnels),
            'ouvrage_non_fonctionnel_source': non_fonctionnels[0]['source'] if non_fonctionnels else None,
        }

    # Regroupement : une ligne par VILLAGE (un village avec plusieurs demandes n'apparaît qu'une fois)
    def ajouter(liste, valeur):
        if valeur not in (None, '') and valeur not in liste:
            liste.append(valeur)

    villages = {}
    for row in demandes:
        (demande_id, statut, type_travaux, _detail, _main_oeuvre, semestre, nom_village, fokontany,
         commune, district, region, type_ouvrage, village_id, nb_validations) = row
        v = villages.get(village_id)
        if v is None:
            v = villages[village_id] = {
                'village_id': village_id, 'village': nom_village, 'fokontany': fokontany,
                'commune': commune, 'district': district, 'region': region,
                'types_ouvrage': [], 'types_travaux': [], 'statuts': [], 'semestres': [],
                'demande_ids': [], 'nb_validations': nb_validations or 0,
                'desserte': desserte_info(village_id),
            }
        v['demande_ids'].append(demande_id)
        ajouter(v['types_ouvrage'], type_ouvrage)
        ajouter(v['types_travaux'], type_travaux)
        ajouter(v['statuts'], statut)
        ajouter(v['semestres'], semestre)
    for v in villages.values():
        v['nb_demandes'] = len(v['demande_ids'])
    lignes = list(villages.values())

    # Statistiques : en VILLAGES avec demande, sur les mêmes filtres que le tableau.
    # Un village = une seule catégorie de desserte : s'il a un ouvrage géré, on ignore les autres.
    stats = {
        'total': len(lignes),
        'nb_demandes': len(demandes),
        'en_demande': sum(1 for v in lignes if 'demande' in v['statuts']),
        'source_validee': sum(1 for v in lignes if 'source_validee' in v['statuts']),
        'en_preparation': sum(1 for v in lignes if 'preparation_en_cours' in v['statuts']),
        'acheves': sum(1 for v in lignes if 'acheve' in v['statuts']),
        'avec_validations': sum(1 for v in lignes if v['nb_validations'] > 0),
        'total_validations': sum(v['nb_validations'] for v in lignes),
        'avec_ouvrage_gere': sum(1 for v in lignes if v['desserte']['gere']),
        'avec_ouvrage_non_gere': sum(
            1 for v in lignes if not v['desserte']['gere'] and v['desserte']['non_gere']),
        'sans_ouvrage': sum(1 for v in lignes if not v['desserte']['dessert']),
    }
    demandes = lignes

    # Listes pour les dropdowns
    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.all().order_by('nom_district')
    # Communes proposées : limitées au district (ou à la région) choisi
    communes = Commune.objects.all()
    if district_filter != 'all':
        communes = communes.filter(district_id=district_filter)
    elif region_filter != 'all':
        communes = communes.filter(district__region_id=region_filter)
    communes = communes.order_by('nom_commune')
    fokontanys = Fokontany.objects.all().order_by('nom_fokontany')
    types_ouvrages = TypeOuvrage.objects.all().order_by('nom_type_ouvrage')

    context = {
        'demandes': demandes,
        'stats': stats,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'fokontanys': fokontanys,
        'types_ouvrages': types_ouvrages,
        'search': search,
        'region_filter': region_filter,
        'district_filter': district_filter,
        'commune_filter': commune_filter,
        'fokontany_filter': fokontany_filter,
        'statut_filter': statut_filter,
        'type_travaux_filter': type_travaux_filter,
        'statut_choices': DemandeVillage.STATUT_CHOICES,
        'type_travaux_choices': DemandeVillage.TYPE_TRAVAUX_CHOICES,
    }

    return render(request, 'chantiers/demandes_list.html', context)


@require_http_methods(["GET"])
def api_villages_search(request):
    """Autocomplétion village : nom + fokontany + commune pour lever les homonymes."""
    q = (request.GET.get('q') or '').strip()
    commune = request.GET.get('commune') or ''

    qs = Village.objects.select_related('fokontany__commune')
    if commune and commune != 'all':
        qs = qs.filter(fokontany__commune_id=commune)
    if q:
        qs = qs.filter(
            Q(nom_village__icontains=q)
            | Q(fokontany__nom_fokontany__icontains=q)
            | Q(fokontany__commune__nom_commune__icontains=q)
        )
    qs = qs.order_by('nom_village', 'fokontany__nom_fokontany')[:30]

    results = []
    for v in qs:
        fok = v.fokontany
        com = fok.commune if fok else None
        results.append({
            'id': str(v.id),
            'village': v.nom_village,
            'fokontany': fok.nom_fokontany if fok else '',
            'commune': com.nom_commune if com else '',
        })
    return JsonResponse({'success': True, 'results': results})


@csrf_exempt
@require_http_methods(["POST"])
def demande_create(request):
    """Créer une nouvelle demande"""
    try:
        data = json.loads(request.body)
        demande = DemandeVillage.objects.create(
            village_id=data['village_id'],
            statut=data.get('statut', 'demande'),
            type_travaux=data['type_travaux'],
            type_ouvrage_id=data['type_ouvrage_id'],
            detail_travaux=data['detail_travaux'],
            type_main_oeuvre=data['type_main_oeuvre'],
            semestre_prevision=data['semestre_prevision']
        )
        return JsonResponse({'success': True, 'id': demande.id})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@require_http_methods(["GET"])
def demande_detail(request, demande_id):
    """Données d'une demande pour pré-remplir le modal de modification."""
    d = get_object_or_404(
        DemandeVillage.objects.select_related('village__fokontany__commune'), id=demande_id)
    fok = d.village.fokontany
    return JsonResponse({
        'success': True,
        'demande': {
            'id': str(d.id),
            'village_id': str(d.village_id),
            'village_label': f"{d.village.nom_village} ({fok.nom_fokontany} - {fok.commune.nom_commune})",
            'statut': d.statut,
            'type_travaux': d.type_travaux,
            'type_ouvrage_id': str(d.type_ouvrage_id),
            'detail_travaux': d.detail_travaux,
            'type_main_oeuvre': d.type_main_oeuvre,
            'semestre_prevision': d.semestre_prevision,
        },
    })


@csrf_exempt
@require_http_methods(["POST"])
def demande_update(request, demande_id):
    """Mettre à jour une demande"""
    try:
        demande = get_object_or_404(DemandeVillage, id=demande_id)
        data = json.loads(request.body)

        for field in ['village_id', 'statut', 'type_travaux', 'type_ouvrage_id',
                      'detail_travaux', 'type_main_oeuvre', 'semestre_prevision']:
            if field in data:
                setattr(demande, field, data[field])

        demande.save()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def demande_delete(request, demande_id):
    """Supprimer une demande"""
    try:
        demande = get_object_or_404(DemandeVillage, id=demande_id)
        demande.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


# ==================== VALIDATIONS SOURCES ====================

def validations_list(request):
    """Page consolidée des validations de sources"""

    # Paramètres de filtrage
    search = request.GET.get('search', '')
    region_filter = request.GET.get('region', 'all')
    district_filter = request.GET.get('district', 'all')
    commune_filter = request.GET.get('commune', 'all')
    fokontany_filter = request.GET.get('fokontany', 'all')
    village_filter = request.GET.get('village', 'all')
    donation_filter = request.GET.get('donation', 'all')

    cursor = connection.cursor()

    query = """
        SELECT
            vs.id,
            vs.nom_source,
            vs.debit_source,
            vs.date_mesure_debit,
            vs.latitude,
            vs.longitude,
            vs.distance_source_km,
            vs.denivele_m,
            vs.donation_source,
            vs.proposition,
            vs.nombre_sources_visitees,
            v.nom_village,
            f.nom_fokontany,
            c.nom_commune,
            di.nom_district,
            r.nom_region,
            e.surnom,
            e.prenom,
            e.nom as emp_nom,
            vs.fiche_validation_id,
            fiche.supabase_path as fiche_path
        FROM validations_sources vs
        JOIN villages v ON vs.village_id = v.id
        JOIN fokontanys f ON v.fokontany_id = f.id
        JOIN communes c ON f.commune_id = c.id
        JOIN districts di ON c.district_id = di.id
        JOIN regions r ON di.region_id = r.id
        JOIN employes e ON vs.visiteur_id = e.id
        LEFT JOIN filemanager_fileupload fiche ON vs.fiche_validation_id = fiche.id
        WHERE 1=1
    """

    params = []

    if search:
        query += " AND (v.nom_village ILIKE %s OR vs.nom_source ILIKE %s)"
        params.extend([f'%{search}%', f'%{search}%'])

    if region_filter != 'all':
        query += " AND r.id = %s"
        params.append(region_filter)
    if district_filter != 'all':
        query += " AND di.id = %s"
        params.append(district_filter)
    if commune_filter != 'all':
        query += " AND c.id = %s"
        params.append(commune_filter)
    if fokontany_filter != 'all':
        query += " AND f.id = %s"
        params.append(fokontany_filter)
    if village_filter != 'all':
        query += " AND v.id = %s"
        params.append(village_filter)
    if donation_filter != 'all':
        query += " AND vs.donation_source = %s"
        params.append(donation_filter)

    query += " ORDER BY vs.date_mesure_debit DESC"

    cursor.execute(query, params)
    validations = cursor.fetchall()

    # Statistiques
    stats_query = """
        SELECT
            COUNT(*) as total,
            AVG(vs.debit_source) as debit_moyen,
            SUM(CASE WHEN vs.donation_source = 'realisee' THEN 1 ELSE 0 END) as donations_realisees,
            SUM(CASE WHEN vs.donation_source = 'en_cours' THEN 1 ELSE 0 END) as donations_en_cours,
            SUM(vs.nombre_sources_visitees) as total_sources_visitees
        FROM validations_sources vs
        WHERE 1=1
    """

    cursor.execute(stats_query)
    stats_row = cursor.fetchone()

    stats = {
        'total': stats_row[0] or 0,
        'debit_moyen': float(stats_row[1]) if stats_row[1] else 0,
        'donations_realisees': stats_row[2] or 0,
        'donations_en_cours': stats_row[3] or 0,
        'total_sources_visitees': stats_row[4] or 0,
    }

    # Listes pour dropdowns
    regions = Region.objects.all().order_by('nom_region')
    districts = District.objects.all().order_by('nom_district')
    communes = Commune.objects.all().order_by('nom_commune')
    fokontanys = Fokontany.objects.all().order_by('nom_fokontany')
    villages = Village.objects.all().order_by('nom_village')
    employes = Employe.objects.filter(date_debauchage__isnull=True).order_by('nom', 'prenom')

    context = {
        'validations': validations,
        'stats': stats,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'fokontanys': fokontanys,
        'villages': villages,
        'employes': employes,
        'search': search,
        'region_filter': region_filter,
        'district_filter': district_filter,
        'commune_filter': commune_filter,
        'fokontany_filter': fokontany_filter,
        'donation_filter': donation_filter,
        'donation_choices': ValidationSource.DONATION_CHOICES,
    }

    return render(request, 'chantiers/validations_list.html', context)


@csrf_exempt
@require_http_methods(["POST"])
def validation_create(request):
    """Créer une nouvelle validation de source"""
    try:
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST

        # Gérer l'upload de photo
        photo_id = None
        if 'photo' in request.FILES:
            # Upload via FileUpload model
            photo_file = request.FILES['photo']
            file_upload = FileUpload.objects.create(
                file=photo_file,
                original_filename=photo_file.name,
                file_type='image',
                uploaded_by=request.user if request.user.is_authenticated else None
            )
            photo_id = file_upload.id

        # Gérer l'upload de fiche de validation
        fiche_id = None
        if 'fiche_validation' in request.FILES:
            fiche_file = request.FILES['fiche_validation']
            file_type = 'pdf' if fiche_file.content_type == 'application/pdf' else 'image'
            file_upload_fiche = FileUpload.objects.create(
                file=fiche_file,
                original_filename=fiche_file.name,
                file_type=file_type,
                uploaded_by=request.user if request.user.is_authenticated else None
            )
            fiche_id = file_upload_fiche.id

        validation = ValidationSource.objects.create(
            village_id=data['village_id'],
            nom_source=data['nom_source'],
            debit_source=data['debit_source'],
            date_mesure_debit=data['date_mesure_debit'],
            latitude=data.get('latitude'),
            longitude=data.get('longitude'),
            photo_source_id=photo_id,
            fiche_validation_id=fiche_id,
            distance_source_km=data['distance_source_km'],
            denivele_m=data['denivele_m'],
            commentaires=data.get('commentaires', ''),
            donation_source=data.get('donation_source', 'en_cours'),
            proposition=data.get('proposition') if data.get('proposition') else None,
            nombre_sources_visitees=data.get('nombre_sources_visitees', 1),
            visiteur_id=data['visiteur_id']
        )
        return JsonResponse({'success': True, 'id': validation.id})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def validation_update(request, validation_id):
    """Mettre à jour une validation"""
    try:
        validation = get_object_or_404(ValidationSource, id=validation_id)
        data = json.loads(request.body) if request.content_type == 'application/json' else request.POST

        # Gérer l'upload de fiche de validation
        if 'fiche_validation' in request.FILES:
            fiche_file = request.FILES['fiche_validation']
            file_type = 'pdf' if fiche_file.content_type == 'application/pdf' else 'image'
            file_upload_fiche = FileUpload.objects.create(
                file=fiche_file,
                original_filename=fiche_file.name,
                file_type=file_type,
                uploaded_by=request.user if request.user.is_authenticated else None
            )
            validation.fiche_validation_id = file_upload_fiche.id

        for field in ['village_id', 'nom_source', 'debit_source', 'date_mesure_debit',
                      'latitude', 'longitude', 'distance_source_km', 'denivele_m',
                      'commentaires', 'donation_source', 'proposition', 'nombre_sources_visitees', 'visiteur_id']:
            if field in data:
                value = data[field] if data[field] else None
                setattr(validation, field, value)

        validation.save()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def validation_delete(request, validation_id):
    """Supprimer une validation"""
    try:
        validation = get_object_or_404(ValidationSource, id=validation_id)
        validation.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def validation_clear_all(request):
    """Vider complètement la table des validations de sources"""
    try:
        count = ValidationSource.objects.count()
        ValidationSource.objects.all().delete()
        return JsonResponse({
            'success': True,
            'count': count,
            'message': f'{count} validation(s) supprimée(s)'
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


def validation_get(request, validation_id):
    """Récupérer les données d'une validation pour édition"""
    try:
        validation = get_object_or_404(ValidationSource, id=validation_id)
        data = {
            'id': str(validation.id),
            'village_id': str(validation.village.id),
            'nom_source': validation.nom_source,
            'debit_source': float(validation.debit_source),
            'date_mesure_debit': validation.date_mesure_debit.strftime('%Y-%m-%d'),
            'distance_source_km': float(validation.distance_source_km) if validation.distance_source_km else '',
            'denivele_m': float(validation.denivele_m) if validation.denivele_m else '',
            'visiteur_id': str(validation.visiteur.id),
            'proposition': validation.proposition if validation.proposition else '',
            'nombre_sources_visitees': validation.nombre_sources_visitees,
            'donation_source': validation.donation_source,
            'commentaires': validation.commentaires,
            'latitude': float(validation.latitude) if validation.latitude else '',
            'longitude': float(validation.longitude) if validation.longitude else '',
        }
        if validation.fiche_validation:
            data['fiche_validation_url'] = validation.fiche_validation.supabase_path
        return JsonResponse(data)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)


# ==================== CHANTIERS ====================

def chantiers_list(request):
    """Page consolidée des chantiers"""
    search = request.GET.get('search', '')
    employe_filter = request.GET.get('employe', 'all')
    tacheron_filter = request.GET.get('tacheron', 'all')
    etat_filter = request.GET.get('etat', 'all')

    # Requête SQL pour récupérer les chantiers avec relations
    cursor = connection.cursor()

    query = """
        SELECT
            c.id,
            o.nom_ouvrage,
            o.id as ouvrage_id,
            e.nom,
            e.prenom,
            t.nom as tacheron_nom,
            c.longueur_adduction_m,
            c.date_fin_chantier,
            c.cout_estimatif_materiaux,
            c.cout_estimatif_main_oeuvre,
            c.pourcentage_apports_locaux,
            c.date_ag_explication,
            c.date_reception_definitive
        FROM caracteristiques_chantiers c
        JOIN ouvrages o ON c.ouvrage_id = o.id
        JOIN employes e ON c.employe_id = e.id
        LEFT JOIN tacherons t ON c.tacheron_id = t.id
        WHERE 1=1
    """

    params = []

    if search:
        query += " AND (o.nom_ouvrage ILIKE %s OR e.nom ILIKE %s OR e.prenom ILIKE %s OR t.nom ILIKE %s)"
        params.extend([f'%{search}%', f'%{search}%', f'%{search}%', f'%{search}%'])

    if employe_filter != 'all':
        query += " AND e.id = %s"
        params.append(employe_filter)

    if tacheron_filter != 'all':
        query += " AND t.id = %s"
        params.append(tacheron_filter)

    query += " ORDER BY c.created_at DESC"

    cursor.execute(query, params)
    chantiers_raw = cursor.fetchall()

    # Calculer les états et filtrer
    chantiers = []
    for row in chantiers_raw:
        # Déterminer l'état (logique simplifiée basée sur dates)
        etat = 'non_demarre'
        if row[12]:  # date_reception_definitive
            etat = 'reception_definitive'
        elif row[11]:  # date_ag_explication
            etat = 'en_cours'

        if etat_filter == 'all' or etat == etat_filter:
            chantiers.append(row + (etat,))

    # Statistiques
    stats_query = """
        SELECT
            COUNT(*) as total,
            SUM(CASE WHEN c.date_reception_definitive IS NOT NULL THEN 1 ELSE 0 END) as acheves,
            SUM(CASE WHEN c.date_ag_explication IS NOT NULL AND c.date_reception_definitive IS NULL THEN 1 ELSE 0 END) as en_cours,
            AVG(c.cout_estimatif_materiaux + c.cout_estimatif_main_oeuvre) as cout_moyen
        FROM caracteristiques_chantiers c
    """
    cursor.execute(stats_query)
    stats_row = cursor.fetchone()

    stats = {
        'total': stats_row[0] or 0,
        'acheves': stats_row[1] or 0,
        'en_cours': stats_row[2] or 0,
        'cout_moyen': stats_row[3] or 0,
    }

    # Listes pour dropdowns
    employes = Employe.objects.filter(date_debauchage__isnull=True).order_by('nom')
    tacherons = Tacheron.objects.all().order_by('nom')

    context = {
        'chantiers': chantiers,
        'stats': stats,
        'employes': employes,
        'tacherons': tacherons,
        'search': search,
        'employe_filter': employe_filter,
        'tacheron_filter': tacheron_filter,
        'etat_filter': etat_filter,
    }

    return render(request, 'chantiers/chantiers_list.html', context)


def chantiers_kanban(request):
    """Vue Kanban pour les chantiers"""
    # Note: relation ouvrage -> point d'eau -> parcelle -> village
    # Pour simplifier, on charge juste ouvrage et employe
    # L'accès au village se fera via les properties si nécessaire
    chantiers = CaracteristiqueChantier.objects.select_related(
        'ouvrage', 'ouvrage__type_ouvrage', 'employe'
    ).all()

    # Grouper par état
    grouped = {
        'non_demarre': [],
        'ag_explication': [],
        'animation_suspendue': [],
        'animation_reprise': [],
        'preparation_chantier': [],
        'chantier_ouvert': [],
        'chantier_suspendu': [],
        'chantier_rouvert': [],
        'reception_provisoire': [],
        'reception_definitive': [],
    }

    for chantier in chantiers:
        etat = chantier.etat_actuel
        if etat in grouped:
            grouped[etat].append(chantier)

    return render(request, 'chantiers/chantiers_kanban.html', {'grouped': grouped})


@csrf_exempt
@require_http_methods(["POST"])
def chantier_update_etat(request, chantier_id):
    """Mettre à jour l'état d'un chantier (utilisé par le Kanban)"""
    try:
        chantier = get_object_or_404(CaracteristiqueChantier, id=chantier_id)
        data = json.loads(request.body)

        nouvel_etat = data.get('nouvel_etat')
        date_str = data.get('date')

        # Mapper l'état à la date correspondante
        etat_date_map = {
            'ag_explication': 'date_ag_explication',
            'animation_suspendue': 'date_suspension_animation',
            'animation_reprise': 'date_reprise_animation',
            'preparation_chantier': 'date_preparation_chantier',
            'chantier_ouvert': 'date_ouverture_chantier',
            'chantier_suspendu': 'date_suspension_chantier',
            'chantier_rouvert': 'date_reouverture_chantier',
            'reception_provisoire': 'date_reception_provisoire',
            'reception_definitive': 'date_reception_definitive',
        }

        if nouvel_etat in etat_date_map:
            field_name = etat_date_map[nouvel_etat]
            setattr(chantier, field_name, datetime.strptime(date_str, '%Y-%m-%d').date())
            chantier.save()
            return JsonResponse({'success': True})
        else:
            return JsonResponse({'success': False, 'error': 'État inconnu'}, status=400)

    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


# ==================== IMPORT/EXPORT EXCEL ====================

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def export_demandes_excel(request):
    """Exporter les demandes en Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Demandes Villages"

    # En-têtes
    headers = ['ID', 'Village', 'Fokontany', 'Commune', 'District', 'Région', 
               'Type Ouvrage', 'Type Travaux', 'Statut', 'Main d\'œuvre', 
               'Semestre', 'Détail Travaux']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    # Données
    cursor = connection.cursor()
    cursor.execute("""
        SELECT d.id, v.nom_village, f.nom_fokontany, c.nom_commune, di.nom_district, r.nom_region,
               typ.nom_type_ouvrage, d.type_travaux, d.statut, d.type_main_oeuvre,
               d.semestre_prevision, d.detail_travaux
        FROM demandes_villages d
        JOIN villages v ON d.village_id = v.id
        JOIN fokontanys f ON v.fokontany_id = f.id
        JOIN communes c ON f.commune_id = c.id
        JOIN districts di ON c.district_id = di.id
        JOIN regions r ON di.region_id = r.id
        JOIN type_ouvrages typ ON d.type_ouvrage_id = typ.id
        ORDER BY d.created_at DESC
    """)

    for row_idx, row in enumerate(cursor.fetchall(), 2):
        for col_idx, value in enumerate(row, 1):
            # Convert UUID to string for Excel compatibility
            if hasattr(value, 'hex'):  # Check if it's a UUID
                value = str(value)
            ws.cell(row_idx, col_idx, value)

    # Ajuster largeur colonnes
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=demandes_villages.xlsx'
    wb.save(response)
    return response


def export_validations_excel(request):
    """Exporter les validations en Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Validations Sources"

    headers = ['ID', 'Village', 'Source', 'Débit (L/s)', 'Date Visite', 'Distance (km)',
               'Dénivelé (m)', 'Latitude', 'Longitude', 'Donation', 'Proposition',
               'Nb Sources Visitées', 'Visiteur']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    cursor = connection.cursor()
    cursor.execute("""
        SELECT vs.id, v.nom_village, vs.nom_source, vs.debit_source, vs.date_mesure_debit,
               vs.distance_source_km, vs.denivele_m, vs.latitude, vs.longitude,
               vs.donation_source, vs.proposition, vs.nombre_sources_visitees,
               COALESCE(e.surnom, e.prenom || ' ' || e.nom)
        FROM validations_sources vs
        JOIN villages v ON vs.village_id = v.id
        JOIN employes e ON vs.visiteur_id = e.id
        ORDER BY vs.date_mesure_debit DESC
    """)

    for row_idx, row in enumerate(cursor.fetchall(), 2):
        for col_idx, value in enumerate(row, 1):
            # Convert UUID to string for Excel compatibility
            if hasattr(value, 'hex'):  # Check if it's a UUID
                value = str(value)
            ws.cell(row_idx, col_idx, value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=validations_sources.xlsx'
    wb.save(response)
    return response


def export_chantiers_excel(request):
    """Exporter les chantiers en Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chantiers"

    headers = ['ID', 'Ouvrage', 'Type Ouvrage', 'Employé', 'Longueur (m)', 
               'Tâcheron', 'Date AG', 'Date Préparation', 'Date Ouverture', 
               'Date Réception Provisoire', 'Date Réception Définitive', 
               'Coût Matériaux', 'Coût Main d\'œuvre', '% Apports Locaux']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    chantiers = CaracteristiqueChantier.objects.select_related(
        'ouvrage', 'ouvrage__type_ouvrage', 'employe'
    ).all()
    
    for row_idx, ch in enumerate(chantiers, 2):
        ws.cell(row_idx, 1, str(ch.id))  # Convert UUID to string
        ws.cell(row_idx, 2, ch.ouvrage.nom_ouvrage)
        ws.cell(row_idx, 3, ch.ouvrage.type_ouvrage.nom_type_ouvrage)
        ws.cell(row_idx, 4, ch.employe.surnom or f"{ch.employe.prenom} {ch.employe.nom}")
        ws.cell(row_idx, 5, float(ch.longueur_adduction_m) if ch.longueur_adduction_m else None)
        ws.cell(row_idx, 6, ch.nom_tacheron)
        ws.cell(row_idx, 7, ch.date_ag_explication)
        ws.cell(row_idx, 8, ch.date_preparation_chantier)
        ws.cell(row_idx, 9, ch.date_ouverture_chantier)
        ws.cell(row_idx, 10, ch.date_reception_provisoire)
        ws.cell(row_idx, 11, ch.date_reception_definitive)
        ws.cell(row_idx, 12, float(ch.cout_estimatif_materiaux) if ch.cout_estimatif_materiaux else None)
        ws.cell(row_idx, 13, float(ch.cout_estimatif_main_oeuvre) if ch.cout_estimatif_main_oeuvre else None)
        ws.cell(row_idx, 14, float(ch.pourcentage_apports_locaux) if ch.pourcentage_apports_locaux else None)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=chantiers.xlsx'
    wb.save(response)
    return response

@csrf_exempt
def import_demandes_excel(request):
    """Importer des demandes depuis Excel"""
    if request.method == 'POST' and request.FILES.get('fichier'):
        try:
            fichier = request.FILES['fichier']
            wb = load_workbook(fichier)
            ws = wb.active

            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:  # Skip empty rows
                    continue

                # Rechercher le village par UUID
                try:
                    village = Village.objects.get(id=row[0])
                except (Village.DoesNotExist, ValueError):
                    continue

                # Rechercher le type d'ouvrage
                type_ouvrage = None
                if row[3]:  # Vérifier que row[3] n'est pas None ou vide
                    type_ouvrage = TypeOuvrage.objects.filter(nom_type_ouvrage__icontains=row[3]).first()

                if not type_ouvrage:
                    continue

                DemandeVillage.objects.create(
                    village=village,
                    statut=row[1] if row[1] else 'demande',
                    type_travaux=row[2] if row[2] else 'construction',
                    type_ouvrage=type_ouvrage,
                    detail_travaux=row[4] or '',
                    type_main_oeuvre=row[5] if row[5] else 'communautaire',
                    semestre_prevision=row[6] or ''
                )
                count += 1

            return JsonResponse({'success': True, 'count': count})
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

    return JsonResponse({'error': 'Fichier manquant'}, status=400)

@csrf_exempt
def import_validations_excel(request):
    """Importer des validations depuis Excel"""
    if request.method == 'POST' and request.FILES.get('fichier'):
        try:
            fichier = request.FILES['fichier']
            wb = load_workbook(fichier)
            ws = wb.active

            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:
                    continue

                # Rechercher le village par UUID
                try:
                    village = Village.objects.get(id=row[0])
                except (Village.DoesNotExist, ValueError):
                    continue

                # Colonnes : 0=Village ID, 1=Nom source, 2=Débit, 3=Date visite,
                #            4=Lat, 5=Lng, 6=Distance, 7=Dénivelé, 8=Donation, 9=Proposition, 10=Nb sources, 11=Visiteur, 12=Commentaires
                visiteur = Employe.objects.filter(nom_employe__icontains=row[11]).first() if row[11] else Employe.objects.filter(date_debauchage__isnull=True).first()
                if not visiteur:
                    continue

                # Parser la date de visite (format JJ/MM/AAAA ou objet datetime)
                from datetime import datetime, date
                date_visite = None
                if row[3]:
                    if isinstance(row[3], (datetime, date)):
                        date_visite = row[3] if isinstance(row[3], date) else row[3].date()
                    elif isinstance(row[3], str):
                        try:
                            date_visite = datetime.strptime(row[3], '%d/%m/%Y').date()
                        except ValueError:
                            try:
                                date_visite = datetime.strptime(row[3], '%Y-%m-%d').date()
                            except ValueError:
                                date_visite = date.today()
                    else:
                        date_visite = date.today()
                else:
                    date_visite = date.today()

                # Parser la proposition (Construction/Réhabilitation/Réparation)
                proposition = None
                if row[9]:
                    proposition_str = str(row[9]).lower()
                    if 'construction' in proposition_str:
                        proposition = 'construction'
                    elif 'réhabilitation' in proposition_str or 'rehabilitation' in proposition_str:
                        proposition = 'rehabilitation'
                    elif 'réparation' in proposition_str or 'reparation' in proposition_str:
                        proposition = 'reparation'

                ValidationSource.objects.create(
                    village=village,
                    nom_source=row[1] or '',
                    debit_source=row[2] if row[2] else 0,
                    date_mesure_debit=date_visite,
                    latitude=row[4] if row[4] else 0,
                    longitude=row[5] if row[5] else 0,
                    distance_source_km=row[6] if row[6] else 0,
                    denivele_m=row[7] if row[7] else 0,
                    donation_source=bool(row[8]) if row[8] else False,
                    proposition=proposition,
                    nombre_sources_visitees=row[10] if row[10] else 1,
                    visiteur=visiteur,
                    commentaires=row[12] or ''
                )
                count += 1

            return JsonResponse({'success': True, 'count': count})
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

    return JsonResponse({'error': 'Fichier manquant'}, status=400)

@csrf_exempt
def import_chantiers_excel(request):
    """Importer des chantiers depuis Excel"""
    if request.method == 'POST' and request.FILES.get('fichier'):
        try:
            fichier = request.FILES['fichier']
            wb = load_workbook(fichier)
            ws = wb.active

            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:
                    continue

                # Rechercher l'ouvrage par nom
                ouvrage = Ouvrage.objects.filter(nom_ouvrage__iexact=row[0]).first()
                if not ouvrage:
                    continue

                # Rechercher l'employé
                employe = Employe.objects.filter(nom_employe__icontains=row[1]).first() if row[1] else Employe.objects.filter(date_debauchage__isnull=True).first()
                if not employe:
                    continue

                # Rechercher le tâcheron par nom
                tacheron = None
                if row[2]:
                    tacheron = Tacheron.objects.filter(nom__iexact=row[2]).first()

                CaracteristiqueChantier.objects.create(
                    ouvrage=ouvrage,
                    employe=employe,
                    tacheron=tacheron,
                    longueur_adduction_m=row[3] if row[3] else 0,
                    date_fin_chantier=row[4] if row[4] else None,
                    cout_estimatif_materiaux=row[5] if row[5] else 0,
                    cout_estimatif_main_oeuvre=row[6] if row[6] else 0,
                    pourcentage_apports_locaux=row[7] if row[7] else 0
                )
                count += 1

            return JsonResponse({'success': True, 'count': count})
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

    return JsonResponse({'error': 'Fichier manquant'}, status=400)

# ==================== TEMPLATES EXCEL ====================

def export_template_demandes(request):
    """Exporter un template Excel vide pour l'import de demandes"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Template Demandes"

    headers = [
        'Village ID (UUID)',
        'Statut (demande/source_validee/preparation_en_cours/suspendu/acheve)',
        'Type travaux (construction/rehabilitation/reparation)',
        'Type ouvrage (nom exact)',
        'Détail travaux',
        'Type main oeuvre (communautaire/paye)',
        'Semestre prévision (ex: S1 2025)'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    # Ajouter liste des villages avec leurs UUID en commentaire
    from apps.geography.models import Village
    villages = Village.objects.select_related('fokontany__commune__district__region').all()[:5]

    exemples = [
        str(villages[0].id) if villages else '550e8400-e29b-41d4-a716-446655440000',
        'demande',
        'construction',
        'Adduction gravitaire',
        'Construction nouvelle adduction avec 3 bornes fontaines',
        'communautaire',
        'S2 2025'
    ]

    # Ajouter feuille de référence villages
    ws_ref = wb.create_sheet("Référence Villages")
    ws_ref.cell(1, 1, "UUID Village").font = Font(bold=True)
    ws_ref.cell(1, 2, "Nom Village").font = Font(bold=True)
    ws_ref.cell(1, 3, "Fokontany").font = Font(bold=True)
    ws_ref.cell(1, 4, "Commune").font = Font(bold=True)

    for idx, village in enumerate(Village.objects.select_related('fokontany__commune').all()[:100], 2):
        ws_ref.cell(idx, 1, str(village.id))
        ws_ref.cell(idx, 2, village.nom_village)
        ws_ref.cell(idx, 3, village.fokontany.nom_fokontany if village.fokontany else '')
        ws_ref.cell(idx, 4, village.fokontany.commune.nom_commune if village.fokontany and village.fokontany.commune else '')

    for col, exemple in enumerate(exemples, 1):
        cell = ws.cell(2, col, exemple)
        cell.font = Font(italic=True, color="666666")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 25

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=template_demandes.xlsx'
    wb.save(response)
    return response

def export_template_validations(request):
    """Exporter un template Excel vide pour l'import de validations"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from apps.geography.models import Village

    wb = Workbook()
    ws = wb.active
    ws.title = "Template Validations"

    headers = [
        'Village ID (UUID)',
        'Nom source',
        'Débit source (L/s)',
        'Date visite (JJ/MM/AAAA)',
        'Latitude (décimal)',
        'Longitude (décimal)',
        'Distance source (km)',
        'Dénivelé (m)',
        'Donation source (OUI/NON)',
        'Proposition (Construction/Réhabilitation/Réparation)',
        'Nb sources visitées',
        'Visiteur (nom employé)',
        'Commentaires'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    villages = Village.objects.all()[:5]
    exemples = [
        str(villages[0].id) if villages else '550e8400-e29b-41d4-a716-446655440000',
        'Source Mandraka',
        15.5,
        '15/01/2025',
        -18.912345,
        47.523456,
        2.5,
        120,
        'OUI',
        'Construction',
        3,
        'Rakoto Jean',
        'Source pérenne, bon débit toute l\'année'
    ]

    # Ajouter feuille de référence villages
    ws_ref = wb.create_sheet("Référence Villages")
    ws_ref.cell(1, 1, "UUID Village").font = Font(bold=True)
    ws_ref.cell(1, 2, "Nom Village").font = Font(bold=True)
    ws_ref.cell(1, 3, "Fokontany").font = Font(bold=True)
    ws_ref.cell(1, 4, "Commune").font = Font(bold=True)

    for idx, village in enumerate(Village.objects.select_related('fokontany__commune').all()[:100], 2):
        ws_ref.cell(idx, 1, str(village.id))
        ws_ref.cell(idx, 2, village.nom_village)
        ws_ref.cell(idx, 3, village.fokontany.nom_fokontany if village.fokontany else '')
        ws_ref.cell(idx, 4, village.fokontany.commune.nom_commune if village.fokontany and village.fokontany.commune else '')

    for col, exemple in enumerate(exemples, 1):
        cell = ws.cell(2, col, exemple)
        cell.font = Font(italic=True, color="666666")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=template_validations.xlsx'
    wb.save(response)
    return response

def export_template_chantiers(request):
    """Exporter un template Excel vide pour l'import de chantiers"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from apps.organization.models import Ouvrage
    from .models import Tacheron

    wb = Workbook()
    ws = wb.active
    ws.title = "Template Chantiers"

    headers = [
        'Nom ouvrage',
        'Employé responsable (nom)',
        'Tâcheron (nom)',
        'Longueur adduction (m)',
        'Date fin (JJ/MM/AAAA)',
        'Coût matériel (Ar)',
        'Coût main oeuvre (Ar)',
        '% Apports locaux'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")

    ouvrages = Ouvrage.objects.all()[:5]
    tacherons = Tacheron.objects.all()[:5]

    exemples = [
        ouvrages[0].nom_ouvrage if ouvrages else 'AEG-001',
        'Rakoto Jean',
        tacherons[0].nom if tacherons else 'Entreprise SAHALA',
        1500,
        '31/12/2025',
        15000000,
        8000000,
        25
    ]

    # Ajouter feuille de référence ouvrages
    ws_ref_ouv = wb.create_sheet("Référence Ouvrages")
    ws_ref_ouv.cell(1, 1, "Nom Ouvrage").font = Font(bold=True)
    ws_ref_ouv.cell(1, 2, "Type").font = Font(bold=True)
    ws_ref_ouv.cell(1, 3, "Statut").font = Font(bold=True)

    for idx, ouvrage in enumerate(Ouvrage.objects.select_related('type_ouvrage').all()[:100], 2):
        ws_ref_ouv.cell(idx, 1, ouvrage.nom_ouvrage)
        ws_ref_ouv.cell(idx, 2, ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else '')
        ws_ref_ouv.cell(idx, 3, ouvrage.statut_ouvrage if hasattr(ouvrage, 'statut_ouvrage') else '')

    # Ajouter feuille de référence tâcherons
    ws_ref_tach = wb.create_sheet("Référence Tâcherons")
    ws_ref_tach.cell(1, 1, "Nom Tâcheron").font = Font(bold=True)
    ws_ref_tach.cell(1, 2, "N° CIN").font = Font(bold=True)
    ws_ref_tach.cell(1, 3, "N° NIF").font = Font(bold=True)

    for idx, tacheron in enumerate(Tacheron.objects.all()[:100], 2):
        ws_ref_tach.cell(idx, 1, tacheron.nom)
        ws_ref_tach.cell(idx, 2, tacheron.numero_cin)
        ws_ref_tach.cell(idx, 3, tacheron.numero_nif)

    for col, exemple in enumerate(exemples, 1):
        cell = ws.cell(2, col, exemple)
        cell.font = Font(italic=True, color="666666")

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=template_chantiers.xlsx'
    wb.save(response)
    return response

# ==================== TÂCHERONS ====================

def tacherons_list(request):
    """Page consolidée des tâcherons"""
    search = request.GET.get('search', '')
    
    tacherons = Tacheron.objects.all()
    
    if search:
        tacherons = tacherons.filter(
            Q(nom__icontains=search) |
            Q(numero_cin__icontains=search) |
            Q(numero_nif__icontains=search) |
            Q(numero_stat__icontains=search)
        )
    
    tacherons = tacherons.order_by('nom')
    
    stats = {
        'total': Tacheron.objects.count(),
        'avec_cin': Tacheron.objects.exclude(numero_cin='').count(),
        'avec_nif': Tacheron.objects.exclude(numero_nif='').count(),
    }
    
    context = {
        'tacherons': tacherons,
        'stats': stats,
        'search': search,
    }
    
    return render(request, 'chantiers/tacherons_list.html', context)


@csrf_exempt
@require_http_methods(["POST"])
def tacheron_create(request):
    """Créer un nouveau tâcheron"""
    try:
        data = json.loads(request.body)
        tacheron = Tacheron.objects.create(
            nom=data['nom'],
            numero_cin=data.get('numero_cin', ''),
            date_delivrance_cin=data.get('date_delivrance_cin'),
            lieu_delivrance_cin=data.get('lieu_delivrance_cin', ''),
            numero_nif=data.get('numero_nif', ''),
            numero_stat=data.get('numero_stat', '')
        )
        return JsonResponse({'success': True, 'id': str(tacheron.id)})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def tacheron_update(request, tacheron_id):
    """Mettre à jour un tâcheron"""
    try:
        tacheron = get_object_or_404(Tacheron, id=tacheron_id)
        data = json.loads(request.body)
        
        for field in ['nom', 'numero_cin', 'date_delivrance_cin', 'lieu_delivrance_cin', 'numero_nif', 'numero_stat']:
            if field in data:
                setattr(tacheron, field, data[field])
        
        tacheron.save()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def tacheron_delete(request, tacheron_id):
    """Supprimer un tâcheron"""
    try:
        tacheron = get_object_or_404(Tacheron, id=tacheron_id)
        tacheron.delete()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)


def export_tacherons_excel(request):
    """Exporter les tâcherons en Excel"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Tâcherons"
    
    headers = ['Nom', 'N° CIN', 'Date CIN', 'Lieu CIN', 'N° NIF', 'N° STAT']
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
    
    for idx, tacheron in enumerate(Tacheron.objects.all(), 2):
        ws.cell(idx, 1, tacheron.nom)
        ws.cell(idx, 2, tacheron.numero_cin)
        ws.cell(idx, 3, tacheron.date_delivrance_cin.strftime('%d/%m/%Y') if tacheron.date_delivrance_cin else '')
        ws.cell(idx, 4, tacheron.lieu_delivrance_cin)
        ws.cell(idx, 5, tacheron.numero_nif)
        ws.cell(idx, 6, tacheron.numero_stat)
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=tacherons.xlsx'
    wb.save(response)
    return response


def export_template_tacherons(request):
    """Exporter un template Excel vide pour l'import de tâcherons"""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Template Tâcherons"
    
    headers = [
        'Nom (requis)',
        'N° CIN',
        'Date délivrance CIN (JJ/MM/AAAA)',
        'Lieu délivrance CIN',
        'N° NIF',
        'N° STAT'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4788", end_color="1F4788", fill_type="solid")
    
    exemples = [
        'Entreprise SAHALA',
        '101234567890',
        '15/01/2020',
        'Antananarivo',
        '4000123456',
        '12345678901234'
    ]
    
    for col, exemple in enumerate(exemples, 1):
        cell = ws.cell(2, col, exemple)
        cell.font = Font(italic=True, color="666666")
    
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 25
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=template_tacherons.xlsx'
    wb.save(response)
    return response


@csrf_exempt
@require_http_methods(["POST"])
def import_tacherons_excel(request):
    """Importer des tâcherons depuis Excel"""
    if request.method == 'POST' and request.FILES.get('fichier'):
        try:
            from openpyxl import load_workbook
            from datetime import datetime
            
            fichier = request.FILES['fichier']
            wb = load_workbook(fichier)
            ws = wb.active
            
            count = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:  # Skip if no name
                    continue
                
                # Parse date if provided
                date_cin = None
                if row[2]:
                    if isinstance(row[2], str):
                        try:
                            date_cin = datetime.strptime(row[2], '%d/%m/%Y').date()
                        except:
                            pass
                    else:
                        date_cin = row[2]
                
                Tacheron.objects.create(
                    nom=row[0],
                    numero_cin=row[1] or '',
                    date_delivrance_cin=date_cin,
                    lieu_delivrance_cin=row[3] or '',
                    numero_nif=row[4] or '',
                    numero_stat=row[5] or ''
                )
                count += 1
            
            return JsonResponse({'success': True, 'count': count})
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)}, status=400)
    
    return JsonResponse({'success': False, 'error': 'Aucun fichier fourni'}, status=400)

# Import des vues pour enquêtes et animations
from .views_enquetes import (
    menages_list, menage_create, menage_detail,
    enquete_form,
    animations_list, animation_create, animation_participants, animation_presence,
    chantier_dashboard
)
