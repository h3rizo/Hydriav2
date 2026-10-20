from __future__ import annotations

from decimal import Decimal
from pathlib import Path
import unicodedata

from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from apps.chantiers.models_workflow import (
    BaremeMainOeuvre,
    CiterneParametrageReference,
    CiterneVolumeConfiguration,
    PosteMainOeuvre,
    WorkflowModeleEquipement,
    WorkflowModeleEquipementLigne,
)
from apps.infrastructure.models import TypeEquipement
from apps.planning.models import Designation


DEFAULT_SOURCE_XLSX = Path(r"c:\Dropbox\Dossiers techniques\source Dossier technique 2.xlsx")
DEFAULT_V3_XLSM = Path(r"c:\Dropbox\Dossiers techniques\Dossier technique V3.xlsm")


class Command(BaseCommand):
    help = (
        "Importe les catalogues workflow chantier depuis les fichiers Excel "
        "(modèles équipements, lignes de désignations, barèmes main d'oeuvre)."
    )

    def add_arguments(self, parser):
        parser.add_argument("--source-xlsx", default=str(DEFAULT_SOURCE_XLSX))
        parser.add_argument("--v3-xlsm", default=str(DEFAULT_V3_XLSM))
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Appliquer les changements (sinon dry-run).",
        )
        parser.add_argument(
            "--clear-existing",
            action="store_true",
            help="Supprimer les données workflow existantes avant import.",
        )

    def handle(self, *args, **options):
        source_path = Path(options["source_xlsx"])
        v3_path = Path(options["v3_xlsm"])
        apply_changes = options["apply"]
        clear_existing = options["clear_existing"]

        if not source_path.exists():
            self.stdout.write(self.style.ERROR(f"Fichier introuvable: {source_path}"))
            return

        source_wb = load_workbook(source_path, data_only=True, read_only=True)
        v3_wb = load_workbook(v3_path, data_only=True, read_only=True) if v3_path.exists() else None

        id_to_designation_name = self._read_designation_dict(source_wb)
        stats = {
            "modele_created": 0,
            "ligne_created": 0,
            "poste_created": 0,
            "bareme_created": 0,
            "citerne_volume_upsert": 0,
            "skipped_designation": 0,
            "skipped_type_equipement": 0,
        }

        action_label = "APPLY" if apply_changes else "DRY-RUN"
        self.stdout.write(f"{action_label}: import workflow depuis {source_path}")
        if v3_wb:
            self.stdout.write(f"{action_label}: enrichissement depuis {v3_path}")

        with transaction.atomic():
            if clear_existing:
                WorkflowModeleEquipementLigne.objects.all().delete()
                WorkflowModeleEquipement.objects.all().delete()
                BaremeMainOeuvre.objects.all().delete()
                PosteMainOeuvre.objects.all().delete()

            self._ensure_citerne_default(apply_changes)
            self._import_citerne_dimensions(source_wb, apply_changes, clear_existing, stats)
            self._import_catalogue_ouvrages(
                source_wb,
                id_to_designation_name,
                apply_changes,
                stats,
            )
            self._import_catalogue_mo(source_wb, apply_changes, stats)
            if v3_wb:
                self._import_liste_mo_pour_contrat(v3_wb, apply_changes, stats)

            if not apply_changes:
                transaction.set_rollback(True)

        self.stdout.write(
            f"{action_label}: modeles={stats['modele_created']}, "
            f"lignes={stats['ligne_created']}, "
            f"postes_mo={stats['poste_created']}, "
            f"baremes_mo={stats['bareme_created']}, "
            f"citerne_volumes={stats['citerne_volume_upsert']}, "
            f"skip_designation={stats['skipped_designation']}, "
            f"skip_type_equipement={stats['skipped_type_equipement']}"
        )

    def _read_designation_dict(self, wb):
        sheet = wb["Designations"]
        data = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            data[str(row[0]).strip()] = (str(row[2]).strip() if row[2] else "")
        return data

    def _find_type_equipement(self, label):
        if not label:
            return None
        txt = str(label).strip()
        te = TypeEquipement.objects.filter(nom_type__iexact=txt).first()
        if te:
            return te
        te = TypeEquipement.objects.filter(nom_type__icontains=txt).first()
        if te:
            return te
        normalized = self._normalize(txt)
        aliases = [
            ("boite de captage", "Captage"),
            ("captage", "Captage"),
            ("brise charge", "Brise-charge"),
            ("citerne", "Citerne"),
            ("borne fontaine", "Distribution"),
            ("bf ", "Distribution"),
            ("conduite", "Adduction"),
            ("passage galva", "Adduction"),
            ("adduction", "Adduction"),
            ("decanteur", "Decanteur"),
            ("filtre", "Filtre"),
            ("repartiteur", "Répartiteur"),
            ("vidange", "Vidange"),
            ("bureau", "Bureau"),
            ("smr", "SMR"),
            ("stock", "Stock"),
            ("pompe", "Pompe test"),
        ]
        for pattern, target in aliases:
            if pattern in normalized:
                te = TypeEquipement.objects.filter(nom_type__iexact=target).first()
                if te:
                    return te
        for te in TypeEquipement.objects.all():
            if self._normalize(te.nom_type) in normalized or normalized in self._normalize(te.nom_type):
                return te
        return None

    def _normalize(self, value: str) -> str:
        value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
        value = value.lower().replace("-", " ").replace("_", " ")
        value = " ".join(value.split())
        return value

    def _q3(self, value):
        return Decimal(str(value)).quantize(Decimal("0.001"))

    def _find_designation(self, article_id, id_to_name):
        key = str(article_id).strip()
        nom = id_to_name.get(key)
        if nom:
            obj = Designation.objects.filter(nom__iexact=nom).first()
            if obj:
                return obj
        return Designation.objects.filter(code_materiel__iexact=key).first()

    def _ensure_citerne_default(self, apply_changes):
        if not apply_changes:
            return
        CiterneParametrageReference.objects.get_or_create(
            nom="DEFAULT",
            defaults={
                "epaisseur_paroi_m": Decimal("0.10"),
                "hauteur_visible_m": Decimal("0.20"),
                "epaisseur_blocage_m": Decimal("0.20"),
                "epaisseur_beton_arme_m": Decimal("0.10"),
                "hauteur_aeration_m": Decimal("0.10"),
                "epaisseur_dalle_couverture_m": Decimal("0.10"),
            },
        )

    def _import_citerne_dimensions(self, wb, apply_changes, clear_existing, stats):
        if "dim citerne" not in wb.sheetnames:
            return

        sheet = wb["dim citerne"]
        if not apply_changes:
            reference = CiterneParametrageReference(
                nom="DEFAULT",
                epaisseur_paroi_m=Decimal("0.10"),
                hauteur_visible_m=Decimal("0.20"),
                epaisseur_blocage_m=Decimal("0.20"),
                epaisseur_beton_arme_m=Decimal("0.10"),
            )
        else:
            reference, _ = CiterneParametrageReference.objects.get_or_create(
                nom="DEFAULT",
                defaults={
                    "epaisseur_paroi_m": Decimal("0.10"),
                    "hauteur_visible_m": Decimal("0.20"),
                    "epaisseur_blocage_m": Decimal("0.20"),
                    "epaisseur_beton_arme_m": Decimal("0.10"),
                    "hauteur_aeration_m": Decimal("0.10"),
                    "epaisseur_dalle_couverture_m": Decimal("0.10"),
                },
            )
            epaisseur_paroi = sheet.cell(row=8, column=4).value
            hauteur_visible = sheet.cell(row=9, column=4).value
            epaisseur_blocage = sheet.cell(row=10, column=4).value
            epaisseur_ba = sheet.cell(row=11, column=4).value
            hauteur_aeration = sheet.cell(row=13, column=4).value
            epaisseur_dalle = sheet.cell(row=15, column=4).value
            if epaisseur_paroi not in (None, ""):
                reference.epaisseur_paroi_m = self._q3(epaisseur_paroi)
            if hauteur_visible not in (None, ""):
                reference.hauteur_visible_m = self._q3(hauteur_visible)
            if epaisseur_blocage not in (None, ""):
                reference.epaisseur_blocage_m = self._q3(epaisseur_blocage)
            if epaisseur_ba not in (None, ""):
                reference.epaisseur_beton_arme_m = self._q3(epaisseur_ba)
            if hauteur_aeration not in (None, ""):
                reference.hauteur_aeration_m = self._q3(hauteur_aeration)
            if epaisseur_dalle not in (None, ""):
                reference.epaisseur_dalle_couverture_m = self._q3(epaisseur_dalle)
            reference.save(
                update_fields=[
                    "epaisseur_paroi_m",
                    "hauteur_visible_m",
                    "epaisseur_blocage_m",
                    "epaisseur_beton_arme_m",
                    "hauteur_aeration_m",
                    "epaisseur_dalle_couverture_m",
                ]
            )

        if apply_changes and clear_existing:
            CiterneVolumeConfiguration.objects.filter(reference=reference).delete()

        max_col = sheet.max_column or 0
        for col in range(4, max_col + 1):
            header = sheet.cell(row=2, column=col).value
            if not header:
                continue
            if "citerne" not in str(header).lower():
                continue

            volume = sheet.cell(row=4, column=col).value
            diametre_interieur = sheet.cell(row=7, column=col).value
            hauteur_trop_plein = sheet.cell(row=14, column=col).value
            if volume in (None, "") or diametre_interieur in (None, "") or hauteur_trop_plein in (None, ""):
                continue

            payload = {
                "diametre_interieur_m": self._q3(diametre_interieur),
                "hauteur_trop_plein_m": self._q3(hauteur_trop_plein),
                "diametre_exterieur_m": self._q3(sheet.cell(row=5, column=col).value) if sheet.cell(row=5, column=col).value not in (None, "") else None,
                "hauteur_interieure_m": self._q3(sheet.cell(row=12, column=col).value) if sheet.cell(row=12, column=col).value not in (None, "") else None,
                "profondeur_fouille_m": self._q3(sheet.cell(row=6, column=col).value) if sheet.cell(row=6, column=col).value not in (None, "") else None,
                "hauteur_utile_m": self._q3(sheet.cell(row=25, column=col).value) if sheet.cell(row=25, column=col).value not in (None, "") else None,
                "actif": True,
            }
            volume_decimal = self._q3(volume)

            if apply_changes:
                CiterneVolumeConfiguration.objects.update_or_create(
                    reference=reference,
                    volume_cible_m3=volume_decimal,
                    defaults=payload,
                )
            stats["citerne_volume_upsert"] += 1

    def _import_catalogue_ouvrages(self, wb, id_to_designation_name, apply_changes, stats):
        sheet = wb["Catalogue_ouvrages"]
        headers = [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if len(headers) < 3:
            return

        modeles_cache = {}
        for col in range(2, len(headers)):
            equip_name = headers[col]
            if not equip_name:
                continue
            type_eq = self._find_type_equipement(equip_name)
            if not type_eq:
                stats["skipped_type_equipement"] += 1
                continue
            if type_eq.id not in modeles_cache:
                if apply_changes:
                    modele, created = WorkflowModeleEquipement.objects.get_or_create(
                        code=f"CAT-{type_eq.nom_type[:40]}",
                        defaults={"type_equipement": type_eq},
                    )
                    if created:
                        stats["modele_created"] += 1
                else:
                    modele = type("Obj", (), {"id": str(type_eq.id)})()
                    stats["modele_created"] += 1
                modeles_cache[col] = modele

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            designation = self._find_designation(row[0], id_to_designation_name)
            if not designation:
                stats["skipped_designation"] += 1
                continue
            for col, modele in modeles_cache.items():
                if col >= len(row):
                    continue
                value = row[col]
                if value in (None, "", 0, 0.0):
                    continue
                try:
                    quantite = Decimal(str(value))
                except Exception:
                    continue
                if quantite <= 0:
                    continue
                if apply_changes:
                    _, created = WorkflowModeleEquipementLigne.objects.get_or_create(
                        modele=modele,
                        designation=designation,
                        defaults={"mode_quantite": "fixe", "quantite_unitaire": quantite, "ordre": 1},
                    )
                    if created:
                        stats["ligne_created"] += 1
                else:
                    stats["ligne_created"] += 1

    def _upsert_bareme(self, type_eq_label, poste_libelle, unite, pu, apply_changes, stats):
        type_eq = self._find_type_equipement(type_eq_label)
        if not type_eq:
            stats["skipped_type_equipement"] += 1
            return
        if not poste_libelle or pu in (None, ""):
            return
        try:
            pu_decimal = Decimal(str(pu))
        except Exception:
            return
        if pu_decimal <= 0:
            return

        poste_code = f"MO-{str(poste_libelle).strip().upper()[:60]}"
        if apply_changes:
            poste, created_poste = PosteMainOeuvre.objects.get_or_create(
                code=poste_code,
                defaults={"libelle": str(poste_libelle).strip(), "unite": str(unite or "U").strip() or "U"},
            )
            if created_poste:
                stats["poste_created"] += 1
            _, created_bareme = BaremeMainOeuvre.objects.get_or_create(
                poste=poste,
                type_equipement=type_eq,
                type_point_eau=None,
                defaults={"quantite_unitaire": Decimal("1"), "prix_unitaire": pu_decimal},
            )
            if created_bareme:
                stats["bareme_created"] += 1
        else:
            stats["poste_created"] += 1
            stats["bareme_created"] += 1

    def _import_catalogue_mo(self, wb, apply_changes, stats):
        sheet = wb["Catalogue_MO"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            type_eq_label = row[3] if len(row) > 3 else None
            poste_libelle = row[4] if len(row) > 4 else None
            unite = row[5] if len(row) > 5 else "U"
            pu = row[6] if len(row) > 6 else None
            self._upsert_bareme(type_eq_label, poste_libelle, unite, pu, apply_changes, stats)

    def _import_liste_mo_pour_contrat(self, wb, apply_changes, stats):
        if "listeMOPourContrat" not in wb.sheetnames:
            return
        sheet = wb["listeMOPourContrat"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            type_eq_label = row[1] if len(row) > 1 else None
            poste_libelle = row[2] if len(row) > 2 else None
            pu = row[3] if len(row) > 3 else None
            self._upsert_bareme(type_eq_label, poste_libelle, "U", pu, apply_changes, stats)
