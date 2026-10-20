from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.db import DatabaseError, IntegrityError, transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook, load_workbook

from apps.billing.models import PointEau, TypePointEau
from apps.infrastructure.models import TypeEquipement
from apps.organization.models import Ouvrage
from apps.planning.models import Designation
from .models import Tacheron
from .models_workflow import (
    ChantierBesoinLigne,
    ChantierDocument,
    ChantierMouvement,
    ChantierSelectionEquipement,
    ChantierSurplusLigne,
    ChantierWorkflow,
    CiterneDimensionnement,
    CiterneVolumeConfiguration,
    PosteMainOeuvre,
    build_citerne_dimensions_from_volume,
    get_or_create_default_citerne_reference,
    generer_besoins_chantier,
)


def _d(value, default="0"):
    if value in (None, ""):
        return Decimal(default)
    return Decimal(str(value).replace(",", "."))


def _d_opt(value):
    if value in (None, ""):
        return None
    return Decimal(str(value).replace(",", "."))


def _bool_from_value(value, default=True):
    if value in (None, ""):
        return default
    txt = str(value).strip().lower()
    return txt in {"1", "true", "oui", "yes", "y", "vrai"}


def _q3(value):
    return Decimal(str(value)).quantize(Decimal("0.001"))


def _compute_citerne_dimensions(volume_cible_m3, diametre_interieur_m, reference):
    volume = _d(volume_cible_m3)
    d_int = _d(diametre_interieur_m)
    if d_int <= 0:
        raise ValueError("Diametre interieur invalide.")

    pi = Decimal("3.141592653589793")
    surface = (pi * d_int * d_int) / Decimal("4")
    if surface <= 0:
        raise ValueError("Surface de base invalide.")

    h_trop_plein = volume / surface
    hauteur_aeration = Decimal(getattr(reference, "hauteur_aeration_m", Decimal("0.10")))
    epaisseur_paroi = Decimal(getattr(reference, "epaisseur_paroi_m", Decimal("0.10")))
    epaisseur_blocage = Decimal(getattr(reference, "epaisseur_blocage_m", Decimal("0.20")))
    epaisseur_ba = Decimal(getattr(reference, "epaisseur_beton_arme_m", Decimal("0.10")))
    epaisseur_dalle = Decimal(getattr(reference, "epaisseur_dalle_couverture_m", Decimal("0.10")))
    hauteur_visible = Decimal(getattr(reference, "hauteur_visible_m", Decimal("0.20")))

    h_interieure = h_trop_plein + hauteur_aeration
    d_exterieur = d_int + (epaisseur_paroi * Decimal("2"))
    h_exterieure = (
        h_interieure
        + epaisseur_blocage
        + epaisseur_ba
        + epaisseur_dalle
    )
    profondeur_fouille = h_exterieure - hauteur_visible
    hauteur_utile = h_trop_plein

    return {
        "volume_cible_m3": _q3(volume),
        "diametre_interieur_m": _q3(d_int),
        "hauteur_trop_plein_m": _q3(h_trop_plein),
        "diametre_exterieur_m": _q3(d_exterieur),
        "profondeur_fouille_m": _q3(profondeur_fouille),
        "hauteur_interieure_m": _q3(h_interieure),
        "hauteur_utile_m": _q3(hauteur_utile),
    }


def _suivi_fields(data):
    """Type de travaux / description / statut du chantier (affichés dans le diagramme de flux)."""
    type_travaux = (data.get("type_travaux") or "").strip()
    if type_travaux not in dict(ChantierWorkflow.TYPE_TRAVAUX_CHOICES):
        type_travaux = ""
    statut = (data.get("statut_chantier") or "").strip()
    if statut not in dict(ChantierWorkflow.STATUT_CHANTIER_CHOICES):
        statut = "en_cours"
    return {
        "type_travaux": type_travaux,
        "description": (data.get("description") or "").strip(),
        "statut_chantier": statut,
    }


def _annoter_workflows_gestion_grading_montant(workflows):
    """Annote en place une liste de ChantierWorkflow avec, pour l'ouvrage lié :
    - deja_gere_organisation : nom de l'organisation gestionnaire active (ou None)
    - grading_actuel : dernier OuvrageGradingAnnuel (grade + année) ou None
    - montant_collecte : Decimal déjà encaissé (forfaitaire + volumétrique pointé DiameX)
    """
    from apps.billing.models import ModaliteGestion
    from apps.billing.services_recouvrement_ouvrage import montant_collecte_par_ouvrage
    from apps.organization.models import OuvrageGradingAnnuel

    ouvrage_ids = [w.ouvrage_id for w in workflows if w.ouvrage_id]
    if not ouvrage_ids:
        return

    gestion_par_ouvrage = dict(
        ModaliteGestion.objects.filter(
            ouvrage_id__in=ouvrage_ids, statut='actif'
        ).select_related('gestionnaire').values_list('ouvrage_id', 'gestionnaire__nom_organisation')
    )

    grading_par_ouvrage = {}
    for g in OuvrageGradingAnnuel.objects.filter(ouvrage_id__in=ouvrage_ids).order_by('ouvrage_id', '-annee'):
        grading_par_ouvrage.setdefault(g.ouvrage_id, g)

    montants_par_ouvrage = montant_collecte_par_ouvrage(ouvrage_ids)

    for w in workflows:
        w.deja_gere_organisation = gestion_par_ouvrage.get(w.ouvrage_id)
        w.grading_actuel = grading_par_ouvrage.get(w.ouvrage_id)
        w.montant_collecte = montants_par_ouvrage.get(w.ouvrage_id, Decimal('0'))


@require_http_methods(["POST"])
def workflow_ui_update_suivi(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    for field, value in _suivi_fields(request.POST).items():
        setattr(workflow, field, value)
    workflow.save(update_fields=["type_travaux", "description", "statut_chantier"])
    messages.success(request, "Suivi du chantier enregistré.")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_cloturer(request, workflow_id):
    """Clôture le chantier (il sort du diagramme de flux « chantiers en cours »), ou le rouvre."""
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    if request.POST.get("action") == "rouvrir":
        if workflow.statut == "cloture":
            workflow.statut = "recollement"
            workflow.save(update_fields=["statut"])
            messages.success(request, f"Chantier {workflow.code} rouvert (statut : Recollement).")
    elif workflow.statut != "cloture":
        workflow.statut = "cloture"
        workflow.save(update_fields=["statut"])
        messages.success(request, f"Chantier {workflow.code} clôturé.")
    next_url = request.POST.get("next")
    if next_url == "list":
        return redirect("chantiers:workflows_ui_list")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


def workflows_ui_list(request):
    prefill_notes = ""
    prefill_ouvrage_id = ""

    if request.method == "POST":
        prefill_notes = request.POST.get("notes", "")
        prefill_ouvrage_id = request.POST.get("ouvrage_id") or ""
        try:
            ouvrage_id = prefill_ouvrage_id
            if not ouvrage_id:
                raise ValueError("Ouvrage obligatoire.")

            wf = None
            for _ in range(5):
                try:
                    wf = ChantierWorkflow.objects.create(
                        code=ChantierWorkflow.next_code(),
                        ouvrage_id=ouvrage_id,
                        notes=prefill_notes,
                        **_suivi_fields(request.POST),
                    )
                    break
                except IntegrityError:
                    # Retry if two users generate the same code concurrently.
                    continue

            if wf is None:
                raise ValueError("Impossible de generer un code workflow unique.")

            messages.success(request, f"Workflow {wf.code} cree.")
            return redirect("chantiers:workflow_ui_detail", workflow_id=wf.id)
        except Exception as e:
            messages.error(request, f"Erreur creation workflow: {e}")

    workflows = list(ChantierWorkflow.objects.select_related("ouvrage").order_by("-id"))
    ouvrages = Ouvrage.objects.select_related("type_ouvrage").order_by("nom_ouvrage")

    # Annoter chaque workflow : gestion active, dernier grading, montant deja collecte
    _annoter_workflows_gestion_grading_montant(workflows)

    localisations_by_ouvrage = {}
    pointseaux = (
        PointEau.objects.filter(ouvrage_id__in=ouvrages.values_list("id", flat=True))
        .select_related("parcelle__village__fokontany__commune__district__region")
        .order_by("ouvrage__nom_ouvrage", "parcelle__village__fokontany__commune__nom_commune")
    )

    for pe in pointseaux:
        if not pe.ouvrage_id:
            continue

        ouvrage_id_str = str(pe.ouvrage_id)
        if ouvrage_id_str in localisations_by_ouvrage:
            continue

        parcelle = pe.parcelle
        village = parcelle.village if parcelle and parcelle.village else None
        fokontany = village.fokontany if village and village.fokontany else None
        commune = fokontany.commune if fokontany and fokontany.commune else None
        district = commune.district if commune and commune.district else None
        region = district.region if district and district.region else None

        localisations_by_ouvrage[ouvrage_id_str] = {
            "commune_id": str(commune.id) if commune else None,
            "commune_nom": commune.nom_commune if commune else None,
            "district_id": str(district.id) if district else None,
            "district_nom": district.nom_district if district else None,
            "region_id": str(region.id) if region else None,
            "region_nom": region.nom_region if region else None,
        }

    ouvrages_data = []
    for ouvrage in ouvrages:
        ouvrage_id_str = str(ouvrage.id)
        localisation = localisations_by_ouvrage.get(
            ouvrage_id_str,
            {
                "commune_id": None,
                "commune_nom": None,
                "district_id": None,
                "district_nom": None,
                "region_id": None,
                "region_nom": None,
            },
        )

        commune_nom = localisation.get("commune_nom")
        label = ouvrage.nom_ouvrage
        if commune_nom:
            label = f"{ouvrage.nom_ouvrage} ({commune_nom})"

        ouvrages_data.append(
            {
                "id": ouvrage_id_str,
                "nom_ouvrage": ouvrage.nom_ouvrage,
                "label": label,
                "type_ouvrage_id": str(ouvrage.type_ouvrage_id) if ouvrage.type_ouvrage_id else None,
                "type_ouvrage_nom": ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else "",
                **localisation,
            }
        )

    return render(
        request,
        "chantiers/workflows_list.html",
        {
            "workflows": workflows,
            "ouvrages_data": ouvrages_data,
            "prefill_notes": prefill_notes,
            "prefill_ouvrage_id": prefill_ouvrage_id,
        },
    )


def _build_recollement(workflow):
    mouvements = workflow.mouvements.select_related("besoin")
    by_besoin = {}
    for m in mouvements:
        if not m.besoin_id:
            continue
        by_besoin.setdefault(
            m.besoin_id,
            {
                "commande": Decimal("0"),
                "livraison": Decimal("0"),
                "utilisation": Decimal("0"),
                "retour": Decimal("0"),
                "recollement": Decimal("0"),
            },
        )
        by_besoin[m.besoin_id][m.type_mouvement] += Decimal(m.quantite)

    lines = []
    total_prevu = Decimal("0")
    total_utilise = Decimal("0")
    total_valeur = Decimal("0")
    for b in workflow.besoins.select_related("designation", "poste_main_oeuvre").all():
        agg = by_besoin.get(
            b.id,
            {
                "commande": Decimal("0"),
                "livraison": Decimal("0"),
                "utilisation": Decimal("0"),
                "retour": Decimal("0"),
                "recollement": Decimal("0"),
            },
        )
        utilise_reel = agg["recollement"] if agg["recollement"] > 0 else agg["utilisation"]
        ecart = utilise_reel - Decimal(b.quantite)
        valeur = Decimal("0")
        if b.prix_unitaire is not None:
            valeur = utilise_reel * Decimal(b.prix_unitaire)
        lines.append(
            {
                "besoin": b,
                "agg": agg,
                "utilise_reel": utilise_reel,
                "ecart": ecart,
                "valeur": valeur,
            }
        )
        total_prevu += Decimal(b.quantite)
        total_utilise += utilise_reel
        total_valeur += valeur
    return lines, total_prevu, total_utilise, total_valeur


def workflow_ui_detail(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow.objects.select_related("ouvrage"), id=workflow_id)
    selections = workflow.selections.select_related("type_equipement", "type_point_eau", "citerne_dimensionnement").order_by("-id")
    besoins = workflow.besoins.select_related("designation", "poste_main_oeuvre", "selection__type_equipement").order_by("-id")
    surplus = workflow.surplus.select_related("designation", "poste_main_oeuvre").order_by("-id")
    documents = workflow.documents.select_related("tacheron", "parent").order_by("-id")
    mouvements = workflow.mouvements.select_related("besoin__designation", "besoin__poste_main_oeuvre").order_by("-id")
    recollement_lines, total_prevu, total_utilise, total_valeur = _build_recollement(workflow)
    citerne_rows = []
    for s in selections:
        c = s.citerne_dimensionnement
        if not c:
            continue
        ref = c.reference
        citerne_rows.append(
            {
                "selection": s,
                "citerne": c,
                "reference": ref,
                "formule_diametre_exterieur": "D_ext = D_int + (2 x e_paroi)",
                "formule_volume_reel": "V_reel = pi x D_int² x H_eau / 4",
            }
        )

    try:
        citerne_configs = CiterneVolumeConfiguration.objects.filter(
            reference=get_or_create_default_citerne_reference(),
            actif=True,
        ).order_by("volume_cible_m3")
    except DatabaseError:
        citerne_configs = []

    context = {
        "workflow": workflow,
        "selections": selections,
        "besoins": besoins,
        "surplus": surplus,
        "documents": documents,
        "mouvements": mouvements,
        "recollement_lines": recollement_lines,
        "recollement_total_prevu": total_prevu,
        "recollement_total_utilise": total_utilise,
        "recollement_total_valeur": total_valeur,
        "citerne_rows": citerne_rows,
        "types_equipements": TypeEquipement.objects.order_by("nom_type"),
        "types_points_eau": TypePointEau.objects.order_by("nom_type"),
        "designations": Designation.objects.order_by("nom")[:300],
        "postes_mo": PosteMainOeuvre.objects.order_by("libelle"),
        "tacherons": Tacheron.objects.order_by("nom"),
        "contrats_docs": documents.filter(type_document="contrat"),
        "citerne_configs": citerne_configs,
        "type_travaux_choices": ChantierWorkflow.TYPE_TRAVAUX_CHOICES,
        "statut_chantier_choices": ChantierWorkflow.STATUT_CHANTIER_CHOICES,
    }
    return render(request, "chantiers/workflow_detail.html", context)


@require_http_methods(["POST"])
def workflow_ui_add_selection(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        with transaction.atomic():
            citerne = None
            if request.POST.get("is_citerne") == "1":
                dims = build_citerne_dimensions_from_volume(request.POST.get("volume_cible_m3"))
                citerne = CiterneDimensionnement.objects.create(
                    reference=dims["reference"],
                    volume_cible_m3=dims["volume_cible_m3"],
                    diametre_interieur_m=dims["diametre_interieur_m"],
                    hauteur_eau_m=dims["hauteur_eau_m"],
                )
            ChantierSelectionEquipement.objects.create(
                chantier=workflow,
                type_equipement_id=request.POST.get("type_equipement_id"),
                type_point_eau_id=request.POST.get("type_point_eau_id") or None,
                quantite=_d(request.POST.get("quantite"), default="1"),
                citerne_dimensionnement=citerne,
            )
        messages.success(request, "Sélection ajoutée.")
    except Exception as e:
        messages.error(request, f"Erreur ajout sélection: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["GET", "POST"])
def workflow_ui_citerne_config_list(request):
    reference = get_or_create_default_citerne_reference()

    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        try:
            if action == "update_reference":
                reference.epaisseur_paroi_m = _d(
                    request.POST.get("epaisseur_paroi_m"),
                    default=str(reference.epaisseur_paroi_m),
                )
                reference.hauteur_visible_m = _d(
                    request.POST.get("hauteur_visible_m"),
                    default=str(reference.hauteur_visible_m),
                )
                reference.epaisseur_blocage_m = _d(
                    request.POST.get("epaisseur_blocage_m"),
                    default=str(reference.epaisseur_blocage_m),
                )
                reference.epaisseur_beton_arme_m = _d(
                    request.POST.get("epaisseur_beton_arme_m"),
                    default=str(reference.epaisseur_beton_arme_m),
                )
                if hasattr(reference, "hauteur_aeration_m"):
                    reference.hauteur_aeration_m = _d(
                        request.POST.get("hauteur_aeration_m"),
                        default=str(reference.hauteur_aeration_m),
                    )
                if hasattr(reference, "epaisseur_dalle_couverture_m"):
                    reference.epaisseur_dalle_couverture_m = _d(
                        request.POST.get("epaisseur_dalle_couverture_m"),
                        default=str(reference.epaisseur_dalle_couverture_m),
                    )
                reference.save()

                # Recalcule toutes les configurations existantes avec les nouveaux paramètres.
                updated = 0
                for cfg in CiterneVolumeConfiguration.objects.filter(reference=reference):
                    computed = _compute_citerne_dimensions(cfg.volume_cible_m3, cfg.diametre_interieur_m, reference)
                    CiterneVolumeConfiguration.objects.filter(id=cfg.id).update(
                        hauteur_trop_plein_m=computed["hauteur_trop_plein_m"],
                        diametre_exterieur_m=computed["diametre_exterieur_m"],
                        hauteur_interieure_m=computed["hauteur_interieure_m"],
                        profondeur_fouille_m=computed["profondeur_fouille_m"],
                        hauteur_utile_m=computed["hauteur_utile_m"],
                    )
                    updated += 1

                messages.success(
                    request,
                    f"Parametres de calcul mis a jour. {updated} configuration(s) recalculee(s).",
                )
            elif action == "delete":
                config_id = request.POST.get("config_id")
                CiterneVolumeConfiguration.objects.filter(id=config_id, reference=reference).delete()
                messages.success(request, "Configuration citerne supprimee.")
            else:
                config_id = request.POST.get("config_id")
                volume_cible = _d(request.POST.get("volume_cible_m3"))
                d_int = _d(request.POST.get("diametre_interieur_m"))
                computed = _compute_citerne_dimensions(volume_cible, d_int, reference)

                payload = {
                    "volume_cible_m3": volume_cible,
                    "diametre_interieur_m": d_int,
                    "hauteur_trop_plein_m": computed["hauteur_trop_plein_m"],
                    "diametre_exterieur_m": computed["diametre_exterieur_m"],
                    "hauteur_interieure_m": computed["hauteur_interieure_m"],
                    "profondeur_fouille_m": computed["profondeur_fouille_m"],
                    "hauteur_utile_m": computed["hauteur_utile_m"],
                    "actif": request.POST.get("actif") == "1",
                }
                if config_id:
                    CiterneVolumeConfiguration.objects.filter(id=config_id, reference=reference).update(**payload)
                    messages.success(request, "Configuration citerne mise a jour.")
                else:
                    CiterneVolumeConfiguration.objects.update_or_create(
                        reference=reference,
                        volume_cible_m3=payload["volume_cible_m3"],
                        defaults=payload,
                    )
                    messages.success(request, "Configuration citerne enregistree.")
        except Exception as e:
            messages.error(request, f"Erreur configuration citerne: {e}")
        return redirect("chantiers:workflow_ui_citerne_config_list")

    try:
        configs = CiterneVolumeConfiguration.objects.filter(reference=reference).order_by("volume_cible_m3")
    except DatabaseError:
        messages.error(request, "La table de configuration citerne n'est pas encore migree.")
        configs = []
    return render(
        request,
        "chantiers/citerne_configs.html",
        {
            "reference": reference,
            "configs": configs,
        },
    )


@require_http_methods(["GET"])
def workflow_ui_citerne_config_template(request):
    reference = get_or_create_default_citerne_reference()
    configs = CiterneVolumeConfiguration.objects.filter(reference=reference).order_by("volume_cible_m3")

    wb = Workbook()
    ws_params = wb.active
    ws_params.title = "parametres_reference"

    ws_params["A1"] = "parametre"
    ws_params["B1"] = "valeur"
    ws_params.append(["epaisseur_paroi_m", float(reference.epaisseur_paroi_m)])
    ws_params.append(["hauteur_partie_visible_m", float(reference.hauteur_visible_m)])
    ws_params.append(["epaisseur_blocage_fond_m", float(reference.epaisseur_blocage_m)])
    ws_params.append(["epaisseur_ba_fond_m", float(reference.epaisseur_beton_arme_m)])
    ws_params.append(["hauteur_aeration_m", float(getattr(reference, "hauteur_aeration_m", Decimal("0.10")))])
    ws_params.append(["epaisseur_dalle_couverture_m", float(getattr(reference, "epaisseur_dalle_couverture_m", Decimal("0.10")))])

    ws = wb.create_sheet(title="citerne_config")

    headers = [
        "volume_cible_m3",
        "diametre_interieur_m",
        "hauteur_trop_plein_m (calc)",
        "diametre_exterieur_m (calc)",
        "profondeur_fouille_m (calc)",
        "hauteur_interieure_m (calc)",
        "hauteur_utile_m (calc)",
        "actif",
    ]
    ws.append(headers)

    def set_formula_row(idx, volume, d_int, actif=True):
        ws.cell(row=idx, column=1, value=volume)
        ws.cell(row=idx, column=2, value=d_int)
        ws.cell(row=idx, column=3, value=f"=A{idx}/(PI()*B{idx}^2/4)")
        ws.cell(row=idx, column=4, value=f"=B{idx}+(parametres_reference!$B$2*2)")
        ws.cell(row=idx, column=6, value=f"=C{idx}+parametres_reference!$B$6")
        ws.cell(
            row=idx,
            column=5,
            value=f"=F{idx}+parametres_reference!$B$4+parametres_reference!$B$5+parametres_reference!$B$7-parametres_reference!$B$3",
        )
        ws.cell(row=idx, column=7, value=f"=C{idx}")
        ws.cell(row=idx, column=8, value=1 if actif else 0)

    if configs:
        for row_idx, cfg in enumerate(configs, start=2):
            set_formula_row(row_idx, float(cfg.volume_cible_m3), float(cfg.diametre_interieur_m), actif=cfg.actif)
    else:
        set_formula_row(2, 1.0, 1.5, actif=True)

    notes = wb.create_sheet(title="instructions")
    notes["A1"] = "Mode d'emploi"
    notes["A2"] = "1) Ajuster les valeurs de la feuille 'parametres_reference' si besoin."
    notes["A3"] = "2) Dans 'citerne_config', saisir seulement volume_cible_m3 et diametre_interieur_m."
    notes["A4"] = "3) Les autres colonnes sont calculees automatiquement par formule Excel."
    notes["A5"] = "4) actif: 1/oui/true = actif, 0/non/false = inactif."
    notes["A6"] = "5) Enregistrer le fichier puis importer."

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="template_configuration_citerne.xlsx"'
    return response


@require_http_methods(["POST"])
def workflow_ui_citerne_config_import(request):
    reference = get_or_create_default_citerne_reference()
    uploaded = request.FILES.get("excel_file")
    if not uploaded:
        messages.error(request, "Aucun fichier Excel fourni.")
        return redirect("chantiers:workflow_ui_citerne_config_list")

    try:
        wb = load_workbook(uploaded, data_only=True)
    except Exception as e:
        messages.error(request, f"Fichier Excel invalide: {e}")
        return redirect("chantiers:workflow_ui_citerne_config_list")

    if "citerne_config" not in wb.sheetnames:
        messages.error(request, "Feuille 'citerne_config' introuvable dans le fichier.")
        return redirect("chantiers:workflow_ui_citerne_config_list")

    # Update reference parameters when provided.
    if "parametres_reference" in wb.sheetnames:
        ws_params = wb["parametres_reference"]
        params_map = {}
        for prow in ws_params.iter_rows(min_row=2, values_only=True):
            if not prow:
                continue
            key = (prow[0] or "").strip() if isinstance(prow[0], str) else prow[0]
            if not key:
                continue
            params_map[str(key).strip()] = prow[1]

        try:
            fields_to_update = [
                "epaisseur_paroi_m",
                "hauteur_visible_m",
                "epaisseur_blocage_m",
                "epaisseur_beton_arme_m",
            ]
            if params_map.get("epaisseur_paroi_m") not in (None, ""):
                reference.epaisseur_paroi_m = _d(params_map.get("epaisseur_paroi_m"))
            if params_map.get("hauteur_partie_visible_m") not in (None, ""):
                reference.hauteur_visible_m = _d(params_map.get("hauteur_partie_visible_m"))
            if params_map.get("epaisseur_blocage_fond_m") not in (None, ""):
                reference.epaisseur_blocage_m = _d(params_map.get("epaisseur_blocage_fond_m"))
            if params_map.get("epaisseur_ba_fond_m") not in (None, ""):
                reference.epaisseur_beton_arme_m = _d(params_map.get("epaisseur_ba_fond_m"))
            if hasattr(reference, "hauteur_aeration_m") and params_map.get("hauteur_aeration_m") not in (None, ""):
                reference.hauteur_aeration_m = _d(params_map.get("hauteur_aeration_m"))
                fields_to_update.append("hauteur_aeration_m")
            if hasattr(reference, "epaisseur_dalle_couverture_m") and params_map.get("epaisseur_dalle_couverture_m") not in (None, ""):
                reference.epaisseur_dalle_couverture_m = _d(params_map.get("epaisseur_dalle_couverture_m"))
                fields_to_update.append("epaisseur_dalle_couverture_m")
            reference.save(update_fields=fields_to_update)
        except Exception as e:
            messages.error(request, f"Erreur parametres_reference: {e}")

    ws = wb["citerne_config"]

    created_or_updated = 0
    skipped = 0
    errors = []

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            skipped += 1
            continue

        volume = row[0] if len(row) > 0 else None
        d_int = row[1] if len(row) > 1 else None
        if volume in (None, "") and d_int in (None, ""):
            skipped += 1
            continue

        if volume in (None, "") or d_int in (None, ""):
            errors.append(f"Ligne {row_index}: colonnes obligatoires manquantes (volume, diametre_interieur).")
            continue

        try:
            volume_dec = _d(volume)
            d_int_dec = _d(d_int)
            computed = _compute_citerne_dimensions(volume_dec, d_int_dec, reference)
            actif = _bool_from_value(row[7] if len(row) > 7 else None, default=True)

            CiterneVolumeConfiguration.objects.update_or_create(
                reference=reference,
                volume_cible_m3=volume_dec,
                defaults={
                    "diametre_interieur_m": computed["diametre_interieur_m"],
                    "hauteur_trop_plein_m": computed["hauteur_trop_plein_m"],
                    "diametre_exterieur_m": computed["diametre_exterieur_m"],
                    "profondeur_fouille_m": computed["profondeur_fouille_m"],
                    "hauteur_interieure_m": computed["hauteur_interieure_m"],
                    "hauteur_utile_m": computed["hauteur_utile_m"],
                    "actif": actif,
                },
            )
            created_or_updated += 1
        except Exception as e:
            errors.append(f"Ligne {row_index}: {e}")

    if created_or_updated:
        messages.success(request, f"Import citerne termine: {created_or_updated} ligne(s) enregistree(s), {skipped} vide(s).")
    if errors:
        preview = " | ".join(errors[:5])
        suffix = "" if len(errors) <= 5 else f" (+{len(errors) - 5} autres)"
        messages.error(request, f"Erreurs import: {preview}{suffix}")
    if not created_or_updated and not errors:
        messages.warning(request, "Aucune donnee importee.")

    return redirect("chantiers:workflow_ui_citerne_config_list")


@require_http_methods(["POST"])
def workflow_ui_generate_besoins(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        created = generer_besoins_chantier(workflow, reset_auto=True)
        messages.success(request, f"Besoins générés: {created}.")
    except Exception as e:
        messages.error(request, f"Erreur génération besoins: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_add_surplus(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        type_ligne = request.POST.get("type_ligne")
        ChantierSurplusLigne.objects.create(
            chantier=workflow,
            type_ligne=type_ligne,
            designation_id=request.POST.get("designation_id") or None,
            poste_main_oeuvre_id=request.POST.get("poste_main_oeuvre_id") or None,
            quantite=_d(request.POST.get("quantite")),
            prix_unitaire=_d(request.POST.get("prix_unitaire")) if request.POST.get("prix_unitaire") else None,
            justification=request.POST.get("justification", ""),
        )
        messages.success(request, "Surplus ajouté.")
    except Exception as e:
        messages.error(request, f"Erreur ajout surplus: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_integrate_surplus(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        with transaction.atomic():
            workflow.besoins.filter(origine="surplus").delete()
            created = 0
            for s in workflow.surplus.all():
                ChantierBesoinLigne.objects.create(
                    chantier=workflow,
                    type_ligne=s.type_ligne,
                    designation=s.designation,
                    poste_main_oeuvre=s.poste_main_oeuvre,
                    quantite=s.quantite,
                    prix_unitaire=s.prix_unitaire,
                    origine="surplus",
                )
                created += 1
        messages.success(request, f"Surplus intégrés: {created}.")
    except Exception as e:
        messages.error(request, f"Erreur intégration surplus: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_generate_devis(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        total = Decimal("0")
        for b in workflow.besoins.all():
            if b.prix_unitaire is not None:
                total += Decimal(b.quantite) * Decimal(b.prix_unitaire)
        idx = workflow.documents.filter(type_document="devis").count() + 1
        doc = ChantierDocument.objects.create(
            chantier=workflow,
            type_document="devis",
            numero=f"DEV-{workflow.code}-{idx:03d}",
            montant=total,
            statut="brouillon",
        )
        workflow.statut = "devis"
        workflow.save(update_fields=["statut"])
        messages.success(request, f"Devis généré: {doc.numero}.")
    except Exception as e:
        messages.error(request, f"Erreur devis: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_create_contrat(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        doc = ChantierDocument.objects.create(
            chantier=workflow,
            type_document="contrat",
            numero=(request.POST.get("numero") or "").strip(),
            tacheron_id=request.POST.get("tacheron_id") or None,
            montant=_d(request.POST.get("montant"), default="0"),
            statut="brouillon",
        )
        workflow.statut = "contrat"
        workflow.save(update_fields=["statut"])
        messages.success(request, f"Contrat créé: {doc.numero}.")
    except Exception as e:
        messages.error(request, f"Erreur contrat: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_create_avenant(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        parent = get_object_or_404(ChantierDocument, id=request.POST.get("contrat_document_id"), chantier=workflow, type_document="contrat")
        doc = ChantierDocument.objects.create(
            chantier=workflow,
            type_document="avenant",
            numero=(request.POST.get("numero") or "").strip(),
            parent=parent,
            tacheron=parent.tacheron,
            montant=_d(request.POST.get("montant"), default="0"),
            statut="brouillon",
        )
        messages.success(request, f"Avenant créé: {doc.numero}.")
    except Exception as e:
        messages.error(request, f"Erreur avenant: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_create_facture_tranche(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        parent = get_object_or_404(ChantierDocument, id=request.POST.get("parent_document_id"), chantier=workflow)
        doc = ChantierDocument.objects.create(
            chantier=workflow,
            type_document="facture_tranche",
            numero=(request.POST.get("numero") or "").strip(),
            parent=parent,
            tacheron=parent.tacheron,
            montant=_d(request.POST.get("montant"), default="0"),
            statut="brouillon",
        )
        messages.success(request, f"Facture tranche créée: {doc.numero}.")
    except Exception as e:
        messages.error(request, f"Erreur facture tranche: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_add_mouvement(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        ChantierMouvement.objects.create(
            chantier=workflow,
            besoin_id=request.POST.get("besoin_id"),
            type_mouvement=request.POST.get("type_mouvement"),
            quantite=_d(request.POST.get("quantite")),
            reference_document=request.POST.get("reference_document", ""),
            date_mouvement=timezone.now(),
        )
        messages.success(request, "Mouvement ajouté.")
    except Exception as e:
        messages.error(request, f"Erreur mouvement: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)


@require_http_methods(["POST"])
def workflow_ui_generate_devis_final(request, workflow_id):
    workflow = get_object_or_404(ChantierWorkflow, id=workflow_id)
    try:
        recollement_lines, _tp, _tu, total_valeur = _build_recollement(workflow)
        if total_valeur == 0:
            for line in recollement_lines:
                b = line["besoin"]
                if b.prix_unitaire is not None:
                    total_valeur += Decimal(b.quantite) * Decimal(b.prix_unitaire)
        idx = workflow.documents.filter(type_document="devis_final").count() + 1
        doc = ChantierDocument.objects.create(
            chantier=workflow,
            type_document="devis_final",
            numero=f"DFIN-{workflow.code}-{idx:03d}",
            montant=total_valeur,
            statut="brouillon",
        )
        workflow.statut = "recollement"
        workflow.save(update_fields=["statut"])
        messages.success(request, f"Devis final généré: {doc.numero}.")
    except Exception as e:
        messages.error(request, f"Erreur devis final: {e}")
    return redirect("chantiers:workflow_ui_detail", workflow_id=workflow.id)

