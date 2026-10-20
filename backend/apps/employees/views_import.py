from datetime import datetime
from django.shortcuts import render, redirect
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.utils.safestring import mark_safe
from .models import Poste, Employe, Contrat, GrilleSalaire, Salaire
from apps.organization.models import Organisation, Base

# Import pour Excel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment


def parse_date_flexible(date_str):
    """
    Parse une date depuis différents formats acceptés.
    Retourne un objet date ou None en cas d'échec.
    Formats acceptés : DD/MM/YYYY, YYYY-MM-DD, YYYY-MM-DD HH:MM:SS
    """
    if not date_str or not str(date_str).strip():
        return None

    # Si c'est déjà un objet datetime (depuis Excel)
    if isinstance(date_str, datetime):
        return date_str.date()

    date_str = str(date_str).strip()

    # Essayer le format avec timestamp : YYYY-MM-DD HH:MM:SS
    try:
        return datetime.strptime(date_str, '%Y-%m-%d %H:%M:%S').date()
    except ValueError:
        pass

    # Essayer le format DD/MM/YYYY (format français)
    try:
        return datetime.strptime(date_str, '%d/%m/%Y').date()
    except ValueError:
        pass

    # Essayer le format YYYY-MM-DD (format ISO)
    try:
        return datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        pass

    return None






def export_excel_data(request, model_name):
    """Export complet des données employés en Excel"""
    models_export = {
        'postes': {
            'model': Poste,
            'fields': ['id', 'nom_poste', 'organisation__nom_organisation', 'organisation_id', 'tarif_indemnite', 'prefixe', 'compte_salaire__nom_compte', 'compte_salaire_id', 'compte_irsa__nom_compte', 'compte_irsa_id', 'compte_cnaps_fmfp_sante__nom_compte', 'compte_cnaps_fmfp_sante_id'],
            'filename': 'export_postes.xlsx',
            'sheet_name': 'Postes'
        },
        'employes': {
            'model': Employe,
            'fields': ['id', 'nom', 'prenom', 'surnom', 'base__organisation__nom_organisation', 'base__organisation_id', 'base__nom_base', 'base_id', 'poste__nom_poste', 'poste_id', 'numero_employe', 'date_naissance', 'lieu_naissance', 'telephone', 'email', 'adresse', 'cin', 'date_embauche', 'date_debauchage', 'personnes_charge', 'actif'],
            'filename': 'export_employes.xlsx',
            'sheet_name': 'Employés'
        },
        'contrats': {
            'model': Contrat,
            'fields': ['id', 'employe_id', 'debut_contrat', 'fin_contrat', 'categorie_professionnel', 'type_contrat'],
            'filename': 'export_contrats.xlsx',
            'sheet_name': 'Contrats'
        },
        'grilles_salaire': {
            'model': GrilleSalaire,
            'fields': [
                'id',
                'contrat__employe__nom',
                'contrat__employe__prenom',
                'contrat_id',
                'date_salaire',
                'personnes_charge',
                'nb_jours_presence',
                'nb_jours_feries_gardien',
                'nb_heures_remplacement_gardien',
                'salaire_brut_base',
                'salaire_brut_total',  # Calculé automatiquement
                'charge_cnaps_employe',  # Calculé automatiquement (1% du salaire brut)
                'irsa',  # Calculé automatiquement (selon barème IRSA)
                'mutuelle_sante',  # Calculé automatiquement (1% de base imposable)
                'salaire_net'  # Calculé automatiquement
            ],
            'filename': 'export_grilles_salaire.xlsx',
            'sheet_name': 'Grilles Salaire',
            'help_texts': {
                'salaire_brut_total': 'JOURNALIER: brut×jours | GARDIEN: brut+majorations | AUTRE: brut seul',
                'charge_cnaps_employe': '1% × min(brut_total, 8×sal_min) - Plafond CNaPS appliqué',
                'irsa': 'Barème progressif sur (brut_total - CNaPS - déductions)',
                'mutuelle_sante': '1% × (brut_total - CNaPS) = 1% de la base imposable',
                'salaire_net': 'brut_total - CNaPS - IRSA - Mutuelle Santé'
            },
            'show_formulas_sheet': True  # Afficher une feuille de documentation
        },
        'salaires': {
            'model': Salaire,
            'fields': [
                'id',
                'grille_salaire__contrat__employe__nom',
                'grille_salaire__contrat__employe__prenom',
                'grille_salaire_id',
                'date_paiement',
                'salaire_net_paye',
                'irsa_paye',
                'cnaps_paye',
                'fmfp_paye',
                'mutuelle_sante_paye',
                'numero_diamex_salaire',
                'base_diamex_salaire',
                'numero_diamex_irsa',
                'base_diamex_irsa',
                'numero_diamex_cnaps',
                'base_diamex_cnaps',
                'numero_diamex_fmfp',
                'base_diamex_fmfp',
                'numero_diamex_mutuelle_sante',
                'base_diamex_mutuelle_sante'
            ],
            'filename': 'export_salaires.xlsx',
            'sheet_name': 'Salaires',
            'help_texts': {
                'salaire_net_paye': 'Calculé depuis salaire net arrondi de la grille (arrondi à 100 Ar près)',
                'irsa_paye': 'Calculé depuis la grille',
                'cnaps_paye': '14 fois la charge CNaPS employé de la grille',
                'fmfp_paye': 'Égal à la charge CNaPS employé de la grille',
                'mutuelle_sante_paye': 'Calculé depuis la grille'
            }
        }
    }

    if model_name not in models_export:
        return JsonResponse({'error': 'Type d\'export non reconnu'}, status=404)

    config = models_export[model_name]
    model_class = config['model']
    fields = config['fields']

    # Créer un workbook Excel
    wb = Workbook()
    ws = wb.active
    ws.title = config['sheet_name']

    # Créer les en-têtes avec des noms plus clairs
    headers = []
    for field in fields:
        if field == 'id':
            headers.append('ID')
        elif '__' in field:
            if field.endswith('__nom_organisation'):
                headers.append('Organisation')
            elif field.endswith('__nom_base'):
                headers.append('Base')
            elif field.endswith('__nom_poste'):
                headers.append('Poste')
            elif field.endswith('__nom'):
                headers.append('Nom Employé')
            elif field.endswith('__prenom'):
                headers.append('Prénom Employé')
            elif field.endswith('__nom_compte'):
                field_name = field.split('__')[0].replace('_', ' ').title()
                headers.append(f'{field_name} (Nom)')
            else:
                headers.append(field.replace('__', ' - ').replace('_', ' ').title())
        elif field.endswith('_id'):
            field_name = field.replace('_id', '').replace('_', ' ').title()
            headers.append(f'{field_name} (ID)')
        else:
            headers.append(field.replace('_', ' ').title())

    # Styliser les en-têtes
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Ajouter les en-têtes avec commentaires explicatifs
    help_texts = config.get('help_texts', {})

    for col_num, (header, field) in enumerate(zip(headers, fields), 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

        # Ajouter un commentaire explicatif si disponible
        if field in help_texts:
            comment = Comment(help_texts[field], "Système")
            cell.comment = comment

    # Exporter toutes les données
    try:
        if model_name == 'postes':
            queryset = model_class.objects.select_related('organisation', 'compte_salaire', 'compte_irsa', 'compte_cnaps_fmfp_sante').all()
        elif model_name == 'employes':
            queryset = model_class.objects.select_related('base__organisation', 'base', 'poste').all()
        elif model_name == 'contrats':
            queryset = model_class.objects.select_related('employe').all()
        elif model_name == 'grilles_salaire':
            queryset = model_class.objects.select_related('contrat__employe').all()
        elif model_name == 'salaires':
            queryset = model_class.objects.select_related('grille_salaire__contrat__employe').all()

        row_num = 2
        for obj in queryset:
            for col_num, field in enumerate(fields, 1):
                try:
                    if '__' in field:
                        field_parts = field.split('__')
                        value = obj
                        for part in field_parts:
                            if value is None:
                                break
                            value = getattr(value, part, None)
                        ws.cell(row=row_num, column=col_num, value=value if value is not None else '')
                    else:
                        value = getattr(obj, field, '')
                        if value is None:
                            ws.cell(row=row_num, column=col_num, value='')
                        elif isinstance(value, bool):
                            ws.cell(row=row_num, column=col_num, value='Oui' if value else 'Non')
                        elif hasattr(value, 'strftime'):  # Date ou datetime
                            ws.cell(row=row_num, column=col_num, value=value.strftime('%d/%m/%Y'))
                        else:
                            ws.cell(row=row_num, column=col_num, value=str(value))
                except Exception:
                    ws.cell(row=row_num, column=col_num, value='')
            row_num += 1

        # Auto-ajuster la largeur des colonnes
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Ajouter une feuille de documentation pour grilles_salaire
        if model_name == 'grilles_salaire':
            ws_doc = wb.create_sheet("📘 Formules de calcul")
            ws_doc.sheet_properties.tabColor = "4472C4"  # Bleu

            # Style du titre
            title_font = Font(bold=True, size=14, color="FFFFFF")
            title_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

            # Titre
            ws_doc['A1'] = '📘 DOCUMENTATION DES FORMULES DE CALCUL - GRILLES DE SALAIRE'
            ws_doc['A1'].font = title_font
            ws_doc['A1'].fill = title_fill
            ws_doc.merge_cells('A1:D1')

            # Contenu de la documentation
            doc_lines = [
                '',
                '1️⃣ SALAIRE BRUT TOTAL (colonne K):',
                '   • JOURNALIER: salaire_brut_base × nb_jours_presence',
                '   • GARDIEN: salaire_brut_base + majorations (détails ci-dessous)',
                '       - Taux horaire gardien: salaire brut ÷ 240',
                '       - Majoration nuit: taux horaire × heures de nuit × 30%',
                '       - Majoration dimanche: taux horaire × heures du dimanche × 40%',
                '       - Jours fériés: taux horaire × heures fériées × 50%',
                '       - 1er mai: taux horaire × heures travaillées × 100%',
                '       - Remplacement: taux horaire × heures de remplacement × 30%',
                '   • AUTRE POSTE AVEC MAJORATIONS PLANNING:',
                '       - Taux horaire: salaire brut ÷ 173,33',
                '       - Dimanche: 8h × taux horaire × 40%',
                '       - Jour férié: 8h × taux horaire × 50%',
                '       - 1er mai: 8h × taux horaire × 100%',
                '',
                '2️⃣ CHARGE CNaPS EMPLOYÉ (colonne L):',
                '   • Formule: 1% × min(salaire_brut_total, 8 × salaire_minimum)',
                '   • Le salaire minimum est configuré dans ConfigSalaire',
                '   • Il y a un plafond CNaPS = 8 fois le salaire minimum',
                '',
                '3️⃣ IRSA - Impôt sur Revenu Salarié (colonne M):',
                '   • Base imposable = salaire_brut_total - charge_cnaps_employe',
                '   • Déductions = nb_personnes_charge × déduction_par_personne',
                '   • Base taxable = base_imposable - déductions',
                '   • IRSA = application du barème progressif par tranches (voir ConfigTrancheIRSA)',
                '   • Exemple barème: 0% jusqu\'à X, puis 5%, 10%, 15%, 20%',
                '',
                '4️⃣ MUTUELLE SANTÉ (colonne N):',
                '   • Formule: 1% × (salaire_brut_total - charge_cnaps_employe)',
                '   • Équivalent: 1% × base_imposable',
                '',
                '5️⃣ SALAIRE NET (colonne O):',
                '   • Formule: salaire_brut_total - charge_cnaps_employe - irsa - mutuelle_sante',
                '',
                '⚙️ CONFIGURATION SYSTÈME (voir table ConfigSalaire):',
                '   • salaire_minimum: Salaire minimum légal de référence',
                '   • deduction_irsa: Montant de déduction IRSA par personne à charge',
                '   • Barème IRSA: Tranches progressives dans ConfigTrancheIRSA',
                '',
                '⚠️ IMPORTANT:',
                '   • Les colonnes K, L, M, N, O sont CALCULÉES AUTOMATIQUEMENT',
                '   • Ne modifiez PAS ces valeurs manuellement dans Excel',
                '   • Elles seront recalculées lors de l\'enregistrement dans le système',
                '   • Modifiez uniquement les colonnes de saisie (F à J)',
                '',
                '📊 COLONNES DE SAISIE (à remplir):',
                '   F - Personnes à charge: Nombre de personnes pour déduction IRSA',
                '   G - Nb jours journalier: Si contrat journalier, sinon 0',
                '   H - Nb jours fériés gardien: Si gardien travaille les jours fériés',
                '   I - Nb heures remplacement gardien: Heures de remplacement',
                '   J - Salaire brut de base: Salaire de référence (mensuel ou journalier)'
            ]

            row = 3
            for line in doc_lines:
                ws_doc[f'A{row}'] = line
                if line.startswith(('1️⃣', '2️⃣', '3️⃣', '4️⃣')):
                    ws_doc[f'A{row}'].font = Font(bold=True, size=12, color="2F5496")
                elif line.startswith('⚙️') or line.startswith('⚠️') or line.startswith('📊'):
                    ws_doc[f'A{row}'].font = Font(bold=True, size=11, color="C65911")
                row += 1

            # Largeur de la colonne
            ws_doc.column_dimensions['A'].width = 100

    except Exception as e:
        return JsonResponse({'error': f'Erreur lors de l\'export: {str(e)}'}, status=500)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{config["filename"]}"'
    wb.save(response)
    return response


def download_excel_template(request, model_name):
    """Téléchargement des templates Excel pour les employés"""
    templates = {
        'postes': {
            'fields': ['nom_poste', 'organisation', 'tarif_indemnite', 'prefixe', 'compte_salaire_nom', 'compte_irsa_nom', 'compte_cnaps_nom'],
            'examples': [
                ['Directeur', 'JIRAMA', '500000', 'DIR', 'Salaires personnel', 'IRSA', 'CNAPS'],
                ['Technicien', 'JIRAMA', '250000', 'TECH', 'Salaires personnel', 'IRSA', 'CNAPS']
            ]
        },
        'employes': {
            'fields': ['nom', 'prenom', 'surnom', 'organisation', 'base', 'poste', 'numero_employe', 'date_naissance', 'lieu_naissance', 'telephone', 'email', 'adresse', 'cin', 'date_embauche', 'date_debauchage', 'personnes_charge', 'actif'],
            'examples': [
                ['Rakoto', 'Jean', 'Koto', 'JIRAMA', 'Base Technique', 'Directeur', 'EMP001', '15/01/1980', 'Antananarivo', '0321234567', 'rakoto@email.com', 'Lot II Ambatobe', '123456789', '01/03/2020', '', '3', 'Oui'],
                ['Rasoa', 'Marie', 'Soa', 'JIRAMA', 'Base Technique', 'Technicien', 'EMP002', '20/05/1985', 'Fianarantsoa', '0347654321', 'rasoa@email.com', 'Rue de la paix', '987654321', '15/06/2021', '31/12/2023', '0', 'Non']
            ]
        },
        'contrats': {
            'fields': ['employe_id', 'debut_contrat', 'fin_contrat', 'categorie_professionnel', 'type_contrat'],
            'examples': [
                ['12345678-1234-1234-1234-123456789012', '01/03/2020', '', '3A', 'Temps plein'],
                ['23456789-2345-2345-2345-234567890123', '15/06/2021', '14/06/2024', 'M1', 'Temps partiel'],
                ['34567890-3456-3456-3456-345678901234', '01/01/2023', '', 'OP2B', 'Mi-temps'],
                ['45678901-4567-4567-4567-456789012345', '10/05/2024', '30/09/2024', 'OP3', 'Journalier'],
                ['56789012-5678-5678-5678-567890123456', '15/02/2023', '', 'HC', 'Temps plein']
            ]
        },
        'grilles_salaire': {
            'fields': ['contrat_id', 'date_salaire', 'personnes_charge', 'nb_jours_presence', 'nb_jours_feries_gardien', 'nb_heures_remplacement_gardien', 'salaire_brut_base'],
            'examples': [
                ['12345678-1234-1234-1234-123456789012', '01/01/2024', '2', '0', '0', '0', '500000'],
                ['23456789-2345-2345-2345-234567890123', '01/01/2024', '0', '22', '2', '8.5', '250000']
            ]
        },
        'salaires': {
            'fields': [
                'grille_salaire_id', 'date_paiement',
                'numero_diamex_salaire', 'base_diamex_salaire',
                'numero_diamex_irsa', 'base_diamex_irsa',
                'numero_diamex_cnaps', 'base_diamex_cnaps',
                'numero_diamex_fmfp', 'base_diamex_fmfp',
                'numero_diamex_mutuelle_sante', 'base_diamex_mutuelle_sante'
            ],
            'examples': [
                ['12345678-1234-1234-1234-123456789012', '31/01/2025', '', '', '', '', '', '', '', '', '', ''],
                ['23456789-2345-2345-2345-234567890123', '28/02/2025', '', '', '', '', '', '', '', '', '', '']
            ]
        }
    }

    if model_name not in templates:
        return JsonResponse({'error': 'Template non trouvé'}, status=404)

    template = templates[model_name]

    wb = Workbook()
    ws = wb.active
    ws.title = f'Template {model_name.title()}'

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = []
    for field in template['fields']:
        headers.append(field.replace('_', ' ').title())

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_num, example in enumerate(template['examples'], 2):
        for col_num, value in enumerate(example, 1):
            ws.cell(row=row_num, column=col_num, value=value)

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="template_{model_name}.xlsx"'
    wb.save(response)
    return response


def import_excel_view(request, model_name):
    """Vue générique pour l'import Excel des employés"""
    models = {
        'postes': Poste,
        'employes': Employe,
        'contrats': Contrat,
        'grilles_salaire': GrilleSalaire,
        'salaires': Salaire
    }

    if model_name not in models:
        messages.error(request, f'Type de données non reconnu: {model_name}')
        return redirect('employees:list')

    if request.method == 'POST':
        if 'excel_file' not in request.FILES:
            messages.error(request, 'Aucun fichier sélectionné.')
            return redirect(request.META.get('HTTP_REFERER', 'employees:list'))

        excel_file = request.FILES['excel_file']

        if not (excel_file.name.endswith('.xlsx') or excel_file.name.endswith('.xls')):
            messages.error(request, 'Le fichier doit être un fichier Excel.')
            return redirect(request.META.get('HTTP_REFERER', 'employees:list'))

        try:
            wb = load_workbook(excel_file)
            ws = wb.active

            headers = []
            for cell in ws[1]:
                headers.append(cell.value.lower().replace(' ', '_') if cell.value else '')

            created_count = 0
            error_count = 0
            error_details = []  # Collecter les détails des erreurs

            for row_num in range(2, ws.max_row + 1):
                try:
                    row_data = {}
                    for col_num, header in enumerate(headers, 1):
                        cell_value = ws.cell(row=row_num, column=col_num).value
                        row_data[header] = str(cell_value).strip() if cell_value is not None else ''

                    if model_name == 'postes':
                        organisation = Organisation.objects.get(nom_organisation=row_data.get('organisation', '').strip())
                        Poste.objects.get_or_create(
                            nom_poste=row_data.get('nom_poste', '').strip(),
                            organisation=organisation,
                            defaults={
                                'tarif_indemnite': float(row_data.get('tarif_indemnite', 0)) if row_data.get('tarif_indemnite', '').strip() else None,
                                'prefixe': row_data.get('prefixe', '').strip(),
                            }
                        )
                    elif model_name == 'employes':
                        organisation = Organisation.objects.get(nom_organisation=row_data.get('organisation', '').strip())
                        base = Base.objects.get(nom_base=row_data.get('base', '').strip()) if row_data.get('base', '').strip() else None
                        poste = Poste.objects.get(nom_poste=row_data.get('poste', '').strip()) if row_data.get('poste', '').strip() else None

                        date_naissance = parse_date_flexible(row_data.get('date_naissance', ''))
                        date_embauche = parse_date_flexible(row_data.get('date_embauche', ''))
                        date_debauchage = parse_date_flexible(row_data.get('date_debauchage', ''))

                        # Parser le champ actif (Oui/Non ou True/False)
                        actif_str = row_data.get('actif', '').strip().lower()
                        actif = actif_str in ['oui', 'true', '1', 'yes', 'vrai'] if actif_str else True

                        # Parser le champ personnes_charge
                        personnes_charge_str = row_data.get('personnes_charge', '0').strip()
                        try:
                            personnes_charge = int(personnes_charge_str) if personnes_charge_str else 0
                        except ValueError:
                            personnes_charge = 0

                        Employe.objects.get_or_create(
                            numero_employe=row_data.get('numero_employe', '').strip(),
                            defaults={
                                'nom': row_data.get('nom', '').strip(),
                                'prenom': row_data.get('prenom', '').strip(),
                                'surnom': row_data.get('surnom', '').strip(),
                                'organisation': organisation,
                                'base': base,
                                'poste': poste,
                                'date_naissance': date_naissance,
                                'lieu_naissance': row_data.get('lieu_naissance', '').strip(),
                                'telephone': row_data.get('telephone', '').strip(),
                                'email': row_data.get('email', '').strip(),
                                'adresse': row_data.get('adresse', '').strip(),
                                'cin': row_data.get('cin', '').strip(),
                                'date_embauche': date_embauche,
                                'date_debauchage': date_debauchage,
                                'personnes_charge': personnes_charge,
                                'actif': actif,
                            }
                        )
                    elif model_name == 'contrats':
                        # Debug : afficher les clés disponibles pour la première ligne
                        if row_num == 2:
                            print(f"DEBUG: Headers trouvés = {headers}")
                            print(f"DEBUG: Données première ligne = {list(row_data.keys())}")
                            print(f"DEBUG: Valeur employe_id = '{row_data.get('employe_id', 'NOT FOUND')}'")

                        # Récupérer l'employé par ID
                        employe_id = row_data.get('employe_id', '').strip()
                        if not employe_id:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: employe_id manquant (colonnes disponibles: {', '.join(headers)})")
                            continue  # Passer cette ligne si pas d'employe_id

                        try:
                            employe = Employe.objects.get(id=employe_id)
                        except Employe.DoesNotExist:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Employé '{employe_id}' introuvable en base")
                            continue  # Passer si employé n'existe pas
                        except Exception as e:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Erreur validation ID '{employe_id}': {str(e)}")
                            continue

                        # Parser les dates
                        debut_contrat = parse_date_flexible(row_data.get('debut_contrat', ''))
                        fin_contrat = parse_date_flexible(row_data.get('fin_contrat', ''))

                        if not debut_contrat:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Date début manquante ou invalide (valeur: '{row_data.get('debut_contrat', '')}')")
                            continue  # Date de début obligatoire

                        # Créer ou récupérer le contrat
                        Contrat.objects.get_or_create(
                            employe=employe,
                            debut_contrat=debut_contrat,
                            defaults={
                                'fin_contrat': fin_contrat,
                                'categorie_professionnel': row_data.get('categorie_professionnel', '').strip(),
                                'type_contrat': row_data.get('type_contrat', '').strip(),
                            }
                        )
                        created_count += 1

                    elif model_name == 'grilles_salaire':
                        # Debug : afficher les clés disponibles pour la première ligne
                        if row_num == 2:
                            print(f"DEBUG GRILLES: Headers trouvés = {headers}")
                            print(f"DEBUG GRILLES: Données première ligne = {list(row_data.keys())}")
                            print(f"DEBUG GRILLES: Valeur contrat_id = '{row_data.get('contrat_id', 'NOT FOUND')}'")

                        # Récupérer le contrat par ID
                        contrat_id = row_data.get('contrat_id', '').strip()
                        if not contrat_id:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: contrat_id manquant (colonnes disponibles: {', '.join(headers)})")
                            continue

                        try:
                            contrat = Contrat.objects.get(id=contrat_id)
                        except Contrat.DoesNotExist:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Contrat '{contrat_id}' introuvable en base")
                            continue
                        except Exception as e:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Erreur validation contrat_id '{contrat_id}': {str(e)}")
                            continue

                        # Parser la date de salaire
                        date_salaire = parse_date_flexible(row_data.get('date_salaire', ''))
                        if not date_salaire:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Date salaire manquante ou invalide (valeur: '{row_data.get('date_salaire', '')}')")
                            continue

                        # Parser les champs numériques
                        try:
                            personnes_charge = int(row_data.get('personnes_charge', '0').strip() or '0')
                        except ValueError:
                            personnes_charge = 0

                        try:
                            nb_jours_presence = int(row_data.get('nb_jours_presence', '0').strip() or '0')
                        except ValueError:
                            nb_jours_presence = 0

                        try:
                            nb_jours_feries_gardien = int(row_data.get('nb_jours_feries_gardien', '0').strip() or '0')
                        except ValueError:
                            nb_jours_feries_gardien = 0

                        try:
                            nb_heures_remplacement_gardien = float(row_data.get('nb_heures_remplacement_gardien', '0').strip() or '0')
                        except ValueError:
                            nb_heures_remplacement_gardien = 0.0

                        try:
                            salaire_brut_base = float(row_data.get('salaire_brut_base', '0').strip() or '0')
                        except ValueError:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Salaire brut base invalide (valeur: '{row_data.get('salaire_brut_base', '')}')")
                            continue

                        # Créer ou récupérer la grille de salaire
                        GrilleSalaire.objects.get_or_create(
                            contrat=contrat,
                            date_salaire=date_salaire,
                            defaults={
                                'personnes_charge': personnes_charge,
                                'nb_jours_presence': nb_jours_presence,
                                'nb_jours_feries_gardien': nb_jours_feries_gardien,
                                'nb_heures_remplacement_gardien': nb_heures_remplacement_gardien,
                                'salaire_brut_base': salaire_brut_base,
                            }
                        )
                        created_count += 1

                    elif model_name == 'salaires':
                        # Debug : afficher les clés disponibles pour la première ligne
                        if row_num == 2:
                            print(f"DEBUG SALAIRES: Headers trouvés = {headers}")
                            print(f"DEBUG SALAIRES: Données première ligne = {list(row_data.keys())}")
                            print(f"DEBUG SALAIRES: Valeur grille_salaire_id = '{row_data.get('grille_salaire_id', 'NOT FOUND')}'")

                        # Récupérer la grille de salaire par ID
                        grille_salaire_id = row_data.get('grille_salaire_id', '').strip()
                        if not grille_salaire_id:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: grille_salaire_id manquant (colonnes disponibles: {', '.join(headers)})")
                            continue

                        try:
                            grille = GrilleSalaire.objects.get(id=grille_salaire_id)
                        except GrilleSalaire.DoesNotExist:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Grille salaire '{grille_salaire_id}' introuvable en base")
                            continue
                        except Exception as e:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Erreur validation grille_salaire_id '{grille_salaire_id}': {str(e)}")
                            continue

                        # Parser la date de paiement
                        date_paiement = parse_date_flexible(row_data.get('date_paiement', ''))
                        if not date_paiement:
                            error_count += 1
                            error_details.append(f"Ligne {row_num}: Date paiement manquante ou invalide (valeur: '{row_data.get('date_paiement', '')}')")
                            continue

                        # Calculer automatiquement les montants depuis la grille
                        from decimal import Decimal
                        import math

                        # Salaire net arrondi à 100 Ar près (arrondi supérieur)
                        salaire_net_arrondi = Decimal(math.ceil(float(grille.salaire_net) / 100) * 100)

                        # IRSA depuis la grille
                        irsa_paye = grille.irsa

                        # CNaPS = 14 × charge CNaPS employé
                        cnaps_paye = grille.charge_cnaps_employe * Decimal('14')

                        # FMFP = charge CNaPS employé
                        fmfp_paye = grille.charge_cnaps_employe

                        # Mutuelle santé depuis la grille
                        mutuelle_sante_paye = grille.mutuelle_sante

                        # Récupérer les champs DiameX optionnels (peuvent être vides)
                        numero_diamex_salaire = row_data.get('numero_diamex_salaire', '').strip() or None
                        base_diamex_salaire = row_data.get('base_diamex_salaire', '').strip() or None
                        numero_diamex_irsa = row_data.get('numero_diamex_irsa', '').strip() or None
                        base_diamex_irsa = row_data.get('base_diamex_irsa', '').strip() or None
                        numero_diamex_cnaps = row_data.get('numero_diamex_cnaps', '').strip() or None
                        base_diamex_cnaps = row_data.get('base_diamex_cnaps', '').strip() or None
                        numero_diamex_fmfp = row_data.get('numero_diamex_fmfp', '').strip() or None
                        base_diamex_fmfp = row_data.get('base_diamex_fmfp', '').strip() or None
                        numero_diamex_mutuelle_sante = row_data.get('numero_diamex_mutuelle_sante', '').strip() or None
                        base_diamex_mutuelle_sante = row_data.get('base_diamex_mutuelle_sante', '').strip() or None

                        # Créer ou récupérer le salaire
                        Salaire.objects.get_or_create(
                            grille_salaire=grille,
                            date_paiement=date_paiement,
                            defaults={
                                'salaire_net_paye': salaire_net_arrondi,
                                'irsa_paye': irsa_paye,
                                'cnaps_paye': cnaps_paye,
                                'fmfp_paye': fmfp_paye,
                                'mutuelle_sante_paye': mutuelle_sante_paye,
                                'numero_diamex_salaire': numero_diamex_salaire,
                                'base_diamex_salaire': base_diamex_salaire,
                                'numero_diamex_irsa': numero_diamex_irsa,
                                'base_diamex_irsa': base_diamex_irsa,
                                'numero_diamex_cnaps': numero_diamex_cnaps,
                                'base_diamex_cnaps': base_diamex_cnaps,
                                'numero_diamex_fmfp': numero_diamex_fmfp,
                                'base_diamex_fmfp': base_diamex_fmfp,
                                'numero_diamex_mutuelle_sante': numero_diamex_mutuelle_sante,
                                'base_diamex_mutuelle_sante': base_diamex_mutuelle_sante,
                            }
                        )
                        created_count += 1

                    else:
                        # Pour les autres modèles non gérés
                        created_count += 1

                except Exception as e:
                    error_count += 1
                    import traceback
                    error_msg = f"Ligne {row_num}: Exception {str(e)}"
                    error_details.append(error_msg)
                    print(error_msg)
                    traceback.print_exc()

            if created_count > 0:
                messages.success(request, f'{created_count} enregistrement(s) importé(s) avec succès depuis Excel.')
            if error_count > 0:
                # Afficher les 10 premières erreurs
                error_sample = error_details[:10]
                error_message = f'{error_count} erreur(s) rencontrée(s) lors de l\'import Excel.<br><br>Détails des erreurs:<br>' + '<br>'.join(error_sample)
                if len(error_details) > 10:
                    error_message += f'<br>... et {len(error_details) - 10} autres erreurs'
                messages.warning(request, mark_safe(error_message))
            if created_count == 0 and error_count == 0:
                messages.info(request, 'Aucun enregistrement trouvé dans le fichier Excel.')

        except Exception as e:
            messages.error(request, f'Erreur lors de la lecture du fichier Excel: {str(e)}')

    return redirect(request.META.get('HTTP_REFERER', 'employees:list'))
