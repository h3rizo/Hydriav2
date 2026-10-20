"""
Vues pour la gestion des paiements de mutuelle santé (AMIT/OSIEFSoVa)
"""
from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from django.db.models import Sum
from decimal import Decimal
import json
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.http import HttpResponse

from .models import PaiementMutuelleSante, PaiementMutuelleSanteDetail, Employe, Salaire
from .salary_consistency import sync_salaire_with_grille
from apps.organization.models import Organisation


@login_required
def paiements_mutuelle_liste(request):
    """
    Page de liste des paiements de mutuelle santé
    Filtre automatiquement par l'organisation de l'utilisateur connecté
    """
    # Récupérer l'organisation de l'utilisateur connecté
    user_organisation = None
    if hasattr(request.user, 'employe') and request.user.employe:
        if request.user.employe.poste:
            user_organisation = request.user.employe.poste.organisation

    annee_courante = timezone.now().year
    annees = list(range(2020, annee_courante + 2))  # De 2020 à année courante + 1

    context = {
        'user_organisation': user_organisation,
        'annee_courante': annee_courante,
        'annees': annees,
    }

    return render(request, 'employees/paiements_mutuelle_liste.html', context)


@login_required
def api_paiements_mutuelle_list(request):
    """API pour lister les paiements de mutuelle santé"""
    organisation_id = request.GET.get('organisation_id')
    annee = request.GET.get('annee')
    trimestre = request.GET.get('trimestre')

    paiements = PaiementMutuelleSante.objects.all().select_related('organisation', 'created_by')

    if organisation_id:
        paiements = paiements.filter(organisation_id=organisation_id)
    if annee:
        paiements = paiements.filter(annee=int(annee))
    if trimestre:
        paiements = paiements.filter(trimestre=int(trimestre))

    paiements_list = []
    total_amit = Decimal('0')
    total_osiefsova = Decimal('0')

    for p in paiements:
        nb_employes = p.details.count()
        total_amit += p.montant_total_amit
        total_osiefsova += p.montant_total_osiefsova

        paiements_list.append({
            'id': str(p.id),
            'organisation': p.organisation.nom_organisation,
            'organisation_id': str(p.organisation.id),
            'trimestre': p.trimestre,
            'annee': p.annee,
            'periode_label': p.periode_label,
            'montant_total_amit': float(p.montant_total_amit),
            'montant_total_osiefsova': float(p.montant_total_osiefsova),
            'montant_total': float(p.montant_total),
            'statut': p.statut,
            'nb_employes': nb_employes,
            'facture_generee': p.facture_generee,
            'numero_diamex_amit': p.numero_diamex_amit or '',
            'numero_diamex_osiefsova': p.numero_diamex_osiefsova or '',
            'created_at': p.created_at.strftime('%d/%m/%Y %H:%M'),
            'created_by': p.created_by.get_full_name() if p.created_by else '',
        })

    return JsonResponse({
        'success': True,
        'paiements': paiements_list,
        'stats': {
            'total': len(paiements_list),
            'total_amit': float(total_amit),
            'total_osiefsova': float(total_osiefsova),
        }
    })


@csrf_exempt
@require_http_methods(["POST"])
def api_generer_paiement_mutuelle(request):
    """API pour générer un paiement de mutuelle santé pour un trimestre"""
    try:
        data = json.loads(request.body)
        organisation_id = data.get('organisation_id')
        trimestre = int(data.get('trimestre'))
        annee = int(data.get('annee'))

        organisation = get_object_or_404(Organisation, id=organisation_id)

        # Vérifier si un paiement existe déjà
        existing = PaiementMutuelleSante.objects.filter(
            organisation=organisation,
            trimestre=trimestre,
            annee=annee
        ).first()

        if existing:
            return JsonResponse({
                'success': False,
                'error': f'Un paiement existe déjà pour le {existing.periode_label}'
            }, status=400)

        # Créer le paiement
        paiement = PaiementMutuelleSante.objects.create(
            organisation=organisation,
            trimestre=trimestre,
            annee=annee,
            created_by=request.user
        )

        # Calculer les mois du trimestre
        mois_trimestre = paiement.mois_trimestre

        # Récupérer tous les employés de l'organisation avec mutuelle santé
        employes = Employe.objects.filter(
            base__organisation=organisation,
            mutuelle_sante__isnull=False
        ).exclude(mutuelle_sante='').select_related('poste', 'base')

        total_amit = Decimal('0')
        total_osiefsova = Decimal('0')
        nb_details_crees = 0

        for employe in employes:
            # Récupérer les salaires des 3 mois du trimestre
            # Filtrer via la relation: Salaire → GrilleSalaire → Contrat → Employe
            salaires_trimestre = Salaire.objects.filter(
                grille_salaire__contrat__employe=employe,
                date_paiement__year=annee,
                date_paiement__month__in=mois_trimestre,
                mutuelle_sante_paye__isnull=False  # Doit avoir un montant mutuelle payé
            ).select_related('grille_salaire').order_by('date_paiement')

            if not salaires_trimestre.exists():
                continue

            # Calculer les montants par mois
            mois_1_montant = Decimal('0')
            mois_2_montant = Decimal('0')
            mois_3_montant = Decimal('0')

            for salaire in salaires_trimestre:
                sync_salaire_with_grille(salaire)
                montant = salaire.mutuelle_sante_paye or Decimal('0')
                mois_salaire = salaire.date_paiement.month

                if mois_salaire == mois_trimestre[0]:
                    mois_1_montant = montant
                elif mois_salaire == mois_trimestre[1]:
                    mois_2_montant = montant
                elif mois_salaire == mois_trimestre[2]:
                    mois_3_montant = montant

            total_trimestre = mois_1_montant + mois_2_montant + mois_3_montant

            if total_trimestre > 0:
                # Créer le détail
                PaiementMutuelleSanteDetail.objects.create(
                    paiement=paiement,
                    employe=employe,
                    employe_nom_complet=employe.nom_complet,
                    employe_fonction=employe.poste.nom_poste if employe.poste else '',
                    mutuelle=employe.mutuelle_sante,
                    mois_1=mois_1_montant,
                    mois_2=mois_2_montant,
                    mois_3=mois_3_montant,
                    total_trimestre=total_trimestre
                )

                # Cumuler selon la mutuelle
                if employe.mutuelle_sante == 'AMIT':
                    total_amit += total_trimestre
                elif employe.mutuelle_sante == 'OSIEFSoVa':
                    total_osiefsova += total_trimestre

                nb_details_crees += 1

        # Vérifier qu'au moins un employé a été traité
        if nb_details_crees == 0:
            paiement.delete()  # Supprimer le paiement vide
            return JsonResponse({
                'success': False,
                'error': f'Aucun employé avec mutuelle santé trouvé pour cette organisation et période. Vérifiez que les employés ont une mutuelle renseignée et des salaires payés pour les mois du trimestre.'
            }, status=400)

        # Mettre à jour les totaux
        paiement.montant_total_amit = total_amit
        paiement.montant_total_osiefsova = total_osiefsova
        paiement.save()

        return JsonResponse({
            'success': True,
            'message': f'Paiement généré avec succès : {nb_details_crees} employé(s)',
            'paiement_id': str(paiement.id),
            'montant_total_amit': float(total_amit),
            'montant_total_osiefsova': float(total_osiefsova)
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@csrf_exempt
@require_http_methods(["POST"])
def api_valider_paiement_mutuelle(request, paiement_id):
    """API pour valider un paiement de mutuelle santé"""
    try:
        paiement = get_object_or_404(PaiementMutuelleSante, id=paiement_id)

        if paiement.statut == 'exporte':
            return JsonResponse({
                'success': False,
                'error': 'Ce paiement a déjà été exporté vers DiameX'
            }, status=400)

        paiement.statut = 'valide'
        paiement.save()

        return JsonResponse({
            'success': True,
            'message': 'Paiement validé avec succès'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
def generer_facture_mutuelle(request, paiement_id):
    """Génère la facture Excel pour un paiement de mutuelle santé"""
    paiement = get_object_or_404(PaiementMutuelleSante, id=paiement_id)

    # Créer le workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mutuelle Santé"

    # Styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Titre
    ws['A1'] = f"PAIEMENT MUTUELLE SANTÉ - {paiement.organisation.nom_organisation}"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:H1')

    ws['A2'] = f"Période : {paiement.periode_label}"
    ws['A2'].font = Font(bold=True, size=12)
    ws.merge_cells('A2:H2')

    # En-têtes
    headers = ['N°', 'Nom complet', 'Fonction', 'Mutuelle', 'Mois 1', 'Mois 2', 'Mois 3', 'Total Trimestre']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    # Données - Groupe par mutuelle
    details_amit = paiement.details.filter(mutuelle='AMIT').order_by('employe_fonction', 'employe_nom_complet')
    details_osiefsova = paiement.details.filter(mutuelle='OSIEFSoVa').order_by('employe_fonction', 'employe_nom_complet')

    row = 5
    num = 1

    # Section AMIT
    if details_amit.exists():
        ws.cell(row=row, column=1).value = "AMIT"
        ws.cell(row=row, column=1).font = Font(bold=True, size=11)
        ws.merge_cells(f'A{row}:H{row}')
        row += 1

        for detail in details_amit:
            ws.cell(row=row, column=1).value = num
            ws.cell(row=row, column=2).value = detail.employe_nom_complet
            ws.cell(row=row, column=3).value = detail.employe_fonction
            ws.cell(row=row, column=4).value = detail.mutuelle
            ws.cell(row=row, column=5).value = float(detail.mois_1)
            ws.cell(row=row, column=6).value = float(detail.mois_2)
            ws.cell(row=row, column=7).value = float(detail.mois_3)
            ws.cell(row=row, column=8).value = float(detail.total_trimestre)

            # Bordures
            for col in range(1, 9):
                ws.cell(row=row, column=col).border = border

            row += 1
            num += 1

        # Total AMIT
        ws.cell(row=row, column=7).value = "Total AMIT:"
        ws.cell(row=row, column=7).font = Font(bold=True)
        ws.cell(row=row, column=8).value = float(paiement.montant_total_amit)
        ws.cell(row=row, column=8).font = Font(bold=True)
        ws.cell(row=row, column=8).fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        row += 2

    # Section OSIEFSoVa
    if details_osiefsova.exists():
        ws.cell(row=row, column=1).value = "OSIEFSoVa"
        ws.cell(row=row, column=1).font = Font(bold=True, size=11)
        ws.merge_cells(f'A{row}:H{row}')
        row += 1

        for detail in details_osiefsova:
            ws.cell(row=row, column=1).value = num
            ws.cell(row=row, column=2).value = detail.employe_nom_complet
            ws.cell(row=row, column=3).value = detail.employe_fonction
            ws.cell(row=row, column=4).value = detail.mutuelle
            ws.cell(row=row, column=5).value = float(detail.mois_1)
            ws.cell(row=row, column=6).value = float(detail.mois_2)
            ws.cell(row=row, column=7).value = float(detail.mois_3)
            ws.cell(row=row, column=8).value = float(detail.total_trimestre)

            # Bordures
            for col in range(1, 9):
                ws.cell(row=row, column=col).border = border

            row += 1
            num += 1

        # Total OSIEFSoVa
        ws.cell(row=row, column=7).value = "Total OSIEFSoVa:"
        ws.cell(row=row, column=7).font = Font(bold=True)
        ws.cell(row=row, column=8).value = float(paiement.montant_total_osiefsova)
        ws.cell(row=row, column=8).font = Font(bold=True)
        ws.cell(row=row, column=8).fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        row += 2

    # Total général
    ws.cell(row=row, column=7).value = "TOTAL GÉNÉRAL:"
    ws.cell(row=row, column=7).font = Font(bold=True, size=12)
    ws.cell(row=row, column=8).value = float(paiement.montant_total)
    ws.cell(row=row, column=8).font = Font(bold=True, size=12)
    ws.cell(row=row, column=8).fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

    # Ajuster les largeurs
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 15

    # Marquer comme générée
    paiement.facture_generee = True
    paiement.date_generation_facture = timezone.now()
    paiement.save()

    # Préparer la réponse
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'Mutuelle_Sante_{paiement.organisation.nom_organisation}_T{paiement.trimestre}_{paiement.annee}.xlsx'
    response['Content-Disposition'] = f'attachment; filename={filename}'

    wb.save(response)
    return response


@csrf_exempt
@require_http_methods(["POST"])
def api_supprimer_paiement_mutuelle(request, paiement_id):
    """API pour supprimer un paiement de mutuelle santé"""
    try:
        paiement = get_object_or_404(PaiementMutuelleSante, id=paiement_id)

        if paiement.statut == 'exporte':
            return JsonResponse({
                'success': False,
                'error': 'Impossible de supprimer un paiement déjà exporté vers DiameX'
            }, status=400)

        # Supprimer (les détails sont supprimés en cascade)
        paiement.delete()

        return JsonResponse({
            'success': True,
            'message': 'Paiement supprimé avec succès'
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
def api_detail_paiement_mutuelle(request, paiement_id):
    """API pour obtenir les détails d'un paiement de mutuelle santé"""
    try:
        paiement = get_object_or_404(PaiementMutuelleSante, id=paiement_id)

        # Récupérer les détails par mutuelle
        details_amit = paiement.details.filter(mutuelle='AMIT').select_related('employe', 'employe__base')
        details_osiefsova = paiement.details.filter(mutuelle='OSIEFSoVa').select_related('employe', 'employe__base')

        # Formater les détails AMIT
        employes_amit = []
        for detail in details_amit:
            employes_amit.append({
                'id': str(detail.id),
                'employe': detail.employe.nom_complet,
                'base': detail.employe.base.nom_base if detail.employe.base else '',
                'montant_mutuelle': float(detail.total_trimestre) if detail.total_trimestre else 0,
                'mois_1': float(detail.mois_1) if detail.mois_1 else 0,
                'mois_2': float(detail.mois_2) if detail.mois_2 else 0,
                'mois_3': float(detail.mois_3) if detail.mois_3 else 0,
            })

        # Formater les détails OSIEFSoVa
        employes_osiefsova = []
        for detail in details_osiefsova:
            employes_osiefsova.append({
                'id': str(detail.id),
                'employe': detail.employe.nom_complet,
                'base': detail.employe.base.nom_base if detail.employe.base else '',
                'montant_mutuelle': float(detail.total_trimestre) if detail.total_trimestre else 0,
                'mois_1': float(detail.mois_1) if detail.mois_1 else 0,
                'mois_2': float(detail.mois_2) if detail.mois_2 else 0,
                'mois_3': float(detail.mois_3) if detail.mois_3 else 0,
            })

        # Noms des mois selon le trimestre
        mois_par_trimestre = {
            1: ['Janvier', 'Février', 'Mars'],
            2: ['Avril', 'Mai', 'Juin'],
            3: ['Juillet', 'Août', 'Septembre'],
            4: ['Octobre', 'Novembre', 'Décembre']
        }
        noms_mois = mois_par_trimestre.get(paiement.trimestre, ['Mois 1', 'Mois 2', 'Mois 3'])

        return JsonResponse({
            'success': True,
            'paiement': {
                'id': str(paiement.id),
                'organisation': paiement.organisation.nom_organisation,
                'trimestre': paiement.trimestre,
                'annee': paiement.annee,
                'periode_label': paiement.periode_label,
                'statut': paiement.statut,
                'montant_total_amit': float(paiement.montant_total_amit),
                'montant_total_osiefsova': float(paiement.montant_total_osiefsova),
                'montant_total': float(paiement.montant_total),
                'facture_generee': paiement.facture_generee,
                'numero_diamex_amit': paiement.numero_diamex_amit or '',
                'base_diamex_amit': paiement.base_diamex_amit or '',
                'numero_diamex_osiefsova': paiement.numero_diamex_osiefsova or '',
                'base_diamex_osiefsova': paiement.base_diamex_osiefsova or '',
                'created_at': paiement.created_at.strftime('%d/%m/%Y %H:%M'),
                'created_by': paiement.created_by.get_full_name() if paiement.created_by else '',
                'noms_mois': noms_mois,
            },
            'employes_amit': employes_amit,
            'employes_osiefsova': employes_osiefsova,
            'stats': {
                'nb_amit': len(employes_amit),
                'nb_osiefsova': len(employes_osiefsova),
                'total_employes': len(employes_amit) + len(employes_osiefsova)
            }
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)
