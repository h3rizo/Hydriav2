"""
Vues pour l'import Excel des relations Employés-Ouvrages
"""
import openpyxl
from datetime import datetime
from io import BytesIO
from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from decimal import Decimal

from .models import Employe, EmployeOuvrages
from apps.organization.models import Ouvrage


def import_employe_ouvrages_view(request):
    """Page d'import des relations Employés-Ouvrages"""
    return render(request, 'employees/import_employe_ouvrages.html')


@csrf_exempt
@require_http_methods(["POST"])
def api_import_employe_ouvrages(request):
    """
    API pour importer les relations Employés-Ouvrages depuis un fichier Excel.

    Format attendu du fichier Excel:
    - Colonne A: Surnom Employé
    - Colonne B: Nom Ouvrage
    - Colonne C: Date début (format DD/MM/YYYY ou YYYY-MM-DD)
    - Colonne D: Date fin (optionnel, format DD/MM/YYYY ou YYYY-MM-DD)
    - Colonne E: Responsable principal (OUI/NON ou TRUE/FALSE)
    - Colonne F: Recevoir alertes (OUI/NON ou TRUE/FALSE, défaut OUI)
    - Colonne G: Notes (optionnel)
    """
    try:
        if 'file' not in request.FILES:
            return JsonResponse({
                'success': False,
                'error': 'Aucun fichier fourni'
            }, status=400)

        file = request.FILES['file']

        # Charger le fichier Excel
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active

        nb_imported = 0
        nb_updated = 0
        errors = []

        # Parcourir les lignes (ignorer l'en-tête)
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                # Extraire les données
                surnom_employe = row[0]
                nom_ouvrage = row[1]
                date_debut_str = row[2]
                date_fin_str = row[3] if len(row) > 3 else None
                responsable_principal_str = row[4] if len(row) > 4 else 'NON'
                recevoir_alertes_str = row[5] if len(row) > 5 else 'OUI'
                notes = row[6] if len(row) > 6 else None

                # Validation des champs obligatoires
                if not surnom_employe or not nom_ouvrage or not date_debut_str:
                    errors.append(f"Ligne {row_num}: Champs obligatoires manquants (Surnom, Ouvrage, Date début)")
                    continue

                # Rechercher l'employé par surnom
                try:
                    employe = Employe.objects.get(surnom=surnom_employe)
                except Employe.DoesNotExist:
                    errors.append(f"Ligne {row_num}: Employé '{surnom_employe}' introuvable")
                    continue

                # Rechercher l'ouvrage par nom
                try:
                    ouvrage = Ouvrage.objects.get(nom_ouvrage=nom_ouvrage)
                except Ouvrage.DoesNotExist:
                    errors.append(f"Ligne {row_num}: Ouvrage '{nom_ouvrage}' introuvable")
                    continue
                except Ouvrage.MultipleObjectsReturned:
                    errors.append(f"Ligne {row_num}: Plusieurs ouvrages trouvés avec le nom '{nom_ouvrage}'")
                    continue

                # Parser la date de début
                try:
                    if isinstance(date_debut_str, datetime):
                        date_debut = date_debut_str.date()
                    else:
                        # Essayer format DD/MM/YYYY
                        try:
                            date_debut = datetime.strptime(str(date_debut_str), '%d/%m/%Y').date()
                        except ValueError:
                            # Essayer format YYYY-MM-DD
                            date_debut = datetime.strptime(str(date_debut_str), '%Y-%m-%d').date()
                except Exception as e:
                    errors.append(f"Ligne {row_num}: Format de date début invalide '{date_debut_str}'")
                    continue

                # Parser la date de fin (optionnelle)
                date_fin = None
                if date_fin_str:
                    try:
                        if isinstance(date_fin_str, datetime):
                            date_fin = date_fin_str.date()
                        else:
                            # Essayer format DD/MM/YYYY
                            try:
                                date_fin = datetime.strptime(str(date_fin_str), '%d/%m/%Y').date()
                            except ValueError:
                                # Essayer format YYYY-MM-DD
                                date_fin = datetime.strptime(str(date_fin_str), '%Y-%m-%d').date()
                    except Exception as e:
                        errors.append(f"Ligne {row_num}: Format de date fin invalide '{date_fin_str}'")
                        continue

                # Parser responsable principal
                responsable_principal = False
                if responsable_principal_str:
                    resp_str = str(responsable_principal_str).upper()
                    if resp_str in ('OUI', 'TRUE', '1', 'VRAI', 'YES'):
                        responsable_principal = True

                # Parser recevoir alertes
                recevoir_alertes = True  # Défaut
                if recevoir_alertes_str:
                    alert_str = str(recevoir_alertes_str).upper()
                    if alert_str in ('NON', 'FALSE', '0', 'FAUX', 'NO'):
                        recevoir_alertes = False

                # Vérifier si la relation existe déjà
                existing = EmployeOuvrages.objects.filter(
                    employe=employe,
                    ouvrage=ouvrage,
                    date_debut=date_debut
                ).first()

                if existing:
                    # Mettre à jour
                    existing.date_fin = date_fin
                    existing.est_responsable_principal = responsable_principal
                    existing.recevoir_alertes = recevoir_alertes
                    existing.notes = notes
                    existing.save()
                    nb_updated += 1
                else:
                    # Créer nouvelle relation
                    EmployeOuvrages.objects.create(
                        employe=employe,
                        ouvrage=ouvrage,
                        date_debut=date_debut,
                        date_fin=date_fin,
                        est_responsable_principal=responsable_principal,
                        recevoir_alertes=recevoir_alertes,
                        notes=notes
                    )
                    nb_imported += 1

            except Exception as e:
                errors.append(f"Ligne {row_num}: Erreur - {str(e)}")
                continue

        return JsonResponse({
            'success': True,
            'nb_imported': nb_imported,
            'nb_updated': nb_updated,
            'errors': errors
        })

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@require_http_methods(["GET"])
def template_excel_employe_ouvrages(request):
    """Télécharger un template Excel pour l'import des relations Employés-Ouvrages"""

    # Créer un nouveau workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relations Employés-Ouvrages"

    # En-têtes
    headers = [
        'Surnom Employé*',
        'Nom Ouvrage*',
        'Date Début*',
        'Date Fin',
        'Responsable Principal',
        'Recevoir Alertes',
        'Notes'
    ]
    ws.append(headers)

    # Style des en-têtes
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Exemples de données
    exemples = [
        ['Victorien', 'Ambanja BP', '01/01/2025', '', 'OUI', 'OUI', 'Responsable principal'],
        ['Ando', 'Vavatenina AEPA', '15/01/2025', '31/12/2025', 'NON', 'OUI', 'Responsable secondaire'],
        ['Jean', 'Antalaha AEPA', '01/02/2025', '', 'OUI', 'NON', ''],
    ]

    for exemple in exemples:
        ws.append(exemple)

    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 30

    # Ajouter une feuille avec la liste des employés
    ws_employes = wb.create_sheet(title="Employés")
    ws_employes.append(['Surnom', 'Nom', 'Prénom'])

    for cell in ws_employes[1]:
        cell.fill = header_fill
        cell.font = header_font

    employes = Employe.objects.all().order_by('surnom')[:500]
    for emp in employes:
        ws_employes.append([emp.surnom, emp.nom, emp.prenom or ''])

    ws_employes.column_dimensions['A'].width = 20
    ws_employes.column_dimensions['B'].width = 25
    ws_employes.column_dimensions['C'].width = 25

    # Ajouter une feuille avec la liste des ouvrages
    ws_ouvrages = wb.create_sheet(title="Ouvrages")
    ws_ouvrages.append(['Nom', 'Type', 'Village'])

    for cell in ws_ouvrages[1]:
        cell.fill = header_fill
        cell.font = header_font

    ouvrages = Ouvrage.objects.select_related('type_ouvrage').all().order_by('nom_ouvrage')[:1000]
    for ouvr in ouvrages:
        ws_ouvrages.append([
            ouvr.nom_ouvrage,
            ouvr.type_ouvrage.nom_type_ouvrage if ouvr.type_ouvrage else '',
            ''  # Village - non disponible via relation directe
        ])

    ws_ouvrages.column_dimensions['A'].width = 30
    ws_ouvrages.column_dimensions['B'].width = 20
    ws_ouvrages.column_dimensions['C'].width = 25

    # Ajouter une feuille d'instructions
    ws_instructions = wb.create_sheet(title="Instructions")
    instructions = [
        ["Import des Relations Employés-Ouvrages"],
        [""],
        ["Format attendu:"],
        ["- Colonne A: Surnom Employé* (obligatoire) - Doit correspondre exactement à un surnom existant"],
        ["- Colonne B: Nom Ouvrage* (obligatoire) - Doit correspondre exactement à un nom d'ouvrage existant"],
        ["- Colonne C: Date Début* (obligatoire) - Format DD/MM/YYYY ou YYYY-MM-DD"],
        ["- Colonne D: Date Fin (optionnel) - Format DD/MM/YYYY ou YYYY-MM-DD - Laisser vide si responsabilité active"],
        ["- Colonne E: Responsable Principal (optionnel) - OUI/NON ou TRUE/FALSE - Défaut: NON"],
        ["- Colonne F: Recevoir Alertes (optionnel) - OUI/NON ou TRUE/FALSE - Défaut: OUI"],
        ["- Colonne G: Notes (optionnel) - Texte libre"],
        [""],
        ["Notes:"],
        ["- Les champs marqués * sont obligatoires"],
        ["- Utilisez les feuilles 'Employés' et 'Ouvrages' pour trouver les noms/surnoms exacts"],
        ["- Si une relation existe déjà (même employé, même ouvrage, même date début), elle sera mise à jour"],
        ["- Sinon, une nouvelle relation sera créée"],
        [""],
        ["Exemples de valeurs:"],
        ["- Responsable Principal: OUI, NON, TRUE, FALSE, 1, 0"],
        ["- Recevoir Alertes: OUI, NON, TRUE, FALSE, 1, 0"],
        ["- Date: 01/01/2025, 2025-01-01, 15/03/2025, 2025-03-15"],
    ]

    for instruction in instructions:
        ws_instructions.append(instruction)

    ws_instructions.column_dimensions['A'].width = 100

    # Enregistrer dans un buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Créer la réponse HTTP
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=template_employe_ouvrages.xlsx'

    return response


@require_http_methods(["GET"])
def export_employe_ouvrages(request):
    """Exporter toutes les relations Employés-Ouvrages vers Excel"""

    # Créer un nouveau workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relations Employés-Ouvrages"

    # En-têtes
    headers = [
        'ID',
        'Surnom Employé',
        'Nom Employé',
        'Prénom Employé',
        'Nom Ouvrage',
        'Type Ouvrage',
        'Village',
        'Date Début',
        'Date Fin',
        'Responsable Principal',
        'Recevoir Alertes',
        'Notes',
        'Créé le',
        'Modifié le'
    ]
    ws.append(headers)

    # Style des en-têtes
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Récupérer toutes les relations
    relations = EmployeOuvrages.objects.select_related(
        'employe', 'ouvrage', 'ouvrage__type_ouvrage'
    ).all().order_by('employe__surnom', 'ouvrage__nom_ouvrage')

    # Ajouter les données
    for rel in relations:
        ws.append([
            str(rel.id),
            rel.employe.surnom,
            rel.employe.nom,
            rel.employe.prenom or '',
            rel.ouvrage.nom_ouvrage,
            rel.ouvrage.type_ouvrage.nom_type_ouvrage if rel.ouvrage.type_ouvrage else '',
            '',  # Village - non disponible via relation directe
            rel.date_debut.strftime('%d/%m/%Y') if rel.date_debut else '',
            rel.date_fin.strftime('%d/%m/%Y') if rel.date_fin else '',
            'OUI' if rel.est_responsable_principal else 'NON',
            'OUI' if rel.recevoir_alertes else 'NON',
            rel.notes or '',
            rel.created_at.strftime('%d/%m/%Y %H:%M') if rel.created_at else '',
            rel.updated_at.strftime('%d/%m/%Y %H:%M') if rel.updated_at else ''
        ])

    # Ajuster la largeur des colonnes
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N']:
        ws.column_dimensions[col].width = 20

    # Enregistrer dans un buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Créer la réponse HTTP
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=export_employe_ouvrages.xlsx'

    return response
