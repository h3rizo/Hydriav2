"""
Vue pour la génération de facture Excel IRI
"""
import os
from datetime import datetime
from decimal import Decimal
from django.http import JsonResponse, FileResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, Side, Alignment, Font

from apps.employees.models import PaiementIRI
from apps.organization.models import Organisation
from apps.organization.scoping import get_user_accessible_gestionnaire_ids


def _get_accessible_paiements_iri_queryset(user):
    queryset = PaiementIRI.objects.select_related('organisation')
    allowed_org_ids = get_user_accessible_gestionnaire_ids(user)
    if allowed_org_ids is None:
        return queryset
    if not allowed_org_ids:
        return queryset.none()
    return queryset.filter(organisation_id__in=allowed_org_ids)


def safe_write(ws, row, col, value, border_style=None, font_style=None, alignment_style=None):
    """
    Écrit une valeur dans une cellule en ignorant les cellules fusionnées
    """
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        cell.value = value
        if border_style:
            cell.border = border_style
        if font_style:
            cell.font = font_style
        if alignment_style:
            cell.alignment = alignment_style


def get_adresse_organisation(organisation):
    """Helper pour construire l'adresse complète"""
    if not organisation:
        return ''

    adresse_parts = []
    if organisation.adresse1:
        adresse_parts.append(organisation.adresse1)
    if organisation.adresse2:
        adresse_parts.append(organisation.adresse2)
    if organisation.localite:
        adresse_parts.append(organisation.localite)

    return ', '.join(adresse_parts) if adresse_parts else ''


@csrf_exempt
@require_http_methods(["POST"])
def generer_facture_iri(request, paiement_id):
    """
    Génère la facture Excel IRI à partir du modèle
    """
    try:
        paiement = _get_accessible_paiements_iri_queryset(request.user).get(id=paiement_id)

        # Vérifier que le paiement est validé
        if paiement.statut != 'valide':
            return JsonResponse({
                'success': False,
                'error': 'Seuls les paiements validés peuvent générer une facture'
            }, status=400)

        # Récupérer les détails groupés par ACTIV
        details_qs = paiement.details.all().order_by('activ', 'date_achat')

        if not details_qs.exists():
            return JsonResponse({
                'success': False,
                'error': 'Aucun détail trouvé pour ce paiement'
            }, status=400)

        # Charger le modèle Excel
        template_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
            'Modèle Facture IRI.xlsx'
        )

        if not os.path.exists(template_path):
            return JsonResponse({
                'success': False,
                'error': f'Modèle non trouvé: {template_path}'
            }, status=500)

        wb = load_workbook(template_path)

        # Générer les deux onglets
        erreurs = []

        # 1. Onglet Annexe IRI
        if 'Annexe IRI' in wb.sheetnames:
            ws_annexe = wb['Annexe IRI']

            # En-tête
            safe_write(ws_annexe, 4, 5, paiement.organisation.nom_organisation if paiement.organisation else '')  # E4 - Nom organisation
            safe_write(ws_annexe, 5, 5, paiement.organisation.nif if paiement.organisation else '')  # E5 - NIF
            safe_write(ws_annexe, 6, 5, get_adresse_organisation(paiement.organisation))  # E6 - Adresse

            # Format mois : MMMM-YY (ex: Janvier-25)
            mois_noms = {
                1: 'Janvier', 2: 'Février', 3: 'Mars', 4: 'Avril',
                5: 'Mai', 6: 'Juin', 7: 'Juillet', 8: 'Août',
                9: 'Septembre', 10: 'Octobre', 11: 'Novembre', 12: 'Décembre'
            }
            mois_label = f"{mois_noms.get(paiement.mois, '')}-{str(paiement.annee)[2:]}"
            safe_write(ws_annexe, 7, 5, mois_label)  # E7 - Mois

            # Effacer les anciennes données du modèle (à partir de la ligne 12)
            # On efface jusqu'à la ligne 1000 pour être sûr
            from openpyxl.styles import Border
            no_border = Border()  # Bordure vide

            for row_idx in range(12, 1000):
                for col_idx in range(1, 11):  # Colonnes A à J
                    cell = ws_annexe.cell(row=row_idx, column=col_idx)
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                        cell.border = no_border  # Effacer les bordures aussi

            # Tableau des détails
            # En-tête se termine à la ligne 11
            # Données à partir de la ligne 12
            # Colonnes: A=N° reçu, B=Fournisseur, C=CIN, D=Adresse, E=Date, F=Libellé,
            #           G=Unité, H=Prix unitaire, I=Montant, J=IRI 5%

            start_row = 12
            current_row = start_row

            # Style des bordures
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Font gras pour les sous-totaux
            bold_font = Font(bold=True)

            # Grouper par ACTIV
            grouped_by_activ = {}
            for detail in details_qs:
                activ = detail.activ or 'Sans ACTIV'
                if activ not in grouped_by_activ:
                    grouped_by_activ[activ] = []
                grouped_by_activ[activ].append(detail)

            # Remplir les données groupées
            for activ in sorted(grouped_by_activ.keys()):
                details_group = grouped_by_activ[activ]
                subtotal_montant = Decimal('0')
                subtotal_iri = Decimal('0')

                # Ligne d'en-tête de groupe ACTIV
                safe_write(ws_annexe, current_row, 1, f'ACTIV: {activ}', border_style=thin_border, font_style=bold_font)
                # Fusionner les colonnes A à J pour l'en-tête de groupe
                ws_annexe.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
                current_row += 1

                # Détails du groupe
                for detail in details_group:
                    safe_write(ws_annexe, current_row, 1, detail.numero_diamex_source or '', border_style=thin_border)  # A - N° reçu
                    safe_write(ws_annexe, current_row, 2, detail.fournisseur_nom or '', border_style=thin_border)  # B - Fournisseur
                    safe_write(ws_annexe, current_row, 3, detail.fournisseur_cin or '', border_style=thin_border)  # C - CIN
                    safe_write(ws_annexe, current_row, 4, detail.fournisseur_adresse or '', border_style=thin_border)  # D - Adresse
                    safe_write(ws_annexe, current_row, 5,
                              detail.date_achat.strftime('%d/%m/%Y') if detail.date_achat else '',
                              border_style=thin_border)  # E - Date
                    safe_write(ws_annexe, current_row, 6, detail.libelle or '', border_style=thin_border)  # F - Libellé
                    safe_write(ws_annexe, current_row, 7, detail.unite or '', border_style=thin_border)  # G - Unité
                    safe_write(ws_annexe, current_row, 8, float(detail.prix_unitaire) if detail.prix_unitaire else 0, border_style=thin_border)  # H - Prix unitaire
                    safe_write(ws_annexe, current_row, 9, float(detail.montant_achat) if detail.montant_achat else 0, border_style=thin_border)  # I - Montant
                    safe_write(ws_annexe, current_row, 10, float(detail.montant_iri) if detail.montant_iri else 0, border_style=thin_border)  # J - IRI 5%

                    subtotal_montant += detail.montant_achat if detail.montant_achat else Decimal('0')
                    subtotal_iri += detail.montant_iri if detail.montant_iri else Decimal('0')
                    current_row += 1

                # Ligne de sous-total par ACTIV
                safe_write(ws_annexe, current_row, 1, f'Sous-total {activ}', border_style=thin_border, font_style=bold_font)
                ws_annexe.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
                safe_write(ws_annexe, current_row, 9, float(subtotal_montant), border_style=thin_border, font_style=bold_font)  # I - Total montant
                safe_write(ws_annexe, current_row, 10, float(subtotal_iri), border_style=thin_border, font_style=bold_font)  # J - Total IRI
                current_row += 1

            # Ligne TOTAL GÉNÉRAL
            total_montant = sum(
                detail.montant_achat if detail.montant_achat else Decimal('0')
                for detail in details_qs
            )
            total_iri = sum(
                detail.montant_iri if detail.montant_iri else Decimal('0')
                for detail in details_qs
            )

            safe_write(ws_annexe, current_row, 1, 'TOTAL GÉNÉRAL', border_style=thin_border, font_style=bold_font)
            ws_annexe.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
            safe_write(ws_annexe, current_row, 9, float(total_montant), border_style=thin_border, font_style=bold_font)  # I - Total montant
            safe_write(ws_annexe, current_row, 10, float(total_iri), border_style=thin_border, font_style=bold_font)  # J - Total IRI

            # Définir la zone d'impression pour s'arrêter après TOTAL GÉNÉRAL
            # Format: $A$1:$J$current_row (de la ligne 1 à la ligne du total général, colonnes A à J)
            ws_annexe.print_area = f'$A$1:$J${current_row}'

        else:
            erreurs.append("Onglet 'Annexe IRI' non trouvé dans le modèle")

        # 2. Onglet Déclaration IRI
        if 'Déclaration IRI' in wb.sheetnames:
            ws_declaration = wb['Déclaration IRI']

            # Informations organisation
            safe_write(ws_declaration, 19, 5, paiement.organisation.nom_organisation if paiement.organisation else '')  # E19
            safe_write(ws_declaration, 23, 5, get_adresse_organisation(paiement.organisation))  # E23
            safe_write(ws_declaration, 25, 5, paiement.organisation.nif if paiement.organisation else '')  # E25

            # Format mois : MMM-YY (ex: Jan-25)
            mois_court = {
                1: 'Jan', 2: 'Fév', 3: 'Mar', 4: 'Avr',
                5: 'Mai', 6: 'Juin', 7: 'Juil', 8: 'Août',
                9: 'Sept', 10: 'Oct', 11: 'Nov', 12: 'Déc'
            }
            mois_label_court = f"{mois_court.get(paiement.mois, '')}-{str(paiement.annee)[2:]}"
            safe_write(ws_declaration, 27, 5, mois_label_court)  # E27

            # Total des montants - utiliser le total stocké dans le paiement
            safe_write(ws_declaration, 31, 7, float(paiement.montant_total_achats))  # G31

            # G33 = G31 * G32 (formule)
            safe_write(ws_declaration, 33, 7, f'=G31*G32')  # G33
        else:
            erreurs.append("Onglet 'Déclaration IRI' non trouvé dans le modèle")

        # Sauvegarder le fichier généré
        media_dir = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
            'media', 'factures_iri'
        )
        os.makedirs(media_dir, exist_ok=True)

        nom_org = paiement.organisation.nom_organisation.replace(' ', '_') if paiement.organisation else 'Organisation'
        filename = f'Facture_IRI_{nom_org}_{paiement.mois:02d}_{paiement.annee}.xlsx'
        output_path = os.path.join(media_dir, filename)

        wb.save(output_path)

        # Marquer le paiement comme ayant une facture générée
        paiement.facture_generee = True
        paiement.save()

        return JsonResponse({
            'success': True,
            'message': 'Facture IRI générée avec succès',
            'filename': filename,
            'path': output_path,
            'erreurs': erreurs
        })

    except PaiementIRI.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Paiement IRI introuvable'
        }, status=404)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f'Erreur lors de la génération: {str(e)}'
        }, status=500)


@require_http_methods(["GET"])
def telecharger_facture_iri(request, paiement_id):
    """
    Télécharge la facture IRI générée
    """
    try:
        paiement = _get_accessible_paiements_iri_queryset(request.user).get(id=paiement_id)

        if not paiement.facture_generee:
            return JsonResponse({
                'success': False,
                'error': 'La facture n\'a pas encore été générée'
            }, status=400)

        # Construire le chemin du fichier
        media_dir = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
            'media', 'factures_iri'
        )

        nom_org = paiement.organisation.nom_organisation.replace(' ', '_') if paiement.organisation else 'Organisation'
        filename = f'Facture_IRI_{nom_org}_{paiement.mois:02d}_{paiement.annee}.xlsx'
        file_path = os.path.join(media_dir, filename)

        if not os.path.exists(file_path):
            return JsonResponse({
                'success': False,
                'error': 'Fichier de facture introuvable'
            }, status=404)

        response = FileResponse(
            open(file_path, 'rb'),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response

    except PaiementIRI.DoesNotExist:
        return JsonResponse({
            'success': False,
            'error': 'Paiement IRI introuvable'
        }, status=404)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f'Erreur lors du téléchargement: {str(e)}'
        }, status=500)
