from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, render, redirect
from django.contrib import messages
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods

from apps.billing.models import (
    ChangementCompteurTete,
    CompteurTete,
    PointEau,
    Releve,
    ReleveCompteurTete,
)
from apps.organization.scoping import (
    get_user_accessible_gestionnaire_ids,
    get_user_organisation,
    get_user_responsable_ouvrage_ids,
)


GAP_SIGNIFICANCE_THRESHOLD = Decimal('0.01')
GAP_PERCENT_TOLERANCE = Decimal('10')


def _get_allowed_ouvrages_for_user(user):
    from .views import _get_allowed_ouvrages_for_user as _source
    return _source(user)


def _user_is_director(user):
    from .views import user_is_director as _source
    return _source(user)


def _get_scoped_ouvrages_for_releves(user):
    ouvrages = _get_allowed_ouvrages_for_user(user).select_related('type_ouvrage')

    gestionnaire_ids = get_user_accessible_gestionnaire_ids(user)
    current_organisation = get_user_organisation(user)

    if gestionnaire_ids == []:
        ouvrages = ouvrages.none()
    elif gestionnaire_ids:
        ouvrages = ouvrages.filter(
            modalites_gestion__date_fin__isnull=True,
            modalites_gestion__gestionnaire_id__in=gestionnaire_ids,
            modalites_gestion__modalitepaiement__in=['au_volume', 'mixte'],
        )
    elif current_organisation:
        ouvrages = ouvrages.filter(
            modalites_gestion__date_fin__isnull=True,
            modalites_gestion__gestionnaire=current_organisation,
            modalites_gestion__modalitepaiement__in=['au_volume', 'mixte'],
        )

    responsable_ids = get_user_responsable_ouvrage_ids(user)
    if responsable_ids is None:
        pass
    elif not responsable_ids:
        ouvrages = ouvrages.none()
    else:
        ouvrages = ouvrages.filter(id__in=responsable_ids)

    return ouvrages.filter(
        compteurs_tete__isnull=False
    ).distinct().order_by('nom_ouvrage')


def _parse_date(value, field_name='date_releve'):
    raw_value = (value or '').strip()
    if not raw_value:
        raise ValidationError({field_name: "La date du releve est obligatoire."})
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(raw_value, fmt).date()
        except ValueError:
            continue
    try:
        return ReleveCompteurTete._meta.get_field('date_releve').to_python(raw_value)
    except Exception as exc:
        raise ValidationError({field_name: "Format de date invalide."}) from exc


def _parse_decimal(value, field_name, label):
    raw_value = (value or '').strip()
    if not raw_value:
        raise ValidationError({field_name: f"{label} est obligatoire."})
    try:
        return Decimal(raw_value.replace(',', '.'))
    except (InvalidOperation, AttributeError) as exc:
        raise ValidationError({field_name: f"{label} invalide."}) from exc


def _decimal_or_zero(value):
    try:
        return Decimal(str(value or 0))
    except Exception:
        return Decimal('0')


def _format_number_with_spaces(value, decimals=0):
    try:
        number = float(value or 0)
    except Exception:
        number = 0.0
    return f"{number:,.{decimals}f}".replace(',', ' ')


def _get_initial_releve_ids(queryset):
    initial_ids = set()
    current_compteur_id = None
    ordered_qs = queryset.order_by('compteur_tete_id', 'date_releve', 'created_at', 'id')
    for releve in ordered_qs.only('id', 'compteur_tete_id'):
        if releve.compteur_tete_id != current_compteur_id:
            initial_ids.add(releve.id)
            current_compteur_id = releve.compteur_tete_id
    return initial_ids


def _get_latest_releve_ids_by_compteur(queryset):
    latest_ids = set()
    seen_compteur_ids = set()
    ordered_qs = queryset.order_by('compteur_tete_id', '-date_releve', '-created_at', '-id')
    for releve in ordered_qs.only('id', 'compteur_tete_id'):
        if releve.compteur_tete_id in seen_compteur_ids:
            continue
        latest_ids.add(releve.id)
        seen_compteur_ids.add(releve.compteur_tete_id)
    return latest_ids


def _allowed_compteurs_queryset(user):
    return CompteurTete.objects.select_related('ouvrage').filter(
        ouvrage__in=_get_scoped_ouvrages_for_releves(user)
    )


def _latest_compteur_index(compteur):
    dernier_releve = compteur.releves_compteur_tete.order_by('-date_releve', '-created_at').first()
    if dernier_releve:
        return _decimal_or_zero(dernier_releve.nouvel_index)
    return _decimal_or_zero(compteur.index_initial)


def _linked_points_queryset(compteur):
    return PointEau.objects.filter(
        compteur_tete=compteur,
        datebranchement__isnull=False,
    ).select_related(
        'parcelle__village'
    ).order_by('nomabonne', 'code_point_eau')


def _allowed_points_queryset(user):
    return PointEau.objects.filter(
        ouvrage__in=_get_scoped_ouvrages_for_releves(user),
        datebranchement__isnull=False,
    ).select_related(
        'ouvrage',
        'compteur_tete',
        'parcelle__village',
    )


def _serialize_point_for_assignment(point):
    village = ''
    if point.parcelle and point.parcelle.village:
        village = point.parcelle.village.nom_village or ''

    return {
        'id': str(point.id),
        'nom': point.nomabonne or '',
        'code': point.code_point_eau or '',
        'village': village,
        'ouvrage_nom': point.ouvrage.nom_ouvrage if point.ouvrage else '',
        'compteur_id': str(point.compteur_tete_id) if point.compteur_tete_id else '',
        'compteur_numero': point.compteur_tete.numerocompteur if point.compteur_tete else '',
    }


def _get_allowed_point_for_compteur(user, compteur, point_id):
    point_id = (point_id or '').strip()
    if not point_id:
        raise ValidationError({'point_id': ["Le point d'eau est obligatoire."]})

    point = get_object_or_404(_allowed_points_queryset(user), id=point_id)
    if point.ouvrage_id != compteur.ouvrage_id:
        raise ValidationError({
            'point_id': ["Le point d'eau doit appartenir au meme ouvrage que le compteur."]
        })
    return point


def _serialize_compteur_option(compteur):
    return {
        'id': str(compteur.id),
        'numero': compteur.numerocompteur,
        'ouvrage_nom': compteur.ouvrage.nom_ouvrage if compteur.ouvrage else '',
    }


def _build_compteur_points_payload(user, compteur):
    linked_points = list(_linked_points_queryset(compteur))
    available_points = list(
        _allowed_points_queryset(user).filter(
            ouvrage=compteur.ouvrage,
        ).exclude(
            compteur_tete=compteur,
        ).order_by('nomabonne', 'code_point_eau')[:300]
    )
    target_compteurs = list(
        _allowed_compteurs_queryset(user).filter(
            ouvrage=compteur.ouvrage,
        ).exclude(
            id=compteur.id,
        ).select_related('ouvrage').order_by('numerocompteur')
    )

    return {
        'compteur': _serialize_compteur_option(compteur),
        'linked_points': [_serialize_point_for_assignment(point) for point in linked_points],
        'available_points': [_serialize_point_for_assignment(point) for point in available_points],
        'target_compteurs': [_serialize_compteur_option(item) for item in target_compteurs],
    }


def _default_calculation_start(date_fin):
    """Periode de repli pour le tout premier releve du compteur de tete."""
    first_day = date_fin.replace(day=1)
    if first_day < date_fin:
        return first_day
    return (first_day - timedelta(days=1)).replace(day=1)


def _first_compteur_calculation_start(compteur, date_fin):
    """Utilise la mise en service comme origine du premier relevé si elle est connue."""
    date_installation = getattr(compteur, 'date_installation', None)
    if date_installation and date_installation < date_fin:
        return date_installation
    return _default_calculation_start(date_fin)


def _downstream_volume_for_period(releves, date_debut, date_fin):
    """Estime le volume aval sur la periode exacte par prorata journalier."""
    ordered = sorted(releves, key=lambda item: (item.datereleve, str(item.id)))
    if not date_debut or not date_fin or date_debut >= date_fin:
        return {
            'volume': Decimal('0'),
            'volume_reel_index': Decimal('0'),
            'nb_releves': 0,
            'date_debut_source': None,
            'date_fin_source': None,
            'estimation_utilisee': False,
            'prorata_utilise': False,
            'methode': 'Periode indisponible',
            'duree_jours': 0,
        }

    target_days = (date_fin - date_debut).days
    exact_dates = {item.datereleve for item in ordered}
    used_releves = set()
    source_dates = []
    total = Decimal('0')
    intervals = []
    source_volumes = {}

    for previous, current in zip(ordered, ordered[1:]):
        interval_days = (current.datereleve - previous.datereleve).days
        if interval_days <= 0:
            continue
        # La comparaison reseau porte sur le volume physique lu au compteur,
        # pas sur le volume facture qui peut exclure des volumes sociaux/non factures.
        actual_volume = _decimal_or_zero(current.consommation)
        daily_volume = actual_volume / Decimal(interval_days)
        intervals.append((previous, current, daily_volume, actual_volume))

        overlap_start = max(date_debut, previous.datereleve)
        overlap_end = min(date_fin, current.datereleve)
        overlap_days = (overlap_end - overlap_start).days
        if overlap_days <= 0:
            continue

        total += daily_volume * Decimal(overlap_days)
        used_releves.update((previous.id, current.id))
        source_volumes[current.id] = actual_volume
        source_dates.extend((previous.datereleve, current.datereleve))

    estimation_utilisee = date_debut not in exact_dates or date_fin not in exact_dates
    methode = 'Dates aval identiques'

    if intervals:
        first_previous, first_current, first_daily_volume, first_actual_volume = intervals[0]
        last_previous, last_current, last_daily_volume, last_actual_volume = intervals[-1]

        uncovered_before_end = min(date_fin, first_previous.datereleve)
        uncovered_before_days = (uncovered_before_end - date_debut).days
        if uncovered_before_days > 0:
            total += first_daily_volume * Decimal(uncovered_before_days)
            used_releves.update((first_previous.id, first_current.id))
            source_volumes[first_current.id] = first_actual_volume
            source_dates.extend((first_previous.datereleve, first_current.datereleve))
            estimation_utilisee = True

        uncovered_after_start = max(date_debut, last_current.datereleve)
        uncovered_after_days = (date_fin - uncovered_after_start).days
        if uncovered_after_days > 0:
            total += last_daily_volume * Decimal(uncovered_after_days)
            used_releves.update((last_previous.id, last_current.id))
            source_volumes[last_current.id] = last_actual_volume
            source_dates.extend((last_previous.datereleve, last_current.datereleve))
            estimation_utilisee = True

        if estimation_utilisee:
            methode = 'Prorata journalier'
    elif len(ordered) == 1 and date_debut <= ordered[0].datereleve <= date_fin:
        # Sans deux dates aval, aucun prorata fiable n'est possible. On conserve
        # le volume du releve unique et on signale explicitement la limite.
        only_releve = ordered[0]
        total = _decimal_or_zero(only_releve.consommation)
        source_volumes[only_releve.id] = total
        used_releves.add(only_releve.id)
        source_dates.append(only_releve.datereleve)
        estimation_utilisee = True
        methode = 'Estimation limitee (un seul releve aval)'
    elif not ordered:
        methode = 'Aucun releve aval'
    else:
        methode = 'Donnees aval insuffisantes'

    return {
        'volume': total.quantize(Decimal('0.01')),
        'volume_reel_index': sum(source_volumes.values(), Decimal('0')).quantize(Decimal('0.01')),
        'nb_releves': len(used_releves),
        'date_debut_source': min(source_dates) if source_dates else None,
        'date_fin_source': max(source_dates) if source_dates else None,
        'estimation_utilisee': estimation_utilisee,
        'prorata_utilise': methode == 'Prorata journalier',
        'methode': methode,
        'duree_jours': target_days,
    }


def _get_downstream_breakdown(compteur, date_debut, date_fin):
    points = list(_linked_points_queryset(compteur))
    point_ids = [point.id for point in points]
    releves_by_point = {str(point_id): [] for point_id in point_ids}

    for point in points:
        village = ''
        if point.parcelle and point.parcelle.village:
            village = point.parcelle.village.nom_village or ''
        point._comparison_village = village

    if not point_ids:
        return Decimal('0'), []

    releves = Releve.objects.filter(
        point_eau_id__in=point_ids,
    ).select_related('point_eau').order_by('point_eau__nomabonne', 'datereleve')

    for releve in releves:
        point_key = str(releve.point_eau_id)
        if point_key in releves_by_point:
            releves_by_point[point_key].append(releve)

    details = []
    total = Decimal('0')
    for point in points:
        estimation = _downstream_volume_for_period(
            releves_by_point.get(str(point.id), []),
            date_debut,
            date_fin,
        )
        total += estimation['volume']
        details.append({
            'id': str(point.id),
            'nom': point.nomabonne or '',
            'code': point.code_point_eau or '',
            'village': getattr(point, '_comparison_village', ''),
            'nb_releves': estimation['nb_releves'],
            'volume_aval': float(estimation['volume']),
            'volume_reel_index': float(estimation['volume_reel_index']),
            'date_debut_source': estimation['date_debut_source'].strftime('%d/%m/%Y') if estimation['date_debut_source'] else '',
            'date_fin_source': estimation['date_fin_source'].strftime('%d/%m/%Y') if estimation['date_fin_source'] else '',
            'estimation_utilisee': estimation['estimation_utilisee'],
            'prorata_utilise': estimation['prorata_utilise'],
            'methode_estimation': estimation['methode'],
        })

    return total, details


def _get_previous_releve(compteur, date_releve, exclude_id=None):
    queryset = ReleveCompteurTete.objects.filter(
        compteur_tete=compteur,
        date_releve__lt=date_releve,
    ).order_by('-date_releve', '-created_at')
    if exclude_id:
        queryset = queryset.exclude(id=exclude_id)
    return queryset.first()


def _get_next_releve(compteur, date_releve, exclude_id=None):
    queryset = ReleveCompteurTete.objects.filter(
        compteur_tete=compteur,
        date_releve__gt=date_releve,
    ).order_by('date_releve', 'created_at')
    if exclude_id:
        queryset = queryset.exclude(id=exclude_id)
    return queryset.first()


def _build_snapshot(compteur, date_releve, nouvel_index, exclude_releve_id=None):
    ancien_releve = _get_previous_releve(compteur, date_releve, exclude_id=exclude_releve_id)
    ancien_index = _decimal_or_zero(ancien_releve.nouvel_index if ancien_releve else compteur.index_initial)

    if nouvel_index < ancien_index:
        raise ValidationError({
            'nouvel_index': "Le nouvel index doit etre superieur ou egal a l'index precedent."
        })

    prochain_releve = _get_next_releve(compteur, date_releve, exclude_id=exclude_releve_id)
    if prochain_releve and _decimal_or_zero(prochain_releve.nouvel_index) < nouvel_index:
        raise ValidationError({
            'nouvel_index': (
                "Le nouvel index est incompatible avec le releve suivant du "
                f"{prochain_releve.date_releve.strftime('%d/%m/%Y')}."
            )
        })

    date_debut_calcul = (
        ancien_releve.date_releve
        if ancien_releve else
        _first_compteur_calculation_start(compteur, date_releve)
    )
    volume_tete = nouvel_index - ancien_index
    volume_aval, points_details = _get_downstream_breakdown(compteur, date_debut_calcul, date_releve)
    ecart_volume = volume_tete - volume_aval

    return {
        'ancien_index': ancien_index,
        'volume_tete': volume_tete,
        'volume_aval': volume_aval,
        'ecart_volume': ecart_volume,
        'nb_points_eau_aval': len(points_details),
        'points_details': points_details,
        'date_debut_calcul': date_debut_calcul,
        'date_fin_calcul': date_releve,
        'aval_estime': any(detail['estimation_utilisee'] for detail in points_details),
    }


def _get_releve_gap_classification(releve):
    """Return the canonical head-meter gap class used by pages and widgets."""
    ecart = _decimal_or_zero(releve.ecart_volume)
    volume_tete = _decimal_or_zero(releve.volume_tete)
    ecart_percent = Decimal('0')
    if abs(volume_tete) > 0:
        ecart_percent = (abs(ecart) / abs(volume_tete)) * Decimal('100')

    severity = 'good'
    if ecart_percent > Decimal('20'):
        severity = 'critical'
    elif ecart_percent > Decimal('10'):
        severity = 'investigate'
    elif ecart_percent >= Decimal('5'):
        severity = 'acceptable'

    return severity, ecart, volume_tete, ecart_percent


def _serialize_releve(releve):
    severity, ecart, volume_tete, ecart_percent = _get_releve_gap_classification(releve)
    explication = (releve.explication_ecart or '').strip()

    explication_required = _gap_explanation_required(ecart, volume_tete=volume_tete) and not explication
    date_debut_calcul = getattr(releve, 'date_debut_calcul', None)
    if not date_debut_calcul:
        previous = _get_previous_releve(releve.compteur_tete, releve.date_releve, exclude_id=releve.id)
        date_debut_calcul = (
            previous.date_releve
            if previous else
            _first_compteur_calculation_start(releve.compteur_tete, releve.date_releve)
        )

    return {
        'id': str(releve.id),
        'compteur_id': str(releve.compteur_tete_id),
        'compteur_numero': releve.compteur_tete.numerocompteur,
        'ouvrage_nom': releve.compteur_tete.ouvrage.nom_ouvrage if releve.compteur_tete.ouvrage else '',
        'date_releve': releve.date_releve.strftime('%Y-%m-%d'),
        'date_releve_display': releve.date_releve.strftime('%d/%m/%Y'),
        'date_debut_calcul': date_debut_calcul.strftime('%Y-%m-%d'),
        'date_debut_calcul_display': date_debut_calcul.strftime('%d/%m/%Y'),
        'date_fin_calcul': releve.date_releve.strftime('%Y-%m-%d'),
        'date_fin_calcul_display': releve.date_releve.strftime('%d/%m/%Y'),
        'duree_calcul_jours': max((releve.date_releve - date_debut_calcul).days, 0),
        'ancien_index': float(_decimal_or_zero(releve.ancien_index)),
        'nouvel_index': float(_decimal_or_zero(releve.nouvel_index)),
        'volume_tete': float(volume_tete),
        'volume_aval': float(_decimal_or_zero(releve.volume_aval)),
        'ecart_volume': float(ecart),
        'ecart_percent': float(ecart_percent),
        'ecart_percent_rounded': int(round(float(ecart_percent))),
        'ecart_progress_percent': min(int(round(float(ecart_percent))), 100),
        'severity': severity,
        'nb_points_eau_aval': int(releve.nb_points_eau_aval or 0),
        'explication_ecart': releve.explication_ecart or '',
        'has_significant_gap': abs(ecart) > GAP_SIGNIFICANCE_THRESHOLD,
        'explication_required': explication_required,
        'updated_at': releve.updated_at.strftime('%d/%m/%Y %H:%M') if releve.updated_at else '',
    }


def _serialize_changement_compteur(changement):
    ancien = changement.ancien_compteur
    nouveau = changement.nouveau_compteur
    ouvrage = ancien.ouvrage if ancien else None
    auteur = changement.created_by
    return {
        'id': str(changement.id),
        'date_changement': changement.date_changement.strftime('%Y-%m-%d'),
        'date_changement_display': changement.date_changement.strftime('%d/%m/%Y'),
        'ouvrage_nom': ouvrage.nom_ouvrage if ouvrage else '',
        'ancien_compteur_id': str(changement.ancien_compteur_id),
        'ancien_compteur_numero': ancien.numerocompteur,
        'nouveau_compteur_id': str(changement.nouveau_compteur_id),
        'nouveau_compteur_numero': nouveau.numerocompteur,
        'index_retrait': float(changement.index_retrait),
        'index_initial_nouveau': float(changement.index_initial_nouveau),
        'motif': changement.motif,
        'motif_display': changement.get_motif_display(),
        'explication': changement.explication,
        'nb_points_transferes': changement.nb_points_transferes,
        'effectue_par': auteur.get_full_name() or auteur.get_username() if auteur else '',
        'created_at': changement.created_at.strftime('%d/%m/%Y %H:%M') if changement.created_at else '',
    }


def _resequence_releves(compteur):
    releves = list(
        ReleveCompteurTete.objects.filter(compteur_tete=compteur).order_by('date_releve', 'created_at', 'id')
    )

    previous_index = _decimal_or_zero(compteur.index_initial)
    previous_date = None
    changed = []
    now_value = timezone.now()

    for releve in releves:
        nouvel_index = _decimal_or_zero(releve.nouvel_index)
        if nouvel_index < previous_index:
            raise ValidationError(
                f"Le releve du {releve.date_releve.strftime('%d/%m/%Y')} est inferieur a l'index precedent."
            )

        date_debut_calcul = previous_date or _first_compteur_calculation_start(compteur, releve.date_releve)
        volume_aval, points_details = _get_downstream_breakdown(
            compteur,
            date_debut_calcul,
            releve.date_releve,
        )
        releve.ancien_index = previous_index
        releve.volume_tete = nouvel_index - previous_index
        releve.volume_aval = volume_aval
        releve.ecart_volume = releve.volume_tete - volume_aval
        releve.nb_points_eau_aval = len(points_details)
        releve.updated_at = now_value
        changed.append(releve)
        previous_index = nouvel_index
        previous_date = releve.date_releve

    if changed:
        ReleveCompteurTete.objects.bulk_update(
            changed,
            ['ancien_index', 'volume_tete', 'volume_aval', 'ecart_volume', 'nb_points_eau_aval', 'updated_at'],
        )


def _gap_explanation_required(ecart_volume, volume_tete=None):
    ecart = abs(_decimal_or_zero(ecart_volume))
    if ecart <= GAP_SIGNIFICANCE_THRESHOLD:
        return False

    volume_tete_decimal = abs(_decimal_or_zero(volume_tete))
    if volume_tete_decimal <= 0:
        return True

    gap_percent = (ecart / volume_tete_decimal) * Decimal('100')
    return gap_percent > GAP_PERCENT_TOLERANCE


def _validate_gap_explanation(ecart_volume, explication, volume_tete=None, require_explanation=True):
    explication = (explication or '').strip()
    if require_explanation and _gap_explanation_required(ecart_volume, volume_tete=volume_tete) and not explication:
        raise ValidationError({
            'explication_ecart': "Une explication est requise lorsque l'ecart depasse 10%."
        })
    return explication or None


def _validation_errors(exc):
    if hasattr(exc, 'message_dict'):
        return exc.message_dict
    return {'__all__': list(getattr(exc, 'messages', [str(exc)]))}


def _selected_ouvrage_from_request(user, ouvrage_id=None):
    ouvrages = _get_scoped_ouvrages_for_releves(user)

    if ouvrage_id:
        selected_ouvrage = get_object_or_404(ouvrages, pk=ouvrage_id)
    else:
        selected_ouvrage = None
    return ouvrages, selected_ouvrage


def _get_compteur_from_request(user, compteur_id):
    compteur_id = (compteur_id or '').strip()
    if not compteur_id:
        raise ValidationError({'compteur_id': ["Le compteur est obligatoire."]})
    return get_object_or_404(_allowed_compteurs_queryset(user), id=compteur_id)


@login_required
def compteurs_releves_list(request, ouvrage_id=None):
    ouvrages, selected_ouvrage = _selected_ouvrage_from_request(request.user, ouvrage_id=ouvrage_id)

    compteurs = CompteurTete.objects.none()
    selected_compteur = None
    releves_queryset = ReleveCompteurTete.objects.none()

    compteurs = _allowed_compteurs_queryset(request.user).select_related(
        'ouvrage',
        'changement_sortant',
    ).prefetch_related('releves_compteur_tete').order_by(
        'ouvrage__nom_ouvrage',
        'numerocompteur',
    )
    if selected_ouvrage:
        compteurs = compteurs.filter(ouvrage=selected_ouvrage)

    compteur_filter = (request.GET.get('compteur') or '').strip()
    if compteur_filter:
        selected_compteur = compteurs.filter(id=compteur_filter).first()

    releves_queryset = ReleveCompteurTete.objects.select_related(
        'compteur_tete',
        'compteur_tete__ouvrage',
    ).filter(compteur_tete__in=compteurs)

    if selected_compteur:
        releves_queryset = releves_queryset.filter(compteur_tete=selected_compteur)

    releves_queryset = releves_queryset.order_by('-date_releve', '-created_at')
    initial_releve_ids = _get_initial_releve_ids(releves_queryset)
    base_queryset = releves_queryset.exclude(id__in=initial_releve_ids) if initial_releve_ids else releves_queryset

    requested_view = (
        request.GET.get('vue_releves')
        or request.GET.get('portee_releves')
        or 'groupee'
    ).strip()
    view_mode = 'derniers' if requested_view == 'derniers' else 'groupee'
    ecart_filter = (request.GET.get('ecart_releve') or '').strip()

    if view_mode == 'derniers':
        latest_ids = _get_latest_releve_ids_by_compteur(base_queryset)
        working_queryset = base_queryset.filter(id__in=latest_ids) if latest_ids else base_queryset.none()
    else:
        working_queryset = base_queryset

    overview_rows = [_serialize_releve(releve) for releve in working_queryset]

    if ecart_filter == 'a_expliquer':
        filtered_rows = [row for row in overview_rows if row['explication_required']]
    elif ecart_filter in ['a_investiguer', 'anormal']:
        target_severity = 'investigate' if ecart_filter == 'a_investiguer' else 'critical'
        filtered_rows = [row for row in overview_rows if row['severity'] == target_severity]
    else:
        ecart_filter = ''
        filtered_rows = overview_rows

    grouped_rows = []
    grouped_ouvrages = []
    if view_mode == 'groupee':
        groups_by_compteur = {}
        for row in filtered_rows:
            group = groups_by_compteur.get(row['compteur_id'])
            if group is None:
                group = {
                    'compteur_id': row['compteur_id'],
                    'compteur_numero': row['compteur_numero'],
                    'ouvrage_nom': row['ouvrage_nom'],
                    'nb_points_eau_aval': row['nb_points_eau_aval'],
                    'latest': row,
                    'rows': [],
                }
                groups_by_compteur[row['compteur_id']] = group
                grouped_rows.append(group)
            group['rows'].append(row)
        for group in grouped_rows:
            group['nb_releves'] = len(group['rows'])

        grouped_rows.sort(key=lambda group: (
            (group['ouvrage_nom'] or '').lower(),
            (group['compteur_numero'] or '').lower(),
        ))
        ouvrages_groups_map = {}
        all_ouvrage_groups = []
        for compteur_group in grouped_rows:
            ouvrage_key = compteur_group['ouvrage_nom'] or 'Ouvrage non renseigne'
            ouvrage_group = ouvrages_groups_map.get(ouvrage_key)
            if ouvrage_group is None:
                ouvrage_group = {
                    'ouvrage_nom': ouvrage_key,
                    'compteurs': [],
                    'nb_compteurs': 0,
                    'nb_releves': 0,
                    'good_count': 0,
                    'acceptable_count': 0,
                    'investigate_count': 0,
                    'critical_count': 0,
                }
                ouvrages_groups_map[ouvrage_key] = ouvrage_group
                all_ouvrage_groups.append(ouvrage_group)
            ouvrage_group['compteurs'].append(compteur_group)
            ouvrage_group['nb_compteurs'] += 1
            ouvrage_group['nb_releves'] += compteur_group['nb_releves']
            severity_count_key = {
                'good': 'good_count',
                'acceptable': 'acceptable_count',
                'investigate': 'investigate_count',
                'critical': 'critical_count',
            }.get(compteur_group['latest']['severity'])
            if severity_count_key:
                ouvrage_group[severity_count_key] += 1

        paginator = Paginator(all_ouvrage_groups, 5)
        page_obj = paginator.get_page(request.GET.get('page', 1))
        grouped_ouvrages = list(page_obj.object_list)
        grouped_releves = [
            compteur_group
            for ouvrage_group in grouped_ouvrages
            for compteur_group in ouvrage_group['compteurs']
        ]
        releves_rows = [row for group in grouped_releves for row in group['rows']]
    else:
        paginator = Paginator(filtered_rows, 25)
        page_obj = paginator.get_page(request.GET.get('page', 1))
        releves_rows = list(page_obj.object_list)
        grouped_releves = []
    severity_rows = filtered_rows
    total_releves_count = len(filtered_rows)
    total_volume_tete = sum((_decimal_or_zero(row.get('volume_tete')) for row in filtered_rows), Decimal('0'))
    total_volume_aval = sum((_decimal_or_zero(row.get('volume_aval')) for row in filtered_rows), Decimal('0'))
    ecarts_sans_explication = sum(1 for releve in severity_rows if releve['explication_required'])
    investigate_count = sum(1 for releve in severity_rows if releve['severity'] == 'investigate')
    critical_count = sum(1 for releve in severity_rows if releve['severity'] == 'critical')
    investigate_percent = (Decimal(investigate_count) / Decimal(total_releves_count) * Decimal('100')) if total_releves_count else Decimal('0')
    critical_percent = (Decimal(critical_count) / Decimal(total_releves_count) * Decimal('100')) if total_releves_count else Decimal('0')

    # Un compteur ne doit apparaitre que dans une seule statistique : son relevé
    # le plus recent dans le perimetre courant definit sa mention.
    latest_status_by_compteur = {}
    for row in overview_rows:
        latest_status_by_compteur.setdefault(row['compteur_id'], row['severity'])
    compteur_status_counts = {
        severity: sum(1 for value in latest_status_by_compteur.values() if value == severity)
        for severity in ('good', 'acceptable', 'investigate', 'critical')
    }
    total_compteurs_classes = len(latest_status_by_compteur)
    compteur_status_percents = {
        severity: (
            Decimal(count) / Decimal(total_compteurs_classes) * Decimal('100')
            if total_compteurs_classes else Decimal('0')
        )
        for severity, count in compteur_status_counts.items()
    }

    compteurs_options = []
    for compteur in compteurs:
        deja_remplace = hasattr(compteur, 'changement_sortant')
        compteurs_options.append({
            'id': str(compteur.id),
            'numero': compteur.numerocompteur,
            'ouvrage_nom': compteur.ouvrage.nom_ouvrage if compteur.ouvrage else '',
            'index_courant': float(_latest_compteur_index(compteur)),
            'nb_releves': compteur.releves_compteur_tete.count(),
            'nb_points_eau': _linked_points_queryset(compteur).count(),
            'statut': compteur.statut,
            'deja_remplace': deja_remplace,
            'peut_remplacer': compteur.statut != 'hors_service' and not deja_remplace,
        })
    total_compteurs_count = len(compteurs_options)

    changements_queryset = ChangementCompteurTete.objects.select_related(
        'ancien_compteur__ouvrage',
        'nouveau_compteur',
        'created_by',
    ).filter(ancien_compteur__in=compteurs)
    if selected_compteur:
        changements_queryset = changements_queryset.filter(
            Q(ancien_compteur=selected_compteur) | Q(nouveau_compteur=selected_compteur)
        )
    changements_compteurs = [
        _serialize_changement_compteur(changement)
        for changement in changements_queryset[:50]
    ]

    context = {
        'title': 'Releves compteurs de tete',
        'ouvrages': ouvrages,
        'selected_ouvrage': selected_ouvrage,
        'compteurs_options': compteurs_options,
        'changements_compteurs': changements_compteurs,
        'motifs_changement_compteur': ChangementCompteurTete.MOTIF_CHOICES,
        'selected_compteur_id': str(selected_compteur.id) if selected_compteur else '',
        'selected_vue_releves': view_mode,
        'selected_portee_releves': 'derniers' if view_mode == 'derniers' else '',
        'selected_ecart_releve': ecart_filter,
        'filter_querystring': (
            (f"&compteur={selected_compteur.id}" if selected_compteur else "")
            + f"&vue_releves={view_mode}"
            + (f"&ecart_releve={ecart_filter}" if ecart_filter else "")
        ),
        'page_obj': page_obj,
        'releves_rows': releves_rows,
        'grouped_releves': grouped_releves,
        'grouped_ouvrages': grouped_ouvrages,
        'stats': {
            'total_releves': total_releves_count,
            'total_volume_tete': float(total_volume_tete),
            'total_volume_aval': float(total_volume_aval),
            'total_ecart': float(total_volume_tete - total_volume_aval),
            'ecarts_sans_explication': ecarts_sans_explication,
            'investigate_count': investigate_count,
            'critical_count': critical_count,
            'investigate_percent': float(investigate_percent),
            'critical_percent': float(critical_percent),
            'total_releves_display': _format_number_with_spaces(total_releves_count, 0),
            'total_volume_tete_display': _format_number_with_spaces(total_volume_tete, 2),
            'total_volume_aval_display': _format_number_with_spaces(total_volume_aval, 2),
            'total_ecart_display': _format_number_with_spaces(total_volume_tete - total_volume_aval, 2),
            'ecarts_sans_explication_display': _format_number_with_spaces(ecarts_sans_explication, 0),
            'investigate_count_display': _format_number_with_spaces(investigate_count, 0),
            'critical_count_display': _format_number_with_spaces(critical_count, 0),
            'compteurs_classes': total_compteurs_classes,
            'compteurs_classes_display': _format_number_with_spaces(total_compteurs_classes, 0),
            'total_compteurs': total_compteurs_count,
            'total_compteurs_display': _format_number_with_spaces(total_compteurs_count, 0),
            'good_compteurs_display': _format_number_with_spaces(compteur_status_counts['good'], 0),
            'acceptable_compteurs_display': _format_number_with_spaces(compteur_status_counts['acceptable'], 0),
            'investigate_compteurs_display': _format_number_with_spaces(compteur_status_counts['investigate'], 0),
            'critical_compteurs_display': _format_number_with_spaces(compteur_status_counts['critical'], 0),
            'good_compteurs_percent': float(compteur_status_percents['good']),
            'acceptable_compteurs_percent': float(compteur_status_percents['acceptable']),
            'investigate_compteurs_percent': float(compteur_status_percents['investigate']),
            'critical_compteurs_percent': float(compteur_status_percents['critical']),
        },
        'est_directeur': _user_is_director(request.user),
    }
    return render(request, 'organization/compteurs_releves_list.html', context)


@login_required
@require_http_methods(["POST"])
def api_compteurs_releves_changement_compteur(request):
    """Remplace un compteur et cloture l'ancien sans produire de volume negatif."""
    try:
        ancien_compteur = _get_compteur_from_request(
            request.user,
            request.POST.get('ancien_compteur_id'),
        )
        if not (request.POST.get('date_changement') or '').strip():
            raise ValidationError({'date_changement': 'La date du changement est obligatoire.'})
        date_changement = _parse_date(request.POST.get('date_changement'), 'date_changement')
        index_retrait = _parse_decimal(
            request.POST.get('index_retrait'),
            'index_retrait',
            "L'index de retrait",
        )
        index_initial_nouveau = _parse_decimal(
            request.POST.get('index_initial_nouveau'),
            'index_initial_nouveau',
            "L'index initial du nouveau compteur",
        )
        nouveau_numero = (request.POST.get('nouveau_numero') or '').strip()
        marque = (request.POST.get('nouvelle_marque') or '').strip() or None
        modele = (request.POST.get('nouveau_modele') or '').strip() or None
        motif = (request.POST.get('motif') or '').strip()
        explication = (request.POST.get('explication') or '').strip()

        if not nouveau_numero:
            raise ValidationError({'nouveau_numero': 'Le numero du nouveau compteur est obligatoire.'})
        if motif not in dict(ChangementCompteurTete.MOTIF_CHOICES):
            raise ValidationError({'motif': 'Selectionnez un motif valide.'})
        if len(explication) < 5:
            raise ValidationError({
                'explication': 'Expliquez le remplacement en au moins 5 caracteres.'
            })
        if index_retrait < 0 or index_initial_nouveau < 0:
            raise ValidationError({'__all__': 'Les index ne peuvent pas etre negatifs.'})
        if nouveau_numero.casefold() == (ancien_compteur.numerocompteur or '').strip().casefold():
            raise ValidationError({'nouveau_numero': 'Le nouveau compteur doit avoir un autre numero.'})
        if CompteurTete.objects.filter(numerocompteur__iexact=nouveau_numero).exists():
            raise ValidationError({'nouveau_numero': 'Ce numero de compteur existe deja.'})

        with transaction.atomic():
            ancien_compteur = CompteurTete.objects.select_for_update().get(id=ancien_compteur.id)
            if ChangementCompteurTete.objects.select_for_update().filter(
                ancien_compteur=ancien_compteur
            ).exists():
                raise ValidationError({
                    'ancien_compteur_id': 'Ce compteur a deja ete remplace.'
                })
            if ancien_compteur.statut == 'hors_service':
                raise ValidationError({
                    'ancien_compteur_id': 'Ce compteur est deja hors service.'
                })
            if CompteurTete.objects.filter(numerocompteur__iexact=nouveau_numero).exists():
                raise ValidationError({'nouveau_numero': 'Ce numero de compteur existe deja.'})
            if ancien_compteur.date_installation and date_changement < ancien_compteur.date_installation:
                raise ValidationError({
                    'date_changement': "La date precede l'installation de l'ancien compteur."
                })

            dernier_releve = ancien_compteur.releves_compteur_tete.select_for_update().order_by(
                '-date_releve', '-created_at', '-id'
            ).first()
            if dernier_releve and date_changement < dernier_releve.date_releve:
                raise ValidationError({
                    'date_changement': (
                        'La date doit etre posterieure ou egale au dernier releve du '
                        f"{dernier_releve.date_releve.strftime('%d/%m/%Y')}."
                    )
                })

            releve_final = ancien_compteur.releves_compteur_tete.select_for_update().filter(
                date_releve=date_changement
            ).first()
            snapshot = _build_snapshot(
                ancien_compteur,
                date_changement,
                index_retrait,
                exclude_releve_id=str(releve_final.id) if releve_final else None,
            )
            motif_display = dict(ChangementCompteurTete.MOTIF_CHOICES)[motif]
            note_remplacement = f"Changement de compteur - {motif_display}. {explication}"

            if releve_final:
                ancienne_explication = (releve_final.explication_ecart or '').strip()
                releve_final.ancien_index = snapshot['ancien_index']
                releve_final.nouvel_index = index_retrait
                releve_final.volume_tete = snapshot['volume_tete']
                releve_final.volume_aval = snapshot['volume_aval']
                releve_final.ecart_volume = snapshot['ecart_volume']
                releve_final.nb_points_eau_aval = snapshot['nb_points_eau_aval']
                releve_final.explication_ecart = (
                    f"{ancienne_explication}\n{note_remplacement}" if ancienne_explication else note_remplacement
                )
                releve_final.save()
            else:
                releve_final = ReleveCompteurTete.objects.create(
                    compteur_tete=ancien_compteur,
                    date_releve=date_changement,
                    ancien_index=snapshot['ancien_index'],
                    nouvel_index=index_retrait,
                    volume_tete=snapshot['volume_tete'],
                    volume_aval=snapshot['volume_aval'],
                    ecart_volume=snapshot['ecart_volume'],
                    nb_points_eau_aval=snapshot['nb_points_eau_aval'],
                    explication_ecart=note_remplacement,
                )
            _resequence_releves(ancien_compteur)

            nouveau_compteur = CompteurTete.objects.create(
                ouvrage=ancien_compteur.ouvrage,
                numerocompteur=nouveau_numero,
                marque=marque,
                modele=modele,
                date_installation=date_changement,
                index_initial=index_initial_nouveau,
                statut='operationnel',
            )
            nb_points_transferes = PointEau.objects.filter(
                compteur_tete=ancien_compteur
            ).update(compteur_tete=nouveau_compteur)

            ancien_compteur.statut = 'hors_service'
            ancien_compteur.save(update_fields=['statut', 'updatedat'])

            changement = ChangementCompteurTete.objects.create(
                ancien_compteur=ancien_compteur,
                nouveau_compteur=nouveau_compteur,
                releve_final_ancien=releve_final,
                date_changement=date_changement,
                index_retrait=index_retrait,
                index_initial_nouveau=index_initial_nouveau,
                motif=motif,
                explication=explication,
                nb_points_transferes=nb_points_transferes,
                created_by=request.user,
            )
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)
    except IntegrityError:
        return JsonResponse({
            'success': False,
            'errors': {'__all__': ['Le changement entre en conflit avec des donnees existantes.']},
        }, status=400)

    return JsonResponse({
        'success': True,
        'message': (
            f"Compteur {ancien_compteur.numerocompteur} remplace par "
            f"{nouveau_compteur.numerocompteur}."
        ),
        'changement': _serialize_changement_compteur(
            ChangementCompteurTete.objects.select_related(
                'ancien_compteur__ouvrage', 'nouveau_compteur', 'created_by'
            ).get(id=changement.id)
        ),
    })


@login_required
@require_GET
def api_compteurs_releves_preview(request):
    compteur_id = (request.GET.get('compteur_id') or '').strip()
    if not compteur_id:
        return JsonResponse({'success': False, 'error': 'Compteur requis.'}, status=400)

    compteur = get_object_or_404(_allowed_compteurs_queryset(request.user), id=compteur_id)

    try:
        date_releve = _parse_date(request.GET.get('date_releve'))
        nouvel_index = _parse_decimal(request.GET.get('nouvel_index'), 'nouvel_index', 'Le nouvel index')
        exclude_releve_id = (request.GET.get('releve_id') or '').strip() or None
        snapshot = _build_snapshot(compteur, date_releve, nouvel_index, exclude_releve_id=exclude_releve_id)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)

    return JsonResponse({
        'success': True,
        'data': {
            'compteur_id': str(compteur.id),
            'compteur_numero': compteur.numerocompteur,
            'ancien_index': float(snapshot['ancien_index']),
            'nouvel_index': float(nouvel_index),
            'volume_tete': float(snapshot['volume_tete']),
            'volume_aval': float(snapshot['volume_aval']),
            'ecart_volume': float(snapshot['ecart_volume']),
            'nb_points_eau_aval': snapshot['nb_points_eau_aval'],
            'points_details': snapshot['points_details'],
            'date_debut_calcul': snapshot['date_debut_calcul'].strftime('%Y-%m-%d'),
            'date_debut_calcul_display': snapshot['date_debut_calcul'].strftime('%d/%m/%Y'),
            'date_fin_calcul': snapshot['date_fin_calcul'].strftime('%Y-%m-%d'),
            'date_fin_calcul_display': snapshot['date_fin_calcul'].strftime('%d/%m/%Y'),
            'duree_calcul_jours': max((snapshot['date_fin_calcul'] - snapshot['date_debut_calcul']).days, 0),
            'aval_estime': snapshot['aval_estime'],
            'explication_required': abs(snapshot['ecart_volume']) > GAP_SIGNIFICANCE_THRESHOLD,
        }
    })


@login_required
@require_GET
def api_compteurs_releves_detail(request, releve_id):
    releve = get_object_or_404(
        ReleveCompteurTete.objects.select_related('compteur_tete', 'compteur_tete__ouvrage').filter(
            compteur_tete__in=_allowed_compteurs_queryset(request.user)
        ),
        id=releve_id,
    )
    previous_releve = _get_previous_releve(
        releve.compteur_tete,
        releve.date_releve,
        exclude_id=releve.id,
    )
    date_debut_calcul = (
        previous_releve.date_releve
        if previous_releve else
        _default_calculation_start(releve.date_releve)
    )
    _, points_details = _get_downstream_breakdown(
        releve.compteur_tete,
        date_debut_calcul,
        releve.date_releve,
    )
    serialized_releve = _serialize_releve(releve)
    serialized_releve['aval_estime'] = any(
        detail['estimation_utilisee'] for detail in points_details
    )
    return JsonResponse({
        'success': True,
        'releve': serialized_releve,
        'points_details': points_details,
        'point_assignment': _build_compteur_points_payload(request.user, releve.compteur_tete),
    })


@login_required
@require_GET
def api_compteurs_points_management(request):
    compteur = _get_compteur_from_request(request.user, request.GET.get('compteur_id'))
    return JsonResponse({
        'success': True,
        'data': _build_compteur_points_payload(request.user, compteur),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_points_attach(request):
    try:
        compteur = _get_compteur_from_request(request.user, request.POST.get('compteur_id'))
        point = _get_allowed_point_for_compteur(request.user, compteur, request.POST.get('point_id'))
        old_compteur = point.compteur_tete

        with transaction.atomic():
            point.compteur_tete = compteur
            point.save(update_fields=['compteur_tete'])
            if old_compteur and old_compteur.id != compteur.id:
                _resequence_releves(old_compteur)
            _resequence_releves(compteur)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)

    return JsonResponse({
        'success': True,
        'message': "Point d'eau affecte au compteur.",
        'data': _build_compteur_points_payload(request.user, compteur),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_points_detach(request, point_id):
    point = get_object_or_404(_allowed_points_queryset(request.user), id=point_id)
    old_compteur = point.compteur_tete

    if old_compteur_id := (request.POST.get('compteur_id') or '').strip():
        if not old_compteur or str(old_compteur.id) != old_compteur_id:
            return JsonResponse({
                'success': False,
                'errors': {'point_id': ["Le point d'eau n'est plus rattache a ce compteur."]},
            }, status=400)

    if not old_compteur:
        return JsonResponse({
            'success': False,
            'errors': {'point_id': ["Le point d'eau n'est rattache a aucun compteur."]},
        }, status=400)

    try:
        with transaction.atomic():
            point.compteur_tete = None
            point.save(update_fields=['compteur_tete'])
            _resequence_releves(old_compteur)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)

    return JsonResponse({
        'success': True,
        'message': "Point d'eau retire du compteur.",
        'data': _build_compteur_points_payload(request.user, old_compteur),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_points_move(request, point_id):
    point = get_object_or_404(_allowed_points_queryset(request.user), id=point_id)
    old_compteur = point.compteur_tete

    try:
        target_compteur = _get_compteur_from_request(request.user, request.POST.get('target_compteur_id'))
        if not old_compteur:
            raise ValidationError({'point_id': ["Le point d'eau n'est rattache a aucun compteur."]})
        if point.ouvrage_id != target_compteur.ouvrage_id:
            raise ValidationError({
                'target_compteur_id': ["Le nouveau compteur doit appartenir au meme ouvrage."]
            })
        if old_compteur.id == target_compteur.id:
            raise ValidationError({
                'target_compteur_id': ["Choisissez un autre compteur pour effectuer le changement."]
            })

        with transaction.atomic():
            point.compteur_tete = target_compteur
            point.save(update_fields=['compteur_tete'])
            _resequence_releves(old_compteur)
            _resequence_releves(target_compteur)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)

    return JsonResponse({
        'success': True,
        'message': "Point d'eau deplace vers un autre compteur.",
        'data': _build_compteur_points_payload(request.user, old_compteur),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_releves_create(request):
    try:
        compteur = _get_compteur_from_request(request.user, request.POST.get('compteur_id'))
        date_releve = _parse_date(request.POST.get('date_releve'))
        nouvel_index = _parse_decimal(request.POST.get('nouvel_index'), 'nouvel_index', 'Le nouvel index')
        snapshot = _build_snapshot(compteur, date_releve, nouvel_index)
        explication_ecart = _validate_gap_explanation(
            snapshot['ecart_volume'],
            request.POST.get('explication_ecart'),
            volume_tete=snapshot['volume_tete'],
        )

        with transaction.atomic():
            releve = ReleveCompteurTete.objects.create(
                compteur_tete=compteur,
                date_releve=date_releve,
                ancien_index=snapshot['ancien_index'],
                nouvel_index=nouvel_index,
                volume_tete=snapshot['volume_tete'],
                volume_aval=snapshot['volume_aval'],
                ecart_volume=snapshot['ecart_volume'],
                nb_points_eau_aval=snapshot['nb_points_eau_aval'],
                explication_ecart=explication_ecart,
            )
            _resequence_releves(compteur)
            releve.refresh_from_db()
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)
    except IntegrityError:
        return JsonResponse({
            'success': False,
            'errors': {'date_releve': ["Un releve existe deja pour ce compteur a cette date."]},
        }, status=400)

    return JsonResponse({
        'success': True,
        'message': 'Releve enregistre avec succes.',
        'releve': _serialize_releve(releve),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_releves_update(request, releve_id):
    releve = get_object_or_404(
        ReleveCompteurTete.objects.select_related('compteur_tete').filter(
            compteur_tete__in=_allowed_compteurs_queryset(request.user)
        ),
        id=releve_id,
    )

    ancien_compteur = releve.compteur_tete

    try:
        compteur = _get_compteur_from_request(request.user, request.POST.get('compteur_id'))
        date_releve = _parse_date(request.POST.get('date_releve'))
        nouvel_index = _parse_decimal(request.POST.get('nouvel_index'), 'nouvel_index', 'Le nouvel index')
        snapshot = _build_snapshot(compteur, date_releve, nouvel_index, exclude_releve_id=str(releve.id))
        explication_ecart = _validate_gap_explanation(
            snapshot['ecart_volume'],
            request.POST.get('explication_ecart'),
            volume_tete=snapshot['volume_tete'],
        )

        with transaction.atomic():
            releve.compteur_tete = compteur
            releve.date_releve = date_releve
            releve.nouvel_index = nouvel_index
            releve.explication_ecart = explication_ecart
            releve.ancien_index = snapshot['ancien_index']
            releve.volume_tete = snapshot['volume_tete']
            releve.volume_aval = snapshot['volume_aval']
            releve.ecart_volume = snapshot['ecart_volume']
            releve.nb_points_eau_aval = snapshot['nb_points_eau_aval']
            releve.save()

            if ancien_compteur_id := getattr(ancien_compteur, 'id', None):
                if ancien_compteur_id != compteur.id:
                    _resequence_releves(ancien_compteur)
            _resequence_releves(compteur)
            releve.refresh_from_db()
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': _validation_errors(exc)}, status=400)
    except IntegrityError:
        return JsonResponse({
            'success': False,
            'errors': {'date_releve': ["Un releve existe deja pour ce compteur a cette date."]},
        }, status=400)

    return JsonResponse({
        'success': True,
        'message': 'Releve modifie avec succes.',
        'releve': _serialize_releve(releve),
    })


@login_required
@require_http_methods(["POST"])
def api_compteurs_releves_delete(request, releve_id):
    releve = get_object_or_404(
        ReleveCompteurTete.objects.select_related('compteur_tete').filter(
            compteur_tete__in=_allowed_compteurs_queryset(request.user)
        ),
        id=releve_id,
    )
    compteur = releve.compteur_tete

    try:
        with transaction.atomic():
            releve.delete()
            _resequence_releves(compteur)
    except ValidationError as exc:
        return JsonResponse({'success': False, 'errors': {'__all__': exc.messages}}, status=400)

    return JsonResponse({
        'success': True,
        'message': 'Releve supprime avec succes.',
    })


@login_required
@require_GET
def compteurs_releves_template_excel(request, ouvrage_id=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Template Releves'

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = ['ID Compteur', 'Numero Compteur', 'Date Releve (JJ/MM/AAAA)', 'Nouvel Index', 'Explication (si ecart)']
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    ouvrages, selected_ouvrage = _selected_ouvrage_from_request(request.user, ouvrage_id)
    compteurs_qs = _allowed_compteurs_queryset(request.user)
    if selected_ouvrage:
         compteurs_qs = compteurs_qs.filter(ouvrage=selected_ouvrage)

    for row_num, compteur in enumerate(compteurs_qs.order_by('ouvrage__nom_ouvrage', 'numerocompteur')[:50], 2):
        ws.cell(row=row_num, column=1, value=str(compteur.id))
        ws.cell(row=row_num, column=2, value=str(compteur.numerocompteur))
        ws.cell(row=row_num, column=3, value=timezone.now().strftime('%d/%m/%Y'))
        ws.cell(row=row_num, column=4, value='')
        ws.cell(row=row_num, column=5, value='')

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="template_releves_compteurs.xlsx"'
    wb.save(response)
    return response


@login_required
@require_GET
def compteurs_releves_export_excel(request, ouvrage_id=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ouvrages, selected_ouvrage = _selected_ouvrage_from_request(request.user, ouvrage_id)
    
    releves_qs = ReleveCompteurTete.objects.select_related(
        'compteur_tete', 'compteur_tete__ouvrage'
    ).filter(
        compteur_tete__in=_allowed_compteurs_queryset(request.user)
    )

    if selected_ouvrage:
        releves_qs = releves_qs.filter(compteur_tete__ouvrage=selected_ouvrage)

    releves_qs = releves_qs.order_by('compteur_tete__numerocompteur', '-date_releve')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Releves Compteurs'

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        'ID Releve', 'ID Compteur', 'Numero Compteur', 'Nom Ouvrage', 'Date Releve', 
        'Ancien Index', 'Nouvel Index', 'Volume Tete', 'Volume Aval', 'Ecart Volume', 
        'Nb Points Eau Aval', 'Explication Ecart'
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row_num, releve in enumerate(releves_qs, 2):
        ws.cell(row=row_num, column=1, value=str(releve.id))
        ws.cell(row=row_num, column=2, value=str(releve.compteur_tete.id))
        ws.cell(row=row_num, column=3, value=releve.compteur_tete.numerocompteur)
        ws.cell(row=row_num, column=4, value=releve.compteur_tete.ouvrage.nom_ouvrage if releve.compteur_tete.ouvrage else '')
        ws.cell(row=row_num, column=5, value=releve.date_releve.strftime('%d/%m/%Y'))
        ws.cell(row=row_num, column=6, value=float(_decimal_or_zero(releve.ancien_index)))
        ws.cell(row=row_num, column=7, value=float(_decimal_or_zero(releve.nouvel_index)))
        ws.cell(row=row_num, column=8, value=float(_decimal_or_zero(releve.volume_tete)))
        ws.cell(row=row_num, column=9, value=float(_decimal_or_zero(releve.volume_aval)))
        ws.cell(row=row_num, column=10, value=float(_decimal_or_zero(releve.ecart_volume)))
        ws.cell(row=row_num, column=11, value=int(releve.nb_points_eau_aval or 0))
        ws.cell(row=row_num, column=12, value=releve.explication_ecart or '')

    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    
    filename = "export_releves_compteurs.xlsx"
    if selected_ouvrage:
        import re
        safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', selected_ouvrage.nom_ouvrage)
        filename = f"export_releves_compteurs_{safe_name}.xlsx"
    
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@require_http_methods(["POST"])
def compteurs_releves_import_excel(request, ouvrage_id=None):
    from openpyxl import load_workbook
    from datetime import datetime

    redirect_url = 'organization:compteurs_releves_list_ouvrage' if ouvrage_id else 'organization:compteurs_releves_list'
    kwargs = {'ouvrage_id': ouvrage_id} if ouvrage_id else {}

    if 'excel_file' not in request.FILES:
        messages.error(request, 'Aucun fichier sélectionné.')
        return redirect(redirect_url, **kwargs)

    excel_file = request.FILES['excel_file']
    if not (excel_file.name.endswith('.xlsx') or excel_file.name.endswith('.xls')):
        messages.error(request, 'Le fichier doit être un fichier Excel (.xlsx ou .xls).')
        return redirect(redirect_url, **kwargs)

    try:
        wb = load_workbook(excel_file, data_only=True)
        ws = wb.active

        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value or '').lower().replace(' ', '_').strip())

        if not headers or 'nouvel_index' not in headers:
             messages.error(request, "Format de fichier invalide. Utilisez le template Excel recommandé.")
             return redirect(redirect_url, **kwargs)

        created_count = 0
        updated_count = 0
        error_messages = []

        allowed_compteurs = {str(c.id): c for c in _allowed_compteurs_queryset(request.user)}

        for row_num in range(2, ws.max_row + 1):
            row_data = {}
            for col_num, header in enumerate(headers, 1):
                if not header: continue
                val = ws.cell(row=row_num, column=col_num).value
                if isinstance(val, datetime):
                     val = val.strftime('%d/%m/%Y')
                row_data[header] = str(val).strip() if val is not None else ''

            id_compteur = row_data.get('id_compteur', '')
            numero = row_data.get('numero_compteur', '')
            nouvel_index_str = row_data.get('nouvel_index', '')
            date_releve_str = row_data.get('date_releve_(jj/mm/aaaa)', '') or row_data.get('date_releve', '')
            explication = row_data.get('explication_(si_ecart)', '') or row_data.get('explication_ecart', '') or row_data.get('explication', '')

            if not nouvel_index_str and not date_releve_str:
                continue

            try:
                compteur = None
                if id_compteur in allowed_compteurs:
                    compteur = allowed_compteurs[id_compteur]
                else:
                    found = [c for c in allowed_compteurs.values() if c.numerocompteur == numero]
                    if found: 
                        compteur = found[0]

                if not compteur:
                    raise Exception(f"Compteur introuvable ou non autorisé (ID: '{id_compteur}', numero: '{numero}')")

                try:
                     date_releve = _parse_date(date_releve_str)
                except ValidationError:
                     raise Exception(f"Date de relevé invalide ou manquante: {date_releve_str}")

                try:
                     nouvel_index = _parse_decimal(nouvel_index_str, 'nouvel_index', 'Le nouvel index')
                except ValidationError:
                     raise Exception(f"Nouvel index invalide ou manquant: {nouvel_index_str}")

                existing_releve = ReleveCompteurTete.objects.filter(compteur_tete=compteur, date_releve=date_releve).first()
                if existing_releve:
                    snapshot = _build_snapshot(compteur, date_releve, nouvel_index, exclude_releve_id=str(existing_releve.id))
                    explic = _validate_gap_explanation(
                        snapshot['ecart_volume'],
                        explication,
                        volume_tete=snapshot['volume_tete'],
                        require_explanation=False,
                    )
                    with transaction.atomic():
                        existing_releve.nouvel_index = nouvel_index
                        existing_releve.ancien_index = snapshot['ancien_index']
                        existing_releve.volume_tete = snapshot['volume_tete']
                        existing_releve.volume_aval = snapshot['volume_aval']
                        existing_releve.ecart_volume = snapshot['ecart_volume']
                        existing_releve.nb_points_eau_aval = snapshot['nb_points_eau_aval']
                        existing_releve.explication_ecart = explic
                        existing_releve.save()
                        _resequence_releves(compteur)
                    updated_count += 1
                else:
                    snapshot = _build_snapshot(compteur, date_releve, nouvel_index)
                    explic = _validate_gap_explanation(
                        snapshot['ecart_volume'],
                        explication,
                        volume_tete=snapshot['volume_tete'],
                        require_explanation=False,
                    )
                    with transaction.atomic():
                        ReleveCompteurTete.objects.create(
                            compteur_tete=compteur,
                            date_releve=date_releve,
                            ancien_index=snapshot['ancien_index'],
                            nouvel_index=nouvel_index,
                            volume_tete=snapshot['volume_tete'],
                            volume_aval=snapshot['volume_aval'],
                            ecart_volume=snapshot['ecart_volume'],
                            nb_points_eau_aval=snapshot['nb_points_eau_aval'],
                            explication_ecart=explic,
                        )
                        _resequence_releves(compteur)
                    created_count += 1

            except Exception as e:
                msg = getattr(e, 'messages', [str(e)])[0] if hasattr(e, 'messages') else str(e)
                error_messages.append(f"Ligne {row_num}: {msg}")

        if created_count > 0:
            messages.success(request, f"{created_count} relevé(s) importé(s) avec succès.")
        if updated_count > 0:
            messages.success(request, f"{updated_count} relevé(s) mis à jour avec succès.")
            
        if error_messages:
            for error in error_messages[:5]:
                messages.error(request, error)
            if len(error_messages) > 5:
                messages.warning(request, f"...et {len(error_messages) - 5} autres erreurs.")
        elif created_count == 0 and updated_count == 0:
            messages.warning(request, "Aucune donnée n'a été importée.")

    except Exception as e:
        messages.error(request, f"Erreur système lors de l'import : {str(e)}")

    return redirect(redirect_url, **kwargs)
