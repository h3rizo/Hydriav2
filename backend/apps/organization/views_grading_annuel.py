from __future__ import annotations

from collections import defaultdict
from datetime import datetime
import re
import unicodedata

from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Exists, OuterRef, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.billing.models import ModaliteGestion, PointEau
from apps.employees.permissions import (
    employe_required,
    is_directeur_poste,
    is_visiteur,
)
from apps.geography.models import Commune, District, Region
from apps.planning.repair_classification import get_planning_repair_family
from apps.planning.models_reparations import Reparation

from .forms import OuvrageGradingAnnuelForm
from .models import Organisation, Ouvrage, OuvrageGradingAnnuel
from .scoping import current_modalite_q, get_user_accessible_gestionnaire_ids

GRADE_ORDER = ['A', 'B', 'B_PANNE', 'C', 'D', 'D_PANNE']
GRADE_COLORS = {
    'A': '#3b82f6',
    'B': '#22c55e',
    'B_PANNE': '#ef4444',
    'C': '#64748b',
    'D': '#f59e0b',
    'D_PANNE': '#ef4444',
}


def _normalize_header(value: object) -> str:
    txt = str(value or '').strip().lower()
    txt = unicodedata.normalize('NFKD', txt)
    txt = ''.join(ch for ch in txt if not unicodedata.combining(ch))
    txt = re.sub(r'[^a-z0-9]+', '_', txt).strip('_')
    return txt


def _normalize_grade(value: object) -> str:
    raw = str(value or '').strip().upper()
    raw = raw.replace('-', '_').replace(' ', '_')
    raw = re.sub(r'[^A-Z0-9_]+', '', raw)
    raw = re.sub(r'_+', '_', raw).strip('_')

    aliases = {
        'BPANNE': 'B_PANNE',
        'DPANNE': 'D_PANNE',
        'BPANNEE': 'B_PANNE',
        'DPANNEE': 'D_PANNE',
    }
    raw = aliases.get(raw, raw)
    return raw


def _allowed_grades() -> set[str]:
    return {choice[0] for choice in OuvrageGradingAnnuel.GRADE_CHOICES}


def _normalize_filter_value(value: str) -> str:
    v = (value or '').strip()
    return '' if v == 'all' else v


def _extract_ouvrage_ids(ouvrages) -> list:
    if hasattr(ouvrages, 'values_list'):
        return list(ouvrages.values_list('id', flat=True))

    ids = []
    for item in ouvrages or []:
        ids.append(getattr(item, 'id', item))
    return ids


def _apply_geo_filters_to_ouvrages(qs, region_id='', district_id='', commune_id=''):
    if not any([region_id, district_id, commune_id]):
        return qs

    point_eaux = PointEau.objects.filter(ouvrage_id=OuterRef('pk'))
    if region_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune__district__region_id=region_id
        )
    if district_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune__district_id=district_id
        )
    if commune_id:
        point_eaux = point_eaux.filter(
            parcelle__village__fokontany__commune_id=commune_id
        )

    return qs.annotate(
        _matches_geo_filters=Exists(point_eaux)
    ).filter(_matches_geo_filters=True)


def _apply_organisation_filter_to_ouvrages(qs, organisation_id=''):
    if organisation_id:
        qs = qs.annotate(
            _has_selected_organisation=Exists(
                ModaliteGestion.objects.filter(
                    current_modalite_q(),
                    ouvrage_id=OuterRef('pk'),
                    gestionnaire_id=organisation_id,
                    statut='actif',
                )
            )
        ).filter(_has_selected_organisation=True)
    return qs


def _get_allowed_organisation_ids(user):
    gestionnaire_ids = get_user_accessible_gestionnaire_ids(user)
    if gestionnaire_ids is None:
        return None
    return [str(_id) for _id in gestionnaire_ids]


def _resolve_organisation_filter_value(request, allowed_ids):
    if allowed_ids is None:
        default_value = 'all'
    elif len(allowed_ids) == 1:
        default_value = allowed_ids[0]
    else:
        default_value = 'all'

    organisation_raw = request.GET.get('organisation_id', default_value)
    organisation_id = _normalize_filter_value(organisation_raw)

    if organisation_id and allowed_ids is not None:
        allowed_set = set(allowed_ids)
        if organisation_id not in allowed_set:
            organisation_id = _normalize_filter_value(default_value)

    return organisation_id


def _grade_label(grade_value: str) -> str:
    for value, label in OuvrageGradingAnnuel.GRADE_CHOICES:
        if value == grade_value:
            return label
    return grade_value


def _accessible_ouvrages(user):
    qs = Ouvrage.objects.filter(
        type_ouvrage__existe=True
    ).select_related('type_ouvrage')

    if getattr(user, 'is_superuser', False) or is_visiteur(user):
        return qs.order_by('nom_ouvrage')

    employe = getattr(user, 'employe', None)
    if not employe:
        return Ouvrage.objects.none()

    if not getattr(employe, 'base', None):
        return qs.order_by('nom_ouvrage')

    gestionnaire_ids = get_user_accessible_gestionnaire_ids(user)
    if gestionnaire_ids is None:
        return qs.order_by('nom_ouvrage')
    if not gestionnaire_ids:
        return Ouvrage.objects.none()

    matching_modalites = ModaliteGestion.objects.filter(
        current_modalite_q(),
        ouvrage_id=OuterRef('pk'),
        gestionnaire_id__in=gestionnaire_ids,
    )
    unassigned_modalites = ModaliteGestion.objects.filter(
        current_modalite_q(),
        ouvrage_id=OuterRef('pk'),
        gestionnaire__isnull=True,
    )
    any_modalite = ModaliteGestion.objects.filter(
        current_modalite_q(),
        ouvrage_id=OuterRef('pk'),
    )

    qs = qs.annotate(
        _has_matching_modalite=Exists(matching_modalites),
        _has_unassigned_modalite=Exists(unassigned_modalites),
        _has_any_modalite=Exists(any_modalite),
    ).filter(
        Q(_has_matching_modalite=True) |
        Q(_has_unassigned_modalite=True) |
        Q(_has_any_modalite=False)
    )
    return qs.order_by('nom_ouvrage')


def _can_manage_grading_annuel(user) -> bool:
    return bool(is_directeur_poste(user))


def _apply_list_filters(qs, request, include_grade=True):
    search = request.GET.get('q', '').strip()
    annee = request.GET.get('annee', '').strip()
    grade = request.GET.get('grade', '').strip()
    region_id = request.GET.get('region_id', '').strip()
    district_id = request.GET.get('district_id', '').strip()
    commune_id = request.GET.get('commune_id', '').strip()

    region_id = _normalize_filter_value(region_id)
    district_id = _normalize_filter_value(district_id)
    commune_id = _normalize_filter_value(commune_id)

    if search:
        qs = qs.filter(
            Q(ouvrage__nom_ouvrage__icontains=search) |
            Q(ouvrage__type_ouvrage__nom_type_ouvrage__icontains=search) |
            Q(observation__icontains=search)
        )
    if region_id:
        qs = qs.filter(
            ouvrage__pointeau__parcelle__village__fokontany__commune__district__region_id=region_id
        )
    if district_id:
        qs = qs.filter(
            ouvrage__pointeau__parcelle__village__fokontany__commune__district_id=district_id
        )
    if commune_id:
        qs = qs.filter(
            ouvrage__pointeau__parcelle__village__fokontany__commune_id=commune_id
        )
    if annee:
        try:
            qs = qs.filter(annee=int(annee))
        except ValueError:
            pass
    if include_grade and grade:
        qs = qs.filter(grade=grade)

    return qs.distinct()


def _geo_filter_options(ouvrages_accessibles, region_id='', district_id=''):
    ouvrage_ids = _extract_ouvrage_ids(ouvrages_accessibles)
    if not ouvrage_ids:
        return Region.objects.none(), District.objects.none(), Commune.objects.none()

    point_eaux = PointEau.objects.filter(ouvrage_id__in=ouvrage_ids)

    regions = Region.objects.filter(
        id__in=point_eaux.exclude(
            parcelle__village__fokontany__commune__district__region_id__isnull=True
        ).values_list(
            'parcelle__village__fokontany__commune__district__region_id',
            flat=True,
        )
    ).order_by('nom_region')

    districts_source = point_eaux
    if region_id:
        districts_source = districts_source.filter(
            parcelle__village__fokontany__commune__district__region_id=region_id
        )

    districts = District.objects.filter(
        id__in=districts_source.exclude(
            parcelle__village__fokontany__commune__district_id__isnull=True
        ).values_list(
            'parcelle__village__fokontany__commune__district_id',
            flat=True,
        )
    ).select_related('region').order_by('nom_district')

    communes_source = districts_source
    if district_id:
        communes_source = communes_source.filter(
            parcelle__village__fokontany__commune__district_id=district_id
        )

    communes = Commune.objects.filter(
        id__in=communes_source.exclude(
            parcelle__village__fokontany__commune_id__isnull=True
        ).values_list(
            'parcelle__village__fokontany__commune_id',
            flat=True,
        )
    ).select_related('district').order_by('nom_commune')

    return regions, districts, communes


def _organisation_filter_options(ouvrages_accessibles, region_id='', district_id='', commune_id=''):
    ouvrage_ids = _extract_ouvrage_ids(ouvrages_accessibles)
    if not ouvrage_ids:
        return Organisation.objects.none()

    ouvrages_scope = _apply_geo_filters_to_ouvrages(
        Ouvrage.objects.filter(id__in=ouvrage_ids),
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )
    ouvrages_scope_ids = _extract_ouvrage_ids(ouvrages_scope)
    if not ouvrages_scope_ids:
        return Organisation.objects.none()

    organisation_ids = (
        ModaliteGestion.objects.filter(
            current_modalite_q(),
            ouvrage_id__in=ouvrages_scope_ids,
            statut='actif',
            gestionnaire_id__isnull=False,
        )
        .values_list('gestionnaire_id', flat=True)
        .distinct()
    )
    return Organisation.objects.filter(id__in=organisation_ids).order_by('nom_organisation')


def _scoped_organisation_filter_options(ouvrages_accessibles, allowed_ids, region_id='', district_id='', commune_id=''):
    organisations = _organisation_filter_options(
        ouvrages_accessibles,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )
    if allowed_ids is not None:
        if allowed_ids:
            organisations = organisations.filter(id__in=allowed_ids)
        else:
            organisations = Organisation.objects.none()
    return organisations.order_by('nom_organisation')


def _location_by_ouvrage_ids(ouvrage_ids):
    if not ouvrage_ids:
        return {}

    point_eaux = (
        PointEau.objects.filter(ouvrage_id__in=ouvrage_ids)
        .select_related('parcelle__village__fokontany__commune__district__region')
        .order_by('ouvrage_id', 'id')
    )

    mapping = {}
    for pe in point_eaux:
        if pe.ouvrage_id in mapping:
            continue

        region_nom = district_nom = commune_nom = 'Non defini'
        parcelle = getattr(pe, 'parcelle', None)
        village = getattr(parcelle, 'village', None) if parcelle else None
        fokontany = getattr(village, 'fokontany', None) if village else None
        commune = getattr(fokontany, 'commune', None) if fokontany else None
        district = getattr(commune, 'district', None) if commune else None
        region = getattr(district, 'region', None) if district else None

        if commune and commune.nom_commune:
            commune_nom = commune.nom_commune
        if district and district.nom_district:
            district_nom = district.nom_district
        if region and region.nom_region:
            region_nom = region.nom_region

        mapping[pe.ouvrage_id] = (region_nom, district_nom, commune_nom)

    return mapping


def _build_grading_cards(gradings_items):
    cards_map = {}

    for item in gradings_items:
        card = cards_map.get(item.ouvrage_id)
        if not card:
            card = {
                'ouvrage_id': str(item.ouvrage_id),
                'nom_ouvrage': item.ouvrage.nom_ouvrage if item.ouvrage else 'Sans nom',
                'type_ouvrage': (
                    item.ouvrage.type_ouvrage.nom_type_ouvrage
                    if item.ouvrage and item.ouvrage.type_ouvrage
                    else 'Non defini'
                ),
                'saisonnier': bool(item.ouvrage.saisonnier) if item.ouvrage else False,
                'region_nom': getattr(item, 'region_nom', 'Non defini'),
                'district_nom': getattr(item, 'district_nom', 'Non defini'),
                'commune_nom': getattr(item, 'commune_nom', 'Non defini'),
                'gradings': [],
            }
            cards_map[item.ouvrage_id] = card

        card['gradings'].append({
            'annee': item.annee,
            'grade': item.grade,
            'source_display': item.get_source_display(),
            'updated_at': item.updated_at,
        })

    cards = list(cards_map.values())
    for card in cards:
        card['gradings'].sort(
            key=lambda x: -int(x['annee'] or 0)
        )
        card['latest_grading'] = card['gradings'][0] if card['gradings'] else None

    cards.sort(
        key=lambda c: (
            c['region_nom'] or '',
            c['district_nom'] or '',
            c['commune_nom'] or '',
            (c['nom_ouvrage'] or '').lower(),
        )
    )
    return cards


def _build_ouvrage_rows(ouvrage_ids, grade_by_ouvrage=None):
    if not ouvrage_ids:
        return []

    ids = list(ouvrage_ids)
    location_map = _location_by_ouvrage_ids(ids)
    ouvrages = (
        Ouvrage.objects.filter(id__in=ids)
        .select_related('type_ouvrage')
        .order_by('nom_ouvrage')
    )

    rows = []
    for ouvrage in ouvrages:
        region_nom, district_nom, commune_nom = location_map.get(
            ouvrage.id,
            ('Non defini', 'Non defini', 'Non defini'),
        )
        row = {
            'ouvrage_id': str(ouvrage.id),
            'nom_ouvrage': ouvrage.nom_ouvrage or 'Sans nom',
            'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
            'saisonnier': bool(ouvrage.saisonnier),
            'region_nom': region_nom,
            'district_nom': district_nom,
            'commune_nom': commune_nom,
        }
        if grade_by_ouvrage is not None and ouvrage.id in grade_by_ouvrage:
            grade_value = grade_by_ouvrage[ouvrage.id]
            row['grade_value'] = grade_value
            row['grade_label'] = _grade_label(grade_value)
        rows.append(row)

    rows.sort(
        key=lambda r: (
            r['region_nom'] or '',
            r['district_nom'] or '',
            r['commune_nom'] or '',
            (r['nom_ouvrage'] or '').lower(),
        )
    )
    return rows


@employe_required
def ouvrage_grading_annuel_list(request):
    ouvrages_accessibles = _accessible_ouvrages(request.user)
    allowed_organisation_ids = _get_allowed_organisation_ids(request.user)
    organisation_id = _resolve_organisation_filter_value(request, allowed_organisation_ids)
    region_id = _normalize_filter_value(request.GET.get('region_id', ''))
    district_id = _normalize_filter_value(request.GET.get('district_id', ''))
    commune_id = _normalize_filter_value(request.GET.get('commune_id', ''))

    base_qs = OuvrageGradingAnnuel.objects.filter(
        ouvrage__in=ouvrages_accessibles
    ).select_related('ouvrage', 'ouvrage__type_ouvrage')
    if organisation_id:
        base_qs = base_qs.filter(
            ouvrage__modalites_gestion__gestionnaire_id=organisation_id,
            ouvrage__modalites_gestion__statut='actif',
        ).distinct()

    stats_qs = _apply_list_filters(base_qs, request, include_grade=False)
    gradings_qs = _apply_list_filters(base_qs, request, include_grade=True)
    gradings_items = list(gradings_qs)

    location_map = _location_by_ouvrage_ids([item.ouvrage_id for item in gradings_items])
    for item in gradings_items:
        region_nom, district_nom, commune_nom = location_map.get(
            item.ouvrage_id,
            ('Non defini', 'Non defini', 'Non defini'),
        )
        item.region_nom = region_nom
        item.district_nom = district_nom
        item.commune_nom = commune_nom

    gradings_items.sort(
        key=lambda x: (
            x.region_nom or '',
            x.district_nom or '',
            x.commune_nom or '',
            (x.ouvrage.nom_ouvrage or '').lower(),
            -int(x.annee or 0),
        )
    )
    grading_cards = _build_grading_cards(gradings_items)

    paginator = Paginator(gradings_items, 80)
    page_obj = paginator.get_page(request.GET.get('page', 1))

    annees = (
        OuvrageGradingAnnuel.objects.filter(ouvrage__in=ouvrages_accessibles)
        .values_list('annee', flat=True)
        .distinct()
        .order_by('-annee')
    )

    grade_counts = {
        row['grade']: row['total']
        for row in stats_qs.order_by().values('grade').annotate(total=Count('ouvrage_id', distinct=True))
    }
    grade_stats = [
        {
            'value': value,
            'label': label,
            'count': grade_counts.get(value, 0),
        }
        for value, label in OuvrageGradingAnnuel.GRADE_CHOICES
    ]

    ouvrages_for_geo_options = _apply_organisation_filter_to_ouvrages(
        ouvrages_accessibles,
        organisation_id=organisation_id,
    )
    regions, districts, communes = _geo_filter_options(
        ouvrages_for_geo_options,
        region_id=region_id,
        district_id=district_id,
    )
    organisations = _scoped_organisation_filter_options(
        ouvrages_accessibles,
        allowed_organisation_ids,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )

    context = {
        'page_obj': page_obj,
        'gradings': page_obj.object_list,
        'grade_choices': OuvrageGradingAnnuel.GRADE_CHOICES,
        'grade_stats': grade_stats,
        'stats_total_ouvrages': stats_qs.values('ouvrage_id').distinct().count(),
        'stats_total_lignes': len(gradings_items),
        'grading_cards': grading_cards,
        'annees': annees,
        'organisations': organisations,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'filters': {
            'q': request.GET.get('q', '').strip(),
            'organisation_id': organisation_id,
            'annee': request.GET.get('annee', '').strip(),
            'grade': request.GET.get('grade', '').strip(),
            'region_id': region_id,
            'district_id': district_id,
            'commune_id': commune_id,
        },
        'title': 'Gradings annuels des ouvrages',
        'is_directeur': _can_manage_grading_annuel(request.user),
    }
    return render(request, 'organization/ouvrage_grading_annuel_list.html', context)


def build_ouvrage_grading_annuel_analyse_payload(request, filters=None):
    filters = filters or {}
    ouvrages_accessibles = _accessible_ouvrages(request.user)

    allowed_organisation_ids = _get_allowed_organisation_ids(request.user)
    if 'organisation_id' in filters:
        organisation_id = _normalize_filter_value(filters.get('organisation_id', ''))
        if allowed_organisation_ids is not None and organisation_id:
            allowed_ids_as_str = {str(item) for item in allowed_organisation_ids}
            if organisation_id not in allowed_ids_as_str:
                organisation_id = ''
    else:
        organisation_id = _resolve_organisation_filter_value(request, allowed_organisation_ids)

    region_id = _normalize_filter_value(filters.get('region_id', request.GET.get('region_id', '')))
    district_id = _normalize_filter_value(filters.get('district_id', request.GET.get('district_id', '')))
    commune_id = _normalize_filter_value(filters.get('commune_id', request.GET.get('commune_id', '')))

    ouvrages_scope = _apply_geo_filters_to_ouvrages(
        ouvrages_accessibles,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )
    ouvrages_scope = _apply_organisation_filter_to_ouvrages(
        ouvrages_scope,
        organisation_id=organisation_id,
    )
    ouvrage_ids_scoped = _extract_ouvrage_ids(ouvrages_scope)
    ouvrages_filtres_count = len(ouvrage_ids_scoped)

    if ouvrage_ids_scoped:
        gradings_qs = OuvrageGradingAnnuel.objects.filter(
            ouvrage_id__in=ouvrage_ids_scoped
        )
        annees_disponibles = list(
            gradings_qs.values_list('annee', flat=True)
            .distinct()
            .order_by('-annee')
        )
    else:
        gradings_qs = OuvrageGradingAnnuel.objects.none()
        annees_disponibles = []

    annee_from_raw = str(filters.get('annee_from', request.GET.get('annee_from', ''))).strip()
    annee_to_raw = str(filters.get('annee_to', request.GET.get('annee_to', ''))).strip()

    if not annee_to_raw and annees_disponibles:
        annee_to_raw = str(annees_disponibles[0])
    if not annee_from_raw and len(annees_disponibles) > 1:
        annee_from_raw = str(annees_disponibles[1])
    elif not annee_from_raw and annees_disponibles:
        annee_from_raw = str(annees_disponibles[0])

    try:
        annee_from = int(annee_from_raw) if annee_from_raw else None
    except ValueError:
        annee_from = None
    try:
        annee_to = int(annee_to_raw) if annee_to_raw else None
    except ValueError:
        annee_to = None

    grades_from_by_ouvrage = {}
    grades_to_by_ouvrage = {}
    dist_from = {g: 0 for g in GRADE_ORDER}
    dist_to = {g: 0 for g in GRADE_ORDER}

    if annee_from is not None:
        for row in gradings_qs.filter(annee=annee_from).values('ouvrage_id', 'grade'):
            if row['grade'] in dist_from:
                grades_from_by_ouvrage[row['ouvrage_id']] = row['grade']
        for grade_val in grades_from_by_ouvrage.values():
            dist_from[grade_val] += 1

    if annee_to is not None:
        for row in gradings_qs.filter(annee=annee_to).values('ouvrage_id', 'grade'):
            if row['grade'] in dist_to:
                grades_to_by_ouvrage[row['ouvrage_id']] = row['grade']
        for grade_val in grades_to_by_ouvrage.values():
            dist_to[grade_val] += 1

    transitions = defaultdict(lambda: defaultdict(int))
    ouvrages_annee_from = set(grades_from_by_ouvrage.keys())
    ouvrages_annee_to = set(grades_to_by_ouvrage.keys())
    ouvrages_communs = ouvrages_annee_from & ouvrages_annee_to
    ouvrages_uniquement_depart = ouvrages_annee_from - ouvrages_annee_to
    ouvrages_uniquement_arrivee = ouvrages_annee_to - ouvrages_annee_from

    ouvrages_uniquement_depart_rows = _build_ouvrage_rows(
        ouvrages_uniquement_depart,
        grade_by_ouvrage=grades_from_by_ouvrage,
    )
    ouvrages_uniquement_arrivee_rows = _build_ouvrage_rows(
        ouvrages_uniquement_arrivee,
        grade_by_ouvrage=grades_to_by_ouvrage,
    )

    for ouvrage_id in ouvrages_communs:
        g_from = grades_from_by_ouvrage[ouvrage_id]
        g_to = grades_to_by_ouvrage[ouvrage_id]
        transitions[g_from][g_to] += 1

    matrix_rows = []
    transition_top = []
    for g_from in GRADE_ORDER:
        row_cells = []
        row_total = 0
        for g_to in GRADE_ORDER:
            c = int(transitions[g_from][g_to])
            row_cells.append({
                'to_value': g_to,
                'to_label': _grade_label(g_to),
                'count': c,
            })
            row_total += c
            if c > 0:
                transition_top.append({
                    'from_value': g_from,
                    'from': _grade_label(g_from),
                    'to_value': g_to,
                    'to': _grade_label(g_to),
                    'count': c,
                })
        matrix_rows.append({
            'from_value': g_from,
            'from_label': _grade_label(g_from),
            'cells': row_cells,
            'total': row_total,
        })

    transition_top.sort(key=lambda x: (-x['count'], x['from'], x['to']))
    transition_top = transition_top[:20]
    transition_top_max = max((t['count'] for t in transition_top), default=0)
    matrix_rows_nonzero = [row for row in matrix_rows if row['total'] > 0]

    matrix_to_values_with_data = [
        g_to for g_to in GRADE_ORDER
        if any(int(transitions[g_from][g_to]) > 0 for g_from in GRADE_ORDER)
    ]
    matrix_grade_columns = [
        {'value': g, 'label': _grade_label(g)} for g in matrix_to_values_with_data
    ]
    visible_to_values = set(matrix_to_values_with_data)
    for row in matrix_rows_nonzero:
        row['cells_filtered'] = [
            cell for cell in row['cells'] if cell['to_value'] in visible_to_values
        ]
    matrix_empty_colspan = len(matrix_grade_columns) + 2

    chart_labels = [_grade_label(g) for g in GRADE_ORDER]
    chart_datasets = []
    for g_to in GRADE_ORDER:
        chart_datasets.append({
            'label': f"Vers {_grade_label(g_to)}",
            'data': [int(transitions[g_from][g_to]) for g_from in GRADE_ORDER],
            'backgroundColor': GRADE_COLORS.get(g_to, '#64748b'),
        })

    # Analyse multi-annees: utilise toutes les annees de la plage [annee_from, annee_to].
    multi_start = None
    multi_end = None
    if annees_disponibles:
        if annee_from is not None and annee_to is not None:
            multi_start = min(annee_from, annee_to)
            multi_end = max(annee_from, annee_to)
        elif annee_to is not None:
            multi_start = annee_to
            multi_end = annee_to
        elif annee_from is not None:
            multi_start = annee_from
            multi_end = annee_from
        else:
            multi_start = min(annees_disponibles)
            multi_end = max(annees_disponibles)

    multi_years = []
    if multi_start is not None and multi_end is not None:
        multi_years = [
            y for y in sorted(annees_disponibles)
            if multi_start <= y <= multi_end
        ]

    multi_grade_counts_by_year = {
        y: {g: 0 for g in GRADE_ORDER} for y in multi_years
    }
    multi_grade_by_ouvrage = defaultdict(dict)

    if multi_years:
        for row in gradings_qs.filter(annee__in=multi_years).values('ouvrage_id', 'annee', 'grade'):
            grade_val = row['grade']
            annee_val = row['annee']
            if grade_val not in GRADE_ORDER:
                continue
            if annee_val not in multi_grade_counts_by_year:
                continue

            multi_grade_counts_by_year[annee_val][grade_val] += 1
            multi_grade_by_ouvrage[row['ouvrage_id']][annee_val] = grade_val

    multi_grade_columns_values = [
        g for g in GRADE_ORDER
        if any(multi_grade_counts_by_year.get(y, {}).get(g, 0) > 0 for y in multi_years)
    ]

    multi_distribution_columns = [
        {'value': g, 'label': _grade_label(g)}
        for g in multi_grade_columns_values
    ]
    multi_distribution_rows = []
    for y in multi_years:
        counts = multi_grade_counts_by_year.get(y, {})
        multi_distribution_rows.append({
            'annee': y,
            'total': int(sum(counts.values())),
            'grades': [
                {
                    'value': g,
                    'label': _grade_label(g),
                    'count': int(counts.get(g, 0)),
                }
                for g in multi_grade_columns_values
            ],
        })
    multi_distribution_rows_nonzero = [
        row for row in multi_distribution_rows
        if int(row['total']) > 0
    ]

    multi_chart_labels = [str(y) for y in multi_years]
    multi_chart_datasets = []
    for g in multi_grade_columns_values:
        grade_color = GRADE_COLORS.get(g, '#64748b')
        multi_chart_datasets.append({
            'label': _grade_label(g),
            'data': [int(multi_grade_counts_by_year[y].get(g, 0)) for y in multi_years],
            'borderColor': grade_color,
            'backgroundColor': grade_color,
            'tension': 0.25,
            'fill': False,
        })

    multi_stability_counts = {g: 0 for g in GRADE_ORDER}
    multi_complete_count = 0
    multi_unstable_count = 0
    required_years_count = len(multi_years)

    if required_years_count > 0:
        for grade_by_year in multi_grade_by_ouvrage.values():
            if len(grade_by_year) != required_years_count:
                continue

            multi_complete_count += 1
            unique_grades = set(grade_by_year.values())
            if len(unique_grades) == 1:
                stable_grade = next(iter(unique_grades))
                if stable_grade in multi_stability_counts:
                    multi_stability_counts[stable_grade] += 1
            else:
                multi_unstable_count += 1

    multi_stability_total = sum(multi_stability_counts.values())
    multi_incomplete_count = max(ouvrages_filtres_count - multi_complete_count, 0)
    multi_stable_rate = (
        round((multi_stability_total * 100) / multi_complete_count, 1)
        if multi_complete_count else 0
    )

    multi_stability_rows = []
    for g in GRADE_ORDER:
        count_val = int(multi_stability_counts.get(g, 0))
        pct_of_complete = (
            round((count_val * 100) / multi_complete_count, 1)
            if multi_complete_count else 0
        )
        multi_stability_rows.append({
            'value': g,
            'label': _grade_label(g),
            'count': count_val,
            'pct_of_complete': pct_of_complete,
        })
    multi_stability_rows_nonzero = [
        row for row in multi_stability_rows
        if int(row['count']) > 0
    ]

    multi_stability_max = max((row['count'] for row in multi_stability_rows_nonzero), default=0)
    multi_stability_chart_labels = [_grade_label(g) for g in GRADE_ORDER]
    multi_stability_chart_data = [int(multi_stability_counts.get(g, 0)) for g in GRADE_ORDER]
    multi_stability_chart_colors = [GRADE_COLORS.get(g, '#64748b') for g in GRADE_ORDER]

    ouvrages_for_geo_options = _apply_organisation_filter_to_ouvrages(
        ouvrages_accessibles,
        organisation_id=organisation_id,
    )
    ouvrage_ids_for_geo_options = _extract_ouvrage_ids(ouvrages_for_geo_options)
    regions, districts, communes = _geo_filter_options(
        ouvrage_ids_for_geo_options,
        region_id=region_id,
        district_id=district_id,
    )
    organisations = _scoped_organisation_filter_options(
        _extract_ouvrage_ids(ouvrages_accessibles),
        allowed_organisation_ids,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )

    context = {
        'title': 'Analyse des gradings',
        'organisations': organisations,
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'annees_disponibles': annees_disponibles,
        'grade_order': [{'value': g, 'label': _grade_label(g)} for g in GRADE_ORDER],
        'matrix_rows': matrix_rows,
        'matrix_rows_nonzero': matrix_rows_nonzero,
        'matrix_grade_columns': matrix_grade_columns,
        'matrix_empty_colspan': matrix_empty_colspan,
        'transition_top': transition_top,
        'chart_labels': chart_labels,
        'chart_datasets': chart_datasets,
        'dist_from': dist_from,
        'dist_to': dist_to,
        'dist_from_total': sum(dist_from.values()),
        'dist_to_total': sum(dist_to.values()),
        'dist_from_rows': [{'value': g, 'label': _grade_label(g), 'count': dist_from.get(g, 0)} for g in GRADE_ORDER],
        'dist_to_rows': [{'value': g, 'label': _grade_label(g), 'count': dist_to.get(g, 0)} for g in GRADE_ORDER],
        'transition_top_max': transition_top_max,
        'stats': {
            'ouvrages_filtres': ouvrages_filtres_count,
            'avec_grading_from': len(grades_from_by_ouvrage),
            'avec_grading_to': len(grades_to_by_ouvrage),
            'avec_deux_annees': len(ouvrages_communs),
            'uniquement_depart': len(ouvrages_uniquement_depart_rows),
            'uniquement_arrivee': len(ouvrages_uniquement_arrivee_rows),
        },
        'multi_period': {
            'start': multi_start,
            'end': multi_end,
            'years_count': len(multi_years),
        },
        'multi_distribution_columns': multi_distribution_columns,
        'multi_distribution_rows': multi_distribution_rows_nonzero,
        'multi_distribution_empty_colspan': len(multi_distribution_columns) + 2,
        'multi_chart_labels': multi_chart_labels,
        'multi_chart_datasets': multi_chart_datasets,
        'multi_stability_rows': multi_stability_rows_nonzero,
        'multi_stability_max': multi_stability_max,
        'multi_stats': {
            'complete_count': multi_complete_count,
            'stable_count': multi_stability_total,
            'unstable_count': multi_unstable_count,
            'incomplete_count': multi_incomplete_count,
            'stable_rate': multi_stable_rate,
        },
        'multi_stability_chart_labels': multi_stability_chart_labels,
        'multi_stability_chart_data': multi_stability_chart_data,
        'multi_stability_chart_colors': multi_stability_chart_colors,
        'ouvrages_uniquement_depart': ouvrages_uniquement_depart_rows,
        'ouvrages_uniquement_arrivee': ouvrages_uniquement_arrivee_rows,
        'filters': {
            'organisation_id': organisation_id,
            'region_id': region_id,
            'district_id': district_id,
            'commune_id': commune_id,
            'annee_from': annee_from_raw,
            'annee_to': annee_to_raw,
        },
    }
    return context


@employe_required
def ouvrage_grading_annuel_analyse(request):
    context = build_ouvrage_grading_annuel_analyse_payload(request)
    return render(request, 'organization/ouvrage_grading_annuel_analyse.html', context)


@employe_required
def api_ouvrage_grading_annuel_transition_details(request):
    ouvrages_accessibles = _accessible_ouvrages(request.user)

    allowed_organisation_ids = _get_allowed_organisation_ids(request.user)
    organisation_id = _resolve_organisation_filter_value(request, allowed_organisation_ids)
    region_id = _normalize_filter_value(request.GET.get('region_id', ''))
    district_id = _normalize_filter_value(request.GET.get('district_id', ''))
    commune_id = _normalize_filter_value(request.GET.get('commune_id', ''))

    annee_from_raw = request.GET.get('annee_from', '').strip()
    annee_to_raw = request.GET.get('annee_to', '').strip()
    grade_from = request.GET.get('grade_from', '').strip().upper()
    grade_to = request.GET.get('grade_to', '').strip().upper()

    try:
        annee_from = int(annee_from_raw)
        annee_to = int(annee_to_raw)
    except (TypeError, ValueError):
        return JsonResponse({
            'success': False,
            'error': "Annees invalides. Parametres 'annee_from' et 'annee_to' obligatoires.",
        }, status=400)

    allowed_grades = set(GRADE_ORDER)
    if grade_from not in allowed_grades or grade_to not in allowed_grades:
        return JsonResponse({
            'success': False,
            'error': "Grades invalides. Parametres 'grade_from' et 'grade_to' obligatoires.",
        }, status=400)

    ouvrages_scope = _apply_geo_filters_to_ouvrages(
        ouvrages_accessibles,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )
    ouvrages_scope = _apply_organisation_filter_to_ouvrages(
        ouvrages_scope,
        organisation_id=organisation_id,
    )
    ouvrage_ids_scoped = _extract_ouvrage_ids(ouvrages_scope)

    if not ouvrage_ids_scoped:
        return JsonResponse({
            'success': True,
            'rows': [],
            'count': 0,
            'truncated': False,
            'grade_from': grade_from,
            'grade_to': grade_to,
            'grade_from_label': _grade_label(grade_from),
            'grade_to_label': _grade_label(grade_to),
            'annee_from': annee_from,
            'annee_to': annee_to,
        })

    has_from = OuvrageGradingAnnuel.objects.filter(
        ouvrage_id=OuterRef('pk'),
        annee=annee_from,
        grade=grade_from,
    )
    has_to = OuvrageGradingAnnuel.objects.filter(
        ouvrage_id=OuterRef('pk'),
        annee=annee_to,
        grade=grade_to,
    )

    detail_qs = (
        Ouvrage.objects.filter(id__in=ouvrage_ids_scoped)
        .annotate(_has_from=Exists(has_from), _has_to=Exists(has_to))
        .filter(_has_from=True, _has_to=True)
        .select_related('type_ouvrage')
        .order_by('nom_ouvrage')
    )

    total = detail_qs.count()
    limit = 1000
    detail_items = list(detail_qs[:limit])
    location_map = _location_by_ouvrage_ids([obj.id for obj in detail_items])

    rows = []
    for ouvrage in detail_items:
        region_nom, district_nom, commune_nom = location_map.get(
            ouvrage.id,
            ('Non defini', 'Non defini', 'Non defini'),
        )
        rows.append({
            'ouvrage_id': str(ouvrage.id),
            'nom_ouvrage': ouvrage.nom_ouvrage or 'Sans nom',
            'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
            'saisonnier': bool(ouvrage.saisonnier),
            'region_nom': region_nom,
            'district_nom': district_nom,
            'commune_nom': commune_nom,
            'grade_from': grade_from,
            'grade_to': grade_to,
            'annee_from': annee_from,
            'annee_to': annee_to,
        })

    rows.sort(
        key=lambda r: (
            r['region_nom'] or '',
            r['district_nom'] or '',
            r['commune_nom'] or '',
            (r['nom_ouvrage'] or '').lower(),
        )
    )

    return JsonResponse({
        'success': True,
        'rows': rows,
        'count': total,
        'truncated': total > limit,
        'grade_from': grade_from,
        'grade_to': grade_to,
        'grade_from_label': _grade_label(grade_from),
        'grade_to_label': _grade_label(grade_to),
        'annee_from': annee_from,
        'annee_to': annee_to,
    })


@employe_required
def api_ouvrage_reparations_non_terminees(request, ouvrage_id):
    ouvrages_accessibles = _accessible_ouvrages(request.user)

    allowed_organisation_ids = _get_allowed_organisation_ids(request.user)
    organisation_id = _resolve_organisation_filter_value(request, allowed_organisation_ids)
    region_id = _normalize_filter_value(request.GET.get('region_id', ''))
    district_id = _normalize_filter_value(request.GET.get('district_id', ''))
    commune_id = _normalize_filter_value(request.GET.get('commune_id', ''))

    ouvrages_scope = _apply_geo_filters_to_ouvrages(
        ouvrages_accessibles,
        region_id=region_id,
        district_id=district_id,
        commune_id=commune_id,
    )
    ouvrages_scope = _apply_organisation_filter_to_ouvrages(
        ouvrages_scope,
        organisation_id=organisation_id,
    )

    ouvrage = ouvrages_scope.filter(id=ouvrage_id).select_related('type_ouvrage').first()
    if not ouvrage:
        return JsonResponse({
            'success': False,
            'error': "Ouvrage introuvable ou non autorise."
        }, status=404)

    rows = []

    # Source 1: workflow reparations (table reparations)
    reparations_qs = (
        Reparation.objects
        .filter(ouvrage_id=ouvrage.id)
        .exclude(statut__in=['terminee', 'validee', 'annulee'])
        .select_related('assigne_a')
        .order_by('-date_signalement')
    )
    for rep in reparations_qs:
        assigne_nom = ''
        if rep.assigne_a:
            assigne_nom = rep.assigne_a.surnom or rep.assigne_a.nom_employe or ''
        rows.append({
            'id': f"reparation:{rep.id}",
            'titre': rep.titre or '',
            'description': rep.description or '',
            'type_reparation': rep.type_reparation or '',
            'type_reparation_display': rep.get_type_reparation_display(),
            'statut': rep.statut or '',
            'statut_display': rep.get_statut_display(),
            'priorite': rep.priorite or '',
            'priorite_display': rep.get_priorite_display(),
            'date_signalement': rep.date_signalement.isoformat() if rep.date_signalement else '',
            'date_prevue': rep.date_prevue.isoformat() if rep.date_prevue else '',
            'assigne_a': assigne_nom,
            'est_en_retard': bool(rep.est_en_retard),
        })

    # Source 2: workflow planning (reparation/rehabilitation non validees)
    from django.utils import timezone
    from apps.planning.models import Planning

    activites_rehabilitation = Activite.objects.filter(
        Q(nom__icontains='habilitation')
    ).exclude(
        Q(nom__icontains='préparation') | Q(nom__icontains='preparation')
    )

    activites_reparation = Activite.objects.filter(
        Q(nom__icontains='paration')
    ).exclude(
        Q(nom__icontains='préparation') |
        Q(nom__icontains='preparation') |
        Q(nom__icontains='au cas où') |
        Q(nom__icontains='au cas ou') |
        Q(nom__icontains='habilitation')
    )

    rehab_activity_ids = set(activites_rehabilitation.values_list('id', flat=True))
    repair_activity_ids = set(activites_reparation.values_list('id', flat=True))
    rehab_type_ids = set(activites_rehabilitation.values_list('type_activite_id', flat=True))
    repair_type_ids = set(activites_reparation.values_list('type_activite_id', flat=True))

    planning_qs = (
        Planning.objects
        .filter(ouvrage_id=ouvrage.id)
        .filter(Q(estvalide__isnull=True) | Q(estvalide=False))
        .select_related('fiche_poste', 'employe', 'fiche_poste__type_activites', 'type_activite')
        .prefetch_related('fiche_poste__activites_postes')
        .order_by('-date')
    )

    today = timezone.localdate()
    for planning in planning_qs:
        fiche_poste = planning.fiche_poste
        if not fiche_poste:
            continue

        type_activite_id = fiche_poste.type_activites_id
        activite_ids_fiche = {
            ap.activite_id for ap in fiche_poste.activites_postes.all() if ap.activite_id
        }

        is_rehab = bool(
            (type_activite_id and type_activite_id in rehab_type_ids)
            or (activite_ids_fiche & rehab_activity_ids)
        )
        is_repair = bool(
            (type_activite_id and type_activite_id in repair_type_ids)
            or (activite_ids_fiche & repair_activity_ids)
        )

        if not (is_rehab or is_repair):
            text = f"{(fiche_poste.tache or '').lower()} {(planning.remarque or '').lower()}"
            is_rehab = (
                'habilitation' in text
                and 'preparation' not in text
            )
            is_repair = (
                not is_rehab
                and 'paration' in text
                and 'preparation' not in text
                and 'au cas ou' not in text
            )
        family = get_planning_repair_family(planning)
        if family:
            is_rehab = family == 'rehabilitation'
            is_repair = family in ('reparation', 'panne')
        if not (is_rehab or is_repair):
            continue

        date_plan = planning.date.date() if getattr(planning, 'date', None) else None
        assigne_nom = ''
        if planning.employe:
            assigne_nom = planning.employe.surnom or planning.employe.nom_employe or ''

        type_display = 'Rehabilitation' if is_rehab else 'Reparation'
        statut_display = 'Planifiee' if date_plan and date_plan >= today else 'En cours'
        est_en_retard = bool(date_plan and date_plan < today)

        rows.append({
            'id': f"planning:{planning.id}",
            'titre': fiche_poste.tache or type_display,
            'description': planning.remarque or fiche_poste.tache or '',
            'type_reparation': 'rehabilitation' if is_rehab else 'reparation',
            'type_reparation_display': type_display,
            'statut': 'planifiee' if statut_display == 'Planifiee' else 'en_cours',
            'statut_display': statut_display,
            'priorite': '',
            'priorite_display': '-',
            'date_signalement': planning.created_at.isoformat() if getattr(planning, 'created_at', None) else '',
            'date_prevue': planning.date.isoformat() if getattr(planning, 'date', None) else '',
            'assigne_a': assigne_nom,
            'est_en_retard': est_en_retard,
        })

    rows.sort(key=lambda r: (r.get('date_prevue') or r.get('date_signalement') or ''), reverse=True)

    return JsonResponse({
        'success': True,
        'ouvrage': {
            'id': str(ouvrage.id),
            'nom_ouvrage': ouvrage.nom_ouvrage or 'Sans nom',
            'type_ouvrage': ouvrage.type_ouvrage.nom_type_ouvrage if ouvrage.type_ouvrage else 'Non defini',
        },
        'count': len(rows),
        'rows': rows,
    })


@employe_required
def ouvrage_grading_annuel_create(request):
    ouvrages_accessibles = _accessible_ouvrages(request.user)
    form = OuvrageGradingAnnuelForm(request.POST or None)
    form.fields['ouvrage'].queryset = ouvrages_accessibles

    if request.method == 'POST' and form.is_valid():
        grading = form.save(commit=False)
        grading.source = 'manuel'
        grading.save()
        messages.success(request, 'Grading annuel enregistre avec succes.')
        return redirect('organization:ouvrage_grading_annuel_list')

    return render(request, 'organization/ouvrage_grading_annuel_form.html', {
        'form': form,
        'mode': 'create',
        'title': 'Nouveau grading annuel',
    })


@employe_required
def ouvrage_grading_annuel_update(request, pk):
    if not _can_manage_grading_annuel(request.user):
        messages.error(
            request,
            'Acces refuse: seuls les Directeurs peuvent modifier un grading annuel.'
        )
        return redirect('organization:ouvrage_grading_annuel_list')

    ouvrages_accessibles = _accessible_ouvrages(request.user)
    grading = get_object_or_404(
        OuvrageGradingAnnuel.objects.select_related('ouvrage'),
        pk=pk,
        ouvrage__in=ouvrages_accessibles
    )

    form = OuvrageGradingAnnuelForm(request.POST or None, instance=grading)
    form.fields['ouvrage'].queryset = ouvrages_accessibles

    if request.method == 'POST' and form.is_valid():
        grading = form.save(commit=False)
        grading.source = 'manuel'
        grading.save()
        messages.success(request, 'Grading annuel mis a jour.')
        return redirect('organization:ouvrage_grading_annuel_list')

    return render(request, 'organization/ouvrage_grading_annuel_form.html', {
        'form': form,
        'grading': grading,
        'mode': 'edit',
        'title': 'Modifier grading annuel',
    })


@employe_required
def ouvrage_grading_annuel_delete(request, pk):
    if not _can_manage_grading_annuel(request.user):
        messages.error(
            request,
            'Acces refuse: seuls les Directeurs peuvent supprimer un grading annuel.'
        )
        return redirect('organization:ouvrage_grading_annuel_list')

    ouvrages_accessibles = _accessible_ouvrages(request.user)
    grading = get_object_or_404(
        OuvrageGradingAnnuel.objects.select_related('ouvrage'),
        pk=pk,
        ouvrage__in=ouvrages_accessibles
    )

    if request.method == 'POST':
        grading.delete()
        messages.success(request, 'Grading annuel supprime.')
        return redirect('organization:ouvrage_grading_annuel_list')

    return render(request, 'organization/ouvrage_grading_annuel_confirm_delete.html', {
        'grading': grading,
        'title': 'Supprimer grading annuel',
    })


@employe_required
def ouvrage_grading_annuel_export_excel(request):
    ouvrages_accessibles = _accessible_ouvrages(request.user)
    allowed_organisation_ids = _get_allowed_organisation_ids(request.user)
    organisation_id = _resolve_organisation_filter_value(request, allowed_organisation_ids)
    qs = OuvrageGradingAnnuel.objects.filter(
        ouvrage__in=ouvrages_accessibles
    ).select_related('ouvrage', 'ouvrage__type_ouvrage')
    if organisation_id:
        qs = qs.filter(
            ouvrage__modalites_gestion__gestionnaire_id=organisation_id,
            ouvrage__modalites_gestion__statut='actif',
        ).distinct()
    qs = _apply_list_filters(qs, request).order_by('-annee', 'ouvrage__nom_ouvrage')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Gradings annuels'

    headers = [
        'id',
        'ouvrage_id',
        'ouvrage',
        'type_ouvrage',
        'annee',
        'grade',
        'observation',
        'source',
        'updated_at',
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_idx, item in enumerate(qs, start=2):
        ws.cell(row=row_idx, column=1, value=str(item.id))
        ws.cell(row=row_idx, column=2, value=str(item.ouvrage_id))
        ws.cell(row=row_idx, column=3, value=item.ouvrage.nom_ouvrage if item.ouvrage else '')
        ws.cell(
            row=row_idx,
            column=4,
            value=item.ouvrage.type_ouvrage.nom_type_ouvrage if item.ouvrage and item.ouvrage.type_ouvrage else ''
        )
        ws.cell(row=row_idx, column=5, value=item.annee)
        ws.cell(row=row_idx, column=6, value=item.grade)
        ws.cell(row=row_idx, column=7, value=item.observation or '')
        ws.cell(row=row_idx, column=8, value=item.source)
        ws.cell(row=row_idx, column=9, value=item.updated_at.strftime('%Y-%m-%d %H:%M:%S'))

    for column_cells in ws.columns:
        max_len = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = '' if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 60)

    filename = f"gradings_ouvrages_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@employe_required
def ouvrage_grading_annuel_template_excel(request):
    ouvrages_accessibles = _accessible_ouvrages(request.user)[:10]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Template'

    headers = ['ouvrage_id', 'ouvrage', 'annee', 'grade', 'observation']
    ws.append(headers)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')

    row_idx = 2
    for ouvrage in ouvrages_accessibles:
        ws.cell(row=row_idx, column=1, value=str(ouvrage.id))
        ws.cell(row=row_idx, column=2, value=ouvrage.nom_ouvrage)
        ws.cell(row=row_idx, column=3, value='')
        ws.cell(row=row_idx, column=4, value='')
        ws.cell(row=row_idx, column=5, value='')
        row_idx += 1

    help_ws = wb.create_sheet('Aide')
    help_ws.append(['Regles'])
    help_ws.append(['- Colonnes requises: ouvrage_id (ou ouvrage), annee, grade'])
    help_ws.append(['- Grades autorises: A, B, B_PANNE, C, D, D_PANNE'])
    help_ws.append(['- Cle unique: (ouvrage, annee). Import = creation ou mise a jour'])

    for sheet in (ws, help_ws):
        for column_cells in sheet.columns:
            max_len = 0
            letter = column_cells[0].column_letter
            for cell in column_cells:
                value = '' if cell.value is None else str(cell.value)
                if len(value) > max_len:
                    max_len = len(value)
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 14), 70)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="template_gradings_ouvrages.xlsx"'
    wb.save(response)
    return response


@employe_required
def ouvrage_grading_annuel_import_excel(request):
    if request.method != 'POST':
        return redirect('organization:ouvrage_grading_annuel_list')

    excel_file = request.FILES.get('excel_file')
    if not excel_file:
        messages.error(request, 'Aucun fichier Excel fourni.')
        return redirect('organization:ouvrage_grading_annuel_list')

    if not excel_file.name.lower().endswith('.xlsx'):
        messages.error(request, 'Le fichier doit etre au format .xlsx')
        return redirect('organization:ouvrage_grading_annuel_list')

    ouvrages_accessibles = _accessible_ouvrages(request.user)
    ouvrages_by_id = {str(o.id): o for o in ouvrages_accessibles}
    ouvrages_by_name = {}
    for ouvrage in ouvrages_accessibles:
        ouvrages_by_name.setdefault(ouvrage.nom_ouvrage.strip().lower(), []).append(ouvrage)

    try:
        wb = load_workbook(excel_file, data_only=True)
        ws = wb.active
    except Exception as exc:
        messages.error(request, f'Lecture Excel impossible: {exc}')
        return redirect('organization:ouvrage_grading_annuel_list')

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        messages.error(request, 'Le fichier est vide.')
        return redirect('organization:ouvrage_grading_annuel_list')

    headers = [_normalize_header(v) for v in rows[0]]
    idx = {header: i for i, header in enumerate(headers) if header}

    ouvrage_id_idx = None
    ouvrage_name_idx = None
    annee_idx = None
    grade_idx = None
    obs_idx = None

    for alias in ('ouvrage_id', 'id_ouvrage', 'ouvrageid', 'uuid_ouvrage'):
        if alias in idx:
            ouvrage_id_idx = idx[alias]
            break
    for alias in ('ouvrage', 'nom_ouvrage', 'nomouvrage'):
        if alias in idx:
            ouvrage_name_idx = idx[alias]
            break
    for alias in ('annee', 'year'):
        if alias in idx:
            annee_idx = idx[alias]
            break
    for alias in ('grade', 'grading'):
        if alias in idx:
            grade_idx = idx[alias]
            break
    for alias in ('observation', 'observations', 'commentaire', 'commentaires', 'note'):
        if alias in idx:
            obs_idx = idx[alias]
            break

    if annee_idx is None or grade_idx is None or (ouvrage_id_idx is None and ouvrage_name_idx is None):
        messages.error(
            request,
            'Colonnes requises manquantes: ouvrage_id (ou ouvrage), annee, grade.'
        )
        return redirect('organization:ouvrage_grading_annuel_list')

    allowed_grades = _allowed_grades()
    created_count = 0
    updated_count = 0
    errors = []

    with transaction.atomic():
        for excel_row, values in enumerate(rows[1:], start=2):
            if not values:
                continue
            if all((v is None or str(v).strip() == '') for v in values):
                continue

            try:
                ouvrage = None

                if ouvrage_id_idx is not None and ouvrage_id_idx < len(values):
                    ouvrage_id = str(values[ouvrage_id_idx] or '').strip()
                    if ouvrage_id:
                        ouvrage = ouvrages_by_id.get(ouvrage_id)

                if ouvrage is None and ouvrage_name_idx is not None and ouvrage_name_idx < len(values):
                    ouvrage_name = str(values[ouvrage_name_idx] or '').strip().lower()
                    if ouvrage_name:
                        matched = ouvrages_by_name.get(ouvrage_name, [])
                        if len(matched) == 1:
                            ouvrage = matched[0]
                        elif len(matched) > 1:
                            raise ValueError(f"Ouvrage ambigu: {values[ouvrage_name_idx]}")

                if ouvrage is None:
                    raise ValueError('Ouvrage introuvable ou non autorise')

                raw_annee = values[annee_idx] if annee_idx < len(values) else None
                annee = int(raw_annee)
                if annee < 1900 or annee > 2200:
                    raise ValueError(f'Annee invalide: {annee}')

                raw_grade = values[grade_idx] if grade_idx < len(values) else None
                grade = _normalize_grade(raw_grade)
                if grade not in allowed_grades:
                    raise ValueError(f'Grade invalide: {raw_grade}')

                observation = ''
                if obs_idx is not None and obs_idx < len(values):
                    observation = '' if values[obs_idx] is None else str(values[obs_idx]).strip()

                obj, created = OuvrageGradingAnnuel.objects.update_or_create(
                    ouvrage=ouvrage,
                    annee=annee,
                    defaults={
                        'grade': grade,
                        'observation': observation,
                        'source': 'import_excel',
                    }
                )
                if created:
                    created_count += 1
                else:
                    updated_count += 1

            except Exception as exc:
                errors.append(f"Ligne {excel_row}: {exc}")

    if created_count or updated_count:
        messages.success(
            request,
            f'Import termine: {created_count} creation(s), {updated_count} mise(s) a jour.'
        )
    if errors:
        preview = '; '.join(errors[:5])
        more = '' if len(errors) <= 5 else f' (+{len(errors) - 5} autres erreurs)'
        messages.error(request, f'{len(errors)} erreur(s): {preview}{more}')

    return redirect('organization:ouvrage_grading_annuel_list')
