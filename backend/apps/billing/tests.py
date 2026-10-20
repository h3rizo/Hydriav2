import json
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from datetime import date, datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.db import IntegrityError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.billing.models import (
    ChangementCompteurTete,
    ChangementCompteurAbonne,
    CompteurAbonne,
    CompteurTete,
    ConsoCarteSociale,
    CoupureEau,
    DepenseService,
    Fournisseur,
    ModaliteGestion,
    PaiementForfaitaire,
    PointEau,
    Releve,
    ReleveCompteurTete,
    SegmentConsommationReleve,
    TypePointEau,
)
from apps.infrastructure.models import CarteSociale
from apps.billing.views_facturation import _get_allowed_ouvrage_ids as get_facturation_allowed_ouvrage_ids
from apps.employees.models import Employe, EmployeOuvrages, Poste
from apps.employees.models_evaluation import FichePoste
from apps.organization.models import Base, Organisation, Ouvrage, TypeOuvrage
from apps.planning.models import Designation, Planning, PlanningDetail, TypeActivite, TypeDesignation


class HistoriqueCoupuresEauTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='agent_historique_coupures',
            password='testpass123',
        )
        self.organisation = Organisation.objects.create(nom_organisation='Gestionnaire historique')
        self.base = Base.objects.create(nom_base='Base historique', organisation=self.organisation)
        self.poste = Poste.objects.create(
            nom_poste='Agent de recouvrement',
            organisation=self.organisation,
        )
        self.employe = Employe.objects.create(
            user=self.user,
            nom_employe='Agent Historique',
            nom='Historique',
            base=self.base,
            poste=self.poste,
        )
        type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='AEPG historique')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage historique',
            type_ouvrage=type_ouvrage,
        )
        EmployeOuvrages.objects.create(
            employe=self.employe,
            ouvrage=self.ouvrage,
            date_debut=date(2025, 1, 1),
        )
        ModaliteGestion.objects.create(
            ouvrage=self.ouvrage,
            gestionnaire=self.organisation,
            type_gestion='privee',
            modalitepaiement='au_volume',
            date_debut=date(2025, 1, 1),
        )
        self.point_remis = PointEau.objects.create(
            code_point_eau='PE-REMIS',
            nomabonne='Abonne remis',
            ouvrage=self.ouvrage,
        )
        self.point_coupe = PointEau.objects.create(
            code_point_eau='PE-COUPE',
            nomabonne='Abonne coupe',
            ouvrage=self.ouvrage,
        )
        bp_type = TypePointEau.objects.create(nom_type='BP')
        other_type = TypePointEau.objects.create(nom_type='Borne fontaine')
        self.point_coupe.type_point_eau = bp_type
        self.point_coupe.save(update_fields=['type_point_eau'])
        self.point_remis.type_point_eau = other_type
        self.point_remis.save(update_fields=['type_point_eau'])
        CoupureEau.objects.create(
            point_eau=self.point_remis,
            type_evenement='coupure',
            date_evenement=date(2026, 1, 2),
            cree_par=self.user,
        )
        CoupureEau.objects.create(
            point_eau=self.point_remis,
            type_evenement='remise',
            date_evenement=date(2026, 1, 3),
            cree_par=self.user,
        )
        CoupureEau.objects.create(
            point_eau=self.point_coupe,
            type_evenement='coupure',
            date_evenement=date(2026, 1, 4),
            cree_par=self.user,
        )
        self.client.force_login(self.user)

    def test_dashboard_defaults_to_connected_organisation_and_computes_states(self):
        response = self.client.get(reverse('billing:historique_coupures_eau'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['filters']['organisation'], str(self.organisation.id))
        self.assertEqual(response.context['stats']['total_evenements'], 3)
        self.assertEqual(response.context['stats']['points_coupes_historique'], 2)
        self.assertEqual(response.context['stats']['points_coupes_historique_pct'], 100.0)
        self.assertEqual(response.context['stats']['encore_coupes'], 1)
        self.assertEqual(response.context['stats']['encore_coupes_pct'], 50.0)
        self.assertEqual(response.context['stats']['remis'], 1)
        self.assertEqual(response.context['stats']['remis_pct'], 50.0)
        self.assertEqual(response.context['stats']['total_points'], 2)
        self.assertEqual(response.context['stats_bp']['total_points'], 1)
        self.assertEqual(response.context['stats_bp']['encore_coupes'], 1)
        self.assertEqual(response.context['stats_bp']['encore_coupes_pct'], 100.0)
        self.assertEqual(response.context['stats_bp']['remis'], 0)
        self.assertEqual(response.context['stats_autres']['total_points'], 1)
        self.assertEqual(response.context['stats_autres']['encore_coupes'], 0)
        self.assertEqual(response.context['stats_autres']['remis'], 1)
        self.assertEqual(response.context['stats_autres']['remis_pct'], 100.0)

    def test_current_cut_filter_keeps_history_of_currently_cut_points(self):
        response = self.client.get(
            reverse('billing:historique_coupures_eau'),
            {'vue': 'encore_coupes'},
        )

        rows = list(response.context['page_obj'].object_list)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['point'].id, self.point_coupe.id)
        self.assertTrue(rows[0]['current_cut'])

    def test_percentages_use_the_filtered_water_point_population(self):
        response = self.client.get(
            reverse('billing:historique_coupures_eau'),
            {'search': 'PE-COUPE'},
        )

        stats = response.context['stats']
        self.assertEqual(stats['total_points'], 1)
        self.assertEqual(stats['points_coupes_historique_pct'], 100.0)
        self.assertEqual(stats['encore_coupes_pct'], 100.0)
        self.assertEqual(stats['remis_pct'], 0.0)

    def test_water_point_category_filter_separates_bp_from_other_types(self):
        response = self.client.get(
            reverse('billing:historique_coupures_eau'),
            {'categorie_point_eau': 'bp'},
        )

        rows = list(response.context['page_obj'].object_list)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['point'].id, self.point_coupe.id)
        self.assertEqual(response.context['stats_bp']['total_points'], 1)
        self.assertEqual(response.context['stats_autres']['total_points'], 0)
        self.assertEqual(response.context['filters']['categorie_point_eau'], 'bp')


class ConsolidatedPointsEauEnhancementsTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser(
            username='admin_points_eau_consolides',
            email='points@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)
        self.organisation = Organisation.objects.create(nom_organisation='Organisation PE')
        self.base = Base.objects.create(nom_base='Base PE', organisation=self.organisation)
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='AEPG PE')
        self.ouvrage_a = Ouvrage.objects.create(nom_ouvrage='Ouvrage A', type_ouvrage=self.type_ouvrage)
        self.ouvrage_b = Ouvrage.objects.create(nom_ouvrage='Ouvrage B', type_ouvrage=self.type_ouvrage)
        self.employe = Employe.objects.create(
            nom_employe='Agent Actif',
            nom='Actif',
            prenom='Agent',
            surnom='Agent A',
            base=self.base,
        )
        self.ancien_employe = Employe.objects.create(
            nom_employe='Agent Debauche',
            nom='Debauche',
            base=self.base,
            date_debauchage=date(2026, 1, 1),
        )
        EmployeOuvrages.objects.create(
            employe=self.employe,
            ouvrage=self.ouvrage_a,
            date_debut=date(2025, 1, 1),
        )
        EmployeOuvrages.objects.create(
            employe=self.ancien_employe,
            ouvrage=self.ouvrage_b,
            date_debut=date(2025, 1, 1),
        )
        self.point_a = PointEau.objects.create(
            nomabonne='Abonne A',
            ouvrage=self.ouvrage_a,
            latitude=None,
            longitude=None,
        )
        self.point_b = PointEau.objects.create(
            nomabonne='Abonne B',
            ouvrage=self.ouvrage_b,
            latitude=-18.1,
            longitude=47.2,
        )

    def test_initial_page_waits_for_explicit_load(self):
        response = self.client.get(reverse('billing:consolidated_points_eau'))

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['results_loaded'])
        self.assertEqual(list(response.context['page_obj'].object_list), [])
        self.assertContains(response, "Charger tous les points d'eau")

    def test_load_parameter_displays_results(self):
        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'load': '1'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['results_loaded'])
        self.assertEqual(
            set(response.context['page_obj'].object_list),
            {self.point_a, self.point_b},
        )

    def test_modal_api_updates_subscriber_name(self):
        response = self.client.post(
            reverse('billing:update_point_eau_api', args=[self.point_a.id]),
            {'nomabonne': '  Nouveau nom  '},
        )

        self.assertEqual(response.status_code, 200)
        self.point_a.refresh_from_db()
        self.assertEqual(self.point_a.nomabonne, 'Nouveau nom')

    def test_details_api_returns_unlocalized_flat_rate(self):
        self.point_a.tarifforfaitaire = Decimal('75000.00')
        self.point_a.save(update_fields=['tarifforfaitaire'])

        response = self.client.get(
            reverse('billing:point_eau_details_api', args=[self.point_a.id]),
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['tarifforfaitaire'], '75000.00')

    def test_details_api_returns_gps_and_photos_for_edit_modal(self):
        self.point_a.latitude = -18.8792
        self.point_a.longitude = 47.5079
        self.point_a.photos = [{
            'filename': 'point.jpg',
            'content_type': 'image/jpeg',
            'data': 'ZmFrZQ==',
        }]
        self.point_a.save(update_fields=['latitude', 'longitude', 'photos'])

        response = self.client.get(
            reverse('billing:point_eau_details_api', args=[self.point_a.id]),
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload['latitude'], -18.8792)
        self.assertEqual(payload['longitude'], 47.5079)
        self.assertEqual(payload['photos'][0]['filename'], 'point.jpg')

    def test_point_type_catalog_uses_real_crud_view_and_lists_types(self):
        point_type = TypePointEau.objects.create(
            nom_type='Borne fontaine catalogue',
            description='Type visible dans le catalogue',
        )

        response = self.client.get(reverse('billing:typepointeau_list'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(point_type, response.context['types'])
        self.assertContains(response, 'Borne fontaine catalogue')
        self.assertContains(response, 'fa-faucet')

    def test_employee_filter_uses_active_responsible_ouvrage(self):
        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'employe': str(self.employe.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context['page_obj'].object_list), [self.point_a])
        self.assertIn(self.employe, response.context['employes'])
        self.assertNotIn(self.ancien_employe, response.context['employes'])

    def test_flat_rate_filter_keeps_matching_points(self):
        self.point_a.tarifforfaitaire = 75000
        self.point_a.save(update_fields=['tarifforfaitaire'])
        self.point_b.tarifforfaitaire = 50000
        self.point_b.save(update_fields=['tarifforfaitaire'])

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'tarifforfaitaire': '75000.0'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context['page_obj'].object_list), [self.point_a])
        self.assertContains(response, 'Tarif forfaitaire')

    def test_point_type_filter_keeps_matching_points(self):
        type_bf = TypePointEau.objects.create(nom_type='BF filtre')
        type_pmh = TypePointEau.objects.create(nom_type='PMH filtre')
        self.point_a.type_point_eau = type_bf
        self.point_a.save(update_fields=['type_point_eau'])
        self.point_b.type_point_eau = type_pmh
        self.point_b.save(update_fields=['type_point_eau'])

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'type_point_eau': str(type_bf.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context['page_obj'].object_list), [self.point_a])

    def test_connection_date_range_uses_madagascar_days_inclusive(self):
        import pytz
        from datetime import datetime as dt
        tz_mg = pytz.timezone('Indian/Antananarivo')
        # 00:30 à Madagascar le 31/03 = 30/03 21:30 UTC : doit compter pour le 31/03
        self.point_a.datebranchement = tz_mg.localize(dt(2026, 3, 31, 0, 30))
        self.point_a.save(update_fields=['datebranchement'])
        self.point_b.datebranchement = tz_mg.localize(dt(2026, 4, 1, 0, 30))
        self.point_b.save(update_fields=['datebranchement'])

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'branchement_du': '2026-03-31', 'branchement_au': '2026-03-31'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['results_loaded'])
        self.assertEqual(list(response.context['page_obj'].object_list), [self.point_a])
        self.assertContains(response, 'value="2026-03-31"')

    def test_repartition_counts_filtered_points_by_commune_and_type(self):
        type_bf = TypePointEau.objects.create(nom_type='BF repartition')
        self.point_a.type_point_eau = type_bf
        self.point_a.datebranchement = timezone.now()
        self.point_a.save(update_fields=['type_point_eau', 'datebranchement'])
        PointEau.objects.create(nomabonne='Abonne C', ouvrage=self.ouvrage_a, type_point_eau=type_bf,
                                datebranchement=timezone.now())
        # Sans date de branchement : exclu de la répartition
        PointEau.objects.create(nomabonne='Abonne sans date', ouvrage=self.ouvrage_a, type_point_eau=type_bf)

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'ouvrage': str(self.ouvrage_a.id)},
        )

        self.assertEqual(response.status_code, 200)
        repartition = response.context['repartition_branchements']
        self.assertEqual(len(repartition), 1)
        self.assertEqual(repartition[0]['type'], 'BF repartition')
        self.assertEqual(repartition[0]['commune'], 'Commune non renseignée')
        self.assertEqual(repartition[0]['n'], 2)
        self.assertEqual(repartition[0]['ouvrage'], 'Ouvrage A')
        self.assertEqual(repartition[0]['ouvrage_id'], str(self.ouvrage_a.id))
        self.assertContains(response, 'repartition-branchements-data')

    def test_repartition_excel_export_follows_filters_and_selected_types(self):
        from io import BytesIO
        from openpyxl import load_workbook
        type_bf = TypePointEau.objects.create(nom_type='BF export')
        type_bp = TypePointEau.objects.create(nom_type='BP export')
        now = timezone.now()
        self.point_a.type_point_eau = type_bf
        self.point_a.datebranchement = now
        self.point_a.save(update_fields=['type_point_eau', 'datebranchement'])
        self.point_b.type_point_eau = type_bf
        self.point_b.datebranchement = now
        self.point_b.save(update_fields=['type_point_eau', 'datebranchement'])
        PointEau.objects.create(nomabonne='Abonne C', ouvrage=self.ouvrage_b, type_point_eau=type_bf, datebranchement=now)
        PointEau.objects.create(nomabonne='Abonne D', ouvrage=self.ouvrage_a, type_point_eau=type_bp, datebranchement=now)

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'load': '1', 'export': 'repartition', 'types': ['BF export']},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        ws = load_workbook(BytesIO(response.content)).active
        header = [c.value for c in ws[5]]
        self.assertEqual(header, ['Région / District / Commune / Ouvrage', 'BF export', 'Total'])
        labels = [ws.cell(row=r, column=1).value for r in range(6, ws.max_row + 1)]
        totals = [ws.cell(row=r, column=3).value for r in range(6, ws.max_row + 1)]
        # Ouvrage B (2 BF) avant Ouvrage A (1 BF, son BP est exclu)
        self.assertEqual(labels[-3:], ['Ouvrage B', 'Ouvrage A', 'Total'])
        self.assertEqual(totals[-3:], [2, 1, 3])
        self.assertIn('Types exclus : BP export', ws.cell(row=3, column=1).value)

    def test_stats_count_only_active_points_and_set_apart_cut_and_undated(self):
        now = timezone.now()
        type_bf = TypePointEau.objects.create(nom_type='BF statut')
        self.point_a.datebranchement = now
        self.point_a.type_point_eau = type_bf
        self.point_a.save(update_fields=['datebranchement', 'type_point_eau'])
        # point_b : sans date de branchement ; point coupé : date d'arrêt
        self.point_b.datearretbranchement = now  # sans date + arrêt -> reste "sans date"
        self.point_b.save(update_fields=['datearretbranchement'])
        coupe = PointEau.objects.create(nomabonne='Coupé', ouvrage=self.ouvrage_a, type_point_eau=type_bf,
                                        datebranchement=now, datearretbranchement=now)

        response = self.client.get(reverse('billing:consolidated_points_eau'), {'load': '1'})

        self.assertEqual(response.context['points_eau_actifs'], 1)
        self.assertEqual(response.context['points_eau_coupes'], 1)
        self.assertEqual(response.context['points_eau_sans_date'], 1)
        self.assertEqual([t['count'] for t in response.context['types_stats']], [1])
        rows = list(response.context['page_obj'].object_list)
        self.assertEqual(rows[-1], self.point_b)  # sans date mis en fin de liste
        statuts = {r['statut'] for r in response.context['repartition_branchements']}
        self.assertEqual(statuts, {'actif', 'coupe'})

        response = self.client.get(reverse('billing:consolidated_points_eau'), {'statut': 'coupe'})
        self.assertEqual(list(response.context['page_obj'].object_list), [coupe])

    def test_point_coupe_depuis_plus_d_un_an_est_abandonne_et_non_coupe(self):
        from datetime import timedelta
        from .points_eau_statut import (
            STATUT_ABANDONNE, STATUT_COUPE, compter_statuts, statut_of,
        )
        now = timezone.now()
        recent = PointEau.objects.create(nomabonne='Coupé récent', ouvrage=self.ouvrage_a,
                                         datebranchement=now, datearretbranchement=now - timedelta(days=100))
        limite = PointEau.objects.create(nomabonne='Coupé 364 j', ouvrage=self.ouvrage_a,
                                         datebranchement=now, datearretbranchement=now - timedelta(days=364))
        ancien = PointEau.objects.create(nomabonne='Coupé 2 ans', ouvrage=self.ouvrage_a,
                                         datebranchement=now, datearretbranchement=now - timedelta(days=730))

        self.assertEqual(statut_of(recent), STATUT_COUPE)
        self.assertEqual(statut_of(limite), STATUT_COUPE)
        self.assertEqual(statut_of(ancien), STATUT_ABANDONNE)

        qs = PointEau.objects.filter(id__in=[recent.id, limite.id, ancien.id])
        counts = compter_statuts(qs)
        self.assertEqual((counts['coupes'], counts['abandonnes'], counts['actifs']), (2, 1, 0))

        response = self.client.get(reverse('billing:consolidated_points_eau'),
                                   {'load': '1', 'statut': 'abandonne'})
        self.assertEqual(response.context['points_eau_abandonnes'], 1)
        self.assertEqual(list(response.context['page_obj'].object_list), [ancien])

    def test_api_type_point_eau_regles_kiosque_et_cloture_preservee(self):
        kiosque = TypePointEau.objects.create(nom_type='Kiosque')
        bf = TypePointEau.objects.create(nom_type='BF api type')
        self.point_a.besoin_en_cloture = True
        self.point_a.tarifforfaitaire = 5000
        self.point_a.save(update_fields=['besoin_en_cloture', 'tarifforfaitaire'])
        url = reverse('billing:update_point_eau_type_api', args=[self.point_a.id])

        # Kiosque sans type d'abonné : refusé
        response = self.client.post(url, {'type_point_eau': kiosque.id, 'typeabonne': ''})
        self.assertEqual(response.status_code, 400)

        # Kiosque + boutique : accepté, tarif vidé (non applicable), clôture conservée
        response = self.client.post(url, {'type_point_eau': kiosque.id, 'typeabonne': 'boutique'})
        self.assertEqual(response.status_code, 200)
        self.point_a.refresh_from_db()
        self.assertEqual((self.point_a.type_point_eau_id, self.point_a.typeabonne), (kiosque.id, 'boutique'))
        self.assertIsNone(self.point_a.tarifforfaitaire)
        self.assertTrue(self.point_a.besoin_en_cloture)

        # Autre type : le type d'abonné est vidé
        response = self.client.post(url, {'type_point_eau': bf.id, 'typeabonne': 'boutique'})
        self.assertEqual(response.status_code, 200)
        self.point_a.refresh_from_db()
        self.assertEqual((self.point_a.type_point_eau_id, self.point_a.typeabonne), (bf.id, None))

    def _points_statuts_mixtes(self):
        # Les pages /organisations/ exigent un employé lié (@employe_required)
        self.employe.user = self.user
        self.employe.save(update_fields=['user'])
        now = timezone.now()
        self.point_a.datebranchement = now
        self.point_a.save(update_fields=['datebranchement'])
        # point_b reste sans date de branchement
        PointEau.objects.create(nomabonne='Coupé A', ouvrage=self.ouvrage_a,
                                datebranchement=now, datearretbranchement=now)
        PointEau.objects.create(nomabonne='Sans date A', ouvrage=self.ouvrage_a)

    def test_ouvrage_points_page_sets_apart_cut_and_undated_points(self):
        self._points_statuts_mixtes()

        response = self.client.get(
            reverse('organization:points_eau_list_ouvrage', args=[self.ouvrage_a.id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['points_eau_actifs_count'], 1)
        self.assertEqual(response.context['points_eau_coupes_count'], 1)
        self.assertEqual(response.context['points_eau_sans_date_count'], 1)
        noms = [p.nomabonne for p in response.context['points_eau']]
        self.assertEqual(noms, ['Abonne A', 'Coupé A', 'Sans date A'])
        self.assertContains(response, 'non comptés comme points d')

    def test_ouvrage_list_counts_only_active_points(self):
        self._points_statuts_mixtes()

        response = self.client.get(reverse('organization:ouvrage_list'), {'hide_inexistants': '0'})

        self.assertEqual(response.status_code, 200)
        ouvrage = next(o for o in response.context['page_obj'].object_list if o.pk == self.ouvrage_a.pk)
        self.assertEqual(ouvrage.nb_points_eau, 1)
        self.assertEqual(ouvrage.nb_points_eau_coupes, 1)
        self.assertEqual(ouvrage.nb_points_eau_sans_date, 1)

    def test_ouvrage_list_filters_non_fonctionnels(self):
        self.ouvrage_b.marque_non_fonctionnel = True
        self.ouvrage_b.save(update_fields=['marque_non_fonctionnel'])
        url = reverse('organization:ouvrage_list')

        response = self.client.get(url, {
            'hide_inexistants': '0',
            'fonctionnalite': 'non_fonctionnel',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['fonctionnalite_filter'], 'non_fonctionnel')
        self.assertEqual(
            {ouvrage.id for ouvrage in response.context['page_obj'].object_list},
            {self.ouvrage_b.id},
        )

        response = self.client.get(url, {
            'hide_inexistants': '0',
            'fonctionnalite': 'fonctionnel',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            {ouvrage.id for ouvrage in response.context['page_obj'].object_list},
            {self.ouvrage_a.id},
        )

    def test_widget_stats_count_active_points_and_expose_cut_and_undated(self):
        from unittest import mock
        self._points_statuts_mixtes()

        # Le périmètre (organisation / modalités) n'est pas l'objet du test : on le fixe
        with mock.patch(
            'apps.core.views_widget._resolve_scoped_ouvrage_ids',
            return_value=[self.ouvrage_a.id, self.ouvrage_b.id],
        ):
            response = self.client.get(reverse('core:api_widget_stats'))

        self.assertEqual(response.status_code, 200)
        stats = response.json()['stats']
        # Actif : Abonne A ; coupé : Coupé A ; sans date : Abonne B + Sans date A
        self.assertEqual(
            (stats['points_eau'], stats['points_eau_coupes'], stats['points_eau_sans_date']),
            (1, 1, 2),
        )
        self.assertEqual(
            stats['points_eau_forfaitaires'] + stats['points_eau_volumetriques'],
            stats['points_eau'],
        )

    def test_filtre_plusieurs_communes_et_organisation_derniere_modalite_active(self):
        from apps.geography.models import Commune, District, Fokontany, Parcelle, Region, Village
        from .models import ModaliteGestion
        region = Region.objects.create(nom_region='Analanjirofo test')
        district = District.objects.create(nom_district='District test', region=region)
        points_par_commune = {}
        for nom in ('Commune 1', 'Commune 2', 'Commune 3'):
            commune = Commune.objects.create(nom_commune=nom, district=district)
            fokontany = Fokontany.objects.create(nom_fokontany=f'Fkt {nom}', commune=commune)
            village = Village.objects.create(nom_village=f'Village {nom}', fokontany=fokontany)
            parcelle = Parcelle.objects.create(reference_parcelle=f'P {nom}', village=village)
            points_par_commune[nom] = (commune, PointEau.objects.create(
                nomabonne=f'Abonne {nom}', ouvrage=self.ouvrage_a, parcelle=parcelle,
            ))

        url = reverse('billing:consolidated_points_eau')
        c1, p1 = points_par_commune['Commune 1']
        c2, p2 = points_par_commune['Commune 2']
        # Deux communes (+ la valeur technique __multi__ envoyée par la liste)
        response = self.client.get(url, {'commune': ['__multi__', str(c1.id), str(c2.id)]})
        self.assertEqual(set(response.context['page_obj'].object_list), {p1, p2})
        self.assertEqual(response.context['commune_filter'], 'multi')
        self.assertContains(response, '2 communes (filtre enregistré)')

        # Organisation : gestionnaire de la DERNIÈRE modalité active de l'ouvrage
        autre_org = Organisation.objects.create(nom_organisation='Autre gestionnaire')
        ModaliteGestion.objects.create(ouvrage=self.ouvrage_a, type_gestion='privee', gestionnaire=autre_org,
                                       date_debut=date(2020, 1, 1))
        ModaliteGestion.objects.create(ouvrage=self.ouvrage_a, type_gestion='privee', gestionnaire=self.organisation,
                                       date_debut=date(2024, 1, 1))
        ModaliteGestion.objects.create(ouvrage=self.ouvrage_b, type_gestion='privee', gestionnaire=autre_org,
                                       date_debut=date(2023, 1, 1), date_fin=date(2024, 1, 1))
        response = self.client.get(url, {'organisation': str(self.organisation.id)})
        self.assertEqual(
            {p.ouvrage_id for p in response.context['page_obj'].object_list}, {self.ouvrage_a.id}
        )
        response = self.client.get(url, {'organisation': str(autre_org.id)})
        self.assertEqual(list(response.context['page_obj'].object_list), [])
        self.assertIn(self.organisation, response.context['organisations'])

    def test_filtres_enregistres_de_la_page(self):
        from apps.core.models import FiltreSauvegarde
        FiltreSauvegarde.objects.create(page_key='points_eau_consolides', nom='Mon filtre',
                                        parametres={'commune': ['x']})
        FiltreSauvegarde.objects.create(page_key='villages_desserte', nom='Autre page', parametres={})

        response = self.client.get(reverse('billing:consolidated_points_eau'))

        noms = [f['nom'] for f in response.context['filtres_sauvegardes_data']]
        self.assertIn('Mon filtre', noms)
        self.assertNotIn('Autre page', noms)
        self.assertContains(response, 'Enregistrer le filtre actuel')

    def test_sans_date_option_lists_points_without_connection_date(self):
        self.point_a.datebranchement = timezone.now()
        self.point_a.save(update_fields=['datebranchement'])

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'annee': 'sans_date'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(list(response.context['page_obj'].object_list), [self.point_b])
        self.assertContains(response, 'value="sans_date" selected')

    def test_invalid_connection_date_is_ignored(self):
        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'branchement_du': 'pas-une-date'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['branchement_du'], '')
        self.assertEqual(
            set(response.context['page_obj'].object_list),
            {self.point_a, self.point_b},
        )

    def test_non_applicable_point_type_clears_flat_rate(self):
        type_bp = TypePointEau.objects.create(nom_type='BP')
        self.point_a.type_point_eau = type_bp
        self.point_a.tarifforfaitaire = 75000
        self.point_a.save(update_fields=['type_point_eau', 'tarifforfaitaire'])

        response = self.client.post(
            reverse('billing:update_point_eau_api', args=[self.point_a.id]),
            {'tarifforfaitaire': '80000'},
        )

        self.assertEqual(response.status_code, 200)
        self.point_a.refresh_from_db()
        self.assertIsNone(self.point_a.tarifforfaitaire)

    def test_print_defaults_to_missing_gps_or_subscriber_name(self):
        point_without_name = PointEau.objects.create(
            nomabonne='',
            ouvrage=self.ouvrage_a,
            latitude=-18.2,
            longitude=47.3,
        )
        point_with_numeric_name = PointEau.objects.create(
            nomabonne='  12345  ',
            ouvrage=self.ouvrage_a,
            latitude=-18.3,
            longitude=47.4,
        )

        response = self.client.get(
            reverse('billing:consolidated_points_eau'),
            {'print': '1', 'missing': 'all', 'employe': str(self.employe.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'billing/consolidated_points_eau_print.html')
        printed_ids = {point.id for point in response.context['points_eau']}
        self.assertEqual(
            printed_ids,
            {self.point_a.id, point_without_name.id, point_with_numeric_name.id},
        )
        self.assertNotIn(self.point_b.id, printed_ids)
        self.assertEqual(response.context['employe_label'], 'Agent A')
        self.assertEqual(response.context['commune_label'], 'Toutes les communes')
        self.assertEqual(response.context['ouvrages_count'], 1)
        self.assertEqual(response.context['points_eau_count'], 3)


class ReconcilierCoupuresSansHistoriqueTests(TestCase):
    def setUp(self):
        import pytz
        from datetime import datetime as dt
        self.tz = pytz.timezone('Indian/Antananarivo')
        self.mg = lambda y, m, d: self.tz.localize(dt(y, m, d))
        type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='AEPG coupures')
        self.ouvrage = Ouvrage.objects.create(nom_ouvrage='Ouvrage coupures', type_ouvrage=type_ouvrage)

    def _point(self, nom, arret=None, tarif=None):
        return PointEau.objects.create(
            nomabonne=nom, ouvrage=self.ouvrage, datebranchement=self.mg(2020, 1, 1),
            datearretbranchement=arret, tarifforfaitaire=tarif,
        )

    def _releve(self, point, jour, conso):
        return Releve.objects.create(
            point_eau=point, datereleve=jour, ancienindex=0, nouvelindex=conso, consommation=conso,
        )

    def test_reconcilie_selon_historique_et_activite(self):
        from io import StringIO
        from django.core.management import call_command
        from .models import CoupureEau

        actif = self._point('Actif en realite', arret=self.mg(2024, 12, 1))
        self._releve(actif, date(2026, 5, 20), 12)
        # Consomme après l'arrêt, puis n'est plus relevé alors que l'ouvrage l'est
        coupe_plus_tard = self._point('Coupe plus tard', arret=self.mg(2024, 12, 1))
        self._releve(coupe_plus_tard, date(2025, 3, 10), 8)
        self._releve(self._point('Voisin releve'), date(2026, 6, 1), 5)
        cloture = self._point('Releve de cloture', arret=self.mg(2024, 12, 1), tarif=Decimal('5000'))
        self._releve(cloture, date(2024, 12, 20), 3)
        confirme = self._point('Confirme', arret=self.mg(2025, 3, 1))
        CoupureEau.objects.bulk_create([CoupureEau(point_eau=confirme, type_evenement='coupure',
                                                   date_evenement=date(2025, 3, 1))])
        sans_arret = self._point('Coupure non reportee')
        CoupureEau.objects.bulk_create([CoupureEau(point_eau=sans_arret, type_evenement='coupure',
                                                   date_evenement=date(2025, 8, 12))])

        out = StringIO()
        call_command('reconcilier_coupures_sans_historique', stdout=out)  # dry-run
        actif.refresh_from_db()
        self.assertIsNotNone(actif.datearretbranchement)
        self.assertIn('Dry-run', out.getvalue())

        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            call_command('reconcilier_coupures_sans_historique', apply=True, backup_dir=tmp, stdout=StringIO())

        for p in (actif, cloture, confirme, sans_arret, coupe_plus_tard):
            p.refresh_from_db()
        self.assertIsNone(actif.datearretbranchement)
        # Coupé plus tard : date d'arrêt reportée au dernier relevé + événement à cette date
        self.assertEqual(coupe_plus_tard.datearretbranchement.astimezone(self.tz).date(), date(2025, 3, 10))
        self.assertEqual(
            list(CoupureEau.objects.filter(point_eau=coupe_plus_tard).values_list('date_evenement', flat=True)),
            [date(2025, 3, 10)],
        )
        # Relevé de clôture (< 1 mois) : reste coupé, événement reconstitué, tarif conservé
        self.assertIsNotNone(cloture.datearretbranchement)
        evt = CoupureEau.objects.get(point_eau=cloture)
        self.assertEqual((evt.type_evenement, evt.date_evenement, evt.raison), ('coupure', date(2024, 12, 1), 'autre'))
        self.assertEqual(cloture.tarifforfaitaire, Decimal('5000'))
        # Confirmé par l'historique : aucun nouvel événement
        self.assertEqual(CoupureEau.objects.filter(point_eau=confirme).count(), 1)
        self.assertEqual(
            sans_arret.datearretbranchement.astimezone(self.tz).date(), date(2025, 8, 12)
        )


class ReleveCarteSocialeDateSyncTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='directeur_releve_cs',
            password='testpass123',
        )
        directeur_group, _ = Group.objects.get_or_create(name='Directeur')
        self.user.groups.add(directeur_group)
        self.client.force_login(self.user)

        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='AEPG CS')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Carte Sociale',
            type_ouvrage=self.type_ouvrage,
        )
        self.type_point_eau = TypePointEau.objects.create(nom_type='Kiosque CS')
        self.point_eau = PointEau.objects.create(
            nomabonne='Kiosque social',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_point_eau,
            code_point_eau='CS-001',
        )
        self.carte = CarteSociale.objects.create(
            numero_carte='CARTE-001',
            nom_beneficiaire='Beneficiaire Test',
            ouvrage=self.ouvrage,
            point_eau=self.point_eau,
        )
        self.releve = Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 7, 1),
            ancienindex=0,
            nouvelindex=10,
            consommation=10,
            volume_carte_sociale=0.2,
            volume_non_facture=0,
            volume_facture=9.8,
            statut='impaye',
        )
        self.conso = ConsoCarteSociale.objects.create(
            releve=self.releve,
            carte_sociale=self.carte,
            quantite_bidons=10,
        )

    def test_conso_creation_uses_linked_releve_date(self):
        self.assertEqual(self.conso.date, date(2026, 7, 1))

    def test_modifying_releve_date_updates_linked_social_consumption_date(self):
        response = self.client.post(
            reverse('billing:api_modifier_releve'),
            data=json.dumps({
                'releve_id': str(self.releve.id),
                'nouvel_index': 10,
                'date_releve': '2026-06-30',
                'volume_non_facture': 0,
                'remarque': '',
                'cartes_sociales': [{
                    'id': str(self.carte.id),
                    'quantite_bidons': 10,
                    'volume_total': 0.2,
                }],
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        self.releve.refresh_from_db()
        self.conso.refresh_from_db()
        self.assertEqual(self.releve.datereleve, date(2026, 6, 30))
        self.assertEqual(self.conso.date, date(2026, 6, 30))

    def test_management_page_moves_releve_and_all_its_linked_consumptions(self):
        second_card = CarteSociale.objects.create(
            numero_carte='CARTE-002',
            nom_beneficiaire='Deuxieme Beneficiaire',
            ouvrage=self.ouvrage,
            point_eau=self.point_eau,
        )
        second_conso = ConsoCarteSociale.objects.create(
            releve=self.releve,
            carte_sociale=second_card,
            quantite_bidons=5,
        )

        response = self.client.post(
            reverse('billing:gestion_consos_cartes_sociales'),
            data={
                'consommation_ids': [str(self.conso.id)],
                'date_cible': '2026-06-30',
                'confirmation': 'DEPLACER',
            },
        )

        self.assertEqual(
            response.status_code,
            302,
            [str(message) for message in response.context['messages']] if response.context else None,
        )
        self.releve.refresh_from_db()
        self.conso.refresh_from_db()
        second_conso.refresh_from_db()
        self.assertEqual(self.releve.datereleve, date(2026, 6, 30))
        self.assertEqual(self.conso.date, date(2026, 6, 30))
        self.assertEqual(second_conso.date, date(2026, 6, 30))

    def test_management_page_blocks_diamex_releve(self):
        self.releve.numeroDiameX = 'DX-123'
        self.releve.save(update_fields=['numeroDiameX'])

        response = self.client.post(
            reverse('billing:gestion_consos_cartes_sociales'),
            data={
                'consommation_ids': [str(self.conso.id)],
                'date_cible': '2026-06-30',
                'confirmation': 'DEPLACER',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.releve.refresh_from_db()
        self.conso.refresh_from_db()
        self.assertEqual(self.releve.datereleve, date(2026, 7, 1))
        self.assertEqual(self.conso.date, date(2026, 7, 1))
        self.assertContains(response, 'DiameX')

    def test_management_page_groups_by_ouvrage_point_and_date_with_subscriber(self):
        recent_releve = Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 8, 1),
            ancienindex=10,
            nouvelindex=15,
            consommation=5,
            volume_carte_sociale=0.1,
            volume_non_facture=0,
            volume_facture=4.9,
            statut='impaye',
        )
        ConsoCarteSociale.objects.create(
            releve=recent_releve,
            carte_sociale=self.carte,
            quantite_bidons=5,
        )

        response = self.client.get(reverse('billing:gestion_consos_cartes_sociales'))

        self.assertEqual(response.status_code, 200)
        groups = response.context['grouped_ouvrages']
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]['ouvrage'], self.ouvrage)
        self.assertEqual(groups[0]['points'][0]['point_eau'], self.point_eau)
        self.assertEqual(groups[0]['points'][0]['nom_abonne'], 'Kiosque social')
        dates = groups[0]['points'][0]['dates']
        self.assertEqual([group['date_releve'] for group in dates], [
            date(2026, 8, 1),
            date(2026, 7, 1),
        ])
        self.assertContains(response, 'Ouvrage Carte Sociale')
        self.assertContains(response, 'Abonne : Kiosque social')
        self.assertContains(response, 'Date de releve :')


class DemandesBPAutoBranchementTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser(
            username='admin_bp',
            email='adminbp@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)

        self.organisation = Organisation.objects.create(
            nom_organisation='Org Test',
            type_organisation=Organisation.TYPE_GESTIONNAIRE,
        )
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='Forage')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Test',
            type_ouvrage=self.type_ouvrage,
        )
        ModaliteGestion.objects.create(
            ouvrage=self.ouvrage,
            type_gestion='privee',
            gestionnaire=self.organisation,
            date_debut=date(2025, 1, 1),
            modalitepaiement='au_volume',
        )
        self.type_bp = TypePointEau.objects.create(nom_type='Branchement Prive')

    @staticmethod
    def _aware_datetime(year, month, day):
        return timezone.make_aware(datetime(year, month, day, 12, 0, 0))

    def test_api_list_auto_regularises_missing_datebranchement_from_first_releve(self):
        point_eau = PointEau.objects.create(
            nomabonne='Abonne Test',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_bp,
            code_point_eau='BP-001',
        )
        Releve.objects.create(
            point_eau=point_eau,
            abonne=point_eau.nomabonne,
            datereleve=date(2025, 6, 1),
            ancienindex=0,
            nouvelindex=12,
            consommation=12,
            statut='impaye',
        )
        Releve.objects.create(
            point_eau=point_eau,
            abonne=point_eau.nomabonne,
            datereleve=date(2025, 7, 1),
            ancienindex=12,
            nouvelindex=20,
            consommation=8,
            statut='impaye',
        )

        response = self.client.get(
            reverse('billing:api_demandes_bp_list'),
            {'statut': 'branche'},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])
        self.assertEqual(payload['normalisations_appliquees'], 1)
        self.assertEqual(len(payload['demandes']), 1)
        self.assertEqual(payload['demandes'][0]['etape'], 'complete')
        self.assertEqual(payload['demandes'][0]['date_branchement'], '01/06/2025')

        point_eau.refresh_from_db()
        self.assertIsNotNone(point_eau.datebranchement)
        self.assertEqual(point_eau.datebranchement.date(), date(2025, 6, 1))

    def test_manual_branchement_regularises_existing_releves_without_duplicate_initial_releve(self):
        point_eau = PointEau.objects.create(
            nomabonne='Abonne Releve',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_bp,
            code_point_eau='BP-002',
            datepaiementdemande=self._aware_datetime(2025, 5, 10),
            datepaiementdevis=self._aware_datetime(2025, 5, 20),
        )
        Releve.objects.create(
            point_eau=point_eau,
            abonne=point_eau.nomabonne,
            datereleve=date(2025, 5, 25),
            ancienindex=0,
            nouvelindex=4,
            consommation=4,
            statut='impaye',
        )

        response = self.client.post(
            reverse('billing:api_enregistrer_branchement', args=[point_eau.id]),
            data=json.dumps({
                'date_branchement': '2025-05-26',
                'index_initial': 0,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])
        self.assertTrue(payload['auto_regularise'])
        self.assertEqual(payload['date_branchement'], '25/05/2025')

        point_eau.refresh_from_db()
        self.assertEqual(point_eau.datebranchement.date(), date(2025, 5, 25))
        self.assertEqual(point_eau.releves.count(), 1)


class PaiementsForfaitairesPlanningValidationTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser(
            username='admin_forfaitaire',
            email='adminforfaitaire@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)

        self.organisation = Organisation.objects.create(
            nom_organisation='Org Forfaitaire',
            type_organisation=Organisation.TYPE_GESTIONNAIRE,
        )
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='Forage')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Forfaitaire',
            type_ouvrage=self.type_ouvrage,
        )
        ModaliteGestion.objects.create(
            ouvrage=self.ouvrage,
            type_gestion='privee',
            gestionnaire=self.organisation,
            date_debut=date(2025, 1, 1),
            modalitepaiement='forfaitaire',
        )
        self.type_point_eau = TypePointEau.objects.create(nom_type='BF')
        self.point_eau = PointEau.objects.create(
            nomabonne='Abonne Forfaitaire',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_point_eau,
            code_point_eau='PF-001',
            tarifforfaitaire=5000,
        )
        self.type_activite = TypeActivite.objects.create(
            nomtype='Collecte forfaitaire',
        )
        self.fiche_poste = FichePoste.objects.create(
            tache='Collecte des paiements forfaitaires',
            type_activites=self.type_activite,
        )

    @staticmethod
    def _aware_datetime(year, month, day):
        return timezone.make_aware(datetime(year, month, day, 8, 0, 0))

    def test_api_paiement_create_auto_validates_linked_planning(self):
        response = self.client.post(
            reverse('billing:api_paiement_create'),
            data=json.dumps({
                'point_eau_id': str(self.point_eau.id),
                'date_facture': '2025-04-01',
                'quantite': 1,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])

        paiement = PaiementForfaitaire.objects.get(id=payload['paiement_id'])
        self.assertIsNotNone(paiement.planning_id)
        paiement.planning.refresh_from_db()
        self.assertTrue(paiement.planning.estvalide)

    def test_api_paiement_update_validates_existing_invalid_planning(self):
        planning = Planning.objects.create(
            date=self._aware_datetime(2025, 4, 1),
            remarque='Planning forfaitaire test',
            estvalide=False,
            statut='en_cours',
            ouvrage=self.ouvrage,
            fiche_poste=self.fiche_poste,
            type_activite=self.type_activite,
        )
        paiement = PaiementForfaitaire.objects.create(
            point_eau=self.point_eau,
            planning=planning,
            date_facture=date(2025, 4, 1),
            quantite=Decimal('1'),
            prix_unitaire=Decimal('5000'),
            montant=Decimal('5000'),
            paye=False,
        )

        response = self.client.post(
            reverse('billing:api_paiement_update', args=[paiement.id]),
            data=json.dumps({
                'quantite': 2,
                'prix_unitaire': 5000,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])

        planning.refresh_from_db()
        self.assertTrue(planning.estvalide)


class RetourFraisDeplacementDefaultsTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='raf_retour_frais',
            email='raf@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)

        self.organisation = Organisation.objects.create(
            nom_organisation='Org Retour',
            type_organisation=Organisation.TYPE_GESTIONNAIRE,
        )
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='Forage')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Retour',
            type_ouvrage=self.type_ouvrage,
        )
        self.base_user = Base.objects.create(
            nom_base='Base A',
            organisation=self.organisation,
            ouvrage=self.ouvrage,
        )
        self.base_other = Base.objects.create(
            nom_base='Base B',
            organisation=self.organisation,
        )
        self.poste_raf = Poste.objects.create(
            nom_poste='Responsable administratif et financier',
            organisation=self.organisation,
        )
        self.employe_user = Employe.objects.create(
            user=self.user,
            nom_employe='Responsable Admin',
            nom='Responsable',
            prenom='Admin',
            surnom='RAF',
            poste=self.poste_raf,
            base=self.base_user,
        )
        self.employe_other = Employe.objects.create(
            nom_employe='Agent Base B',
            nom='Agent',
            prenom='BaseB',
            surnom='ABB',
            poste=self.poste_raf,
            base=self.base_other,
        )
        self.employe_no_frais = Employe.objects.create(
            nom_employe='Agent Sans Frais',
            nom='Sans',
            prenom='Frais',
            surnom='ASF',
            poste=self.poste_raf,
            base=self.base_user,
        )
        self.type_designation = TypeDesignation.objects.create(nom='Depenses')
        self.designation = Designation.objects.create(
            nom='Frais de déplacement',
            type_designation=self.type_designation,
        )

        planning_user = Planning.objects.create(
            date=timezone.now(),
            employe=self.employe_user,
            ouvrage=self.ouvrage,
        )
        planning_other = Planning.objects.create(
            date=timezone.now(),
            employe=self.employe_other,
            ouvrage=self.ouvrage,
        )
        PlanningDetail.objects.create(
            planning=planning_user,
            designation=self.designation,
            quantity=1,
            pu=Decimal('1000'),
            montant=Decimal('1000'),
            retour=False,
        )
        PlanningDetail.objects.create(
            planning=planning_other,
            designation=self.designation,
            quantity=1,
            pu=Decimal('2000'),
            montant=Decimal('2000'),
            retour=False,
        )

    def test_retour_frais_defaults_to_connected_user_base_even_with_multiple_visible_bases(self):
        response = self.client.get(reverse('billing:retour_frais_deplacement'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['default_base_id'], str(self.base_user.id))

        employes_dropdown_data = response.context['employes_dropdown_data']
        self.assertEqual(
            {item['id'] for item in employes_dropdown_data},
            {str(self.employe_user.id), str(self.employe_other.id), str(self.employe_no_frais.id)}
        )
        self.assertEqual({item['base_id'] for item in employes_dropdown_data}, {str(self.base_user.id), str(self.base_other.id)})


class DepensesServicesWorkflowPermissionsTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='agent_depense_service',
            email='agent.depense@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)
        self.fournisseur = Fournisseur.objects.create(
            nom_fournisseur='Prestataire Test',
        )

    def test_regular_authenticated_user_can_create_depense_service(self):
        response = self.client.post(
            reverse('billing:api_depenses_services_create'),
            data=json.dumps({
                'date_facture': '2026-04-01',
                'fournisseur_id': str(self.fournisseur.id),
                'designation': 'Facture de service test',
                'unite': 'Mois',
                'prix_unitaire': 2,
                'montant': 15000,
            }),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])

        depense = DepenseService.objects.get(id=payload['depense']['id'])
        self.assertEqual(depense.statut, DepenseService.STATUT_BROUILLON)


class FacturationWorkflowScopeTests(SimpleTestCase):
    @patch('apps.billing.views_facturation.apply_user_responsable_ouvrage_scope')
    @patch('apps.billing.views_facturation.get_user_accessible_gestionnaire_ids')
    @patch('apps.billing.views_facturation._is_directeur_user', return_value=False)
    @patch('apps.billing.views_facturation.ModaliteGestion')
    def test_allowed_ouvrages_are_limited_by_responsable_scope(
        self,
        modalite_gestion_mock,
        _is_directeur_mock,
        gestionnaire_ids_mock,
        responsable_scope_mock,
    ):
        user = SimpleNamespace(is_authenticated=True, is_superuser=False)
        gestionnaire_ids_mock.return_value = ['gest-1']

        modalites_qs = MagicMock()
        modalites_filtered_qs = MagicMock()
        modalites_excluded_qs = MagicMock()
        ouvrage_ids_qs = MagicMock()

        modalite_gestion_mock.objects.filter.return_value = modalites_qs
        modalites_qs.filter.return_value = modalites_filtered_qs
        modalites_filtered_qs.exclude.return_value = modalites_excluded_qs
        modalites_excluded_qs.values_list.return_value = ouvrage_ids_qs
        ouvrage_ids_qs.distinct.return_value = ['ouvrage-a', 'ouvrage-b']

        responsable_scope_mock.return_value = {'ouvrage-b'}

        scoped_ids = get_facturation_allowed_ouvrage_ids(user)

        self.assertEqual(scoped_ids, ['ouvrage-b'])
        responsable_scope_mock.assert_called_once_with({'ouvrage-a', 'ouvrage-b'}, user)


class RelevesCompteursTeteCrudTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser(
            username='admin_compteur_tete',
            email='admin.compteur@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)

        self.organisation = Organisation.objects.create(
            nom_organisation='Org Compteur',
            type_organisation=Organisation.TYPE_GESTIONNAIRE,
        )
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='AEPG')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Compteur',
            type_ouvrage=self.type_ouvrage,
        )
        ModaliteGestion.objects.create(
            ouvrage=self.ouvrage,
            type_gestion='privee',
            gestionnaire=self.organisation,
            date_debut=date(2026, 1, 1),
            modalitepaiement='au_volume',
        )
        self.compteur = CompteurTete.objects.create(
            ouvrage=self.ouvrage,
            numerocompteur='CT-001',
            index_initial=Decimal('0'),
        )
        self.type_point_eau = TypePointEau.objects.create(nom_type='BF')
        self.point_eau = PointEau.objects.create(
            nomabonne='Point Aval',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_point_eau,
            compteur_tete=self.compteur,
            code_point_eau='PE-001',
            datebranchement=timezone.now(),
        )
        Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 4, 15),
            ancienindex=0,
            nouvelindex=40,
            consommation=40,
            volume_facture=40,
            statut='impaye',
        )

    def test_create_requires_explanation_when_gap_exists(self):
        response = self.client.post(
            reverse('organization:api_compteurs_releves_create'),
            data={
                'compteur_id': str(self.compteur.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '50',
                'explication_ecart': '',
            },
        )

        self.assertEqual(response.status_code, 400)
        payload = response.json()
        self.assertFalse(payload['success'])
        self.assertIn('explication_ecart', payload['errors'])
        self.assertEqual(ReleveCompteurTete.objects.count(), 0)

    def test_create_persists_snapshot_with_downstream_comparison(self):
        response = self.client.post(
            reverse('organization:api_compteurs_releves_create'),
            data={
                'compteur_id': str(self.compteur.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '50',
                'explication_ecart': 'Perte reseau constatee pendant le mois.',
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])

        releve = ReleveCompteurTete.objects.get(id=payload['releve']['id'])
        self.assertEqual(releve.ancien_index, Decimal('0'))
        self.assertEqual(releve.nouvel_index, Decimal('50'))
        self.assertEqual(releve.volume_tete, Decimal('50'))
        self.assertEqual(releve.volume_aval, Decimal('40'))
        self.assertEqual(releve.ecart_volume, Decimal('10'))
        self.assertEqual(releve.nb_points_eau_aval, 1)

    def test_update_rejects_sequence_break_against_next_reading(self):
        releve_1 = ReleveCompteurTete.objects.create(
            compteur_tete=self.compteur,
            date_releve=date(2026, 4, 30),
            ancien_index=Decimal('0'),
            nouvel_index=Decimal('50'),
            volume_tete=Decimal('50'),
            volume_aval=Decimal('40'),
            ecart_volume=Decimal('10'),
            nb_points_eau_aval=1,
            explication_ecart='Ecart avril',
        )
        ReleveCompteurTete.objects.create(
            compteur_tete=self.compteur,
            date_releve=date(2026, 5, 31),
            ancien_index=Decimal('50'),
            nouvel_index=Decimal('70'),
            volume_tete=Decimal('20'),
            volume_aval=Decimal('0'),
            ecart_volume=Decimal('20'),
            nb_points_eau_aval=1,
            explication_ecart='Ecart mai',
        )

        response = self.client.post(
            reverse('organization:api_compteurs_releves_update', args=[releve_1.id]),
            data={
                'compteur_id': str(self.compteur.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '80',
                'explication_ecart': 'Correction invalide',
            },
        )

        self.assertEqual(response.status_code, 400)
        payload = response.json()
        self.assertFalse(payload['success'])
        self.assertIn('nouvel_index', payload['errors'])

    def test_create_prorates_downstream_volume_on_head_meter_period(self):
        Releve.objects.filter(point_eau=self.point_eau).delete()
        Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 4, 1),
            ancienindex=0,
            nouvelindex=0,
            consommation=0,
            volume_facture=0,
            statut='impaye',
        )
        Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 4, 21),
            ancienindex=0,
            nouvelindex=40,
            consommation=40,
            volume_facture=35,
            statut='impaye',
        )
        Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 5, 1),
            ancienindex=40,
            nouvelindex=60,
            consommation=20,
            volume_facture=15,
            statut='impaye',
        )
        ReleveCompteurTete.objects.create(
            compteur_tete=self.compteur,
            date_releve=date(2026, 4, 6),
            ancien_index=Decimal('0'),
            nouvel_index=Decimal('10'),
            volume_tete=Decimal('10'),
            volume_aval=Decimal('0'),
            ecart_volume=Decimal('10'),
            nb_points_eau_aval=1,
            explication_ecart='Releve initial',
        )

        response = self.client.post(
            reverse('organization:api_compteurs_releves_create'),
            data={
                'compteur_id': str(self.compteur.id),
                'date_releve': '2026-04-26',
                'nouvel_index': '50',
                'explication_ecart': '',
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        releve = ReleveCompteurTete.objects.get(id=payload['releve']['id'])
        self.assertEqual(releve.volume_tete, Decimal('40'))
        self.assertEqual(releve.volume_aval, Decimal('40'))
        self.assertEqual(releve.ecart_volume, Decimal('0'))

        detail_response = self.client.get(
            reverse('organization:api_compteurs_releves_detail', args=[releve.id])
        )
        detail_payload = detail_response.json()
        self.assertEqual(detail_payload['releve']['date_debut_calcul_display'], '06/04/2026')
        self.assertEqual(detail_payload['releve']['date_fin_calcul_display'], '26/04/2026')
        self.assertEqual(detail_payload['releve']['duree_calcul_jours'], 20)
        self.assertTrue(detail_payload['releve']['aval_estime'])
        self.assertEqual(
            detail_payload['points_details'][0]['methode_estimation'],
            'Prorata journalier',
        )
        self.assertTrue(detail_payload['points_details'][0]['prorata_utilise'])
        self.assertEqual(detail_payload['points_details'][0]['volume_reel_index'], 60.0)

        page_response = self.client.get(
            reverse('organization:compteurs_releves_list')
        )
        self.assertEqual(page_response.context['stats']['compteurs_classes'], 1)
        self.assertEqual(page_response.context['stats']['total_compteurs'], 1)
        self.assertEqual(page_response.context['stats']['good_compteurs_percent'], 100.0)
        self.assertEqual(page_response.context['selected_vue_releves'], 'groupee')
        self.assertEqual(len(page_response.context['grouped_releves']), 1)
        self.assertEqual(len(page_response.context['grouped_ouvrages']), 1)
        self.assertEqual(
            page_response.context['grouped_ouvrages'][0]['ouvrage_nom'],
            'Ouvrage Compteur',
        )
        self.assertEqual(page_response.context['grouped_ouvrages'][0]['nb_compteurs'], 1)
        self.assertEqual(page_response.context['grouped_ouvrages'][0]['good_count'], 1)
        self.assertEqual(page_response.context['grouped_ouvrages'][0]['acceptable_count'], 0)
        self.assertEqual(page_response.context['grouped_ouvrages'][0]['investigate_count'], 0)
        self.assertEqual(page_response.context['grouped_ouvrages'][0]['critical_count'], 0)
        self.assertContains(page_response, 'Periode comparee')
        self.assertContains(page_response, '06/04/2026')
        self.assertContains(page_response, '26/04/2026')
        self.assertContains(page_response, '20 jour(s)')
        self.assertContains(page_response, 'Reseau bon')
        self.assertContains(page_response, 'Ecart tolerable')
        self.assertContains(page_response, 'de la repartition')

        ReleveCompteurTete.objects.create(
            compteur_tete=self.compteur,
            date_releve=date(2026, 5, 26),
            ancien_index=Decimal('50'),
            nouvel_index=Decimal('80'),
            volume_tete=Decimal('30'),
            volume_aval=Decimal('25'),
            ecart_volume=Decimal('5'),
            nb_points_eau_aval=1,
            explication_ecart='Ecart mai',
        )
        latest_response = self.client.get(
            reverse('organization:compteurs_releves_list'),
            {'vue_releves': 'derniers'},
        )
        self.assertEqual(latest_response.context['selected_vue_releves'], 'derniers')
        self.assertEqual(len(latest_response.context['releves_rows']), 1)
        self.assertEqual(latest_response.context['releves_rows'][0]['date_fin_calcul'], '2026-05-26')
        self.assertContains(latest_response, 'Derniers releves')

    def test_changement_compteur_closes_old_meter_and_transfers_points(self):
        response = self.client.post(
            reverse('organization:api_compteurs_releves_changement_compteur'),
            data={
                'ancien_compteur_id': str(self.compteur.id),
                'date_changement': '2026-04-30',
                'index_retrait': '50',
                'nouveau_numero': 'CT-002',
                'index_initial_nouveau': '3',
                'nouvelle_marque': 'Hydria',
                'nouveau_modele': 'H2',
                'motif': 'defectueux',
                'explication': 'Le compteur se bloque regulierement.',
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['success'])

        changement = ChangementCompteurTete.objects.select_related(
            'ancien_compteur', 'nouveau_compteur', 'releve_final_ancien'
        ).get()
        self.assertEqual(changement.ancien_compteur_id, self.compteur.id)
        self.assertEqual(changement.index_retrait, Decimal('50'))
        self.assertEqual(changement.index_initial_nouveau, Decimal('3'))
        self.assertEqual(changement.nb_points_transferes, 1)
        self.assertEqual(changement.created_by, self.user)

        self.compteur.refresh_from_db()
        self.point_eau.refresh_from_db()
        self.assertEqual(self.compteur.statut, 'hors_service')
        self.assertEqual(self.point_eau.compteur_tete_id, changement.nouveau_compteur_id)
        self.assertEqual(changement.nouveau_compteur.index_initial, Decimal('3'))
        self.assertEqual(changement.nouveau_compteur.date_installation, date(2026, 4, 30))
        self.assertEqual(changement.releve_final_ancien.nouvel_index, Decimal('50'))
        self.assertEqual(changement.releve_final_ancien.volume_tete, Decimal('50'))

        next_response = self.client.post(
            reverse('organization:api_compteurs_releves_create'),
            data={
                'compteur_id': str(changement.nouveau_compteur_id),
                'date_releve': '2026-05-31',
                'nouvel_index': '23',
                'explication_ecart': 'Comparaison apres installation du nouveau compteur.',
            },
        )
        self.assertEqual(next_response.status_code, 200)
        next_payload = next_response.json()
        next_releve = ReleveCompteurTete.objects.get(id=next_payload['releve']['id'])
        self.assertEqual(next_releve.ancien_index, Decimal('3'))
        self.assertEqual(next_releve.volume_tete, Decimal('20'))
        self.assertEqual(next_payload['releve']['date_debut_calcul'], '2026-04-30')

        page_response = self.client.get(reverse('organization:compteurs_releves_list'))
        self.assertContains(page_response, 'Changements de compteur')
        self.assertContains(page_response, 'CT-001')
        self.assertContains(page_response, 'CT-002')
        self.assertContains(page_response, 'Compteur defectueux')

    @patch(
        'apps.organization.views_compteurs_releves.ChangementCompteurTete.objects.create',
        side_effect=IntegrityError,
    )
    def test_changement_compteur_is_atomic(self, changement_create_mock):
        response = self.client.post(
            reverse('organization:api_compteurs_releves_changement_compteur'),
            data={
                'ancien_compteur_id': str(self.compteur.id),
                'date_changement': '2026-04-30',
                'index_retrait': '50',
                'nouveau_numero': 'CT-ROLLBACK',
                'index_initial_nouveau': '0',
                'motif': 'autre',
                'explication': 'Test de retour arriere atomique.',
            },
        )

        self.assertEqual(response.status_code, 400)
        self.compteur.refresh_from_db()
        self.point_eau.refresh_from_db()
        self.assertEqual(self.compteur.statut, 'operationnel')
        self.assertEqual(self.point_eau.compteur_tete_id, self.compteur.id)
        self.assertFalse(CompteurTete.objects.filter(numerocompteur='CT-ROLLBACK').exists())
        self.assertFalse(ReleveCompteurTete.objects.filter(compteur_tete=self.compteur).exists())
        self.assertFalse(ChangementCompteurTete.objects.exists())


class ChangementCompteurAbonneTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_superuser(
            username='admin_compteur_abonne',
            email='admin.compteur.abonne@example.com',
            password='testpass123',
        )
        self.client.force_login(self.user)
        self.organisation = Organisation.objects.create(
            nom_organisation='Org Compteur Abonne',
            type_organisation=Organisation.TYPE_GESTIONNAIRE,
        )
        self.type_ouvrage = TypeOuvrage.objects.create(nom_type_ouvrage='Reseau test')
        self.ouvrage = Ouvrage.objects.create(
            nom_ouvrage='Ouvrage Compteur Abonne', type_ouvrage=self.type_ouvrage
        )
        ModaliteGestion.objects.create(
            ouvrage=self.ouvrage,
            type_gestion='privee',
            gestionnaire=self.organisation,
            date_debut=date(2026, 1, 1),
            modalitepaiement='au_volume',
        )
        self.type_point_eau = TypePointEau.objects.create(nom_type='BP compteur')
        self.point_eau = PointEau.objects.create(
            nomabonne='Abonne remplacement',
            ouvrage=self.ouvrage,
            type_point_eau=self.type_point_eau,
            code_point_eau='PE-REMPLACEMENT',
            datebranchement=timezone.make_aware(datetime(2026, 1, 1)),
        )
        self.ancien_releve = Releve.objects.create(
            point_eau=self.point_eau,
            abonne=self.point_eau.nomabonne,
            datereleve=date(2026, 4, 1),
            ancienindex=80,
            nouvelindex=100,
            consommation=20,
            volume_facture=20,
            montant=Decimal('5000'),
            statut='impaye',
        )

    def test_change_then_next_invoice_uses_two_meter_segments(self):
        change_response = self.client.post(
            reverse('billing:api_changement_compteur_abonne'),
            data=json.dumps({
                'point_eau_id': str(self.point_eau.id),
                'ancien_numero': 'CA-OLD-001',
                'date_changement': '2026-04-15',
                'index_retrait': '118',
                'nouveau_numero': 'CA-NEW-001',
                'index_initial_nouveau': '3',
                'motif': 'defectueux',
                'explication': 'Compteur remplace apres controle terrain.',
            }),
            content_type='application/json',
        )
        self.assertEqual(change_response.status_code, 200, change_response.content)
        self.ancien_releve.refresh_from_db()
        self.assertEqual(self.ancien_releve.consommation, 20)
        self.assertEqual(self.ancien_releve.montant, Decimal('5000'))

        save_response = self.client.post(
            reverse('billing:api_save_releve'),
            data={
                'point_eau_id': str(self.point_eau.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '11',
                'volume_non_facture': '0',
                'cartes_sociales': '[]',
            },
        )
        self.assertEqual(save_response.status_code, 200, save_response.content)
        releve = Releve.objects.get(id=save_response.json()['releve_id'])
        self.assertEqual(Decimal(str(releve.consommation)), Decimal('26'))
        self.assertEqual(releve.origine_calcul, 'multi_compteur')
        self.assertEqual(releve.compteur_abonne.numero, 'CA-NEW-001')
        segments = list(releve.segments_consommation.order_by('ordre'))
        self.assertEqual([segment.volume for segment in segments], [Decimal('18'), Decimal('8')])
        changement = ChangementCompteurAbonne.objects.get()
        self.assertEqual(changement.releve_facture_id, releve.id)

    def test_negative_non_billed_volume_is_rejected_for_new_reading(self):
        response = self.client.post(
            reverse('billing:api_save_releve'),
            data={
                'point_eau_id': str(self.point_eau.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '110',
                'volume_non_facture': '-10',
                'cartes_sociales': '[]',
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Changer le compteur', response.json()['error'])

    def test_web_reading_accepts_optional_naive_terrain_capture_date(self):
        response = self.client.post(
            reverse('billing:api_save_releve'),
            data={
                'point_eau_id': str(self.point_eau.id),
                'date_releve': '2026-04-30',
                'nouvel_index': '110',
                'volume_non_facture': '0',
                'cartes_sociales': '[]',
                'terrain_captured_at': '2026-04-30T08:15:00',
            },
        )

        self.assertEqual(response.status_code, 200, response.content)
        releve = Releve.objects.get(id=response.json()['releve_id'])
        self.assertIsNotNone(releve.terrain_captured_at)
        self.assertTrue(timezone.is_aware(releve.terrain_captured_at))

    def test_paid_invoice_cannot_be_rewritten(self):
        self.ancien_releve.numeroDiameX = 'DX-001'
        self.ancien_releve.statut = 'paye'
        self.ancien_releve.save(update_fields=['numeroDiameX', 'statut'])
        response = self.client.post(
            reverse('billing:api_modifier_releve'),
            data=json.dumps({
                'releve_id': str(self.ancien_releve.id),
                'nouvel_index': 101,
                'volume_non_facture': 0,
            }),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 409)
