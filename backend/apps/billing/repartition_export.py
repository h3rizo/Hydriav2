"""Export Excel du tableau « Répartition des branchements par commune »
de la page Points d'eau consolidés.

Reproduit la hiérarchie et le tri de l'écran : Région > District > Commune >
Ouvrage, régions et districts par ordre alphabétique, communes et ouvrages par
total décroissant (types retenus uniquement), égalités par ordre alphabétique.
"""
from datetime import datetime

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _type_totals(rows):
    totals = {}
    for row in rows:
        totals[row['type']] = totals.get(row['type'], 0) + row['n']
    return totals


def _build_tree(rows, types):
    """Retourne la liste des régions triées, avec districts, communes et ouvrages."""
    tree = {}
    for row in rows:
        if row['type'] not in types:
            continue
        districts = tree.setdefault(row['region'], {})
        communes = districts.setdefault(row['district'], {})
        commune = communes.setdefault(
            row['commune_id'] or row['commune'],
            {'nom': row['commune'], 'par_type': {}, 'ouvrages': {}},
        )
        commune['par_type'][row['type']] = commune['par_type'].get(row['type'], 0) + row['n']
        ouvrage = commune['ouvrages'].setdefault(
            row['ouvrage_id'] or row['ouvrage'],
            {'nom': row['ouvrage'], 'par_type': {}},
        )
        ouvrage['par_type'][row['type']] = ouvrage['par_type'].get(row['type'], 0) + row['n']

    def total(par_type):
        return sum(par_type.values())

    def sum_into(items):
        acc = {}
        for item in items:
            for type_nom, n in item['par_type'].items():
                acc[type_nom] = acc.get(type_nom, 0) + n
        return acc

    def by_total_then_name(items):
        # Deux tris stables : nom croissant, puis total décroissant
        items = sorted(items, key=lambda item: item['nom'].casefold())
        return sorted(items, key=lambda item: -item['total'])

    regions = []
    for region_nom in sorted(tree, key=str.casefold):
        districts = []
        for district_nom in sorted(tree[region_nom], key=str.casefold):
            communes = []
            for commune in tree[region_nom][district_nom].values():
                ouvrages = [
                    {**ouvrage, 'total': total(ouvrage['par_type'])}
                    for ouvrage in commune['ouvrages'].values()
                ]
                communes.append({
                    'nom': commune['nom'],
                    'par_type': commune['par_type'],
                    'total': total(commune['par_type']),
                    'ouvrages': by_total_then_name(ouvrages),
                })
            communes = by_total_then_name(communes)
            par_type = sum_into(communes)
            districts.append({'nom': district_nom, 'communes': communes,
                              'par_type': par_type, 'total': total(par_type)})
        par_type = sum_into(districts)
        regions.append({'nom': region_nom, 'districts': districts,
                        'par_type': par_type, 'total': total(par_type)})
    return regions


def export_repartition_excel(rows, selected_types=None, filtres_description=''):
    """Construit la réponse HTTP .xlsx.

    rows : lignes agrégées {region, district, commune_id, commune, ouvrage_id,
           ouvrage, type, n} (celles envoyées à la page).
    selected_types : types cochés à l'écran ; vide/None = tous les types.
    """
    type_totals = _type_totals(rows)
    all_types = sorted(type_totals, key=lambda t: -type_totals[t])
    types = [t for t in all_types if not selected_types or t in selected_types]
    regions = _build_tree(rows, set(types))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Répartition branchements'

    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    fills = {
        'region': PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid'),
        'district': PatternFill(start_color='E7ECF6', end_color='E7ECF6', fill_type='solid'),
        'commune': PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid'),
        'total': PatternFill(start_color='BFBFBF', end_color='BFBFBF', fill_type='solid'),
    }
    last_col = len(types) + 2

    ws.cell(row=1, column=1, value='Répartition des branchements par commune').font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Généré le {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
                                   + (f" — {filtres_description}" if filtres_description else ''))
    ws.cell(row=2, column=1).font = Font(italic=True, color='666666')
    if selected_types and len(types) < len(all_types):
        exclus = [t for t in all_types if t not in types]
        ws.cell(row=3, column=1, value='Types exclus : ' + ', '.join(exclus)).font = Font(italic=True, color='666666')

    header_row = 5
    headers = ['Région / District / Commune / Ouvrage'] + types + ['Total']
    for col, label in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='left' if col == 1 else 'center', vertical='center', wrap_text=True)
        cell.border = border

    row_num = header_row + 1

    def write_row(label, par_type, level, outline=0):
        nonlocal row_num
        indent = {'region': 0, 'district': 1, 'commune': 2, 'ouvrage': 4, 'total': 0}[level]
        label_cell = ws.cell(row=row_num, column=1, value=label)
        label_cell.alignment = Alignment(indent=indent)
        values = [par_type.get(t) or None for t in types] + [sum(par_type.get(t, 0) for t in types)]
        for col, value in enumerate(values, 2):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            if level != 'ouvrage':
                cell.font = Font(bold=True)
                cell.fill = fills[level]
        if outline:
            ws.row_dimensions[row_num].outlineLevel = outline
        row_num += 1

    grand = {}
    for region in regions:
        for t, n in region['par_type'].items():
            grand[t] = grand.get(t, 0) + n
        write_row(region['nom'], region['par_type'], 'region')
        for district in region['districts']:
            nb = len(district['communes'])
            write_row(f"{district['nom']} ({nb} commune{'s' if nb > 1 else ''})", district['par_type'], 'district', 1)
            for commune in district['communes']:
                nb = len(commune['ouvrages'])
                write_row(f"{commune['nom']} ({nb} ouvrage{'s' if nb > 1 else ''})", commune['par_type'], 'commune', 2)
                for ouvrage in commune['ouvrages']:
                    write_row(ouvrage['nom'], ouvrage['par_type'], 'ouvrage', 3)
    write_row('Total', grand, 'total')

    ws.column_dimensions['A'].width = 48
    for col in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.sheet_properties.outlinePr.summaryBelow = False

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"repartition_branchements_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
