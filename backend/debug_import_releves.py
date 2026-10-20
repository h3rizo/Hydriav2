"""
Script de diagnostic pour l'import de relevés Excel
Affiche les 20 premières lignes avec tous les détails

Usage:
    python debug_import_releves.py chemin/vers/fichier.xlsx
"""

import sys
from openpyxl import load_workbook
from datetime import datetime

def parse_date(date_val):
    if not date_val:
        return None
    if isinstance(date_val, datetime):
        return date_val.date()
    date_str = str(date_val).strip()
    try:
        return datetime.strptime(date_str, '%d/%m/%Y').date()
    except:
        try:
            return datetime.strptime(date_str, '%Y-%m-%d').date()
        except:
            return None

def debug_excel(filepath):
    wb = load_workbook(filepath)
    ws = wb.active

    print(f"\n{'='*120}")
    print(f"DIAGNOSTIC IMPORT RELEVÉS : {filepath}")
    print(f"{'='*120}\n")

    # Header
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value or ''))

    print(f"En-têtes détectés : {headers}\n")
    print(f"{'Ligne':<6} | {'Point Eau ID':<38} | {'Date':<12} | {'Index':<8} | {'Vol NF':<8} | {'Montant':<12} | {'Num DiameX':<12} | {'Base DiameX':<15}")
    print('-' * 150)

    # Données (20 premières lignes)
    for row_num, row in enumerate(ws.iter_rows(min_row=2, max_row=21, values_only=True), 2):
        if not any(row):
            continue

        point_eau_id = str(row[0]).strip() if row[0] else 'VIDE'
        date_releve = parse_date(row[1]) if len(row) > 1 else None
        nouvel_index = row[2] if len(row) > 2 else None
        volume_non_facture = row[3] if len(row) > 3 else None
        montant = row[4] if len(row) > 4 else None
        numero_diamex = row[5] if len(row) > 5 else None
        base_diamex = row[6] if len(row) > 6 else None

        # Format d'affichage
        pe_display = point_eau_id[:36] + '..' if len(point_eau_id) > 38 else point_eau_id
        date_display = str(date_releve) if date_releve else 'INVALIDE'
        index_display = str(nouvel_index) if nouvel_index not in (None, '') else 'VIDE'
        vol_nf_display = str(volume_non_facture) if volume_non_facture not in (None, '') else '0'
        montant_display = str(montant) if montant not in (None, '') else 'VIDE'
        num_diamex_display = str(numero_diamex) if numero_diamex not in (None, '') else 'VIDE'
        base_diamex_display = str(base_diamex) if base_diamex not in (None, '') else 'VIDE'

        print(f"{row_num:<6} | {pe_display:<38} | {date_display:<12} | {index_display:<8} | {vol_nf_display:<8} | {montant_display:<12} | {num_diamex_display:<12} | {base_diamex_display:<15}")

    print('\n' + '='*120)

    # Analyser les doublons
    print("\nRECHERCHE DE DOUBLONS (même Point d'eau + même Date) :\n")

    dates_by_pe = {}
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(row):
            continue

        point_eau_id = str(row[0]).strip() if row[0] else None
        date_releve = parse_date(row[1]) if len(row) > 1 else None

        if point_eau_id and date_releve:
            key = (point_eau_id, date_releve)
            if key not in dates_by_pe:
                dates_by_pe[key] = []
            dates_by_pe[key].append(row_num)

    doublons_trouves = False
    for (pe_id, date), lignes in dates_by_pe.items():
        if len(lignes) > 1:
            doublons_trouves = True
            print(f"⚠️  Point d'eau {pe_id[:8]}... | Date {date} : {len(lignes)} lignes trouvées (lignes {lignes})")

    if not doublons_trouves:
        print("✅ Aucun doublon détecté")

    print('\n' + '='*120 + '\n')

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python debug_import_releves.py chemin/vers/fichier.xlsx")
        sys.exit(1)

    filepath = sys.argv[1]
    debug_excel(filepath)
