import os
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'hydria.settings')
django.setup()

from apps.employees.models import PaiementCNaPSFMFP
from openpyxl import load_workbook
from django.conf import settings

# ID du paiement
paiement_id = '5e70939a-aa39-4da1-aa98-5518c9d3923d'

print("Tentative de recuperation du paiement...")
try:
    paiement = PaiementCNaPSFMFP.objects.select_related('organisation').get(id=paiement_id)
    print(f"[OK] Paiement trouve: {paiement}")
    print(f"  Organisation: {paiement.organisation.nom_organisation}")
    print(f"  Trimestre: {paiement.trimestre} - Annee: {paiement.annee}")
    
    # Test chargement du modèle
    modele_path = os.path.join(settings.BASE_DIR, 'Modèle Facture CNaPS FMFP.xlsx')
    print(f"\nChemin modele: {modele_path}")
    print(f"Existe: {os.path.exists(modele_path)}")
    
    print("\nTentative de chargement du modele Excel...")
    wb = load_workbook(modele_path)
    print(f"[OK] Modele charge avec succes")
    print(f"  Feuilles: {wb.sheetnames}")
    
    # Test récupération des détails
    print("\nRecuperation des details...")
    details = paiement.details.select_related('employe', 'employe__poste').order_by('employe_fonction', 'employe_nom_complet')
    print(f"[OK] Nombre de details: {details.count()}")
    
    for detail in details[:3]:
        print(f"  - {detail.employe_nom_complet}")
    
except PaiementCNaPSFMFP.DoesNotExist:
    print("[ERREUR] Paiement introuvable")
except Exception as e:
    import traceback
    print(f"[ERREUR] {e}")
    traceback.print_exc()
