from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView
from django.contrib.auth import login, REDIRECT_FIELD_NAME
from django.contrib.auth.models import User
from django.utils.http import url_has_allowed_host_and_scheme
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.db.models import Count, Sum, Q, Case, When, Value, CharField, F
from django.db import transaction
from django.contrib import messages
from django.conf import settings
from django.urls import reverse, NoReverseMatch
from django.core.exceptions import ObjectDoesNotExist, ValidationError
from django.core.validators import validate_email
from django.utils.crypto import get_random_string
from datetime import date
import zlib
from apps.geography.models import Region, District, Commune, Parcelle
from apps.organization.models import Ouvrage
from apps.billing.models import PointEau, TypePointEau, ModaliteGestion, CompteurTete, Fournisseur, TarifVolumetrique, CoupureEau
from apps.billing.payment_modalities import (
    build_payment_modality_scope,
    build_point_usage_modality_breakdown,
    is_private_point_type_name,
    private_point_type_ids,
)
from apps.billing.points_eau_statut import filtrer_points_annee
from apps.infrastructure.models import EquipementConstitutif, Adresse, CarteSociale
from apps.employees.permissions import (
    employe_required,
    filter_queryset_by_permissions,
    has_admin_access,
    has_compta_home_access,
    has_compta_access,
    is_visiteur,
)
from apps.organization.scoping import (
    apply_user_responsable_ouvrage_scope,
    get_user_accessible_gestionnaire_ids,
    filter_ouvrages_by_gestionnaire_period,
    filter_modalites_by_gestionnaire_period,
)
from apps.system.visiteur import ensure_visiteur_profile

class CustomLoginView(LoginView):
    """
    Vue de connexion personnalisée:
    - redirige les clients actifs vers le détail de leur point d'eau après connexion.
    - conserve le comportement standard pour les autres utilisateurs.
    """

    template_name = 'registration/login.html'

    @staticmethod
    def _safe_related(instance, attr_name):
        try:
            return getattr(instance, attr_name)
        except ObjectDoesNotExist:
            return None

    def get_success_url(self):
        request = self.request

        redirect_to = request.POST.get(
            REDIRECT_FIELD_NAME, request.GET.get(REDIRECT_FIELD_NAME, '')
        )
        if redirect_to and url_has_allowed_host_and_scheme(
            url=redirect_to,
            allowed_hosts={request.get_host()},
            require_https=request.is_secure(),
        ):
            return redirect_to

        user = request.user

        candidat = self._safe_related(user, 'candidat')
        if candidat and candidat.statut == 'test_termine':
            latest_session = (
                candidat.sessions_test
                .filter(statut__in=['termine', 'expire'])
                .order_by('-date_fin', '-date_debut', '-created_at')
                .first()
            )
            if latest_session:
                return reverse('recrutement:resultats', kwargs={'session_id': latest_session.id})

        client = self._safe_related(user, 'client')
        if client and getattr(client, 'actif', False):
            if getattr(client, 'point_eau_id', None):
                return reverse(
                    'core:search_detail',
                    kwargs={
                        'entity_type': 'point_eau',
                        'entity_id': str(client.point_eau_id),
                    }
                )
            return reverse('client_home')

        return reverse('home_launcher')


def _exclude_dashboard_non_operational_ouvrages(ouvrages_qs):
    return ouvrages_qs.exclude(
        Q(type_ouvrage__nom_type_ouvrage__icontains='suspend') |
        Q(type_ouvrage__nom_type_ouvrage__icontains='potentiel')
    ).distinct()


def _pct_label(part, total):
    """Pourcentage entier en libellé : « 50 % », « < 1 % » pour une part non nulle, « 0 % »."""
    if not total or not part:
        return '0 %'
    pct = round(part / total * 100)
    return '< 1 %' if pct < 1 else f'{pct} %'


def _top_ouvrages(compteur, n=3):
    """Les ``n`` ouvrages comptant le plus de points : ``[{'nom': 'Ambato', 'count': 5}, ...]``."""
    return [
        {'nom': nom, 'count': effectif}
        for nom, effectif in sorted(compteur.items(), key=lambda kv: (-kv[1], kv[0]))[:n]
    ]


def _ventiler_par_type(points_qs, point_ids=None):
    """Répartition par type de point d'eau, avec les 3 premiers ouvrages de chaque type.

    ``[{'label': 'BF institutionnel', 'count': 3, 'ouvrages': [{'nom': 'Ambato', 'count': 2}, ...]}, ...]``

    BF / BP / Kiosque / PMH… ; un branchement institutionnel (BF ou BP coché
    « institutionnel ») est compté à part, avec le suffixe « institutionnel ».
    ``point_ids=None`` : tout le queryset. Trié par effectif décroissant.
    """
    if point_ids is not None:
        if not point_ids:
            return []
        points_qs = points_qs.filter(id__in=point_ids)
    effectifs, ouvrages = {}, {}
    # L'id reste dans le SELECT : le périmètre est DISTINCT, sans lui deux points de même type seraient fusionnés
    for _id, nom_type, est_inst, nom_ouvrage in points_qs.order_by().values_list(
        'id', 'type_point_eau__nom_type', 'estbranchementinstitutionnel', 'ouvrage__nom_ouvrage'
    ):
        type_nom = (nom_type or '').strip()
        type_cle = type_nom.upper()
        if not type_nom:
            label = 'Sans type'
        elif type_cle in ('BF', 'BP'):
            label = f'{type_cle} institutionnel' if est_inst else type_cle
        elif 'KIOSQUE' in type_cle:
            label = 'Kiosque'
        elif type_cle in ('PMH', 'PMHI'):
            label = 'PMH'
        else:
            label = type_nom
        effectifs[label] = effectifs.get(label, 0) + 1
        par_ouvrage = ouvrages.setdefault(label, {})
        nom = nom_ouvrage or 'Sans ouvrage'
        par_ouvrage[nom] = par_ouvrage.get(nom, 0) + 1
    return [
        {'label': l, 'count': n, 'ouvrages': _top_ouvrages(ouvrages[l])}
        for l, n in sorted(effectifs.items(), key=lambda kv: (-kv[1], kv[0]))
    ]


def _build_dashboard_service_scope(points_eau_qs, ouvrages_qs, annee=None, avec_coupures=True):
    """Périmètre des KPI du dashboard.

    Retourne ``(ouvrages en service, points du périmètre, points actifs, ids coupés, ids abandonnés)``.
    Actif = branché ET ni coupé ni abandonné à la date de référence (fin d'année du filtre,
    plafonnée à aujourd'hui) : même règle que ``apps.billing.points_eau_statut``, appliquée à
    l'historique CoupureEau. ``avec_coupures=False`` évite ce calcul quand seul le périmètre sert.
    """
    ouvrages_service_qs = _exclude_dashboard_non_operational_ouvrages(ouvrages_qs)
    points_service_base_qs = points_eau_qs.filter(
        ouvrage_id__in=ouvrages_service_qs.values_list('id', flat=True)
    ).distinct()
    points_service_actifs_qs = points_service_base_qs.filter(datebranchement__isnull=False)
    coupes_ids, abandonnes_ids = set(), set()
    if avec_coupures:
        coupes_ids, abandonnes_ids = _classer_coupes_point_ids(points_service_base_qs, annee=annee)
        if coupes_ids or abandonnes_ids:
            points_service_actifs_qs = points_service_actifs_qs.exclude(id__in=coupes_ids | abandonnes_ids)
    return ouvrages_service_qs, points_service_base_qs, points_service_actifs_qs, coupes_ids, abandonnes_ids


@require_http_methods(["POST"])
def visitor_email_login(request):
    """
    Connexion visiteur via email sans preconfiguration manuelle.
    - cree le compte si inexistant
    - force le role visiteur en lecture seule
    - cree/lie un employe pour permettre la capture GPS
    """
    email = (request.POST.get('visitor_email') or '').strip().lower()
    next_url = request.POST.get('next') or request.GET.get('next') or ''

    if getattr(request, 'is_accounting_subdomain', False) or getattr(request, 'is_accounting_workflow_subdomain', False):
        messages.error(
            request,
            "La connexion visiteur n'est pas autorisee sur l'accueil comptable."
        )
        return redirect('login')

    if not email:
        messages.error(request, "Veuillez saisir une adresse email.")
        return redirect('login')

    try:
        validate_email(email)
    except ValidationError:
        messages.error(request, "Adresse email invalide.")
        return redirect('login')

    try:
        with transaction.atomic():
            user = User.objects.filter(email__iexact=email).first()
            if user is None:
                user = User.objects.filter(username__iexact=email).first()

            if user and not is_visiteur(user):
                messages.error(
                    request,
                    "Un compte non-visiteur existe deja avec cet email. Utilisez la connexion standard."
                )
                return redirect('login')

            if user is None:
                user = User.objects.create_user(
                    username=email,
                    email=email,
                    first_name='Visiteur',
                    last_name='',
                    password=get_random_string(24)
                )

            ensure_visiteur_profile(user, email=email)

        login(request, user, backend='django.contrib.auth.backends.ModelBackend')
        messages.success(request, "Connexion visiteur reussie.")
    except Exception as exc:
        messages.error(request, f"Connexion visiteur impossible: {exc}")
        return redirect('login')

    if next_url.startswith('/'):
        return redirect(next_url)
    return redirect('home_launcher')


@login_required
def client_home(request):
    """Page d'accueil fallback pour les utilisateurs de type client."""
    client = getattr(request.user, 'client', None)
    if not client:
        return render(request, 'client/home.html', {'client': None})

    return render(request, 'client/home.html', {'client': client})


def _can_access_admin_sensitive_buttons(user):
    """Autoriser uniquement certains postes aux boutons sensibles sur /accueil/."""
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    if not hasattr(user, 'employe') or not user.employe:
        return False

    poste = getattr(user.employe, 'poste', None)
    if not poste:
        return False

    nom_poste = getattr(poste, 'nom_poste', str(poste)) or ''
    nom_normalise = nom_poste.strip().lower()

    return nom_normalise in {
        'directeur',
        'directeur adjoint',
        'responsable administratif et financier',
    }


def _safe_reverse_url(route_name):
    try:
        return reverse(route_name)
    except NoReverseMatch:
        return None


def _launcher_item(label, route_name, icon, description='', badge='', query_string='', primary=False):
    url = _safe_reverse_url(route_name)
    if not url:
        return None
    if query_string:
        url = f'{url}?{query_string.lstrip("?")}'
    return {
        'label': label,
        'url': url,
        'icon': icon,
        'description': description,
        'badge': badge,
        'primary': primary,
    }


def _compact_items(items):
    return [item for item in items if item]


def _payroll_url(path):
    base_url = (getattr(settings, 'PAYROLL_SUBDOMAIN_URL', 'https://paie.hydria.mg') or 'https://paie.hydria.mg').rstrip('/')
    if not path.startswith('/'):
        path = f'/{path}'
    return f'{base_url}{path}'


def _build_launcher_workspaces(user):
    is_employee = hasattr(user, 'employe') and bool(user.employe)
    is_admin_sensitive = _can_access_admin_sensitive_buttons(user)
    can_compta = has_compta_access(user) or is_admin_sensitive or getattr(user, 'is_superuser', False)
    can_admin = has_admin_access(user) or is_admin_sensitive or getattr(user, 'is_superuser', False)
    client = getattr(user, 'client', None)

    workspaces = []

    if is_employee:
        workspaces.extend([
            {
                'key': 'facturation',
                'title': 'Gerer les releves et factures',
                'subtitle': 'Saisir les releves, generer les factures et encaisser.',
                'icon': 'file-text',
                'tone': 'workspace-tone-info',
                'items': _compact_items([
                    _launcher_item('Releves compteurs', 'billing:releves', 'gauge', 'Saisie'),
                    _launcher_item('Facturation volume', 'billing:workflow_facturation', 'file-text', 'Generation'),
                    _launcher_item('Encaissements volume', 'billing:encaissement_factures', 'banknote', 'Paiement'),
                ]),
            },
            {
                'key': 'ouvrages',
                'title': 'Saisir des nouveaux ouvrages',
                'subtitle': 'Creer et retrouver les ouvrages, points d eau et demandes terrain.',
                'icon': 'droplets',
                'tone': 'workspace-tone-primary',
                'items': _compact_items([
                    _launcher_item(
                        'Saisir un nouvel ouvrage',
                        'organization:ouvrage_assistant',
                        'plus-circle',
                        "Assistant guide : ouvrage, points d'eau, gestion, responsable",
                        primary=True,
                    ),
                    _launcher_item(
                        'Gerer un ouvrage existant',
                        'organization:ouvrage_assistant',
                        'settings-2',
                        'Completer un ouvrage deja cree (points d eau, gestion, responsable)',
                        query_string='mode=existing',
                        primary=True,
                    ),
                    _launcher_item('Ouvrages', 'ouvrages:dashboard', 'droplets', 'Base ouvrages'),
                    _launcher_item('Recherche globale', 'core:search', 'search', 'Retrouver un ouvrage'),
                    _launcher_item('Demandes BP', 'billing:demandes_bp', 'file-plus-2', 'Branchements'),
                ]),
            },
        ])

    if is_employee or can_compta:
        workspaces.append({
            'key': 'suivi_paiements',
            'title': 'Suivi des paiements',
            'subtitle': 'Consulter le recouvrement forfaitaire et volumetrique.',
            'icon': 'circle-dollar-sign',
            'tone': 'workspace-tone-success',
            'items': _compact_items([
                _launcher_item(
                    'Paiements forfaitaires',
                    'core:dashboard_payments_flat',
                    'coins',
                    'Score et ecarts de paiement',
                    query_string='score_ecart=1',
                ),
                _launcher_item(
                    'Paiements au volume',
                    'accounting:dashboard_paiements_volume',
                    'waves',
                    'Suivi du recouvrement',
                ),
            ]),
        })

    if can_compta:
        workspaces.append({
            'key': 'dashboards',
            'title': 'Consulter les tableaux de bord',
            'subtitle': 'Piloter recouvrement, recettes, cartes, rapports et comptabilite.',
            'icon': 'bar-chart-3',
            'tone': 'workspace-tone-secondary',
            'items': _compact_items([
                _launcher_item('Dashboard principal', 'core:dashboard', 'trending-up', 'Global'),
                _launcher_item('Paiements volume', 'accounting:dashboard_paiements_volume', 'chart-pie', 'Recouvrement'),
                _launcher_item('Paiements forfaitaires', 'accounting:dashboard_paiements_forfaitaire', 'chart-column', 'Recouvrement'),
            ]),
        })

    if client and getattr(client, 'actif', False) and getattr(client, 'point_eau_id', None):
        workspaces.insert(0, {
            'key': 'client',
            'title': 'Consulter et payer mes factures',
            'subtitle': 'Acceder au point d eau et aux paiements disponibles.',
            'icon': 'smartphone',
            'tone': 'workspace-tone-success',
            'items': [{
                'label': 'Mon point d eau',
                'url': reverse('core:search_detail', kwargs={
                    'entity_type': 'point_eau',
                    'entity_id': str(client.point_eau_id),
                }),
                'icon': 'droplet',
                'description': 'Factures et MVola',
                'badge': '',
            }],
        })

    return [workspace for workspace in workspaces if workspace['items']]


def _get_dashboard_allowed_organisation_ids(user):
    """
    IDs des organisations gestionnaires accessibles pour le dashboard.
    - None: pas de restriction (superuser/visiteur)
    - []: aucun acces
    - [id...]: gestionnaires accessibles
    """
    if not user or not getattr(user, 'is_authenticated', False):
        return []
    if getattr(user, 'is_superuser', False) or is_visiteur(user):
        return None

    gestionnaire_ids = get_user_accessible_gestionnaire_ids(user)
    if gestionnaire_ids is None:
        return None
    return [str(_id) for _id in gestionnaire_ids]


def _resolve_dashboard_organisation_filter(user, requested_organisation_id):
    """
    Normalise le filtre organisation pour qu'un profil Partenaire
    utilise les gestionnaires qui lui sont affectes.
    """
    requested = (requested_organisation_id or '').strip()
    if requested == 'all':
        requested = ''

    allowed_ids = _get_dashboard_allowed_organisation_ids(user)
    if allowed_ids is None:
        return requested

    if requested and requested not in allowed_ids:
        return ''

    if not requested and len(allowed_ids) == 1:
        return allowed_ids[0]

    return requested


def _classer_coupes_point_ids(points_eau_qs, annee=None):
    """
    Retourne ``(coupes_ids, abandonnes_ids)`` : IDs des points d'eau en état "coupé"
    en se basant prioritairement sur les événements CoupureEau/Remise (fallback legacy :
    datearretbranchement si aucun événement).

    Un point dont la dernière coupure (sans remise depuis) remonte à ``ABANDON_JOURS`` jours
    ou plus, à la date de référence (fin d'année du filtre, plafonnée à aujourd'hui), est
    « abandonné » et n'est plus compté parmi les coupés.
    """
    from apps.billing.points_eau_statut import est_abandonne_depuis

    cutoff = date(int(annee), 12, 31) if annee else date.today()
    reference_abandon = min(cutoff, date.today())

    # Sans date de branchement : jamais coupé ni abandonné (statut « sans date » prioritaire,
    # comme dans points_eau_statut). Sous-requêtes : pas de liste d'IDs chargée en mémoire.
    branches_ids = points_eau_qs.filter(datebranchement__isnull=False).order_by().values('id')
    events_qs = CoupureEau.objects.filter(
        point_eau_id__in=branches_ids,
        date_evenement__isnull=False,
        date_evenement__lte=cutoff
    )
    events = events_qs.order_by('point_eau_id', 'date_evenement', 'created_at').values(
        'point_eau_id', 'type_evenement', 'date_evenement'
    )

    # Dernier événement connu : (est une coupure ?, date de l'événement)
    state_by_point = {}
    for evt in events:
        state_by_point[evt['point_eau_id']] = (evt.get('type_evenement') == 'coupure', evt['date_evenement'])

    coupes_ids = set()
    abandonnes_ids = set()
    for pe_id, (is_coupe, date_evt) in state_by_point.items():
        if not is_coupe:
            continue
        if est_abandonne_depuis(date_evt, reference_abandon):
            abandonnes_ids.add(pe_id)
        else:
            coupes_ids.add(pe_id)

    # Fallback pour les anciens enregistrements sans événements
    legacy_qs = PointEau.objects.filter(
        id__in=branches_ids,
        datearretbranchement__isnull=False,
    ).exclude(id__in=events_qs.values('point_eau_id'))
    if annee:
        legacy_qs = legacy_qs.filter(datearretbranchement__year__lte=int(annee))
    for pe_id, date_arret in legacy_qs.values_list('id', 'datearretbranchement'):
        if est_abandonne_depuis(date_arret, reference_abandon):
            abandonnes_ids.add(pe_id)
        else:
            coupes_ids.add(pe_id)

    return coupes_ids, abandonnes_ids


def _active_coupes_point_ids(points_eau_qs, annee=None):
    """IDs des points d'eau coupés depuis moins d'un an (les abandonnés sont à part)."""
    return _classer_coupes_point_ids(points_eau_qs, annee=annee)[0]

@login_required
def home_launcher(request):
    """Page d'accueil avec launcher d'applications"""
    workspaces = _build_launcher_workspaces(request.user)
    active_workspace_key = request.GET.get('workspace') or ''
    valid_workspace_keys = {workspace['key'] for workspace in workspaces}
    show_all_interface = active_workspace_key == 'all'
    if active_workspace_key not in valid_workspace_keys and not show_all_interface:
        active_workspace_key = workspaces[0]['key'] if workspaces else 'all'
        show_all_interface = not workspaces
    active_workspace = next(
        (workspace for workspace in workspaces if workspace['key'] == active_workspace_key),
        None
    )
    return render(request, 'home_launcher.html', {
        'can_access_admin_sensitive_buttons': _can_access_admin_sensitive_buttons(request.user),
        'launcher_workspaces': workspaces,
        'active_workspace_key': active_workspace_key,
        'active_workspace': active_workspace,
        'show_all_interface': show_all_interface,
    })

@employe_required
def home(request):
    """Dashboard principal de l'application Hydria"""
    from datetime import datetime
    from django.db.models import Min
    from apps.organization.models import Organisation

    # Récupérer l'employé connecté et appliquer les filtres de permissions
    employe = request.employe

    # Récupérer et normaliser le filtre organisation selon le scope de l'utilisateur
    organisation_filter = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation', '')
    )
    allowed_organisation_ids = _get_dashboard_allowed_organisation_ids(request.user)

    # Filtrer les ouvrages accessibles selon les permissions
    ouvrages_accessibles = filter_queryset_by_permissions(Ouvrage.objects.all(), request.user)
    scoped_ouvrage_ids = apply_user_responsable_ouvrage_scope(
        ouvrages_accessibles.values_list('id', flat=True),
        request.user,
    )
    if scoped_ouvrage_ids is not None:
        ouvrages_accessibles = ouvrages_accessibles.filter(id__in=scoped_ouvrage_ids)

    # Récupérer uniquement les régions/districts/communes qui ont des points d'eau liés aux ouvrages accessibles
    # Ouvrage -> PointEau -> Parcelle -> Village -> Fokontany -> Commune
    from apps.billing.models import PointEau
    from apps.geography.models import Village, Fokontany

    points_eau_accessibles = PointEau.objects.filter(ouvrage__in=ouvrages_accessibles)

    communes_accessibles = Commune.objects.filter(
        fokontanys__villages__parcelles__pointeau__in=points_eau_accessibles
    ).distinct()

    districts_accessibles = District.objects.filter(
        communes__in=communes_accessibles
    ).distinct()

    regions_accessibles = Region.objects.filter(
        districts__in=districts_accessibles
    ).distinct()

    # Pour les filtres, n'afficher que les données accessibles
    regions = regions_accessibles
    districts = districts_accessibles
    communes = communes_accessibles

    # Organisations gestionnaires accessibles pour le filtre dashboard
    if request.user.is_superuser or is_visiteur(request.user):
        organisations = Organisation.objects.filter(
            type_organisation=Organisation.TYPE_GESTIONNAIRE
        ).order_by('nom_organisation')
    elif allowed_organisation_ids is None:
        organisations = Organisation.objects.filter(
            type_organisation=Organisation.TYPE_GESTIONNAIRE
        ).order_by('nom_organisation')
    elif allowed_organisation_ids:
        organisations = Organisation.objects.filter(
            id__in=allowed_organisation_ids,
            type_organisation=Organisation.TYPE_GESTIONNAIRE
        ).order_by('nom_organisation')
    else:
        organisations = Organisation.objects.none()

    # Générer la liste des années disponibles (depuis la première année d'ouvrage accessible jusqu'Ã  aujourd'hui)
    annee_min_data = ouvrages_accessibles.aggregate(
        min_realisation=Min('annee_realisation'),
        min_avant=Min('annee_realisation_avant')
    )
    annee_min = min(
        filter(None, [annee_min_data['min_realisation'], annee_min_data['min_avant']])
    ) if any([annee_min_data['min_realisation'], annee_min_data['min_avant']]) else datetime.now().year - 50

    annee_max = datetime.now().year
    annees = list(range(int(annee_min), annee_max + 1))
    annees.reverse()  # Du plus récent au plus ancien

    # Calculer les statistiques globales avec filtre organisation ET permissions
    stats = calculate_dashboard_stats_with_permissions(
        ouvrages_accessibles,
        organisation_id=organisation_filter
    )
    stats_par_commune = calculate_dashboard_stats_par_commune_with_permissions(
        ouvrages_accessibles,
        organisation_id=organisation_filter
    )

    context = {
        'regions': regions,
        'districts': districts,
        'communes': communes,
        'organisations': organisations,
        'annees': annees,
        'stats': stats,
        'stats_par_commune': stats_par_commune,
        'organisation_filter': organisation_filter,
        'title': 'Dashboard Principal'
    }

    return render(request, 'dashboard/main.html', context)


def calculate_dashboard_stats(region_id=None, district_id=None, commune_id=None, fokontany_id=None, village_id=None, annee=None, organisation_id=None):
    """Calculer les statistiques du dashboard avec filtres optionnels

    Args:
        annee: Année cumulative - affiche les ouvrages créés en cette année ou avant
        organisation_id: Filtre par organisation (via modalités actives)
    """
    from django.db.models import Case, When, IntegerField
    from apps.geography.models import Fokontany, Village

    # Filtres géographiques
    parcelles_qs = Parcelle.objects.all()
    # Compter uniquement les ouvrages qui ont au moins un point d'eau
    ouvrages_qs = Ouvrage.objects.filter(pointeau__isnull=False).distinct()
    points_eau_qs = PointEau.objects.all()


def calculate_dashboard_stats_with_permissions(ouvrages_accessibles, region_id=None, district_id=None, commune_id=None, fokontany_id=None, village_id=None, annee=None, organisation_id=None, return_scope=False, scope_only=False):
    """Calculer les statistiques du dashboard avec permissions et filtres optionnels

    Args:
        ouvrages_accessibles: QuerySet des ouvrages accessibles selon les permissions
        annee: Année cumulative - affiche les ouvrages créés en cette année ou avant
        organisation_id: Filtre par organisation (via modalités actives)
    """
    from django.db.models import Case, When, IntegerField
    from apps.geography.models import Fokontany, Village

    # Filtres géographiques
    parcelles_qs = Parcelle.objects.all()
    # Utiliser les ouvrages accessibles fournis + filtre points d'eau
    ouvrages_qs = ouvrages_accessibles.filter(pointeau__isnull=False).distinct()
    # Filtrer les points d'eau par les ouvrages accessibles
    ouvrage_ids_accessibles = ouvrages_accessibles.values_list('id', flat=True)
    points_eau_qs = PointEau.objects.filter(ouvrage_id__in=ouvrage_ids_accessibles)

    # Sans filtre organisation, conserver tous les ouvrages accessibles sans
    # distinction de modalite de gestion. Pour une organisation selectionnee,
    # une affectation future ne doit pas encore entrer dans ses compteurs.
    if organisation_id and organisation_id != 'all':
        ouvrages_qs = filter_ouvrages_by_gestionnaire_period(
            ouvrages_qs,
            organisation_id,
            annee,
            limit_current_year_to_today=True,
        )
        points_eau_qs = points_eau_qs.filter(
            ouvrage_id__in=ouvrages_qs.values_list('id', flat=True)
        )

    # Filtre par année (cumulatif) : ouvrages créés en année â‰¤ annee
    # On prend le minimum entre annee_realisation et annee_realisation_avant (en ignorant NULL)
    if annee:
        try:
            annee_int = int(annee)
            from django.db.models.functions import ExtractYear
            from datetime import date

            ouvrages_qs = ouvrages_qs.filter(
                Q(annee_realisation__lte=annee_int) |
                Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee_int)
            )
            # Points branchés en fin d'année ou avant (+ sans date, comptés à part). L'état de
            # coupure au 31/12 est déterminé par _build_dashboard_service_scope : un point coupé
            # ne peut plus être à la fois dans « Total » et dans « Coupés ».
            points_eau_qs = filtrer_points_annee(points_eau_qs, annee_int)
        except (ValueError, TypeError):
            pass  # Ignorer si annee n'est pas un entier valide

    # QuerySets pour les nouvelles tables
    equipements_qs = EquipementConstitutif.objects.all()
    compteurs_qs = CompteurTete.objects.all()
    adresses_qs = Adresse.objects.all()
    cartes_sociales_qs = CarteSociale.objects.all()
    tarifs_qs = TarifVolumetrique.objects.all()
    modalites_qs = ModaliteGestion.objects.all()
    fournisseurs_qs = Fournisseur.objects.filter(actif=True)

    modalites_qs = filter_modalites_by_gestionnaire_period(
        modalites_qs,
        organisation_id if organisation_id and organisation_id != 'all' else None,
        annee,
        limit_current_year_to_today=True,
    )

    if village_id:
        village = Village.objects.get(pk=village_id)
        # Parcelle -> Village
        parcelles_qs = parcelles_qs.filter(village=village)
        points_eau_qs = points_eau_qs.filter(parcelle__village=village)
        # Filtrer les ouvrages par village
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village=village).distinct()
        # Filtrer équipements, compteurs et tarifs via les ouvrages de la zone
        # Un ouvrage appartient géographiquement aux zones qu'il dessert via ses points d'eau
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements_qs = equipements_qs.filter(ouvrage_id__in=ouvrage_ids)
        compteurs_qs = compteurs_qs.filter(ouvrage_id__in=ouvrage_ids)
        tarifs_qs = tarifs_qs.filter(ouvrage_id__in=ouvrage_ids)

        # Filtrer cartes sociales par village
        cartes_sociales_qs = cartes_sociales_qs.filter(point_eau__parcelle__village=village)

        communes_count = 1
        districts_count = 1
    elif fokontany_id:
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        # Parcelle -> Village -> Fokontany
        parcelles_qs = parcelles_qs.filter(village__fokontany=fokontany)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany=fokontany)
        # Filtrer les ouvrages par fokontany
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany=fokontany).distinct()
        # Filtrer équipements, compteurs et tarifs via les ouvrages de la zone
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements_qs = equipements_qs.filter(ouvrage_id__in=ouvrage_ids)
        compteurs_qs = compteurs_qs.filter(ouvrage_id__in=ouvrage_ids)
        tarifs_qs = tarifs_qs.filter(ouvrage_id__in=ouvrage_ids)

        # Filtrer cartes sociales par fokontany
        cartes_sociales_qs = cartes_sociales_qs.filter(point_eau__parcelle__village__fokontany=fokontany)

        communes_count = 1
        districts_count = 1
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        # Parcelle -> Village -> Fokontany -> Commune
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune=commune)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune=commune)
        # Filtrer les ouvrages par commune
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune=commune).distinct()
        # Filtrer équipements, compteurs et tarifs via les ouvrages de la zone
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements_qs = equipements_qs.filter(ouvrage_id__in=ouvrage_ids)
        compteurs_qs = compteurs_qs.filter(ouvrage_id__in=ouvrage_ids)
        tarifs_qs = tarifs_qs.filter(ouvrage_id__in=ouvrage_ids)

        # Filtrer les nouvelles tables par commune via les relations
        # Adresses -> via ouvrage (pas de lien géographique direct)
        # Cartes sociales -> via point d'eau -> parcelle
        cartes_sociales_qs = cartes_sociales_qs.filter(point_eau__parcelle__village__fokontany__commune=commune)

        communes_count = 1
        districts_count = 1
    elif district_id:
        district = District.objects.get(pk=district_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district=district)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district=district)
        # Filtrer les ouvrages par district
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district=district).distinct()
        # Filtrer équipements, compteurs et tarifs via les ouvrages de la zone
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements_qs = equipements_qs.filter(ouvrage_id__in=ouvrage_ids)
        compteurs_qs = compteurs_qs.filter(ouvrage_id__in=ouvrage_ids)
        tarifs_qs = tarifs_qs.filter(ouvrage_id__in=ouvrage_ids)

        # Filtrer cartes sociales par district
        cartes_sociales_qs = cartes_sociales_qs.filter(point_eau__parcelle__village__fokontany__commune__district=district)

        communes_count = Commune.objects.filter(district=district).count()
        districts_count = 1
    elif region_id:
        region = Region.objects.get(pk=region_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district__region=region)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district__region=region)
        # Filtrer les ouvrages par région
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district__region=region).distinct()
        # Filtrer équipements, compteurs et tarifs via les ouvrages de la zone
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements_qs = equipements_qs.filter(ouvrage_id__in=ouvrage_ids)
        compteurs_qs = compteurs_qs.filter(ouvrage_id__in=ouvrage_ids)
        tarifs_qs = tarifs_qs.filter(ouvrage_id__in=ouvrage_ids)

        # Filtrer cartes sociales par région
        cartes_sociales_qs = cartes_sociales_qs.filter(point_eau__parcelle__village__fokontany__commune__district__region=region)

        communes_count = Commune.objects.filter(district__region=region).count()
        districts_count = District.objects.filter(region=region).count()
    else:
        communes_count = Commune.objects.count()
        districts_count = District.objects.count()

    # Si un filtre organisation est appliqué, le compteur communes doit suivre ce périmètre.
    # On compte les communes réellement couvertes par les points d'eau filtrés.
    if organisation_id and organisation_id != 'all':
        communes_count = points_eau_qs.filter(
            parcelle__village__fokontany__commune__isnull=False
        ).values(
            'parcelle__village__fokontany__commune_id'
        ).distinct().count()

    (
        ouvrages_qs_non_suspendu, points_eau_kpi_qs, points_eau_actifs_qs,
        coupes_ids_kpi, abandonnes_ids_kpi,
    ) = _build_dashboard_service_scope(
        points_eau_qs,
        ouvrages_qs,
        annee=annee,
        avec_coupures=not scope_only,
    )

    if scope_only:
        # Analyses dérivées : seul le périmètre est utile, on évite le calcul complet des stats
        return None, points_eau_kpi_qs, ouvrages_qs_non_suspendu

    # Calculer les statistiques
    # Privés = BP ou type « privé » (même règle que la répartition par modalité), publics = le reste.
    # Une seule requête pour le total, les privés et les publics.
    types_prives_ids = private_point_type_ids()
    comptes_actifs = points_eau_actifs_qs.order_by().aggregate(
        total=Count('id', distinct=True),
        prives=Count('id', distinct=True, filter=Q(type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False)),
    )
    points_eau_total_count = comptes_actifs['total']
    points_eau_prives_count = comptes_actifs['prives']
    points_eau_publics_count = points_eau_total_count - points_eau_prives_count

    # Points d'eau sans date de branchement : comptés à part, ce ne sont pas encore des points d'eau.
    # Coupés / abandonnés : historique CoupureEau (calculés avec le périmètre).
    points_eau_sans_date_count = points_eau_kpi_qs.filter(datebranchement__isnull=True).count()
    points_eau_coupes_count = len(coupes_ids_kpi)
    points_eau_abandonnes_count = len(abandonnes_ids_kpi)

    # Calcul des habitants
    total_habitants = parcelles_qs.aggregate(total=Sum('nombre_habitants'))['total'] or 0

    # Habitants desservis : habitants des parcelles liées aux points d'eau actifs
    # déjÃ  filtrés (permissions, géographie, année, organisation).
    parcelles_avec_points_eau = parcelles_qs.filter(
        pointeau__in=points_eau_actifs_qs
    ).distinct()
    habitants_desservis = parcelles_avec_points_eau.aggregate(total=Sum('nombre_habitants'))['total'] or 0

    # Habitants sur points d'eau privés : 4 personnes par point d'eau privé
    habitants_prives = points_eau_prives_count * 4

    # Habitants sur points d'eau publics : (habitants desservis - habitants privés) / nombre de points d'eau publics
    habitants_publics_total = max(0, habitants_desservis - habitants_prives)
    habitants_par_point_eau_public = habitants_publics_total / points_eau_publics_count if points_eau_publics_count > 0 else 0

    # Calculer la médiane du nombre d'habitants par point d'eau public
    # Pour cela, on calcule le ratio pour chaque parcelle desservie par un point d'eau public
    import statistics
    habitants_par_point_eau_public_mediane = 0
    if points_eau_publics_count > 0:
        # Habitants de la parcelle de chaque point d'eau public, en une requête
        # (l'id est gardé dans le SELECT pour que le DISTINCT ne fusionne pas deux points de même valeur)
        ratios = [
            habitants
            for _id, habitants in points_eau_actifs_qs.exclude(
                type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False
            ).filter(parcelle__nombre_habitants__gt=0).order_by().values_list(
                'id', 'parcelle__nombre_habitants'
            )
        ]

        if ratios:
            habitants_par_point_eau_public_mediane = statistics.median(ratios)

    # Comptage des villages
    from apps.geography.models import Village
    villages_qs = Village.objects.all()

    # Appliquer les filtres géographiques aux villages
    if commune_id:
        villages_qs = villages_qs.filter(fokontany__commune=commune_id)
    elif district_id:
        villages_qs = villages_qs.filter(fokontany__commune__district=district_id)
    elif region_id:
        villages_qs = villages_qs.filter(fokontany__commune__district__region=region_id)

    # Si un filtre organisation est applique, garder uniquement les villages
    # relies a au moins un point d'eau du perimetre filtre.
    if organisation_id and organisation_id != 'all':
        villages_qs = villages_qs.filter(
            parcelles__pointeau__in=points_eau_qs
        ).distinct()

    villages_count = villages_qs.count()

    # Villages desservis : villages reliés à un point d'eau actif dans le scope KPI dashboard.
    villages_desservis_count = villages_qs.filter(
        parcelles__pointeau__in=points_eau_actifs_qs
    ).distinct().count()

    stats = {
        'communes': communes_count,
        'districts': districts_count,
        'villages': villages_count,
        'villages_desservis': villages_desservis_count,
        'ouvrages': ouvrages_qs_non_suspendu.count(),
        'habitants': total_habitants,
        'habitants_desservis': habitants_desservis,  # Nouveau compteur
        'points_eau_total': points_eau_total_count,
        'points_eau_prives': points_eau_prives_count,
        'points_eau_publics': points_eau_publics_count,
        'points_eau_coupes': points_eau_coupes_count,
        'points_eau_abandonnes': points_eau_abandonnes_count,
        'points_eau_sans_date': points_eau_sans_date_count,

        # Nouveaux compteurs d'habitants
        'habitants_prives': habitants_prives,  # 4 * nombre de BP
        'habitants_par_point_eau_public': round(habitants_par_point_eau_public, 1),  # (Hab desservis - hab privés) / nb points d'eau publics
        'habitants_par_point_eau_public_mediane': round(habitants_par_point_eau_public_mediane, 1),  # Médiane du nombre d'habitants par point d'eau public
        'habitants_publics': habitants_publics_total,
        # Parts du total des habitants (pourcentages entiers, règle d'affichage du projet ;
        # « < 1 % » plutôt que « 0 % » pour une part non nulle)
        'pct_habitants_desservis': _pct_label(habitants_desservis, total_habitants),
        'pct_habitants_prives': _pct_label(habitants_prives, total_habitants),
        'pct_habitants_publics': _pct_label(habitants_publics_total, total_habitants),

        # Autres compteurs
        'equipements': equipements_qs.count(),
        'compteurs_tete': compteurs_qs.count(),
        'adresses': adresses_qs.count(),
        'cartes_sociales': cartes_sociales_qs.count(),
        'tarifs_volumetriques': tarifs_qs.filter(actif=True).count(),
        'modalites_gestion': modalites_qs.filter(statut='actif').count(),
        'fournisseurs': fournisseurs_qs.count(),
    }

    points_eau_modalites_qs = points_eau_actifs_qs
    payment_breakdown = build_point_usage_modality_breakdown(points_eau_modalites_qs, annee=annee)
    payment_scope = payment_breakdown['payment_scope']
    payment_breakdown_counts = payment_breakdown['counts']

    ouvrages_ids_forfaitaire = payment_scope['ouvrage_ids']['forfaitaire']
    ouvrages_ids_volume = payment_scope['ouvrage_ids']['au_volume']
    ouvrages_ids_mixte = payment_scope['ouvrage_ids']['mixte']
    ouvrages_ids_cotisation = payment_scope['ouvrage_ids']['cotisation_annuelle']

    stats['ouvrages_forfaitaire'] = len(ouvrages_ids_forfaitaire)
    stats['ouvrages_volume'] = len(ouvrages_ids_volume)
    stats['ouvrages_mixte'] = len(ouvrages_ids_mixte)
    stats['ouvrages_cotisation'] = len(ouvrages_ids_cotisation)

    stats['points_eau_forfaitaire'] = len(payment_scope['point_ids']['forfaitaire'])
    stats['points_eau_volume'] = len(payment_scope['point_ids']['au_volume'])
    stats['points_eau_mixte'] = len(payment_scope['point_ids']['mixte'])
    stats['points_eau_cotisation'] = len(payment_scope['point_ids']['cotisation_annuelle'])
    stats['points_eau_prives_au_volume'] = payment_breakdown_counts['prives_au_volume']
    stats['points_eau_prives_forfaitaire'] = payment_breakdown_counts['prives_forfaitaire']
    stats['points_eau_prives_sans_modalite'] = payment_breakdown_counts['prives_sans_modalite']
    stats['points_eau_publics_au_volume'] = payment_breakdown_counts['publics_au_volume']

    # Publics au volume : kiosques séparés par type d'abonné, BI volume, et le reste en anomalie
    # (un point public au volume ne peut être qu'un kiosque ou un BI volume)
    abonnes_publics_volume = {
        'particuliers': 0, 'boutiques': 0, 'kiosques_non_renseignes': 0, 'bi_volume': 0, 'anomalies': 0,
    }
    ouvrages_publics_volume = {cle: {} for cle in abonnes_publics_volume}
    publics_volume_ids = payment_breakdown['category_point_ids']['publics_au_volume']
    if publics_volume_ids:
        # L'id reste dans le SELECT : le périmètre est DISTINCT, sans lui deux kiosques
        # « particulier » seraient fusionnés en une seule ligne (sous-comptage)
        for _id, typeabonne, nom_type, est_inst, nom_ouvrage in points_eau_actifs_qs.filter(
            id__in=publics_volume_ids
        ).order_by().values_list(
            'id', 'typeabonne', 'type_point_eau__nom_type', 'estbranchementinstitutionnel', 'ouvrage__nom_ouvrage'
        ):
            abonne = (typeabonne or '').strip().lower()
            type_nom = (nom_type or '').strip().lower()
            if type_nom == 'bp' and est_inst:  # ancien type « BI volume »
                cle = 'bi_volume'
            elif type_nom != 'kiosque':
                cle = 'anomalies'
            elif abonne == 'particulier':
                cle = 'particuliers'
            elif abonne == 'boutique':
                cle = 'boutiques'
            else:
                cle = 'kiosques_non_renseignes'
            abonnes_publics_volume[cle] += 1
            nom = nom_ouvrage or 'Sans ouvrage'
            ouvrages_publics_volume[cle][nom] = ouvrages_publics_volume[cle].get(nom, 0) + 1
    for cle, valeur in abonnes_publics_volume.items():
        stats[f'points_eau_publics_au_volume_{cle}'] = valeur
    stats['points_eau_publics_au_volume_ouvrages'] = {
        cle: _top_ouvrages(compteur) for cle, compteur in ouvrages_publics_volume.items()
    }
    stats['points_eau_publics_forfaitaire'] = payment_breakdown_counts['publics_forfaitaire']

    # Publics forfaitaires : BF ordinaires séparés des BF institutionnels (ex-« BI »)
    publics_forfait_ids = payment_breakdown['category_point_ids']['publics_forfaitaire']
    forfait_publics = {'bf': 0, 'bf_institutionnel': 0, 'autres': 0}
    ouvrages_forfait_publics = {cle: {} for cle in forfait_publics}
    if publics_forfait_ids:
        # L'id reste dans le SELECT : le périmètre est DISTINCT (cf. publics au volume)
        for _id, nom_type, est_inst, nom_ouvrage in points_eau_actifs_qs.filter(
            id__in=publics_forfait_ids
        ).order_by().values_list(
            'id', 'type_point_eau__nom_type', 'estbranchementinstitutionnel', 'ouvrage__nom_ouvrage'
        ):
            type_nom = (nom_type or '').strip().lower()
            if est_inst and type_nom in ('bf', 'bp'):  # un BP institutionnel au forfait = BF institutionnel
                cle = 'bf_institutionnel'
            elif type_nom == 'bf':
                cle = 'bf'
            else:
                cle = 'autres'
            forfait_publics[cle] += 1
            nom = nom_ouvrage or 'Sans ouvrage'
            ouvrages_forfait_publics[cle][nom] = ouvrages_forfait_publics[cle].get(nom, 0) + 1
    for cle, valeur in forfait_publics.items():
        stats[f'points_eau_publics_forfaitaire_{cle}'] = valeur
    stats['points_eau_publics_forfaitaire_ouvrages'] = {
        cle: _top_ouvrages(compteur) for cle, compteur in ouvrages_forfait_publics.items()
    }

    # Autres tuiles : répartition par type (BF, BF institutionnel, BP, Kiosque…)
    categories_ids = payment_breakdown['category_point_ids']
    stats['types_prives_au_volume'] = _ventiler_par_type(points_eau_kpi_qs, categories_ids['prives_au_volume'])
    stats['types_prives_sans_modalite'] = _ventiler_par_type(points_eau_kpi_qs, categories_ids['prives_sans_modalite'])
    stats['types_publics_sans_modalite'] = _ventiler_par_type(points_eau_kpi_qs, categories_ids['publics_sans_modalite'])
    stats['types_coupes'] = _ventiler_par_type(points_eau_kpi_qs, coupes_ids_kpi)
    stats['types_abandonnes'] = _ventiler_par_type(points_eau_kpi_qs, abandonnes_ids_kpi)
    stats['types_sans_date'] = _ventiler_par_type(points_eau_kpi_qs.filter(datebranchement__isnull=True))
    stats['points_eau_publics_sans_modalite'] = payment_breakdown_counts['publics_sans_modalite']

    if return_scope:
        # Périmètre exact des KPI (ouvrages en service + points d'eau) pour les analyses dérivées
        return stats, points_eau_kpi_qs, ouvrages_qs_non_suspendu
    return stats


def calculate_dashboard_stats_par_commune_with_permissions(
    ouvrages_accessibles,
    region_id=None,
    district_id=None,
    commune_id=None,
    fokontany_id=None,
    village_id=None,
    annee=None,
    organisation_id=None
):
    """Calcule un sous-ensemble des KPIs du dashboard, agrégés par commune."""
    from django.db.models.functions import ExtractYear
    from apps.geography.models import Fokontany, Village

    parcelles_qs = Parcelle.objects.all()
    ouvrages_qs = ouvrages_accessibles.filter(pointeau__isnull=False).distinct()
    ouvrage_ids_accessibles = ouvrages_accessibles.values_list('id', flat=True)
    points_eau_qs = PointEau.objects.filter(ouvrage_id__in=ouvrage_ids_accessibles)
    villages_qs = Village.objects.all()

    if organisation_id and organisation_id != 'all':
        ouvrages_qs = filter_ouvrages_by_gestionnaire_period(
            ouvrages_qs,
            organisation_id,
            annee,
            limit_current_year_to_today=True,
        )
        points_eau_qs = points_eau_qs.filter(
            ouvrage_id__in=ouvrages_qs.values_list('id', flat=True)
        )

    if annee:
        try:
            annee_int = int(annee)
            ouvrages_qs = ouvrages_qs.filter(
                Q(annee_realisation__lte=annee_int) |
                Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee_int)
            )
            points_eau_qs = filtrer_points_annee(points_eau_qs, annee_int)
        except (TypeError, ValueError):
            pass

    if village_id:
        village = Village.objects.get(pk=village_id)
        parcelles_qs = parcelles_qs.filter(village=village)
        points_eau_qs = points_eau_qs.filter(parcelle__village=village)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village=village).distinct()
        villages_qs = villages_qs.filter(id=village.id)
    elif fokontany_id:
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany=fokontany)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany=fokontany)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany=fokontany).distinct()
        villages_qs = villages_qs.filter(fokontany=fokontany)
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune=commune)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune=commune)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune=commune).distinct()
        villages_qs = villages_qs.filter(fokontany__commune=commune)
    elif district_id:
        district = District.objects.get(pk=district_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district=district)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district=district)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district=district).distinct()
        villages_qs = villages_qs.filter(fokontany__commune__district=district)
    elif region_id:
        region = Region.objects.get(pk=region_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district__region=region)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district__region=region)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village__fokontany__commune__district__region=region
        ).distinct()
        villages_qs = villages_qs.filter(fokontany__commune__district__region=region)

    # Aligner le nombre de villages par commune avec le filtre organisation.
    if organisation_id and organisation_id != 'all':
        villages_qs = villages_qs.filter(
            parcelles__pointeau__in=points_eau_qs
        ).distinct()

    commune_ids = list(
        points_eau_qs.filter(
            parcelle__village__fokontany__commune__isnull=False
        ).values_list('parcelle__village__fokontany__commune_id', flat=True).distinct()
    )
    if not commune_ids:
        return []

    (
        ouvrages_qs_non_suspendu, points_eau_kpi_qs, points_eau_actifs_qs,
        points_eau_coupes_ids, points_eau_abandonnes_ids,
    ) = _build_dashboard_service_scope(points_eau_qs, ouvrages_qs, annee=annee)

    key_points = 'parcelle__village__fokontany__commune_id'
    key_villages = 'fokontany__commune_id'
    key_parcelles = 'village__fokontany__commune_id'

    def to_count_map(rows, key_field, value_field='total'):
        mapped = {}
        for row in rows:
            key = row.get(key_field)
            if key:
                mapped[str(key)] = int(row.get(value_field) or 0)
        return mapped

    def to_sum_map(rows, key_field, value_field='total'):
        mapped = {}
        for row in rows:
            key = row.get(key_field)
            if key:
                mapped[str(key)] = int(row.get(value_field) or 0)
        return mapped

    # Total / privés / publics par commune en une seule requête (privés : même règle que les KPI)
    types_prives_ids = private_point_type_ids()
    points_par_commune_rows = list(
        points_eau_actifs_qs.order_by().values(key_points).annotate(
            total=Count('id', distinct=True),
            prives=Count('id', distinct=True, filter=Q(type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False)),
        )
    )
    points_total_by_commune = to_count_map(points_par_commune_rows, key_points)
    points_prives_by_commune = to_count_map(points_par_commune_rows, key_points, 'prives')
    points_publics_by_commune = {
        key: total - points_prives_by_commune.get(key, 0)
        for key, total in points_total_by_commune.items()
    }
    points_coupes_by_commune = to_count_map(
        points_eau_kpi_qs.filter(id__in=points_eau_coupes_ids).values(key_points).annotate(total=Count('id')),
        key_points
    )
    points_abandonnes_by_commune = to_count_map(
        points_eau_kpi_qs.filter(id__in=points_eau_abandonnes_ids).values(key_points).annotate(total=Count('id')),
        key_points
    )

    villages_by_commune = to_count_map(
        villages_qs.order_by().values(key_villages).annotate(total=Count('id', distinct=True)),
        key_villages
    )
    villages_desservis_by_commune = to_count_map(
        villages_qs.filter(
            parcelles__pointeau__in=points_eau_actifs_qs
        ).order_by().values(key_villages).annotate(total=Count('id', distinct=True)),
        key_villages
    )

    habitants_by_commune = to_sum_map(
        parcelles_qs.order_by().values(key_parcelles).annotate(total=Sum('nombre_habitants')),
        key_parcelles
    )

    parcelles_desservies_ids = parcelles_qs.filter(
        pointeau__in=points_eau_actifs_qs
    ).values_list('id', flat=True).distinct()
    habitants_desservis_by_commune = to_sum_map(
        Parcelle.objects.filter(id__in=parcelles_desservies_ids).order_by().values(key_parcelles).annotate(total=Sum('nombre_habitants')),
        key_parcelles
    )

    # Compter les ouvrages par commune via les points d'eau actifs (date de branchement renseignée)
    # pour aligner le KPI avec l'activité réelle.
    ouvrages_by_commune = to_count_map(
        points_eau_actifs_qs.filter(
            parcelle__village__fokontany__commune__isnull=False
        ).values(key_points).annotate(
            total=Count('ouvrage_id', distinct=True)
        ),
        key_points
    )

    payment_scope = build_payment_modality_scope(
        points_eau_kpi_qs,
        annee=annee
    )
    ouvrages_forfaitaire_ids = payment_scope['ouvrage_ids']['forfaitaire']
    ouvrages_volume_ids = payment_scope['ouvrage_ids']['au_volume']
    ouvrages_mixte_ids = payment_scope['ouvrage_ids']['mixte']

    ouvrages_forfaitaire_by_commune = to_count_map(
        points_eau_kpi_qs.filter(
            ouvrage_id__in=ouvrages_forfaitaire_ids,
            parcelle__village__fokontany__commune__isnull=False
        ).values(key_points).annotate(
            total=Count('ouvrage_id', distinct=True)
        ),
        key_points
    )
    ouvrages_volume_by_commune = to_count_map(
        points_eau_kpi_qs.filter(
            ouvrage_id__in=ouvrages_volume_ids,
            parcelle__village__fokontany__commune__isnull=False
        ).values(key_points).annotate(
            total=Count('ouvrage_id', distinct=True)
        ),
        key_points
    )
    ouvrages_mixte_by_commune = to_count_map(
        points_eau_kpi_qs.filter(
            ouvrage_id__in=ouvrages_mixte_ids,
            parcelle__village__fokontany__commune__isnull=False
        ).values(key_points).annotate(
            total=Count('ouvrage_id', distinct=True)
        ),
        key_points
    )

    communes_qs = Commune.objects.filter(
        id__in=commune_ids
    ).select_related('district__region').order_by(
        'district__nom_district', 'nom_commune'
    )

    stats_par_commune = []
    for commune in communes_qs:
        commune_key = str(commune.id)
        points_prives = points_prives_by_commune.get(commune_key, 0)
        points_publics = points_publics_by_commune.get(commune_key, 0)
        habitants_desservis = habitants_desservis_by_commune.get(commune_key, 0)
        habitants_prives = points_prives * 4
        habitants_publics_total = max(0, habitants_desservis - habitants_prives)
        habitants_par_point_public = (
            round(habitants_publics_total / points_publics, 1)
            if points_publics > 0 else 0
        )

        stats_par_commune.append({
            'commune_id': str(commune.id),
            'commune': commune.nom_commune,
            'district': commune.district.nom_district if commune.district else '',
            'region': commune.district.region.nom_region if commune.district and commune.district.region else '',
            'villages': villages_by_commune.get(commune_key, 0),
            'villages_desservis': villages_desservis_by_commune.get(commune_key, 0),
            'ouvrages': ouvrages_by_commune.get(commune_key, 0),
            'ouvrages_forfaitaire': ouvrages_forfaitaire_by_commune.get(commune_key, 0),
            'ouvrages_volume': ouvrages_volume_by_commune.get(commune_key, 0),
            'ouvrages_mixte': ouvrages_mixte_by_commune.get(commune_key, 0),
            'habitants': habitants_by_commune.get(commune_key, 0),
            'habitants_desservis': habitants_desservis_by_commune.get(commune_key, 0),
            'points_eau_total': points_total_by_commune.get(commune_key, 0),
            'points_eau_prives': points_prives,
            'points_eau_publics': points_publics,
            'points_eau_coupes': points_coupes_by_commune.get(commune_key, 0),
            'points_eau_abandonnes': points_abandonnes_by_commune.get(commune_key, 0),
            'habitants_prives': habitants_prives,
            'habitants_par_point_eau_public': habitants_par_point_public,
        })

    return stats_par_commune


@employe_required
def dashboard_api(request):
    """API pour récupérer les statistiques du dashboard avec filtres et permissions"""
    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')
    organisation_id = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation_id')
    )
    annee = request.GET.get('annee')

    # Filtrer les données selon les permissions avant de calculer les stats
    ouvrages_accessibles = filter_queryset_by_permissions(Ouvrage.objects.all(), request.user)
    scoped_ouvrage_ids = apply_user_responsable_ouvrage_scope(
        ouvrages_accessibles.values_list('id', flat=True),
        request.user,
    )
    if scoped_ouvrage_ids is not None:
        ouvrages_accessibles = ouvrages_accessibles.filter(id__in=scoped_ouvrage_ids)

    # Calculer les stats avec filtre géographique
    stats = calculate_dashboard_stats_with_permissions(
        ouvrages_accessibles, region_id, district_id, commune_id,
        fokontany_id, village_id, annee, organisation_id
    )

    return JsonResponse(stats)


def _calculer_continuite(ouvrages_qs, points_qs, annee, mode):
    """Coupures / pannes par ouvrage (voir dashboard_continuite_api pour les règles). Retourne un dict."""
    from collections import defaultdict
    from datetime import datetime, time, timedelta
    import pytz
    from apps.planning.models import Planning

    # Fenêtre d'observation
    aujourd_hui = date.today()
    try:
        annee_int = int(annee) if annee else None
    except (TypeError, ValueError):
        annee_int = None
    debut_donnees = date(1900, 1, 1)
    if annee_int:
        fin = min(date(annee_int, 12, 31), aujourd_hui)
        debut = debut_donnees if mode == 'cumul' else date(annee_int, 1, 1)
    else:
        fin, debut = aujourd_hui, debut_donnees
    debut = min(debut, fin)

    def _en_date(valeur):
        return valeur.date() if hasattr(valeur, 'date') and callable(valeur.date) else valeur

    def _fmt(d):
        return d.strftime('%d/%m/%Y') if d else ''

    ouvrages = {
        str(o['id']): {
            'id': str(o['id']), 'nom': o['nom_ouvrage'] or '-',
            'type': o['type_ouvrage__nom_type_ouvrage'] or '',
            'nb_points': 0, 'nb_points_coupes': 0, 'nb_coupures': 0, 'derniere_coupure': None,
            'coupure_depuis_avant': False, 'nb_pannes': 0, 'derniere_panne': None,
        }
        for o in ouvrages_qs.values('id', 'nom_ouvrage', 'type_ouvrage__nom_type_ouvrage')
    }

    points = list(points_qs.values(
        'id', 'nomabonne', 'code_point_eau', 'idxls', 'ouvrage_id', 'datebranchement', 'datearretbranchement',
    ))
    evenements = defaultdict(list)
    for point_id, typ, dte in CoupureEau.objects.filter(
        point_eau_id__in=[p['id'] for p in points], date_evenement__isnull=False, date_evenement__lte=fin,
    ).order_by('point_eau_id', 'date_evenement', 'created_at').values_list('point_eau_id', 'type_evenement', 'date_evenement'):
        evenements[point_id].append((typ, dte))

    points_coupes = []
    for p in points:
        ouv = ouvrages.get(str(p['ouvrage_id']))
        if ouv is None:
            continue
        ouv['nb_points'] += 1
        evts = evenements.get(p['id'], [])
        coupe_au_debut, nb, derniere = False, 0, None
        source = 'evenements'
        detail_evts = []
        for typ, dte in evts:
            detail_evts.append({
                'type': typ, 'date': _fmt(dte), 'avant_periode': dte < debut,
                'compte': dte >= debut and typ == 'coupure',
            })
            if dte < debut:
                coupe_au_debut = (typ == 'coupure')
            elif typ == 'coupure':
                nb += 1
                derniere = dte
        if not evts:
            arret = _en_date(p['datearretbranchement'])
            if arret and p['datebranchement'] and arret <= fin:
                nb, derniere = 1, arret
                source = 'legacy'
                detail_evts.append({'type': 'coupure', 'date': _fmt(arret), 'avant_periode': False, 'compte': True})
        if coupe_au_debut or nb:
            ouv['nb_points_coupes'] += 1
            ouv['nb_coupures'] += nb
            ouv['coupure_depuis_avant'] = ouv['coupure_depuis_avant'] or coupe_au_debut
            if derniere and (ouv['derniere_coupure'] is None or derniere > ouv['derniere_coupure']):
                ouv['derniere_coupure'] = derniere
            points_coupes.append({
                'id': str(p['id']), 'nom': p['nomabonne'] or p['code_point_eau'] or '-',
                'idxls': str(p['idxls'] or ''), 'ouvrage': ouv['nom'], 'ouvrage_id': ouv['id'],
                'nb_coupures': nb, 'derniere_coupure': _fmt(derniere), 'depuis_avant': coupe_au_debut,
                'source': source, 'evenements': detail_evts,
                'date_arret_branchement': _fmt(_en_date(p['datearretbranchement'])),
            })

    tz = pytz.timezone('Indian/Antananarivo')
    debut_dt = tz.localize(datetime.combine(debut, time.min))
    fin_dt = tz.localize(datetime.combine(fin + timedelta(days=1), time.min))
    pannes = Planning.objects.filter(
        ouvrage_id__in=list(ouvrages.keys()), date__gte=debut_dt, date__lt=fin_dt,
    ).exclude(statut='annule').filter(
        Q(fiche_poste__tache__icontains='panne') | Q(type_activite__nomtype__icontains='panne')
    ).order_by('date').values_list('ouvrage_id', 'date', 'fiche_poste__tache', 'type_activite__nomtype')
    for ouvrage_id, dt, tache, type_act in pannes:
        ouv = ouvrages.get(str(ouvrage_id))
        if ouv is None:
            continue
        jour = dt.astimezone(tz).date()
        ouv.setdefault('pannes_detail', []).append({
            'date': _fmt(jour), 'activite': tache or type_act or '-',
        })
        ouv['nb_pannes'] += 1
        if ouv['derniere_panne'] is None or jour > ouv['derniere_panne']:
            ouv['derniere_panne'] = jour

    lignes = []
    for ouv in sorted(ouvrages.values(), key=lambda o: o['nom'].lower()):
        ouv.setdefault('pannes_detail', [])
        lignes.append({
            **ouv,
            'coupure': ouv['nb_points_coupes'] > 0,
            'panne': ouv['nb_pannes'] > 0,
            'derniere_coupure': _fmt(ouv['derniere_coupure']),
            'derniere_panne': _fmt(ouv['derniere_panne']),
        })
    return {
        'periode': {'mode': mode if annee_int else 'cumul', 'debut': _fmt(debut) if debut > debut_donnees else '', 'fin': _fmt(fin)},
        'lignes': lignes,
        'points_coupes': points_coupes,
        'nb_points': len(points),
    }


@employe_required
def dashboard_continuite_api(request):
    """
    Continuité de service : ouvrages / points d'eau avec ou sans coupure d'eau et panne.

    Périmètre = celui des KPI du dashboard (mêmes filtres, filtre année cumulatif : ouvrages
    existants à la fin de l'année choisie). Période observée :
      - periode=annee (défaut) : l'année civile choisie (jusqu'à aujourd'hui si c'est l'année en cours)
      - periode=cumul          : depuis le début des données jusqu'à la fin de l'année choisie
    Coupure : événement CoupureEau 'coupure' dans la période, ou point déjà coupé au début de la
    période et pas encore remis en eau ; repli legacy sur datearretbranchement.
    Panne : planning non annulé de la période, rattaché à l'ouvrage, dont l'activité
    (fiche de poste ou type d'activité) contient « panne ».
    """
    from collections import defaultdict
    from datetime import datetime, time, timedelta
    import pytz
    from apps.planning.models import Planning

    organisation_id = _resolve_dashboard_organisation_filter(request.user, request.GET.get('organisation_id'))
    annee = request.GET.get('annee')
    mode = request.GET.get('periode', 'annee')

    ouvrages_accessibles = filter_queryset_by_permissions(Ouvrage.objects.all(), request.user)
    scoped_ouvrage_ids = apply_user_responsable_ouvrage_scope(
        ouvrages_accessibles.values_list('id', flat=True), request.user,
    )
    if scoped_ouvrage_ids is not None:
        ouvrages_accessibles = ouvrages_accessibles.filter(id__in=scoped_ouvrage_ids)

    _stats, points_qs, ouvrages_qs = calculate_dashboard_stats_with_permissions(
        ouvrages_accessibles, request.GET.get('region_id'), request.GET.get('district_id'),
        request.GET.get('commune_id'), request.GET.get('fokontany_id'), request.GET.get('village_id'),
        annee, organisation_id, scope_only=True,
    )

    res = _calculer_continuite(ouvrages_qs, points_qs, annee, mode)
    lignes, points_coupes = res['lignes'], res['points_coupes']
    total = len(lignes)
    avec_coupure = sum(1 for o in lignes if o['coupure'])
    avec_panne = sum(1 for o in lignes if o['panne'])
    avec_les_deux = sum(1 for o in lignes if o['coupure'] and o['panne'])
    return JsonResponse({
        'periode': res['periode'],
        'compteurs': {
            'total_ouvrages': total,
            'sans_coupure': total - avec_coupure,
            'sans_panne': total - avec_panne,
            'sans_coupure_ni_panne': total - (avec_coupure + avec_panne - avec_les_deux),
            'avec_coupure': avec_coupure,
            'avec_panne': avec_panne,
            'avec_coupure_ou_panne': avec_coupure + avec_panne - avec_les_deux,
            'points_eau_coupes': len(points_coupes),
            'points_eau_total': res['nb_points'],
        },
        'ouvrages': lignes,
        'points_coupes': sorted(points_coupes, key=lambda x: (x['ouvrage'].lower(), x['nom'].lower())),
    })


@employe_required
def dashboard_stats_par_commune_api(request):
    """API dédiée au tableau stats par commune pour éviter de bloquer les KPIs globaux."""
    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')
    organisation_id = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation_id')
    )
    annee = request.GET.get('annee')

    ouvrages_accessibles = filter_queryset_by_permissions(Ouvrage.objects.all(), request.user)
    scoped_ouvrage_ids = apply_user_responsable_ouvrage_scope(
        ouvrages_accessibles.values_list('id', flat=True),
        request.user,
    )
    if scoped_ouvrage_ids is not None:
        ouvrages_accessibles = ouvrages_accessibles.filter(id__in=scoped_ouvrage_ids)

    stats_par_commune = calculate_dashboard_stats_par_commune_with_permissions(
        ouvrages_accessibles,
        region_id,
        district_id,
        commune_id,
        fokontany_id,
        village_id,
        annee,
        organisation_id
    )

    return JsonResponse({
        'stats_par_commune': stats_par_commune
    })


@employe_required
def dashboard_export_stats_par_commune_excel(request):
    """Export Excel des statistiques par commune du dashboard principal."""
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')
    organisation_id = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation_id')
    )
    annee = request.GET.get('annee')

    ouvrages_accessibles = filter_queryset_by_permissions(Ouvrage.objects.all(), request.user)

    stats_globaux = calculate_dashboard_stats_with_permissions(
        ouvrages_accessibles,
        region_id,
        district_id,
        commune_id,
        fokontany_id,
        village_id,
        annee,
        organisation_id
    )
    rows = calculate_dashboard_stats_par_commune_with_permissions(
        ouvrages_accessibles,
        region_id,
        district_id,
        commune_id,
        fokontany_id,
        village_id,
        annee,
        organisation_id
    )

    rows_sorted = sorted(
        rows,
        key=lambda r: (str(r.get('district') or ''), str(r.get('commune') or ''))
    )
    grouped_by_district = {}
    for row in rows_sorted:
        district_name = row.get('district') or 'Sans district'
        grouped_by_district.setdefault(district_name, []).append(row)

    numeric_fields = [
        'villages',
        'villages_desservis',
        'ouvrages',
        'ouvrages_forfaitaire',
        'ouvrages_volume',
        'ouvrages_mixte',
        'habitants',
        'habitants_desservis',
        'points_eau_total',
        'points_eau_prives',
        'points_eau_publics',
        'points_eau_coupes'
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Stats communes"

    ws.append(["Dashboard principal - Statistiques par commune"])
    ws.append([f"Export du {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])

    header = [
        "Type ligne",
        "District",
        "Commune",
        "Villages",
        "Villages desservis",
        "Ouvrages",
        "Ouvrages forfaitaire",
        "Ouvrages au volume",
        "Ouvrages mixte",
        "Habitants",
        "Habitants desservis",
        "Points d'eau total",
        "Points d'eau privés",
        "Points d'eau publics",
        "Points d'eau coupés",
    ]
    ws.append(header)
    header_row = ws.max_row

    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True, color="FFFFFF")
    strong_font = Font(bold=True)
    title_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    district_fill = PatternFill(start_color="E8F1FB", end_color="E8F1FB", fill_type="solid")
    subtotal_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A2"].font = Font(italic=True, size=10)

    for col in range(1, len(header) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    grand_totals = {field: 0 for field in numeric_fields}

    for district_name, district_rows in grouped_by_district.items():
        ws.append(["District", district_name] + [""] * (len(header) - 2))
        district_header_row = ws.max_row
        for col in range(1, len(header) + 1):
            cell = ws.cell(row=district_header_row, column=col)
            cell.fill = district_fill
            cell.font = strong_font
        ws.cell(row=district_header_row, column=2).alignment = Alignment(horizontal='left')

        district_totals = {field: 0 for field in numeric_fields}
        for row in district_rows:
            values = [int(row.get(field) or 0) for field in numeric_fields]
            ws.append([
                "Commune",
                district_name,
                row.get('commune') or '',
                *values
            ])
            for field, value in zip(numeric_fields, values):
                district_totals[field] += value
                grand_totals[field] += value

        district_totals_values = [district_totals[field] for field in numeric_fields]
        ws.append([
            "Sous-total",
            district_name,
            "",
            *district_totals_values,
        ])
        subtotal_row = ws.max_row
        for col in range(1, len(header) + 1):
            cell = ws.cell(row=subtotal_row, column=col)
            cell.fill = subtotal_fill
            cell.font = strong_font

    # Aligner le grand total des ouvrages sur le KPI global (mÃªme logique que l'UI)
    grand_totals['ouvrages'] = int(stats_globaux.get('ouvrages') or 0)

    grand_totals_values = [grand_totals[field] for field in numeric_fields]
    ws.append([
        "Grand total",
        "",
        "",
        *grand_totals_values,
    ])
    grand_total_row = ws.max_row
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=grand_total_row, column=col)
        cell.fill = total_fill
        cell.font = strong_font

    # Alignements + format numérique
    for row_idx in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal='left')
        ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal='left')
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal='left')
        for col_idx in range(4, 4 + len(numeric_fields)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')

    column_widths = {
        'A': 14,
        'B': 26,
        'C': 24,
        'D': 10,
        'E': 18,
        'F': 10,
        'G': 16,
        'H': 16,
        'I': 16,
        'J': 14,
        'K': 18,
        'L': 16,
        'M': 17,
        'N': 18,
        'O': 18,
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(header))}{header_row}"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"dashboard_stats_communes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def get_districts_by_region(request):
    """API pour récupérer les districts d'une région"""
    region_id = request.GET.get('region_id')

    if region_id:
        districts = District.objects.filter(region_id=region_id).values('id', 'nom_district')
        return JsonResponse({'districts': list(districts)})

    return JsonResponse({'districts': []})


def get_communes_by_district(request):
    """API pour récupérer les communes d'un district"""
    district_id = request.GET.get('district_id')

    if district_id:
        communes = Commune.objects.filter(district_id=district_id).values('id', 'nom_commune')
        return JsonResponse({'communes': list(communes)})

    return JsonResponse({'communes': []})


def get_fokontany_by_commune(request):
    """API pour récupérer les fokontany d'une commune"""
    from apps.geography.models import Fokontany
    commune_id = request.GET.get('commune_id')

    if commune_id:
        fokontany = Fokontany.objects.filter(commune_id=commune_id).values('id', 'nom_fokontany')
        return JsonResponse({'fokontany': list(fokontany)})

    return JsonResponse({'fokontany': []})


def get_villages_by_fokontany(request):
    """API pour récupérer les villages d'un fokontany"""
    from apps.geography.models import Village
    fokontany_id = request.GET.get('fokontany_id')

    if fokontany_id:
        villages = Village.objects.filter(fokontany_id=fokontany_id).values('id', 'nom_village')
        return JsonResponse({'villages': list(villages)})

    return JsonResponse({'villages': []})


@employe_required
def ouvrages_par_annee_api(request):
    """API pour récupérer le nombre cumulé d'ouvrages par année avec filtres géographiques"""
    from django.db.models.functions import Coalesce
    from apps.geography.models import Fokontany, Village

    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')

    # Filtrer les ouvrages selon les permissions
    ouvrages_qs = filter_queryset_by_permissions(
        Ouvrage.objects.filter(pointeau__isnull=False).distinct(),
        request.user
    )

    # Appliquer les filtres géographiques
    if village_id:
        village = Village.objects.get(pk=village_id)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village=village
        ).distinct()
    elif fokontany_id:
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village__fokontany=fokontany
        ).distinct()
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village__fokontany__commune=commune
        ).distinct()
    elif district_id:
        district = District.objects.get(pk=district_id)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village__fokontany__commune__district=district
        ).distinct()
    elif region_id:
        region = Region.objects.get(pk=region_id)
        ouvrages_qs = ouvrages_qs.filter(
            pointeau__parcelle__village__fokontany__commune__district__region=region
        ).distinct()

    # Récupérer toutes les années (min entre annee_realisation et annee_realisation_avant)
    ouvrages_avec_annee = ouvrages_qs.annotate(
        annee=Coalesce('annee_realisation', 'annee_realisation_avant')
    ).filter(annee__isnull=False).values_list('annee', flat=True)

    annees = sorted(set(ouvrages_avec_annee))

    # Calculer le nombre cumulé d'ouvrages par année
    data = {'labels': [], 'values': []}

    for annee in annees:
        count = ouvrages_qs.filter(
            Q(annee_realisation__lte=annee) |
            Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee)
        ).count()

        data['labels'].append(str(annee))
        data['values'].append(count)

    return JsonResponse(data)


@employe_required
def points_eau_par_annee_api(request):
    """API pour récupérer le nombre cumulé de points d'eau par année avec filtres géographiques"""
    from django.db.models.functions import ExtractYear
    from apps.geography.models import Fokontany, Village

    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')

    # Filtrer les points d'eau selon les permissions
    points_eau_qs = filter_queryset_by_permissions(
        PointEau.objects.filter(datebranchement__isnull=False),
        request.user
    )

    # Appliquer les filtres géographiques
    if village_id:
        village = Village.objects.get(pk=village_id)
        points_eau_qs = points_eau_qs.filter(parcelle__village=village)
    elif fokontany_id:
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany=fokontany)
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune=commune)
    elif district_id:
        district = District.objects.get(pk=district_id)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district=district)
    elif region_id:
        region = Region.objects.get(pk=region_id)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district__region=region)

    # Extraire les années de datebranchement
    annees = points_eau_qs.annotate(
        annee=ExtractYear('datebranchement')
    ).values_list('annee', flat=True).distinct()

    annees = sorted(set(filter(None, annees)))

    # Calculer le nombre cumulé de points d'eau par année
    data = {'labels': [], 'values': []}

    for annee in annees:
        count = points_eau_qs.annotate(
            annee_branchement=ExtractYear('datebranchement')
        ).filter(annee_branchement__lte=annee).count()

        data['labels'].append(str(annee))
        data['values'].append(count)

    return JsonResponse(data)


@employe_required
def dashboard_details_api(request, detail_type):
    """API pour récupérer les détails d'une statistique du dashboard"""
    region_id = request.GET.get('region_id')
    district_id = request.GET.get('district_id')
    commune_id = request.GET.get('commune_id')
    fokontany_id = request.GET.get('fokontany_id')
    village_id = request.GET.get('village_id')
    organisation_id = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation_id')
    )
    annee = request.GET.get('annee')

    # Filtres géographiques de base avec permissions
    parcelles_qs = filter_queryset_by_permissions(Parcelle.objects.all(), request.user)
    # Compter uniquement les ouvrages qui ont au moins un point d'eau
    ouvrages_qs = filter_queryset_by_permissions(
        Ouvrage.objects.filter(pointeau__isnull=False).distinct(),
        request.user
    )
    points_eau_qs = filter_queryset_by_permissions(PointEau.objects.all(), request.user)
    points_eau_actifs_qs = filter_queryset_by_permissions(
        PointEau.objects.filter(datebranchement__isnull=False),
        request.user
    )
    communes_qs = Commune.objects.all()

    # Filtre géographique avec priorité : village > fokontany > commune > district > region
    if village_id:
        from apps.geography.models import Village
        village = Village.objects.get(pk=village_id)
        parcelles_qs = parcelles_qs.filter(village=village)
        points_eau_qs = points_eau_qs.filter(parcelle__village=village)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(parcelle__village=village)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village=village).distinct()
        communes_qs = communes_qs.filter(id=village.fokontany.commune.id)
    elif fokontany_id:
        from apps.geography.models import Fokontany
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany=fokontany)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany=fokontany)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(parcelle__village__fokontany=fokontany)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany=fokontany).distinct()
        communes_qs = communes_qs.filter(id=fokontany.commune.id)
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune=commune)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune=commune)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(parcelle__village__fokontany__commune=commune)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune=commune).distinct()
        communes_qs = communes_qs.filter(id=commune_id)
    elif district_id:
        district = District.objects.get(pk=district_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district=district)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district=district)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(parcelle__village__fokontany__commune__district=district)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district=district).distinct()
        communes_qs = communes_qs.filter(district=district)
    elif region_id:
        region = Region.objects.get(pk=region_id)
        parcelles_qs = parcelles_qs.filter(village__fokontany__commune__district__region=region)
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district__region=region)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(parcelle__village__fokontany__commune__district__region=region)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district__region=region).distinct()
        communes_qs = communes_qs.filter(district__region=region)

    # Meme perimetre temporel que les compteurs principaux lorsqu'une
    # organisation est choisie. "Toutes organisations" reste un total brut.
    if organisation_id and organisation_id != 'all':
        ouvrages_qs = filter_ouvrages_by_gestionnaire_period(
            ouvrages_qs,
            organisation_id,
            annee,
            limit_current_year_to_today=True,
        )
        ouvrage_ids_org = ouvrages_qs.values_list('id', flat=True)
        points_eau_qs = points_eau_qs.filter(ouvrage_id__in=ouvrage_ids_org)
        points_eau_actifs_qs = points_eau_actifs_qs.filter(ouvrage_id__in=ouvrage_ids_org)

    # Exclure les ouvrages de type "suspendu" et "potentiel" pour la cohérence des stats
    ouvrages_qs = ouvrages_qs.exclude(
        Q(type_ouvrage__nom_type_ouvrage__icontains='suspend') |
        Q(type_ouvrage__nom_type_ouvrage__icontains='potentiel')
    )
    points_eau_qs = points_eau_qs.filter(ouvrage__in=ouvrages_qs)

    # Mêmes règles que les tuiles : actif = branché, ni coupé ni abandonné, ouvrage en service ;
    # privé = BP ou type « privé ».
    points_eau_actifs_qs = points_eau_actifs_qs.filter(ouvrage__in=ouvrages_qs)
    if detail_type != 'points-eau-coupes':
        coupes_detail_ids, abandonnes_detail_ids = _classer_coupes_point_ids(points_eau_qs, annee=annee)
        if coupes_detail_ids or abandonnes_detail_ids:
            points_eau_actifs_qs = points_eau_actifs_qs.exclude(id__in=coupes_detail_ids | abandonnes_detail_ids)
    types_prives_ids = private_point_type_ids()

    if detail_type == 'communes':
        # Détails des communes avec nombre d'ouvrages et points d'eau par commune
        data = []
        chart_data = {'labels': [], 'ouvrages': [], 'points_eau': []}

        for commune in communes_qs[:10]:  # Limiter Ã  10 pour l'affichage
            # Filtrer ouvrages par commune via les points d'eau
            ouvrages_count = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune=commune).distinct().count()
            points_eau_count = points_eau_actifs_qs.filter(parcelle__village__fokontany__commune=commune).count()

            data.append({
                'nom': commune.nom_commune,
                'district': commune.district.nom_district,
                'ouvrages': ouvrages_count,
                'points_eau': points_eau_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['ouvrages'].append(ouvrages_count)
            chart_data['points_eau'].append(points_eau_count)

    elif detail_type == 'ouvrages':
        # Détails des ouvrages avec leurs points d'eau
        data = []
        chart_data = {'labels': [], 'points_eau': []}

        for ouvrage in ouvrages_qs:
            points_eau_ouvrage = points_eau_actifs_qs.filter(ouvrage=ouvrage)
            points_eau_count = points_eau_ouvrage.count()
            communes = list(
                points_eau_ouvrage
                .values_list('parcelle__village__fokontany__commune__nom_commune', flat=True)
                .exclude(parcelle__village__fokontany__commune__nom_commune__isnull=True)
                .distinct()
            )
            if not communes:
                commune_label = 'N/A'
            elif len(communes) == 1:
                commune_label = communes[0]
            else:
                commune_label = f"{communes[0]} (+{len(communes)-1})"

            data.append({
                'nom': ouvrage.nom_ouvrage,
                'type': str(ouvrage.type_ouvrage),
                'commune': commune_label,
                'points_eau': points_eau_count
            })

            chart_data['labels'].append(ouvrage.nom_ouvrage[:20] + '...' if len(ouvrage.nom_ouvrage) > 20 else ouvrage.nom_ouvrage)
            chart_data['points_eau'].append(points_eau_count)

    elif detail_type == 'habitants':
        # Répartition des habitants par commune
        data = []
        chart_data = {'labels': [], 'habitants': []}

        for commune in communes_qs[:10]:
            habitants_count = parcelles_qs.filter(village__fokontany__commune=commune).aggregate(
                total=Sum('nombre_habitants')
            )['total'] or 0

            data.append({
                'commune': commune.nom_commune,
                'district': commune.district.nom_district,
                'habitants': habitants_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['habitants'].append(habitants_count)

    elif detail_type == 'points-eau':
        # Détails des points d'eau par type
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Grouper par type de point d'eau
        from django.db.models import Value, CharField
        points_eau_by_type = points_eau_qs.values('type_point_eau__nom_type').annotate(
            count=Count('id'),
            type_name=Value('type_point_eau__nom_type', output_field=CharField())
        ).order_by('-count')[:10]

        for item in points_eau_by_type:
            type_name = item['type_point_eau__nom_type'] or 'Type non défini'
            data.append({
                'type': type_name,
                'count': item['count']
            })

            chart_data['labels'].append(type_name)
            chart_data['counts'].append(item['count'])

    elif detail_type == 'points-eau-prives':
        # Points d'eau privés (type BP) par commune
        data = []
        chart_data = {'labels': [], 'points_eau_prives': []}

        for commune in communes_qs[:10]:
            prives_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune,
                type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False
            ).count()

            data.append({
                'commune': commune.nom_commune,
                'points_eau_prives': prives_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['points_eau_prives'].append(prives_count)

    elif detail_type == 'points-eau-publics':
        # Points d'eau publics (tous sauf BP) par commune
        data = []
        chart_data = {'labels': [], 'points_eau_publics': []}

        for commune in communes_qs[:10]:
            publics_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune
            ).exclude(
                type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False
            ).count()

            data.append({
                'commune': commune.nom_commune,
                'points_eau_publics': publics_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['points_eau_publics'].append(publics_count)

    elif detail_type == 'points-eau-coupes':
        # Points d'eau coupés par commune
        data = []
        chart_data = {'labels': [], 'points_eau_coupes': []}
        annee = request.GET.get('annee')
        active_coupes_ids = _active_coupes_point_ids(points_eau_qs, annee=annee)

        for commune in communes_qs[:10]:
            coupes_count = points_eau_qs.filter(
                id__in=active_coupes_ids,
                parcelle__village__fokontany__commune=commune,
                datebranchement__isnull=False
            ).count()

            data.append({
                'commune': commune.nom_commune,
                'points_eau_coupes': coupes_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['points_eau_coupes'].append(coupes_count)

    elif detail_type == 'equipements':
        # Détails des équipements par ouvrage
        from apps.infrastructure.models import EquipementConstitutif
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Filtrer les équipements par les ouvrages géographiquement filtrés
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        equipements = EquipementConstitutif.objects.filter(ouvrage_id__in=ouvrage_ids).select_related('ouvrage')[:20]

        for equipement in equipements:
            data.append({
                'ouvrage': equipement.ouvrage.nom_ouvrage if equipement.ouvrage else 'N/A',
                'equipement': equipement.nom or 'Non spécifié',
                'quantite': 1  # Removed quantite field that doesn't exist
            })

        # Compter par type d'équipement
        equipement_counts = EquipementConstitutif.objects.filter(ouvrage_id__in=ouvrage_ids).values('nom').annotate(
            count=Count('id')
        ).order_by('-count')[:10]

        for item in equipement_counts:
            chart_data['labels'].append(item['nom'] or 'Non spécifié')
            chart_data['counts'].append(item['count'])

    elif detail_type == 'compteurs':
        # Détails des compteurs de tÃªte
        from apps.billing.models import CompteurTete
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Filtrer les compteurs par les ouvrages géographiquement filtrés
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        compteurs = CompteurTete.objects.filter(ouvrage_id__in=ouvrage_ids).select_related('ouvrage')[:20]

        for compteur in compteurs:
            data.append({
                'numero': compteur.numerocompteur or 'N/A',
                'ouvrage': compteur.ouvrage.nom_ouvrage if compteur.ouvrage else 'N/A',
                'marque': compteur.marque or 'N/A'
            })

        # Compter par ouvrage
        compteur_counts = CompteurTete.objects.filter(ouvrage_id__in=ouvrage_ids, ouvrage__isnull=False).values(
            'ouvrage__nom_ouvrage'
        ).annotate(count=Count('id')).order_by('-count')[:10]

        for item in compteur_counts:
            chart_data['labels'].append(item['ouvrage__nom_ouvrage'])
            chart_data['counts'].append(item['count'])

    elif detail_type == 'habitants-desservis':
        # Habitants desservis par commune
        data = []
        chart_data = {'labels': [], 'habitants': []}

        for commune in communes_qs[:10]:
            habitants_desservis = parcelles_qs.filter(
                village__fokontany__commune=commune,
                pointeau__in=points_eau_actifs_qs
            ).distinct().aggregate(total=Sum('nombre_habitants'))['total'] or 0

            points_eau_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune
            ).count()

            data.append({
                'commune': commune.nom_commune,
                'habitants_desservis': habitants_desservis,
                'points_eau': points_eau_count
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['habitants'].append(habitants_desservis)

    elif detail_type == 'habitants-prives':
        # Habitants sur points d'eau privés par commune
        data = []
        chart_data = {'labels': [], 'habitants': []}

        for commune in communes_qs[:10]:
            prives_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune,
                type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False
            ).count()
            habitants_prives = prives_count * 4

            data.append({
                'commune': commune.nom_commune,
                'points_eau': prives_count,
                'habitants': habitants_prives
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['habitants'].append(habitants_prives)

    elif detail_type == 'habitants-publics':
        # Habitants par point d'eau public par commune
        data = []
        chart_data = {'labels': [], 'ratios': []}

        for commune in communes_qs[:10]:
            # Habitants desservis de la commune (parcelles avec points d'eau)
            habitants_desservis = parcelles_qs.filter(
                village__fokontany__commune=commune,
                pointeau__in=points_eau_actifs_qs
            ).distinct().aggregate(total=Sum('nombre_habitants'))['total'] or 0

            prives_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune,
                type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False
            ).count()

            publics_count = points_eau_actifs_qs.filter(
                parcelle__village__fokontany__commune=commune
            ).exclude(type_point_eau_id__in=types_prives_ids, estbranchementinstitutionnel=False).count()

            habitants_prives = prives_count * 4
            habitants_publics_total = max(0, habitants_desservis - habitants_prives)
            ratio = habitants_publics_total / publics_count if publics_count > 0 else 0

            data.append({
                'commune': commune.nom_commune,
                'habitants_desservis': habitants_desservis,
                'habitants_publics': habitants_publics_total,
                'points_eau_publics': publics_count,
                'ratio': round(ratio, 1)
            })

            chart_data['labels'].append(commune.nom_commune)
            chart_data['ratios'].append(round(ratio, 1))

    elif detail_type == 'tarifs':
        # Détails des tarifs volumétriques actifs
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Filtrer les tarifs par les ouvrages géographiquement filtrés
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        tarifs = TarifVolumetrique.objects.filter(actif=True, ouvrage_id__in=ouvrage_ids).select_related('ouvrage')[:20]

        for tarif in tarifs:
            data.append({
                'ouvrage': tarif.ouvrage.nom_ouvrage if tarif.ouvrage else 'N/A',
                'date_debut': tarif.datedebuttarif.strftime('%d/%m/%Y') if tarif.datedebuttarif else 'N/A',
                'prix_abonnement': tarif.prixabonnement
            })

        # Compter par ouvrage
        tarif_counts = TarifVolumetrique.objects.filter(actif=True, ouvrage_id__in=ouvrage_ids, ouvrage__isnull=False).values(
            'ouvrage__nom_ouvrage'
        ).annotate(count=Count('id')).order_by('-count')[:10]

        for item in tarif_counts:
            chart_data['labels'].append(item['ouvrage__nom_ouvrage'])
            chart_data['counts'].append(item['count'])

    elif detail_type == 'modalites':
        # Détails des modalités de gestion actives
        from apps.billing.models import ModaliteGestion
        data = []
        chart_data = {'labels': [], 'counts': []}

        modalites = filter_modalites_by_gestionnaire_period(
            ModaliteGestion.objects.all(),
            organisation_id if organisation_id and organisation_id != 'all' else None,
            annee,
            limit_current_year_to_today=True,
        )

        for modalite in modalites:
            # Filtrer ouvrages par modalite ET par les ouvrages géographiquement filtrés
            ouvrages_count = ouvrages_qs.filter(modalites_gestion=modalite).count()
            data.append({
                'modalite': str(modalite),  # Use __str__ method
                'ouvrages': ouvrages_count
            })

            chart_data['labels'].append(str(modalite))
            chart_data['counts'].append(ouvrages_count)

    elif detail_type == 'fournisseurs':
        # Détails des fournisseurs actifs
        data = []
        chart_data = {'labels': [], 'counts': []}

        fournisseurs = Fournisseur.objects.filter(actif=True)[:20]

        for fournisseur in fournisseurs:
            data.append({
                'nom': fournisseur.nom_fournisseur,
                'telephone': fournisseur.telephone or 'N/A',
                'email': fournisseur.email or 'N/A'
            })

        chart_data['labels'] = ['Fournisseurs actifs']
        chart_data['counts'] = [fournisseurs.count()]

    elif detail_type == 'villages':
        # Détails des villages avec nombre d'habitants
        from apps.geography.models import Village
        data = []
        chart_data = {'labels': [], 'habitants': []}

        # Récupérer les villages filtrés géographiquement
        villages_qs = Village.objects.all()

        if village_id:
            villages_qs = villages_qs.filter(id=village_id)
        elif fokontany_id:
            villages_qs = villages_qs.filter(fokontany_id=fokontany_id)
        elif commune_id:
            villages_qs = villages_qs.filter(fokontany__commune_id=commune_id)
        elif district_id:
            villages_qs = villages_qs.filter(fokontany__commune__district_id=district_id)
        elif region_id:
            villages_qs = villages_qs.filter(fokontany__commune__district__region_id=region_id)

        if organisation_id and organisation_id != 'all':
            villages_qs = villages_qs.filter(
                parcelles__pointeau__in=points_eau_qs
            ).distinct()

        for village in villages_qs[:10]:
            habitants_count = parcelles_qs.filter(village=village).aggregate(
                total=Sum('nombre_habitants')
            )['total'] or 0

            data.append({
                'village': village.nom_village,
                'fokontany': village.fokontany.nom_fokontany,
                'commune': village.fokontany.commune.nom_commune,
                'habitants': habitants_count
            })

            chart_data['labels'].append(village.nom_village)
            chart_data['habitants'].append(habitants_count)

    elif detail_type == 'villages-desservis':
        # Détails des villages desservis (avec points d'eau actifs)
        from apps.geography.models import Village
        data = []
        chart_data = {'labels': [], 'points_eau': []}

        # Récupérer les villages filtrés géographiquement qui ont au moins un point d'eau actif
        villages_qs = Village.objects.filter(
            parcelles__pointeau__datebranchement__isnull=False
        ).distinct()

        if village_id:
            villages_qs = villages_qs.filter(id=village_id)
        elif fokontany_id:
            villages_qs = villages_qs.filter(fokontany_id=fokontany_id)
        elif commune_id:
            villages_qs = villages_qs.filter(fokontany__commune_id=commune_id)
        elif district_id:
            villages_qs = villages_qs.filter(fokontany__commune__district_id=district_id)
        elif region_id:
            villages_qs = villages_qs.filter(fokontany__commune__district__region_id=region_id)

        if organisation_id and organisation_id != 'all':
            villages_qs = villages_qs.filter(
                parcelles__pointeau__in=points_eau_qs
            ).distinct()

        for village in villages_qs[:10]:
            points_eau_count = points_eau_actifs_qs.filter(parcelle__village=village).count()
            habitants_desservis = parcelles_qs.filter(
                village=village,
                pointeau__in=points_eau_actifs_qs
            ).distinct().aggregate(total=Sum('nombre_habitants'))['total'] or 0

            data.append({
                'village': village.nom_village,
                'fokontany': village.fokontany.nom_fokontany,
                'commune': village.fokontany.commune.nom_commune,
                'points_eau': points_eau_count,
                'habitants_desservis': habitants_desservis
            })

            chart_data['labels'].append(village.nom_village)
            chart_data['points_eau'].append(points_eau_count)

    elif detail_type == 'adresses':
        # Détails des adresses d'ouvrages
        from apps.infrastructure.models import Adresse
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Filtrer les adresses par les ouvrages géographiquement filtrés
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        adresses = Adresse.objects.filter(ouvrage_id__in=ouvrage_ids, is_active=True).select_related('ouvrage')[:20]

        for adresse in adresses:
            data.append({
                'ouvrage': adresse.ouvrage.nom_ouvrage if adresse.ouvrage else 'N/A',
                'nom': adresse.nomadresse or 'Non spécifié',
                'type': adresse.type_adresse or 'N/A'
            })

        # Compter par type d'adresse
        adresse_counts = Adresse.objects.filter(ouvrage_id__in=ouvrage_ids, is_active=True).values('type_adresse').annotate(
            count=Count('id')
        ).order_by('-count')[:10]

        for item in adresse_counts:
            chart_data['labels'].append(item['type_adresse'] or 'Non spécifié')
            chart_data['counts'].append(item['count'])

    elif detail_type == 'cartes-sociales':
        # Détails des cartes sociales actives
        from apps.infrastructure.models import CarteSociale
        data = []
        chart_data = {'labels': [], 'counts': []}

        # Filtrer les cartes sociales par les ouvrages géographiquement filtrés
        ouvrage_ids = ouvrages_qs.values_list('id', flat=True)
        cartes = CarteSociale.objects.filter(ouvrage_id__in=ouvrage_ids, statut='actif').select_related('ouvrage', 'point_eau')[:20]

        for carte in cartes:
            data.append({
                'numero': carte.numero_carte,
                'beneficiaire': carte.nom_beneficiaire,
                'ouvrage': carte.ouvrage.nom_ouvrage if carte.ouvrage else 'N/A',
                'habitants': carte.nombre_habitants or 0
            })

        # Compter par ouvrage
        carte_counts = CarteSociale.objects.filter(ouvrage_id__in=ouvrage_ids, statut='actif').values('ouvrage__nom_ouvrage').annotate(
            count=Count('id')
        ).order_by('-count')[:10]

        for item in carte_counts:
            chart_data['labels'].append(item['ouvrage__nom_ouvrage'] or 'Non spécifié')
            chart_data['counts'].append(item['count'])

    else:
        return JsonResponse({'error': 'Type de détail non supporté'}, status=400)

    return JsonResponse({
        'table_data': data,
        'chart_data': chart_data,
        'total_items': len(data)
    })


def build_dashboard_commune_ouvrages_stats_payload(user, commune, annee=None, organisation_id=None):
    """
    Construit le detail des KPIs par ouvrage pour une commune.
    Source partagee entre le dashboard et le rapport partenaires afin
    d'assurer exactement le meme perimetre de comptage.
    """
    from django.db.models.functions import ExtractYear

    ouvrages_qs = filter_queryset_by_permissions(
        Ouvrage.objects.filter(
            pointeau__isnull=False,
            pointeau__parcelle__village__fokontany__commune=commune
        ).distinct(),
        user
    )

    points_eau_qs = filter_queryset_by_permissions(
        PointEau.objects.filter(
            ouvrage__in=ouvrages_qs,
            parcelle__village__fokontany__commune=commune
        ),
        user
    )

    if organisation_id and organisation_id != 'all':
        ouvrages_qs = filter_ouvrages_by_gestionnaire_period(
            ouvrages_qs,
            organisation_id,
            annee,
            limit_current_year_to_today=True,
        )
        points_eau_qs = points_eau_qs.filter(ouvrage__in=ouvrages_qs)

    ouvrages_qs = ouvrages_qs.exclude(
        Q(type_ouvrage__nom_type_ouvrage__icontains='suspend') |
        Q(type_ouvrage__nom_type_ouvrage__icontains='potentiel')
    ).distinct()
    points_eau_qs = points_eau_qs.filter(ouvrage__in=ouvrages_qs)

    if annee:
        try:
            annee_int = int(annee)
            ouvrages_qs = ouvrages_qs.filter(
                Q(annee_realisation__lte=annee_int) |
                Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee_int)
            )
            points_eau_qs = filtrer_points_annee(points_eau_qs, annee_int)
        except (TypeError, ValueError):
            pass

    ouvrage_ids = list(ouvrages_qs.values_list('id', flat=True))
    payload = {
        'commune': {
            'id': str(commune.id),
            'nom': commune.nom_commune,
            'district': commune.district.nom_district if commune.district else '',
            'region': commune.district.region.nom_region if commune.district and commune.district.region else '',
        },
        'rows': [],
        'total_ouvrages': 0,
    }
    if not ouvrage_ids:
        return payload

    points_coupes_ids, points_abandonnes_ids = _classer_coupes_point_ids(points_eau_qs, annee=annee)
    modalites_qs = filter_modalites_by_gestionnaire_period(
        ModaliteGestion.objects.filter(ouvrage_id__in=ouvrage_ids),
        organisation_id if organisation_id and organisation_id != 'all' else None,
        annee,
        limit_current_year_to_today=True,
    ).order_by('ouvrage_id', '-date_debut', '-id')

    def modalite_label(code):
        if code == 'forfaitaire':
            return 'Forfaitaire'
        if code == 'au_volume':
            return 'Au volume'
        if code == 'mixte':
            return 'Mixte'
        if code == 'cotisation_annuelle':
            return 'Cotisation annuelle'
        return 'Non défini'

    modalite_by_ouvrage = {}
    for modalite in modalites_qs:
        oid = str(modalite.ouvrage_id)
        if oid not in modalite_by_ouvrage:
            modalite_by_ouvrage[oid] = modalite_label(modalite.modalitepaiement)

    ouvrages_meta = {
        str(ouvrage.id): {
            'ouvrage_id': str(ouvrage.id),
            'ouvrage': ouvrage.nom_ouvrage,
            'modalite_paiement': modalite_by_ouvrage.get(str(ouvrage.id), 'Non défini'),
            'villages_set': set(),
            'villages_desservis_set': set(),
            'parcelles_set': set(),
            'parcelles_desservies_set': set(),
            'points_eau_total': 0,
            'points_eau_prives': 0,
            'points_eau_publics': 0,
            'points_eau_coupes': 0,
            'points_eau_abandonnes': 0,
        }
        for ouvrage in ouvrages_qs
    }

    for point in points_eau_qs.values(
        'id',
        'ouvrage_id',
        'parcelle_id',
        'parcelle__village_id',
        'datebranchement',
        'type_point_eau__nom_type',
    ):
        oid = str(point.get('ouvrage_id'))
        if oid not in ouvrages_meta:
            continue
        target = ouvrages_meta[oid]
        village_id = point.get('parcelle__village_id')
        parcelle_id = point.get('parcelle_id')
        point_id = point.get('id')
        # Actif = branché, ni coupé ni abandonné (mêmes règles que les tuiles du dashboard)
        is_active = (
            bool(point.get('datebranchement'))
            and point_id not in points_coupes_ids
            and point_id not in points_abandonnes_ids
        )
        is_private = is_private_point_type_name(
            point.get('type_point_eau__nom_type'), point.get('estbranchementinstitutionnel')
        )

        if village_id:
            target['villages_set'].add(str(village_id))
            if is_active:
                target['villages_desservis_set'].add(str(village_id))

        if parcelle_id:
            target['parcelles_set'].add(str(parcelle_id))
            if is_active:
                target['parcelles_desservies_set'].add(str(parcelle_id))

        if is_active:
            target['points_eau_total'] += 1
            if is_private:
                target['points_eau_prives'] += 1
            else:
                target['points_eau_publics'] += 1

        if point_id in points_coupes_ids:
            target['points_eau_coupes'] += 1
        elif point_id in points_abandonnes_ids:
            target['points_eau_abandonnes'] += 1

    parcelles_habitants = {
        str(row['id']): int(row['nombre_habitants'] or 0)
        for row in Parcelle.objects.filter(
            id__in={
                *[pid for item in ouvrages_meta.values() for pid in item['parcelles_set']],
                *[pid for item in ouvrages_meta.values() for pid in item['parcelles_desservies_set']],
            }
        ).values('id', 'nombre_habitants')
    }

    for item in ouvrages_meta.values():
        habitants = sum(parcelles_habitants.get(pid, 0) for pid in item['parcelles_set'])
        habitants_desservis = sum(parcelles_habitants.get(pid, 0) for pid in item['parcelles_desservies_set'])
        payload['rows'].append({
            'ouvrage_id': item['ouvrage_id'],
            'ouvrage': item['ouvrage'],
            'modalite_paiement': item['modalite_paiement'],
            'villages': len(item['villages_set']),
            'villages_desservis': len(item['villages_desservis_set']),
            'habitants': habitants,
            'habitants_desservis': habitants_desservis,
            'points_eau_total': item['points_eau_total'],
            'points_eau_prives': item['points_eau_prives'],
            'points_eau_publics': item['points_eau_publics'],
            'points_eau_coupes': item['points_eau_coupes'],
            'points_eau_abandonnes': item['points_eau_abandonnes'],
        })

    payload['rows'].sort(key=lambda row: (row.get('ouvrage') or '').lower())
    payload['total_ouvrages'] = len(payload['rows'])
    return payload


@employe_required
def dashboard_commune_ouvrages_stats_api(request, commune_id):
    """
    API de détail des KPIs par ouvrage pour une commune donnée.
    Colonnes alignées avec le tableau "Statistiques par commune".
    """
    from django.shortcuts import get_object_or_404

    annee = request.GET.get('annee')
    organisation_id = _resolve_dashboard_organisation_filter(
        request.user,
        request.GET.get('organisation_id')
    )

    commune = get_object_or_404(
        Commune.objects.select_related('district__region'),
        pk=commune_id
    )
    payload = build_dashboard_commune_ouvrages_stats_payload(
        request.user,
        commune,
        annee=annee,
        organisation_id=organisation_id,
    )

    return JsonResponse({
        'success': True,
        'commune': payload['commune'],
        'rows': payload['rows'],
        'total_ouvrages': payload['total_ouvrages'],
    })


DASHBOARD_POINT_TYPE_ORDER = ['BF', 'BI', 'Kiosque', 'PMH', 'BP', 'Sans type']
DASHBOARD_POINT_TYPE_COLORS = {
    'BF': '#3B82F6',
    'BI': '#10B981',
    'Kiosque': '#8B5CF6',
    'PMH': '#F59E0B',
    'BP': '#EF4444',
    'Sans type': '#6B7280',
}
DASHBOARD_OUVRAGE_TYPE_COLORS = {
    'AEPG': '#06B6D4',
    'PMH': '#F59E0B',
    'BF': '#3B82F6',
    'BP': '#EF4444',
    'Kiosque': '#8B5CF6',
}
DASHBOARD_MODALITES_FALLBACK_COLORS = [
    '#06B6D4',
    '#14B8A6',
    '#6366F1',
    '#EC4899',
    '#84CC16',
    '#F97316',
    '#A855F7',
    '#0EA5E9',
]
DASHBOARD_OUVRAGE_MODALITE_ORDER = [
    ('forfaitaire', 'Forfaitaire'),
    ('au_volume', 'Au volume'),
    ('mixte', 'Mixte'),
    ('cotisation_annuelle', 'Cotisation'),
]
DASHBOARD_POINT_MODALITE_ORDER = [
    ('forfaitaire', 'Forfaitaire'),
    ('au_volume', 'Au volume'),
    ('sans_modalite', 'Sans modalité'),
]


def _dashboard_modalites_normalize_label(value, fallback='Type non défini'):
    value = (value or '').strip()
    return value or fallback


def _dashboard_modalites_sort_labels(labels, preferred_order):
    order_map = {label: index for index, label in enumerate(preferred_order)}
    return sorted(
        labels,
        key=lambda label: (order_map.get(label, len(order_map)), label.lower())
    )


def _dashboard_modalites_stable_color(label, fixed_colors=None):
    normalized = _dashboard_modalites_normalize_label(label)
    if fixed_colors and normalized in fixed_colors:
        return fixed_colors[normalized]
    return DASHBOARD_MODALITES_FALLBACK_COLORS[
        zlib.crc32(normalized.encode('utf-8')) % len(DASHBOARD_MODALITES_FALLBACK_COLORS)
    ]


def _dashboard_modalites_pick_unique_color(label, used_colors, fixed_colors=None):
    normalized = _dashboard_modalites_normalize_label(label)
    if fixed_colors and normalized in fixed_colors:
        color = fixed_colors[normalized]
        used_colors.add(color)
        return color

    for color in DASHBOARD_MODALITES_FALLBACK_COLORS:
        if color not in used_colors:
            used_colors.add(color)
            return color

    color = _dashboard_modalites_stable_color(normalized, fixed_colors)
    used_colors.add(color)
    return color


def _dashboard_modalites_normalize_ouvrage_type_label(value):
    label = _dashboard_modalites_normalize_label(value)
    compact_upper = ''.join(label.upper().split())
    if compact_upper in {'FPMH', 'PMH', 'PPMH'}:
        return 'PMH'
    return label


def _build_dashboard_ouvrages_modalites_chart_data(filtered_ouvrages_qs, ouvrage_ids_by_modalite):
    ouvrage_modalite_keys = [key for key, _label in DASHBOARD_OUVRAGE_MODALITE_ORDER]
    type_labels = set()
    count_map = {}
    for modalite_key in ouvrage_modalite_keys:
        rows = list(
            filtered_ouvrages_qs.filter(
                id__in=ouvrage_ids_by_modalite.get(modalite_key, set())
            ).values('type_ouvrage__nom_type_ouvrage').annotate(total=Count('id', distinct=True))
        )
        for row in rows:
            type_label = _dashboard_modalites_normalize_ouvrage_type_label(row['type_ouvrage__nom_type_ouvrage'])
            type_labels.add(type_label)
            count_key = (modalite_key, type_label)
            count_map[count_key] = count_map.get(count_key, 0) + int(row['total'] or 0)

    type_labels = _dashboard_modalites_sort_labels(type_labels, [])

    chart_data = {
        'labels': [label for _key, label in DASHBOARD_OUVRAGE_MODALITE_ORDER],
        'datasets': [],
    }
    used_colors = set()

    for type_label in type_labels:
        data = [
            count_map.get((modalite_key, type_label), 0)
            for modalite_key, _label in DASHBOARD_OUVRAGE_MODALITE_ORDER
        ]
        if not any(data):
            continue

        color = _dashboard_modalites_pick_unique_color(type_label, used_colors, DASHBOARD_OUVRAGE_TYPE_COLORS)
        chart_data['datasets'].append({
            'label': type_label,
            'data': data,
            'backgroundColor': color,
            'borderColor': color,
            'borderWidth': 1,
        })

    keep_indexes = []
    for idx in range(len(chart_data['labels'])):
        total_for_modalite = sum(
            int(dataset['data'][idx] or 0)
            for dataset in chart_data['datasets']
            if idx < len(dataset['data'])
        )
        if total_for_modalite > 0:
            keep_indexes.append(idx)

    chart_data['labels'] = [chart_data['labels'][i] for i in keep_indexes]
    for dataset in chart_data['datasets']:
        dataset['data'] = [dataset['data'][i] for i in keep_indexes]

    if not chart_data['datasets']:
        chart_data['labels'] = []

    return chart_data


def _build_dashboard_points_modalites_chart_data(filtered_points_eau_qs, point_ids_by_modalite):
    point_modalite_keys = [key for key, _label in DASHBOARD_POINT_MODALITE_ORDER]
    points_eau_chart_qs = filtered_points_eau_qs.filter(
        ouvrage__isnull=False
    ).annotate(
        type_point_eau_calcule=Case(
            When(type_point_eau__isnull=True, then=Value('Sans type')),
            When(type_point_eau__nom_type__isnull=True, then=Value('Sans type')),
            When(
                Q(estbranchementinstitutionnel=True, type_point_eau__nom_type__iregex=r'^\s*(BF|BP)\s*$'),
                then=Value('BI'),
            ),
            When(type_point_eau__nom_type__iregex=r'^\s*BF\s*$', then=Value('BF')),
            When(type_point_eau__nom_type__icontains='kiosque', then=Value('Kiosque')),
            When(type_point_eau__nom_type__iregex=r'^\s*PMHI?\s*$', then=Value('PMH')),
            When(type_point_eau__nom_type__iregex=r'^\s*BP\s*$', then=Value('BP')),
            default=F('type_point_eau__nom_type'),
            output_field=CharField(),
        )
    )

    point_type_labels = _dashboard_modalites_sort_labels(
        {
            _dashboard_modalites_normalize_label(label)
            for label in points_eau_chart_qs.values_list('type_point_eau_calcule', flat=True).distinct()
        },
        DASHBOARD_POINT_TYPE_ORDER
    )

    counts_by_modalite = {}
    for modalite_key in point_modalite_keys:
        counts_by_modalite[modalite_key] = {
            _dashboard_modalites_normalize_label(row['type_point_eau_calcule']): int(row['total'] or 0)
            for row in points_eau_chart_qs.filter(
                id__in=point_ids_by_modalite.get(modalite_key, set())
            ).values('type_point_eau_calcule').annotate(total=Count('id'))
        }

    chart_data = {
        'labels': [label for _key, label in DASHBOARD_POINT_MODALITE_ORDER],
        'datasets': []
    }

    for type_label in point_type_labels:
        data = [
            counts_by_modalite.get(modalite_key, {}).get(type_label, 0)
            for modalite_key, _label in DASHBOARD_POINT_MODALITE_ORDER
        ]
        if not any(data):
            continue

        color = _dashboard_modalites_stable_color(type_label, DASHBOARD_POINT_TYPE_COLORS)
        chart_data['datasets'].append({
            'label': type_label,
            'data': data,
            'backgroundColor': color,
            'borderColor': color,
            'borderWidth': 1
        })

    keep_indexes = []
    for idx in range(len(chart_data['labels'])):
        total_for_modalite = sum(
            int(dataset['data'][idx] or 0)
            for dataset in chart_data['datasets']
            if idx < len(dataset['data'])
        )
        if total_for_modalite > 0:
            keep_indexes.append(idx)

    chart_data['labels'] = [chart_data['labels'][i] for i in keep_indexes]
    for dataset in chart_data['datasets']:
        dataset['data'] = [dataset['data'][i] for i in keep_indexes]

    if not chart_data['datasets']:
        chart_data['labels'] = []

    return chart_data


def get_dashboard_modalites_payment_chart_payload(
    user,
    region_id='',
    district_id='',
    commune_id='',
    fokontany_id='',
    village_id='',
    annee='',
    organisation_id='',
):
    """Construit les donnees de repartition des modalites pour le dashboard principal."""
    from apps.geography.models import Fokontany, Village

    ouvrages_qs = filter_queryset_by_permissions(
        Ouvrage.objects.filter(type_ouvrage__existe=True),
        user
    )
    points_eau_qs = filter_queryset_by_permissions(
        PointEau.objects.filter(ouvrage__type_ouvrage__existe=True),
        user
    )

    if organisation_id and organisation_id != 'all':
        ouvrages_qs = filter_ouvrages_by_gestionnaire_period(
            ouvrages_qs,
            organisation_id,
            annee,
            limit_current_year_to_today=True,
        )
        points_eau_qs = points_eau_qs.filter(ouvrage__in=ouvrages_qs)

    if annee:
        try:
            annee_int = int(annee)
            ouvrages_qs = ouvrages_qs.filter(
                Q(annee_realisation__lte=annee_int) |
                Q(annee_realisation__isnull=True, annee_realisation_avant__lte=annee_int)
            )
            points_eau_qs = filtrer_points_annee(points_eau_qs, annee_int)
        except (TypeError, ValueError):
            pass

    if village_id:
        village = Village.objects.get(pk=village_id)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village=village).distinct()
        points_eau_qs = points_eau_qs.filter(parcelle__village=village)
    elif fokontany_id:
        fokontany = Fokontany.objects.get(pk=fokontany_id)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany=fokontany).distinct()
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany=fokontany)
    elif commune_id:
        commune = Commune.objects.get(pk=commune_id)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune=commune).distinct()
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune=commune)
    elif district_id:
        district = District.objects.get(pk=district_id)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district=district).distinct()
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district=district)
    elif region_id:
        region = Region.objects.get(pk=region_id)
        ouvrages_qs = ouvrages_qs.filter(pointeau__parcelle__village__fokontany__commune__district__region=region).distinct()
        points_eau_qs = points_eau_qs.filter(parcelle__village__fokontany__commune__district__region=region)

    ouvrages_qs = ouvrages_qs.exclude(
        Q(type_ouvrage__nom_type_ouvrage__icontains='suspend') |
        Q(type_ouvrage__nom_type_ouvrage__icontains='potentiel')
    )
    points_eau_qs = points_eau_qs.filter(ouvrage__in=ouvrages_qs)

    # Actifs = branchés, ni coupés ni abandonnés (mêmes règles que les tuiles)
    points_eau_actifs_qs = points_eau_qs.filter(datebranchement__isnull=False)
    coupes_chart_ids, abandonnes_chart_ids = _classer_coupes_point_ids(points_eau_qs, annee=annee)
    if coupes_chart_ids or abandonnes_chart_ids:
        points_eau_actifs_qs = points_eau_actifs_qs.exclude(id__in=coupes_chart_ids | abandonnes_chart_ids)
    point_payment_scope = build_payment_modality_scope(points_eau_actifs_qs, annee=annee)
    ouvrage_payment_scope = build_payment_modality_scope(points_eau_qs, annee=annee)
    chart_data = _build_dashboard_points_modalites_chart_data(
        points_eau_actifs_qs,
        point_payment_scope['point_ids'],
    )

    return {
        'chart_data': chart_data,
        'ouvrages_chart_data': _build_dashboard_ouvrages_modalites_chart_data(
            ouvrages_qs,
            ouvrage_payment_scope['ouvrage_ids'],
        ),
        'total_modalites': len(chart_data['labels']),
        'total_types': len(chart_data['datasets']),
    }


@employe_required
def modalites_payment_chart_api(request):
    """API pour recuperer les donnees des graphiques de modalites du dashboard."""
    payload = get_dashboard_modalites_payment_chart_payload(
        request.user,
        region_id=request.GET.get('region_id', ''),
        district_id=request.GET.get('district_id', ''),
        commune_id=request.GET.get('commune_id', ''),
        fokontany_id=request.GET.get('fokontany_id', ''),
        village_id=request.GET.get('village_id', ''),
        annee=request.GET.get('annee', ''),
        organisation_id=_resolve_dashboard_organisation_filter(
            request.user,
            request.GET.get('organisation_id')
        ),
    )
    return JsonResponse(payload)
