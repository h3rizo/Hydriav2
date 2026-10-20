"""
Vue pour l'export Excel des écritures partie double avec scans
"""
from django.http import HttpResponse, JsonResponse, StreamingHttpResponse
from django.views.decorators.http import require_http_methods
from .models_partie_double import EcriturePartieDouble
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import zipfile
from io import BytesIO
from datetime import datetime
import json
import time


def generate_progress_stream(request):
    """
    Générateur pour le streaming de progression via Server-Sent Events
    Export Excel SANS les scans (trop lourd)
    """
    # Récupérer les mêmes filtres que l'API de liste
    annee = request.GET.get('annee')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    compte = request.GET.get('compte')
    type_conversion = request.GET.get('type_conversion')
    search = request.GET.get('search')

    try:
        # Yield initial message
        yield f"data: {json.dumps({'step': 'init', 'message': 'Initialisation...', 'progress': 0})}\n\n"
        time.sleep(0.1)

        # Requête de base
        ecritures_qs = EcriturePartieDouble.objects.select_related(
            'entree_diamex',
            'entree_diamex__imported_by'
        )

        # Appliquer les filtres
        if annee:
            try:
                ecritures_qs = ecritures_qs.filter(date__year=int(annee))
            except (ValueError, TypeError):
                pass
        if date_debut:
            ecritures_qs = ecritures_qs.filter(date__gte=date_debut)
        if date_fin:
            ecritures_qs = ecritures_qs.filter(date__lte=date_fin)
        if compte:
            ecritures_qs = ecritures_qs.filter(compte=compte)
        if type_conversion:
            ecritures_qs = ecritures_qs.filter(type_conversion=type_conversion)
        if search:
            from django.db.models import Q
            ecritures_qs = ecritures_qs.filter(
                Q(libelle__icontains=search) |
                Q(compte__icontains=search) |
                Q(fournisseur__icontains=search) |
                Q(detail__icontains=search)
            )

        # PAS de limite - exporter tout
        total_ecritures = ecritures_qs.count()

        yield f"data: {json.dumps({'step': 'count', 'message': f'{total_ecritures} écritures trouvées', 'progress': 5})}\n\n"
        time.sleep(0.1)

        # Créer le workbook Excel
        yield f"data: {json.dumps({'step': 'excel', 'message': 'Création du fichier Excel...', 'progress': 10})}\n\n"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Écritures Partie Double"

        # Styles
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # En-têtes de colonnes
        headers = [
            "N° Paire", "N° Ligne", "Date", "Rec", "Libellé", "Compte",
            "Débit", "Crédit", "Fournisseur", "Détail", "Info", "Chemin Scan",
            "Type Conversion", "Équilibrée", "Total Débit Paire", "Total Crédit Paire", "Écart"
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border_thin

        # Grouper par id_lien
        yield f"data: {json.dumps({'step': 'grouping', 'message': 'Regroupement par paires...', 'progress': 15})}\n\n"

        paires = {}
        scans_info = {}  # {mois: count} pour statistiques

        for ecriture in ecritures_qs:
            id_lien = ecriture.id_lien or f'sans-lien-{ecriture.id}'
            if id_lien not in paires:
                paires[id_lien] = []
            paires[id_lien].append(ecriture)

        total_paires = len(paires)
        yield f"data: {json.dumps({'step': 'pairs', 'message': f'{total_paires} paires identifiées', 'progress': 20})}\n\n"

        # Remplir le fichier Excel
        current_row = 2
        pair_index = 0

        for id_lien, lignes in paires.items():
            pair_index += 1

            # Calculer progression
            progress = 20 + int((pair_index / total_paires) * 60)  # 20% à 80%
            if pair_index % 100 == 0 or pair_index == total_paires:
                yield f"data: {json.dumps({'step': 'processing', 'message': f'Traitement paire {pair_index}/{total_paires}', 'progress': progress, 'pairs_done': pair_index, 'pairs_total': total_paires})}\n\n"

            # Calculer totaux de la paire
            total_debit = sum(float(l.debit) for l in lignes)
            total_credit = sum(float(l.credit) for l in lignes)
            ecart = abs(total_debit - total_credit)
            est_equilibree = ecart < 0.01

            # Écrire chaque ligne de la paire
            for ligne_index, ecriture in enumerate(lignes, 1):
                # Récupérer le scan (chemin seulement, pas de copie)
                scan_path = ""
                if ecriture.scan:
                    scan_source = None
                    if ecriture.entree_diamex.imported_by:
                        employe = ecriture.entree_diamex.imported_by
                        employe_base = employe.bases_diamex.filter(
                            base_diamex=ecriture.entree_diamex.fichier_source
                        ).first()

                        if employe_base:
                            chemin_base = employe_base.chemin_base_diamex
                            scan_relatif = ecriture.scan
                            scan_clean = scan_relatif.lstrip('.')

                            if not scan_relatif.startswith(chemin_base):
                                scan_source = chemin_base + scan_clean
                            else:
                                scan_source = scan_relatif

                    if scan_source and os.path.exists(scan_source):
                        scan_path = scan_source
                        # Compter les scans par mois pour statistiques
                        if ecriture.date:
                            mois = ecriture.date.strftime('%Y-%m')
                            scans_info[mois] = scans_info.get(mois, 0) + 1

                # Écrire la ligne
                ws.cell(row=current_row, column=1, value=pair_index)
                ws.cell(row=current_row, column=2, value=ligne_index)
                date_cell = ws.cell(row=current_row, column=3, value=ecriture.date if ecriture.date else None)
                if ecriture.date:
                    date_cell.number_format = 'DD/MM/YYYY'
                ws.cell(row=current_row, column=4, value=ecriture.rec)
                ws.cell(row=current_row, column=5, value=ecriture.libelle)
                ws.cell(row=current_row, column=6, value=ecriture.compte)
                ws.cell(row=current_row, column=7, value=float(ecriture.debit))
                ws.cell(row=current_row, column=8, value=float(ecriture.credit))
                ws.cell(row=current_row, column=9, value=ecriture.fournisseur)
                ws.cell(row=current_row, column=10, value=ecriture.detail)
                ws.cell(row=current_row, column=11, value=ecriture.info)
                ws.cell(row=current_row, column=12, value=scan_path)  # Chemin complet

                ws.cell(row=current_row, column=13, value=ecriture.type_conversion)

                # Totaux de la paire (uniquement sur la première ligne)
                if ligne_index == 1:
                    ws.cell(row=current_row, column=14, value="Oui" if est_equilibree else "Non")
                    ws.cell(row=current_row, column=15, value=total_debit)
                    ws.cell(row=current_row, column=16, value=total_credit)
                    ws.cell(row=current_row, column=17, value=ecart)

                # Appliquer les bordures
                for col in range(1, 18):
                    ws.cell(row=current_row, column=col).border = border_thin

                current_row += 1

        # Ajuster la largeur des colonnes
        yield f"data: {json.dumps({'step': 'formatting', 'message': 'Formatage du fichier Excel...', 'progress': 85})}\n\n"

        for col in range(1, 18):
            ws.column_dimensions[get_column_letter(col)].width = 15
        ws.column_dimensions[get_column_letter(12)].width = 50  # Colonne scan plus large

        # Sauvegarder le fichier Excel
        yield f"data: {json.dumps({'step': 'saving_excel', 'message': 'Sauvegarde du fichier Excel...', 'progress': 90})}\n\n"

        temp_filename = f'ecritures_partie_double_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        temp_path = os.path.join('media', 'temp_exports', temp_filename)

        # Créer le dossier si nécessaire
        os.makedirs(os.path.dirname(temp_path), exist_ok=True)

        wb.save(temp_path)

        yield f"data: {json.dumps({'step': 'complete', 'message': 'Export terminé !', 'progress': 100, 'download_url': f'/media/temp_exports/{temp_filename}', 'filename': temp_filename, 'scans_by_month': scans_info})}\n\n"

    except Exception as e:
        yield f"data: {json.dumps({'step': 'error', 'message': f'Erreur: {str(e)}', 'progress': 0})}\n\n"


@require_http_methods(["GET"])
def api_export_partie_double_progress(request):
    """
    API pour streamer la progression de l'export via Server-Sent Events
    """
    response = StreamingHttpResponse(
        generate_progress_stream(request),
        content_type='text/event-stream'
    )
    response['Cache-Control'] = 'no-cache'
    response['X-Accel-Buffering'] = 'no'
    return response


@require_http_methods(["GET"])
def api_export_scans_mois(request):
    """
    API pour exporter les scans d'un mois spécifique en ZIP
    """
    try:
        mois = request.GET.get('mois')  # Format: YYYY-MM
        if not mois:
            return JsonResponse({'success': False, 'error': 'Mois requis'}, status=400)

        # Récupérer les mêmes filtres que l'export principal
        date_debut = request.GET.get('date_debut')
        date_fin = request.GET.get('date_fin')
        compte = request.GET.get('compte')
        type_conversion = request.GET.get('type_conversion')
        search = request.GET.get('search')

        # Parser le mois
        annee, mois_num = map(int, mois.split('-'))

        # Requête de base
        ecritures_qs = EcriturePartieDouble.objects.select_related(
            'entree_diamex',
            'entree_diamex__imported_by'
        ).filter(
            date__year=annee,
            date__month=mois_num
        ).exclude(scan__isnull=True).exclude(scan='')

        # Appliquer les autres filtres
        if date_debut:
            ecritures_qs = ecritures_qs.filter(date__gte=date_debut)
        if date_fin:
            ecritures_qs = ecritures_qs.filter(date__lte=date_fin)
        if compte:
            ecritures_qs = ecritures_qs.filter(compte=compte)
        if type_conversion:
            ecritures_qs = ecritures_qs.filter(type_conversion=type_conversion)
        if search:
            from django.db.models import Q
            ecritures_qs = ecritures_qs.filter(
                Q(libelle__icontains=search) |
                Q(compte__icontains=search) |
                Q(fournisseur__icontains=search) |
                Q(detail__icontains=search)
            )

        # Collecter les scans uniques
        scans_a_copier = {}  # {chemin_source: nom_fichier_destination}

        for ecriture in ecritures_qs:
            if ecriture.scan:
                scan_source = None
                if ecriture.entree_diamex.imported_by:
                    employe = ecriture.entree_diamex.imported_by
                    employe_base = employe.bases_diamex.filter(
                        base_diamex=ecriture.entree_diamex.fichier_source
                    ).first()

                    if employe_base:
                        chemin_base = employe_base.chemin_base_diamex
                        scan_relatif = ecriture.scan
                        scan_clean = scan_relatif.lstrip('.')

                        if not scan_relatif.startswith(chemin_base):
                            scan_source = chemin_base + scan_clean
                        else:
                            scan_source = scan_relatif

                if scan_source and os.path.exists(scan_source):
                    nom_fichier = os.path.basename(scan_source)
                    scans_a_copier[scan_source] = nom_fichier

        if not scans_a_copier:
            return JsonResponse({'success': False, 'error': 'Aucun scan trouvé pour ce mois'}, status=404)

        # Créer le ZIP
        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for scan_source, nom_fichier in scans_a_copier.items():
                try:
                    with open(scan_source, 'rb') as f:
                        zip_file.writestr(nom_fichier, f.read())
                except Exception as e:
                    print(f"Erreur copie {scan_source}: {e}")
                    continue

        zip_buffer.seek(0)

        # Retourner le ZIP
        response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="scans_{mois}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip"'
        return response

    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@require_http_methods(["GET"])
def api_export_partie_double_excel(request):
    """
    API pour exporter les écritures partie double en Excel avec scans
    Version synchrone (sans streaming) - DEPRECATED, utiliser api_export_partie_double_progress
    """
    try:
        # Récupérer les mêmes filtres que l'API de liste
        annee = request.GET.get('annee')
        date_debut = request.GET.get('date_debut')
        date_fin = request.GET.get('date_fin')
        compte = request.GET.get('compte')
        type_conversion = request.GET.get('type_conversion')
        search = request.GET.get('search')

        # Requête de base
        ecritures_qs = EcriturePartieDouble.objects.select_related(
            'entree_diamex',
            'entree_diamex__imported_by'
        )

        # Appliquer les filtres
        if annee:
            try:
                ecritures_qs = ecritures_qs.filter(date__year=int(annee))
            except (ValueError, TypeError):
                pass
        if date_debut:
            ecritures_qs = ecritures_qs.filter(date__gte=date_debut)
        if date_fin:
            ecritures_qs = ecritures_qs.filter(date__lte=date_fin)
        if compte:
            ecritures_qs = ecritures_qs.filter(compte=compte)
        if type_conversion:
            ecritures_qs = ecritures_qs.filter(type_conversion=type_conversion)
        if search:
            from django.db.models import Q
            ecritures_qs = ecritures_qs.filter(
                Q(libelle__icontains=search) |
                Q(compte__icontains=search) |
                Q(fournisseur__icontains=search) |
                Q(detail__icontains=search)
            )

        # Limiter à 1000 résultats
        ecritures_qs = ecritures_qs[:1000]

        # Créer le workbook Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Écritures Partie Double"

        # Styles
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        pair_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        pair_header_font = Font(bold=True)
        border_thin = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # En-têtes de colonnes
        headers = [
            "N° Paire", "N° Ligne", "Date", "Rec", "Libellé", "Compte",
            "Débit", "Crédit", "Fournisseur", "Détail", "Info", "Scan",
            "Type Conversion", "Équilibrée", "Total Débit Paire", "Total Crédit Paire", "Écart"
        ]

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border_thin

        # Grouper par id_lien
        paires = {}
        scans_a_copier = {}  # {chemin_source: nom_fichier_destination}

        for ecriture in ecritures_qs:
            id_lien = ecriture.id_lien or f'sans-lien-{ecriture.id}'
            if id_lien not in paires:
                paires[id_lien] = []
            paires[id_lien].append(ecriture)

        # Remplir le fichier Excel
        current_row = 2
        pair_index = 0

        for id_lien, lignes in paires.items():
            pair_index += 1

            # Calculer totaux de la paire
            total_debit = sum(float(l.debit) for l in lignes)
            total_credit = sum(float(l.credit) for l in lignes)
            ecart = abs(total_debit - total_credit)
            est_equilibree = ecart < 0.01

            # Écrire chaque ligne de la paire
            for ligne_index, ecriture in enumerate(lignes, 1):
                # Récupérer le scan et préparer la copie
                scan_link = ""
                if ecriture.scan:
                    # Reconstruire le chemin absolu du scan
                    scan_source = None
                    if ecriture.entree_diamex.imported_by:
                        employe = ecriture.entree_diamex.imported_by
                        employe_base = employe.bases_diamex.filter(
                            base_diamex=ecriture.entree_diamex.fichier_source
                        ).first()

                        if employe_base:
                            chemin_base = employe_base.chemin_base_diamex
                            scan_relatif = ecriture.scan
                            scan_clean = scan_relatif.lstrip('.')

                            if not scan_relatif.startswith(chemin_base):
                                scan_source = chemin_base + scan_clean
                            else:
                                scan_source = scan_relatif

                    # Extraire le nom du fichier
                    if scan_source and os.path.exists(scan_source):
                        nom_fichier = os.path.basename(scan_source)
                        scans_a_copier[scan_source] = nom_fichier
                        scan_link = f"./scans/{nom_fichier}"

                # Écrire la ligne
                ws.cell(row=current_row, column=1, value=pair_index)
                ws.cell(row=current_row, column=2, value=ligne_index)
                date_cell = ws.cell(row=current_row, column=3, value=ecriture.date if ecriture.date else None)
                if ecriture.date:
                    date_cell.number_format = 'DD/MM/YYYY'
                ws.cell(row=current_row, column=4, value=ecriture.rec)
                ws.cell(row=current_row, column=5, value=ecriture.libelle)
                ws.cell(row=current_row, column=6, value=ecriture.compte)
                ws.cell(row=current_row, column=7, value=float(ecriture.debit))
                ws.cell(row=current_row, column=8, value=float(ecriture.credit))
                ws.cell(row=current_row, column=9, value=ecriture.fournisseur)
                ws.cell(row=current_row, column=10, value=ecriture.detail)
                ws.cell(row=current_row, column=11, value=ecriture.info)

                # Lien hypertexte vers le scan
                if scan_link:
                    cell_scan = ws.cell(row=current_row, column=12)
                    cell_scan.value = os.path.basename(scan_link)
                    cell_scan.hyperlink = scan_link
                    cell_scan.font = Font(color="0563C1", underline="single")

                ws.cell(row=current_row, column=13, value=ecriture.type_conversion)

                # Totaux de la paire (uniquement sur la première ligne)
                if ligne_index == 1:
                    ws.cell(row=current_row, column=14, value="Oui" if est_equilibree else "Non")
                    ws.cell(row=current_row, column=15, value=total_debit)
                    ws.cell(row=current_row, column=16, value=total_credit)
                    ws.cell(row=current_row, column=17, value=ecart)

                # Appliquer les bordures
                for col in range(1, 18):
                    ws.cell(row=current_row, column=col).border = border_thin

                current_row += 1

        # Ajuster la largeur des colonnes
        for col in range(1, 18):
            ws.column_dimensions[get_column_letter(col)].width = 15

        # Sauvegarder le fichier Excel en mémoire
        excel_buffer = BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        # Créer un fichier ZIP contenant Excel + scans
        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            # Ajouter le fichier Excel
            zip_file.writestr('ecritures_partie_double.xlsx', excel_buffer.getvalue())

            # Ajouter les scans
            for scan_source, nom_fichier in scans_a_copier.items():
                if os.path.exists(scan_source):
                    with open(scan_source, 'rb') as f:
                        zip_file.writestr(f'scans/{nom_fichier}', f.read())

        zip_buffer.seek(0)

        # Retourner le fichier ZIP
        response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="ecritures_partie_double_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip"'
        return response

    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)
